"""Keep the real conversion screen responsive during slow file validation."""

import json
import os
import tempfile
import time
from pathlib import Path
from threading import Event

from pptx import Presentation
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, Qt, QTimer
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import input_documents
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import control, pump, until


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    out = Path('reports/rc84/input-selection')
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-selection-', ignore_cleanup_errors=True) as folder:
        folder = Path(folder)
        os.environ['BILICLASS_DATA'] = str(folder / 'library')
        library = Library(folder / 'library')
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings, ticks, refreshes = [], [], []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT / 'qml/Main.qml')
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        root.setProperty('page', 'new')
        beat = QTimer()
        beat.setInterval(10)
        beat.timeout.connect(lambda: ticks.append(time.monotonic()))
        original = input_documents.document_selection
        gate, entered = Event(), Event()

        def select(name):
            started = time.monotonic()
            QMetaObject.invokeMethod(root, 'acceptInputDocuments', Qt.DirectConnection,
                                     Q_ARG('QVariant', [(folder / name).as_uri()]))
            return time.monotonic() - started

        def slow_validate(urls):
            entered.set()
            assert gate.wait(5)
            return original(urls)

        try:
            deck = Presentation()
            for _ in range(51):
                deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(folder / 'Bài gốc.pptx')
            deck.save(folder / 'Bài mới.pptx')
            until(lambda: bridge._catalog.worker is None)
            pump(250)
            bridge.changed.connect(lambda: refreshes.append(1))
            beat.start()
            input_documents.document_selection = slow_validate
            duration = select('Bài gốc.pptx')
            assert duration < .1, duration
            until(entered.is_set)
            assert bridge.inputDocuments.busy
            assert not control(root, 'createLessonButton').isEnabled()
            pump(600)
            # A teacher can still edit lesson configuration while the disk waits.
            title = control(root, 'lessonTitle')
            title.setProperty('text', 'Tên giáo viên tự đặt')
            assert title.isEnabled()
            gate.set()
            until(lambda: not bridge.inputDocuments.busy)
            assert root.property('selectedFileName') == 'Bài gốc.pptx'
            assert title.property('text') == 'Tên giáo viên tự đặt'
            input_documents.document_selection = original
            title.setProperty('text', 'Bài gốc')
            select('Bài mới.pptx')
            until(lambda: not bridge.inputDocuments.busy)
            assert title.property('text') == 'Bài mới'
            select('Missing.pptx')
            until(lambda: not bridge.inputDocuments.busy)
            assert root.property('inputDocumentError') and not root.property('selectedFile')
            assert not control(root, 'createLessonButton').isEnabled()
            gate.clear()
            entered.clear()
            input_documents.document_selection = slow_validate
            select('Bài gốc.pptx')
            until(entered.is_set)
            QMetaObject.invokeMethod(root, 'removeInputDocument', Qt.DirectConnection)
            gate.set()
            until(lambda: bridge.inputDocuments.worker is None)
            assert not root.property('selectedFile') and not root.property('inputDocumentError')
            assert not refreshes, 'File selection must not refresh every library/settings binding'
            max_gap = max(b - a for a, b in zip(ticks, ticks[1:]))
            assert max_gap < .2, max_gap
            assert not warnings, warnings
            root.grabWindow().save(str(out / 'conversion.png'))
            result = dict(status='passed', selection_slot_seconds=duration,
                          slow_disk_simulation_seconds=.6, max_gui_tick_seconds=max_gap,
                          preserves_custom_title=True, replacement_title=True,
                          clear_discards_late_result=True, missing_file_rejected=True,
                          broad_refreshes=len(refreshes), warnings=warnings)
            (out / 'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result), flush=True)
        finally:
            gate.set()
            input_documents.document_selection = original
            beat.stop()
            bridge.inputDocuments.shutdown()
            for task in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                         bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                         bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
                task.shutdown()
            root.close()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
