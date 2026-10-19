"""Exercise teacher subject defaults and persisted class selection in real Qt."""

import json
import logging
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, Qt
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import click, control, pump, until

from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    out = Path('reports/rc90/defaults')
    out.mkdir(parents=True, exist_ok=True)
    warnings, checks = [], []
    with tempfile.TemporaryDirectory(prefix='bcl-conversion-defaults-', ignore_cleanup_errors=True) as temp:
        for phase in ('first', 'reopened'):
            library = Library(Path(temp) / 'data')
            bridge = Bridge(library)
            engine = QQmlApplicationEngine()
            engine.warnings.connect(lambda values: warnings.extend(str(v) for v in values))
            engine.rootContext().setContextProperty('bridge', bridge)
            engine.load(RESOURCE_ROOT / 'qml/Main.qml')
            assert engine.rootObjects(), warnings
            window = engine.rootObjects()[0]
            window.resize(1366, 850)
            def c(name):
                return control(window, name)

            def choose(name, index):
                item = c(name)
                item.setProperty('currentIndex', index)
                assert QMetaObject.invokeMethod(item, 'activated', Qt.DirectConnection, Q_ARG(int, index))
                pump(80)

            try:
                pump(200)
                if phase == 'reopened':
                    assert bridge.defaultTeachingSubject == 'STEM – dự án, thực hành'
                    assert c('lessonSubject').property('currentText') == 'Vật lý'
                    assert bridge.conversionSubjectSelection == 'Vật lý'
                    assert bridge.mascotAccessorySubject == 'Toán'
                    assert c('creationEducation').property('currentText') == 'THCS'
                    assert c('creationGrade').property('currentText') == '8'
                    assert c('settingsSubjects').property('currentText') == 'Khác'
                    assert c('settingsCustomSubject').property('text') == bridge.defaultTeachingSubject
                    checks.append('Subject and THCS/8 restored after closing and reopening library and QML')
                    continue

                assert c('lessonSubject').property('currentIndex') == -1
                assert c('creationEducation').property('currentText') == 'THPT'
                assert c('creationGrade').property('currentText') == '10'
                assert library.setting('conversion_class_selection', None) is None
                window.setProperty('page', 'settings')
                choose('settingsSubjects', bridge.conversionSubjects.index('Tin học'))
                assert not bridge.settings['subjects'], 'An unsaved choice must not change the profile'
                click(window, c('saveProfile'))
                assert bridge.defaultTeachingSubject == 'Tin học'
                assert c('lessonSubject').property('currentText') == 'Tin học'
                assert bridge.mascotAccessorySubject == 'Tin học'
                assert c('mascotAccessorySubject').property('currentText') == 'Tin học'
                window.grabWindow().save(str(out / 'general-subject.png'))
                checks.append('Settings dropdown shares conversion subjects; default changes only after Save')

                window.setProperty('page', 'new')
                choose('lessonSubject', bridge.conversionSubjects.index('Vật lý'))
                choose('creationGrade', 1)
                assert bridge.conversionClassSelection == {'education': 'THPT', 'grade': '11'}
                # Exercise the actual QML submit button and Bridge dispatch,
                # writing only disposable requests; no service/API is called.
                from app.chatgpt_handoff import inspect_returned_deck
                from app.conversion_input import creation_config
                from app.google_conversion import prepare

                source = Path(temp) / 'Bài thử môn Vật lý.pptx'
                deck = Presentation()
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = 'VI: Cơ năng\nEN: Mechanical energy'
                deck.save(source)
                QMetaObject.invokeMethod(window, 'acceptInputDocuments', Qt.DirectConnection,
                                         Q_ARG('QVariant', [source.as_uri()]))
                until(lambda: not bridge.inputDocuments.busy)
                assert c('lessonSubject').property('currentText') == 'Vật lý'
                sent = []
                original_prepare = bridge._prepare_google

                def capture(title, subject, education, grade, text, file_url, mode, preset, style):
                    config = creation_config(title, subject, education, grade, mode, preset, style)
                    job = prepare(library.directory, config, source)
                    saved = library.create_external_lesson(config, library.store_source(source), inspect_returned_deck(source))
                    assert job['config']['subject'] == saved['subject'] == 'Vật lý'
                    assert job['config']['grade'] == saved['grade'] == '11'
                    sent.append(config)

                assert bridge.googleTranslation
                bridge._prepare_google = capture
                try:
                    for style in ('source', 'template'):
                        click(window, c('creationKeepOriginal' if style == 'source' else 'creationUseTemplate'))
                        for mode in ('parallel_columns', 'english_only'):
                            click(window, c('formatChoice_' + mode))
                            assert c('lessonSubject').property('currentText') == 'Vật lý'
                            prompt = window.property('creationPrompt')
                            assert (prompt.toVariant() if hasattr(prompt, 'toVariant') else prompt)['ready']
                            click(window, c('createLessonButton'))
                            assert sent[-1]['style'] == style and sent[-1]['conversion_format'] == mode
                    assert len(sent) == 4
                finally:
                    bridge._prepare_google = original_prepare
                checks.append('Chosen subject/class survive file/design/mode choices and reach submit, saved request and lesson unchanged')
                window.setProperty('page', 'settings')
                full_name = 'Hoạt động trải nghiệm, hướng nghiệp'
                choose('settingsSubjects', bridge.conversionSubjects.index(full_name))
                click(window, c('saveProfile'))
                assert bridge.settings['subjects'] == [full_name], 'Do not split a curriculum subject at its comma'
                window.setProperty('page', 'new')
                assert c('lessonSubject').property('currentText') == 'Vật lý'
                assert c('creationGrade').property('currentText') == '11'
                assert bridge.mascotAccessorySubject == full_name
                c('settingsPage').setProperty('activeTab', 3)
                window.setProperty('page', 'settings')
                choices = c('settingsPage').property('accessorySubjects')
                choices = choices.toVariant() if hasattr(choices, 'toVariant') else choices
                choose('mascotAccessorySubject', choices.index('Toán'))
                click(window, c('saveMascot'))
                until(lambda: not bridge.mascotSaving)
                assert bridge.mascotAccessorySubject == 'Toán'
                c('settingsPage').setProperty('activeTab', 0)
                bridge.navigate.emit('voice')
                pump(100)
                assert c('lessonSubject').property('currentText') == 'Vật lý'
                assert c('creationGrade').property('currentText') == '11'
                checks.append('Latest conversion subject survives new drafts; accessory starts from General and saves its own override')

                window.setProperty('page', 'new')
                for index, education, grade in ((2, 'Tiểu học', '4'), (0, 'THPT', '12'), (1, 'THCS', '8')):
                    choose('creationEducation', index)
                    grades = bridge.educationGrades(education)
                    assert c('creationGrade').property('currentText') == grades[0]
                    choose('creationGrade', grades.index(grade))
                    assert bridge.conversionClassSelection == {'education': education, 'grade': grade}
                bridge.rememberConversionClass('THCS', '12')
                assert bridge.conversionClassSelection == {'education': 'THCS', 'grade': '8'}
                checks.append('All three education levels reset to a valid grade and persist the latest pair')

                window.setProperty('page', 'settings')
                choose('settingsSubjects', bridge.conversionSubjects.index('Khác'))
                c('settingsCustomSubject').setProperty('text', '')
                assert not c('saveProfile').isEnabled()
                c('settingsCustomSubject').setProperty('text', 'STEM – dự án, thực hành')
                click(window, c('saveProfile'))
                assert bridge.settings['subjects'] == ['STEM – dự án, thực hành']
                assert c('lessonSubject').property('currentText') == 'Vật lý'
                assert bridge.mascotAccessorySubject == 'Toán'
                checks.append('Custom subjects, including commas, remain intact; empty custom subject cannot be saved')
            finally:
                window.hide()
                for worker in (bridge.inputDocuments, bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog, bridge.browserAI,
                               bridge.english_reader, bridge.narration_preparation, bridge.classroom, bridge.voice_controls,
                               bridge._companion_preferences):
                    worker.shutdown()
                engine.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                app.processEvents()
                library.close()
                logging.shutdown()
        assert not warnings, warnings
        result = {'status': 'passed', 'checks': checks, 'qml_warnings': warnings}
        (out / 'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
        print(json.dumps(result), flush=True)


if __name__ == '__main__':
    main()
