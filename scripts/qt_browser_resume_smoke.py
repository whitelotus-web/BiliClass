"""Check actual Qt startup timers against durable job/cancel checkpoints."""

import json
import logging
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from PySide6.QtCore import QTimer  # noqa: E402
from PySide6.QtGui import QGuiApplication  # noqa: E402

from app.browser_accounts import BrowserAccounts  # noqa: E402
from app.browser_automation import write_record  # noqa: E402
from app.chatgpt_handoff import prepare_request  # noqa: E402
from app.library import Library  # noqa: E402
from app.ui import Bridge  # noqa: E402


def main():
    application = QGuiApplication(sys.argv)
    calls, libraries, bridges = [], [], []

    class StartupBridge(Bridge):
        def _run_web_conversion(self):
            calls.append(self.library.directory.name)

    with tempfile.TemporaryDirectory(prefix="biliclass-qt-resume-") as temporary:
        root = Path(temporary)
        source = root / "source.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(source)
        for mode in ("resume", "stopped", "stop-before-timer", "saved", "unconfirmed"):
            library = Library(root / mode)
            libraries.append(library)
            store = BrowserAccounts(library.directory)
            account = store.add("Fixture")
            store.observe(account["id"], "free")
            request = prepare_request(library.directory, {
                "title": "Resume fixture", "subject": "Toán", "education_level": "THPT", "grade": "11",
                "level": 2, "layout": "split_view", "preset": "standard", "style": "source", "mode": "level",
                "provider": "browser_web"}, source)
            record = {"account_id": account["id"], "state": "prepared", "url": "", "followups": 0,
                      "error": "browser", "auto_resume": mode != "unconfirmed"}
            if mode == "saved":
                record["lesson_id"] = "fixture-already-saved"
            write_record(request["folder"], record)
            library.set_setting("chatgpt_request_folder", request["folder"])
            if mode == "stopped":
                (Path(request["folder"]) / "conversion-stopped").touch()
            bridge = StartupBridge(library)
            bridges.append(bridge)
            if mode == "stop-before-timer":
                bridge._browser_running = True
                bridge._busy = True
                bridge.cancelJob()
                assert bridge.cancel_event.is_set()
                bridge._busy = bridge._browser_running = False
        QTimer.singleShot(1200, application.quit)
        application.exec()
        for bridge in bridges:
            bridge.browserAI.shutdown()
        for library in libraries:
            library.close()
        logging.shutdown()
        assert calls == ["resume"], calls
        report = {"status": "passed", "calls": calls,
                  "scope": "Actual Qt timer; stopped, cancel-before-timer, saved and legacy jobs stay idle; no real browser/ChatGPT submission"}
        target = ROOT / "reports/app/browser-startup-resume.json"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps(report))


if __name__ == "__main__":
    main()
