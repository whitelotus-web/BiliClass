"""Real Office cell fitting starts automatically; the teacher only opens the deck."""

import hashlib
import json
import shutil
import tempfile
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import speech
from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.library import Library
from app.office_tasks import render_quality
from app.paths import RESOURCE_ROOT
from app.slide_repair import failure_keys

from .qt_result_preview_mascot_smoke import control, pump, until


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    report = Path('reports/rc144/automatic-completion')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='biliclass-auto-fit-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        path = folder / 'source.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table_shape = slide.shapes.add_table(2, 2, Inches(.5), Inches(1), Inches(9), Inches(1.6))
        for row, text in [(0, 'A triangle has three vertices. Its three angles sum to 180 degrees.'),
                          (1, 'Tam giác có ba đỉnh. Tổng ba góc trong tam giác bằng 180 độ.')]:
            cell = table_shape.table.cell(row, 0)
            cell.text = text
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
        table_shape.table.cell(1, 1).text = 'Unrelated cell stays unchanged.'
        table_shape.table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        math = etree.fromstring(b'<a14:m xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>x+1</m:t></m:r></m:oMath></a14:m>')
        table_shape.table.cell(0, 1).text_frame.paragraphs[0]._p.append(math)
        table_shape.table.cell(0, 1).fill.solid()
        table_shape.table.cell(0, 1).fill.fore_color.rgb = RGBColor(255, 255, 255)
        table_shape.table.cell(0, 1).text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        label = slide.shapes.add_textbox(Inches(.5), Inches(3.4), Inches(6), Inches(.6))
        label.text = 'A clear label remains readable.'
        label.fill.solid()
        label.fill.fore_color.rgb = RGBColor(255, 255, 255)
        label.text_frame.paragraphs[0].font.size = Pt(20)
        label.text_frame.paragraphs[0].font.color.rgb = RGBColor(224, 224, 224)
        deck.save(path)
        shutil.copyfile(path, report / 'source-fixture.pptx')
        original = path.read_bytes()
        inspection = inspect_returned_deck(path)
        render_quality(path, folder / 'data', inspection['slide_quality'])
        (report / 'initial-quality.json').write_text(json.dumps(inspection['slide_quality'], indent=2))
        before = failure_keys(inspection['slide_quality'])
        assert len(before) >= 2, before
        library = Library(folder / 'data')
        lesson = library.create_external_lesson({'title': 'Automatic cell fitting', 'subject': 'Toán', 'style': 'source',
            'education_level': 'THPT', 'grade': '11', 'preset': 'standard', 'conversion_format': 'parallel_columns'},
            library.store_source(path), inspection)
        from app.ui import Bridge

        with patch.object(speech, 'list_voices', return_value=[dict(speech.PAIRS[0])]):
            bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT / 'qml/Main.qml')
        root = engine.rootObjects()[0]
        try:
            root.resize(1366, 850)
            bridge._browser_state.update(phase='completed', lessonId=lesson['id'])
            # No click on Repair: the stale, failed layout starts completion itself.
            bridge._show_lesson(lesson, external_preview(lesson, library.directory),
                                destination='chatgpt', load_preview=False)
            until(lambda: bridge.busy, timeout=10)
            until(lambda: not bridge.busy, timeout=120)
            (report / 'applied-quality.json').write_text(json.dumps(bridge.lesson.get('slide_quality'), indent=2))
            assert bridge.lesson['revision'] > lesson['revision'], (bridge._message, bridge._error, bridge.lesson.get('slide_quality'))
            final = library.get(lesson['id'])
            assert not failure_keys(final['slide_quality']), final['slide_quality']
            cells = [row for row in final['slide_quality']['local_repairs'] if row['operation'] == 'fit_table_cell']
            assert len(cells) == 2 and {tuple(row['cell']) for row in cells} == {(0, 0), (1, 0)}
            assert final['slide_quality']['completion_version'] >= 7
            assert bridge.lessonDiagnostics['attention'] == 0
            assert any(row['operation'] == 'improve_text_contrast' for row in final['slide_quality']['local_repairs'])
            output = Presentation(library.directory / 'sources' / final['source']['file'])
            before_table, after_table = table_shape.table, output.slides[0].shapes[0].table
            assert [cell.text for row in after_table.rows for cell in row.cells] == [cell.text for row in before_table.rows for cell in row.cells]
            assert etree.tostring(after_table.cell(0, 1)._tc) == etree.tostring(before_table.cell(0, 1)._tc)
            assert etree.tostring(after_table.cell(1, 1)._tc) == etree.tostring(before_table.cell(1, 1)._tc)
            assert path.read_bytes() == original
            pump(400)
            assert not bridge.busy, 'Completion must not start a repeated repair loop'
            assert not warnings, warnings
            assert root.grabWindow().save(str(report / 'completed.png'))
            result = dict(status='passed', initial_cell_errors=len(before), repaired_cells=len(cells),
                          final_errors=0, automatic_without_repair_click=True, original_preserved=True,
                          solid_fill_contrast_corrected=True,
                          table_grid_other_cells_and_math_preserved=True, qml_warnings=warnings,
                          original_sha256=hashlib.sha256(original).hexdigest())
            (report / 'result.json').write_text(json.dumps(result, indent=2))
            print(json.dumps(result), flush=True)
        finally:
            control(root, 'companionWindow').setProperty('disposing', True)
            root.hide()
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview', '_catalog',
                         'browserAI', 'english_reader', 'narration_preparation', 'narrationPreferences',
                         '_companion_preferences', 'classroom', 'voice_controls', 'inputDocuments'):
                getattr(bridge, name).shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
