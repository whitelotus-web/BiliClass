"""Real Qt + local PowerPoint; API fixture only, no keys or teacher data sent."""
import json
import os
import tempfile
import time
from pathlib import Path
from threading import Event

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PySide6.QtCore import (
    Q_ARG,
    QCoreApplication,
    QEvent,
    QEventLoop,
    QMetaObject,
    QObject,
    QPointF,
    Qt,
    QTimer,
    QUrl,
)
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from app import google_conversion
from app.branding import configure_application
from app.google_translation import Translator
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def pump(ms=30):
    loop = QEventLoop()
    QTimer.singleShot(ms, loop.quit)
    loop.exec()


def until(check, timeout=120):
    end = time.monotonic() + timeout
    while not check():
        assert time.monotonic() < end, 'Google Qt flow timed out'
        pump()


app = QGuiApplication([])
configure_application(app)
QQuickStyle.setStyle('Basic')
for path in (RESOURCE_ROOT / 'assets').glob('*.ttf'):
    QFontDatabase.addApplicationFont(str(path))
app.setFont(QFont('Be Vietnam Pro', 10))
reports = Path(os.environ.get('BILICLASS_GOOGLE_REPORT', 'reports/google-api'))
reports.mkdir(parents=True, exist_ok=True)
calls = []
release_response = Event()


def transport(texts, key, source, target):
    calls.append(len(texts))
    assert release_response.wait(30), 'Fixture response was never released'
    return [text.replace('Cơ năng', 'Mechanical energy').replace('Vật có vận tốc', 'The object has a velocity of') for text in texts]


google_conversion.Translator = lambda directory, budget: Translator(directory, budget, transport=transport)
with tempfile.TemporaryDirectory(prefix='biliclass-google-ui-', ignore_cleanup_errors=True) as directory:
    directory = Path(directory)
    source = directory / 'source.pptx'
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for number in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255) if number == 1 else RGBColor(241, 247, 255)
        for text, y, h, size in [('Cơ năng', .4, 1, 32), ('Vật có vận tốc 12 m/s.', 2, 4, 24)]:
            box = slide.shapes.add_textbox(Inches(.6), Inches(y), Inches(12), Inches(h))
            box.text = text
            box.text_frame.paragraphs[0].font.size = Pt(size)
            box.text_frame.paragraphs[0].font.name = 'Arial'
    deck.save(source)
    library = Library(directory / 'library')
    library.set_setting('chatgpt_request_folder', 'old-browser-job-must-not-run')
    bridge = Bridge(library)
    bridge._browser_ai.signIn = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError('Browser must stay disabled'))
    assert bridge.googleTranslation and not bridge._browser_ai._health_timer.isActive()
    assert not bridge.chatgptRequest
    assert [row['id'] for row in bridge.conversionFormats] == ['parallel_columns','english_only']
    warnings = []
    engine = QQmlApplicationEngine()
    engine.warnings.connect(lambda rows: warnings.extend(str(row) for row in rows))
    engine.rootContext().setContextProperty('bridge', bridge)
    engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
    assert engine.rootObjects(), warnings
    window = engine.rootObjects()[0]
    window.resize(1366, 850)

    def control(name):
        queue = [window.contentItem()]
        while queue:
            item = queue.pop()
            if item.objectName() == name:
                return item
            queue.extend(item.childItems())
        raise AssertionError(name)

    def click(name):
        item = control(name)
        parent = item.parentItem()
        while parent is not None:
            if parent.metaObject().indexOfProperty('contentY') >= 0:
                top = item.mapToItem(parent, QPointF(0, 0)).y()
                bottom = item.mapToItem(parent, QPointF(0, item.height())).y()
                if bottom > parent.height() or top < 0:
                    parent.setProperty('contentY', max(0, min(parent.property('contentHeight') - parent.height(),
                        parent.property('contentY') + top - 100)))
                    pump(100)
            parent = parent.parentItem()
        point = item.mapToScene(QPointF(item.width()/2, item.height()/2)).toPoint()
        assert item.isVisible() and item.isEnabled(), name
        assert 0 <= point.y() < window.height(), (name, point)
        QTest.mouseClick(window, Qt.LeftButton, Qt.NoModifier, point)
        pump(70)

    window.setProperty('page', 'settings')
    pump(200)
    click('settingsTab6')
    assert not control('browserServiceChatGPT').isVisible()
    assert not control('browserServiceGemini').isVisible()
    control('googleApiKeys').setProperty('text', 'Project A | ' + 'AIza' + 'a'*35 + '\nProject B | ' + 'AIza' + 'b'*35)
    click('saveGoogleApiKeys')
    until(lambda: not bridge.busy)
    assert not control('googleApiKeys').property('text')
    assert len(bridge.googleSettings.state['keys']) == 2
    assert not hasattr(bridge.googleSettings, 'syncGitHub')
    for name in ('receiveGoogleKeysGitHub', 'sendGoogleKeysGitHub', 'googleCharacterBudget',
                 'preparePictureOcr', 'settingsTranslationUsage', 'toggleGoogleAdvanced'):
        assert window.findChild(QObject, name) is None, name
    assert window.grabWindow().save(str(reports / 'settings.png'))
    beats = []
    timer = QTimer()
    timer.setInterval(50)
    timer.timeout.connect(lambda: beats.append(time.monotonic()))
    timer.start()
    window.setProperty('page', 'new')
    QMetaObject.invokeMethod(window, 'acceptInputDocuments', Qt.DirectConnection,
                             Q_ARG('QVariant', [source.as_uri()]))
    until(lambda: not bridge.inputDocuments.busy)
    control('lessonTitle').setProperty('text', 'Cơ năng')
    control('lessonSubject').setProperty('currentIndex', bridge.conversionSubjects.index('Vật lý'))
    control('creationGrade').setProperty('currentIndex', bridge.educationGrades('THPT').index('11'))
    selected_file = window.property('selectedFile')
    click('createLessonButton')
    until(lambda: calls)
    first_calls = calls.copy()
    assert bridge.browserState['running']
    assert window.property('page') == 'new', 'Conversion must stay in its configuration workspace'
    assert window.property('selectedFile') == selected_file
    assert control('lessonTitle').property('text') == 'Cơ năng', 'Starting work must preserve the entered fields'
    assert control('creationConfiguration').isVisible() and not control('conversionConfigurationFields').isEnabled()
    assert not control('chooseInputDocumentButton').isEnabled()
    assert not control('createLessonButton').isEnabled()
    for width, height in ((1080, 700), (1366, 850), (1920, 1080)):
        window.resize(width, height)
        pump(150)
        pump(100)
        bar = control('browserConversionProgressBar')
        y = bar.mapToScene(QPointF(0, 0)).y()
        footer_y = control('creationFooter').mapToScene(QPointF(0, 0)).y()
        assert 0 <= y < footer_y - bar.height(), (width, height, y, footer_y)
        cancel = control('browserConversionCancel')
        point = cancel.mapToScene(QPointF(cancel.width()/2, cancel.height()/2))
        assert footer_y <= point.y() < window.height() and 0 <= point.x() < window.width()
        assert window.grabWindow().save(str(reports / f'inline-{width}.png'))
    window.resize(1366, 850)
    bridge.setApplicationTheme('dark')
    pump(150)
    assert window.grabWindow().save(str(reports / 'inline-dark.png'))
    window.setProperty('page', 'library')
    pump(80)
    assert control('applicationStatusMessage').property('returnable')
    control('applicationStatusMessage').clicked.emit()
    pump(120)
    assert window.property('page') == 'new' and not control('applicationStatusMessage').property('returnable')
    assert window.property('selectedFile') == selected_file and control('lessonTitle').property('text') == 'Cơ năng'
    pump(1100)
    assert bridge.conversionClock['seconds'] >= 1
    assert control('browserConversionPercent').property('text').endswith('%')
    assert not control('browserConversionSettings').isVisible()
    assert window.grabWindow().save(str(reports / 'loading.png'))
    folder = bridge.chatgptRequest['folder']
    conversion_id = bridge.browserState['conversionId']
    assert control('browserConversionCancel').property('text') == 'Dừng'
    assert control('browserConversionCancel').property('symbol') == '■'
    click('browserConversionCancel')
    assert bridge.browserState['stopping'] and bridge.cancel_event.is_set()
    assert not control('browserConversionCancel').isEnabled()
    assert bridge.conversionClock['status'] == 'Đang dừng…'
    assert window.grabWindow().save(str(reports / 'stopping.png'))
    release_response.set()
    until(lambda: not bridge._busy)
    assert bridge.browserState['phase'] == 'cancelled' and bridge.chatgptRequest['folder'] == folder
    assert window.property('page') == 'new' and not library.list_lessons()
    assert library.setting('google_request_folder', '') == folder
    assert not (Path(folder) / 'conversion-discarded').exists()
    assert google_conversion.load(library.directory, folder)['folder'] == folder
    assert control('conversionConfigurationFields').isEnabled() and control('browserConversionResume').isEnabled()
    assert window.grabWindow().save(str(reports / 'paused.png'))
    click('browserConversionResume')
    until(lambda: bridge._lesson and not bridge._busy)
    lesson = bridge._lesson
    assert lesson['external_deck']['provider'] == 'google_translate'
    assert lesson['narration']['mode'] == 'none'
    assert bridge.browserState['phase'] == 'completed'
    until(lambda: not bridge.previewBusy and bool(bridge.quickResult.get('image')))
    until(lambda: control('conversionCompletionImage').property('paintedWidth') > 0)
    assert window.property('page') == 'chatgpt', 'Completion must wait for the teacher to choose the next step'
    assert control('conversionCompletion').isVisible()
    action_names = ['openConvertedVoiceButton', 'downloadCompletedConversion', 'teachCompletedConversion', 'libraryAfterConversion']
    for name in action_names:
        assert control(name).isVisible() and control(name).isEnabled(), name
    assert not control('browserConversionCancel').isVisible()
    assert control('conversionCompletionCounter').property('text') == 'Slide 1 / 2'
    assert not control('conversionCompletionPrevious').isEnabled()
    assert control('conversionCompletionNext').isEnabled()
    first_image = bridge.quickResult['image']
    click('conversionCompletionNext')
    until(lambda: not bridge.previewBusy and bridge.quickResult['slide'] == 2 and bridge.quickResult['image'] != first_image)
    assert control('conversionCompletionCounter').property('text') == 'Slide 2 / 2'
    assert not control('conversionCompletionNext').isEnabled()
    assert window.property('page') == 'chatgpt' and calls == first_calls, (window.property('page'), calls, first_calls)
    click('conversionCompletionPrevious')
    until(lambda: not bridge.previewBusy and bridge.quickResult['slide'] == 1 and bridge.quickResult['image'] == first_image)
    assert not any(row['approved'] for row in library.get(lesson['id'])['segments'])
    for width, height in ((1080, 700), (1366, 850), (1920, 1080)):
        window.resize(width, height)
        pump(150)
        actions = [control(name) for name in action_names]
        assert max(a.width() for a in actions) - min(a.width() for a in actions) <= 1
        for item in actions:
            point = item.mapToScene(QPointF())
            assert point.x() >= 0 and point.x() + item.width() <= width
            assert point.y() >= 0 and point.y() + item.height() < height
        navigation = control('conversionCompletionNavigation')
        counter = control('conversionCompletionCounter')
        assert abs(navigation.mapToScene(QPointF(navigation.width()/2, 0)).x() - counter.mapToScene(QPointF(counter.width()/2, 0)).x()) <= 1
        assert window.grabWindow().save(str(reports / f'completed-{width}.png'))
    window.resize(1366, 850)
    pump(100)
    assert window.grabWindow().save(str(reports / 'completed.png'))
    # Both navigation choices stay explicit; merely finishing must not start
    # narration, teaching or a second conversion. Inline preview is automatic.
    click('libraryAfterConversion')
    assert window.property('page') == 'library'
    window.setProperty('page', 'chatgpt')
    pump(100)
    assert not bridge.browserState['running'] and bridge.conversionClock['status'] == 'Đã hoàn tất'
    assert control('browserConversionClockState').property('text') == 'Đã hoàn tất'
    completed_id = bridge.lesson['id']
    other = library.create('Another lesson', 'Tin học', 'THPT', '10', [('Slide 1', 'Máy tính.')])
    bridge._show_lesson(other, destination='editor', auto_repair=False, load_preview=False)
    window.setProperty('page', 'chatgpt')
    until(lambda: not bridge.busy and bridge.lesson['id'] == completed_id)
    assert bridge.quickResult['lesson_id'] == completed_id, 'Completion actions must never target a different browsed lesson'
    until(lambda: not bridge.previewBusy and bool(bridge.quickResult.get('image')))
    assert control('conversionCompletionPreview').property('current')
    bridge.browserConversionProgress.emit({'stage': 'quality', 'percent': 72, 'message': 'late', 'runId': 'older-run'})
    pump(100)
    assert bridge.browserState['phase'] == 'completed' and bridge.browserState['percent'] == 100
    window.setProperty('page', 'chatgpt')
    pump(100)
    click('openConvertedVoiceButton')
    assert window.property('page') == 'voice' and bridge.lesson['id'] == completed_id
    guide = control('voiceWorkflowGuide')
    heading = control('voicePageHeading')
    assert guide.isVisible() and 'giọng đọc và ngôn ngữ' in guide.property('text')
    assert guide.mapToScene(QPointF()).y() >= heading.mapToScene(QPointF()).y() + heading.height()
    assert window.grabWindow().save(str(reports / 'voice-instructions.png'))
    # A worker may remain attached after Office stops presenting; preview still works.
    bridge.powerpoint_worker = object()
    bridge._quick_result.update(image='')
    bridge.quickResultChanged.emit()
    click('toggleVoiceSlidePreview')
    until(lambda: not bridge.previewBusy and bool(bridge.quickResult.get('image')))
    until(lambda: control('convertedSlideImage').property('paintedWidth') > 0)
    assert control('lessonSlidePreview').isVisible()
    assert window.grabWindow().save(str(reports / 'voice-slide-preview.png'))
    bridge.powerpoint_worker = None
    save = window.findChild(QObject, 'presentationNoVoiceSaveDialog')
    QMetaObject.invokeMethod(save, 'openSuggested')
    pump(250)
    suggested = Path(save.property('selectedFile').toLocalFile()).name
    assert source.stem in suggested and '(Song ngữ VI-EN)' in suggested, suggested
    save.setProperty('visible', False)
    window.setProperty('page', 'new')
    pump(100)
    assert control('resumeCreationTask').property('text') == 'Mở bài vừa chuyển đổi'
    click('resumeCreationTask')
    assert window.property('page') == 'chatgpt' and calls == first_calls, 'Reopening must never translate again or start the voice step'
    until(lambda: not bridge._busy)
    output = directory / 'download.pptx'
    bridge.exportPowerPoint(QUrl.fromLocalFile(str(output)).toString())
    until(lambda: not bridge._busy)
    assert output.is_file() and len(Presentation(output).slides) == 2
    assert calls == first_calls
    assert all('key' not in row for row in bridge.googleSettings.state['keys'])
    for size in [(1366, 850), (1080, 700)]:
        window.resize(*size)
        window.setProperty('page', 'settings')
        pump(150)
        assert window.grabWindow().save(str(reports / f'settings-{size[0]}.png'))
    preparing, release_preparation = Event(), Event()
    original_prepare = google_conversion.prepare

    def held_prepare(*args, **kwargs):
        result = original_prepare(*args, **kwargs)
        preparing.set()
        assert release_preparation.wait(15)
        return result

    google_conversion.prepare = held_prepare
    bridge.convertDocument('Chuẩn bị rồi dừng', 'Vật lý', 'THPT', '11', '', QUrl.fromLocalFile(str(source)).toString(),
                           'parallel_columns', 'standard', 'source')
    until(preparing.is_set)
    bridge.discardConversion(conversion_id)
    assert not bridge.cancel_event.is_set(), 'A stale discard request must not stop a different conversion'
    click('browserConversionCancel')
    release_preparation.set()
    until(lambda: not bridge._busy)
    assert bridge.browserState['phase'] == 'cancelled' and bridge.chatgptRequest['folder']
    assert library.setting('google_request_folder', '') == bridge.chatgptRequest['folder']
    assert google_conversion.load(library.directory, bridge.chatgptRequest['folder'])['config']['title'] == 'Chuẩn bị rồi dừng'
    assert {row['id'] for row in library.list_lessons(verify_audio=False)} == {completed_id, other['id']}
    assert library.get(completed_id)['id'] == completed_id
    google_conversion.prepare = original_prepare
    # Clearing the input after a failed/paused run must also forget its durable
    # resume state. Merely hiding its selected filename leaves a stale progress UI.
    paused_folder=bridge.chatgptRequest['folder']
    usage=bridge.googleSettings.state['keys']
    window.setProperty('selectedFile',QUrl.fromLocalFile(str(source)).toString())
    window.setProperty('selectedFileName',source.name)
    window.setProperty('previousInputTitle',source.stem)
    control('lessonTitle').setProperty('text',source.stem)
    window.setProperty('page','new')
    pump(100)
    assert control('creationConversionProgress').isVisible()
    click('removeInputDocumentButton')
    pump(100)
    assert not window.property('selectedFile') and not control('lessonTitle').property('text')
    assert not bridge.chatgptRequest and not library.setting('google_request_folder','')
    assert bridge.conversionClock['seconds']==0 and bridge.browserState['percent']==0
    assert not control('creationConversionProgress').isVisible() and not control('browserConversionResume').isVisible()
    assert (Path(paused_folder)/'conversion-discarded').exists()
    assert bridge.googleSettings.state['keys']==usage
    assert {row['id'] for row in library.list_lessons(verify_audio=False)}=={completed_id,other['id']}
    bridge.browserConversionProgress.emit({'percent':54,'message':'late','runId':'dismissed-run'})
    pump(60)
    assert not control('creationConversionProgress').isVisible()
    # A fresh choice is independent; it never resumes the dismissed run.
    assert QMetaObject.invokeMethod(window,'applyInputSelection',Qt.DirectConnection,
        Q_ARG('QVariant',{'valid':True,'url':QUrl.fromLocalFile(str(source)).toString(),'name':source.name}),
        Q_ARG('QVariant',''))
    pump(80)
    assert window.property('selectedFile') and not control('creationConversionProgress').isVisible()
    assert control('createLessonButton').isEnabled()
    assert window.grabWindow().save(str(reports/'cleared-and-reselected.png'))
    timer.stop()
    assert len(beats) > 10, 'UI timer did not keep ticking'
    assert not warnings, warnings
    result = {'status': 'passed', 'api': 'fixture; not a live Google benchmark', 'api_calls': len(calls),
              'checks': ['multi-key save and clear', 'browser controls hidden; no startup browser job',
                         'native source → text API → local editable PPTX', 'clock and percentage during request',
                         'inline conversion retains configuration and input; progress/cancel visible at 1080/1366/1920; task return stays inline',
                         'Stop uses square icon, disables while stopping, preserves work, resumes without another API call',
                         'stale discard guard; cached response and usage retained',
                         'stop during preparation before UI knows folder; prepared request and earlier completed lessons retained',
                         'dismiss paused input clears persisted job/progress/clock/error; selecting again is fresh; saved lessons/usage remain',
                         'completion immediately shows real slides; previous/next without leaving completion',
                         'four balanced actions at 1080/1366/1920; restore exact completed lesson',
                         'voice preview works with inactive attached Office worker; download keeps original filename',
                         'separate voice step; download without voice', 'no QML warnings'],
              'qml_warnings': warnings, 'max_ui_gap_seconds': round(max(b-a for a,b in zip(beats, beats[1:])), 3)}
    (reports / 'qt-flow.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))
    window.close()
    bridge.slidePreview.shutdown()
    bridge.googleSettings.shutdown()
    bridge._library_preferences.shutdown()
    bridge._catalog.shutdown()
    bridge.classroom.shutdown()
    bridge.libraryThumbnails.shutdown()
    bridge.voice_controls.shutdown()
    engine.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    app.processEvents()
    bridge.browserAI.shutdown()
    bridge.english_reader.shutdown()
    bridge.narration_preparation.shutdown()
    library.close()
