"""Exercise the real teaching button, native Office slideshow, voice and mascot.

Always run through this guarded entry: Windows helpers must not reopen the UI.
The presentation and library are disposable; no teacher account is used.
"""

import hashlib
import json
import shutil
import sys
import tempfile
import time
from multiprocessing import freeze_support
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main():
    import pythoncom
    import win32api
    import win32com.client
    import win32gui
    from pptx import Presentation
    from pptx.util import Inches
    from PySide6.QtCore import QEventLoop, QMetaObject, QObject, QPoint, QPointF, Qt, QTimer, QUrl
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtQml import QQmlApplicationEngine
    from PySide6.QtQuickControls2 import QQuickStyle
    from PySide6.QtTest import QTest

    from app import speech, ui
    from app.chatgpt_handoff import inspect_returned_deck
    from app.library import Library
    from app.paths import RESOURCE_ROOT
    from app.ui import Bridge

    app = QGuiApplication([])
    QQuickStyle.setStyle("Basic")
    reports = RESOURCE_ROOT.parent / "reports/app"
    reports.mkdir(parents=True, exist_ok=True)
    trace, warnings = [], []

    def pump(ms=80):
        loop = QEventLoop()
        QTimer.singleShot(ms, loop.quit)
        loop.exec()

    def wait(predicate, timeout=55, check_errors=True):
        until = time.monotonic() + timeout
        while not predicate():
            pump()
            if check_errors:
                assert not bridge.error, bridge.message
            assert time.monotonic() < until, "Teaching timeout: " + bridge.message

    with tempfile.TemporaryDirectory(prefix="biliclass-teaching-", ignore_cleanup_errors=True) as folder:
        directory = Path(folder)
        path = directory / "Teaching smoke.pptx"
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
        pairs = [
            ("Chào cả lớp.", "Welcome to our bilingual lesson."),
            ("Hãy giải thích câu trả lời.", "Explain your answer to the class."),
            ("", ""),  # Image-only slides must not read stale narration.
        ]
        for index, (vi, en) in enumerate(pairs, 1):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            if vi:
                slide.shapes.add_textbox(Inches(.6), Inches(.6), Inches(10), Inches(2)).text = vi + "\n" + en
                slide.notes_slide.notes_text_frame.text = f"VI: {vi}\nEN: {en}"
        deck.save(path)
        original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        library = Library(directory / "library")
        config = {"title": "Teaching smoke", "subject": "Liên môn", "education_level": "THPT",
                  "grade": "10", "level": 2, "layout": "split_view", "preset": "standard",
                  "style": "source", "mode": "level"}
        lesson = library.create_external_lesson(config, library.store_source(path), inspect_returned_deck(path))
        bridge = Bridge(library)
        plays = []
        real_play = speech.play

        def recorded_play(audio):
            plays.append(str(audio))
            real_play(audio)

        speech.play = recorded_play
        engine = QQmlApplicationEngine()
        engine.warnings.connect(lambda entries: warnings.extend(map(str, entries)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        companion = window.findChild(QObject, "companionWindow")

        def click(name):
            item = window.findChild(QObject, name)
            assert item and item.property("enabled"), name
            assert QMetaObject.invokeMethod(item, "clicked", Qt.DirectConnection)

        try:
            voice = bridge.voiceSettings["en"]
            assert voice, "English voice unavailable"
            samples = [speech.synthesize(en, voice, bridge.voiceSettings["rate"], reports / "teaching-voice-cache")["path"]
                     for _, en in pairs[:2]]
            (library.directory / "audio").mkdir(exist_ok=True)
            audio = [shutil.copy2(sample, library.directory / "audio" / Path(sample).name) for sample in samples]
            print("English audio prepared", flush=True)
            bridge.openLesson(lesson["id"])
            wait(lambda: not bridge.busy)
            assert not bridge.error, bridge.message
            click("useConvertedLessonButton")
            assert window.findChild(QObject, "wholeLessonReview").property("visible")
            click("confirmWholeLessonButton")
            assert bridge.powerpoint_worker is not None, (bridge.message, bridge.busy, bridge.error)
            print("Teaching session requested", flush=True)
            bridge.powerpoint_worker.failed.connect(lambda message: print("Office error: " + message, flush=True))
            bridge.powerpoint_worker.stateChanged.connect(lambda state: print("Office state: " + json.dumps(state), flush=True))
            wait(lambda: bridge.powerpointState.get("active") and companion.isVisible())
            assert not bridge.error, bridge.message
            assert companion.property("collapsed") and companion.width() <= 130
            assert companion.grabWindow().pixelColor(0, 0).alpha() == 0
            assert not window.findChild(QObject, "companionMenu").isVisible()
            hwnd = bridge.powerpointState["hwnd"]

            def native_navigation(command):
                # Change the actual Office view independently of the BiliClass
                # worker. Avoid global key injection while the user is working.
                pythoncom.CoInitialize()
                office = owned = view = None
                try:
                    office = win32com.client.GetActiveObject("PowerPoint.Application")
                    for number in range(1, office.SlideShowWindows.Count + 1):
                        candidate = office.SlideShowWindows(number)
                        if Path(candidate.Presentation.FullName).resolve() == Path(bridge.quickResult["path"]).resolve():
                            owned = candidate
                            break
                    assert owned is not None, "Fixture slideshow not found"
                    view = owned.View
                    getattr(view, command)()
                finally:
                    candidate = view = owned = office = None
                    pythoncom.CoUninitialize()

            rectangle = win32gui.GetWindowRect(hwnd)
            monitor = win32api.GetMonitorInfo(win32api.MonitorFromWindow(hwnd))["Monitor"]
            assert all(abs(a - b) <= 2 for a, b in zip(rectangle, monitor)), (rectangle, monitor)
            trace.append("Teaching button -> whole-lesson confirmation -> native full-screen PowerPoint + transparent mascot")

            center = QPoint(companion.width() // 2, companion.height() // 2)
            before = (companion.x(), companion.y())
            QTest.mousePress(companion, Qt.LeftButton, pos=center)
            QTest.mouseMove(companion, center + QPoint(-60, -40))
            QTest.mouseRelease(companion, Qt.LeftButton, pos=center + QPoint(-60, -40))
            pump(200)
            assert companion.x() < before[0] - 20 and companion.y() < before[1] - 20
            QTest.mouseClick(companion, Qt.LeftButton, pos=center)
            pump(200)
            assert window.findChild(QObject, "companionMenu").isVisible()
            read = window.findChild(QObject, "companionReadEnglish")
            position = read.mapToScene(QPointF(read.width() / 2, read.height() / 2)).toPoint()
            QTest.mouseClick(companion, Qt.LeftButton, pos=position)
            pump()
            assert bridge.mascotState == "speaking" and plays[-1] == str(audio[0])
            click("companionStopSpeech")
            assert bridge.mascotState == "idle"
            assert companion.grabWindow().save(str(reports / "teaching-mascot.png"))

            # Simulate slow synthesis finishing after the teacher changes slide.
            pending = []
            normal_launch, normal_cache = bridge.launch, ui.available_audio
            try:
                bridge.launch = lambda action, callback: pending.append(callback)
                ui.available_audio = lambda *args: None
                click("companionReadEnglish")
                assert len(pending) == 1
            finally:
                bridge.launch, ui.available_audio = normal_launch, normal_cache
            # Change the real Office view, not just the BiliClass navigation slot.
            native_navigation("Next")
            wait(lambda: bridge.powerpointState.get("slide") == 2)
            assert bridge.segmentIndex == 1 and bridge._reading_text("en") == pairs[1][1]
            previous_plays = len(plays)
            pending[0]({"path": audio[0]})
            assert len(plays) == previous_plays and bridge.mascotState == "idle"
            click("companionReadEnglish")
            assert bridge.mascotState == "speaking" and plays[-1] == str(audio[1])
            click("companionNext")
            wait(lambda: bridge.powerpointState.get("slide") == 3)
            assert bridge.mascotState == "idle", bridge.mascotState
            assert not bridge._reading_text("en"), bridge._reading_text("en")
            assert not read.property("enabled")
            click("companionPrevious")
            wait(lambda: bridge.powerpointState.get("slide") == 2)
            assert bridge.segmentIndex == 1 and read.property("enabled")
            trace.append("Mascot drag/click -> cached English playback; navigation follows slide; delayed old voice discarded; blank slide disables voice")
            native_navigation("Exit")
            wait(lambda: bridge.powerpoint_worker is None)
            assert bridge.mascotState == "idle"
            # The same button must work a second time after ending the show.
            click("useConvertedLessonButton")
            bridge.powerpoint_worker.failed.connect(lambda message: print("Reopen error: " + message, flush=True))
            bridge.powerpoint_worker.stateChanged.connect(lambda state: print("Reopen state: " + json.dumps(state), flush=True))
            wait(lambda: bridge.powerpointState.get("active"))
            bridge.stopPowerPoint()
            wait(lambda: bridge.powerpoint_worker is None)
            assert hashlib.sha256(path.read_bytes()).hexdigest() == original_hash
            assert hashlib.sha256(Path(bridge.quickResult["path"]).read_bytes()).hexdigest() == original_hash
            assert not warnings, warnings
            trace.append("Native Office exit -> voice stops -> reopen/close succeeds; both source copies unchanged; no QML warnings")
        finally:
            speech.play = real_play
            bridge.stopSpeech()
            if bridge.powerpoint_worker:
                bridge.stopPowerPoint()
                wait(lambda: bridge.powerpoint_worker is None, check_errors=False)
            if bridge.worker:
                bridge.worker.wait()
            bridge.browserAI.shutdown()
            companion.hide()
            window.close()
            pump()
            library.close()
    report = {"status": "passed", "trace": trace, "warnings": warnings,
              "scope": "Real Qt, installed PowerPoint and local English voice; disposable 3-slide deck, one screen"}
    (reports / "teaching-smoke.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False), flush=True)
    app.quit()


if __name__ == "__main__":
    freeze_support()
    main()
