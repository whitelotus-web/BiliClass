"""Disposable Qt import/narration/export and relocation controls; no real API or voice model."""

import json
import os
import sys
import tempfile
import wave
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QPointF, QUrl
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import click, control, pump, until

from app import speech
from app.google_translation import KeyStore
from app.library import Library
from app.library_location import cleanup_previous, finalize_transfer, verify_transfer_start
from app.paths import RESOURCE_ROOT, user_data
from app.ui import Bridge


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    out = Path('reports/rc65/storage-voice')
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='biliclass-storage-voice-', ignore_cleanup_errors=True) as folder:
        root = Path(folder)
        os.environ['LOCALAPPDATA'] = str(root / 'control')
        os.environ.pop('BILICLASS_DATA', None)
        library = Library(root / 'original')
        KeyStore(library.directory).add('AIza' + 'a' * 35)
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(str(value) for value in values))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]

        def c(name):
            return control(window, name)

        def fake_synthesis(text, voice, rate, directory, *args, **kwargs):
            path = Path(directory) / (speech.audio_key(text, voice, rate) + '.wav')
            path.parent.mkdir(parents=True, exist_ok=True)
            with wave.open(str(path), 'wb') as audio:
                audio.setnchannels(1)
                audio.setsampwidth(2)
                audio.setframerate(24000)
                audio.writeframes(b'\0\0' * 24000)
            return path

        try:
            window.resize(1366, 850)
            window.setProperty('page', 'voice')
            pump(200)
            assert c('navigationvoice').mapToScene(QPointF()).y() > c('navigationnew').mapToScene(QPointF()).y()
            assert c('navigationvoice').mapToScene(QPointF()).y() < c('navigationlibrary').mapToScene(QPointF()).y()
            assert c('chooseVoicePresentation').isVisible()
            assert window.grabWindow().save(str(out / 'voice-empty.png'))
            source = root / 'Lời đọc có sẵn.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = 'Máy tính xử lý thông tin.'
            deck.save(source)
            before = source.read_bytes()
            bridge.importNarrationPresentation(source.as_uri())
            until(lambda: not bridge.busy and bridge.lesson.get('voice_only'), 30)
            assert window.property('page') == 'voice'
            assert source.read_bytes() == before
            assert c('narrationVietnamese').property('checked')
            assert not c('narrationBoth').isEnabled()
            until(lambda: c('createLessonNarrationButton').isEnabled())
            with patch('app.speech.synthesize_background', side_effect=fake_synthesis), patch('app.speech.narration_workers', return_value=1):
                click(window, c('createLessonNarrationButton'))
                until(lambda: not bridge.narrationBusy and bridge.lessonNarration.get('ready'), 30)
            assert not c('createLessonNarrationButton').isEnabled(), 'Finished voice variant must not be generated twice'
            assert not any(segment['approved'] for segment in bridge.lesson['segments'])
            assert not bridge.powerpointState.get('active')
            for width, height in ((1080, 700), (1366, 850)):
                window.resize(width, height)
                c('resultScroll').property('contentItem').setProperty('contentY', 0)
                pump(150)
                assert window.grabWindow().save(str(out / f'voice-ready-{width}.png'))

            downloads = root / 'exports'
            downloads.mkdir()
            bridge.rememberExportDirectory((downloads / 'anything.pptx').as_uri())
            assert Path(QUrl(bridge.downloadsDirectory).toLocalFile()) == downloads
            assert 'Lời đọc có sẵn' in QUrl(bridge.presentationDownloadSuggestion(True)).toLocalFile()
            assert bridge.libraryDownloadSuggestion([bridge.lesson['id']], True).endswith('.pptx')
            assert bridge.libraryDownloadSuggestion(['one', 'two'], False).endswith('.zip')
            assert bridge.exportSuggestion('bcbackup').endswith('.bcbackup')
            window.setProperty('page', 'settings')
            c('settingsPage').setProperty('activeTab', 5)
            pump(150)
            assert c('relocateLibraryButton').isVisible()
            assert window.grabWindow().save(str(out / 'data-settings.png'))
            # Existing target failure keeps the UI usable and all source data intact.
            (root / 'occupied/BiliClass-Data').mkdir(parents=True)
            bridge.storage.relocate((root / 'occupied').as_uri())
            until(lambda: not bridge.busy)
            assert not bridge.storage.moving and bridge.error
            (root / 'new-drive').mkdir()
            with patch('app.storage_controls.QGuiApplication.quit') as quit_app:
                bridge.storage.relocate((root / 'new-drive').as_uri())
                until(lambda: bridge.storage.transfer_record is not None, 30)
                assert quit_app.call_count == 1 and bridge.storage.moving
            record = bridge.storage.transfer_record
            assert (library.directory / 'library.db').is_file()
            assert not warnings, warnings
        finally:
            window.close()
            for worker in (bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog, bridge.browserAI,
                           bridge.english_reader, bridge.narration_preparation, bridge.classroom, bridge.voice_controls):
                worker.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()
        target = finalize_transfer(record)
        verify_transfer_start(target)
        assert user_data() == target
        replacement = Library(target)
        assert replacement.get(bridge.lesson['id'])['narration']['complete'] > 0
        assert len(KeyStore(target).public()) == 1
        assert replacement.setting('export_directory', '') == str(downloads)
        replacement.close()
        assert cleanup_previous(target) > 0
        assert not (library.directory / 'library.db').exists()
        result = {'status': 'passed', 'qml_warnings': warnings,
                  'scope': 'Real Qt/import/export settings and copied library activation; speech synthesis is a WAV fixture; quit is intercepted, no real user data moved',
                  'checks': ['sidebar order and standalone PowerPoint import', 'same narration controls; one-language guard',
                             'voice ready locks repeat; no auto approve/read', 'suggested names and remembered export directory',
                             'data settings and failed relocation recovery', 'successful copy triggers quit',
                             'post-close finalize/rebase, selected directory, reopened voice/key data and exact old-file cleanup']}
        (out / 'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
        print(json.dumps(result), flush=True)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
