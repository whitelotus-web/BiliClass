"""Generate bounded Word layout fixtures; optionally verify them in native Word."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from app.word_document import NS, M, Package, digest, keyword_annotations


def build(folder):
    folder.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(12)
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Title', 24), ('Heading 1', 16)]:
        document.styles[name].font.name = 'Arial'
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles['Title']._element.findall('.//w:pBdr', NS):
        border.getparent().remove(border)
    document.add_heading('Bài học chuyển động', 0)
    document.add_paragraph('Quan sát các đại lượng vật lý và giữ đúng số liệu, công thức khi trình bày song ngữ.')
    paragraph = document.add_paragraph()
    paragraph.add_run('Vận tốc').bold = True
    paragraph.add_run(' của vật là 12 m/s. Phần hướng dẫn này giữ nguyên tiếng Việt.')
    paragraph = document.add_paragraph('Phương trình ')
    equation = etree.SubElement(paragraph._p, '{' + M + '}oMath')
    run = etree.SubElement(equation, '{' + M + '}r')
    etree.SubElement(run, '{' + M + '}t').text = 'v = s/t'
    paragraph.add_run(' mô tả mối liên hệ giữa vận tốc, quãng đường và thời gian.')
    document.add_heading('Bảng đại lượng', 1)
    table = document.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    for row, values in zip(table.rows, [('Đại lượng', 'Ví dụ'), ('Vận tốc', 'Vận tốc 12 m/s.'), ('Gia tốc', 'Đơn vị là m/s².')], strict=True):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    document.add_paragraph('Nội dung lặp lại cần giữ đúng vị trí.')
    document.add_paragraph('Nội dung lặp lại cần giữ đúng vị trí.')
    section.header.paragraphs[0].text = 'Bài giảng Word · Vật lý'
    section.footer.paragraphs[0].text = 'Tài liệu chỉnh sửa được bằng Microsoft Word'
    source = folder / 'source.docx'
    document.save(source)
    signature = digest(source)
    results = []
    for mode in ('below', 'columns', 'english', 'keywords'):
        package = Package(source)
        regions = []
        for starts, english in [
            ('Vận tốc của vật', 'The object has a velocity of 12 m/s.'),
            ('Phương trình', 'The equation ⟦BLC0⟧ describes the relationship between velocity, distance and time.'),
            ('Vận tốc 12', 'Velocity 12 m/s.'),
        ]:
            row = next(row for row in package.rows if row['text'].startswith(starts))
            end = len('Vận tốc của vật là 12 m/s.') if starts == 'Vận tốc của vật' else len(row['text'])
            region = package.selection([dict(paragraph=row['id'], start=0, end=end)])[0]
            region['en'] = english
            if mode == 'keywords':
                region['annotations'] = keyword_annotations(region, {'vận tốc': 'velocity', 'phương trình': 'equation', 'quãng đường': 'distance', 'thời gian': 'time'})
            regions.append(region)
        output = folder / (mode + '.docx')
        package.save(output, regions, dict(mode=mode, color='1667B1', italic=True))
        results.append(dict(mode=mode, path=str(output.resolve()), sha256=digest(output)))
    assert digest(source) == signature
    return results


def render_word(rows, folder):
    import pypdfium2 as pdfium
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    word.AutomationSecurity = 3
    try:
        for row in rows:
            document = word.Documents.Open(row['path'], ReadOnly=True, AddToRecentFiles=False,
                                           ConfirmConversions=False, OpenAndRepair=False, Visible=False)
            try:
                pdf_path = (folder / (row['mode'] + '.pdf')).resolve()
                document.ExportAsFixedFormat(str(pdf_path), 17)
            finally:
                document.Close(False)
            pdf = pdfium.PdfDocument(str(pdf_path))
            row['images'] = []
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=1.5)
                image_path = folder / (row['mode'] + '-page-' + str(index + 1) + '.png')
                bitmap.to_pil().save(image_path)
                row['images'].append(str(image_path.resolve()))
                bitmap.close()
                page.close()
            row['pages'] = len(pdf)
            pdf.close()
    finally:
        word.Quit(False)
        pythoncom.CoUninitialize()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--report', default='reports/rc169-word/layout')
    parser.add_argument('--render-word', action='store_true')
    args = parser.parse_args()
    folder = Path(args.report)
    rows = build(folder)
    if args.render_word:
        render_word(rows, folder)
    (folder / 'result.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(rows, ensure_ascii=True))


if __name__ == '__main__':
    main()
