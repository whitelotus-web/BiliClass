"""Qt saved-deck voice checks; no API, personal data or model required."""
import json
import sys
import tempfile
import time
from pathlib import Path
from threading import Event
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEvent, QObject, QPointF, QTimer
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import control, pump, until

from app import speech
from app.chatgpt_handoff import inspect_returned_deck
from app.library import Library
from app.ui import Bridge

ROOT = Path(__file__).resolve().parents[1]


def main():
    QQuickStyle.setStyle('Basic')
    app = QGuiApplication([])
    checking, allow_check, synthesis, allow_voice = (Event() for _ in range(4))
    ticks, warnings = [], []
    started = time.monotonic()

    def check(_path):
        checking.set()
        assert allow_check.wait(10), 'GUI did not release cache check'
        return False

    def render(*_args):
        synthesis.set()
        assert allow_voice.wait(10), 'GUI did not release synthesis'
        raise RuntimeError('VieNeu deliberate local voice failure')

    with tempfile.TemporaryDirectory(prefix='bcl-qt-audio-', ignore_cleanup_errors=True) as temporary:
        root = Path(temporary)
        source = root / 'fixture.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(914400, 914400, 6000000, 1200000).text = 'Năng lượng được bảo toàn.\nEnergy is conserved.'
        slide.notes_slide.notes_text_frame.text = 'VI: Năng lượng được bảo toàn.\nEN: Energy is conserved.'
        deck.save(source)
        library = Library(root / 'data')
        config = {'title': 'Voice lifecycle', 'subject': 'Vật lý', 'education_level': 'THPT', 'grade': '11',
                  'conversion_format': 'parallel_columns', 'style': 'source', 'preset': 'standard'}
        lesson = library.create_external_lesson(config, library.store_source(source), inspect_returned_deck(source))
        with patch.object(Bridge, 'previewConvertedSlide'):
            with patch.object(speech, 'list_voices', return_value=[dict(speech.PAIRS[0])]):
                bridge = Bridge(library)
            engine = QQmlApplicationEngine()
            engine.warnings.connect(lambda items: warnings.extend(str(x) for x in items))
            engine.rootContext().setContextProperty('bridge', bridge)
            engine.load(str(ROOT / 'app/qml/Main.qml'))
            assert engine.rootObjects(), warnings
            timer = QTimer()
            timer.setInterval(20)
            timer.timeout.connect(lambda: ticks.append(time.monotonic()))
            timer.start()
            try:
                bridge.openLesson(lesson['id'])
                until(lambda: bool(bridge.lesson) and not bridge.busy)
                pump(150)
                assert not bridge.narrationBusy and bridge.lessonNarration['status'] == 'none'
                with patch.object(speech, 'valid_audio', side_effect=check), patch.object(speech, 'synthesize_background', side_effect=render):
                    bridge.createLessonNarrationWithVoice('en', speech.PAIRS[0]['id'], 0)
                    until(checking.is_set)
                    before = len(ticks)
                    pump(250)
                    cache_ticks = len(ticks) - before
                    assert cache_ticks >= 5 and bridge.narrationBusy
                    allow_check.set()
                    until(synthesis.is_set)
                    before = len(ticks)
                    pump(250)
                    voice_ticks = len(ticks) - before
                    assert voice_ticks >= 5 and not bridge.busy and bridge.narrationBusy
                    assert bridge.quickResult['path']
                    allow_voice.set()
                    until(lambda: not bridge.narrationBusy)
                assert bridge.narrationPreparation['failed'] > 0
                saved = library.get(lesson['id'])
                assert saved['external_deck']['reviewed_revision'] == 0
                assert not any(row['approved'] for row in saved['segments'])
                assert source.read_bytes() == (library.directory / 'sources' / saved['source']['file']).read_bytes()
                assert bridge.error
                assert 'vieneu' not in bridge.message.casefold()
                window = engine.rootObjects()[0]
                window.setProperty('page', 'voice')
                until(lambda: window.property('narrationDraft').get('phase') == 'partial')
                button = window.findChild(QObject, 'createLessonNarrationButton')
                assert button.isVisible() and button.isEnabled()
                assert button.property('text') == 'Tiếp tục tạo giọng'
                failures = bridge.narrationPreparation['failed_segments']
                assert failures[0]['slides'] == [1] and failures[0]['languages'] == ['en']
                pump(50)
                location = control(window, 'failedNarrationLocation')
                assert location is not None and 'Slide 1' in location.property('text')
                assert 'Energy is conserved.' in location.property('text')
                assert 'VieNeu' not in location.property('text')
                for width, height in ((1080, 700), (1366, 850)):
                    window.resize(width, height)
                    pump(100)
                    summary = window.findChild(QObject, 'narrationConfiguration')
                    assert summary.mapToScene(QPointF(summary.width(), 0)).x() + 15 <= button.mapToScene(QPointF()).x()
                    assert button.mapToScene(QPointF(button.width(), 0)).x() <= window.width()-28
                bridge.dismissNarrationLesson()
                pump(100)
                assert not bridge.lesson and not bridge.quickResult and not bridge.narrationPreparation
                assert not bridge.lessonNarration and not bridge.error and not bridge.narrationBusy
                assert engine.rootObjects()[0].property('page')=='voice'
                assert library.get(lesson['id'])==saved, 'Dismissing failed voice must keep the saved lesson and plan'
                assert not warnings, warnings
                report = {'status': 'passed', 'elapsed_seconds': round(time.monotonic() - started, 2),
                          'cache_heartbeat_ticks': cache_ticks, 'synthesis_heartbeat_ticks': voice_ticks,
                          'qml_warnings': warnings, 'checks': ['no automatic voice on open',
                          'Qt responsive during cache validation and synthesis',
                          'voice failure preserves deck and never approves lesson',
                          'failure text hides engine details and partial job offers Continue at both widths']}
                output = ROOT / 'reports/rc58/narration-qt.json'
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
                print(json.dumps(report), flush=True)
            finally:
                allow_check.set()
                allow_voice.set()
                timer.stop()
                bridge.narration_preparation.shutdown()
                bridge.classroom.shutdown()
                bridge.voice_controls.shutdown()
                bridge.libraryThumbnails.shutdown()
                bridge.slidePreview.shutdown()
                bridge._catalog.shutdown()
                bridge.browserAI.shutdown()
                bridge.english_reader.shutdown()
                engine.rootObjects()[0].close()
                engine.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                app.processEvents()
                library.close()


if __name__ == '__main__':
    main()
