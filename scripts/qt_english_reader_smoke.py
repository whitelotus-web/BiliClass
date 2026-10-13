"""Real PowerPoint and offline English WAV playback; isolated data and no web requests."""

import hashlib
import json
import tempfile
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from PySide6.QtCore import QMetaObject, QObject, Qt, QUrl
from PySide6.QtGui import QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import speech
from app.chatgpt_handoff import inspect_returned_deck
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    app = QGuiApplication.instance() or QGuiApplication([])
    QQuickStyle.setStyle("Basic")
    for name in ("BeVietnamPro-Regular.ttf", "BeVietnamPro-SemiBold.ttf", "BeVietnamPro-Bold.ttf"):
        QFontDatabase.addApplicationFont(str(RESOURCE_ROOT / "assets" / name))
    report = {"status": "running", "checks": [], "warnings": [], "scope": "real Office and MCI/VieNeu multilingual on this Windows machine; no uploads"}
    output = Path("reports/english-reader")
    output.mkdir(parents=True, exist_ok=True)

    def wait(condition, timeout=60):
        deadline = time.monotonic()+timeout
        while not condition() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(.02)
        assert condition(), "Timed out waiting for reader/PowerPoint"

    with tempfile.TemporaryDirectory(prefix="biliclass-reader-", ignore_cleanup_errors=True) as temporary:
        root = Path(temporary)
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        paragraphs = ["Energy and motion", "The energy of an object depends on its mass and speed. "
                      "Discuss the result in groups and explain why the total energy is conserved."]
        for index in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            texts = paragraphs if index == 0 else ["Đây là slide chỉ có tiếng Việt."]
            for y, text in enumerate(texts):
                shape = slide.shapes.add_textbox(Inches(1), Inches(1+y), Inches(11), Inches(.8))
                shape.text = text
                shape.name = "BCL_EN_title" if y == 0 and index == 0 else "Content"
                for run in shape.text_frame.paragraphs[0].runs:
                    run.font.size = Pt(24)
            slide.notes_slide.notes_text_frame.text = "VI: Bài thử\nEN: Energy and motion" if index == 0 else "VI: Bài Việt"
        path = root / "reader.pptx"
        deck.save(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        library = Library(root / "library")
        config = {"title": "English Reader smoke", "subject": "Vật lý", "education_level": "THPT", "grade": "11",
                  "conversion_format": "parallel_columns", "style": "source", "preset": "standard"}
        lesson = library.create_external_lesson(config, library.store_source(path), inspect_returned_deck(path))
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        engine.warnings.connect(lambda entries: report["warnings"].extend(map(str, entries)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), report["warnings"]
        window = engine.rootObjects()[0]
        companion = window.findChild(QObject, "companionWindow")
        plays = []
        real_play = bridge.english_reader.player.play

        def play(audio):
            plays.append(str(audio))
            return real_play(audio)

        bridge.english_reader.player.play = play
        try:
            bridge.openLesson(lesson["id"])
            wait(lambda: not bridge.busy)
            voice = bridge.voiceSettings["en"]
            assert voice.startswith("vieneu:") and voice == bridge.voiceSettings["vi"], "Shared multilingual voice missing"
            for text in paragraphs:
                speech.synthesize(text, voice, bridge.voiceSettings["rate"], library.directory / "audio")
            bridge.startPowerPoint()
            wait(lambda: bridge.powerpointState.get("active"))
            assert not plays and bridge.englishReaderState["state"] == "stopped"
            companion.setProperty("collapsed", False)
            companion.setProperty("visible", True)
            wait(lambda: window.findChild(QObject, "companionReadEnglish").property("enabled"))
            assert QMetaObject.invokeMethod(window.findChild(QObject, "companionReadEnglish"), "clicked", Qt.DirectConnection)
            print(json.dumps({"after_click": bridge.englishReaderState, "message": bridge.message,
                              "powerpoint": bridge.powerpointState}), flush=True)
            wait(lambda: bridge.englishReaderState["state"] == "reading")
            assert bridge.mascotState == "speaking"
            assert QMetaObject.invokeMethod(window.findChild(QObject, "companionPauseEnglish"), "clicked", Qt.DirectConnection)
            assert bridge.englishReaderState["state"] == "paused" and bridge.mascotState != "speaking"
            assert QMetaObject.invokeMethod(window.findChild(QObject, "companionPauseEnglish"), "clicked", Qt.DirectConnection)
            assert bridge.englishReaderState["state"] == "reading"
            wait(lambda: len(plays) == 2)
            report["checks"].append("One mascot click reads title and full native English body; real shared VieNeu WAV and MCI pause/resume")
            assert window.grabWindow().save(str(output / "reader-controls.png"))
            companion.grabWindow().save(str(output / "mascot-reader.png"))
            bridge.navigatePowerPoint("next")
            wait(lambda: bridge.powerpointState.get("slide") == 2)
            assert bridge.englishReaderState["state"] == "stopped" and bridge.mascotState != "speaking"
            count = len(plays)
            bridge.readSlideEnglish()
            assert "chưa có nội dung tiếng Anh" in bridge.message
            assert len(plays) == count
            report["checks"].append("Slide change cancels narration; Vietnamese-only slide gives notice and never auto-plays")
            bridge.stopPowerPoint()
            wait(lambda: bridge.powerpoint_worker is None)
            assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
            report["checks"].append("Read-only PowerPoint source unchanged; slideshow shutdown releases reader")
            assert not report["warnings"], report["warnings"]
            report["status"] = "passed"
        except Exception:
            print(json.dumps({"reader": bridge.englishReaderState, "message": bridge.message,
                              "powerpoint": bridge.powerpointState}, ensure_ascii=False), flush=True)
            raise
        finally:
            bridge.stopSpeech()
            bridge.english_reader.shutdown()
            if bridge.powerpoint_worker:
                bridge.stopPowerPoint()
                bridge.powerpoint_worker.wait()
            if bridge.worker:
                bridge.worker.wait()
            window.close()
            library.close()
            for handler in bridge.logger.handlers:
                handler.close()
            del engine
    (output / "native-reader.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
