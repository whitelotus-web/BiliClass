"""Exercise the real template controls and async preview builder, offline."""

import hashlib
import json
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QMetaObject, QObject, QTimer, QUrl
from PySide6.QtGui import QDesktopServices, QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import control


def pump():
    loop = QEventLoop()
    QTimer.singleShot(20, loop.quit)
    loop.exec()


def until(check):
    end = time.monotonic() + 30
    while not check():
        assert time.monotonic() < end, "Template preview timed out"
        pump()


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle("Basic")
    for p in (RESOURCE_ROOT / "assets").glob("*.ttf"):
        QFontDatabase.addApplicationFont(str(p))
    app.setFont(QFont("Be Vietnam Pro", 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else "reports/rc47/qt")
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="biliclass-universal-", ignore_cleanup_errors=True) as directory:
        library = Library(Path(directory))
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        root.setProperty("page", "new")
        root.setProperty("preferredCreationDesign", 1)
        until(lambda: root.findChild(QObject, "creationPreset") is not None)
        combo = root.findChild(QObject, "creationPreset")
        bookends = root.findChild(QObject, "templateBookends")
        assert combo.property("count") == 3 and bookends is None
        assert not bridge.templateBookends
        bridge.setTemplateBookends(True)
        pump()
        assert library.setting("template_bookends", False)
        bridge.setTemplateBookends(False)
        opened = []
        original = bridge._open_teaching_presentation
        bridge._open_teaching_presentation = lambda url, **_: opened.append(QUrl(url).toLocalFile())
        builds = []
        try:
            for index in range(3):
                combo.setProperty("currentIndex", index)
                pump()
                for mode in ("parallel_columns", "english_only"):
                    old = len(opened)
                    bridge.openLocalTemplateSample(("standard", "visual", "practice")[index], mode)
                    until(lambda: len(opened) > old)
                    path = Path(opened[-1])
                    deck = Presentation(path)
                    assert len(deck.slides) >= 4
                    text = "\n".join(
                        s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame
                    )
                    assert "{{" not in text
                    assert "What is the unit of energy?" in text
                    builds.append(
                        {
                            "preset": index,
                            "mode": mode,
                            "pages": len(deck.slides),
                            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                        }
                    )
            # Also invoke the exact QML button handler, not just its Python slot.
            cached = Path(opened[4])
            stamp = cached.stat().st_mtime_ns
            old = len(opened)
            assert QMetaObject.invokeMethod(control(root, "creationPresetView_practice"), "clicked")
            until(lambda: len(opened) > old)
            assert Path(opened[-1]) == cached and cached.stat().st_mtime_ns == stamp
        finally:
            bridge._open_teaching_presentation = original
        for width, height in ((1366, 850), (1080, 700)):
            root.setWidth(width)
            root.setHeight(height)
            pump()
            assert root.grabWindow().save(str(report / f"new-{width}.png"))
        assert not warnings, warnings
        result = {
            "status": "passed",
            "previews": builds,
            "qml_button": True,
            "preview_cache_reused": True,
            "bookends_persisted": True,
            "warnings": warnings,
        }
        (report / "result.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(json.dumps(result, ensure_ascii=False), flush=True)
        root.close()
        bridge.slidePreview.shutdown()
        bridge._catalog.shutdown()
        bridge.browserAI.shutdown()
        bridge.english_reader.shutdown()
        bridge.narration_preparation.shutdown()
        bridge.classroom.disconnect()
        engine.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        app.processEvents()
        library.close()


if __name__ == "__main__":
    main()
