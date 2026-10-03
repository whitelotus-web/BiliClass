import hashlib
import json
import sys
import tempfile
import time
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches, Pt
from PySide6.QtCore import QMetaObject, QObject, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuick import QQuickWindow
from PySide6.QtQuickControls2 import QQuickStyle

from app.input_analysis import assess_blocks
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.quick_conversion import verify_preview
from app.readiness import preparation_current
from app.ui import Bridge

app = QGuiApplication([])
QQuickStyle.setStyle('Basic')
app.setFont(QFont('Arial', 10))
root = Path('reports/quick-convert')
root.mkdir(parents=True, exist_ok=True)
Path('.runtime/quick-convert-smoke').mkdir(parents=True, exist_ok=True)
workspace = Path(tempfile.mkdtemp(prefix='flow-', dir='.runtime/quick-convert-smoke'))
library = Library(workspace / 'library')
texts = [
    'CHƯƠNG II: DÃY SỐ VÀ CẤP SỐ CỘNG | CHAPTER II: SEQUENCES AND ARITHMETIC PROGRESSIONS',
    'Mỗi đầu vào có một giá trị. Hãy thảo luận theo nhóm.',
    'Each input has one value. Explain your answer.',
]
source = workspace / 'Bài thử của thầy cô.pptx'
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
for text in texts:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(.5), Inches(.6), Inches(10.5), Inches(1.4))
    box.text = text
    box.text_frame.word_wrap = True
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
deck.save(source)
original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
bridge = Bridge(library)
engine = QQmlApplicationEngine()
warnings, failures, progress, stages = [], [], [], []
engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
engine.rootContext().setContextProperty('bridge', bridge)
engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
window = engine.rootObjects()[0]
window.setProperty('page', 'new')
window.setProperty('selectedFileName', source.name)
window.setProperty('selectedFile', QUrl.fromLocalFile(str(source.resolve())).toString())
window.findChild(QObject, 'lessonTitle').setProperty('text', 'Bài giảng chuyển đổi một lần')
window.findChild(QObject, 'lessonSubject').setProperty('text', 'Toán')
window.findChild(QObject, 'conversionProvider').setProperty('currentIndex', 1)
bridge.conversionProgress.connect(lambda current, total: progress.append([current, total]))
deadline = time.monotonic() + 240


def safe(action):
    def run():
        try:
            action()
        except Exception:
            failures.append(traceback.format_exc())
            print(failures[-1], flush=True)
            app.exit(1)
    return run


def click(name):
    target = window.findChild(QObject, name)
    assert target and target.property('enabled')
    assert QMetaObject.invokeMethod(target, 'clicked', Qt.DirectConnection)


def wait_idle(action):
    assert time.monotonic() < deadline, 'Flow timed out'
    if bridge.busy:
        QTimer.singleShot(150, safe(lambda: wait_idle(action)))
    else:
        action()


def start():
    assert not window.findChild(QObject, 'creationMode').property('visible')
    assert not window.findChild(QObject, 'sourceLanguage').property('visible')
    assert not window.findChild(QObject, 'creationTemplatesButton').property('visible')
    assert QQuickWindow.grabWindow(window).save(str(root / 'new-simple.png'))
    click('createLessonButton')
    QTimer.singleShot(150, safe(lambda: wait_idle(converted)))


def converted():
    assert not bridge.error, bridge.message
    assert window.property('page') == 'result'
    result = bridge.quickResult
    assert result['path'] and result['image'] and result['total'] >= 3, result
    assert len(progress) == 2
    assert bridge.lesson['level'] == 2 and bridge.lesson['presentation_style'] == 'source'
    assert bridge.lesson['conversion_mode'] == 'level'
    assert not any(s['approved'] for s in bridge.lesson['segments'])
    assert bridge.lesson['segments'][0]['existing_pair']
    assert bridge.lesson['segments'][0]['en'] == texts[0].split('|')[1].strip()
    assert all(s['vi'] and s['en'] for s in bridge.lesson['segments'])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
    assert verify_preview(bridge.lesson, result, library.directory).is_file()
    stages.append('PPTX one button -> assess -> both model directions -> export -> PowerPoint render')
    print(stages[-1], flush=True)
    assert QQuickWindow.grabWindow(window).save(str(root / 'result-pptx.png'))
    bridge.previewConvertedSlide(2)
    QTimer.singleShot(150, safe(lambda: wait_idle(review)))


def review():
    assert bridge.quickResult['slide'] == 2
    click('useConvertedLessonButton')
    assert window.findChild(QObject, 'wholeLessonReview').property('visible')
    assert not any(s['approved'] for s in bridge.lesson['segments'])
    # Verify the final handoff without showing a slideshow over the user's app.
    opened = []
    bridge._start_powerpoint_session = lambda *args: opened.append(args)
    click('confirmWholeLessonButton')
    assert preparation_current(bridge.lesson) and all(s['approved'] for s in bridge.lesson['segments'])
    assert len(opened) == 1 and Path(opened[0][0]).is_file()
    stages.append('One explicit whole-lesson confirmation -> prepared -> slideshow handoff')
    print(stages[-1], flush=True)
    click('editLessonDetailsButton')
    assert window.property('page') == 'editor'
    assert window.findChild(QObject, 'translateMissingButton').property('visible')
    click('backToConvertedLessonButton')
    assert window.property('page') == 'result'
    # The same quick workflow also generates a native presentation from text.
    blocks = [('Ý 1', 'VI: Có một giá trị.\nEN: There is one value.')]
    lesson = library.create('Tài liệu theo mẫu', 'Liên môn', 'THPT', '10', blocks, analysis=assess_blocks(blocks))
    bridge.openLesson(lesson['id'])
    bridge.convertCurrentLesson()
    QTimer.singleShot(150, safe(lambda: wait_idle(template)))


def template():
    assert not bridge.error, bridge.message
    result = bridge.quickResult
    assert result['path'] and result['image'] and result['segment_map']
    assert len(result['segment_map']) == result['total']
    assert all(s == bridge.lesson['segments'][0]['id'] for s in result['segment_map'])
    assert not any(s['approved'] for s in bridge.lesson['segments'])
    assert QQuickWindow.grabWindow(window).save(str(root / 'result-template.png'))
    stages.append('Text -> template slides -> PowerPoint render -> mascot segment mapping')
    assert not warnings, warnings
    app.exit(0)


QTimer.singleShot(750, safe(start))
code = app.exec()
if bridge.worker:
    bridge.cancel_event.set()
    bridge.worker.wait()
report = {'status': 'passed' if code == 0 and not failures and not warnings else 'failed',
          'stages': stages, 'warnings': warnings, 'failures': failures, 'progress': progress,
          'scope': 'Source Qt and actual local translation/PowerPoint rendering; final session handoff intercepted',
          'library': str(workspace)}
(root / 'qt-flow.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
library.close()
print(json.dumps(report, ensure_ascii=False), flush=True)
raise SystemExit(code)
