"""Crash a disposable parent; its Office helper must close only its own deck."""

import hashlib
import json
import multiprocessing
import os
import queue
import tempfile
import time
from pathlib import Path

import psutil
from pptx import Presentation

from app.powerpoint import _office_session


def abrupt_parent(path, commands, progress):
    context = multiprocessing.get_context('spawn')
    events = context.Queue()
    helper = context.Process(target=_office_session, args=(path, 1, commands, events, False, {}), daemon=True)
    helper.start()
    progress.put({'helper_pid': helper.pid})
    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        try:
            state = events.get(timeout=.2)
        except queue.Empty:
            continue
        if state.get('error'):
            progress.put({'error': state['error']})
            progress.close()
            progress.join_thread()
            return
        if state.get('opened'):
            progress.put({'opened': True})
            progress.close()
            progress.join_thread()
            os._exit(0)  # Deliberately skip multiprocessing's normal daemon cleanup.
    progress.put({'error': 'Office did not open before the fixture deadline'})
    progress.close()
    progress.join_thread()


def main():
    context = multiprocessing.get_context('spawn')
    commands, progress = context.Queue(), context.Queue()
    helper = None
    with tempfile.TemporaryDirectory(prefix='bcl-parent-exit-', ignore_cleanup_errors=True) as folder:
        path = Path(folder) / 'BiliClass parent-exit fixture.pptx'
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(path)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        parent = context.Process(target=abrupt_parent, args=(str(path), commands, progress))
        parent.start()
        try:
            opened = False
            deadline = time.monotonic() + 50
            while time.monotonic() < deadline and not opened:
                state = progress.get(timeout=50)
                assert not state.get('error'), state
                if state.get('helper_pid'):
                    helper = psutil.Process(state['helper_pid'])
                opened = bool(state.get('opened'))
            assert opened and helper is not None
            parent.join(5)
            assert parent.exitcode == 0
            started = time.monotonic()
            helper.wait(timeout=12)
            seconds = time.monotonic() - started
            assert hashlib.sha256(path.read_bytes()).hexdigest() == original
            result = dict(status='passed', helper_exit_seconds=round(seconds, 3), source_unchanged=True,
                          scope='Current Windows/Office, synthetic read-only deck, abrupt parent exit; no user Office process terminated.')
            output = Path('reports/rc76/teaching-parent-exit.json')
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result), flush=True)
        finally:
            commands.put(('stop', None))  # Always ask only this private helper to close.
            if helper is not None and helper.is_running():
                helper.wait(timeout=8)
            if parent.is_alive():
                parent.join(5)
            commands.close()
            progress.close()


if __name__ == '__main__':
    multiprocessing.freeze_support()
    main()
