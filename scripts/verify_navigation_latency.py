"""Measure confirmed Office navigation on a disposable, content-heavy deck."""

import argparse
import hashlib
import json
import statistics
import tempfile
import time
from multiprocessing import freeze_support
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from PySide6.QtCore import QCoreApplication, QEventLoop, QTimer

from app.powerpoint import PowerPointSession
from app.slideshow_shortcuts import DEFAULTS


def main():
    parser = argparse.ArgumentParser(__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    app = QCoreApplication([])
    states, errors, ticks = [], [], []
    heartbeat = QTimer()
    heartbeat.setInterval(10)
    heartbeat.timeout.connect(lambda: ticks.append(time.monotonic()))
    heartbeat.start()

    def wait(predicate, timeout=60):
        started = time.monotonic()
        while not predicate():
            loop = QEventLoop()
            QTimer.singleShot(5, loop.quit)
            loop.exec()
            assert not errors, errors
            assert time.monotonic() - started < timeout, states[-1:] or errors

    with tempfile.TemporaryDirectory(prefix='bcl-navigation-', ignore_cleanup_errors=True) as folder:
        path = Path(folder) / 'BiliClass disposable 51 slides.pptx'
        deck = Presentation()
        for number in range(51):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index in range(120):
                box = slide.shapes.add_textbox(Inches((index % 10) * .95), Inches((index // 10) * .55), Inches(.9), Inches(.5))
                box.text = f'Slide {number + 1}: item {index + 1}'
                box.text_frame.paragraphs[0].font.size = Pt(10)
        deck.save(path)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        worker = PowerPointSession(path, shortcuts=DEFAULTS)
        worker.stateChanged.connect(states.append)
        worker.failed.connect(errors.append)
        worker.start()
        latencies = []
        try:
            wait(lambda: bool(states) and states[-1].get('active'))
            for command, expected in [('next_slide', n) for n in range(2, 17)] + [('previous_slide', n) for n in range(15, 1, -1)]:
                started = time.monotonic()
                worker.navigate(command)
                wait(lambda: states[-1].get('active') and states[-1].get('slide') == expected)
                latencies.append(time.monotonic() - started)
            assert len({s['slide_id'] for s in states if s.get('active')}) == 16
            before_burst = len(states)
            for command in ['next_slide'] * 4 + ['previous_slide'] * 4:
                worker.navigate(command)
            wait(lambda: len(states) >= before_burst + 8 and states[-1].get('slide') == 2)
            assert [s['slide'] for s in states[before_burst:]] == [3, 4, 5, 6, 5, 4, 3, 2]
            worker.navigate('goto', 51)
            wait(lambda: states[-1].get('slide') == 51)
            worker.navigate('next_slide')
            worker.navigate('previous_slide')
            wait(lambda: states[-1].get('slide') == 50)
            worker.navigate('goto', 1)
            wait(lambda: states[-1].get('slide') == 1)
            worker.navigate('previous_slide')
            worker.navigate('next_slide')
            wait(lambda: states[-1].get('slide') == 2)
            assert hashlib.sha256(path.read_bytes()).hexdigest() == original
        finally:
            worker.stop()
            wait(lambda: not worker.isRunning(), timeout=12)
        result = {'status': 'passed', 'slides': 51, 'shapes_per_slide': 120,
                  'navigation_count': len(latencies), 'latencies_seconds': latencies,
                  'mean_seconds': statistics.mean(latencies), 'max_seconds': max(latencies),
                  'p95_seconds': sorted(latencies)[int(len(latencies) * .95)],
                  'max_gui_tick_seconds': max(b-a for a, b in zip(ticks, ticks[1:])),
                  'burst_order_and_boundaries': True,
                  'source_unchanged': True}
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(result, indent=2), encoding='utf-8')
        print(json.dumps({k: v for k, v in result.items() if k != 'latencies_seconds'}), flush=True)
    app.quit()


if __name__ == '__main__':
    freeze_support()
    main()
