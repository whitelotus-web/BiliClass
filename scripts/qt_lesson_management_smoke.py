"""Native Qt library/edit/trash/back/download workflow in a disposable library."""

import json
import tempfile
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import Q_ARG, QMetaObject, QObject, Qt, QUrl
from PySide6.QtGui import QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlEngine, QQmlExpression
from PySide6.QtQuickControls2 import QQuickStyle

from app import library as library_module
from app.chatgpt_handoff import inspect_returned_deck
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle("Basic")
    for name in ("BeVietnamPro-Regular.ttf", "BeVietnamPro-SemiBold.ttf", "BeVietnamPro-Bold.ttf"):
        QFontDatabase.addApplicationFont(str(RESOURCE_ROOT / "assets" / name))
    warnings, checks = [], []
    output = Path("reports/lesson-management")
    output.mkdir(parents=True, exist_ok=True)

    def wait(condition, timeout=90):
        deadline = time.monotonic()+timeout
        while not condition() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(.02)
        assert condition(), "UI condition did not complete"

    with tempfile.TemporaryDirectory(prefix="biliclass-library-", ignore_cleanup_errors=True) as folder:
        root = Path(folder)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = "The energy is conserved."
        slide.notes_slide.notes_text_frame.text = "VI: Năng lượng được bảo toàn.\nEN: The energy is conserved."
        source = root / "returned.pptx"
        deck.save(source)
        library = Library(root / "library")
        real_now = library_module.now
        config = {"title": "Bài giảng mẫu", "subject": "Vật lý", "education_level": "THPT", "grade": "11",
                  "conversion_format": "parallel_columns", "style": "source", "preset": "standard"}
        try:
            library_module.now = lambda: "2026-10-01T03:00:00+00:00"
            first = library.create_external_lesson(config, library.store_source(source), inspect_returned_deck(source))
            library_module.now = lambda: "2026-10-12T03:00:00+00:00"
            second = library.create("Bài học mới hơn", "Toán", "THPT", "10", [("Đoạn", "Giải thích kết quả.")])
        finally:
            library_module.now = real_now
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        engine.warnings.connect(lambda entries: warnings.extend(map(str, entries)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]

        def item(name):
            value = window.findChild(QObject, name)
            if value is None:
                def visual(parent):
                    for child in parent.childItems():
                        if child.objectName() == name:
                            return child
                        found = visual(child)
                        if found is not None:
                            return found
                    return None
                value = visual(window.contentItem())
            assert value, name
            return value

        def click(name):
            button = item(name)
            assert button.property("enabled"), name
            assert QMetaObject.invokeMethod(button, "clicked", Qt.DirectConnection)
            app.processEvents()

        def shot(name):
            app.processEvents()
            assert window.grabWindow().save(str(output / name))

        def filtered():
            value = window.property('libraryResult')
            return value.toVariant() if hasattr(value, 'toVariant') else value

        try:
            window.setProperty("page", "library")
            for _ in range(30):
                app.processEvents()
                time.sleep(.02)
            wait(lambda: item("lessonCard_"+first["id"]).isVisible())
            assert bridge.lessons[0]["id"] == second["id"]
            assert "01/10/2026" in item("lessonCreated_"+first["id"]).property("text")
            sort = item("lessonSortOrder")
            sort.setProperty("currentIndex", 1)
            QMetaObject.invokeMethod(sort, "activated", Qt.DirectConnection, Q_ARG(int, 1))
            assert bridge.lessons[0]["id"] == first["id"]
            checks.append("Creation timestamps shown; chronology dropdown changes and persists order")
            item('librarySearch').setProperty('text', 'BAI GIANG MAU')
            app.processEvents()
            value = filtered()
            assert value['total'] == 1 and value['rows'][0]['id'] == first['id']
            click('libraryClearFilters')
            subject = item('librarySubjectFilter')
            subject.setProperty('currentIndex', next(i for i,r in enumerate(bridge.librarySubjects) if r['id']=='Toán'))
            QMetaObject.invokeMethod(subject, 'activated', Qt.DirectConnection, Q_ARG(int, subject.property('currentIndex')))
            app.processEvents()
            assert filtered()['rows'][0]['id'] == second['id']
            state = item('libraryStateFilter'); state.setProperty('currentIndex', 1)
            app.processEvents()
            assert filtered()['total'] == 0
            click('libraryClearFilters')
            assert filtered()['total'] == 2
            window.setProperty('page', 'home'); app.processEvents()
            assert item('homeWorkPanel').isVisible() and not item('libraryFilters').isVisible()
            shot('home.png')
            window.setProperty('page', 'library'); app.processEvents()
            assert not item('homeWorkPanel').isVisible() and item('libraryFilters').isVisible()
            checks.append('Unaccented search; combined subject/state filters; clear; distinct home/library')
            shot("library.png")
            click("lessonEdit_"+first["id"])
            assert item("lessonMetadataDialog").property("visible")
            item("metadataLessonTitle").setProperty("text", "Tên bài đã sửa")
            click("saveLessonMetadataButton")
            assert library.get(first["id"])["title"] == "Tên bài đã sửa"
            assert library.get(first["id"])["created_at"] == first["created_at"]
            checks.append("Row edit saves lesson information without altering creation date")
            click("lessonDelete_"+second["id"])
            click("cancelDeleteLessonButton")
            assert len(bridge.lessons) == 2
            click("lessonDelete_"+second["id"])
            click("confirmDeleteLessonButton")
            wait(lambda: len(bridge.lessons) == 1 and len(bridge.trashedLessons) == 1)
            trash = item("lessonTrashDialog")
            QMetaObject.invokeMethod(trash, "open", Qt.DirectConnection)
            wait(lambda: item("restoreLesson_"+second["id"]).isVisible())
            click("restoreLesson_"+second["id"])
            QMetaObject.invokeMethod(trash, "close", Qt.DirectConnection)
            wait(lambda: len(bridge.lessons) == 2)
            checks.append("Delete confirmation can cancel; trash and restore retain lesson")
            click("lessonOpen_"+first["id"])
            wait(lambda: not bridge.busy and bool(bridge.quickResult.get("image")))
            assert window.property("page") == "result"
            assert item("backToLibraryButton").isVisible() and item("downloadPresentationButton").isVisible()
            assert not item("openConvertedDeckButton").isVisible()
            shot("slide.png")
            window.setProperty("width", 1080)
            app.processEvents()
            shot("slide-1080.png")
            window.setProperty("width", 1366)
            saved = root / "lesson-download.pptx"
            dialog = item("presentationSaveDialog")
            expression = QQmlExpression(QQmlEngine.contextForObject(window), window,
                                       "presentationSave.options = FileDialog.DontUseNativeDialog")
            expression.evaluate()
            assert not expression.hasError(), expression.error().toString()
            click("downloadPresentationButton")
            assert dialog.setProperty("selectedFile", QUrl.fromLocalFile(str(saved)))
            QMetaObject.invokeMethod(dialog, "accepted", Qt.DirectConnection)
            QMetaObject.invokeMethod(dialog, "close", Qt.DirectConnection)
            wait(lambda: not bridge.busy)
            assert saved.read_bytes() == source.read_bytes()
            checks.append("Download button/save dialog preserves exact PPTX; back/result controls fit small window")
            click("backToLibraryButton")
            assert window.property("page") == "library"
            assert not warnings, warnings
        finally:
            if bridge.worker:
                bridge.worker.wait()
            bridge.english_reader.shutdown()
            window.close()
            library.close()
            for handler in bridge.logger.handlers:
                handler.close()
            del engine
    report = {"status": "passed", "checks": checks, "warnings": warnings,
              "scope": "Qt native workflow and Office preview; isolated library; no web requests or teacher data"}
    (output / "ui.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
