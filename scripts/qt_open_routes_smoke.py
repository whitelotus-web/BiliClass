"""Owned Qt/Office fixtures: tool opens restore Mascot and native F5 stays bound."""

import json
import os
import sys
from multiprocessing import freeze_support
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'tests'))

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QObject
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.lesson_narration import variant_id
from app.library import Library
from app.office_retry import call as office_call
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import control, pump, until
from tests.test_teaching_session import prepared_lesson


def main():
    out=Path(sys.argv[1] if len(sys.argv)>1 else 'reports/rc165/open-routes').resolve()
    out.mkdir(parents=True,exist_ok=False)
    lesson,preview,directory,source=prepared_lesson(out)
    original=source.read_bytes()
    input_path=out/'Upload-source.pptx'
    deck=Presentation()
    for text in ('First independent slide.', 'Second independent slide.'):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text=text
    deck.save(input_path)
    input_bytes=input_path.read_bytes()
    os.environ['BILICLASS_DATA']=str(directory)
    app=QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    library=Library(directory)
    bridge=Bridge(library)
    engine=QQmlApplicationEngine()
    warnings=[]
    engine.warnings.connect(lambda values:warnings.extend(map(str,values)))
    engine.rootContext().setContextProperty('bridge',bridge)
    engine.load(RESOURCE_ROOT/'qml/Main.qml')
    assert engine.rootObjects(),warnings
    window=engine.rootObjects()[0]
    companion=control(window,'companionWindow')
    checks=[]
    try:
        until(lambda:not bridge.busy)
        # Reproduces the previously shown -> hidden -> tool Open lifecycle.
        bridge.showCompanion(companion)
        companion.hide()
        for route in ('library','conversion_input','narration'):
            if route=='library':
                bridge.openTeachingLibraryLesson(lesson['id'])
            elif route=='conversion_input':
                bridge.openInputPowerPoint(input_path.as_uri())
            else:
                bridge._show_lesson(library.get(lesson['id']),preview,auto_repair=False,load_preview=False,destination='voice')
                bridge.openNarratedPowerPoint(variant_id(lesson['narration']))
            until(lambda:not bridge.busy and bridge.powerpointState.get('opened') and not bridge.teachingTransitioning,timeout=70)
            assert companion.isVisible(), 'Explicit tool Open must restore a previously hidden Mascot'
            assert not companion.property('collapsed')
            bound=bridge.teaching_session.state
            assert bound['bound'] and not bound.get('active')
            label=companion.findChild(QObject,'companionBoundLesson')
            assert label.property('text')==bound['title']
            if route=='narration':
                assert bound['lesson_id']==lesson['id']
            # Start through PowerPoint's own object, not the Mascot start action.
            import win32com.client

            application=office_call(lambda:win32com.client.GetActiveObject('PowerPoint.Application'))
            owned=Path(bridge.teaching_session.bound['path']).resolve()
            presentation=None
            for candidate in application.Presentations:
                try:
                    if Path(candidate.FullName).resolve()==owned:
                        presentation=candidate
                        break
                except Exception:
                    continue  # An unrelated Explorer preview may deny automation.
            assert presentation is not None, 'Only the exact owned fixture may be presented'
            show=office_call(lambda p=presentation:p.SlideShowSettings.Run())
            until(lambda:bridge.powerpointState.get('active'),timeout=30)
            assert bridge.companionCanRead and bridge.powerpointState.get('hwnd')
            pump(100)
            panel_geometry=(companion.x(),companion.y(),companion.width(),companion.height())
            companion.findChild(QObject,'companionNext').clicked.emit()
            until(lambda:bridge.powerpointState.get('slide')==2,timeout=30)
            assert int(show.View.Slide.SlideIndex)==2
            pump(100)
            assert panel_geometry==(companion.x(),companion.y(),companion.width(),companion.height()), 'Empty slide must not resize or reposition Mascot'
            companion.findChild(QObject,'companionPrevious').clicked.emit()
            until(lambda:bridge.powerpointState.get('slide')==1,timeout=30)
            assert int(show.View.Slide.SlideIndex)==1
            pump(100)
            assert panel_geometry==(companion.x(),companion.y(),companion.width(),companion.height())
            assert companion.grabWindow().save(str(out/(route+'.png')))
            del show,presentation,application
            checks.append(route+': hidden Mascot restored, exact lesson bound, native PowerPoint show and Mascot next/previous')
            print('Passed '+route,flush=True)
            bridge.stopPowerPoint()
            until(lambda:bridge.powerpoint_worker is None,timeout=30)
            companion.hide()
        assert source.read_bytes()==original and input_path.read_bytes()==input_bytes
        assert not warnings,warnings
        result=dict(status='passed',checks=checks,qml_warnings=warnings,source_unchanged=True,
                    scope='Current Windows/Office with isolated library, not the receiving machine')
        (out/'result.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
    finally:
        bridge.stopPowerPoint()
        until(lambda:bridge.powerpoint_worker is None,timeout=30)
        companion.setProperty('disposing',True)
        companion.hide()
        window.hide()
        for worker in (bridge._catalog,bridge.browserAI,bridge.slidePreview,bridge.libraryThumbnails,
                       bridge.narrationPreferences,bridge.narration_preparation,bridge.english_reader,
                       bridge.voice_controls,bridge._companion_preferences,bridge.classroom):
            worker.shutdown()
        engine.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
        pump(20)
        library.close()


if __name__=='__main__':
    freeze_support()
    main()
