"""Isolated native Word multi-file flow; no teacher files, network or Office launch."""

import json
import sys
import tempfile
from pathlib import Path
from threading import Event
from unittest.mock import patch

from docx import Document
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlExpression
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.google_translation import KeyStore
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.word_batch_ui import WordBatch
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT / 'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/word-batch-qt')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-word-batch-', ignore_cleanup_errors=True) as temporary:
        directory = Path(temporary)
        library = Library(directory / 'data')
        KeyStore(library.directory).add('AIza' + 'a' * 35)
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        window.resize(1366, 850)
        window.setProperty('page', 'word')
        window.show()
        window.requestActivate()
        until(window.isActive)
        entered, release = Event(), Event()
        block = [True]
        calls, opened = [], []

        def c(name):
            return control(window, name)

        class Translator:
            def configure_terms(self, *_):
                pass

            def configure_knowledge(self, *_):
                pass

            def translate(self, rows, cancel, progress):
                calls.append(list(rows))
                if block[0] and len(calls) == 2:
                    entered.set()
                    assert release.wait(15)
                if cancel.is_set():
                    raise InterruptedError()
                progress(len(rows), len(rows))
                return [row.replace('Bài học', 'Lesson').replace('Vận tốc', 'Velocity') for row in rows]

        try:
            with patch('app.translation_knowledge.lookup', return_value=None), patch('app.translation_provider.build_translator', lambda *_: Translator()), patch('app.word_ui.QDesktopServices.openUrl', lambda url: opened.append(url.toString()) or True):
                until(lambda: bridge.googleSettings.ready)
                sources = []
                for number in range(2):
                    path = directory / f'Bài học {number + 1}.docx'
                    document = Document()
                    document.add_paragraph(f'Bài học {number + 1}.')
                    document.add_paragraph('Vận tốc 12 m/s.')
                    document.save(path)
                    sources.append(path)
                original = [path.read_bytes() for path in sources]
                picker = QQmlExpression(engine.rootContext(), c('wordFileDialog'), 'fileMode === 1')
                multiple, undefined = picker.evaluate()
                assert multiple and not undefined and not picker.hasError(), picker.error().toString()
                assert QMetaObject.invokeMethod(c('wordSourceArea'), 'filesDropped', Q_ARG('QVariant', [QUrl.fromLocalFile(str(path)) for path in sources]))
                pump(80)
                batch = bridge.wordDocuments.batch
                assert c('wordPage').property('batchMode') and batch.state['phase'] == 'draft'
                assert c('wordWholeDocument').property('checked') and not c('wordWholeDocument').isEnabled()
                assert not c('wordDocumentPanels').isVisible() and c('wordBatchPanel').isVisible()
                assert c('wordFormatPreview').isVisible() and not c('wordTitle').isEnabled()
                assert c('wordSubject').isEnabled() and c('wordPreview').isEnabled()
                assert not c('wordStop').isVisible() and not c('wordSave').isVisible()
                assert not calls and [path.read_bytes() for path in sources] == original
                for width, height in ((1080, 700), (1366, 850), (1760, 990)):
                    window.resize(width, height)
                    pump(80)
                    assert c('wordPreview').mapRectToScene(c('wordPreview').boundingRect()).right() <= window.width()
                    assert window.grabWindow().save(str(report / f'word-batch-draft-{width}.png'))
                window.resize(1366, 850)
                pump(60)
                click(window, c('wordPreview'))
                until(entered.is_set)
                assert batch.state['running'] and bridge.busy
                assert not c('wordSubject').isEnabled() and not c('wordChooseFile').isEnabled()
                assert not c('wordDismissInput').isEnabled() and c('wordStop').isVisible()
                assert batch.state['counts']['done'] == 1
                assert window.grabWindow().save(str(report / 'word-batch-working.png'))
                click(window, c('wordStop'))
                assert batch.state['stopping']
                release.set()
                until(lambda: not bridge.busy)
                assert batch.state['phase'] == 'paused' and batch.state['counts']['waiting'] == 1
                assert batch.state['counts']['done'] == 1 and c('wordPreview').property('text') == 'Tiếp tục chuyển đổi'
                restored = WordBatch(bridge.wordDocuments)
                assert restored.state['counts']['done'] == 1 and restored.state['canResume']
                first_id = batch.state['rows'][0]['word_id']
                block[0] = False
                click(window, c('wordPreview'))
                until(lambda: not bridge.busy)
                assert batch.state['complete'] and batch.state['counts']['done'] == 2
                assert batch.state['rows'][0]['word_id'] == first_id and len(calls) == 3
                assert not c('wordStop').isVisible() and c('wordBatchPanel').isVisible()
                assert c('wordPreview').property('text') == 'Về bài giảng'
                assert not c('wordDocumentPanels').isVisible()
                assert [path.read_bytes() for path in sources] == original
                for path in sources:
                    assert path.with_name(path.stem + ' (Song ngu VI-EN).docx').is_file()
                for theme in ('light', 'dark'):
                    bridge.setApplicationTheme(theme)
                    pump(60)
                    assert window.grabWindow().save(str(report / f'word-batch-complete-{theme}.png'))
                click(window, c('wordDismissInput'))
                pump(60)
                assert not batch.state['total'] and not c('wordPage').property('batchMode')
                assert not c('wordFormatPreview').isVisible()
                assert c('wordContinuousText').property('text') == ''
                assert len(bridge.wordDocuments.library) == 2
                assert QMetaObject.invokeMethod(c('wordPage'), 'openDocuments', Q_ARG('QVariant', [sources[0].as_uri()]))
                until(lambda: not bridge.busy)
                assert not c('wordPage').property('batchMode') and c('wordDocumentPanels').isVisible()
                assert c('wordWholeDocument').isEnabled()
                if c('wordWholeDocument').property('checked'):
                    click(window, c('wordWholeDocument'))
                    until(lambda: not bridge.busy)
                assert c('wordSelectByHighlight').isEnabled() and c('wordSelectByCheckbox').isEnabled()
                click(window, c('wordDismissInput'))
                assert c('wordContinuousText').property('text') == ''
                assert not opened and not warnings, (opened, warnings)
                result = dict(status='passed', api='fixture only', warnings=warnings,
                    checks=['Multi-select locks whole-document mode and replaces manual selection with the Word-only queue',
                        'One-file selection retains manual/whole-document choices',
                        'Shared configuration is editable before Start and pinned during resume/retry',
                        'Sequential conversion; Stop/Resume and restored queue skip completed files',
                        'Completion remains visible; outputs saved beside each source; original bytes unchanged',
                        'Dismiss clears both queue and rendered source; light/dark and three window sizes'])
                (report / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
                print(json.dumps(result, ensure_ascii=False))
        finally:
            release.set()
            window.close()
            for component in (bridge.batchQueue, bridge.slidePreview, bridge.googleSettings, bridge._library_preferences,
                              bridge._catalog, bridge.classroom, bridge.libraryThumbnails, bridge.voice_controls,
                              bridge.browserAI, bridge.english_reader, bridge.narration_preparation, bridge.inputDocuments):
                component.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
