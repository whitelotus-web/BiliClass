"""Show exact issue pages in both conversion and narration workspaces."""

import json
import tempfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge

from .qt_result_preview_mascot_smoke import click, control, pump


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    with tempfile.TemporaryDirectory(prefix='biliclass-checks-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        source = folder/'fixture.pptx'
        deck = Presentation()
        for _ in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = 'Cơ năng\nMechanical energy'
            slide.notes_slide.notes_text_frame.text = 'VI: Cơ năng\nEN: Mechanical energy'
        deck.save(source)
        original = source.read_bytes()
        library = Library(folder/'data')
        inspection = inspect_returned_deck(source)
        inspection['slide_quality'] = {'counts': {'FAIL': 1, 'CHECK': 1, 'PASS': 0}, 'slides': [
            {'slide': 1, 'issues': [{'code': 'possible_overlap', 'severity': 'CHECK', 'excerpt': 'Cơ năng'}]},
            {'slide': 2, 'issues': [{'code': 'text_overflow', 'severity': 'FAIL', 'excerpt': 'Mechanical energy'}]}]}
        lesson = library.create_external_lesson({'title': 'Fixture', 'subject': 'Vật lý', 'style': 'source',
            'education_level': 'THPT', 'grade': '11', 'preset': 'standard', 'conversion_format': 'parallel_columns'},
            library.store_source(source), inspection)
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT/'qml/Main.qml')
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        report = Path('reports/rc138/lesson-checks')
        report.mkdir(parents=True, exist_ok=True)
        try:
            bridge._browser_state.update(phase='completed', lessonId=lesson['id'])
            bridge._show_lesson(lesson, external_preview(lesson, library.directory),
                                auto_repair=False, destination='chatgpt', load_preview=False)
            for page, name in [('chatgpt', 'completionChecks'), ('voice', 'voiceChecks')]:
                root.setProperty('page', page)
                checks = control(root, name)
                checks.setProperty('expanded', True)
                for width, height in ((1080, 700), (1366, 850)):
                    root.resize(width, height)
                    pump(150)
                    assert checks.isVisible() and checks.property('report')['failures'] == 1
                    assert checks.property('page')['slide'] == 2
                    with patch.object(Bridge, 'previewConvertedSlide') as show:
                        click(root, control(root, name+'View'))
                        show.assert_called_once_with(2)
                    assert root.grabWindow().save(str(report/f'{page}-{width}.png'))
            assert source.read_bytes() == original
            assert not warnings, warnings
            (report/'result.json').write_text(json.dumps(dict(status='passed', qml_warnings=warnings)))
            print('Qt exact slide checks in Conversion and Voice passed; no QML warnings.')
        finally:
            control(root, 'companionWindow').setProperty('disposing', True)
            root.hide()
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview',
                         '_catalog', 'browserAI', 'english_reader', 'narration_preparation',
                         'narrationPreferences', '_companion_preferences', 'classroom', 'voice_controls', 'inputDocuments'):
                getattr(bridge, name).shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
