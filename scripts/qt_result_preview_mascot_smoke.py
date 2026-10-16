"""Disposable Qt flow: result order, queued preview, narration and mascot motion.

PNG fixtures test preview caching only; they are never native QA evidence.
No teacher library, API account, model or audio output is used.
"""

import hashlib
import json
import tempfile
import time
import wave
from pathlib import Path
from threading import Event
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QEventLoop, QMetaObject, QObject, QPointF, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from app.branding import configure_application
from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.conversion_input import creation_config
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.slide_fingerprint import environment_key
from app.ui import Bridge


def pump(ms=30):
    loop = QEventLoop()
    QTimer.singleShot(ms, loop.quit)
    loop.exec()


def until(check, timeout=15):
    end = time.monotonic() + timeout
    while not check():
        assert time.monotonic() < end, "Qt result flow timed out"
        pump(10)


def control(root, name):
    found = root.findChild(QObject, name)
    if found is not None:
        return found
    pending = [root.contentItem()]
    while pending:
        item = pending.pop()
        if item.objectName() == name:
            return item
        pending.extend(item.childItems())
    raise AssertionError(name)


def click(window, item):
    parent = item.parentItem()
    while parent is not None:
        if parent.metaObject().indexOfProperty("contentY") >= 0:
            y = item.mapToItem(parent, QPointF(0, 0)).y()
            if y < 0 or y + item.height() > parent.height():
                parent.setProperty(
                    "contentY",
                    max(
                        0,
                        min(
                            parent.property("contentHeight") - parent.height(),
                            parent.property("contentY") + y - 80,
                        ),
                    ),
                )
                pump(100)
        parent = parent.parentItem()
    point = item.mapToScene(QPointF(item.width() / 2, item.height() / 2)).toPoint()
    assert item.isVisible() and item.isEnabled()
    assert 0 <= point.y() < window.height(), point
    QTest.mouseClick(window, Qt.LeftButton, Qt.NoModifier, point)
    pump(80)


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle("Basic")
    for font in (RESOURCE_ROOT / "assets").glob("*.ttf"):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont("Be Vietnam Pro", 10))
    import sys
    reports = Path(sys.argv[1] if len(sys.argv) > 1 else "reports/rc45")
    reports.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="biliclass-result-", ignore_cleanup_errors=True) as folder:
        library = Library(Path(folder))
        source = Path(folder) / "fixture.pptx"
        deck = Presentation()
        for number in (1, 2, 3):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for text, lang, y in [
                ("Năng lượng được bảo toàn.", "vi-VN", 1),
                ("Energy is conserved.", "en-US", 3),
            ]:
                shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(8), Inches(1))
                shape.text = text
                p = shape.text_frame.paragraphs[0]
                p.font.name, p.font.size = "Arial", Pt(24)
                p.font._rPr.set("lang", lang)
        deck.save(source)
        stored = library.store_source(source)
        inspection = inspect_returned_deck(source)
        lesson = library.create_external_lesson(
            creation_config(
                "Bài 6 – Dao Động Tắt Dần, Dao Động Cưỡng Bức & Hiện Tượng Cộng Hưởng · Bài thử giao diện với tên dài", "Vật lý", "THPT", "11", "parallel_columns", "standard", "source"
            ),
            stored,
            inspection,
        )
        # A layout-only failure must not block safe, recognized speech.
        lesson["slide_quality"]["counts"]["FAIL"] = 1
        library.set_setting("teacher", "Cô Nguyễn Diệu Linh")
        bridge = Bridge(library)
        bridge._auto_reviewed = {lesson["id"]}
        bridge._show_lesson(lesson, external_preview(lesson, library.directory))
        digest = stored["sha256"]
        cache = library.directory / "temp/slide-quality" / digest[:24]
        cache.mkdir(parents=True, exist_ok=True)
        entries = {}
        for n, color in enumerate(["#dcecff", "#cce9d8", "#fff1d6"], 1):
            png = cache / f"slide-{n}.png"
            Image.new("RGB", (800, 450), color).save(png)
            entries[str(n)] = {
                "environment": environment_key(),
                "image_sha256": hashlib.sha256(png.read_bytes()).hexdigest(),
            }
        (cache / "render-cache.json").write_text(
            json.dumps({"version": 2, "sha256": digest, "slides": entries})
        )
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        window.resize(1366, 850)
        window.setProperty("page", "result")
        until(lambda: not bridge.previewBusy)
        pump(200)
        def c(name):
            return control(window, name)
        teacher = c("sidebarTeacherName")
        assert teacher.property("text") == "Cô Nguyễn Diệu Linh"
        assert not teacher.property("truncated")
        help_button = c('sidebarHelpButton')
        assert abs(help_button.x()+help_button.width()/2-help_button.parentItem().width()/2) <= 1
        assert [row['label'] for row in bridge.conversionFormats] == ['Song ngữ Việt–Anh','Tiếng Anh 100%']
        assert c("lessonSlidePreview").y() < c("lessonNarrationPanel").y()
        assert not c("lessonDetailsPanel").isVisible()
        assert window.findChild(QObject, "repairConvertedLessonButton") is None
        assert window.grabWindow().save(str(reports / "result.png"))
        responsive = []
        for width, height in ((1080, 700), (1366, 850)):
            window.resize(width, height)
            pump(150)
            for name in ('downloadPresentationButton', 'useConvertedLessonButton', 'createLessonNarrationButton'):
                item = c(name)
                point = item.mapToScene(QPointF(0, 0))
                assert point.x() >= 0 and point.x() + item.width() <= width, (name, point.x(), item.width(), width)
                assert item.isVisible()
            scroll = c('resultScroll')
            assert scroll.property('contentWidth') <= scroll.width()
            status=c('resultSaveStatus');download=c('downloadPresentationButton')
            status_center=status.mapToScene(QPointF(0,status.height()/2)).y()
            action_center=download.mapToScene(QPointF(0,download.height()/2)).y()
            assert abs(status_center-action_center)<=1, (width,status_center,action_center)
            nav=c('resultSlideNavigation');counter=c('resultSlideCounter')
            nav_center=nav.mapToScene(QPointF(nav.width()/2,0)).x()
            counter_center=counter.mapToScene(QPointF(counter.width()/2,0)).x()
            assert abs(nav_center-counter_center)<=1, (width,nav_center,counter_center)
            responsive.append([width, height])
            assert window.grabWindow().save(str(reports / ('result-' + str(width) + '.png')))

        # Deliberately slow the worker to exercise coalescing and UI ticking.
        render = bridge.slidePreview.session.render
        rendered = []

        def delayed(*args):
            time.sleep(0.15)
            rendered.append(args[-1])
            return render(*args)

        bridge.slidePreview.session.render = delayed
        beats = []
        timer = QTimer()
        timer.setInterval(20)
        timer.timeout.connect(lambda: beats.append(time.monotonic()))
        timer.start()
        bridge.previewConvertedSlide(2)
        bridge.previewConvertedSlide(1)
        bridge.previewConvertedSlide(3)
        assert bridge.previewBusy and not bridge.busy
        until(lambda: not bridge.previewBusy)
        assert bridge.quickResult["slide"] == 3 and bridge.quickResult["image"].endswith("slide-3.png")
        assert rendered == [2, 3], rendered
        assert len(beats) >= 8
        timer.stop()
        # A real worker remains blocked inside synthesis. The visible clock must
        # continue even though it cannot emit chunk progress during this call.
        entered,release=Event(),Event()
        def held_synthesis(*args):
            entered.set()
            assert release.wait(10)
            raise InterruptedError('Disposable cancellation fixture')
        with patch('app.speech.synthesize_background',side_effect=held_synthesis):
            bridge.createLessonNarration('en')
            until(entered.is_set)
            pump(1250)
            elapsed=bridge.narrationClock.elapsedSeconds
            assert elapsed>=1 and bridge.lessonNarration['complete']==0
            label=c('lessonNarrationElapsed')
            assert label.isVisible() and '00:00:00' not in label.property('text')
            bridge.cancelLessonNarration();release.set()
            until(lambda: bridge.narration_preparation.worker is None)
            frozen=bridge.narrationClock.elapsedSeconds
            pump(1100)
            assert bridge.narrationClock.elapsedSeconds==frozen
            assert not bridge.narration_preparation.clock.isActive()
        narration = []
        bridge.narration_preparation.start = lambda *args: narration.append(args)
        assert c("createLessonNarrationButton").isEnabled()
        click(window, c("createLessonNarrationButton"))
        assert len(narration) == 1 and narration[0][-1]["total"] > 0
        from app.lesson_narration import make_plan
        plan,_=make_plan(bridge._lesson,'en',bridge.voiceSettings)
        for key in plan['keys']['en']:
            wav=library.directory/'audio'/(key+'.wav')
            wav.parent.mkdir(exist_ok=True)
            with wave.open(str(wav),'wb') as writer:
                writer.setparams((1,2,24000,0,'NONE','not compressed'))
                writer.writeframes(b'\x01\x00'*240)
        bridge._lesson=library.update_narration(lesson['id'],plan)
        bridge.narrationChanged.emit()
        pump(120)
        assert bridge.lessonNarration['ready']
        assert not c('createLessonNarrationButton').isEnabled()
        assert 'Đã tạo' in c('narrationEnglish').property('text')
        assert c('narrationEnglish').isEnabled()
        assert c('narrationReadyNextStep').isVisible()
        assert not c('wholeLessonReview').property('visible'), 'Completion must not open a dialog automatically'
        bridge._lesson['slide_quality']['counts']['FAIL'] = 1
        bridge.changed.emit();pump(50)
        click(window,c('continueAfterNarrationButton'))
        assert c('wholeLessonReview').property('visible')
        assert 'Tiếng Anh' in c('teachingNarrationSummary').property('text')
        assert not c('confirmWholeLessonButton').isEnabled(), 'Layout FAIL must not be silently approved'
        assert window.grabWindow().save(str(reports / 'teaching-preparation.png'))
        QMetaObject.invokeMethod(c('wholeLessonReview'),'close',Qt.DirectConnection);pump(100)
        bridge._lesson['slide_quality']['counts']['FAIL'] = 0
        bridge.changed.emit();pump(50)
        assert c('embedPresentationVoice').isVisible() and c('embedPresentationVoice').property('checked')
        click(window,c('embedPresentationVoice'))
        assert not window.property('includeDownloadVoice')
        click(window,c('embedPresentationVoice'))
        assert window.property('includeDownloadVoice')

        # All six choices use the shipped sample here, without changing defaults.
        from app import speech
        from app.lesson_narration import variant_id
        original_preferences = library.setting('voice_options', None)
        played = []
        with patch.object(bridge, '_play_audio', side_effect=lambda path: played.append(path)):
            for index, voice in enumerate(bridge.multilingualVoices):
                choice = c('lessonNarrationVoiceChoice')
                choice.setProperty('currentIndex', index)
                QMetaObject.invokeMethod(choice, 'activated', Qt.DirectConnection, Q_ARG(int, index))
                assert window.property('narrationVoice') == voice['id']
                click(window, c('previewLessonNarrationVoice'))
                until(lambda: len(played) == index + 1)
                click(window, c('stopLessonNarrationPreview'))
        assert library.setting('voice_options', None) == original_preferences
        click(window, c('createLessonNarrationButton'))
        assert narration[-1][2]['voice'] == bridge.multilingualVoices[-1]['id']
        for mode, voice in [('vi', 'vieneu:Hải Đăng'), ('both', 'vieneu:Mai Anh')]:
            plan, _ = make_plan(bridge._lesson, mode, {'voice': voice, 'rate': 0},
                                playback='both_controls' if mode == 'both' else 'separate')
            for keys in plan['keys'].values():
                for key in keys:
                    path = library.directory/'audio'/(key+'.wav')
                    path.parent.mkdir(exist_ok=True)
                    with wave.open(str(path), 'wb') as writer:
                        writer.setparams((1,2,24000,0,'NONE','not compressed'));writer.writeframes(b'\x01\x00'*240)
            bridge._lesson = library.update_narration(lesson['id'], plan)
        bridge._catalog.refresh(force=True)
        until(lambda: any(v['ready'] and 'riêng và chung' in v['title'] for v in bridge.narrationVariants))
        bridge.narrationChanged.emit();pump(100)
        vi = next(v for v in bridge.narrationVariants if v['title'].startswith('Tiếng Việt'))
        choice = c('lessonNarrationVariantChoice')
        index = next(i for i,v in enumerate(bridge.narrationVariants) if v['id']==vi['id'])
        choice.setProperty('currentIndex', index)
        QMetaObject.invokeMethod(choice,'activated',Qt.DirectConnection,Q_ARG(int,index))
        assert bridge._lesson_voice_settings()['vi'] == 'vieneu:Hải Đăng'
        assert not bridge._lesson_voice_settings()['en']
        assert not c('createLessonNarrationButton').isEnabled() and c('narrationVietnamese').isEnabled()
        click(window,c('narrationEnglish'))
        assert c('createLessonNarrationButton').isEnabled(), 'English has not been made with this voice'
        assert not c('narrationReadyNextStep').isVisible()
        click(window,c('narrationVietnamese'))
        assert not c('createLessonNarrationButton').isEnabled()
        rate=c('lessonNarrationRateChoice');rate.setProperty('currentIndex',4)
        QMetaObject.invokeMethod(rate,'activated',Qt.DirectConnection,Q_ARG(int,4));pump(100)
        assert c('createLessonNarrationButton').isEnabled(), 'New speed is a separate product'
        rate.setProperty('currentIndex',3)
        QMetaObject.invokeMethod(rate,'activated',Qt.DirectConnection,Q_ARG(int,3));pump(100)
        assert not c('createLessonNarrationButton').isEnabled()
        assert 'Tiếng Việt' in c('presentationSaveDialog').property('title')
        assert 'Hải Đăng' in QUrl(bridge.presentationDownloadSuggestion(True)).toLocalFile()
        for width in (1080,1366):
            window.resize(width,850);pump(120)
            click(window,c('stopLessonNarrationPreview'))
            for name in ('lessonNarrationVoiceChoice','lessonNarrationVariantChoice','previewLessonNarrationVoice'):
                target=c(name);p=target.mapToScene(QPointF(0,0))
                assert p.x()>=0 and p.x()+target.width()<=width,(name,width)
            assert window.grabWindow().save(str(reports/f'voice-variants-{width}.png'))
        bridge.stopSpeech()

        from app.template_builtin import create
        template = create(library.directory/'own-template.pptx', 'visual')
        bridge.selectCustomTemplate(QUrl.fromLocalFile(str(template)).toString())
        until(lambda: not bridge.busy)
        assert bridge.customTemplate['name'] == 'own-template.pptx'
        assert len(bridge.customTemplate['pages']) == 4
        bridge.setCustomTemplateRoles(1, 0, 3, 4)
        assert bridge.customTemplate['roles']['intro'] == 0
        window.setProperty('page', 'new')
        pump(100)
        click(window, c('creationUseTemplate'))
        assert c('chooseCustomTemplateButton').isVisible()
        assert c('templateContentPage').property('value') == 3
        bridge.clearCustomTemplate()
        assert not bridge.customTemplate

        window.setProperty("page", "settings")
        pump(100)
        click(window, c('settingsTab6'))
        assert c('receiveGoogleKeysGitHub').isVisible()
        assert c('sendGoogleKeysGitHub').isVisible()
        assert window.grabWindow().save(str(reports / 'private-key-sync.png'))
        click(window, c("settingsTab3"))
        click(window, c("mascotLumi"))
        assert c("mascotPreview").property("character") == "Lumi"
        subject=c('mascotSubject');idx=bridge.conversionSubjects.index('Mĩ thuật')
        subject.setProperty('currentIndex',idx)
        QMetaObject.invokeMethod(subject,'activated',Qt.DirectConnection,Q_ARG(int,idx))
        assert c('mascotSubjectMode').property('currentIndex')==1
        click(window,c('saveMascot'));pump(100)
        assert bridge.mascotSubject=='Mĩ thuật'
        assert c('mascotCompanion').property('subject')=='Mĩ thuật'
        assert not bridge.companionCanRead
        assert not c('companionReadEnglish').isEnabled() and not c('companionReadBoth').isEnabled()
        assert window.findChild(QObject, "mascotSlideLumi") is None
        assert window.findChild(QObject, "mascotSlideMilo") is None
        assert window.grabWindow().save(str(reports / "mascot-settings.png"))
        engine.loadData(
            b"""import QtQuick
import QtQuick.Window
Window { visible: true; width: 700; height: 470; color: "#f3f7fc"
    Mascot { objectName: "miloMotion"; x: 20; y: 20; width: 310; height: 410; character: "Milo"; subject: "V\\u1eadt l\\u00fd" }
    Mascot { objectName: "lumiMotion"; x: 370; y: 20; width: 310; height: 410; character: "Lumi"; subject: "To\\u00e1n h\\u1ecdc" }
}""",
            QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/ResultMotionSmoke.qml")),
        )
        motion = engine.rootObjects()[-1]
        milo, lumi = control(motion, "miloMotion"), control(motion, "lumiMotion")
        until(lambda: milo.property("motionActive") and lumi.property("motionActive"))
        until(lambda: milo.property("blinkProgress") > 0.2, timeout=8)
        for mascot in (milo, lumi):
            mascot.setProperty("reducedMotion", True)
            assert not mascot.property("motionActive") and mascot.property("blinkProgress") == 0
            assert mascot.property("scale") == 1
        pump(150)
        assert motion.grabWindow().save(str(reports / "mascots-open.png"))
        for mascot in (milo, lumi):
            mascot.setProperty("blinkProgress", 1)
        pump(150)
        assert motion.grabWindow().save(str(reports / "mascots-blink.png"))
        milo.setProperty("subject", "Tin học")
        assert milo.findChild(QObject, "mascotSubjectBadge").property("visible")
        for mascot in (milo, lumi):
            mascot.setProperty("blinkProgress", 0)
            mascot.setProperty("gestureFrame", 1)
        pump(180)
        assert motion.grabWindow().save(str(reports / "mascots-wave.png"))
        milo.setProperty("reducedMotion", False)
        milo.setProperty("y", 800)
        until(lambda: not milo.property("motionActive"))
        motion.close()
        assert not warnings, warnings
        report = {
            "status": "passed",
            "qml_warnings": warnings,
            "responsive_result_sizes": responsive,
            "status_and_actions_same_row":True,
            "slide_navigation_centered":True,
            "narration_clock_ticks_without_chunk_progress":True,
            "narration_clock_freezes_after_cancel":True,
            "custom_template_roles": "passed",
            "queued_pages_rendered": rendered,
            "ui_ticks_while_previewing": len(beats),
            "narration_starts_with_layout_warning": True,
            "preview_before_voice": True,
            "full_teacher_name": True,
            "download_voice_checkbox": True,
            'six_local_voice_previews_without_global_save': True,
            'saved_variants_route_selected_voice': True,
            'saved_subject_updates_floating_mascot': True,
            'companion_read_disabled_without_slideshow': True,
            'completed_voice_creation_locked_but_selectable': True,
            'new_voice_or_speed_unlocks_creation': True,
            'explicit_next_step_opens_teaching_preparation': True,
            'completion_does_not_auto_navigate_or_approve': True,
            "mascot_checks": [
                "card selection",
                "blink timer",
                "wave pose",
                "no zoom",
                "subject props",
                "reduced motion",
                "offscreen pause",
                "sample buttons removed",
            ],
        }
        (reports / "result-qt.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(json.dumps(report, ensure_ascii=False))
        window.close()
        bridge.slidePreview.shutdown()
        bridge._catalog.shutdown()
        bridge.browserAI.shutdown()
        bridge.english_reader.shutdown()
        bridge.narration_preparation.shutdown()
        bridge.classroom.disconnect()
        engine.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        app.processEvents()
        library.close()


if __name__ == "__main__":
    main()
