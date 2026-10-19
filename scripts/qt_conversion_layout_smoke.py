"""Exercise the restored upload panel and reachable conversion controls in Qt."""

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlExpression
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import click, control, pump, until

from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle("Basic")
    out = Path(os.environ.get('BILICLASS_CONVERSION_REPORT', 'reports/rc96/conversion'))
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="bcl-compact-conversion-", ignore_cleanup_errors=True) as temp:
        temp = Path(temp)
        library = Library(temp / "data")
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(str(v) for v in values))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(RESOURCE_ROOT / "qml/Main.qml")
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]

        def c(name):
            return control(window, name)

        measurements = []
        def verify_layout(entry, scroll):
            upload = c('creationSourceArea')
            assert upload.property('compact') and not upload.property('condensed')
            assert upload.height() >= 140, 'Keep the illustrated upload panel readable'
            icon = c('chooseInputDocumentButtonIllustration')
            button = c('chooseInputDocumentButton')
            assert icon.width() >= 139 and icon.height() >= 109
            assert button.mapToScene(QPointF(0, 0)).x() > icon.mapToScene(QPointF(icon.width(), 0)).x()
            entry['requires_scroll'] = entry['content'] > entry['viewport'] + .5
            if entry['requires_scroll']:
                # The restored panel retains its original size. Extra fields
                # remain reachable by scrolling, without restyling other choices.
                flick = scroll.property('contentItem')
                flick.setProperty('contentY', entry['content'] - entry['viewport'])
                pump(80)
            else:
                assert not c('creationScrollBar').isVisible(), entry
            bottom = c('creationConfiguration').mapToScene(QPointF(0, c('creationConfiguration').height())).y()
            viewport_bottom = scroll.mapToScene(QPointF(0, scroll.property('availableHeight'))).y()
            assert bottom <= viewport_bottom + 1, entry
            summary = c('conversionConfigurationSummary')
            assert summary.isVisible()
            container = c('creationConfiguration')
            local = summary.mapToItem(container, QPointF())
            assert local.x() >= 10 and local.y() > 0
            assert local.x() + summary.width() <= container.width() - 10
            assert local.y() + summary.height() <= container.height() - 10
            assert summary.mapToScene(QPointF(0, summary.height())).y() <= viewport_bottom + 1, entry
            footer = c('creationFooter')
            convert = c('createLessonButton')
            right = convert.mapToItem(footer, QPointF(convert.width(), 0)).x()
            entry['footer_right_gap'] = footer.width() - right
            assert abs(entry['footer_right_gap'] - 28) <= 1, entry

        def choice_geometry():
            # Switching design/subject may add fields below the cards, but must
            # not change card sizing, outer margins or the preview's placement.
            window.grabWindow()
            pump(30)
            names = ['formatChoice_' + row['id'] for row in bridge.conversionFormats]
            names += ['creationKeepOriginal', 'creationUseTemplate']
            dimensions = []
            previews = []
            for name in names:
                card = c(name + 'Card')
                dimensions.extend((card.width(), card.height()))
                if name.startswith('formatChoice_'):
                    preview = c(name + 'Preview')
                    local = preview.mapToItem(card, QPointF())
                    dimensions.extend((local.x(), local.y(), preview.width(), preview.height()))
                    previews.append(preview.property('text'))
            return dimensions, previews

        def verify_same_choices(expected):
            actual = choice_geometry()
            assert actual[1] == expected[1], (expected, actual)
            assert all(abs(a - b) <= 1 for a, b in zip(actual[0], expected[0], strict=True)), (expected, actual)
        try:
            window.setProperty("page", "new")
            path = temp / "Bài giảng PowerPoint có tên dài để kiểm tra bố cục.pptx"
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(path)
            for width, height in ((1080, 700), (1366, 768), (1366, 850), (1920, 1080)):
                window.resize(width, height)
                pump(200)
                for selected in (False, True):
                    if selected:
                        QMetaObject.invokeMethod(
                            window,
                            "acceptInputDocuments",
                            Qt.DirectConnection,
                            Q_ARG("QVariant", [path.as_uri()]),
                        )
                        until(lambda: not bridge.inputDocuments.busy)
                        c("lessonSubject").setProperty(
                            "currentIndex", bridge.conversionSubjects.index("Tin học")
                        )
                    else:
                        QMetaObject.invokeMethod(window, "removeInputDocument", Qt.DirectConnection)
                    # A restored task exposes the left-hand resume action. It
                    # must not shift the primary conversion action inward.
                    bridge._chatgpt_request = {'folder': str(temp / 'disposable-request')} if selected else {}
                    bridge.changed.emit()
                    geometry = None
                    for template in (False, True, False):
                        click(window, c("creationUseTemplate" if template else "creationKeepOriginal"))
                        pump(130)
                        for name in ('formatChoice_parallel_columns', 'creationUseTemplate',
                                     'creationPresetSelect_standard', 'creationPresetSelect_visual', 'creationPresetSelect_practice'):
                            radio = c(name)
                            indicator = radio.property('indicator')
                            assert indicator.width() == indicator.height() == 18
                            expected = '#0869f9' if radio.property('checked') else '#8c9eb8'
                            expression = QQmlExpression(engine.rootContext(), indicator, 'border.color')
                            color = expression.evaluate()
                            color = color[0] if isinstance(color, tuple) else color
                            assert color.name() == expected
                        scroll = c("creationScroll")
                        flick = scroll.property("contentItem")
                        flick.setProperty("contentY", 0)
                        pump(50)
                        if geometry is None:
                            geometry = choice_geometry()
                        else:
                            verify_same_choices(geometry)
                        entry = {
                            "width": width,
                            "height": height,
                            "selected": selected,
                            "template": template,
                            "content": scroll.property("contentHeight"),
                            "viewport": scroll.property("availableHeight"),
                            "upload": c("creationSourceArea").height(),
                            "config": c("creationConfiguration").height(),
                        }
                        verify_layout(entry, scroll)
                        if template:
                            assert c('creationPreset').property('count') == 3
                            for preset in ('standard', 'visual', 'practice'):
                                preview = c('creationPresetPreview_' + preset)
                                assert QQmlExpression(engine.rootContext(), preview, 'status === 1').evaluate()[0]
                                assert preview.property('paintedWidth') > 100
                            assert window.findChild(type(c('creationPreset')), 'creationTemplateOptionsButton') is None
                        flick.setProperty('contentY', 0)
                        pump(50)
                        measurements.append(entry)
                        window.grabWindow().save(
                            str(out / f"{width}-{height}-{int(selected)}-{int(template)}.png")
                        )
                        c('lessonSubject').setProperty('currentIndex', bridge.conversionSubjects.index('Khác'))
                        c('creationCustomSubject').setProperty('text', 'Môn học riêng')
                        pump(150)
                        custom = {**entry, 'other_subject': True,
                                  'content': scroll.property('contentHeight'), 'viewport': scroll.property('availableHeight')}
                        verify_layout(custom, scroll)
                        verify_same_choices(geometry)
                        measurements.append(custom)
                        c('lessonSubject').setProperty('currentIndex', bridge.conversionSubjects.index('Tin học'))
                        pump(50)
                QMetaObject.invokeMethod(window, 'removeInputDocument', Qt.DirectConnection)
                window.setProperty('page', 'new')
                pump(150)
                creation_height = c('creationSourceArea').height()
                window.setProperty('page', 'voice')
                pump(150)
                voice = c('voiceSourceArea')
                assert abs(voice.height() - creation_height) <= 1, (voice.height(), creation_height)
                assert voice.property('compact') and not voice.property('condensed')
                footer = c('creationFooter')
                generate = c('createLessonNarrationButton')
                right = generate.mapToItem(footer, QPointF(generate.width(), 0)).x()
                assert abs(footer.width() - right - 28) <= 1
                window.grabWindow().save(str(out / f'voice-{width}-{height}.png'))
                window.setProperty('page', 'new')
            assert not warnings, warnings
            result = {"status": "passed", "measurements": measurements, "qml_warnings": warnings}
            (out / "result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
            print(json.dumps(result), flush=True)
        finally:
            window.hide()
            for worker in (
                bridge.inputDocuments,
                bridge.libraryThumbnails,
                bridge.slidePreview,
                bridge._catalog,
                bridge.browserAI,
                bridge.english_reader,
                bridge.narration_preparation,
                bridge.classroom,
                bridge.voice_controls,
            ):
                worker.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()
            logging.shutdown()


if __name__ == "__main__":
    main()
