"""Exercise document drop/chooser, prompt preview and automatic PPTX conversion.

Web response, voices and the final classroom session use fixtures. PowerPoint
preview rendering is real; there are no uploads or teacher-library changes.
"""
import hashlib
import json
import sys
import tempfile
import time
import traceback
from pathlib import Path
from threading import Event

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import (
    QCoreApplication,
    QEvent,
    QMetaObject,
    QMimeData,
    QObject,
    QPoint,
    QPointF,
    Qt,
    QTimer,
    QUrl,
)
from PySide6.QtGui import (
    QDesktopServices,
    QDragEnterEvent,
    QDragLeaveEvent,
    QDropEvent,
    QFont,
    QFontDatabase,
    QGuiApplication,
)
from PySide6.QtQml import QQmlApplicationEngine, QQmlEngine, QQmlExpression
from PySide6.QtQuick import QQuickWindow
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from app import browser_audio, browser_automation
from app.chatgpt_handoff import load_request
from app.document_limits import MAX_POWERPOINT_BYTES
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge

app = QGuiApplication([])
QQuickStyle.setStyle("Basic")
for font in (RESOURCE_ROOT / "assets").glob("*.ttf"):
    QFontDatabase.addApplicationFontFromData(font.read_bytes())
app.setFont(QFont("Be Vietnam Pro", 10))
reports = Path("reports/chatgpt-handoff")
reports.mkdir(parents=True, exist_ok=True)
Path(".runtime/chatgpt-smoke").mkdir(parents=True, exist_ok=True)
workspace = Path(tempfile.mkdtemp(prefix="flow-", dir=".runtime/chatgpt-smoke")).resolve()
library = Library(workspace / "library")
source = workspace / "Bài gốc.pptx"
returned = workspace / "Bài ChatGPT trả về.pptx"
deck = Presentation()
for i in range(2):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2)).text = "Bài giảng của thầy cô"
    slide.notes_slide.notes_text_frame.text = 'VI: Diện tích bằng 12 cm².\nEN: The area is 12 cm².\nQUIZ: ' + json.dumps({
        "kind": "single", "vi": "Diện tích bằng bao nhiêu?", "en": "What is the area?",
        "options": [{"vi": "12 cm²", "en": "12 cm²"}, {"vi": "6 cm²", "en": "6 cm²"}],
        "correct": "A", "rationale_vi": "Theo nội dung slide.", "rationale_en": "As shown on the slide."}, ensure_ascii=False) + "\nCHECK: Kiểm tra hoạt ảnh trước khi dạy."
deck.save(source)
deck.save(returned)
digest = hashlib.sha256(returned.read_bytes()).hexdigest()
bridge = Bridge(library)
bridge.browserAI.saveOptions(False, False)  # New UI must override obsolete manual/audio options.
account = bridge.browserAI.store.add("UI fixture")
bridge.browserAI.store.observe(account["id"], "free", identity={"name": "Cô giáo thử", "email": "fixture@example.test"})
release_web = Event()
sent, narration = [], []


def fake_convert(account, root, folder, cancel, progress, **options):
    sent.append(load_request(folder))
    while not release_web.wait(.05):
        if cancel.is_set():
            raise browser_automation.BrowserProblem("cancelled", "Đã hủy.")
    options["observed"](account["id"], "free", "Fixture")
    target = Path(folder) / "bai-giang-song-ngu.pptx"
    target.write_bytes(returned.read_bytes())
    browser_automation.write_record(folder, {"account_id": account["id"], "state": "completed", "followups": 0,
        "url": "https://chatgpt.com/c/fixture", "sha256": digest})
    return {"path": str(target), "url": "https://chatgpt.com/c/fixture"}


def fake_narration(inspection, voices, directory, cancel, progress):
    narration.extend(inspection["profile"]["units"])
    return {"complete": 4, "total": 4, "failed": 0, "skipped": 0}


browser_automation.convert = fake_convert
browser_audio.prepare_narration = fake_narration
engine = QQmlApplicationEngine()
warnings, failures, stages, opened = [], [], [], []
engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
engine.rootContext().setContextProperty("bridge", bridge)
engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
assert engine.rootObjects(), warnings
window = engine.rootObjects()[0]
window.setProperty("page", "new")
QDesktopServices.openUrl = lambda url: opened.append(url.toString()) or True
deadline = time.monotonic() + 120
preview_prompt = ""
teacher_notes = "Giữ từng bước giải và công thức.\nDùng thuật ngữ phù hợp học sinh lớp 10; không thêm bài tập ngoài nguồn."


def safe(action):
    def run():
        try:
            action()
        except Exception:
            failures.append(traceback.format_exc())
            print(failures[-1], flush=True)
            app.exit(1)
    return run


def click(name):
    target = window.findChild(QObject, name)
    if target is None:
        # Repeater delegates can belong to a QML context outside the QObject tree.
        queue = [window.contentItem()]
        while queue:
            item = queue.pop()
            if item.objectName() == name:
                target = item
                break
            queue.extend(item.childItems())
    assert target and target.property("enabled"), name
    assert QMetaObject.invokeMethod(target, "clicked", Qt.DirectConnection)


def visual_value(name, expression):
    # Return scalars only; do not transfer ownership of QML Repeater delegates.
    code = '''(function() {
        function find(node) {
            if (node.objectName === NAME) return node;
            for (var i = 0; i < node.children.length; i++) {
                var found = find(node.children[i]); if (found) return found;
            }
            return null;
        }
        var target = find(contentItem);
        if (!target) throw new Error("Missing " + NAME);
        return EXPRESSION;
    })()'''.replace("NAME", json.dumps(name)).replace("EXPRESSION", expression)
    query = QQmlExpression(QQmlEngine.contextForObject(window), window, code)
    value = query.evaluate()[0]
    assert not query.hasError(), query.error().toString()
    return value


def mouse_click(name):
    assert visual_value(name, "target.enabled && target.visible"), name
    point = QPoint(int(visual_value(name, "target.mapToItem(null, target.width / 2, target.height / 2).x")),
                   int(visual_value(name, "target.mapToItem(null, target.width / 2, target.height / 2).y")))
    QTest.mouseClick(window, Qt.LeftButton, Qt.NoModifier, point)


def drop_file(paths):
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path)) for path in paths])
    item = window.findChild(QObject, "creationSourceArea")
    point = item.mapToScene(QPointF(30, 70))
    entered = QDragEnterEvent(point.toPoint(), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    QCoreApplication.sendEvent(window, entered)
    if not entered.isAccepted():
        QCoreApplication.sendEvent(window, QDragLeaveEvent())
        return False
    dropped = QDropEvent(point, Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    QCoreApplication.sendEvent(window, dropped)
    accepted = dropped.isAccepted()
    QCoreApplication.sendEvent(window, QDragLeaveEvent())
    return accepted


def wait_idle(action):
    assert time.monotonic() < deadline, "Flow timed out"
    if bridge.busy:
        QTimer.singleShot(100, safe(lambda: wait_idle(action)))
    else:
        action()


def start():
    assert not window.findChild(QObject, "createLessonButton").property("enabled")
    assert visual_value("creationKeepOriginal", "target.enabled && target.checked")
    assert not window.findChild(QObject, "creationEducation").property("editable")
    assert not window.findChild(QObject, "creationGrade").property("editable")
    check_education(0)


def check_education(index):
    window.findChild(QObject, "creationEducation").setProperty("currentIndex", [2, 1, 0][index])
    QTimer.singleShot(100, safe(lambda: education_checked(index)))


def education_checked(index):
    grade = window.findChild(QObject, "creationGrade")
    expected = [list(map(str, range(1, 6))), list(map(str, range(6, 10))), ["10", "11", "12"]][index]
    assert grade.property("count") == len(expected)
    assert grade.property("currentText") == expected[0]
    for position, text in enumerate(expected):
        grade.setProperty("currentIndex", position)
        assert grade.property("currentText") == text
    if index < 2:
        check_education(index + 1)
    else:
        grade.setProperty("currentIndex", 0)
        stages.append("Education filters exact grade ranges; stale grade resets; source design selectable before importing")
        choose_document()


def choose_document():
    click("creationUseTemplate")
    assert drop_file([source]), "Windows-style local document drop was rejected"
    assert window.property("selectedFileName") == source.name
    assert window.findChild(QObject, "lessonTitle").property("text") == source.stem
    assert window.findChild(QObject, "creationWorkflow").property("currentIndex") == 1, "Keep explicit template choice"
    click("creationKeepOriginal")
    assert not drop_file([source, source]), "Multiple documents must not replace the accepted source"
    assert window.property("selectedFileName") == source.name
    too_large = workspace / "Bài vượt dung lượng.pptx"
    with too_large.open("wb") as stream:
        stream.truncate(MAX_POWERPOINT_BYTES + 1)
    assert not drop_file([too_large]), "Oversized source must not replace the accepted document"
    assert window.property("selectedFileName") == source.name
    assert window.findChild(QObject, "creationDocumentError").property("visible")
    assert "200 MB" in window.property("inputDocumentError")
    # File chooser accepted event routes through the same validation path.
    picked = workspace / "Chọn ảnh #1.png"
    picked.write_bytes(b"image fixture")
    dialog = window.findChild(QObject, "creationFileDialog")
    dialog.setProperty("selectedFile", QUrl.fromLocalFile(str(picked)))
    assert QMetaObject.invokeMethod(dialog, "accepted", Qt.DirectConnection)
    QTimer.singleShot(100, safe(lambda: chooser_checked(picked)))


def chooser_checked(picked):
    assert window.property("selectedFileName") == picked.name
    assert not visual_value("creationKeepOriginal", "target.enabled")
    assert window.findChild(QObject, "creationWorkflow").property("currentIndex") == 1
    click("removeInputDocumentButton")
    assert not window.property("selectedFile") and not window.property("selectedFileName")
    assert not window.findChild(QObject, "lessonTitle").property("text"), "Remove filename-derived title with its source"
    assert drop_file([source])
    assert window.findChild(QObject, "lessonTitle").property("text") == source.stem
    assert not bridge.error and source.name in bridge.message, "Valid selection must clear stale input error"
    window.findChild(QObject, "lessonTitle").setProperty("text", "Bài giảng với ChatGPT")
    window.findChild(QObject, "lessonSubject").setProperty("text", "Toán")
    window.findChild(QObject, "creationTeacherNotes").setProperty("text", teacher_notes)
    assert not window.property("inputDocumentError")
    QTimer.singleShot(150, safe(formats))


def formats():
    choice = window.findChild(QObject, "creationFormat")
    assert choice.property("visible") and choice.property("count") == 4
    for index, spec in enumerate(bridge.conversionFormats):
        mouse_click("formatChoice_" + spec["id"])
        assert choice.property("currentIndex") == index
        assert sum(visual_value("formatChoice_" + other["id"], "target.checked") for other in bridge.conversionFormats) == 1
        assert window.findChild(QObject, "creationFormatDescription").property("text") == spec["detail"]
    choice.setProperty("currentIndex", 0)
    assert window.findChild(QObject, "creationLevel") is None
    assert window.findChild(QObject, "creationLayout") is None
    assert window.findChild(QObject, "creationWorkflow").property("visible")
    assert window.findChild(QObject, "creationWorkflow").property("currentIndex") == 0
    assert window.findChild(QObject, "conversionProvider") is None
    assert window.findChild(QObject, "creationManualChatGPT") is None
    assert window.findChild(QObject, "creationAudio") is None
    assert not window.findChild(QObject, "creationTemplatesButton").property("visible")
    assert QQuickWindow.grabWindow(window).save(str(reports / "new-native.png"))
    mouse_click("formatChoice_english_onlyPreview")
    QTimer.singleShot(200, safe(example))


def example(index=3):
    dialog = window.findChild(QObject, "conversionExampleDialog")
    assert dialog.property("visible") and dialog.property("formatIndex") == index
    key = bridge.conversionFormats[index]["id"]
    assert window.findChild(QObject, "conversionExampleSlide").property("conversionFormat") == key
    assert window.findChild(QObject, "conversionExampleSlide").property("englishOnly") == (index == 3)
    assert window.findChild(QObject, "creationFormat").property("currentIndex") == 0, "Preview is separate from selection"
    assert QQuickWindow.grabWindow(window).save(str(reports / ("example-" + key + ".png")))
    if index != 2:
        assert QMetaObject.invokeMethod(dialog, "close", Qt.DirectConnection)
        following = (index + 1) % 4
        mouse_click("formatChoice_" + bridge.conversionFormats[following]["id"] + "Preview")
        QTimer.singleShot(150, safe(lambda: example(following)))
        return
    click("selectConversionExample")
    assert window.findChild(QObject, "creationFormat").property("currentIndex") == 2
    window.findChild(QObject, "creationFormat").setProperty("currentIndex", 0)
    stages.append("Native file drop and chooser; PPTX/non-PPTX defaults; four exclusive selections and independent layout preview")
    window.findChild(QObject, "creationWorkflow").setProperty("currentIndex", 1)
    QTimer.singleShot(150, safe(gallery))


def gallery():
    assert window.findChild(QObject, "creationTemplatesButton").property("visible")
    assert QQuickWindow.grabWindow(window).save(str(reports / "new-template.png"))
    click("creationTemplatesButton")
    QTimer.singleShot(300, safe(gallery_visible))


def gallery_visible():
    dialog = window.findChild(QObject, "templateGallery")
    assert dialog.property("visible")
    plan = window.findChild(QObject, "templateSampleSlide").property("plan")
    assert plan["image"] and not plan["overflow"]
    assert QQuickWindow.grabWindow(window).save(str(reports / "gallery-standard.png"))
    dialog.setProperty("preset", "visual")
    QTimer.singleShot(300, safe(gallery_visual))


def gallery_visual():
    click("templateThumbnail_visual")
    QTimer.singleShot(150, safe(gallery_visual_slide))


def gallery_visual_slide():
    dialog = window.findChild(QObject, "templateGallery")
    plan = window.findChild(QObject, "templateSampleSlide").property("plan")
    assert plan["kind"] == "visual" and plan["image"]
    assert QQuickWindow.grabWindow(window).save(str(reports / "gallery-visual.png"))
    assert QMetaObject.invokeMethod(dialog, "close", Qt.DirectConnection)
    # Keep the configured split view when switching back to the original deck.
    window.findChild(QObject, "creationWorkflow").setProperty("currentIndex", 0)
    notes = window.findChild(QObject, "creationTeacherNotes")
    notes.setProperty("text", "x" * 6001)
    assert not window.findChild(QObject, "createLessonButton").property("enabled"), "Do not truncate oversized notes"
    notes.setProperty("text", teacher_notes)
    click("creationPromptToggle")
    prompt = window.findChild(QObject, "creationPromptText")
    assert prompt.property("visible") and prompt.property("readOnly")
    global preview_prompt
    preview_prompt = prompt.property("text")
    assert "Bài giảng với ChatGPT" in preview_prompt and source.name in preview_prompt
    assert "Giữ theme" in preview_prompt and "VI:" in preview_prompt
    assert preview_prompt.count(teacher_notes) == 1 and "LƯU Ý BỔ SUNG CỦA GIÁO VIÊN" in preview_prompt
    scroll = window.findChild(QObject, "creationScroll").property("contentItem")
    scroll.setProperty("contentY", max(0, scroll.property("contentY") + notes.mapToScene(QPointF(0, notes.height())).y()
                                      - scroll.mapToScene(QPointF(0, scroll.height())).y() + 20))
    QTimer.singleShot(150, safe(notes_shown))


def notes_shown():
    assert QQuickWindow.grabWindow(window).save(str(reports / "teacher-notes.png"))
    QTimer.singleShot(150, safe(lambda: scroll_to_prompt(prompt_shown)))


def scroll_to_prompt(action):
    scroll = window.findChild(QObject, "creationScroll").property("contentItem")
    scroll.setProperty("contentY", max(0, scroll.property("contentHeight") - scroll.property("height")))
    QTimer.singleShot(150, safe(action))


def prompt_shown():
    assert QQuickWindow.grabWindow(window).save(str(reports / "prompt-visible.png"))
    window.setWidth(1080)
    window.setHeight(700)
    QTimer.singleShot(150, safe(lambda: scroll_to_prompt(compact)))


def compact():
    footer = window.findChild(QObject, "createLessonButton")
    point = footer.mapToScene(QPointF(0, 0))
    assert 0 < point.y() < window.height() - footer.height(), "Conversion button must remain on screen"
    assert window.findChild(QObject, "creationEducation").property("width") >= 115
    assert window.findChild(QObject, "creationGrade").property("width") >= 80
    assert QQuickWindow.grabWindow(window).save(str(reports / "compact-prompt.png"))
    click("creationPromptToggle")
    assert not window.findChild(QObject, "creationPromptText").property("visible")
    window.setWidth(1366)
    window.setHeight(850)
    stages.append("Live prompt show/hide with read-only script; conversion button stays visible at 1080x700")
    click("createLessonButton")
    QTimer.singleShot(100, safe(prepared))


def prepared():
    assert time.monotonic() < deadline, "Automatic request was not started"
    if not sent:
        QTimer.singleShot(100, safe(prepared))
        return
    assert not bridge.error, bridge.message
    assert window.property("page") == "chatgpt"
    request = bridge.chatgptRequest
    assert request["config"]["conversion_format"] == "parallel_columns"
    assert request["config"]["layout"] == "split_view" and request["config"]["level"] == 3
    assert request["config"]["education_level"] == "THPT" and request["config"]["grade"] == "10"
    assert request["config"]["teacher_notes"] == teacher_notes
    assert "Hai cột" in request["prompt"] and "Giữ theme" in request["prompt"]
    assert request["prompt"] == preview_prompt == sent[0]["prompt"]
    assert bridge.browserAI.automatic and bridge.browserAI.audio
    assert opened == [], "Automatic conversion must not launch a manual web handoff"
    assert Path(request["bundle"]).is_file()
    assert QQuickWindow.grabWindow(window).save(str(reports / "browser-handoff.png"))
    stages.append("One conversion click overrides old manual/audio settings; sent prompt equals preview; web send/receive fixture runs automatically")
    release_web.set()
    QTimer.singleShot(100, safe(lambda: wait_idle(received)))


def received():
    assert not bridge.error, bridge.message
    assert window.property("page") == "result"
    result = bridge.quickResult
    assert result["external"] and result["draft"] and result["image"]
    assert narration and all(unit["vi"] and unit["en"] for unit in narration)
    assert result["audio"]["complete"] == 4
    assert hashlib.sha256(Path(result["path"]).read_bytes()).hexdigest() == digest
    assert not any(s["approved"] for s in bridge.lesson["segments"])
    assert all(s["vi"] and s["en"] for s in bridge.lesson["segments"])
    assert bridge.lesson["conversion_format"] == "parallel_columns"
    assert len(bridge.lesson["questions"]) == 2 and not any(q["approved"] for q in bridge.lesson["questions"])
    assert all("QUIZ:" not in s["en"] for s in bridge.lesson["segments"])
    assert window.findChild(QObject, "returnedDeckReviewNotes").property("visible")
    assert "Kiểm tra hoạt ảnh" in window.findChild(QObject, "returnedDeckReviewNotes").property("text")
    assert QQuickWindow.grabWindow(window).save(str(reports / "returned-pptx.png"))
    click("useConvertedLessonButton")
    assert window.findChild(QObject, "wholeLessonReview").property("visible")
    sessions = []
    bridge._start_powerpoint_session = lambda *args: sessions.append(args)
    click("confirmWholeLessonButton")
    assert len(sessions) == 1 and all(s["approved"] for s in bridge.lesson["segments"])
    assert not bridge.quickResult["draft"]
    assert not any(q["approved"] for q in bridge.lesson["questions"])
    stages.append("Receive byte-identical PPTX -> real Office preview -> whole-deck confirmation -> mascot mapping")
    lesson_id = bridge.lesson["id"]
    bridge.openLesson(lesson_id)
    assert bridge.quickResult["external"] and not bridge.quickResult["draft"]
    assert bridge.quickResult["segment_map"]
    stages.append("Reopen saved presentation retains returned file, confirmation and slide mapping")
    assert not warnings, warnings
    app.exit(0)


QTimer.singleShot(550, safe(start))
code = app.exec()
release_web.set()
if bridge.worker:
    bridge.cancel_event.set()
    bridge.worker.wait()
report = {"status": "passed" if code == 0 and not warnings and not failures else "failed",
          "stages": stages, "warnings": warnings, "failures": failures,
          "scope": "Qt/native input and real Office preview; web/voices use fixtures, no uploads"}
(reports / "qt-flow.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
window.close()
engine.deleteLater()
app.processEvents()
QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
library.close()
print(json.dumps(report, ensure_ascii=False), flush=True)
raise SystemExit(code)
