"""Open the real Windows file picker and verify its accepted document in QML.

Uses an isolated library; optional source files are read only, never uploaded.
"""

import argparse
import hashlib
import json
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path
from threading import Thread

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import win32con
import win32gui
import win32process
from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEvent, QMetaObject, QObject, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlComponent
from PySide6.QtQuickControls2 import QQuickStyle

from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, help="Also select every PPTX in this directory (read only)")
    parser.add_argument('--theme', choices=('light', 'dark'), default='light')
    args = parser.parse_args()
    app = QGuiApplication([])
    QQuickStyle.setStyle("Basic")
    for font in (RESOURCE_ROOT / "assets").glob("*.ttf"):
        QFontDatabase.addApplicationFontFromData(font.read_bytes())
    app.setFont(QFont("Be Vietnam Pro", 10))
    reports = Path(os.environ.get('BILICLASS_PICKER_REPORT', 'reports/conversion-picker'))
    reports.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="biliclass-picker-", ignore_cleanup_errors=True) as directory:
        workspace = Path(directory)
        source = workspace / "Giáo án #1 – chọn từ máy.PPTX"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(source)
        sources = [source]
        if args.source_dir:
            assert args.source_dir.is_dir(), "Source directory missing"
            sources.extend(sorted(path.resolve() for path in args.source_dir.rglob("*")
                                  if path.is_file() and path.suffix.casefold() == ".pptx"))
            assert len(sources) > 1, "No PPTX in source directory"
        hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in sources}
        library = Library(workspace / "library")
        library.set_setting('application_theme', args.theme)
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings, failures, native, checked = [], [], [], []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty("bridge", bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        window.setProperty("page", "new")
        deadline = time.monotonic() + 30 + 10 * len(sources)
        index = 0
        cancel_checked = False
        save_checked = []
        save_component = QQmlComponent(engine, QUrl.fromLocalFile(str(RESOURCE_ROOT/'qml/SuggestedSaveDialog.qml')))
        save_dialog = save_component.create()
        assert save_dialog is not None, save_component.errors()
        save_dialog.setParent(window)
        save_dialog.setProperty('parentWindow', window)
        save_target = workspace/'Bài gốc #1 (Song ngữ VI-EN).pptx'

        def choose_in_native_dialog(path):
            try:
                while time.monotonic() < deadline:
                    dialogs = []
                    def top(hwnd, _):
                        if (win32process.GetWindowThreadProcessId(hwnd)[1] == os.getpid()
                                and win32gui.GetClassName(hwnd) == "#32770"
                                and win32gui.IsWindowVisible(hwnd)):
                            dialogs.append(hwnd)
                    win32gui.EnumWindows(top, None)
                    if not dialogs:
                        time.sleep(.05)
                        continue
                    dialog = dialogs[0]
                    native.append(dialog)
                    if path is None:
                        win32gui.PostMessage(dialog, win32con.WM_COMMAND, win32con.IDCANCEL, 0)
                        return
                    edits = []
                    def child(hwnd, _):
                        if not win32gui.IsWindowVisible(hwnd):
                            return
                        kind = win32gui.GetClassName(hwnd)
                        parent = win32gui.GetParent(hwnd)
                        identifier = win32gui.GetDlgCtrlID(hwnd)
                        if kind == "Edit" and win32gui.IsWindowEnabled(hwnd):
                            edits.append((hwnd, identifier, win32gui.GetDlgCtrlID(parent)))
                    win32gui.EnumChildWindows(dialog, child, None)
                    filename = next((hwnd for hwnd, identifier, parent_id in edits
                                     if identifier == 1148 or parent_id == 1148), None)
                    if filename is None:
                        filename = next((hwnd for hwnd, _, _ in edits
                                         if win32gui.GetClassName(win32gui.GetParent(hwnd)) in {"ComboBox", "ComboBoxEx32"}), None)
                    if filename is None:
                        # The shell can show its HWND before its controls are ready.
                        time.sleep(.05)
                        continue
                    if isinstance(path, tuple):
                        assert win32gui.GetWindowText(filename) == path[1].name
                        save_checked.append(True)
                        win32gui.PostMessage(dialog, win32con.WM_COMMAND, win32con.IDCANCEL, 0)
                        return
                    win32gui.SendMessage(filename, win32con.WM_SETTEXT, 0, str(path))
                    win32gui.PostMessage(dialog, win32con.WM_COMMAND, win32con.IDOK, 0)
                    return
                raise AssertionError("Real Windows file picker did not open")
            except Exception:
                failures.append(traceback.format_exc())

        def begin(path):
            Thread(target=choose_in_native_dialog, args=(path,), daemon=True).start()
            button = window.findChild(QObject, "chooseInputDocumentButton")
            assert QMetaObject.invokeMethod(button, "clicked", Qt.DirectConnection)
            QTimer.singleShot(100, lambda: inspect(path))

        def inspect(path):
            nonlocal index, cancel_checked
            if failures or time.monotonic() >= deadline:
                if not failures:
                    failures.append("Native picker did not deliver a document to BiliClass")
                app.exit(1)
            elif (window.property("selectedFileName") == path.name if path else
                  not window.findChild(QObject, "creationFileDialog").property("visible")):
                try:
                    assert native and not bridge.error, bridge.message
                    expected = path or sources[-1]
                    if path is not None and bridge.inputDocuments.preview.get('busy'):
                        QTimer.singleShot(100, lambda: inspect(path))
                        return
                    assert Path(QUrl(window.property("selectedFile")).toLocalFile()) == expected
                    assert window.property("selectedFileName") == expected.name
                    assert window.findChild(QObject, "lessonTitle").property("text") == expected.stem
                    assert window.findChild(QObject, "creationWorkflow").property("currentIndex") == 0
                    if path is not None:
                        assert bridge.inputDocuments.preview.get('image'), bridge.inputDocuments.preview
                        assert window.findChild(QObject, 'creationSourceArea').property('condensed')
                        checked.append({"name": path.name, "bytes": path.stat().st_size})
                        print(f"Native picker accepted {len(checked)}/{len(sources)}", flush=True)
                        assert window.grabWindow().save(str(reports / "accepted-pptx.png"))
                        index += 1
                        if index < len(sources):
                            QTimer.singleShot(150, lambda: begin(sources[index]))
                        else:
                            QTimer.singleShot(150, lambda: begin(None))
                    else:
                        cancel_checked = True
                        assert all(hashlib.sha256(p.read_bytes()).hexdigest() == h for p, h in hashes.items())
                        QTimer.singleShot(150, begin_save)
                except Exception:
                    failures.append(traceback.format_exc())
                    app.exit(1)
            else:
                QTimer.singleShot(100, lambda: inspect(path))

        def begin_save():
            save_dialog.setProperty('currentFolder', QUrl.fromLocalFile(str(library.directory)))
            save_dialog.setProperty('suggestedFile', QUrl.fromLocalFile(str(save_target)))
            Thread(target=choose_in_native_dialog, args=(('save-probe', save_target),), daemon=True).start()
            assert QMetaObject.invokeMethod(save_dialog, 'openSuggested', Qt.DirectConnection)
            assert Path(QUrl(save_dialog.property('currentFolder')).toLocalFile()).resolve() == workspace
            QTimer.singleShot(100, inspect_save)

        def inspect_save():
            if failures or time.monotonic() >= deadline:
                if not failures:
                    failures.append('Native Save dialog did not close after cancel')
                app.exit(1)
            elif save_checked and not save_dialog.property('visible'):
                assert not save_target.exists(), 'The Save probe only cancels'
                app.exit(0)
            else:
                QTimer.singleShot(100, inspect_save)

        QTimer.singleShot(300, lambda: begin(sources[0]))
        code = app.exec()
        report = {"status": "passed" if code == 0 and not warnings and not failures else "failed",
                  "native_picker": bool(native), "theme": args.theme, "files": checked, "cancel_keeps_selection": cancel_checked,
                  "native_save_prefills_source_folder_and_unicode_name": bool(save_checked),
                  "warnings": warnings, "failures": failures,
                  "scope": "Real Windows picker, Unicode PPTX selection/replacement/cancel; source thumbnail rendered on a private copy; unchanged SHA256, no browser/uploads"}
        (reports / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        for hwnd in native:
            if win32gui.IsWindow(hwnd) and win32process.GetWindowThreadProcessId(hwnd)[1] == os.getpid():
                win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                       bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
            worker.shutdown()
        bridge.inputDocuments.shutdown()
        companion = window.findChild(QObject, 'companionWindow')
        companion.setProperty('disposing', True)
        companion.hide()
        window.hide()
        engine.deleteLater()
        app.processEvents()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        library.close()
        print(json.dumps(report, ensure_ascii=False), flush=True)
        return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
