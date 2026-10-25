"""RC168 opt-in routing UI and LanguageApp batch checkpoint checks with fake HTTP only."""

import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path
from threading import Event
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.alternative_translation import TRANSPORTS as ALTERNATIVE_TRANSPORTS
from app.alternative_translation import AlternativeTranslator
from app.branding import configure_application
from app.google_conversion import load
from app.google_translation import KeyStore
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT / 'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/rc168-translation')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-provider-', ignore_cleanup_errors=True) as temporary:
        directory = Path(temporary)
        library = Library(directory / 'data')
        KeyStore(library.directory).add('AIza' + 'x' * 35)
        google_file = library.directory / 'translation_api/keys.dpapi'
        original_google = google_file.read_bytes()
        bridge = Bridge(library)
        bridge.previewConvertedSlide = lambda *_: None
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        window.resize(1366, 850)
        window.show()
        entered, release = Event(), Event()
        calls = []

        def c(name):
            return control(window, name)

        def fixture(self, rows, credential, source, target):
            assert credential['key'] == 'a' * 32 and credential['option'].startswith('https://script.google.com/macros/s/')
            calls.append(len(rows))
            entered.set()
            assert release.wait(25), 'Fixture was not released'
            return [row.replace('Vật có ', 'The object has a ') for row in rows]

        try:
            until(lambda: bridge.googleSettings.ready)
            window.setProperty('page', 'settings')
            c('settingsPage').setProperty('activeTab', 6)
            pump(100)
            assert c('translationProviderGoogle').property('checked')
            # All three choices render and save masked credentials through real Qt controls.
            for name, radio, field, save in [('gemini', 'Gemini', 'geminiApiKey', 'saveGemini'), ('deepl', 'DeepL', 'deeplApiKey', 'saveDeepL')]:
                click(window, c('translationProvider' + radio))
                until(lambda: not bridge.busy and bridge.googleSettings.state['provider'] == name)
                c(field).setProperty('text', 'test_secret_do_not_use_1234567890')
                click(window, c(save))
                until(lambda: not bridge.busy)
                assert bridge.googleSettings.state['activeConfigured']
                assert not c(field).property('text')
                assert 'test_secret_do_not_use_1234567890' not in json.dumps(bridge.googleSettings.state)
                def probe(texts, *_):
                    assert texts == ['Xin chào.']
                    return ['Hello.']
                with patch.dict(ALTERNATIVE_TRANSPORTS, {name: probe}):
                    click(window, c('check' + radio))
                    until(lambda: not bridge.busy)
                    assert 'Đã kết nối' in bridge.message
                pump(100)
                assert window.grabWindow().save(str(report / (name + '-settings.png')))
            click(window, c('translationProviderLanguageApp'))
            until(lambda: not bridge.busy and bridge.googleSettings.state['provider'] == 'languageapp')
            assert c('languageappGuide').property('visible'), 'First-time selection opens setup automatically'
            assert c('translationProviderLanguageApp').property('checked')
            assert not c('saveGoogleApiKeys').isVisible()
            assert c('languageappTranslationCard').isVisible()
            assert not bridge.googleSettings.state['activeConfigured']
            assert not c('languageappSetupNext').property('enabled')
            click(window, c('languageappCopySetup'))
            code = QGuiApplication.clipboard().text()
            assert 'function doPost(e)' in code
            setup_secret = re.search(r"const BILICLASS_SETUP_CODE = '([A-Za-z0-9_-]+)';", code)[1]
            assert len(setup_secret) == 32
            click(window, c('languageappCopySetup'))
            assert QGuiApplication.clipboard().text() == code, 'Copying again must keep the same private code'
            assert not bridge.googleSettings.state['activeConfigured'], 'Copy is not a save'
            assert setup_secret not in json.dumps(bridge.googleSettings.state)
            assert window.grabWindow().save(str(report / 'languageapp-setup-guide.png'))
            click(window, c('languageappSetupNext'))
            assert c('languageappGuide').property('step') == 1
            assert window.grabWindow().save(str(report / 'languageapp-setup-deploy.png'))
            click(window, c('languageappSetupNext'))
            assert c('languageappGuide').property('step') == 2
            c('languageappSetupUrl').setProperty('text', 'https://example.com/wrong')
            click(window, c('languageappSetupNext'))
            until(lambda: not bridge.busy)
            assert c('languageappGuide').property('step') == 2
            assert not bridge.googleSettings.state['activeConfigured']
            c('languageappSetupUrl').setProperty('text', 'https://script.google.com/macros/s/' + 'A' * 40 + '/exec')
            from app.google_translation import TranslationProblem

            def missing_deployment(*_):
                raise TranslationProblem('api_deployment', 'Bản triển khai chưa có doPost. Chọn Phiên bản mới.')
            with patch.dict(ALTERNATIVE_TRANSPORTS, {'languageapp': missing_deployment}):
                click(window, c('languageappSetupNext'))
                until(lambda: not bridge.busy)
                assert c('languageappGuide').property('step') == 2
                assert 'doPost' in c('languageappSetupResult').property('text')
                assert window.grabWindow().save(str(report / 'languageapp-setup-error.png'))

            def setup_probe(texts, credential, *_):
                assert texts == ['Xin chào.'] and credential['key'] == setup_secret
                return ['Hello.']
            with patch.dict(ALTERNATIVE_TRANSPORTS, {'languageapp': setup_probe}):
                click(window, c('languageappSetupNext'))
                until(lambda: not bridge.busy)
                assert c('languageappGuide').property('step') == 3
                assert bridge.googleSettings.state['activeConfigured']
                assert setup_secret not in json.dumps(bridge.googleSettings.state)
                assert window.grabWindow().save(str(report / 'languageapp-setup-complete.png'))
            click(window, c('languageappSetupClose'))
            QGuiApplication.clipboard().clear()
            bridge.googleSettings.refresh()
            until(lambda: not bridge.googleSettings.loading)
            assert bridge.googleSettings.state['activeConfigured'], 'Saved configuration survives a refresh'
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                window.resize(860, 650)
                pump(180)
                old_step = c('languageappGuide').property('step')
                click(window, c('languageappSetup'))
                assert c('languageappGuide').property('step') == (0 if old_step == 3 else old_step), 'Reset prior success; retain an unfinished setup'
                for step in range(3):
                    c('languageappGuide').setProperty('step', step)
                    pump(60)
                    footer = c('languageappSetupNext')
                    rect = footer.mapRectToScene(footer.boundingRect())
                    assert rect.left() >= 0 and rect.right() <= window.width()
                    assert rect.top() >= 0 and rect.bottom() <= window.height(), 'Footer must stay reachable on a small laptop'
                    assert window.grabWindow().save(str(report / f'languageapp-guide-{theme}-860-step{step+1}.png'))
                click(window, c('languageappSetupClose'))
            window.resize(1366, 850)
            bridge.setApplicationTheme('light')
            pump(180)
            c('languageappAccess').setProperty('text', 'a' * 32)
            c('languageappUrl').setProperty('text', 'https://script.google.com/macros/s/' + 'A' * 40 + '/exec')
            click(window, c('saveLanguageApp'))
            until(lambda: not bridge.busy)
            assert bridge.googleSettings.state['activeConfigured']
            assert not c('languageappAccess').property('text')
            assert 'a' * 32 not in json.dumps(bridge.googleSettings.state)
            with patch.dict(ALTERNATIVE_TRANSPORTS, {'languageapp': missing_deployment}):
                click(window, c('checkLanguageApp'))
                until(lambda: not bridge.busy)
                assert 'doPost' in bridge.message
                assert bridge.googleSettings.state['provider'] == 'languageapp'
            assert google_file.read_bytes() == original_google
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                for width, height in ((860, 650), (1080, 700), (1366, 850)):
                    window.resize(width, height)
                    # Bring the selection card back into the settings viewport.
                    parent = c('translationProviderCard').parentItem()
                    while parent:
                        if parent.metaObject().indexOfProperty('contentY') >= 0:
                            parent.setProperty('contentY', 0)
                        parent = parent.parentItem()
                    pump(100)
                    assert window.grabWindow().save(str(report / f'languageapp-{theme}-{width}.png'))
            sources = []
            for number in (12, 13):
                deck = Presentation()
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
                shape.text = f'Vật có vận tốc {number} m/s.'
                path = directory / f'Thí nghiệm {number}.pptx'
                deck.save(path)
                sources.append(path)
            source_hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in sources]
            assert bridge.batchQueue.selectFiles('convert', [path.as_uri() for path in sources])
            with patch.object(AlternativeTranslator, '_send', fixture):
                bridge.batchQueue.startConversion(dict(subject='Vật lý', education='THPT', grade='11',
                    format='parallel_columns', preset='standard', style='template'))
                until(entered.is_set, timeout=60)
                assert bridge.batchQueue.running
                assert bridge.batchQueue.state['options']['translation_provider'] == 'languageapp'
                assert not bridge.googleSettings.chooseProvider('google'), 'Provider switched during work'
                bridge.batchQueue.stop()
                release.set()
                until(lambda: not bridge.busy and not bridge.batchQueue.running, timeout=120)
                assert bridge.batchQueue.state['phase'] == 'paused'
                assert bridge.googleSettings.chooseProvider('google')
                until(lambda: not bridge.busy)
                assert bridge.googleSettings.state['provider'] == 'google'
                assert bridge.batchQueue.state['options']['translation_provider'] == 'languageapp'
                bridge.batchQueue.resume()
                until(lambda: not bridge.busy and not bridge.batchQueue.running, timeout=180)
            state = bridge.batchQueue.state
            assert state['counts']['done'] == 2 and state['complete'], state
            for row in state['rows']:
                assert load(library.directory, row['folder'])['config']['translation_provider'] == 'languageapp'
                assert library.get(row['lesson_id'])['external_deck']['translation_provider'] == 'languageapp'
            assert [hashlib.sha256(path.read_bytes()).hexdigest() for path in sources] == source_hashes
            assert KeyStore(library.directory).public()[0]['used'] == 0
            assert google_file.read_bytes() == original_google
            assert not warnings, warnings
            result = dict(status='passed', api='fixture; no live LanguageApp quality/speed claim', calls=len(calls),
                checks=['Google remains default; original key ciphertext unchanged',
                        'LanguageApp selection, masked save and secret clearing',
                        'first-time setup creates a private ready-to-paste script; invalid URL, doPost failure and retry success',
                        'light/dark at 1080 and 1366; zero QML warnings',
                        'provider cannot change while work runs',
                        'two real PPTX inputs; stop/resume pins LanguageApp after Settings selects Google',
                        'completion retained; no Google usage; source PPTX unchanged'], qml_warnings=warnings)
            (report / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
            print(json.dumps(result, ensure_ascii=False))
        finally:
            release.set()
            window.close()
            bridge.batchQueue.shutdown()
            bridge.slidePreview.shutdown()
            bridge.googleSettings.shutdown()
            bridge._library_preferences.shutdown()
            bridge._catalog.shutdown()
            bridge.classroom.shutdown()
            bridge.libraryThumbnails.shutdown()
            bridge.voice_controls.shutdown()
            bridge.browserAI.shutdown()
            bridge.english_reader.shutdown()
            bridge.narration_preparation.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
