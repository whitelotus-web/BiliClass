"""Measure quantity line positions in a disposable native PowerPoint deck."""

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from app.slide_typography import protect_paragraph


def main():
    import pythoncom
    import win32com.client

    folder = Path('reports/rc88/quantities').resolve()
    folder.mkdir(parents=True, exist_ok=True)
    source = folder/'quantities.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    examples = ('Travel 108 km', 'Time is 10 s', 'Time is 10s', 'Speed 108 km/h')
    for index, text in enumerate(examples):
        box = slide.shapes.add_textbox(Pt(50+index*140), Pt(50), Pt(100), Pt(120))
        box.text = text
        box.text_frame.word_wrap = True
        paragraph = box.text_frame.paragraphs[0]
        paragraph.font.size = Pt(18)
        paragraph.font.name = 'Arial'
        protect_paragraph(paragraph._p)
    deck.save(source)
    pythoncom.CoInitialize()
    office = native = None
    checks = []
    try:
        office = win32com.client.DispatchEx('PowerPoint.Application')
        native = office.Presentations.Open(str(source), ReadOnly=True, Untitled=False, WithWindow=False)
        for shape, expected in zip(native.Slides(1).Shapes, examples, strict=True):
            text = shape.TextFrame2.TextRange
            value = text.Text
            begin = next(i for i, char in enumerate(value) if char.isdigit())
            characters = text._oleobj_.GetIDsOfNames('Characters')
            positions = [win32com.client.Dispatch(text._oleobj_.Invoke(
                characters, 0, pythoncom.DISPATCH_PROPERTYGET, True, index+1, 1)).BoundTop
                for index in range(begin, len(value)) if not value[index].isspace()]
            assert max(positions)-min(positions) < 1, (expected, positions)
            assert text.BoundWidth <= shape.Width+1 and text.BoundHeight <= shape.Height+1
            checks.append(expected)
    finally:
        if native is not None:
            native.Close()
        if office is not None and office.Presentations.Count == 0:
            office.Quit()
        pythoncom.CoUninitialize()
    result = {'status': 'passed', 'checks': checks, 'scope': 'native per-character line positions, disposable PPTX'}
    (folder/'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result), flush=True)


if __name__ == '__main__':
    main()
