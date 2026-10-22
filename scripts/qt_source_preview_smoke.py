"""Real Qt controls and native source previews in a disposable library."""
import json
import os
import tempfile
from multiprocessing import freeze_support
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, Qt
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
 app=QGuiApplication([])
 configure_application(app)
 QQuickStyle.setStyle('Basic')
 for f in (RESOURCE_ROOT/'assets').glob('*.ttf'):
  QFontDatabase.addApplicationFont(str(f))
 app.setFont(QFont('Be Vietnam Pro',10))
 out=Path(os.environ.get('BILICLASS_SOURCE_PREVIEW_REPORT','reports/rc123/source-preview'))
 out.mkdir(parents=True,exist_ok=True)
 with tempfile.TemporaryDirectory(prefix='bcl-source-preview-',ignore_cleanup_errors=True) as temp:
  temp=Path(temp)
  os.environ['BILICLASS_DATA']=str(temp/'data')
  deck=Presentation()
  deck.slide_width,deck.slide_height=Inches(13.333),Inches(7.5)
  for n in range(1,4):
   page=deck.slides.add_slide(deck.slide_layouts[6])
   box=page.shapes.add_textbox(Inches(1),Inches(1),Inches(10),Inches(2))
   box.text=f'Slide gốc {n}'
   box.text_frame.paragraphs[0].font.size=Pt(40)
  path=temp/'Bài kiểm thử.pptx'
  deck.save(path)
  library=Library(temp/'data')
  bridge=Bridge(library)
  engine=QQmlApplicationEngine()
  warnings=[]
  engine.warnings.connect(lambda values:warnings.extend(map(str,values)))
  engine.rootContext().setContextProperty('bridge',bridge)
  engine.load(RESOURCE_ROOT/'qml/Main.qml')
  root=engine.rootObjects()[0]
  def c(name): return control(root,name)
  results=[]
  try:
   root.setProperty('page','new')
   root.show()
   root.raise_()
   root.requestActivate()
   QMetaObject.invokeMethod(root,'acceptInputDocuments',Qt.DirectConnection,Q_ARG('QVariant',[path.as_uri()]))
   until(lambda:bool(bridge.inputDocuments.preview.get('image')),timeout=120)
   assert not bridge.inputDocuments.preview.get('error'),bridge.inputDocuments.preview
   for theme in ['light','dark']:
    bridge.setApplicationTheme(theme)
    for w,h in [(1366,850),(1920,1080),(1080,700)]:
     root.resize(w,h)
     pump(200)
     assert c('creationSelectedSlidePreview').isVisible()
     click(root,c('creationPreviewButton'))
     pump(200)
     assert root.property('creationPreviewExpanded')
     assert c('creationExpandedSlidePreview').isVisible() and not c('creationSelectedSlidePreview').isVisible()
     assert not root.findChild(type(c('creationPreviewButton')),'creationInputPreviewDialog')
     assert c('creationSlideCounter').property('text')=='Slide 1 / 3'
     assert not c('creationPreviousSlide').isEnabled()
     assert c('creationConfiguration').isEnabled() and not bridge.busy
     click(root,c('creationNextSlide'))
     until(lambda:bridge.inputDocuments.preview.get('slide')==2 and not bridge.inputDocuments.preview.get('busy'),timeout=120)
     pump(100)
     assert c('creationSlideCounter').property('text')=='Slide 2 / 3'
     first_image=bridge.inputDocuments.preview['image']
     click(root,c('creationNextSlide'))
     until(lambda:bridge.inputDocuments.preview.get('slide')==3 and not bridge.inputDocuments.preview.get('busy'),timeout=120)
     assert not c('creationNextSlide').isEnabled() and first_image!=bridge.inputDocuments.preview['image']
     click(root,c('creationPreviousSlide'))
     until(lambda:not bridge.inputDocuments.preview.get('busy'),timeout=120)
     click(root,c('creationPreviousSlide'))
     until(lambda:not bridge.inputDocuments.preview.get('busy'),timeout=120)
     c('creationScroll').property('contentItem').setProperty('contentY',0)
     pump(100)
     root.grabWindow().save(str(out/f'{theme}-{w}.png'))
     click(root,c('creationPreviewButton'))
     pump(100)
     assert not c('creationSlideNavigation').isVisible() and c('creationSelectedSlidePreview').isVisible()
     results.append(dict(theme=theme,size=[w,h],pages=3,inline=True))
   click(root,c('creationSelectedSlidePreview'))
   pump(100)
   assert root.property('creationPreviewExpanded')
   QMetaObject.invokeMethod(root,'removeInputDocument',Qt.DirectConnection)
   pump(50)
   assert not root.property('creationPreviewExpanded') and not bridge.inputDocuments.preview
   assert not warnings,warnings
   (out/'result.json').write_text(json.dumps(dict(status='passed',cases=results,warnings=warnings),indent=2))
   print(json.dumps(dict(status='passed',cases=len(results),warnings=warnings)))
  finally:
   bridge.inputDocuments.shutdown()
   for worker in (bridge._catalog,bridge.libraryThumbnails,bridge.slidePreview,
                  bridge.browserAI,bridge.english_reader,bridge.narration_preparation,
                  bridge.classroom,bridge.voice_controls):
    worker.shutdown()
   root.close()
   engine.deleteLater()
   QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
   app.processEvents()
   library.close()
if __name__=='__main__':
 freeze_support()
 main()
