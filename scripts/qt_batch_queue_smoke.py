"""Real Qt multi-file flows; disposable lessons and WAVs, no global input or API usage."""

import json
import shutil
import sys
import tempfile
import time
import wave
from pathlib import Path
from threading import Event
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlExpression
from PySide6.QtQuickControls2 import QQuickStyle

from app import speech
from app.batch_queue import BatchQueue
from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.voice_import import import_presentation
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT/'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/batch-queue')
    report.mkdir(parents=True, exist_ok=True)
    (report/'result.json').unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-queue-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        sources = []
        for n in range(3):
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = f'Bài {n+1}: Cơ năng'
            slide.notes_slide.notes_text_frame.text = f'VI: Bài {n+1}. Năng lượng được bảo toàn.\nEN: Lesson {n+1}. Energy is conserved.'
            path = folder/f'Bài {n+1} — Năng lượng.pptx'
            deck.save(path)
            sources.append(path)
        original = {path.name: path.read_bytes() for path in sources}
        library = Library(folder/'data')
        library.set_setting('teaching_subjects', ['Toán'])
        library.set_setting('voice_options', {'voice': 'pair:truc-ly', 'rate': 0})
        bridge = Bridge(library)
        bridge.previewConvertedSlide = lambda *_: None
        engine = QQmlApplicationEngine()
        warnings, ticks, gaps = [], [], []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT/'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        timer = QTimer()
        timer.setInterval(20)
        last_tick = [None]

        def beat():
            if not bridge.batchQueue.running:
                last_tick[0] = None
                return
            now = time.monotonic()
            ticks.append(now)
            if last_tick[0] is not None:
                gaps.append(now-last_tick[0])
            last_tick[0] = now
        timer.timeout.connect(beat)
        entered, release = Event(), Event()
        calls, failing = [], set()
        active = peak = 0

        def c(name):
            return control(window, name)

        def qml(source):
            expression = QQmlExpression(engine.contextForObject(window), window, source)
            value, undefined = expression.evaluate()
            assert not expression.hasError() and not undefined, expression.error().toString()
            return value

        def slow_convert(directory, row, options, cancel, update):
            nonlocal active, peak
            active += 1
            peak = max(active, peak)
            calls.append(row['id'])
            try:
                update({'percent': 30, 'message': 'Đang chuyển đổi · Slide 1', 'folder': 'fixture-checkpoint'}, True)
                entered.set()
                while not release.wait(.02):
                    if cancel.is_set():
                        raise InterruptedError()
                if row['name'] == sources[1].name and row['id'] not in failing:
                    failing.add(row['id'])
                    raise ValueError('Slide 2: một khung chữ chưa đọc được. Phần đã xong được giữ.')
                output = folder/(Path(row['name']).stem+' — kết quả thử.pptx')
                shutil.copyfile(Path(QUrl(row['url']).toLocalFile()), output)
                return {'lesson_id': 'fixture-'+row['id'], 'output': str(output)}
            finally:
                active -= 1

        def accept(method, paths):
            assert QMetaObject.invokeMethod(window, method, Qt.DirectConnection,
                Q_ARG('QVariant', [QUrl.fromLocalFile(str(path)).toString() for path in paths]))
            pump(60)

        def screenshots(label):
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                for width, height in ((1080, 700), (1366, 850)):
                    window.resize(width, height)
                    pump(70)
                    assert window.grabWindow().save(str(report/f'{label}-{theme}-{width}.png'))

        try:
            window.resize(1366, 850)
            window.show()
            window.setProperty('page', 'new')
            pump(150)
            expression = QQmlExpression(engine.contextForObject(window), window,
                'powerPointInputOpen.fileMode === FileDialog.OpenFiles')
            modes, undefined = expression.evaluate()
            assert not expression.hasError() and not undefined and modes, expression.error().toString()
            # selectedFiles is read-only in this installed Qt. Exercise the same
            # QML receiver with selected local URLs; do not inject Windows input.
            accept('acceptInputDocuments', sources)
            until(lambda: bridge.batchQueue.state['total'] == 3)
            assert window.property('creationBatch')
            assert not c('lessonTitle').isEnabled()
            assert c('createLessonButton').isEnabled()
            assert '3' in c('createLessonButton').property('text')
            screenshots('conversion-selected')
            with patch('app.batch_jobs.convert_one', slow_convert):
                timer.start()
                click(window, c('createLessonButton'))
                until(entered.is_set)
                assert bridge.busy and bridge.batchQueue.running
                assert not c('createLessonButton').isEnabled()
                assert c('creationConfiguration').isVisible()
                assert not c('conversionConfigurationFields').isEnabled()
                assert c('conversionBatchPanel').parentItem().parentItem() == c('creationConfiguration')
                pump(1250)
                until(lambda: c('batchTotalTime_convert').property('text') != 'Tổng thời gian: 00:00:00', timeout=3)
                elapsed = bridge.batchQueue.state['elapsed_seconds']
                assert elapsed >= 1
                assert c('batchTotalTime_convert').property('text') != 'Tổng thời gian: 00:00:00'
                assert c('batchPercent_convert').property('text') == '30%'
                screenshots('conversion-running')
                assert len(ticks) >= 8, ticks
                click(window, c('batchStop_convert'))
                until(lambda: not bridge.busy)
                assert bridge.batchQueue.state['phase'] == 'paused'
                assert bridge.batchQueue.state['counts']['waiting'] == 3
                stopped_time = bridge.batchQueue.state['elapsed_seconds']
                pump(1100)
                assert bridge.batchQueue.state['elapsed_seconds'] == stopped_time
                assert c('creationConfiguration').isVisible() and c('batchResume_convert').isEnabled()
                restored = BatchQueue(bridge)
                assert restored.data['id'] == bridge.batchQueue.data['id']
                assert restored.data['options'] == bridge.batchQueue.data['options']
                assert not restored.running
                window.setProperty('page', 'home')
                pump(80)
                assert c('applicationStatusMessage').isEnabled()
                click(window, c('applicationStatusMessage'))
                assert window.property('page') == 'new'
                release.set()
                pump(120)
                click(window, c('batchResume_convert'))
                until(lambda: not bridge.busy)
                assert bridge.batchQueue.state['counts']['done'] == 2, (bridge.batchQueue.state, bridge.message, calls)
                assert bridge.batchQueue.state['counts']['failed'] == 1
                assert peak == 1
                screenshots('conversion-partial')
                failed = next(row for row in bridge.batchQueue.data['rows'] if row['status'] == 'failed')
                assert 'Slide 2' in c('batchStatus_'+failed['id']).property('text')
                row = bridge.batchQueue.data['rows'][0]
                opened, folders = [], []
                with patch.object(bridge, 'openLesson', side_effect=opened.append):
                    click(window, c('batchResult_'+row['id']))
                assert opened == [row['lesson_id']] and not bridge.batchQueue.state['shown']
                bridge.batchQueue.show()
                with patch('PySide6.QtGui.QDesktopServices.openUrl', side_effect=lambda url: folders.append(Path(url.toLocalFile())) or True):
                    click(window, c('batchFolder_'+row['id']))
                assert folders == [folder]
                before = len(calls)
                click(window, c('batchRetry_convert'))
                until(lambda: not bridge.busy)
                assert calls[before:] == [failed['id']]
                assert bridge.batchQueue.state['counts']['done'] == 3
            assert bridge.batchQueue.state['complete'] and not bridge.batchQueue.state['shown']
            assert not c('conversionBatchPanel').isVisible()
            assert not qml('creationBatch') and c('creationSourceArea').isVisible()
            assert BatchQueue(bridge).state['complete'] and not BatchQueue(bridge).state['shown']
            bridge.batchQueue.show()
            pump(60)
            assert not c('conversionBatchPanel').isVisible(), 'Completed queue stays out of input'
            bridge.batchQueue.clear()
            until(lambda: not bridge.busy)
            assert not bridge.batchQueue.data
            assert not BatchQueue(bridge).data, 'Dismissal survives restart'
            assert c('lessonTitle').property('text') != 'Tên theo từng file'

            # A selected single-language lesson must not restrict a new batch or get its position edited.
            old_path = folder/'Bài đang xem.pptx'
            old = Presentation()
            slide = old.slides.add_slide(old.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = 'Chỉ có lời tiếng Việt.'
            old.save(old_path)
            lesson, preview = import_presentation(library.directory, old_path)
            bridge._show_lesson(lesson, preview, destination='voice', auto_repair=False, load_preview=False)
            old_before = library.get(lesson['id'])
            accept('acceptVoiceDocuments', sources[:2])
            assert window.property('voiceBatch')
            assert not c('voiceLessonOverview').isVisible()
            assert c('narrationBoth').isVisible() and c('narrationBoth').isEnabled(), (window.property('page'), bridge.busy, bridge.batchQueue.state, c('lessonNarrationPanel').isVisible(), c('narrationCreationOptions').isVisible())
            click(window, c('narrationBoth'))
            assert window.property('narrationMode') == 'both'
            window.setProperty('narrationVoice', 'pair:truc-ly')
            click(window, c('narrationSpeakerPosition'))
            pump(80)
            click(window, c('speakerPosition_top_left'))
            assert window.property('batchSpeakerPosition') == 'top_left'
            assert library.get(lesson['id']) == old_before
            screenshots('voice-selected')
            generated = []

            def synthesize(text, voice, rate, directory, cancel):
                time.sleep(.06)
                path = Path(directory)/(speech.audio_key(text, voice, rate)+'.wav')
                path.parent.mkdir(parents=True, exist_ok=True)
                with wave.open(str(path), 'wb') as wav:
                    wav.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                    wav.writeframes(b'\1\0'*2400)
                generated.append(path)
                return {'path': str(path), 'cached': False}

            with patch.object(bridge, '_available_voice', return_value=True), patch('app.speech.synthesize_background', synthesize), patch('app.speech.narration_workers', return_value=1):
                click(window, c('createLessonNarrationButton'))
                until(lambda: bridge.batchQueue.running)
                assert c('lessonNarrationPanel').isVisible()
                assert not c('narrationCreationOptions').isEnabled()
                assert c('voiceBatchPanel').parentItem().parentItem() == c('lessonNarrationPanel')
                screenshots('voice-running')
                until(lambda: not bridge.busy, timeout=45)
            assert bridge.batchQueue.state['counts']['done'] == 2, bridge.batchQueue.state
            assert generated and not window.property('voiceResultShown')
            assert library.get(lesson['id']) == old_before
            for row in bridge.batchQueue.data['rows']:
                assert Path(row['output']).is_file() and Path(row['output']).parent == folder
                assert library.get(row['lesson_id'])['speaker_position'] == 'top_left'
            assert bridge.batchQueue.state['complete'] and not bridge.batchQueue.state['shown']
            assert not c('voiceBatchPanel').isVisible() and not qml('voiceBatch')
            screenshots('voice-finished')
            ready_id = bridge.batchQueue.data['rows'][0]['lesson_id']
            bridge.batchQueue.clear()
            until(lambda: not bridge.busy)
            # Publish one imported fixture as a library product so the actual
            # selection controls can queue its existing identity beside a ready lesson.
            pending = library.get(lesson['id'])
            pending['voice_published'] = True
            library.db.execute('UPDATE lessons SET content=? WHERE id=?', (json.dumps(pending), pending['id']))
            library.db.commit()
            bridge._catalog.refresh(force=True)
            window.setProperty('page', 'library')
            until(lambda: any(row['id'] == pending['id'] for row in bridge.lessons))
            until(lambda: bridge._catalog.worker is None)
            window.setProperty('selectedLessons', [pending['id'], ready_id])
            pump(80)
            assert c('libraryNarrateSelected').isVisible() and c('openLibraryPresentation').isVisible()
            opened_local = []
            dialog = c('libraryPresentationDialog')
            assert dialog.setProperty('selectedFile', QUrl.fromLocalFile(str(sources[0])))
            with patch.object(bridge, '_open_teaching_presentation',
                              side_effect=lambda url, **kwargs: opened_local.append((url, kwargs))):
                assert QMetaObject.invokeMethod(dialog, 'accepted', Qt.DirectConnection)
            assert opened_local == [(QUrl.fromLocalFile(str(sources[0])).toString(), {'destination': 'library'})]
            screenshots('library-selected')
            click(window, c('libraryTab_ready'))
            assert qml('selectedLessons.length') == 0
            assert qml('libraryResult.ids.indexOf('+json.dumps(pending['id'])+')') == -1
            click(window, c('libraryTab_none'))
            assert qml('libraryResult.ids.indexOf('+json.dumps(pending['id'])+')') >= 0
            click(window, c('libraryTab_all'))
            window.setProperty('selectedLessons', [pending['id'], ready_id])
            click(window, c('libraryNarrateSelected'))
            until(lambda: not bridge.busy)
            assert window.property('page') == 'voice'
            assert bridge.batchQueue.data['skipped_ready'] == 1
            assert [row['lesson_id'] for row in bridge.batchQueue.data['rows']] == [pending['id']]
            assert c('lessonNarrationPanel').isVisible() and c('narrationCreationOptions').isEnabled()
            assert library.db.execute('SELECT COUNT(*) FROM lessons WHERE id=?', (pending['id'],)).fetchone()[0] == 1
            # A long queue keeps the currently running file visible across progress model updates.
            bridge.batchQueue.clear()
            until(lambda: not bridge.busy)
            window.setProperty('page', 'new')
            long_sources = [folder/f'Bài chờ {n:02d}.pptx' for n in range(25)]
            accept('acceptInputDocuments', long_sources)
            last_entered = Event()

            def long_process(directory, row, options, cancel, update):
                if row['name'] == long_sources[-1].name:
                    last_entered.set()
                    for tick in range(125):
                        update({'message': 'Đang xử lý bài cuối', 'percent': tick % 100})
                        if cancel.wait(.08):
                            raise InterruptedError()
                return {}

            with patch('app.batch_jobs.convert_one', long_process):
                click(window, c('createLessonButton'))
                until(last_entered.is_set)
                until(lambda: bridge.batchQueue.data['rows'][-1]['status'] == 'running')
                pump(250)
                assert bridge.batchQueue.running, bridge.batchQueue.state
                view = c('batchRows_convert')
                last_name = c('batchName_'+bridge.batchQueue.data['rows'][-1]['id'])
                assert view.property('contentY') > 0
                assert 0 <= last_name.mapToItem(view, QPointF()).y() < view.height()
                click(window, c('applicationCancel'))
                until(lambda: not bridge.busy)
            assert bridge.batchQueue.state['counts']['done'] == 24
            assert bridge.batchQueue.state['counts']['waiting'] == 1
            bridge.batchQueue.clear()
            until(lambda: not bridge.busy)
            assert all(path.read_bytes() == original[path.name] for path in sources)
            assert not warnings, warnings
            assert max(gaps, default=0) < .8, max(gaps)
            (report/'result.json').write_text(json.dumps(dict(status='passed',
                multi_file_qml_receiver=True, both_dialogs_openfiles_mode=True, native_dialog_selection_not_automated=True, maximum_simultaneous_lessons=peak,
                stop_resume_checkpoint_and_restart=True, local_failure_continues=True, retry_only_failed=True,
                exact_result_and_folder_buttons=True, narration_real_exports_with_fixture_wavs=True,
                unrelated_selected_lesson_unchanged=True, sources_preserved=True, long_queue_follows_active_file=True,
                finished_queues_hidden_and_stay_hidden_after_restart=True,
                max_gui_gap_seconds=round(max(gaps, default=0), 3), qml_warnings=[])), encoding='utf-8')
            print('Sequential conversion and narration Qt flows passed.', flush=True)
        finally:
            release.set()
            bridge.batchQueue.stop()
            until(lambda: not bridge.busy, timeout=60)
            timer.stop()
            for service in (bridge._catalog, bridge.browserAI, bridge.googleSettings, bridge.slidePreview,
                            bridge.libraryThumbnails, bridge.narrationPreferences, bridge.narration_preparation,
                            bridge.english_reader, bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
                service.shutdown()
            c('companionWindow').setProperty('disposing', True)
            c('companionWindow').hide()
            window.hide()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            pump(60)
            library.close()


if __name__ == '__main__':
    main()
