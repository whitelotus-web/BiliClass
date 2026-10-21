"""Click the description/padding of real voice/conversion cards, using disposable data."""

import json
import tempfile
from pathlib import Path

from pptx import Presentation
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge

from .qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    report = Path('reports/rc138/choice-cards')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='biliclass-choice-', ignore_cleanup_errors=True) as temporary:
        library = Library(Path(temporary))
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT / 'qml/Main.qml')
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        source = Path(temporary) / 'Bài thử.pptx'
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(source)

        def c(name):
            return control(root, name)

        def click_at(item, point):
            point = item.mapToScene(point).toPoint()
            assert 0 <= point.y() < root.height(), point
            QTest.mouseClick(root, Qt.LeftButton, Qt.NoModifier, point)
            pump(70)

        checks = []
        try:
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                for width, height in ((1080, 700), (1366, 850)):
                    root.resize(width, height)
                    root.setProperty('page', 'voice')
                    pump(180)
                    for mode, name in (('vi', 'narrationVietnamese'), ('both', 'narrationBoth'),
                                       ('en', 'narrationEnglish')):
                        click(root, c(name + 'Detail'))
                        assert root.property('narrationMode') == mode, (name, 'description did not select')
                        assert c(name).property('checked')
                        assert sum(c(n).property('checked') for n in
                                   ('narrationEnglish', 'narrationVietnamese', 'narrationBoth')) == 1
                    card = c('narrationVietnameseCard')
                    click_at(card, QPointF(card.width() - 4, card.height() - 4))
                    assert root.property('narrationMode') == 'vi', 'Card padding must select'
                    root.requestActivate()
                    c('narrationEnglish').forceActiveFocus(Qt.TabFocusReason)
                    QTest.keyClick(root, Qt.Key_Space)
                    pump(60)
                    assert root.property('narrationMode') == 'en', 'Keyboard selection must remain usable'
                    bridge._busy = True
                    bridge.changed.emit()
                    pump(40)
                    assert not c('narrationVietnameseDetail').isEnabled()
                    detail = c('narrationVietnameseDetail')
                    click_at(detail, QPointF(detail.width()/2, detail.height()/2))
                    assert root.property('narrationMode') == 'en', 'Disabled description must not select'
                    bridge._busy = False
                    bridge.changed.emit()
                    pump(40)
                    root.grabWindow().save(str(report / f'voice-{theme}-{width}.png'))

                    root.setProperty('page', 'new')
                    QMetaObject.invokeMethod(root, 'acceptInputDocuments', Qt.DirectConnection,
                                             Q_ARG('QVariant', [source.as_uri()]))
                    until(lambda: not bridge.inputDocuments.busy)
                    pump(130)
                    formats = bridge.conversionFormats
                    assert len(formats) >= 2
                    first, second = ['formatChoice_' + row['id'] for row in formats[:2]]
                    click(root, c(second + 'Detail'))
                    assert c('creationFormat').property('currentIndex') == 1
                    click(root, c(first))
                    assert c('creationFormat').property('currentIndex') == 0
                    click(root, c(second + 'Preview'))
                    assert c('conversionExampleDialog').property('visible')
                    assert c('creationFormat').property('currentIndex') == 0, 'Preview must not select the card'
                    QMetaObject.invokeMethod(c('conversionExampleDialog'), 'close', Qt.DirectConnection)
                    pump(130)
                    click(root, c('creationUseTemplateDetail'))
                    assert c('creationWorkflow').property('currentIndex') == 1
                    for index, preset in ((1, 'visual'), (2, 'practice'), (0, 'standard')):
                        click(root, c('creationPresetDetail_' + preset))
                        assert c('creationPreset').property('currentIndex') == index
                        preview, detail = c('creationPresetView_' + preset), c('creationPresetDetail_' + preset)
                        assert abs(preview.mapToScene(QPointF(0, preview.height()/2)).y()
                                   - detail.mapToScene(QPointF(0, detail.height()/2)).y()) <= 1
                    click(root, c('creationPresetView_visual'))
                    assert c('creationTemplateImageDialog').property('visible')
                    assert c('creationPreset').property('currentIndex') == 0, 'Preset preview must not select'
                    QMetaObject.invokeMethod(c('creationTemplateImageDialog'), 'close', Qt.DirectConnection)
                    pump(130)
                    root.grabWindow().save(str(report / f'conversion-{theme}-{width}.png'))
                    click(root, c('creationKeepOriginalDetail'))
                    assert c('creationWorkflow').property('currentIndex') == 0
                    checks.append({'theme': theme, 'size': [width, height]})
            assert not warnings, warnings
            result = {'status': 'passed', 'cases': checks, 'qml_warnings': warnings,
                      'checks': ['voice descriptions and padding select exactly one language',
                                 'keyboard and disabled guards', 'conversion and design descriptions',
                                 'nested preview opens without selecting its card']}
            (report / 'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result), flush=True)
        finally:
            bridge._busy = False
            c('companionWindow').setProperty('disposing', True)
            c('companionWindow').hide()
            root.hide()
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview',
                         '_catalog', 'browserAI', 'english_reader', 'narration_preparation',
                         'narrationPreferences', '_companion_preferences', 'classroom', 'voice_controls',
                         'inputDocuments'):
                getattr(bridge, name).shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
