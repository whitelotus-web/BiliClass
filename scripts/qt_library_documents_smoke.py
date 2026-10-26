"""Shared Word/PowerPoint library and matching conversion headers; disposable fixtures."""

import json
import sys
import tempfile
from pathlib import Path
from threading import Event
from unittest.mock import patch

from docx import Document
from PIL import Image, ImageDraw
from PySide6.QtCore import QCoreApplication, QEvent, QMetaObject, QPointF, Qt, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import word_jobs
from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.word_document import Package
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT / 'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/rc169-library/qt-mixed')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-shared-library-', ignore_cleanup_errors=True) as temporary:
        directory = Path(temporary)
        library = Library(directory / 'data')
        slides = []
        for title in ('Năng lượng trong cuộc sống', 'Dao động điều hòa'):
            lesson = library.create(title, 'Vật lý', 'THPT', '11', [('Slide 1', 'Vận tốc 12 m/s.')])
            slides.append(library.edit_segment(lesson['id'], lesson['segments'][0]['id'], 'Vận tốc 12 m/s.', 'Velocity 12 m/s.', True))
        source = directory / 'Kế hoạch bài dạy - Vận tốc.docx'
        document = Document()
        document.add_heading('Bài 11. Vận tốc và gia tốc', 0)
        document.add_paragraph('Vận tốc 12 m/s. Quan sát chuyển động và giải thích kết quả thí nghiệm.')
        document.save(source)
        job = word_jobs.create(library.directory, source)
        row = Package(source).rows[1]
        region = Package(source).selection([dict(paragraph=row['id'], start=0, end=len('Vận tốc 12 m/s.'))])[0]
        word_jobs.selections(library.directory, job['id'], [region])
        word_jobs.edit(library.directory, job['id'], region['id'], 'Velocity 12 m/s.')
        word_jobs.export(library.directory, job['id'], Event())
        identity = 'word:' + job['id']
        audio_calls = []

        def audio(lesson, *_args, **_kwargs):
            audio_calls.append(lesson['id'])
            ready = lesson['id'] == slides[0]['id']
            return dict(ready=ready, status='ready' if ready else 'none', mode='en' if ready else 'none',
                        label='Có giọng tiếng Anh' if ready else 'Chưa có giọng',
                        variants=[dict(id='fixture-en', ready=True, mode='en', voice='heart', rate=100)] if ready else [])

        with patch('app.library_catalog.audio_status', audio):
            bridge = Bridge(library)
            for index, lesson in enumerate(slides):
                image = Image.new('RGB', (640, 360), '#e3effc')
                draw = ImageDraw.Draw(image)
                draw.rounded_rectangle((22, 22, 618, 338), radius=18, fill='#ffffff', outline='#b7cce5', width=2)
                draw.rectangle((48, 58, 440, 78), fill='#2773c9')
                for line in range(4):
                    draw.rectangle((48, 118 + 32 * line, 360 - 20 * line, 127 + 32 * line), fill='#c1d6ee')
                draw.ellipse((430, 130, 565, 265), fill='#8ab8e9')
                cover = directory / f'cover-{index}.png'
                image.save(cover)
                bridge.libraryThumbnails.results[(lesson['id'], lesson['revision'])] = QUrl.fromLocalFile(str(cover)).toString()
            engine = QQmlApplicationEngine()
            warnings = []
            engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
            engine.rootContext().setContextProperty('bridge', bridge)
            engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
            assert engine.rootObjects(), warnings
            window = engine.rootObjects()[0]
            window.resize(1366, 850)
            window.show()

            def c(name):
                return control(window, name)

            def settled():
                return bridge._catalog.worker is None and bridge._catalog.word_worker is None and not bridge.busy

            try:
                window.setProperty('page', 'word')
                pump(150)
                assert c('wordSourceArea').isVisible() and c('wordUploadChoose').isVisible()
                assert window.grabWindow().save(str(report / 'word-empty.png'))
                window.setProperty('page', 'library')
                until(lambda: settled() and len(bridge.libraryDocuments) == 3)
                assert len(bridge.lessons) == 2, 'Word must not enter the PowerPoint/narration model'
                counts = window.property('libraryResult')['groups']
                assert counts == dict(all=3, none=1, ready=1, word=1), counts
                for group, ids in [('none', [slides[1]['id']]), ('ready', [slides[0]['id']]), ('word', [identity])]:
                    click(window, c('libraryTab_' + group))
                    assert window.property('libraryResult')['ids'] == ids
                assert not c('libraryVoiceFilter').isEnabled()
                click(window, c('lessonDownload_' + identity))
                assert c('libraryDownloadDialog').property('singleWord')
                assert not c('libraryDownloadVoice').isVisible()
                QMetaObject.invokeMethod(c('libraryDownloadDialog'), 'close', Qt.DirectConnection)
                assert bridge.libraryDownloadSuggestion([identity], True).endswith('.docx')
                click(window, c('libraryTab_all'))
                click(window, c('librarySelectAll'))
                assert len(window.property('selectedLessons').toVariant()) == 3
                assert set(window.property('selectedPowerPoints').toVariant()) == {row['id'] for row in slides}
                for theme in ('light', 'dark'):
                    bridge.setApplicationTheme(theme)
                    for view in ('list', 'grid'):
                        bridge.setLibraryView(view)
                        for width, height in ((1080, 700), (1366, 850)):
                            window.resize(width, height)
                            pump(180)
                            for key in ('all', 'none', 'ready', 'word'):
                                tab = c('libraryTab_' + key)
                                right = tab.mapToScene(QPointF(tab.width(), 0)).x()
                                assert right < window.width(), (key, width, right)
                            narration = c('libraryNarrateSelected')
                            assert narration.mapToScene(QPointF(narration.width(), 0)).x() < window.width()
                            assert c('lessonType_' + identity).property('word')
                            assert c('lessonNarrate_' + identity).property('text') == 'Sửa song ngữ'
                            assert window.grabWindow().save(str(report / f'{theme}-{view}-{width}.png'))
                assert all(not key[0].startswith('word:') for key in bridge.libraryThumbnails.wanted)
                until(lambda: bridge._library_preferences.worker is None and not bridge._library_preferences.pending)
                bridge._catalog.refresh()  # Incorporate the last asynchronous view-preference write first.
                until(settled)
                scan_count = len(audio_calls)
                click(window, c('lessonNarrate_' + identity))
                until(lambda: settled() and window.property('page') == 'word')
                assert bridge.wordDocuments.state['id'] == job['id']
                assert c('wordPageHeading').property('text') == 'Chuyển đổi bài giảng song ngữ'
                assert c('wordTabWord').mapToScene(QPointF()).y() > c('wordPageHeading').mapToScene(QPointF()).y()
                assert len(audio_calls) == scan_count, 'Word refresh must not rescan PowerPoint audio'
                assert window.grabWindow().save(str(report / 'word-matching-layout.png'))
                click(window, c('wordTabPowerPoint'))
                assert c('creationPageHeading').property('text') == 'Chuyển đổi bài giảng song ngữ'
                assert c('conversionDocumentPowerPoint').mapToScene(QPointF()).y() > c('creationPageHeading').mapToScene(QPointF()).y()
                assert window.grabWindow().save(str(report / 'powerpoint-matching-layout.png'))
                window.setProperty('page', 'library')
                until(settled)
                ids = [row['id'] for row in bridge.libraryDocuments]
                assert bridge.trashLessons(ids)
                until(lambda: settled() and not bridge.libraryDocuments)
                assert bridge.undoDeleteCount == 3
                assert bridge.undoDeletion()
                until(lambda: settled() and len(bridge.libraryDocuments) == 3)
                assert source.is_file() and len(library.list_lessons()) == 2
                assert not warnings, warnings
                result = dict(status='passed', live_api=False, qml_warnings=warnings,
                    checks=['shared counts and type icons; voice tabs contain only PowerPoint',
                            'light/dark list/cards at 1080 and 1366 with mixed selection',
                            'Word download is docx and never offers voice',
                            'only selected PowerPoints enter narration; Word never starts a slide cover worker',
                            'Word edit resumes correct record; Word updates do not rescan PowerPoint audio',
                            'matching conversion headings and tabs below headings; workflows remain independent',
                            'mixed trash and undo restore all records; teacher originals retained'])
                (report / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
                print(json.dumps(result, ensure_ascii=False))
            finally:
                window.close()
                for component in (bridge.batchQueue, bridge.slidePreview, bridge.googleSettings, bridge._library_preferences,
                                  bridge._catalog, bridge.classroom, bridge.libraryThumbnails, bridge.voice_controls,
                                  bridge.browserAI, bridge.english_reader, bridge.narration_preparation, bridge.inputDocuments):
                    component.shutdown()
                engine.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                app.processEvents()
                library.close()


if __name__ == '__main__':
    main()
