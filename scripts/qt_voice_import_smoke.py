"""Real voice import responsiveness and matching selected-file layouts; isolated data."""

import argparse
import hashlib
import json
import tempfile
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt, QTimer
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from scripts.qt_result_preview_mascot_smoke import control, pump, until


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, help='Optional existing PPTX; read only')
    parser.add_argument('--report', type=Path, default=Path('reports/voice-import'))
    args = parser.parse_args()
    args.report.mkdir(parents=True, exist_ok=True)
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT/'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    with tempfile.TemporaryDirectory(prefix='bcl-voice-import-', ignore_cleanup_errors=True) as folder:
        directory = Path(folder)
        source = args.source.resolve() if args.source else directory/'Bài học #1 – năng lượng.pptx'
        if not args.source:
            deck = Presentation()
            deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
            for number in range(1, 49):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2)).text = f'Năng lượng – bài học {number}'
                slide.notes_slide.notes_text_frame.text = f'VI: Năng lượng được bảo toàn. Slide {number}.'
            deck.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        library = Library(directory/'library')
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT/'qml/Main.qml')
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        root.resize(1366, 850)
        root.setProperty('page', 'voice')
        ticks, timings = [], []
        heartbeat = QTimer()
        heartbeat.setInterval(20)
        heartbeat.timeout.connect(lambda: ticks.append(time.monotonic()))

        def c(name):
            return control(root, name)

        try:
            pump(120)
            for run in ('first', 'repeat'):
                ticks.clear()
                heartbeat.start()
                started = time.monotonic()
                capture_windows = []
                bridge.importNarrationPresentation(source.as_uri())
                assert time.monotonic()-started < .3
                assert bridge.busy and bridge.narrationImportName == source.name
                pump(30)
                if run == 'first':
                    assert c('voiceSourceArea').property('loadingFileName') == source.name
                    capture_start = time.monotonic()
                    root.grabWindow().save(str(args.report/'loading.png'))
                    capture_windows.append((capture_start, time.monotonic()))
                until(lambda: not bridge.busy, timeout=40)
                elapsed = time.monotonic()-started
                heartbeat.stop()
                assert not bridge.error, bridge.message
                assert not bridge.narrationImportName and bridge.quickResult.get('path')
                if not args.source:
                    assert c('narrationVietnamese').isEnabled()
                    assert not c('narrationEnglish').isEnabled() and not c('narrationBoth').isEnabled()
                    assert c('narrationLanguageAvailability').isVisible()
                    assert c('narrationEnglish').property('opacity') < .5
                samples = [started, *ticks, time.monotonic()]
                # Synchronous screenshot readback is test overhead, not GUI work.
                gaps = [b-a for a, b in zip(samples, samples[1:])
                        if not any(a <= end and b >= start for start, end in capture_windows)]
                max_gap = max(gaps, default=0)
                assert max_gap < 1.0, (run, elapsed, max_gap)
                timings.append(dict(run=run, seconds=round(elapsed, 3), gui_max_gap=round(max_gap, 3)))
                until(lambda: bool(bridge.quickResult.get('image')) and not bridge.previewBusy, timeout=90)
            assert len(library.list_lessons(verify_audio=False)) == 1
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                for width, height in ((1080, 700), (1366, 850), (1920, 1080)):
                    root.resize(width, height)
                    root.setProperty('page', 'voice')
                    pump(100)
                    voice = c('voiceSlideThumbnail')
                    voice_width = voice.width()
                    root.grabWindow().save(str(args.report/f'{theme}-voice-{width}.png'))
                    root.setProperty('page', 'new')
                    assert QMetaObject.invokeMethod(root, 'acceptInputDocuments', Qt.DirectConnection,
                                                    Q_ARG('QVariant', [source.as_uri()]))
                    until(lambda: not bridge.inputDocuments.busy and not bridge.inputDocuments.preview.get('busy'), timeout=90)
                    pump(100)
                    assert bridge.inputDocuments.preview.get('image'), bridge.inputDocuments.preview
                    thumbnail = c('creationSelectedSlidePreview')
                    assert abs(thumbnail.width()-voice_width) <= 1
                    details = c('creationLessonDetails')
                    assert thumbnail.mapToScene(QPointF()).x() >= details.mapToScene(QPointF(details.width(), 0)).x()+23
                    assert c('removeInputDocumentButton').isVisible() and not c('creationSourceArea').isVisible()
                    summary = c('conversionConfigurationSummary')
                    assert summary.isVisible() and summary.height() <= 110
                    assert c('narrationConfiguration').property('selectedVoice')
                    root.grabWindow().save(str(args.report/f'{theme}-conversion-{width}.png'))
            assert hashlib.sha256(source.read_bytes()).hexdigest() == original
            assert not warnings, warnings
            result = dict(status='passed', timings=timings, source_bytes=source.stat().st_size,
                          slides=bridge.quickResult['total'], source_unchanged=True, qml_warnings=warnings,
                          matched_layouts=6)
            (args.report/'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result), flush=True)
        finally:
            heartbeat.stop()
            for worker in (bridge.inputDocuments, bridge._catalog, bridge.browserAI, bridge.slidePreview,
                           bridge.libraryThumbnails, bridge.narrationPreferences, bridge.narration_preparation,
                           bridge.english_reader, bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
                worker.shutdown()
            c('companionWindow').setProperty('disposing', True)
            c('companionWindow').hide()
            root.hide()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
