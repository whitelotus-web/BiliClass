"""Exercise the companion against a disposable deck, without opening user files."""

import hashlib
import json
import tempfile
import time
from multiprocessing import freeze_support
from pathlib import Path

from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEventLoop, QTimer

from app.powerpoint import PowerPointSession


def main():
    app = QCoreApplication([])
    states, errors = [], []

    def wait(predicate):
        start = time.monotonic()
        while not predicate():
            loop = QEventLoop()
            QTimer.singleShot(40, loop.quit)
            loop.exec()
            if time.monotonic() - start > 50:
                raise AssertionError(f"PowerPoint timeout: {errors}")

    with tempfile.TemporaryDirectory(prefix="biliclass-companion-") as temporary:
        path = Path(temporary) / "Bài thử đa môn.pptx"
        deck = Presentation()
        for name in ("Mở bài", "Thảo luận", "Tổng kết"):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = name
        deck.save(path)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        worker = PowerPointSession(path)
        worker.stateChanged.connect(states.append)
        worker.failed.connect(errors.append)
        worker.start()
        try:
            for command, expected in ((None, 1), ("next", 2), ("goto", 3), ("previous", 2)):
                if command:
                    worker.navigate(command, expected if command == "goto" else None)
                wait(lambda: bool(states) and states[-1].get("slide") == expected)
            assert len({s["slide_id"] for s in states if s.get("active")}) == 3
            assert hashlib.sha256(path.read_bytes()).hexdigest() == original
            assert not errors, errors
        finally:
            worker.stop()
            wait(lambda: not worker.isRunning())
        report = {"status": "passed", "trace": states, "source_unchanged": True,
                  "scope": "Office on current machine, windowed slideshow; no projector certification"}
        target = Path(__file__).resolve().parents[1] / "reports/app/powerpoint-companion.json"
        target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"status": "passed", "slides": 3}))
    app.quit()


if __name__ == "__main__":
    freeze_support()
    main()
