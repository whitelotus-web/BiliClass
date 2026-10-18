"""Generate eight owned fixture decks and inspect native bilingual output offline.

Use --native for sequential isolated PowerPoint rendering and bounded repair.
Reports/PNGs/PPTX remain local; fixture translations are not a translation audit.
"""

import argparse
import hashlib
import json
import time
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

from app.chatgpt_handoff import inspect_returned_deck
from app.native_bilingual import assemble, extract, leaves
from app.slide_refinement import source_media_preserved
from app.source_deck import NS

CASES = ('plain', 'table', 'formula', 'image', 'animation', 'dense', 'bilingual', 'group')
PAIRS = {
    'Cơ năng': 'Mechanical energy',
    'Năng lượng của vật được bảo toàn.': 'The energy of the object is conserved.',
    'Vận tốc của vật tăng khi thế năng giảm.': 'The velocity increases as potential energy decreases.',
    'Đại lượng x': 'Quantity x',
    'Khối lượng': 'Mass',
    'Vận tốc': 'Velocity',
}


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def text(shapes, value, x, y, w, h, size=24):
    shape = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text = value
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.name, paragraph.font.size = 'Arial', Pt(size)
    return shape


def fixture(path, case):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text(slide.shapes, 'Cơ năng', .5, .3, 12, 1, 32)
    body = text(slide.shapes, 'Năng lượng của vật được bảo toàn.\nVận tốc của vật tăng khi thế năng giảm.',
                .5, 1.8, 11.8, 4.8)
    if case == 'table':
        body.text = ''
        table = slide.shapes.add_table(2, 2, Inches(.5), Inches(2), Inches(11.8), Inches(3)).table
        for row, values in zip(table.rows, (('Khối lượng', '12 kg'), ('Vận tốc', '24 m/s')), strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    elif case == 'formula':
        paragraph = body.text_frame.add_paragraph()
        paragraph.text = 'Đại lượng x'
        run = paragraph.add_run()
        run.text = '2'
        run._r.get_or_add_rPr().set('baseline', '30000')
    elif case == 'image':
        image = Image.new('RGB', (400, 180), 'white')
        canvas = ImageDraw.Draw(image)
        canvas.line((40, 140, 340, 140, 340, 30), fill='#1476b8', width=4)
        canvas.rectangle((120, 70, 230, 140), outline='#1476b8', width=4)
        image.save(path.with_suffix('.png'))
        body.width = Inches(7)
        slide.shapes.add_picture(str(path.with_suffix('.png')), Inches(8), Inches(2), width=Inches(4))
    elif case == 'animation':
        timing = etree.SubElement(slide._element, f'{{{NS["p"]}}}timing')
        sequence = etree.SubElement(timing, f'{{{NS["p"]}}}tnLst')
        parent = etree.SubElement(sequence, f'{{{NS["p"]}}}par')
        clock = etree.SubElement(parent, f'{{{NS["p"]}}}cTn', id='1', nodeType='clickEffect')
        children = etree.SubElement(clock, f'{{{NS["p"]}}}childTnLst')
        behavior = etree.SubElement(children, f'{{{NS["p"]}}}animEffect', transition='in', filter='fade')
        common = etree.SubElement(behavior, f'{{{NS["p"]}}}cBhvr')
        etree.SubElement(common, f'{{{NS["p"]}}}cTn', id='2', dur='500')
        target = etree.SubElement(common, f'{{{NS["p"]}}}tgtEl')
        etree.SubElement(target, f'{{{NS["p"]}}}spTgt', spid=str(body.shape_id))
    elif case == 'dense':
        body.text = '\n'.join(['Năng lượng của vật được bảo toàn.'] * 18)
    elif case == 'bilingual':
        body.text += '\nThe energy of the object is conserved.'
    elif case == 'group':
        body.text = ''
        group = slide.shapes.add_group_shape()
        text(group.shapes, 'Năng lượng của vật được bảo toàn.', 1, 2, 5, 2)
        text(group.shapes, 'Vận tốc của vật tăng khi thế năng giảm.', 7, 2, 5, 2)
    deck.save(path)


def issues(report):
    return [{'slide': row['slide'], **issue} for row in report['slides'] for issue in row['issues']]


def verify_case(folder, case, native):
    from app.native_layout import fit_preserved_frames
    from app.office_tasks import render_quality
    from app.slide_repair import repair_and_inspect

    started = time.monotonic()
    source, target = folder / 'source.pptx', folder / 'bilingual.pptx'
    fixture(source, case)
    before = digest(source)
    manifest = extract(source, 'parallel_columns', preserve_design=True, same_slide=True)
    translations = {}
    for unit in manifest['units']:
        value = unit['text']
        for vi, en in PAIRS.items():
            value = value.replace(vi, en)
        translations[unit['id']] = {'en': value}
    assemble(source, target, manifest, translations, 'parallel_columns')
    fitting = fit_preserved_frames(target, manifest) if native else {'available': None, 'changes': []}
    original = inspect_returned_deck(source)
    inspection = inspect_returned_deck(target, source, 'parallel_columns', 'source',
                                       layout_plan=manifest.get('layout_plan'))
    if native:
        render_quality(source, folder, original['slide_quality'])
        render_quality(target, folder, inspection['slide_quality'])
        checked = repair_and_inspect(target, folder, source=source, mode='parallel_columns',
                                     layout_plan=manifest.get('layout_plan'))
        target, final = Path(checked['path']), checked['inspection']
    else:
        final = inspection
    assert digest(source) == before
    assert source_media_preserved(source, target, 'source')
    # Formula run bytes must survive, not just its flattened text.
    if case == 'formula':
        result = Presentation(target)
        assert any(run._r.rPr is not None and run._r.rPr.get('baseline') == '30000'
                   for slide in result.slides for shape in leaves(slide.shapes) if shape.has_text_frame
                   for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
    if case == 'animation':
        with zipfile.ZipFile(source) as old, zipfile.ZipFile(target) as new:
            old_timing = etree.fromstring(old.read('ppt/slides/slide1.xml')).find('p:timing', NS)
            new_timing = etree.fromstring(new.read('ppt/slides/slide1.xml')).find('p:timing', NS)
            # Same-slide delivery adds matching English effects. Every original
            # effect and its clock/target must survive unchanged.
            effects = [etree.tostring(effect) for effect in new_timing.findall('.//p:animEffect', NS)]
            assert all(etree.tostring(effect) in effects for effect in old_timing.findall('.//p:animEffect', NS))
    for name, data in (('manifest', manifest), ('source-inspection', original),
                       ('before-repair', inspection), ('after-repair', final)):
        (folder / (name + '.json')).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return dict(case=case, source_sha256=before, output_sha256=digest(target),
                source_unchanged=True, media_preserved=True, native=native, fit=fitting,
                source_issues=issues(original['slide_quality']), output_issues=issues(final['slide_quality']),
                counts=final['slide_quality']['counts'], output=str(target),
                rendered=sum(bool(row['rendered']) for row in final['slide_quality']['slides']),
                seconds=round(time.monotonic() - started, 3))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=Path('reports/rc73/quality-corpus'))
    parser.add_argument('--native', action='store_true')
    parser.add_argument('--cases', choices=CASES, nargs='+', default=CASES)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    rows = []
    for case in args.cases:
        folder = args.output / case
        folder.mkdir(exist_ok=True)
        rows.append(verify_case(folder, case, args.native))
        print(json.dumps(rows[-1], ensure_ascii=True), flush=True)
    report = dict(status='structural_checks_passed', cases=rows, network_calls=0,
                  scope='Owned synthetic fixtures with fixed translations; inspect issues and PNGs. '
                        'Not a translation, pronunciation or animation-playback certification.')
    (args.output / 'result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
