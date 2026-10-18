"""Real Qt narration workflow with disposable slides and deterministic WAV fixtures."""

import json
import logging
import os
import sys
import tempfile
import wave
import zipfile
from pathlib import Path
from threading import Event
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import click, control, pump, until

from app import speech
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    out = Path('reports/rc71/workspace')
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='biliclass-voice-workspace-', ignore_cleanup_errors=True) as folder:
        folder = Path(folder)
        os.environ['BILICLASS_DATA'] = str(folder / 'data')
        library = Library(folder / 'data')
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(str(v) for v in values))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT / 'qml/Main.qml')
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]

        def c(name):
            return control(window, name)

        def render(text, voice, rate, directory, cancel):
            target = Path(directory) / (speech.audio_key(text, voice, rate) + '.wav')
            target.parent.mkdir(parents=True, exist_ok=True)
            with wave.open(str(target), 'wb') as audio:
                audio.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                audio.writeframes(b'\1\0' * 2400)
            return {'path': str(target), 'cached': False}

        try:
            window.resize(1366, 850)
            window.setProperty('page', 'voice')
            pump(150)
            assert window.property('narrationMode') == 'en'
            assert c('narrationEnglish').property('checked')
            assert c('narrationBoth').property('text') == 'Việt + Anh'
            placement = c('narrationSpeakerPosition')
            assert placement.property('currentIndex') == 2
            placement.setProperty('currentIndex', 0)
            QMetaObject.invokeMethod(placement, 'activated', Qt.DirectConnection, Q_ARG('int', 0))
            until(lambda: library.setting('narration_speaker_position', '') == 'bottom_left')
            click(window, c('narrationVietnamese'))
            until(lambda: library.setting('narration_last_mode', '') == 'vi')
            click(window, c('narrationBoth'))
            until(lambda: library.setting('narration_last_mode', '') == 'both')
            icon = c('narrationNavigationIcon')
            assert icon.width() == icon.height() == 29

            def blue_pixels():
                image = window.grabWindow()
                ratio = image.devicePixelRatio()
                position = (icon.mapToScene(QPointF()) * ratio).toPoint()
                return sum(color.blue() - color.green() > 70 and color.blue() - color.red() > 100
                           for x in range(position.x(), position.x() + round(29 * ratio))
                           for y in range(position.y(), position.y() + round(29 * ratio))
                           for color in [image.pixelColor(x, y)])

            assert blue_pixels() > 20, 'The selected supplied icon must be blue'
            click(window, c('navigationhome'))
            pump(150)
            assert blue_pixels() < 3, 'The unselected supplied icon must retain its gray color'
            click(window, c('navigationvoice'))
            for width, height in ((1080, 700), (1366, 850)):
                window.resize(width, height)
                c('resultScroll').property('contentItem').setProperty('contentY', 0)
                pump(150)
                for name in ('lessonNarrationVoiceChoice', 'lessonNarrationRateChoice', 'narrationEnglish', 'narrationVietnamese', 'narrationBoth'):
                    item = c(name)
                    assert item.isVisible() and item.isEnabled(), name
                    assert 0 <= item.mapToScene(QPointF()).y() < window.height() - 25, name
                language = c('narrationEnglish')
                placement = c('narrationSpeakerPosition')
                assert abs(language.mapToScene(QPointF(0, language.height()/2)).y()
                           - placement.mapToScene(QPointF(0, placement.height()/2)).y()) < 2
                assert placement.mapToScene(QPointF()).x() > c('narrationBoth').mapToScene(
                    QPointF(c('narrationBoth').width(), 0)).x()
                assert not c('createLessonNarrationButton').isEnabled()
                assert window.grabWindow().save(str(out / f'configure-before-file-{width}.png'))
            voice_choice = c('lessonNarrationVoiceChoice')
            voice_choice.setProperty('currentIndex', 2)
            QMetaObject.invokeMethod(voice_choice, 'activated', Qt.DirectConnection, Q_ARG('int', 2))
            rate_choice = c('lessonNarrationRateChoice')
            rate_choice.setProperty('currentIndex', 4)
            QMetaObject.invokeMethod(rate_choice, 'activated', Qt.DirectConnection, Q_ARG('int', 4))
            click(window, c('narrationBoth'))
            selected_voice = window.property('narrationVoice')
            assert selected_voice == 'vieneu:Hải Đăng' and window.property('narrationRate') == 1
            path = folder / 'Bài giảng ảnh.pptx'
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6])
            deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(path)
            original = path.read_bytes()
            bridge.importNarrationPresentation(path.as_uri())
            until(lambda: not bridge.busy and bool(bridge.quickResult.get('path')))
            pump(250)
            assert window.property('narrationVoice') == selected_voice and window.property('narrationRate') == 1
            assert window.grabWindow().save(str(out / 'imported.png'))
            assert not c('narrationEnglish').isEnabled() and not c('narrationVietnamese').isEnabled()
            assert not c('createLessonNarrationButton').isEnabled()
            assert c('narrationScriptNotice').isVisible()
            assert 'slide thiếu lời' in c('narrationCoverageMessage').property('text')
            assert c('fixMissingNarration').isVisible()
            click(window, c('resultMoreButton'))
            assert c('resultMoreMenu').property('count') == 3
            assert window.grabWindow().save(str(out / 'more-menu.png'))
            click(window, c('reviewNarrationScript'))
            until(lambda: c('narrationScriptDialog').property('visible'))
            pump(150)
            c('narrationScriptVietnamese').setProperty('text', 'Năng lượng được bảo toàn.')
            c('narrationScriptEnglish').setProperty('text', 'Energy is conserved.')
            assert not c('narrationScriptSlide').isEnabled(), 'Unsaved script cannot change slides'
            assert window.grabWindow().save(str(out / 'script-1366.png'))
            click(window, c('saveNarrationScript'))
            assert bridge.lesson['segments'][0]['en'] == 'Energy is conserved.'
            assert not c('narrationScriptDialog').property('dirty')
            click(window, c('discardNarrationScript'))
            click(window, c('narrationBoth'))
            assert c('narrationBoth').property('checked')
            started, calls = Event(), []

            def slow(text, voice, rate, directory, cancel):
                calls.append(text)
                started.set()
                cancel.wait(5)
                return render(text, voice, rate, directory, cancel)

            with patch('app.speech.narration_workers', return_value=1), patch('app.speech.synthesize_background', side_effect=slow):
                click(window, c('createLessonNarrationButton'))
                until(started.is_set)
                assert c('cancelLessonNarrationButton').property('danger')
                assert c('cancelLessonNarrationButton').property('symbol') == '■'
                assert bridge.narrationBusy
                assert not c('changeVoiceSource').isEnabled() and not c('downloadPresentationButton').isEnabled()
                previous = bridge.lesson['id']
                bridge.narrationWorkspace.selectLibrary('different-id')
                assert bridge.lesson['id'] == previous and bridge.error
                click(window, c('cancelLessonNarrationButton'))
                until(lambda: not bridge.narrationBusy)
            assert bridge.lessonNarration['complete'] == 1 and not bridge.lessonNarration['ready']
            assert c('createLessonNarrationButton').property('text') == 'Tạo giọng đọc'

            with patch('app.speech.narration_workers', return_value=1), patch('app.speech.synthesize_background', side_effect=render):
                click(window, c('createLessonNarrationButton'))
                until(lambda: not bridge.narrationBusy and bridge.lessonNarration.get('ready'))
            assert bridge.narrationPreparation['cached'] == 1
            # Placement is presentation metadata: cached voice and revision stay valid.
            before_position = bridge.lesson['revision']
            before_voice = dict(bridge.lesson['narration'])
            bridge.narrationWorkspace.setSpeakerPosition('bottom_center')
            until(lambda: not bridge.busy and bridge.lesson.get('speaker_position') == 'bottom_center')
            assert bridge.lesson['revision'] == before_position
            assert bridge.lesson['narration'] == before_voice and bridge.lessonNarration['ready']
            assert c('createLessonNarrationButton').isEnabled()
            assert not window.findChild(type(c('createLessonNarrationButton')), 'voiceNarrationVariantChoice')
            assert 'hai ngôn ngữ' in c('narrationConfigurationExplanation').property('text')
            click(window, c('createLessonNarrationButton'))
            assert c('existingNarrationDialog').property('visible') and not bridge.narrationBusy
            click(window, c('useExistingNarration'))
            assert not c('existingNarrationDialog').property('visible') and not bridge.narrationBusy
            with patch('app.speech.narration_workers', return_value=1), patch('app.speech.synthesize_background', side_effect=render):
                click(window, c('createLessonNarrationButton'))
                click(window, c('regenerateExistingNarration'))
                until(lambda: not bridge.narrationBusy and bridge.lessonNarration.get('ready'))
                assert bridge.narrationPreparation['cached'] == 0
            assert c('downloadNarrationAudioButton').isVisible()
            assert not c('voiceDownloadWithoutNarration').isVisible()
            assert not c('voiceTeachWithoutNarration').isVisible()
            assert c('continueAfterNarrationButton').isVisible()
            assert c('downloadAfterNarrationButton').isVisible()
            assert c('continueAfterNarrationButton').property('text') == 'Mở PowerPoint'
            assert 'lưu vào Bài giảng' in c('narrationCompletionTitle').property('text')
            assert not c('narrationCreationOptions').isVisible()
            click(window, c('resultMoreButton'))
            assert QMetaObject.invokeMethod(c('resultAdvancedMenu'), 'open')
            pump(100)
            assert c('voiceDownloadWithoutNarrationMenu').isVisible()
            assert QMetaObject.invokeMethod(c('resultAdvancedMenu'), 'close')
            assert QMetaObject.invokeMethod(c('resultMoreMenu'), 'close')
            assert '1 slide thiếu lời Việt' in c('narrationCoverageMessage').property('text')
            assert not any(s['approved'] for s in bridge.lesson['segments'])
            for width, height in ((1080, 700), (1366, 850)):
                window.resize(width, height)
                c('resultScroll').property('contentItem').setProperty('contentY', 0)
                pump(150)
                action = c('downloadAfterNarrationButton')
                viewport = c('resultScroll').property('contentItem')
                overflow = action.mapToScene(QPointF(0, action.height())).y() - c('creationFooter').y() + 16
                if overflow > 0:
                    viewport.setProperty('contentY', overflow)
                    pump(100)
                create = c('createLessonNarrationButton')
                bottom = create.mapToScene(QPointF(0, create.height())).y()
                assert bottom <= window.height() - 30, (width, height, bottom)
                assert window.grabWindow().save(str(out / f'ready-{width}.png'))
            clicked_audio = []
            with patch.object(bridge.english_reader, 'start', side_effect=lambda *a, **k: clicked_audio.append((a, k))):
                c('checkCompletedNarrationButton').clicked.emit()
                until(lambda: not bridge.previewBusy)
                assert window.property('voicePreviewExpanded')
                assert not clicked_audio, 'Opening a preview never starts playback'
                c('listenCreatedEnglish').clicked.emit()
                assert len(clicked_audio) == 1 and clicked_audio[0][1] == {'cache_only': True}
            c('toggleVoiceSlidePreview').clicked.emit()
            assert not window.property('voicePreviewExpanded')
            c('viewCompletedNarrationInLibrary').clicked.emit()
            pump(100)
            assert window.property('page') == 'library'
            assert bridge.lesson['id'] in window.property('selectedLessons').toVariant()
            window.setProperty('page', 'voice')
            c('resultScroll').property('contentItem').setProperty('contentY', 0)
            pump(100)
            window.setProperty('includeDownloadVoice', False)
            click(window, c('downloadAfterNarrationButton'))
            assert window.property('includeDownloadVoice'), 'The voice download automatically includes the chosen voice'
            assert c('presentationSaveDialog').property('visible')
            assert QMetaObject.invokeMethod(c('presentationSaveDialog'), 'reject')
            audio_target = folder / 'Audio.zip'
            bridge.narrationWorkspace.exportAudio(audio_target.as_uri())
            until(lambda: not bridge.busy)
            assert not bridge.error, bridge.message
            with zipfile.ZipFile(audio_target) as archive:
                assert 'Slide-001-both.wav' in archive.namelist()
            ppt_target = folder / 'Có giọng.pptx'
            bridge.downloadPresentationWithVoice(ppt_target.as_uri(), True)
            until(lambda: not bridge.busy)
            assert not bridge.error, bridge.message
            with zipfile.ZipFile(ppt_target) as archive:
                assert any(name.endswith('.wav') for name in archive.namelist())
            assert path.read_bytes() == original
            assert not bridge.powerpointState.get('active')

            # The library route must prepare a preview even for an old text lesson.
            lesson = library.create('Bài thư viện', 'Tin học', 'THPT', '10', [('Slide 1', 'Máy tính xử lý thông tin.')])
            lesson = library.edit_segment(lesson['id'], lesson['segments'][0]['id'], 'Máy tính xử lý thông tin.', 'A computer processes information.')
            bridge.narrationWorkspace.selectLibrary(lesson['id'])
            until(lambda: not bridge.busy and bridge.lesson.get('id') == lesson['id'])
            assert bridge.quickResult['lesson_id'] == lesson['id'] and bridge.quickResult['path']
            assert window.property('page') == 'voice'
            assert c('lessonNarrationPanel').isVisible()
            mono = folder / 'Chỉ Việt.pptx'
            single = Presentation()
            single.slides.add_slide(single.slide_layouts[6]).notes_slide.notes_text_frame.text = 'VI: Xin chào các em.'
            single.save(mono)
            bridge.importNarrationPresentation(mono.as_uri())
            until(lambda: not bridge.busy and bridge.lesson.get('title') == 'Chỉ Việt')
            assert window.property('narrationMode') == 'vi'
            assert window.property('narrationVoice') == selected_voice and window.property('narrationRate') == 1
            assert not c('narrationBoth').isEnabled()
            assert not c('lessonNarrationState').isVisible()
            assert c('narrationEnglish').property('text') == 'Tiếng Anh'
            saved_id = bridge.lesson['id']
            with patch('app.speech.narration_workers', return_value=1), patch('app.speech.synthesize_background', side_effect=slow):
                started.clear()
                click(window, c('createLessonNarrationButton'))
                until(started.is_set)
                assert c('dismissVoiceLesson').property('danger')
                assert c('dismissVoiceLesson').property('symbol') == '×'
                window.setProperty('page', 'new')
                pump(100)
                assert c('creationSourceArea').property('workingTitle') == 'Đang tạo giọng đọc'
                assert c('chooseInputDocumentButton').property('text') == 'Xem tiến độ tạo giọng'
                assert not c('createLessonButton').isEnabled()
                assert not c('resumeCreationTask').isVisible()
                assert c('conversionConfigurationSummary').isVisible()
                c('lessonTitle').setProperty('text', 'Bài mới trong khi tạo giọng')
                c('lessonSubject').setProperty('currentIndex', bridge.conversionSubjects.index('Tin học'))
                window.setProperty('selectedFile', mono.as_uri())
                window.setProperty('selectedFileName', mono.name)
                pump(30)
                preview = window.property('creationPrompt')
                assert (preview.toVariant() if hasattr(preview, 'toVariant') else preview)['ready']
                assert not c('createLessonButton').isEnabled()
                bridge.convertDocument('Bài mới', 'Tin học', 'THPT', '11', '', mono.as_uri(),
                                       'parallel_columns', 'standard', 'source')
                assert bridge.narrationBusy and bridge.lesson['id'] == saved_id
                assert 'Chờ tạo giọng' in bridge.message
                assert 'Giữ thiết kế gốc' in c('conversionConfigurationTitle').property('text')
                c('creationFormat').setProperty('currentIndex', 1)
                pump(30)
                assert 'tiếng Anh' in c('conversionConfigurationExplanation').property('text')
                click(window, c('chooseInputDocumentButton'))
                assert window.property('page') == 'voice'
                assert bridge.narrationBusy
                click(window, c('dismissVoiceLesson'))
                until(lambda: not bridge.narrationBusy and not bridge.quickResult.get('path'))
            assert library.get(saved_id)['narration_variants'], 'Cancel selection retains voice progress'
            assert window.property('page') == 'voice'
            assert not c('dismissVoiceLesson').isVisible()
            assert c('voiceSourceArea').isVisible()
            for name, label in (('narrationEnglish', 'Tiếng Anh'), ('narrationVietnamese', 'Tiếng Việt'),
                                ('narrationBoth', 'Việt + Anh')):
                assert c(name).property('text') == label
            bridge.narrationWorkspace.selectLibrary(saved_id)
            until(lambda: not bridge.busy and bool(bridge.quickResult.get('path')))
            converted = library.get(saved_id)
            converted['voice_only'] = False
            library._write(converted)
            bridge.narrationWorkspace.selectLibrary(saved_id)
            until(lambda: not bridge.busy)
            click(window, c('resultMoreButton'))
            click(window, c('reviewNarrationScript'))
            until(lambda: c('narrationScriptDialog').property('visible'))
            assert not c('narrationScriptEnglish').property('readOnly')
            c('narrationScriptEnglish').setProperty('text', 'Hello, students.')
            click(window, c('saveNarrationScript'))
            assert bridge.narrationInputLanguages['en']
            assert library.get(saved_id)['segments'][0]['en'] == '', 'Narration edit keeps converted slide text intact'
            assert library.get(saved_id)['narration_script'][converted['segments'][0]['id']]['en'] == 'Hello, students.'
            click(window, c('discardNarrationScript'))
            click(window, c('dismissVoiceLesson'))
            assert not bridge.lesson and library.get(saved_id)
            assert not warnings, warnings
            result = {'status': 'passed', 'qml_warnings': warnings,
                      'checks': ['supplied narration icon is 29px with blue/gray selection states',
                                 'voice/rate/three modes visible and selectable before file at two sizes', 'configuration survives file/library choices and mono-language adjustment',
                                 'blank/image slide script editing without source changes', 'dirty script blocks slide switch',
                                 'language availability', 'cancel and resume reuse cached WAV', 'busy source/download guards',
                                 'duplicate configuration offers reuse or explicit regeneration', 'ZIP WAV and embedded PowerPoint export',
                                 'regular library lesson prepares matching preview', 'no auto review or playback'],
                      'new_checks': ['consistent unselected and missing-language status', 'cancel selected lesson while idle or synthesizing preserves library and WAVs', 'edit converted lesson narration directly without changing slide text'],
                      'scope': 'Real Qt and PowerPoint export; synthesized audio uses WAV fixtures'}
            (out / 'result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result), flush=True)
        finally:
            window.close()
            for worker in (bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog, bridge.browserAI,
                           bridge.english_reader, bridge.narration_preparation, bridge.classroom, bridge.voice_controls):
                worker.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()
            logging.shutdown()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
