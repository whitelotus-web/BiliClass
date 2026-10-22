"""Owned Qt input routing and ID-based multiple library selection; no teacher data."""
import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, Qt, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlExpression
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.voice_import import import_presentation
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT/'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/narration-picker')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-picker-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        library = Library(folder/'data')
        sources, lessons = [], []
        for n in range(3):
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = f'Lesson {n+1}: Energy is conserved.'
            slide.notes_slide.notes_text_frame.text = f'EN: Lesson {n+1}. Energy is conserved.'
            source = folder/f'Bài thử {n+1}.pptx'
            deck.save(source)
            sources.append(source)
            lesson, preview = import_presentation(library.directory, source)
            lessons.append((lesson, preview))
        originals = [source.read_bytes() for source in sources]
        bridge = Bridge(library)
        bridge.previewConvertedSlide = lambda *_: None
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT/'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        root.resize(1366, 850)
        root.show()

        def c(name):
            return control(root, name)

        def qml(source):
            expression = QQmlExpression(engine.contextForObject(root), root, source)
            value, undefined = expression.evaluate()
            assert not expression.hasError() and not undefined, expression.error().toString()
            return value

        def invoke(obj, method, values):
            assert QMetaObject.invokeMethod(obj, method, Qt.DirectConnection, Q_ARG('QVariant', values))
            pump(80)

        try:
            until(lambda: len(bridge.narrationSources) == 3 and bridge._catalog.worker is None)
            qml('powerPointInputOpen.options = FileDialog.DontUseNativeDialog')
            root.setProperty('page', 'voice')
            pump(120)
            invoke(c('homeWelcome'), 'documentDropped', [source.as_uri() for source in sources])
            assert not bridge.batchQueue.data and root.property('page') == 'voice', 'Hidden Home cannot receive Voice inputs'
            # Real upload button opens the one chooser with an explicit purpose.
            click(root, c('chooseVoicePresentation'))
            dialog = c('powerPointInputFileDialog')
            assert dialog.property('visible') and dialog.property('purpose') == 'voice'
            assert qml('powerPointInputOpen.fileMode === FileDialog.OpenFiles')
            QMetaObject.invokeMethod(dialog, 'reject', Qt.DirectConnection)
            pump(80)
            # Exercise the accepted URL receiver; Windows native multi-selection is not automated.
            root.setProperty('page', 'new')
            invoke(dialog, 'receiveFiles', [source.as_uri() for source in sources])
            assert bridge.batchQueue.state['kind'] == 'voice' and root.property('page') == 'voice'
            assert c('voiceBatchPanel').isVisible() and not c('conversionBatchPanel').isVisible()
            assert c('createLessonNarrationButton').isEnabled()
            assert not bridge.busy and not bridge.batchQueue.running
            bridge.batchQueue.clear()
            pump(80)
            imported = []
            assert dialog.setProperty('selectedFile', QUrl.fromLocalFile(str(sources[0])))
            with patch.object(bridge, 'importNarrationPresentation', side_effect=imported.append):
                QMetaObject.invokeMethod(dialog, 'accepted', Qt.DirectConnection)
            assert len(imported) == 1 and Path(QUrl(imported[0]).toLocalFile()) == sources[0], (imported, dialog.property('selectedFiles'), 'Single-file accepted event keeps the Voice route')
            root.setProperty('page', 'new')
            click(root, c('chooseInputDocumentButton'))
            assert dialog.property('purpose') == 'convert' and dialog.property('visible')
            QMetaObject.invokeMethod(dialog, 'reject', Qt.DirectConnection)
            pump(80)
            root.setProperty('page', 'voice')
            invoke(dialog, 'receiveFiles', [source.as_uri() for source in sources])
            assert bridge.batchQueue.state['kind'] == 'convert' and root.property('page') == 'new'
            assert c('conversionBatchPanel').isVisible() and not c('voiceBatchPanel').isVisible()
            bridge.batchQueue.clear()
            pump(60)
            root.setProperty('page', 'voice')
            picker = c('voiceLibraryPicker')
            assert not c('selectVoiceLibraryLesson').isEnabled()
            click(root, c('voiceLibraryPickerButton'))
            ids = [lesson['id'] for lesson, _ in lessons]
            for identity in ids[:2]:
                click(root, c('voiceLibraryChoice_'+identity))
            assert qml('voiceLibraryPicker.selectedCount') == 2
            assert c('voiceLibraryPickerPopup').property('visible'), 'Clicking a row keeps multi-picker open'
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                pump(100)
                root.grabWindow().save(str(report/f'picker-{theme}.png'))
            click(root, c('voiceLibraryPickerDone'))
            click(root, c('selectVoiceLibraryLesson'))
            until(lambda: not bridge.busy and bridge.batchQueue.state.get('kind') == 'voice')
            assert [row['lesson_id'] for row in bridge.batchQueue.data['rows']] == ids[:2]
            assert bridge.batchQueue.state['phase'] == 'draft' and not bridge.batchQueue.running
            assert len(library.list_lessons(verify_audio=False)) == 3
            bridge.batchQueue.clear()
            until(lambda: not bridge.busy)
            # Selection survives reorder, and removes a vanished ID rather than selecting its replacement.
            model = list(reversed(bridge.narrationSources))
            picker.setProperty('model', model)
            pump(60)
            assert set(json.loads(qml('JSON.stringify(voiceLibraryPicker.selectedIds)'))) == set(ids[:2])
            picker.setProperty('model', [row for row in model if row['id'] != ids[0]])
            pump(60)
            assert json.loads(qml('JSON.stringify(voiceLibraryPicker.selectedIds)')) == [ids[1]]
            # One selected lesson follows the existing one-lesson workspace.
            with patch('app.voice_import.prepare_library_lesson', return_value=lessons[1]) as prepared:
                click(root, c('selectVoiceLibraryLesson'))
                until(lambda: not bridge.busy and bridge.lesson.get('id') == ids[1])
            assert prepared.call_args.args[1] == ids[1] and not bridge.batchQueue.data
            picker.setProperty('model', [{'id': str(n), 'title': f'Bài {n}'} for n in range(105)])
            root.setProperty('voiceSourceExpanded', True)
            pump(80)
            click(root, c('voiceLibraryPickerButton'))
            click(root, c('voiceLibrarySelectAll'))
            assert qml('voiceLibraryPicker.selectedCount') == 100
            click(root, c('voiceLibraryClearSelection'))
            assert qml('voiceLibraryPicker.selectedCount') == 0
            click(root, c('voiceLibraryPickerDone'))
            root.setProperty('page', 'library')
            pump(100)
            for theme in ('light', 'dark'):
                bridge.setApplicationTheme(theme)
                for key in ('all', 'none', 'ready'):
                    click(root, c('libraryTab_'+key))
                    assert c('libraryTab_'+key).property('selected')
                    assert c('libraryTabIcon_'+key).isVisible()
                    root.grabWindow().save(str(report/f'library-{key}-{theme}.png'))
            assert [source.read_bytes() for source in sources] == originals
            assert not warnings, warnings
            (report/'result.json').write_text(json.dumps(dict(status='passed', explicit_input_routing=True,
                native_multiselect_not_automated=True, multiple_library_choices_stay_open=True,
                selected_ids_survive_reorder=True, removed_ids_pruned=True, one_lesson_uses_existing_workspace=True,
                multiple_lessons_draft_only=True, sources_preserved=True, library_tabs_icons_and_filters=True,
                qml_warnings=[])), encoding='utf-8')
            print('Narration multi-picker and input routing Qt smoke passed.', flush=True)
        finally:
            until(lambda: not bridge.busy)
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview',
                         '_catalog', 'browserAI', 'english_reader', 'narration_preparation',
                         'narrationPreferences', '_companion_preferences', 'classroom', 'voice_controls', 'inputDocuments'):
                getattr(bridge, name).shutdown()
            c('companionWindow').setProperty('disposing', True)
            root.hide()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
