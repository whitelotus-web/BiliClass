"""Disposable Qt preview/final-product flow; WAV fixtures, no API or playback."""

import json
import os
import sys
import tempfile
import traceback
import wave
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QMetaObject, QPointF, Slot
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from qt_result_preview_mascot_smoke import control, pump, until

from app.chatgpt_handoff import external_preview
from app.lesson_narration import make_plan
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.voice_import import import_presentation


class PreviewBridge(Bridge):
    opened = []

    @Slot(str)
    def openTeachingLibraryLesson(self, identity):
        self.opened.append(identity)


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    out = Path('reports/rc88/lesson-final')
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='biliclass-final-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        os.environ['BILICLASS_DATA'] = str(folder/'data')
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = 'Xin chào.\nHello.'
        slide.notes_slide.notes_text_frame.text = 'VI: Xin chào.\nEN: Hello.'
        source = folder/'Bài giảng.pptx'
        deck.save(source)
        lesson, preview = import_presentation(folder/'data', source)
        library = Library(folder/'data')
        bridge = PreviewBridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(str(value) for value in values))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT/'qml/Main.qml')
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        def c(name):
            return control(window, name)
        try:
            bridge._show_lesson(lesson, preview, auto_repair=False, load_preview=False)
            until(lambda: bridge._catalog.selected_narration.worker is None)
            for width, height in ((1080, 700), (1366, 850), (1680, 980)):
                window.resize(width, height)
                pump(120)
                assert c('lessonInformation').isVisible()
                assert c('lessonNoVoiceInformation').isVisible()
                assert not c('lessonNarrationPanel').isVisible() and not c('creationFooter').isVisible()
                assert c('resultSlideNavigation').isVisible()
                for name in ('useConvertedLessonButton', 'downloadPresentationButton', 'resultCreateNarrationButton'):
                    action = c(name)
                    point = action.mapToScene(QPointF(action.width(), action.height()))
                    assert action.isVisible() and point.x() <= width-12 and point.y() < height-20, (name, point)
                assert window.grabWindow().save(str(out/f'preview-{width}.png'))
            c('useConvertedLessonButton').clicked.emit()
            assert bridge.opened == [lesson['id']]
            c('finalPresentationSaveDialog').setProperty('options', 8)
            c('downloadPresentationButton').clicked.emit()
            until(lambda: c('finalPresentationSaveDialog').property('visible'))
            QMetaObject.invokeMethod(c('finalPresentationSaveDialog'), 'close')
            c('resultCreateNarrationButton').clicked.emit()
            until(lambda: window.property('page') == 'voice')
            assert c('narrationEnglish').isVisible() and c('narrationVietnamese').isVisible()

            for mode in ('en', 'both'):
                plan, _ = make_plan(lesson, mode, {'voice': 'vieneu:Hải Đăng', 'rate': 0},
                                    playback='both_controls' if mode == 'both' else 'separate')
                for keys in plan['keys'].values():
                    for key in keys:
                        target = library.directory/'audio'/f'{key}.wav'
                        target.parent.mkdir(exist_ok=True)
                        with wave.open(str(target), 'wb') as audio:
                            audio.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                            audio.writeframes(b'\1\0'*2400)
                lesson = library.update_narration(lesson['id'], plan)
                bridge._show_lesson(lesson, external_preview(lesson, library.directory), auto_repair=False,
                                    load_preview=False, destination='voice')
                until(lambda: bridge.narrationAvailableModes(bridge.narrationVariants).get(mode))
                pump(150)
                assert not c('narrationEnglish').isVisible()
                if mode == 'en':
                    assert c('narrationVietnamese').isVisible() and c('narrationBoth').isVisible()
                    assert window.property('narrationMode') == 'vi'
                else:
                    assert not c('narrationCreationOptions').isVisible()
                    assert not c('createLessonNarrationButton').isEnabled()
                bridge.openLesson(lesson['id'])
                until(lambda: not bridge.busy and window.property('page') == 'result')
                assert not c('lessonNarrationPanel').isVisible() and c('lessonInformation').isVisible()
                assert c('lessonVoiceInformation_'+bridge.selectedNarrationVariant).isVisible()
                assert not c('lessonNoVoiceInformation').isVisible()
                window.grabWindow().save(str(out/f'preview-{mode}.png'))
                c('resultCreateNarrationButton').clicked.emit()
                until(lambda: window.property('page') == 'voice')
            assert not warnings, warnings
            result = {'status': 'passed', 'qml_warnings': warnings, 'sizes': [1080, 1366, 1680],
                      'checks': ['preview has details and three final actions; no voice configuration/footer',
                                 'open dispatches to Mascot teaching entry point', 'download uses final product dialog',
                                 'voice button navigates to dedicated tab; only missing languages visible',
                                 'bilingual hides all three generation choices', 'preview opens from voice tab correctly'],
                      'scope': 'real Qt, disposable lesson and WAV fixtures; no teacher data/API/model/playback'}
            (out/'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
            print(json.dumps(result, ensure_ascii=False), flush=True)
        except BaseException:
            (out/'failure.txt').write_text(traceback.format_exc(), encoding='utf-8')
            raise
        finally:
            window.close()
            for worker in (bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog, bridge.browserAI,
                           bridge.english_reader, bridge.narration_preparation, bridge.classroom, bridge.voice_controls,
                           bridge._companion_preferences):
                worker.shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
