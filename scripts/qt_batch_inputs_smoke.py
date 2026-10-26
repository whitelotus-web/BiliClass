"""Owned queue selection/preview race and layout checks; no teacher documents."""
import json
import sys
import tempfile
from pathlib import Path
from threading import Event
from unittest.mock import patch

from pptx import Presentation
from PySide6.QtCore import Q_ARG, QCoreApplication, QEvent, QMetaObject, QPointF, Qt, QUrl
from PySide6.QtGui import QColor, QFont, QFontDatabase, QGuiApplication, QImage
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app.branding import configure_application
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.voice_import import import_presentation
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT/'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    report = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/batch-inputs')
    report.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-input-queue-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        sources, pictures = [], []
        for n, color in enumerate(('#c64b43', '#3776c5', '#309765', '#b2792a')):
            deck = Presentation()
            for _ in range(n+2):
                deck.slides.add_slide(deck.slide_layouts[6])
            path = folder/f'Bài {n+1} có tên dài để thử hàng đợi, #1.pptx'
            deck.save(path)
            sources.append(path)
            picture = QImage(960, 540, QImage.Format_RGB32)
            picture.fill(QColor(color))
            image = folder/f'preview-{n}.png'
            assert picture.save(str(image))
            pictures.append(image.as_uri())
        original = [source.read_bytes() for source in sources]
        library = Library(folder/'data')
        library.set_setting('teaching_subjects', ['Toán'])
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT/'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        root = engine.rootObjects()[0]
        root.resize(1366, 850)
        root.show()
        entered, release = Event(), Event()
        delayed = [False]

        def renderer(url, directory, cancel, number=1):
            n = sources.index(Path(QUrl(url).toLocalFile()))
            if delayed[0] and n == 1:
                entered.set()
                assert release.wait(8)
                # Deliberately return a stale image despite cancellation.
            return dict(image=pictures[n], total=n+2, slide=number)

        def c(name):
            return control(root, name)

        def accept(method):
            assert QMetaObject.invokeMethod(root, method, Qt.DirectConnection,
                Q_ARG('QVariant', [path.as_uri() for path in sources]))

        inputs = bridge.batchQueue.inputs
        batch_geometry, single_geometry, empty_geometry = {}, {}, {}

        def geometry(names):
            return {key: (c(name).mapToScene(QPointF()).x(), c(name).mapToScene(QPointF()).y(),
                          c(name).width(), c(name).height()) for key, name in names.items()}

        def same_geometry(first, second, scope):
            # Conversion has PowerPoint/Word tabs; Voice starts directly with
            # its guide. Compare the workspace relative to that shared row.
            vertical_offset = second['guide'][1] - first['guide'][1]
            for key in first:
                expected = tuple(value + (vertical_offset if index == 1 and key != 'heading' else 0)
                    for index, value in enumerate(first[key]))
                assert all(abs(a-b) <= 1 for a, b in zip(expected, second[key])), (scope, key, expected, second[key])

        def check_input_actions(kind, source_panel):
            prefix = 'creation' if kind == 'convert' else 'voice'
            choose = c('chooseInputDocumentButton' if kind == 'convert' else 'changeVoiceSource')
            cancel = c('removeInputDocumentButton' if kind == 'convert' else 'dismissVoiceLesson')
            heading, guide = c(prefix + 'PageHeading'), c(prefix + 'WorkflowGuide')
            bounds = [item.mapRectToScene(item.boundingRect()) for item in (choose, cancel, heading, guide, c(source_panel))]
            change_rect, cancel_rect, heading_rect, guide_rect, source_rect = bounds
            assert choose.isVisible() and cancel.isVisible()
            assert guide_rect.right() <= change_rect.left() < cancel_rect.left() and change_rect.right() <= cancel_rect.left()
            assert abs(change_rect.center().y() - guide_rect.center().y()) <= 1
            assert abs(cancel_rect.center().y() - guide_rect.center().y()) <= 1
            assert heading_rect.bottom() < change_rect.top() < cancel_rect.bottom() < source_rect.top()
            assert cancel_rect.right() <= root.width(), bounds

        def resize(width, height):
            root.resize(width, height)
            pump(100)
            if root.width() != width or root.height() != height:
                # Windows can finish applying the initial native frame after the first resize.
                root.resize(width, height)
                pump(100)
            assert (root.width(), root.height()) == (width, height)

        try:
            # Shared upload controls align before a single file or queue exists.
            for kind, page, prefix, upload, button in (
                ('convert', 'new', 'creation', 'creationSourceArea', 'chooseInputDocumentButton'),
                ('voice', 'voice', 'voice', 'voiceSourceArea', 'chooseVoicePresentation'),
            ):
                root.setProperty('page', page)
                for theme in ('light', 'dark'):
                    bridge.setApplicationTheme(theme)
                    for width, height in ((1080, 700), (1366, 850)):
                        resize(width, height)
                        measured = geometry(dict(heading=prefix+'PageHeading', guide=prefix+'WorkflowGuide',
                            upload=upload, button=button, illustration=button+'Illustration'))
                        assert c(upload).isVisible() and c(upload).property('multiple')
                        assert c(button).isEnabled()
                        if kind == 'convert':
                            empty_geometry[theme, width] = measured
                        else:
                            same_geometry(empty_geometry[theme, width], measured, ('empty', theme, width))
                        root.grabWindow().save(str(report/f'empty-{kind}-{theme}-{width}.png'))
            with patch('app.input_documents.input_preview', renderer):
                for kind, method, source_panel, config in (
                    ('convert', 'acceptInputDocuments', 'conversionBatchInputs', 'creationConfiguration'),
                    ('voice', 'acceptVoiceDocuments', 'voiceBatchInputs', 'lessonNarrationPanel'),
                ):
                    root.setProperty('page', 'new' if kind == 'convert' else 'voice')
                    accept(method)
                    until(lambda: len(inputs.details) == 4 and not inputs.documents.preview.get('busy', True))
                    ids = [row['id'] for row in bridge.batchQueue.data['rows']]
                    assert inputs.state['selected_id'] == ids[0]
                    assert inputs.documents.preview['image'] == pictures[0]
                    assert [row['total'] for row in inputs.state['rows']] == [2, 3, 4, 5]
                    assert c(source_panel).isVisible()
                    assert c(source_panel).mapToScene(QPointF()).y() < c(config).mapToScene(QPointF()).y()
                    cancel = c('removeInputDocumentButton' if kind == 'convert' else 'dismissVoiceLesson')
                    assert cancel.isVisible() and cancel.isEnabled()
                    for theme in ('light', 'dark'):
                        bridge.setApplicationTheme(theme)
                        for width, height in ((1080, 700), (1366, 850)):
                            resize(width, height)
                            check_input_actions(kind, source_panel)
                            list_view, thumbnail = c('batchInputRows_'+kind), c('batchInputThumbnail_'+kind)
                            assert thumbnail.mapToScene(QPointF()).x() >= list_view.mapToScene(QPointF(list_view.width(), 0)).x()
                            assert abs(thumbnail.width()/thumbnail.height()-16/9) < .01
                            nav = c('batchInputThumbnailNavigation_'+kind)
                            assert nav.isVisible() and nav.mapToScene(QPointF()).y() >= thumbnail.mapToScene(QPointF(0, thumbnail.height())).y()
                            assert abs(nav.mapToScene(QPointF(nav.width()/2, 0)).x()-thumbnail.mapToScene(QPointF(thumbnail.width()/2, 0)).x()) <= 1
                            assert c('batchInputPreviewToggle_'+kind).property('subtle')
                            hint = c('batchInputHint_'+kind)
                            assert hint.mapToScene(QPointF()).y() >= list_view.mapToScene(QPointF(0, list_view.height())).y()
                            measured = geometry(dict(
                                guide='creationWorkflowGuide' if kind == 'convert' else 'voiceWorkflowGuide',
                                thumbnail='batchInputThumbnail_'+kind, navigation='batchInputThumbnailNavigation_'+kind,
                                rows='batchInputRows_'+kind, hint='batchInputHint_'+kind,
                                configuration=config,
                            ))
                            measured['configuration'] = measured['configuration'][:2]
                            if kind == 'convert':
                                batch_geometry[theme, width] = measured
                            else:
                                same_geometry(batch_geometry[theme, width], measured, ('batch', theme, width))
                            root.grabWindow().save(str(report/f'{kind}-{theme}-{width}.png'))
                    assert QMetaObject.invokeMethod(list_view, 'positionViewAtEnd', Qt.DirectConnection)
                    pump(80)
                    click(root, c('batchInputChoice_'+ids[3]))
                    until(lambda: inputs.documents.preview.get('image') == pictures[3] and not inputs.documents.preview.get('busy'))
                    delayed[0] = True
                    entered.clear()
                    release.clear()
                    click(root, c('batchInputChoice_'+ids[1]))
                    until(entered.is_set)
                    click(root, c('batchInputChoice_'+ids[2]))
                    assert inputs.state['selected_id'] == ids[2]
                    release.set()
                    until(lambda: inputs.documents.preview.get('image') == pictures[2] and not inputs.documents.preview.get('busy'))
                    delayed[0] = False
                    click(root, c('batchInputNext_'+kind))
                    until(lambda: inputs.documents.preview.get('slide') == 2 and not inputs.documents.preview.get('busy'))
                    assert inputs.documents.preview['image'] == pictures[2]
                    click(root, c('batchInputPreviewToggle_'+kind))
                    assert c('batchInputExpanded_'+kind).isVisible()
                    assert c('batchInputExpandedNavigation_'+kind).isVisible()
                    click(root, c('batchInputPreviewToggle_'+kind))
                    assert QMetaObject.invokeMethod(list_view, 'positionViewAtBeginning', Qt.DirectConnection)
                    pump(80)
                    opened = []
                    with patch.object(bridge, 'openInputPowerPoint', side_effect=opened.append):
                        click(root, c('batchInputOpen_'+ids[0]))
                    assert [Path(QUrl(url).toLocalFile()) for url in opened] == [sources[0]], opened
                    assert inputs.state['selected_id'] == ids[2], inputs.state
                    click(root, c('batchInputRemove_'+ids[2]))
                    until(lambda: inputs.state['selected_id'] == ids[0] and inputs.documents.preview.get('image') == pictures[0])
                    assert [row['id'] for row in bridge.batchQueue.data['rows']] == [*ids[:2], ids[3]]
                    click(root, cancel)
                    pump(100)
                    assert not bridge.batchQueue.data and not inputs.state['selected_id'] and not inputs.documents.preview
                    assert not bridge.busy and not bridge.narrationBusy
                assert library.db.execute('SELECT COUNT(*) FROM lessons').fetchone()[0] == 0
                # Single-file conversion can change the thumbnail before expanding it.
                root.setProperty('page', 'new')
                assert QMetaObject.invokeMethod(root, 'acceptInputDocuments', Qt.DirectConnection,
                    Q_ARG('QVariant', [sources[0].as_uri()]))
                until(lambda: not bridge.inputDocuments.preview.get('busy', True) and not bridge.inputDocuments.busy)
                thumbnail, nav = c('creationSelectedSlidePreview'), c('creationThumbnailNavigation')
                assert thumbnail.isVisible() and nav.isVisible() and not root.property('creationPreviewExpanded')
                assert nav.mapToScene(QPointF()).y() >= thumbnail.mapToScene(QPointF(0, thumbnail.height())).y()
                click(root, c('creationNextSlide'))
                until(lambda: bridge.inputDocuments.preview.get('slide') == 2 and not bridge.inputDocuments.preview.get('busy'))
                assert not root.property('creationPreviewExpanded')
                for theme in ('light', 'dark'):
                    bridge.setApplicationTheme(theme)
                    for width, height in ((1080, 700), (1366, 850)):
                        resize(width, height)
                        check_input_actions('convert', 'creationLessonOverview')
                        single_geometry[theme, width] = geometry(dict(
                            guide='creationWorkflowGuide', title='creationSelectedLessonTitle', info='creationSelectedLessonInfo',
                            state='creationInputState', summary='creationSelectedSlideCount', open='creationOpenPowerPoint',
                            thumbnail='creationSelectedSlidePreview', navigation='creationThumbnailNavigation',
                            preview='creationPreviewButton', configuration='creationConfiguration', configTitle='conversionConfigurationTitle',
                        ))
                        root.grabWindow().save(str(report/f'single-conversion-{theme}-{width}.png'))
                QMetaObject.invokeMethod(root, 'removeInputDocument', Qt.DirectConnection)
                # Existing one-lesson Voice preview uses the same navigation placement.
                lesson, preview = import_presentation(library.directory, sources[0])
                with patch('app.slide_preview.PreviewSession.render', return_value=QUrl(pictures[0]).toLocalFile()):
                    bridge._show_lesson(lesson, preview, destination='voice', auto_repair=False, load_preview=False)
                    bridge.previewConvertedSlide(1)
                    until(lambda: not bridge.previewBusy and bool(bridge.quickResult.get('image')))
                    root.setProperty('voicePreviewExpanded', False)
                    pump(100)
                    thumbnail, nav = c('voiceSlideThumbnail'), c('voiceThumbnailNavigation')
                    assert nav.isVisible() and thumbnail.isVisible()
                    assert nav.mapToScene(QPointF()).y() >= thumbnail.mapToScene(QPointF(0, thumbnail.height())).y()
                    click(root, c('voiceThumbnailNextSlide'))
                    until(lambda: not bridge.previewBusy and bridge.quickResult.get('slide') == 2)
                    assert not root.property('voicePreviewExpanded')
                    for theme in ('light', 'dark'):
                        bridge.setApplicationTheme(theme)
                        for width, height in ((1080, 700), (1366, 850)):
                            resize(width, height)
                            check_input_actions('voice', 'voiceLessonOverview')
                            measured = geometry(dict(
                                guide='voiceWorkflowGuide', title='resultLessonTitle', info='voiceSelectedLessonInfo',
                                state='lessonWorkflowState', summary='voiceSourceSummary', open='voiceTeachWithoutNarration',
                                thumbnail='voiceSlideThumbnail', navigation='voiceThumbnailNavigation',
                                preview='toggleVoiceSlidePreview', configuration='lessonNarrationPanel', configTitle='narrationConfigurationTitle',
                            ))
                            # Content-specific text and form heights may differ; common edges and baselines must align.
                            first = {key: value[:2] for key, value in single_geometry[theme, width].items()}
                            second = {key: value[:2] for key, value in measured.items()}
                            same_geometry(first, second, ('single', theme, width))
                            root.grabWindow().save(str(report/f'single-voice-{theme}-{width}.png'))
                assert [source.read_bytes() for source in sources] == original
                assert not warnings, warnings
            (report/'result.json').write_text(json.dumps(dict(status='passed', shared_convert_voice_input_panel=True,
                first_file_preview=True, latest_selection_wins=True, remove_selected_falls_back=True,
                per_file_open_keeps_selection=True, header_cancel_clears_only_inputs=True,
                slide_navigation_and_expand=True, metadata_real_pptx=True, sources_and_library_preserved=True,
                navigation_under_thumbnail_single_and_multiple=True, scroll_to_fourth_input=True,
                matched_empty_single_and_batch_geometry_between_tools=True, queue_hint_below_rows=True,
                input_actions_beside_guide_above_selected_source=True,
                light_dark_1080_1366=True, qml_warnings=[])), encoding='utf-8')
            print('Batch input selection, preview race and layout Qt checks passed.', flush=True)
        finally:
            release.set()
            bridge.batchQueue.shutdown()
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview',
                         '_catalog', 'browserAI', 'english_reader', 'narration_preparation',
                         'narrationPreferences', '_companion_preferences', 'classroom', 'voice_controls', 'inputDocuments'):
                getattr(bridge, name).shutdown()
            c('companionWindow').setProperty('disposing', True)
            root.hide()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
