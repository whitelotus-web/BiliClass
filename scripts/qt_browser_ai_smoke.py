"""Free/Plus account UI and PPTX receive flow; mocked web, real Office preview."""

import hashlib
import json
import sys
import tempfile
import time
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QMetaObject, QObject, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, QQmlEngine, QQmlExpression
from PySide6.QtQuick import QQuickWindow
from PySide6.QtQuickControls2 import QQuickStyle

from app import browser_automation, speech
from app.chatgpt_handoff import load_request
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import RESOURCES, Bridge

Path(".runtime/browser-web-smoke").mkdir(parents=True, exist_ok=True)
workspace = Path(tempfile.mkdtemp(dir=".runtime/browser-web-smoke")).resolve()
reports = Path("reports/browser-web")
reports.mkdir(parents=True, exist_ok=True)
source = workspace / "Nguồn thử.pptx"
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2)).text = "Diện tích · Area"
slide.notes_slide.notes_text_frame.text = "VI: Diện tích bằng 12 cm².\nEN: The area is 12 cm²."
deck.save(source)
attempts, sends, audio, warnings, failures, stages = [], [], [], [], [], []
conversions = []
free_health_failure = False


def fake_login(account, root, cancel, progress, **options):
    attempts.append(account["id"])
    if len(attempts) == 1:
        return {"ready": False}
    if len(attempts) in {2, 4}:
        raise browser_automation.BrowserProblem("verification", browser_automation.VERIFICATION_MESSAGE)
    if len(attempts) == 6:
        assert options["url"] == "https://chatgpt.com/c/fixture" and account["id"] == free_id
    if len(attempts) in {3, 6}:
        options["awaiting_confirmation"](True)
        deadline = time.monotonic() + 15
        while not options["finish"].is_set():
            assert time.monotonic() < deadline, "UI did not confirm the plain Chrome sign-in"
            if cancel.wait(.05):
                raise browser_automation.BrowserProblem("cancelled", "Đã hủy đăng nhập.")
        assert not cancel.is_set(), "Confirmation must verify, not cancel"
        options["awaiting_confirmation"](False)
    plan = "free" if len(attempts) in {3, 6} else "plus"
    return {"ready": True, "plan": plan, "name": "Cô giáo thử " + plan, "email": plan + "@example.test"}


def fake_convert(account, root, folder, cancel, progress, **options):
    request = load_request(folder)
    assert request["config"]["provider"] == "browser_web"
    assert "FILE POWERPOINT" in request["prompt"] and "VI:" in request["prompt"]
    conversions.append(account["id"])
    if account["plan"] == "plus":
        assert len(conversions) == 1 and not sends
        browser_automation.write_record(folder, {"account_id": account["id"], "state": "prepared", "followups": 0,
            "url": "", "error": "limit"})
        raise browser_automation.BrowserProblem("limit", "ChatGPT báo hết lượt dùng.")
    if len(conversions) == 2:
        sends.append(account["id"])
        browser_automation.write_record(folder, {"account_id": account["id"], "state": "waiting", "followups": 0,
            "url": "https://chatgpt.com/c/fixture", "error": "verification"})
        raise browser_automation.BrowserProblem("verification", browser_automation.VERIFICATION_MESSAGE)
    assert len(sends) == 1 and browser_automation.read_record(folder, account["id"])["state"] == "waiting"
    if len(conversions) == 3:
        record = browser_automation.read_record(folder, account["id"])
        record["error"] = "browser"
        browser_automation.write_record(folder, record)
        raise browser_automation.BrowserProblem("browser", "Kết nối tạm gián đoạn.")
    options["observed"](account["id"], "free", "Web model fixture")
    target = Path(folder) / "bai-giang-song-ngu.pptx"
    target.write_bytes(source.read_bytes())
    browser_automation.write_record(folder, {"account_id": account["id"], "state": "completed", "followups": 0,
        "url": "https://chatgpt.com/c/fixture", "sha256": hashlib.sha256(target.read_bytes()).hexdigest()})
    return {"path": str(target), "url": "https://chatgpt.com/c/fixture"}


def fake_check(account, root, cancel, progress):
    if free_health_failure and account["id"] == free_id:
        raise browser_automation.BrowserProblem("login", "Phiên hết hạn. Đăng nhập lại.")
    return {"ready": True, "plan": account["plan"], "name": "Cô giáo thử " + account["plan"]}


browser_automation.login = fake_login
browser_automation.convert = fake_convert
browser_automation.check_session = fake_check
speech.synthesize = lambda text, *_: audio.append(text)
application = QGuiApplication([])
for font in (RESOURCES / "assets").glob("*.ttf"):
    QFontDatabase.addApplicationFontFromData(font.read_bytes())
application.setFont(QFont("Be Vietnam Pro", 10))
QQuickStyle.setStyle("Basic")
library = Library(workspace / "library")
bridge = Bridge(library)
engine = QQmlApplicationEngine()
engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
engine.rootContext().setContextProperty("bridge", bridge)
engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
assert engine.rootObjects(), warnings
window = engine.rootObjects()[0]
window.setProperty("page", "settings")
window.findChild(QObject, "settingsPage").setProperty("activeTab", 6)
phase, deadline, free_id, plus_id, lesson_id = "start", time.monotonic() + 180, "", "", ""


def click(name):
    item = window.findChild(QObject, name)
    assert item and item.property("enabled"), name
    assert QMetaObject.invokeMethod(item, "clicked", Qt.DirectConnection)


def profile_value(account_id, control, *, activate=False):
    # Keep dynamic delegates inside the QML engine. PySide can take incorrect
    # ownership when returning these internal layout objects to Python.
    repeater = window.findChild(QObject, "browserProfiles")
    operation = 'if (!target.enabled) throw new Error("Button disabled"); target.clicked(); return true;' if activate else 'return target.text;'
    code = '''(function() {
        function find(node, name) {
            if (node.objectName === name) return node;
            for (var j = 0; j < node.children.length; j++) {
                var found = find(node.children[j], name); if (found) return found;
            }
            return null;
        }
        for (var i = 0; i < count; i++) {
            var row = itemAt(i);
            if (row.accountId === ACCOUNT) {
                var target = find(row, CONTROL);
                if (!target) throw new Error("Profile control missing");
                OPERATION
            }
        }
        throw new Error("Profile row missing");
    })()'''.replace("ACCOUNT", json.dumps(account_id)).replace("CONTROL", json.dumps(control)).replace("OPERATION", operation)
    expression = QQmlExpression(QQmlEngine.contextForObject(repeater), repeater, code)
    value = expression.evaluate()[0]
    assert not expression.hasError(), expression.error().toString()
    return value


def click_profile(account_id, action):
    assert profile_value(account_id, action, activate=True)


def capture(name):
    assert QQuickWindow.grabWindow(window).save(str(reports / name))


def step():
    global phase, free_id, plus_id, lesson_id, free_health_failure
    try:
        assert time.monotonic() < deadline, "Web UI flow timed out"
        if phase == "start":
            QMetaObject.invokeMethod(window.findChild(QObject, "helpDialog"), "close", Qt.DirectConnection)
            assert window.findChild(QObject, "browserAccountAdd").property("text") == "Đăng nhập ChatGPT"
            capture("login.png")
            click("browserAccountAdd")
            phase = "closed"
        elif phase == "closed" and not bridge.browserAI.loginBusy:
            assert not bridge.browserAI.accountInfo["ready"]
            assert not window.findChild(QObject, "browserAccountDetails").property("visible")
            assert profile_value(bridge.browserAI.activeId, "browserProfileState") == "Chưa đăng nhập"
            click_profile(bridge.browserAI.activeId, "browserProfileReconnect")
            phase = "verification"
        elif phase == "verification" and not bridge.browserAI.loginBusy:
            assert bridge.browserAI.hasError and bridge.browserAI.verificationBlocked
            assert not bridge.browserAI.accountInfo["ready"]
            assert window.findChild(QObject, "browserManualRecovery").property("visible")
            assert "Cloudflare" in window.findChild(QObject, "browserAccountStatus").property("text")
            assert bridge.browserAI.store.get()["last_error"]["code"] == "verification"
            capture("verification-blocked.png")
            # The fallback is an explicit choice and must not claim a connection.
            click("browserManualRecovery")
            assert not bridge.browserAI.automatic and not bridge.browserAI.accountInfo["ready"]
            assert window.property("page") == "new"
            bridge.browserAI.saveOptions(True, True)
            window.setProperty("page", "settings")
            window.findChild(QObject, "settingsPage").setProperty("activeTab", 6)
            stages.append("Verification error persisted; explicit manual fallback does not mark an account ready")
            click_profile(bridge.browserAI.activeId, "browserProfileReconnect")
            phase = "confirm_free"
        elif phase == "confirm_free" and bridge.browserAI.loginAwaitingConfirmation:
            assert not bridge.browserAI.accountInfo["ready"]
            assert window.findChild(QObject, "browserAccountAdd").property("text") == "Kiểm tra và lưu"
            capture("login-confirm.png")
            click("browserAccountAdd")
            assert not bridge.browserAI.loginAwaitingConfirmation
            phase = "free"
        elif phase == "free" and not bridge.browserAI.loginBusy:
            assert bridge.browserAI.accountInfo["ready"] and bridge.browserAI.accountInfo["plan"] == "free"
            assert len(bridge.browserAI.accounts) == 1 and attempts[0] == attempts[1] == attempts[2]
            assert not bridge.browserAI.verificationBlocked and "last_error" not in bridge.browserAI.store.get()
            free_id = bridge.browserAI.activeId
            assert window.findChild(QObject, "browserAccountAdd").property("text") == "Thêm tài khoản"
            assert not window.findChild(QObject, "browserQuotaStatus").property("visible")
            assert not window.findChild(QObject, "browserSelectionPolicy").property("visible")
            assert not window.findChild(QObject, "browserAccountMenu").property("visible")
            assert profile_value(free_id, "browserProfilePlan") == "Free"
            assert profile_value(free_id, "browserProfileDelete") == "Xóa profile"
            capture("free-connected.png")
            stages.append("One primary button confirms plain Chrome sign-in; closing early is not success; verified Free session saved")
            click("browserAccountAdd")
            phase = "plus_failed"
        elif phase == "plus_failed" and not bridge.browserAI.loginBusy:
            assert len(bridge.browserAI.accounts) == 2 and not bridge.browserAI.accountInfo["ready"]
            assert bridge.browserAI.activeId != free_id and bridge.browserAI.store.preferred()["id"] == free_id
            assert window.findChild(QObject, "browserAccountAdd").property("text") == "Thêm tài khoản"
            click_profile(bridge.browserAI.activeId, "browserProfileReconnect")
            phase = "plus"
        elif phase == "plus" and not bridge.browserAI.loginBusy:
            assert len(bridge.browserAI.accounts) == 2 and bridge.browserAI.accountInfo["plan"] == "plus"
            assert attempts[3] == attempts[4], "Retry must reopen the failed new account, not add/reopen a different profile"
            plus_id = bridge.browserAI.activeId
            assert free_id != plus_id
            capture("plus-preferred.png")
            assert bridge.browserAI.store.preferred()["id"] == plus_id
            assert window.findChild(QObject, "browserProfiles").property("count") == 2
            stages.append("Separate account rows, visible profile deletion; no policy menu or unknown quota row")
            bridge.convertDocument("Bài thử", "Toán", "THPT", "11", "", QUrl.fromLocalFile(str(source)).toString(),
                                   "parallel_columns", "standard", "source")
            phase = "blocked_conversion"
        elif phase == "blocked_conversion" and not bridge.busy:
            assert bridge.error and bridge.browserState["needsLogin"]
            assert not bridge.browserAI.store.get(free_id)["ready"]
            assert bridge.browserAI.store.get(plus_id)["ready"] and bridge.browserAI.store.get(plus_id)["quota_limited"]
            assert sends == [free_id], "Limited Plus must switch to Free before sending exactly once"
            assert window.findChild(QObject, "browserConversionResume").property("text") == "Đăng nhập và tiếp tục"
            reopened = Bridge(library)
            assert reopened.browserState["needsLogin"] and reopened.chatgptRequest["folder"] == bridge.chatgptRequest["folder"]
            reopened.browserAI.shutdown()
            reopened.deleteLater()
            capture("conversion-reconnect.png")
            click("browserConversionResume")
            stages.append("Failed new account retries its own profile; conversion verification clears stale ready state and reconnects its pinned account")
            phase = "confirm_resume"
        elif phase == "confirm_resume" and bridge.browserAI.loginAwaitingConfirmation:
            window.setProperty("page", "settings")
            window.findChild(QObject, "settingsPage").setProperty("activeTab", 6)
            assert len(conversions) == 2 and not bridge.browserAI.accountInfo["ready"]
            assert window.findChild(QObject, "browserAccountAdd").property("text") == "Kiểm tra và lưu"
            click("browserAccountAdd")
            phase = "converted"
        elif phase == "converted" and not bridge.busy and not bridge.browserAI.loginBusy and len(conversions) == 4:
            assert not bridge.error, bridge.message
            assert sends == [free_id]
            assert conversions[-2:] == [free_id, free_id], "Transport recovery must not need another user click or resend"
            stages.append("Temporary transport failure recovered automatically on the same conversation, with no duplicate send")
            assert bridge.lesson["external_deck"] and bridge.quickResult["image"]
            assert Path(bridge.quickResult["path"]).read_bytes() == source.read_bytes()
            assert not any(segment["approved"] for segment in bridge.lesson["segments"])
            if bridge.voiceSettings["en"]:
                assert audio, "Requested narration must be prepared"
            assert window.property("page") == "result"
            lesson_id = bridge.lesson["id"]
            capture("result.png")
            assert bridge.browserAI.store.preferred()["id"] == free_id
            assert bridge.browserAI.conversionAccount(bridge.chatgptRequest["folder"])["id"] == free_id
            bridge.browserAI.checkAccount(plus_id)
            phase = "quota_recovered"
        elif phase == "quota_recovered" and not bridge.browserAI.loginBusy:
            assert bridge.browserAI.store.preferred()["id"] == plus_id
            assert bridge.browserAI.conversionAccount(bridge.chatgptRequest["folder"])["id"] == free_id
            free_health_failure = True
            bridge.browserAI.checkAccount(free_id)
            phase = "health_lost"
        elif phase == "health_lost" and not bridge.browserAI.loginBusy:
            assert not bridge.browserAI.store.get(free_id)["ready"]
            assert bridge.browserAI.store.preferred()["id"] == plus_id
            assert profile_value(free_id, "browserProfileState") == "Mất kết nối"
            assert profile_value(free_id, "browserProfileReconnect") == "Đăng nhập lại"
            window.setProperty("page", "settings")
            window.findChild(QObject, "settingsPage").setProperty("activeTab", 6)
            phase = "delete_profiles"
        elif phase == "delete_profiles" and not bridge.browserAI.loginBusy:
            capture("accounts-simple.png")
            click_profile(free_id, "browserProfileDelete")
            assert len(bridge.browserAI.accounts) == 1 and bridge.browserAI.accounts[0]["id"] == plus_id
            click_profile(plus_id, "browserProfileDelete")
            stages.append("Plus quota fallback sends once through Free; resumed job stays on Free; health probe warns of lost session; each visible delete removes only its profile")
            bridge.runBrowserAI()
            phase = "cached"
        elif phase == "cached" and not bridge.busy:
            assert not bridge.error and bridge.lesson["id"] == lesson_id
            assert len(sends) == 1 and not bridge.browserAI.accounts
            stages.append("PPTX/narration received; Office preview rendered; teacher approval kept; cached lesson opens without accounts/resend")
            application.quit()
            return
    except Exception:
        failures.append(traceback.format_exc())
        print(failures[-1], flush=True)
        application.exit(1)
        return
    QTimer.singleShot(100, step)


QTimer.singleShot(250, step)
code = application.exec()
# Destroy QML while its Python bridge is still alive, including queued job
# cleanup; otherwise interpreter teardown can evaluate properties too late.
engine.deleteLater()
QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
bridge.browserAI.shutdown()
library.close()
result = {"status": "passed" if not failures and not warnings else "failed", "stages": stages,
          "warnings": warnings, "failures": failures, "scope": "Fixture web login/conversion; real Qt/Office preview; no real ChatGPT request"}
(reports / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False), flush=True)
sys.exit(code or bool(failures) or bool(warnings))
