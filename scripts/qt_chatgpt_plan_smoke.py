"""Source Qt flow with controlled AI, actual Office render, isolated teacher data."""

import hashlib
import json
import sys
import tempfile
import time
import traceback
from pathlib import Path
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))

import httpx
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QMetaObject, QObject, Qt, QTimer, QUrl
from PySide6.QtGui import QFont, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuick import QQuickWindow
from PySide6.QtQuickControls2 import QQuickStyle
from test_ai_lesson import EchoProvider

from app import chatgpt_plan, speech
from app.browser_ai_ui import BrowserAI
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.quick_conversion import verify_preview
from app.ui import Bridge


class FixtureProvider(EchoProvider):
    def __init__(self, _accounts):
        super().__init__()
        self.client = httpx.Client()


chatgpt_plan.ChatGPTPlanProvider = FixtureProvider
narration = []


def synthetic_audio(text, voice, rate, directory):
    narration.append((text, voice, rate))
    return {"path": str(directory / "fixture.wav")}


speech.synthesize = synthetic_audio
application = QGuiApplication([])
QQuickStyle.setStyle("Basic")
application.setFont(QFont("Arial", 10))
reports = Path("reports/chatgpt-plan")
reports.mkdir(parents=True, exist_ok=True)
Path(".runtime/chatgpt-plan-smoke").mkdir(parents=True, exist_ok=True)
workspace = Path(tempfile.mkdtemp(prefix="qt-", dir=".runtime/chatgpt-plan-smoke")).resolve()
library = Library(workspace / "library")
source = workspace / "Bài gốc.pptx"
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(1.5)).text = "Có một giá trị."
deck.save(source)
digest = hashlib.sha256(source.read_bytes()).hexdigest()
bridge = Bridge(library, browser_ai_factory=BrowserAI)
account_id = str(uuid4())
bridge.browserAI.store.data.update(accounts=[{"id": account_id, "label": "Tài khoản thử", "status": "Kết nối thử",
                                             "plan": "unknown", "ready": True}], active=account_id, audio=True)
bridge.browserAI.changed.emit()
engine = QQmlApplicationEngine()
warnings, failures, stages = [], [], []
engine.warnings.connect(lambda values: warnings.extend(map(str, values)))
engine.rootContext().setContextProperty("bridge", bridge)
engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / "qml/Main.qml")))
assert engine.rootObjects(), warnings
window = engine.rootObjects()[0]
deadline = time.monotonic() + 240


def safe(action):
    def run():
        try:
            action()
        except Exception:
            failures.append(traceback.format_exc())
            print(failures[-1], flush=True)
            application.exit(1)
    return run


def click(name):
    target = window.findChild(QObject, name)
    assert target and target.property("enabled"), name
    assert QMetaObject.invokeMethod(target, "clicked", Qt.DirectConnection)


def wait_idle(action):
    assert time.monotonic() < deadline, "AI flow timed out"
    if bridge.busy:
        QTimer.singleShot(100, safe(lambda: wait_idle(action)))
    else:
        action()


def start():
    QMetaObject.invokeMethod(window.findChild(QObject, "helpDialog"), "close", Qt.DirectConnection)
    window.setProperty("page", "settings")
    window.findChild(QObject, "settingsPage").setProperty("activeTab", 6)
    assert window.findChild(QObject, "browserAccountAdd").property("text") == "Đăng nhập lại"
    assert QQuickWindow.grabWindow(window).save(str(reports / "accounts.png"))
    window.setProperty("page", "new")
    window.setProperty("selectedFile", QUrl.fromLocalFile(str(source)).toString())
    window.setProperty("selectedFileName", source.name)
    window.findChild(QObject, "lessonTitle").setProperty("text", "Bài từ ChatGPT")
    window.findChild(QObject, "lessonSubject").setProperty("text", "Toán")
    window.findChild(QObject, "conversionProvider").setProperty("currentIndex", 0)
    click("createLessonButton")
    QTimer.singleShot(200, safe(lambda: wait_idle(converted)))


def converted():
    assert not bridge.error, bridge.message
    assert window.property("page") == "result"
    result = bridge.quickResult
    assert result["path"] and result["image"] and result["total"] >= 1, result
    assert bridge.lesson["ai_conversion"]["provider"] == "chatgpt_plan"
    assert bridge._reading_text("en") == "One value."
    if bridge.voiceSettings["en"]:
        assert any(item[0] == "One value." for item in narration), narration
    assert not any(segment["approved"] for segment in bridge.lesson["segments"])
    assert not any(item["approved"] for segment in bridge.lesson["segments"] for item in segment["support"])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert verify_preview(bridge.lesson, result, library.directory).is_file()
    assert QQuickWindow.grabWindow(window).save(str(reports / "result.png"))
    stages.append("One click -> source manifest -> controlled AI JSON -> native PowerPoint -> actual Office preview")
    print(stages[-1], flush=True)
    opened = []
    bridge._start_powerpoint_session = lambda *args: opened.append(args)
    click("useConvertedLessonButton")
    assert window.findChild(QObject, "wholeLessonReview").property("visible")
    click("confirmWholeLessonButton")
    assert opened and all(segment["approved"] for segment in bridge.lesson["segments"])
    assert all(item["approved"] for segment in bridge.lesson["segments"] for item in segment["support"])
    stages.append("Whole lesson confirmation -> reviewed AI support -> native slideshow/mascot handoff")
    print(stages[-1], flush=True)
    bridge.browserAI.store.data["accounts"] = []
    bridge.browserAI.store.data["active"] = ""
    bridge.runBrowserAI()
    QTimer.singleShot(100, safe(lambda: wait_idle(reopened)))


def reopened():
    assert not bridge.error, bridge.message
    assert len(library.list_lessons()) == 1
    assert bridge.quickResult["path"] and bridge.quickResult["image"]
    assert all(segment["approved"] for segment in bridge.lesson["segments"])
    stages.append("Resume completed request without a connected account -> same lesson, no new AI call")
    print(stages[-1], flush=True)
    assert not warnings, warnings
    application.exit(0)


QTimer.singleShot(700, safe(start))
code = application.exec()
if bridge.worker:
    bridge.cancel_event.set()
    bridge.worker.wait()
bridge.browserAI.shutdown()
report = {"status": "passed" if code == 0 and not warnings and not failures else "failed", "stages": stages,
          "warnings": warnings, "failures": failures, "library": str(workspace),
          "scope": "Qt + real Office; AI/audio fixtures, no live OAuth/quota/teacher document sent; final slideshow intercepted"}
(reports / "qt-flow.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
del engine
library.close()
print(json.dumps(report, ensure_ascii=False), flush=True)
raise SystemExit(code)
