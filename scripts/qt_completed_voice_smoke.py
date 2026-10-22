"""Completed narration, automatic slide preview, and compact actions in real Qt.

Uses a disposable library, real Office preview and cached fixture WAVs. Does not
send global input or open/close any teacher presentation.
"""
import json
import sys
import tempfile
import wave
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent, QObject, QPointF, Qt, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from app.branding import configure_application
from app.chatgpt_handoff import external_preview
from app.lesson_narration import make_plan, variant_id
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.voice_import import import_presentation
from scripts.qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    configure_application(app)
    QQuickStyle.setStyle('Basic')
    for font in (RESOURCE_ROOT / 'assets').glob('*.ttf'):
        QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Be Vietnam Pro', 10))
    reports = Path(sys.argv[1] if len(sys.argv) > 1 else 'reports/completed-voice')
    reports.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bcl-completed-voice-', ignore_cleanup_errors=True) as folder:
        folder = Path(folder)
        source = folder / 'Bài 4 – Độ dịch chuyển.pptx'
        presentation = Presentation()
        for n in (1, 2, 3):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = (
                f'Bài 4: Độ dịch chuyển\nLesson 4: Displacement\nSlide {n}')
            slide.notes_slide.notes_text_frame.text = f'VI: Bài 4. Slide {n}.\nEN: Lesson 4. Slide {n}.'
        presentation.save(source)
        lesson, _ = import_presentation(folder / 'data', source)
        library = Library(folder / 'data')
        plan, _ = make_plan(lesson, 'both', {'voice': 'pair:truc-ly', 'rate': 0})
        library.set_setting('voice_options', {'voice': plan['voice'], 'rate': plan['rate']})
        for keys in plan['keys'].values():
            for key in keys:
                target = library.directory / 'audio' / (key + '.wav')
                target.parent.mkdir(exist_ok=True)
                with wave.open(str(target), 'wb') as audio:
                    audio.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                    audio.writeframes(b'\1\0' * 2400)
        plan.update(lesson_id=lesson['id'], revision=lesson['revision'], complete=plan['total'], running=False)
        lesson = library.update_narration(lesson['id'], plan)
        lesson = library.select_narration_variant(lesson['id'], variant_id(plan))
        preview = external_preview(lesson, library.directory)
        bridge = Bridge(library)
        initial_powerpoint_state = dict(bridge.powerpointState)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(QUrl.fromLocalFile(str(RESOURCE_ROOT / 'qml/Main.qml')))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]

        def c(name):
            return control(window, name)

        def check_footer_alignment(page, recent_link=False):
            window.setProperty('page', page)
            pump(80)
            summary = c('narrationConfiguration' if page == 'voice' else 'conversionConfigurationSummary')
            button = c('createLessonNarrationButton' if page == 'voice' else 'createLessonButton')
            link = c('resumeNarrationTask' if page == 'voice' else 'resumeCreationTask')
            assert summary.isVisible() and button.isVisible()
            assert abs(summary.mapToScene(QPointF(0, summary.height()/2)).y()
                       - button.mapToScene(QPointF(0, button.height()/2)).y()) < 2
            assert summary.mapToScene(QPointF(summary.width(), 0)).x() + 12 <= button.mapToScene(QPointF()).x()
            assert link.isVisible() == recent_link
            if recent_link:
                assert link.mapToScene(QPointF()).y() >= summary.mapToScene(QPointF(0, summary.height())).y() + 2
                footer = c('creationFooter')
                assert link.mapToItem(footer, QPointF(0, link.height())).y() <= footer.height() - 10

        try:
            window.show()
            for page in ('new', 'voice'):
                window.setProperty('page', page)
                pump(60)
                assert not c('resumeNarrationTask').isVisible()
                assert not c('resumeCreationTask').isVisible()
                check_footer_alignment(page)
            bridge._show_lesson(lesson, preview, destination='voice', auto_repair=False, load_preview=False)
            until(lambda: window.property('narrationComplete'), timeout=30)
            until(lambda: bool(bridge.quickResult.get('image')) and not bridge.previewBusy, timeout=90)
            assert not c('narrationResultPage').isVisible(), 'Opening Voice config must not resurrect an old result'
            window.setProperty('voiceResultDismissed', False)
            bridge.narrationChanged.emit()
            pump(120)
            assert window.property('voicePreviewExpanded')
            assert c('narrationResultSlideImage').property('source').toString()
            assert c('narrationResultPage').isVisible()
            assert not c('resultScroll').isVisible()
            assert window.findChild(QObject, 'checkCompletedNarrationButton') is None
            assert 'Kiểm tra slide và giọng' not in (RESOURCE_ROOT / 'qml/NarrationReadyCard.qml').read_text(encoding='utf-8')
            sizes = ((1080, 700), (1366, 850), (1920, 1080))
            def check_listening_alignment(english=True, vietnamese=True):
                # Exercise the existing visibility cases without creating audio jobs.
                c('narrationResultListenEnglish').setVisible(english)
                c('narrationResultListenVietnamese').setVisible(vietnamese)
                pump(80)
                listening = c('narrationResultListening')
                parent = listening.parentItem()
                assert abs(listening.mapToItem(parent, QPointF()).x() + listening.width()/2 - parent.width()/2) < 2
                buttons = [c(name) for name in ('narrationResultListenEnglish',
                           'narrationResultListenVietnamese') if c(name).isVisible()]
                bounds = [(button.mapToItem(parent, QPointF()).x(),
                           button.mapToItem(parent, QPointF(button.width(), 0)).x()) for button in buttons]
                assert abs((bounds[0][0] + bounds[-1][1])/2 - parent.width()/2) < 2
                assert bounds[0][0] >= 0 and bounds[-1][1] <= parent.width()
                assert all(right + 8 <= following for (_, right), (following, _) in zip(bounds, bounds[1:]))
                assert all(button.width() >= button.implicitWidth() for button in buttons)
                assert all(not button.findChild(QObject, 'actionButtonLabel').property('truncated') for button in buttons)

            assert c('narrationResultDownload').property('text') == 'Tải PowerPoint có giọng'
            for width, height in sizes:
                window.resize(width, height)
                pump(150)
                for name in ('narrationResultPrevious', 'narrationResultCounter', 'narrationResultNext',
                             'narrationResultListenEnglish', 'narrationResultListenVietnamese'):
                    item = c(name)
                    assert item.isVisible(), name
                    assert item.mapToScene(QPointF()).x() >= 0
                    assert item.mapToScene(QPointF(item.width(), 0)).x() <= width
                counter = c('narrationResultCounter')
                actions = c('narrationResultActions')
                parent = actions.parentItem()
                assert abs(actions.mapToItem(parent, QPointF()).x()+actions.width()/2-parent.width()/2) < 2
                viewport = c('narrationResultPage').property('contentItem')
                assert abs(counter.mapToScene(QPointF(counter.width()/2, 0)).x()
                           - c('narrationResultPage').mapToScene(QPointF(c('narrationResultPage').width()/2, 0)).x()) < 5
                viewport.setProperty('contentY', max(0, viewport.property('contentHeight') - viewport.height()))
                pump(180)
                for english, vietnamese in ((True, True), (True, False), (False, True)):
                    check_listening_alignment(english, vietnamese)
                    assert window.grabWindow().save(str(reports / f'listening-{width}-{english}-{vietnamese}.png'))
                check_listening_alignment()
                assert window.grabWindow().save(str(reports / f'completed-{width}.png'))
            bridge.setApplicationTheme('dark')
            pump(120)
            check_listening_alignment()
            assert window.grabWindow().save(str(reports / 'completed-dark.png'))
            bridge.setApplicationTheme('light')
            # Use valid cached WAVs and the real reader/queue, with a controllable
            # output device. Never synthesize or alter the teacher's system audio.
            class PreviewPlayer:
                auto_finish = False

                def play(self, path):
                    self.path = path

                def finished(self):
                    return self.auto_finish

                def stop(self):
                    self.path = None

            player = PreviewPlayer()
            reader = bridge.english_reader
            en, vi = c('narrationResultListenEnglish'), c('narrationResultListenVietnamese')
            assert window.findChild(QObject, 'narrationResultStop') is None
            assert window.findChild(QObject, 'stopCreatedNarration') is None
            widths = [en.width(), vi.width()]
            with patch.object(reader, 'player', player), patch('app.speech.synthesize', side_effect=AssertionError('Cached preview only')):
                click(window, en)
                assert reader.state == 'reading' and bridge.narrationWorkspace.createdPlayback['language'] == 'en'
                assert en.property('text') == 'Dừng nghe' and en.property('symbol') == '■'
                assert vi.property('text') == 'Nghe giọng Việt'
                assert widths == [en.width(), vi.width()]
                click(window, en)
                assert reader.state == 'stopped' and en.property('text') == 'Nghe giọng Anh'
                vi.forceActiveFocus(Qt.OtherFocusReason)
                QTest.keyClick(window, Qt.Key_Space)
                pump(50)
                assert vi.property('text') == 'Dừng nghe' and en.property('text') == 'Nghe giọng Anh'
                click(window, en)
                assert bridge.narrationWorkspace.createdPlayback['language'] == 'en'
                assert vi.property('text') == 'Nghe giọng Việt' and en.property('text') == 'Dừng nghe'
                bridge.stopSpeech()
                pump(50)
                assert not bridge.narrationWorkspace.createdPlayback
                click(window, vi)
                player.auto_finish = True
                until(lambda: reader.state == 'stopped')
                assert vi.property('text') == 'Nghe giọng Việt'
                player.auto_finish = False
                for theme in ('dark', 'light'):
                    bridge.setApplicationTheme(theme)
                    for language, button in (('en', en), ('vi', vi)):
                        bridge.stopSpeech()
                        click(window, button)
                        assert button.property('text') == 'Dừng nghe' and button.property('danger')
                        check_listening_alignment()
                        assert widths == [en.width(), vi.width()]
                        assert window.grabWindow().save(str(reports / f'toggle-{theme}-{language}.png'))
                        click(window, button)
                        assert reader.state == 'stopped'
                with patch.object(player, 'play', side_effect=RuntimeError('Fixture playback device unavailable')):
                    click(window, en)
                    assert reader.state == 'error' and en.property('text') == 'Nghe giọng Anh'
                click(window, en)
                queue, voice, rate = list(reader.queue), reader.voice, reader.rate
                reader.start(queue, voice, rate, cache_only=True)
                pump(50)
                assert reader.state == 'reading' and not bridge.narrationWorkspace.createdPlayback
                assert en.property('text') == 'Nghe giọng Anh', 'Other reader activity must not mark this button as playing'
                bridge.stopSpeech()
                click(window, en)
                click(window, c('narrationResultNext'))
                until(lambda: not bridge.previewBusy, timeout=90)
                assert reader.state == 'stopped' and not bridge.narrationWorkspace.createdPlayback
                assert en.property('text') == 'Nghe giọng Anh'
            bridge._powerpoint_state = {**initial_powerpoint_state, 'active': True, 'starting': False}
            bridge.powerpointStateChanged.emit()
            with patch.object(bridge, 'stopSpeech') as stopped:
                bridge.previewConvertedSlide(2)
                until(lambda: not bridge.previewBusy, timeout=90)
                assert bridge.quickResult['slide'] == 2 and bridge.quickResult['image']
                stopped.assert_not_called()
            assert c('narrationResultNext').isEnabled()
            assert c('narrationResultDownload').isEnabled()
            assert c('narrationResultOpen').isEnabled(), 'Opening an active managed lesson is a safe focus/switch action'
            bridge._powerpoint_state = dict(initial_powerpoint_state)
            bridge.powerpointStateChanged.emit()
            c('narrationResultChange').clicked.emit()
            pump(100)
            assert c('downloadAfterNarrationButton').property('text') == 'Tải PowerPoint có giọng'
            assert c('voiceSlideThumbnail').isVisible()
            c('toggleVoiceSlidePreview').clicked.emit()
            until(lambda: not bridge.previewBusy, timeout=90)
            pump(100)
            assert c('createdNarrationListening').isVisible()
            with patch.object(reader, 'player', player):
                for name in ('listenCreatedEnglish', 'listenCreatedVietnamese'):
                    button = c(name)
                    assert button.isVisible() and button.isEnabled()
                    initial_width = button.width()
                    button.clicked.emit()
                    pump(50)
                    assert button.property('text') == 'Dừng nghe'
                    assert button.width() == initial_width
                    button.clicked.emit()
                    assert reader.state == 'stopped'
            c('toggleVoiceSlidePreview').clicked.emit()
            assert abs(c('voiceSlideThumbnail').width() / c('voiceSlideThumbnail').height() - 16/9) < .01
            c('toggleVoiceSlidePreview').clicked.emit()
            pump(60)
            assert window.property('voicePreviewExpanded')
            c('toggleVoiceSlidePreview').clicked.emit()
            pump(60)
            assert not window.property('voicePreviewExpanded')
            window.setProperty('page', 'library')
            window.setProperty('page', 'voice')
            pump(150)
            assert not c('narrationResultPage').isVisible(), 'Returning to the tool keeps configuration available'
            window.setProperty('voiceResultDismissed', False)
            bridge.narrationChanged.emit()
            pump(100)
            assert c('narrationResultPage').isVisible()
            c('navigationvoice').clicked.emit()
            pump(100)
            assert not c('narrationResultPage').isVisible(), 'Clicking the current Voice tab exits its result too'
            window.setProperty('narrationOptionsPinned', True)
            window.setProperty('narrationVoice', 'pair:north-male')
            bridge.narration_preparation.state = dict(plan, lesson_id=lesson['id'], running=True, complete=0)
            bridge.narrationChanged.emit()
            bridge.narration_preparation.state.update(running=False, complete=6, failed=0, cancelled=False)
            bridge.narrationChanged.emit()
            pump(120)
            assert window.property('voicePreviewExpanded'), 'Completion reveals the slide even if the draft choice differs'
            c('narrationResultChange').clicked.emit()
            window.setProperty('voicePreviewExpanded', False)
            bridge.narrationChanged.emit()
            pump(60)
            assert not window.property('voicePreviewExpanded'), 'Later notifications must respect a deliberate hide'
            c('navigationlibrary').clicked.emit()
            pump(100)
            assert window.property('page') == 'library'
            c('navigationvoice').clicked.emit()
            pump(100)
            assert window.property('page') == 'voice' and not c('narrationResultPage').isVisible()
            assert c('resumeNarrationTask').isVisible() and c('resumeNarrationTask').property('attention')
            recent = dict(bridge.recentNarrationResult)
            assert recent['lesson_id'] == lesson['id']
            window.setProperty('page', 'new')
            pump(60)
            assert not c('resumeNarrationTask').isVisible(), 'Narration link belongs only to Voice'
            assert not c('resumeCreationTask').isVisible(), 'Narration must not invent a converted result'
            window.setProperty('page', 'voice')
            pump(60)
            actions = c('narrationReadyActions')
            parent = actions.parentItem()
            assert abs(actions.mapToItem(parent, QPointF()).x()+actions.width()/2-parent.width()/2) < 2
            window.setProperty('narrationMode', 'en')
            window.setProperty('narrationRate', 1)
            c('resumeNarrationTask').clicked.emit()
            until(lambda: not bridge.busy and c('narrationResultPage').isVisible() and bridge.lessonNarration.get('ready'), timeout=60)
            assert window.property('narrationMode') == recent['mode'] and window.property('narrationRate') == recent['rate']
            assert bridge.selectedNarrationVariant == recent['variant_id']
            bridge._clear_narration_selection()
            pump(150)
            assert c('resumeNarrationTask').isVisible() and not bridge.lesson.get('id')
            c('resumeNarrationTask').clicked.emit()
            until(lambda: not bridge.busy and c('narrationResultPage').isVisible() and bridge.lesson.get('id') == lesson['id'], timeout=60)
            assert bridge.selectedNarrationVariant == recent['variant_id']
            # Keep two different completed lessons: each tab must reopen its own result.
            converted_source = folder / 'Bài 5 – Kết quả chuyển đổi.pptx'
            presentation.slides[0].shapes[0].text = 'Bài 5: Kết quả chuyển đổi\nLesson 5: Converted result'
            presentation.save(converted_source)
            converted, _ = import_presentation(folder / 'data', converted_source)
            assert converted['id'] != lesson['id']
            bridge._chatgpt_request = dict(folder=str(folder), config=dict(title=converted['title']))
            bridge._browser_state = dict(phase='completed', running=False, percent=100, lessonId=converted['id'])
            bridge.changed.emit()
            bridge.browserStateChanged.emit()
            window.setProperty('page', 'new')
            pump(80)
            assert not c('resumeNarrationTask').isVisible()
            assert c('resumeCreationTask').isVisible() and c('resumeCreationTask').property('attention')
            assert c('resumeCreationTask').property('text') == 'Mở bài vừa chuyển đổi'
            c('resumeCreationTask').clicked.emit()
            until(lambda: not bridge.busy and bridge.lesson.get('id') == converted['id'], timeout=60)
            assert window.property('page') == 'chatgpt'
            assert bridge.recentNarrationResult == recent
            c('navigationvoice').clicked.emit()
            pump(80)
            assert c('resumeNarrationTask').isVisible() and c('resumeNarrationTask').property('attention')
            assert not c('resumeCreationTask').isVisible()
            c('resumeNarrationTask').clicked.emit()
            until(lambda: not bridge.busy and c('narrationResultPage').isVisible() and bridge.lesson.get('id') == lesson['id'], timeout=60)
            assert bridge.selectedNarrationVariant == recent['variant_id']
            assert bridge.browserState['lessonId'] == converted['id']
            c('narrationResultChange').clicked.emit()
            pump(80)
            assert window.grabWindow().save(str(reports / 'voice-only-recent-link.png'))
            window.setProperty('page', 'new')
            pump(80)
            assert window.grabWindow().save(str(reports / 'conversion-only-recent-link.png'))
            window.setProperty('narrationRate', 1)
            for width, height in sizes:
                window.resize(width, height)
                for page in ('voice', 'new'):
                    check_footer_alignment(page, recent_link=True)
                    assert window.grabWindow().save(str(reports / f'footer-{page}-{width}.png'))
            assert not warnings, warnings
            (reports / 'result.json').write_text(json.dumps({'status': 'passed', 'automatic_preview': True,
                'resume_exact_completed_variant_after_configuration_or_clear': True, 'centered_actions': True,
                'centered_listening_all_language_modes': True, 'voiced_download_label': True,
                'one_listen_stop_toggle_per_language': True, 'toggle_keyboard_and_stable_width': True,
                'reset_on_finished_error_navigation_external_stop': True, 'reader_ownership_checked': True,
                'tab_specific_recent_results_reopen_distinct_lessons': True,
                'footer_primary_action_aligned_with_summary': True,
                'preview_while_powerpoint_active': True, 'sizes': sizes, 'qml_warnings': warnings}), encoding='utf-8')
            print('Completed narration Qt flow passed.', flush=True)
        finally:
            bridge._powerpoint_state = dict(initial_powerpoint_state)
            for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                           bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                           bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
                worker.shutdown()
            c('companionWindow').setProperty('disposing', True)
            c('companionWindow').hide()
            window.hide()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            pump(80)
            library.close()


if __name__ == '__main__':
    main()
