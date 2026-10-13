"""Time native Office rendering, verified cross-folder reuse and first-slide preview."""

import hashlib
import json
import tempfile
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.chatgpt_handoff import inspect_returned_deck
from app.powerpoint_review import render_slide
from app.slide_quality import render_quality


def main():
    output = Path("reports/conversion-progress")
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="render-cache-", dir=output) as temporary:
        workspace = Path(temporary).resolve()
        source = workspace / "Bài kiểm tra.pptx"
        deck = Presentation()
        for index in range(6):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2))
            box.text = f"Slide {index+1}\nThe force acts on the object."
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.size = Pt(28)
        deck.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        report = inspect_returned_deck(source)["slide_quality"]
        started = time.perf_counter()
        first = render_quality(source, workspace / "web", report)
        first_seconds = time.perf_counter()-started
        assert all(row["rendered"] for row in first["slides"]), "Native Office rendering is required"
        fresh = inspect_returned_deck(source)["slide_quality"]
        started = time.perf_counter()
        second = render_quality(source, workspace / "library", fresh, reuse_directory=workspace / "web")
        second_seconds = time.perf_counter()-started
        assert second["render_cache"] == {"reused": 6, "rendered": 0, "total": 6}
        assert first["counts"] == second["counts"]
        started = time.perf_counter()
        preview = Path(render_slide(source, 1, workspace / "library"))
        preview_seconds = time.perf_counter()-started
        assert preview.is_file() and "slide-quality" in preview.parts
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original
        result = {"status": "passed", "slides": 6, "first_render_seconds": round(first_seconds, 3),
                  "reused_render_seconds": round(second_seconds, 3), "preview_seconds": round(preview_seconds, 3),
                  "scope": "Real PowerPoint on this machine; no web submissions, no teacher-library edits"}
    (output / "render-cache.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
