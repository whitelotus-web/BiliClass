"""Imported source stays in Voice; finishing audio publishes the same library row."""

import json
import tempfile
import wave
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import speech
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge

from .qt_result_preview_mascot_smoke import click, control, pump, until


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    with tempfile.TemporaryDirectory(prefix='biliclass-products-', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        source = folder/'Bài thử.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = 'Energy is conserved.'
        slide.notes_slide.notes_text_frame.text = 'EN: Energy is conserved.'
        deck.save(source)
        original = source.read_bytes()
        library = Library(folder/'data')
        with patch.object(speech, 'list_voices', return_value=[dict(speech.PAIRS[0])]):
            bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda rows: warnings.extend(map(str, rows)))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(RESOURCE_ROOT/'qml/Main.qml')
        root = engine.rootObjects()[0]

        def render(text, voice, rate, directory, cancel):
            target = Path(directory)/(speech.audio_key(text, voice, rate)+'.wav')
            target.parent.mkdir(parents=True, exist_ok=True)
            with wave.open(str(target), 'wb') as wav:
                wav.setparams((1, 2, 48000, 0, 'NONE', 'not compressed'))
                wav.writeframes(b'\1\0'*4800)
            return {'path': str(target), 'cached': False}

        try:
            with patch.object(Bridge, 'previewConvertedSlide'):
                bridge.importNarrationPresentation(source.as_uri())
                until(lambda: not bridge.busy and bool(bridge.lesson))
                identity = bridge.lesson['id']
                until(lambda: bridge._catalog.worker is None and len(bridge.narrationSources) == 1)
                assert not bridge.lessons
                assert bridge.narrationSources[0]['title'].startswith('Đang chuẩn bị · ')
                bridge.dismissNarrationLesson()
                pump(100)
                picker = control(root, 'voiceLibraryPicker')
                assert picker.property('count') == 1
                picker.setProperty('selectedIds', [identity])
                pump(60)
                click(root, control(root, 'selectVoiceLibraryLesson'))
                until(lambda: not bridge.busy and bridge.lesson.get('id') == identity)
                with patch.object(speech, 'synthesize_background', side_effect=render), \
                        patch.object(speech, 'narration_workers', return_value=1):
                    bridge.createLessonNarrationWithVoice('en', speech.PAIRS[0]['id'], 0)
                    until(lambda: not bridge.busy and not bridge.narrationBusy and bridge.lessonNarration.get('ready'))
                until(lambda: bridge._catalog.worker is None and len(bridge.lessons) == 1)
                assert bridge.lessons[0]['id'] == identity
                assert library.get(identity)['voice_published']
                assert len(library.list_lessons(verify_audio=False)) == 1
                assert source.read_bytes() == original
                assert not warnings, warnings
                report = Path('reports/rc138/library-products')
                report.mkdir(parents=True, exist_ok=True)
                (report/'result.json').write_text(json.dumps({'status': 'passed', 'qml_warnings': warnings,
                    'checks': ['source hidden from products but selectable in Voice',
                               'resume source after leaving workspace',
                               'voice completion publishes same ID; no duplicate or source edit']}))
                print('Qt product/source lifecycle passed; no QML warnings.', flush=True)
        finally:
            control(root, 'companionWindow').setProperty('disposing', True)
            root.hide()
            for name in ('libraryThumbnails', 'googleSettings', '_library_preferences', 'slidePreview',
                         '_catalog', 'browserAI', 'english_reader', 'narration_preparation',
                         'narrationPreferences', '_companion_preferences', 'classroom', 'voice_controls',
                         'inputDocuments'):
                getattr(bridge, name).shutdown()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
