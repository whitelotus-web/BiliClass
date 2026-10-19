"""Real Qt voice transitions with slow storage, safe stall logging and worker churn.

Uses a disposable 51-slide deck and silent WAV fixtures; no teacher data or API.
Model inference is verified separately by the packaged voice smoke check.
"""

import json
import tempfile
import time
import wave
from pathlib import Path
from threading import Event, Thread, get_ident
from unittest.mock import patch

from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEvent, QTimer
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle

from app import lesson_narration
from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.library import Library
from app.paths import RESOURCE_ROOT
from app.ui import Bridge
from app.ui_watchdog import UiWatchdog

from .qt_result_preview_mascot_smoke import pump, until


def wav(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with wave.open(str(path), 'wb') as output:
        output.setparams((1, 2, 48000, 0, 'NONE', 'not compressed'))
        output.writeframes(b'\0\0' * 2400)


def main():
    app = QGuiApplication([])
    QQuickStyle.setStyle('Basic')
    main_thread = get_ident()
    with tempfile.TemporaryDirectory(prefix='bcl-voice-stability-', ignore_cleanup_errors=True) as folder:
        directory = Path(folder)
        deck = Presentation()
        for index in range(51):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text = f'VI: Năng lượng được bảo toàn trong bài {index}.\nEN: Energy is conserved in lesson {index}.'
            slide.shapes.add_textbox(914400, 914400, 6500000, 1200000).text = text
            slide.notes_slide.notes_text_frame.text = text
        source = directory / 'fixture.pptx'
        deck.save(source)
        library = Library(directory / 'data')
        config = dict(title='Voice stability', subject='Vật lý', education_level='THPT', grade='11',
                      conversion_format='parallel_columns', style='source', preset='standard')
        lesson = library.create_external_lesson(config, library.store_source(source), inspect_returned_deck(source))
        voice = 'vieneu:Mai Anh'
        plan, _ = lesson_narration.make_plan(lesson, 'en', {'voice': voice, 'rate': 0})
        assert plan['total'] == 51
        for key in plan['keys']['en']:
            wav(library.directory / 'audio' / (key + '.wav'))
        lesson = library.update_narration(lesson['id'], plan)
        bridge = Bridge(library)
        engine = QQmlApplicationEngine()
        warnings = []
        engine.warnings.connect(lambda values: warnings.extend(str(v) for v in values))
        engine.rootContext().setContextProperty('bridge', bridge)
        engine.load(str(RESOURCE_ROOT / 'qml/Main.qml'))
        assert engine.rootObjects(), warnings
        window = engine.rootObjects()[0]
        timer = QTimer()
        timer.setInterval(20)
        ticks = []
        timer.timeout.connect(lambda: ticks.append(time.monotonic()))
        watcher = UiWatchdog(app, library.directory)
        churn_stop = Event()

        def churn():
            while not churn_stop.wait(.001):
                exec(compile('sum(range(10))', '<temporary-code>', 'exec'), {})

        churn_worker = Thread(target=churn)
        try:
            until(lambda: bridge._catalog.worker is None and bool(bridge._catalog.rows()), timeout=30)
            # Reproduce a stalled GUI while another thread creates/destroys code.
            # The RC85 native frame walker could crash here; this writer holds the GIL.
            churn_worker.start()
            time.sleep(6.2)
            pump(80)
            assert watcher.path.exists() and '(GUI)' in watcher.path.read_text(encoding='utf-8')
            churn_stop.set()
            churn_worker.join(2)
            timer.start()
            scan = lesson_narration.valid_key
            scan_threads = set()

            def slow_scan(directory, key):
                scan_threads.add(get_ident())
                assert get_ident() != main_thread, 'QML or click handler scanned WAV files on GUI thread'
                time.sleep(.12)
                return scan(directory, key)

            with patch.object(lesson_narration, 'valid_key', side_effect=slow_scan):
                bridge._browser_state = dict(phase='completed', lessonId=lesson['id'])
                bridge._show_lesson(lesson, external_preview(lesson, library.directory),
                                    destination='chatgpt', auto_repair=False, load_preview=False)
                bridge.openConvertedLesson()
                assert window.property('page') == 'voice'
                before = len(ticks)
                pump(400)
                assert len(ticks) - before >= 10 and bridge.lessonNarration['status'] == 'checking'
                started = time.monotonic()
                for _ in range(1000):
                    bridge.lessonNarration
                getter_seconds = time.monotonic() - started
                assert getter_seconds < .25
                until(lambda: bridge.lessonNarration.get('ready'), timeout=30)
                assert bridge.lessonNarration['complete'] == 51
                confirmations = []
                bridge.narrationAlreadyCreated.connect(confirmations.append)
                bridge.createLessonNarrationWithVoice('en', voice, 0)
                assert bridge.busy and not bridge.narrationBusy
                before = len(ticks)
                pump(400)
                assert len(ticks) - before >= 10
                until(lambda: not bridge.busy, timeout=30)
                assert confirmations and not bridge.narrationBusy
            # Repeated navigation and replacement must never publish an old worker result.
            for _ in range(40):
                window.setProperty('page', 'library')
                bridge._clear_narration_selection()
                bridge._show_lesson(lesson, external_preview(lesson, library.directory),
                                    destination='voice', auto_repair=False, load_preview=False)
                pump(15)
            until(lambda: bridge.lessonNarration.get('ready'), timeout=30)
            gaps = [b-a for a, b in zip(ticks, ticks[1:])]
            assert max(gaps, default=0) < 1, max(gaps)
            assert not warnings, warnings
            assert len(list((library.directory / 'audio').glob('*.wav'))) == 51
            report = dict(status='passed', slides=51, transitions=40,
                          safe_stall_dump=True, storage_delay_per_wav_seconds=.12,
                          audio_scan_threads=len(scan_threads), wav_scans_on_gui=0,
                          thousand_status_reads_seconds=round(getter_seconds, 4),
                          max_ui_gap_seconds=round(max(gaps), 4), qml_warnings=warnings,
                          no_teacher_data=True, no_network=True)
            output = Path('reports/rc86/voice-stability.json')
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(json.dumps(report, indent=2), encoding='utf-8')
            print(json.dumps(report), flush=True)
        finally:
            timer.stop()
            churn_stop.set()
            if churn_worker.ident is not None:
                churn_worker.join(2)
            watcher.stop()
            watcher.thread.join(2)
            for service in (bridge.narration_preparation, bridge.english_reader, bridge.voice_controls,
                            bridge.classroom, bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog,
                            bridge.browserAI, bridge.narrationPreferences):
                service.shutdown()
            window.close()
            engine.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            app.processEvents()
            library.close()


if __name__ == '__main__':
    main()
