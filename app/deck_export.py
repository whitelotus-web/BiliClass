"""Create a separate, editable bilingual PowerPoint from reviewed segments."""
import io
import tempfile
from pathlib import Path

from .lesson_design import paired_pages
from .powerpoint import slide_for_locator
from .presentation_policy import presentation_content


def export_deck(lesson, directory, destination, profile=None, terms=()):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches, Pt
    target = Path(destination).with_suffix(".pptx").resolve()
    source = lesson.get("source")
    if source and target == (Path(directory) / "sources" / source["file"]).resolve():
        raise ValueError("Chọn đường dẫn mới để giữ nguyên tệp nguồn.")
    if not lesson.get("segments") or any(not s.get("approved") for s in lesson["segments"]):
        raise ValueError("Duyệt toàn bộ cặp Việt/Anh trước khi xuất PowerPoint.")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    source_deck = None
    if source and source.get("file", "").lower().endswith(".pptx"):
        from .powerpoint import verified_presentation
        source_deck = Presentation(verified_presentation(lesson, directory))

    def textbox(slide, text, x, y, w, h, size, color, bold=False, italic=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.04)
        for i, line in enumerate(text.splitlines()):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold, p.font.italic = "Arial", Pt(size), bold, italic
            p.font.color.rgb = RGBColor.from_string(color)
        return box

    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                yield shape

    def add_original_visual(source_slide):
        original = source_deck.slides[source_slide - 1]
        images = list(pictures(original.shapes))
        if not images:
            return
        visual = presentation.slides.add_slide(presentation.slide_layouts[6])
        visual.background.fill.solid()
        visual.background.fill.fore_color.rgb = RGBColor.from_string("F3F7FC")
        textbox(visual, lesson["title"][:100], .6, .35, 12, .65, 24, "112650", True)
        textbox(visual, f"Hình từ slide gốc {source_slide}", .65, 6.85, 12, .3, 12, "667997")
        for picture in images:
            # Keep relative positions without stretching or inventing captions.
            scale = min(11.6 / float(source_deck.slide_width), 5.55 / float(source_deck.slide_height))
            visual.shapes.add_picture(
                io.BytesIO(picture.image.blob),
                int(Inches(.85) + picture.left * scale),
                int(Inches(1.05) + picture.top * scale),
                width=int(picture.width * scale),
                height=int(picture.height * scale),
            )

    next_source_slide = 1
    preset = lesson.get("teaching_preset", "standard")
    if preset not in {"standard", "visual", "practice"}:
        raise ValueError("Kiểu dạy không hợp lệ.")
    for segment in lesson["segments"]:
        source_slide = slide_for_locator(segment.get("locator", "")) if source_deck else None
        if source_slide and source_slide <= len(source_deck.slides):
            while next_source_slide <= source_slide:
                add_original_visual(next_source_slide)
                next_source_slide += 1
        layout = lesson.get("layout", "line_pair")
        content = presentation_content(segment, lesson.get("level", 2), layout, terms=terms)
        vi = content["vi"] if content["show_vi"] else ""
        en = content["en"] if content["show_en"] else ""
        try:
            pages = paired_pages(vi, en, limit=340 if layout == "split_view" else 430)
        except ValueError as exc:
            raise ValueError(f"{segment['locator']}: {exc}") from exc
        for index, (vi_text, en_text) in enumerate(pages):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(
                {"standard": "F3F7FC", "visual": "F1FAFA", "practice": "FFF8EE"}[preset]
            )
            logo_path = Path(directory) / "assets" / "school-logo.png"
            header_x = .6
            if profile and profile.get("show_profile", True) and logo_path.is_file():
                from PIL import Image
                with Image.open(logo_path) as logo:
                    ratio = logo.width / logo.height
                width, height = (.68, .68 / ratio) if ratio >= 1 else (.68 * ratio, .68)
                slide.shapes.add_picture(str(logo_path), Inches(.6), Inches(.22), width=Inches(width), height=Inches(height))
                header_x = 1.45
            textbox(slide, "BiliClass  /  " + lesson["subject"] + "  /  Khối " + str(lesson.get("grade", "")), header_x, .35, 12 - header_x, .5, 16, "0869F9", True)
            textbox(slide, lesson["title"][:100], .6, .95, 12, .9, 28, "112650", True)
            if profile and profile.get("show_profile", True):
                byline = " · ".join(filter(None, (
                    profile.get("teacher", ""), profile.get("school", "")
                )))
                if byline:
                    textbox(slide, byline[:150], .65, 1.69, 12, .28, 12, "667997")
            if layout == "split_view":
                textbox(slide, vi_text, .65, 2, 5.85, 4.5, 20, "112650")
                textbox(slide, en_text, 6.85, 2, 5.85, 4.5, 20, "0869F9")
            elif layout == "line_pair":
                textbox(slide, vi_text, .65, 2, 12, 2.1, 23, "112650")
                textbox(slide, en_text, .65, 4.3, 12, 2.1, 21, "0869F9", italic=True)
            elif layout == "keyword_overlay":
                textbox(slide, vi_text, .65, 2, 12, 4.5, 23, "112650")
            else:
                textbox(slide, en_text, .65, 2, 12, 4.5, 24, "0869F9")
            textbox(slide, segment["locator"] + (f" · {index+1}/{len(pages)}" if len(pages) > 1 else ""), .65, 6.8, 12, .4, 12, "667997")
            slide.notes_slide.notes_text_frame.text = "BiliClass: nội dung đã duyệt. Đoạn: " + segment["id"] + "\n" + segment["locator"]
    if source_deck:
        while next_source_slide <= len(source_deck.slides):
            add_original_visual(next_source_slide)
            next_source_slide += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".pptx", delete=False) as f:
        pending = Path(f.name)
    try:
        presentation.save(str(pending))
        pending.replace(target)
    finally:
        pending.unlink(missing_ok=True)
    return target
