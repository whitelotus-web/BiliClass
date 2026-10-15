"""Fit measured overflows on a copy; preserve all unrelated PPTX ZIP parts."""

import copy
import hashlib
import json
import zipfile
from pathlib import Path
from uuid import uuid4

from lxml import etree

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def file_digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def failure_keys(report):
    return {(row["slide"], issue["code"], issue.get("shape_id"))
            for row in report["slides"] for issue in row["issues"] if issue["severity"] == "FAIL"}


def retain_improved_pages(current, candidate, before, after, directory):
    """An unsuccessful page must not discard independently verified page fixes."""
    from .source_deck import source_slide_parts

    wanted = failure_keys(before)
    remaining = failure_keys(after)
    selected = {number for number in range(1,len(before['slides'])+1)
                if {key for key in remaining if key[0]==number} < {key for key in wanted if key[0]==number}}
    if not selected:
        return None, set()
    parts, adjusted = source_slide_parts(current), source_slide_parts(candidate)
    with zipfile.ZipFile(candidate) as archive:
        replacements = {parts[n-1]:archive.read(adjusted[n-1]) for n in selected}
    target = Path(directory)/'temp/slide-repair'/uuid4().hex/'selected-pages.pptx'
    target.parent.mkdir(parents=True)
    _write_copy(current,target,replacements)
    return target,selected


def actionable_rows(report):
    return [row for row in report["slides"] if any(
        i["severity"] == "FAIL" or i["code"] in {"english_body_missing", "translation_coverage"}
        for i in row["issues"])]


def _shape(root, shape_id):
    found = root.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@id=$id]]', namespaces=NS, id=str(shape_id))
    return found[0] if len(found) == 1 else None


def _write_copy(source, target, replacements):
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as after:
        for item in before.infolist():
            after.writestr(item, replacements.get(item.filename, before.read(item)))


def _compact_fixed_spacing(body):
    """Let line spacing follow the font in an already measured overflow.

    Shrinking glyphs alone cannot fit text with inherited 50–70 pt leading.
    Keep ordinary spacing and all wording; touch only excessive fixed gaps.
    """
    sizes = [int(n) for n in body.xpath('.//a:rPr/@sz|.//a:defRPr/@sz', namespaces=NS)]
    size = max(sizes, default=2400)
    for node in body.xpath('.//a:lnSpc/a:spcPts', namespaces=NS):
        if int(node.get('val', '0')) > size * 1.35:
            parent = node.getparent()
            parent.remove(node)
            etree.SubElement(parent, '{' + NS['a'] + '}spcPct', val='110000')
    for node in body.xpath('.//a:spcBef/a:spcPts|.//a:spcAft/a:spcPts', namespaces=NS):
        if int(node.get('val', '0')) > size * .5:
            node.set('val', str(round(size * .15)))


def _separate_title_styles(root, width, height):
    """Style an unambiguous stacked title stored in two distinct text boxes."""
    from .slide_speech import VI_MARKS, english_text

    headings = []
    animated = set(root.xpath('.//p:spTgt/@spid', namespaces=NS))
    for shape in root.findall('p:cSld/p:spTree/p:sp', NS):
        identity = shape.find('p:nvSpPr/p:cNvPr', NS)
        transform = shape.find('p:spPr/a:xfrm', NS)
        paragraphs = shape.findall('p:txBody/a:p', NS)
        if identity is None or transform is None or len(paragraphs) != 1 or identity.get('id') in animated:
            continue
        off, ext = transform.find('a:off', NS), transform.find('a:ext', NS)
        if off is None or ext is None or transform.get('rot', '0') != '0':
            continue
        x, y, w, h = int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))
        text = ''.join(paragraphs[0].xpath('.//a:t/text()', namespaces=NS)).strip()
        if not text or len(text) > 160 or y < 0 or y > height * .22 or w < width * .4:
            continue
        headings.append((identity, paragraphs[0], text, x, y, w, h))
    changed = []
    for identity, paragraph, text, x, y, w, _ in headings:
        if not english_text(text):
            continue
        matches = [row for row in headings if row[2] != text and any(c in VI_MARKS for c in row[2].casefold())
                   and not english_text(row[2]) and row[4]+row[6]-12700 <= y <= row[4]+row[6]+height*.08
                   and min(x+w, row[3]+row[5])-max(x, row[3]) >= min(w, row[5])*.8]
        if len(matches) != 1:
            continue
        source = matches[0][1]
        fonts = source.xpath('.//a:rPr/a:latin|a:pPr/a:defRPr/a:latin', namespaces=NS)
        sizes = source.xpath('.//a:rPr/@sz|a:pPr/a:defRPr/@sz', namespaces=NS)
        if not fonts or not sizes or max(map(int, sizes)) < 2200:
            continue
        before = etree.tostring(paragraph)
        for run in paragraph.findall('a:r', NS):
            props = run.find('a:rPr', NS)
            if props is None:
                props = etree.Element('{' + NS['a'] + '}rPr')
                run.insert(0, props)
            props.set('i', '1')
            size = int(props.get('sz', str(max(map(int, sizes)))))
            props.set('sz', str(max(min(size, 1400), min(size, round(max(map(int, sizes))*.8)))))
            for font in list(props.findall('a:latin', NS)):
                props.remove(font)
            insert_at = next((i for i, n in enumerate(props) if etree.QName(n).localname
                              in {'ea','cs','sym','hlinkClick','hlinkMouseOver','rtl','extLst'}), len(props))
            props.insert(insert_at, copy.deepcopy(fonts[0]))
        if etree.tostring(paragraph) != before:
            changed.append(int(identity.get('id')))
    return changed


def normalize_titles(path, directory, mode):
    """Style short native VI/EN heading pairs without changing their wording."""
    from pptx import Presentation

    from .slide_speech import VI_MARKS, english_text
    from .source_deck import source_slide_parts

    if mode not in {"parallel_columns", "sentence_pairs", "integrated_keywords"}:
        return {"path": str(path), "changes": []}
    deck = Presentation(path)
    parts = source_slide_parts(path)
    replacements, changes = {}, []
    with zipfile.ZipFile(path) as archive:
        for number, slide in enumerate(deck.slides, 1):
            part = parts[number - 1]
            root = etree.fromstring(archive.read(part))
            for shape in root.xpath(".//p:sp", namespaces=NS):
                body = shape.find("p:txBody", NS)
                if body is None or shape.xpath(".//*[local-name()='oMath']"):
                    continue
                paragraphs = body.findall("a:p", NS)
                if len(paragraphs) != 2:
                    continue
                texts = ["".join(p.xpath(".//a:t/text()", namespaces=NS)).strip() for p in paragraphs]
                if (not any(c in VI_MARKS for c in texts[0].casefold()) or len(texts[0]) > 120
                        or len(texts[1]) > 160 or not english_text(texts[1], tagged=True)):
                    continue
                offset = shape.find("p:spPr/a:xfrm/a:off", NS)
                native_ids = {i.shape_id for i in slide.shapes if getattr(i, "has_text_frame", False)}
                if offset is None or int(offset.get("y", "0")) > deck.slide_height * .27 and len(native_ids) > 1:
                    continue
                before = etree.tostring(body)
                vi_fonts = paragraphs[0].xpath(".//a:rPr/a:latin|a:pPr/a:defRPr/a:latin", namespaces=NS)
                sizes = [int(n) for n in paragraphs[0].xpath(".//a:rPr/@sz", namespaces=NS)]
                color = paragraphs[0].xpath(".//a:rPr/a:solidFill/a:srgbClr/@val", namespaces=NS)
                ink = "FFFFFF" if color and color[0].upper() in {"FFFFFF", "FFFFFE", "FEFEFE"} else "111111"
                for run in paragraphs[1].findall("a:r", NS):
                    props = run.find("a:rPr", NS)
                    if props is None:
                        props = etree.Element("{" + NS["a"] + "}rPr")
                        run.insert(0, props)
                    props.set("i", "1")
                    if sizes:
                        existing_size = int(props.get("sz", str(sizes[0])))
                        props.set("sz", str(max(min(1400, existing_size), min(existing_size, round(sizes[0] * .75)))))
                    for node in list(props):
                        if etree.QName(node).localname in {"latin", "solidFill", "noFill", "gradFill", "pattFill", "grpFill"}:
                            props.remove(node)
                    if vi_fonts:
                        insert_at = next((i for i, n in enumerate(props) if etree.QName(n).localname
                                          in {"ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"}), len(props))
                        props.insert(insert_at, copy.deepcopy(vi_fonts[0]))
                    fill = etree.Element("{" + NS["a"] + "}solidFill")
                    etree.SubElement(fill, "{" + NS["a"] + "}srgbClr", val=ink)
                    props.insert(1 if props.find("a:ln", NS) is not None else 0, fill)
                for p in paragraphs:
                    ppr = p.find("a:pPr", NS)
                    if ppr is not None:
                        for node in list(ppr):
                            if etree.QName(node).localname in {"spcBef", "spcAft"}:
                                ppr.remove(node)
                bodypr = body.find("a:bodyPr", NS)
                if bodypr is not None:
                    bodypr.set("anchor", "t")
                if etree.tostring(body) != before:
                    node = shape.find("p:nvSpPr/p:cNvPr", NS)
                    changes.append({"slide": number, "shape_id": int(node.get("id")), "operation": "title_style"})
            changes.extend({'slide': number, 'shape_id': shape_id, 'operation': 'title_style'}
                           for shape_id in _separate_title_styles(root, deck.slide_width, deck.slide_height))
            if changes and changes[-1]["slide"] == number:
                replacements[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if not replacements:
        return {"path": str(path), "changes": []}
    folder = Path(directory) / "temp" / "slide-repair" / uuid4().hex
    folder.mkdir(parents=True)
    target = folder / "titles.pptx"
    _write_copy(path, target, replacements)
    return {"path": str(target), "changes": changes}


def _office_fit(path, target, targets, cancel):
    """Use the native layout engine; only selected text bodies will be retained."""
    import psutil
    import pythoncom
    import win32com.client

    existing = any((p.info["name"] or "").lower() == "powerpnt.exe" for p in psutil.process_iter(["name"]))
    pythoncom.CoInitialize()
    office = deck = None
    try:
        office = win32com.client.DispatchEx("PowerPoint.Application")
        deck = office.Presentations.Open(str(path.resolve()), ReadOnly=False, Untitled=False, WithWindow=False)
        for number, ids in targets.items():
            if cancel and cancel.is_set():
                raise ValueError("Đã dừng sửa bài giảng; bản trước vẫn được giữ.")
            for shape in deck.Slides(number).Shapes:
                if (shape.Id in ids and shape.Type in {1, 14, 17} and shape.Rotation == 0
                        and shape.HasTextFrame and shape.TextFrame2.HasText):
                    shape.TextFrame2.AutoSize = 2  # msoAutoSizeTextToFitShape
                    shape.TextFrame2.WordWrap = -1
                    text = shape.TextFrame2.TextRange
                    for _ in range(4):
                        if text.BoundHeight <= shape.Height + 1 and text.BoundWidth <= shape.Width + 1:
                            break
                        factor = min((shape.Height - 2) / max(text.BoundHeight, 1),
                                     (shape.Width - 2) / max(text.BoundWidth, 1), .95)
                        changed = False
                        # Preserve the VI/EN hierarchy and every run's other styling.
                        runs_id = text._oleobj_.GetIDsOfNames("Runs")
                        for index in range(text.Runs.Count, 0, -1):
                            # Runs is an indexed COM property; calling its default
                            # member returns text instead of a TextRange2 object.
                            run = win32com.client.Dispatch(text._oleobj_.Invoke(
                                runs_id, 0, pythoncom.DISPATCH_PROPERTYGET, True, index, 1))
                            size = float(run.Font.Size)
                            fitted = max(min(14, size), size * factor)
                            if size > 0 and fitted < size - .05:
                                run.Font.Size = fitted
                                changed = True
                        if not changed:
                            break
                    shape.TextFrame2.AutoSize = 0
        deck.SaveAs(str(target.resolve()), 24)  # ppSaveAsOpenXMLPresentation
    finally:
        if deck is not None:
            deck.Close()
        if office is not None and not existing and office.Presentations.Count == 0:
            office.Quit()
        deck = office = None
        pythoncom.CoUninitialize()


def repair_layout(path, directory, report, cancel=None, progress=None):
    """Change only overflowing, horizontal native text. Never rewrite prose."""
    path, directory = Path(path), Path(directory)
    if file_digest(path) != report["sha256"]:
        raise ValueError("Bài giảng đã thay đổi; kiểm tra lại trước khi tự sửa.")
    targets = {}
    for row in report["slides"]:
        ids = {i["shape_id"] for i in row["issues"]
               if i["code"] == "text_overflow" and "shape_id" in i}
        if ids:
            targets[row["slide"]] = ids
    if not targets or cancel and cancel.is_set():
        return {"path": str(path), "changes": []}
    folder = directory / "temp" / "slide-repair" / uuid4().hex
    folder.mkdir(parents=True)
    # Slide part numbers are not necessarily presentation order after edits.
    from .source_deck import source_slide_parts

    ordered_parts = source_slide_parts(path)
    parts = {number: ordered_parts[number - 1] for number in targets}
    replacements, roots = {}, {}
    with zipfile.ZipFile(path) as archive:
        for number, part in parts.items():
            root = etree.fromstring(archive.read(part))
            roots[number] = root
            for shape_id in targets[number]:
                shape = _shape(root, shape_id)
                if shape is None:
                    continue
                body = shape.find("p:txBody", NS)
                if body is None:
                    continue
                _compact_fixed_spacing(body)
                # Empty trailing paragraphs can double a bilingual box's height.
                paragraphs = body.findall("a:p", NS)
                while len(paragraphs) > 1:
                    last = paragraphs[-1]
                    if (last.xpath(".//a:t/text()", namespaces=NS) and
                            any(t.strip() for t in last.xpath(".//a:t/text()", namespaces=NS))
                            or last.xpath(".//a:br|.//a:fld|.//*[local-name()='oMath']", namespaces=NS)):
                        break
                    body.remove(last)
                    paragraphs.pop()
            replacements[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    trimmed, fitted, candidate = folder / "trimmed.pptx", folder / "office-fit.pptx", folder / "repaired.pptx"
    _write_copy(path, trimmed, replacements)
    try:
        _office_fit(trimmed, fitted, targets, cancel)
    except Exception as exc:
        if cancel and cancel.is_set():
            raise ValueError("Đã dừng sửa bài giảng; bản trước vẫn được giữ.")
        (folder / "unavailable.txt").write_text(type(exc).__name__ + ": " + str(exc), encoding="utf-8")
        return {"path": str(path), "changes": [], "unavailable": True}
    changes = []
    fitted_parts = source_slide_parts(fitted)
    with zipfile.ZipFile(fitted) as archive:
        for number, part in parts.items():
            office_root = etree.fromstring(archive.read(fitted_parts[number - 1]))
            for shape_id in targets[number]:
                original, adjusted = _shape(roots[number], shape_id), _shape(office_root, shape_id)
                if original is None or adjusted is None:
                    continue
                before_text = ["".join(p.xpath(".//a:t/text()", namespaces=NS)) for p in original.findall("p:txBody/a:p", NS)]
                after_text = ["".join(p.xpath(".//a:t/text()", namespaces=NS)) for p in adjusted.findall("p:txBody/a:p", NS)]
                if before_text != after_text or original.xpath(".//*[local-name()='oMath']"):
                    continue
                body = original.find("p:txBody/a:bodyPr", NS)
                office_body = adjusted.find("p:txBody/a:bodyPr", NS)
                if body is None or office_body is None or body.get("vert", "horz") != "horz":
                    continue
                original_body = original.find("p:txBody", NS)
                original.replace(original_body, copy.deepcopy(adjusted.find("p:txBody", NS)))
                changes.append({"slide": number, "shape_id": shape_id, "operation": "fit_text"})
            replacements[part] = etree.tostring(roots[number], xml_declaration=True, encoding="UTF-8", standalone=True)
    _write_copy(path, candidate, replacements)
    (folder / "changes.json").write_text(json.dumps(changes, indent=2), encoding="utf-8")
    return {"path": str(candidate), "changes": changes}


def repair_and_inspect(path, directory, *, source=None, mode="", style="source", cancel=None, progress=None,
                       reuse_directory=None, layout_plan=None):
    """At most three local passes; keep the best verified copy and stop on no progress."""
    from .chatgpt_handoff import inspect_returned_deck
    from .slide_quality import render_quality

    current = Path(path)
    inspection = inspect_returned_deck(current, source, mode, style, layout_plan=layout_plan)
    render_quality(current, directory, inspection["slide_quality"], cancel, progress, reuse_directory=reuse_directory)
    history = []
    if not layout_plan and not (cancel and cancel.is_set()):
        title_result = normalize_titles(current, directory, mode)
        if Path(title_result["path"]) != current:
            titled = inspect_returned_deck(title_result["path"], source, mode, style)
            render_quality(title_result["path"], directory, titled["slide_quality"], cancel, progress)
            if (failure_keys(titled["slide_quality"]) <= failure_keys(inspection["slide_quality"])
                    and all(r["rendered"] for r in titled["slide_quality"]["slides"])):
                current, inspection = Path(title_result["path"]), titled
                history.extend(title_result["changes"])
    for _ in range(3):
        if cancel and cancel.is_set():
            break
        before = failure_keys(inspection["slide_quality"])
        if not before:
            break
        result = repair_layout(current, directory, inspection["slide_quality"], cancel, progress)
        candidate = Path(result["path"])
        if candidate == current:
            break
        after = inspect_returned_deck(candidate, source, mode, style, layout_plan=layout_plan)
        render_quality(candidate, directory, after["slide_quality"], cancel, progress)
        remaining = failure_keys(after["slide_quality"])
        changes=result['changes']
        if not remaining < before:
            selected,pages=retain_improved_pages(current,candidate,inspection['slide_quality'],after['slide_quality'],directory)
            if selected is None:
                break
            candidate=selected
            after=inspect_returned_deck(candidate, source, mode, style, layout_plan=layout_plan)
            render_quality(candidate,directory,after['slide_quality'],cancel,progress)
            remaining=failure_keys(after['slide_quality'])
            changes=[row for row in changes if row['slide'] in pages]
        if not remaining < before or any(not r['rendered'] for r in after['slide_quality']['slides']):
            break  # Never publish a regression or a page that has not been rendered.
        history.extend(changes)
        current, inspection = candidate, after
    inspection["slide_quality"]["local_repairs"] = history
    return {"path": str(current), "inspection": inspection, "changes": history}
