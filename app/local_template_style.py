"""Three bundled classroom themes, built locally without generated content."""

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

THEMES = {
    'standard': {'label': 'Chuẩn lớp học', 'ink': '112650', 'accent': '0869F9', 'paper': 'F4F8FE'},
    'visual': {'label': 'Trực quan', 'ink': '123F44', 'accent': '087F83', 'paper': 'EFF8F6'},
    'practice': {'label': 'Luyện tập', 'ink': '49331E', 'accent': 'AC541A', 'paper': 'FFF8EF'},
}


def decorate(deck, config):
    theme = THEMES[config.get('preset', 'standard')]
    ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    for number, slide in enumerate(deck.slides, 1):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['paper'])
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.12), deck.slide_height)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(theme['accent'])
        stripe.line.fill.background()
        # Place decorations behind every text/image, leaving source evidence visible.
        slide.shapes._spTree.remove(stripe._element)
        slide.shapes._spTree.insert(2, stripe._element)
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.color.rgb = RGBColor.from_string(theme['ink'])
                if shape.top < Inches(1.3):
                    paragraph.font.bold = not paragraph.font.italic
        if config['conversion_format'] == 'parallel_columns':
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.55), Inches(1.95),
                                          Inches(.015), Inches(4.75))
            rule.fill.solid()
            rule.fill.fore_color.rgb = RGBColor.from_string(theme['accent'])
            rule.line.fill.background()
            slide.shapes._spTree.remove(rule._element)
            slide.shapes._spTree.insert(2, rule._element)
        footer = slide.shapes.add_textbox(Inches(.55), Inches(7.22), Inches(12.2), Inches(.2))
        footer.name = 'BCL_DECORATION footer'
        paragraph = footer.text_frame.paragraphs[0]
        label = {'standard':'Classroom', 'visual':'Visual', 'practice':'Practice'}[config.get('preset','standard')] if config['conversion_format']=='english_only' else theme['label']
        paragraph.text = f"BiliClass · {label}                                     {number:02d} / {len(deck.slides):02d}"
        paragraph.font.name, paragraph.font.size = 'Arial', Pt(9)
        paragraph.font.color.rgb = RGBColor.from_string(theme['accent'])
        footer.text_frame.margin_top = footer.text_frame.margin_bottom = 0
        transition = etree.Element('{' + ns + '}transition', spd='fast', advClick='1')
        etree.SubElement(transition, '{' + ns + '}fade')
        following = slide._element.find('{' + ns + '}timing')
        if following is None:
            following = slide._element.find('{' + ns + '}extLst')
        if following is None:
            slide._element.append(transition)
        else:
            following.addprevious(transition)
