"""Import an existing PowerPoint for narration without translation or redesign."""

import json
import re
from contextlib import closing
from pathlib import Path

from .chatgpt_handoff import external_preview, inspect_returned_deck
from .library import Library, RevisionConflict, now
from .narration_script import segments as narration_segments


def script_speech(lesson):
    """Explicit slide narration, including notes, is independent of native builds."""
    from .powerpoint import slide_for_locator

    return {'schema_version': 1, 'sha256': lesson['source']['sha256'], 'slides': [
        {'slide': slide_for_locator(segment['locator']), 'slide_id': index + 1,
         'issues': [], 'speech_status': 'PASS', 'englishSpeechSegments': [
             {'id': 'script-' + segment['id'], 'slide': slide_for_locator(segment['locator']),
              'slide_id': index + 1, 'english_text': segment['en'], 'status': 'ready',
              'visibility': {'kind': 'always'}, 'kind': 'body', 'order': 0,
              'source': {'source_type': 'narration_script'}, 'x': 0, 'y': 0}
         ] if segment['en'].strip() else []}
        for index, segment in enumerate(narration_segments(lesson))]}


def prepare_script(library, lesson):
    if lesson.get('voice_only') and not lesson.get('narration_script_version'):
        lesson['source_slide_speech'] = lesson.get('slide_speech', {})
        lesson['slide_speech'] = script_speech(lesson)
        lesson['narration_script_version'] = 1
        library._write(lesson)
    return lesson


def save_script(library, lesson_id, segment_id, revision, vi, en, *, silent=None):
    """Edit narration metadata; never modify the source PowerPoint or approve it."""
    from .chatgpt_handoff import text_signature

    if silent is not None and type(silent) is not bool:
        raise ValueError('Lựa chọn slide không đọc chưa hợp lệ.')
    lesson = library.get(lesson_id)
    if lesson['revision'] != revision:
        raise RevisionConflict('Bài vừa thay đổi. Đóng và mở lại lời đọc trước khi lưu.')
    if not isinstance(vi, str) or not isinstance(en, str) or len(vi) + len(en) > 100_000:
        raise ValueError('Mỗi slide tối đa 100.000 ký tự lời đọc.')
    # Validate the original file before updating the metadata associated with it.
    external_preview(lesson, library.directory)
    segment = next((s for s in lesson['segments'] if s['id'] == segment_id), None)
    if segment is None:
        raise ValueError('Không tìm thấy slide cần sửa.')
    vi, en = vi.strip(), en.strip()
    if sum(len(s['vi']) + len(s['en']) for s in narration_segments(lesson) if s['id'] != segment_id) + len(vi) + len(en) > 1_000_000:
        raise ValueError('Tổng lời đọc của bài tối đa 1.000.000 ký tự.')
    effective = next(s for s in narration_segments(lesson) if s['id'] == segment_id)
    was_silent = effective.get('silent', False)
    if (effective.get('original_vi', effective['vi']), effective.get('original_en', effective['en'])) == (vi, en) and (
            silent is None or silent == was_silent):
        return lesson
    if silent is not None:
        identities = set(lesson.get('narration_silent', []))
        identities.add(segment_id) if silent else identities.discard(segment_id)
        lesson['narration_silent'] = sorted(identities)
    if lesson.get('voice_only'):
        segment.update(vi=vi, en=en, approved=False)
        lesson.setdefault('narration_script', {}).setdefault(segment_id, {})['manual'] = True
    else:
        lesson.setdefault('narration_script', {})[segment_id] = {'vi': vi, 'en': en, 'manual': True}
    lesson.update(revision=lesson['revision'] + 1, updated_at=now(), narration_script_version=2)
    lesson['narration_classification_version'] = 1
    from .narration_script import refresh_speech
    from .powerpoint import slide_for_locator

    refresh_speech(lesson, edited=(slide_for_locator(segment['locator']),))
    lesson['external_deck']['text_signature'] = text_signature(lesson)
    lesson['external_deck']['reviewed_revision'] = 0
    lesson.pop('prepared', None)
    library._write(lesson)
    return lesson


def prepare_library_lesson(directory, identity, profile=None, *, complete=False, cancel=None):
    from .quick_conversion import build_preview

    with closing(Library(directory)) as library:
        lesson = library.get(identity)
        if lesson.get('deleted_at'):
            raise ValueError('Khôi phục bài từ thùng rác trước khi tạo giọng.')
        lesson = prepare_script(library, lesson)
        from .narration_script import prepare

        lesson = prepare(library, lesson)
        lesson = refresh_native_speech(library, lesson, cancel)
        if complete:
            from .lesson_completion import complete_existing

            lesson = prepare(library, complete_existing(library, lesson, cancel))
        if lesson.get('external_deck'):
            return lesson, external_preview(lesson, directory)
        if all(s['vi'].strip() and s['en'].strip() for s in lesson['segments']):
            return lesson, build_preview(lesson, directory, profile)
        source = lesson.get('source') or {}
        if source.get('file', '').lower().endswith('.pptx'):
            from .powerpoint import verified_presentation

            path = verified_presentation(lesson, directory)
        else:
            raise ValueError('Bài này chưa có PowerPoint hoặc chưa đủ nội dung. Hoàn tất bài hoặc chọn file PowerPoint có sẵn.')
    return import_presentation(directory, path)


def refresh_native_speech(library, lesson, cancel=None):
    """Upgrade derived speech once, preserving explicit scripts and valid WAVs."""
    import copy

    from pptx import Presentation

    from .lesson_completion import _retain_unchanged_audio
    from .narration_script import refresh_speech
    from .powerpoint import verified_presentation
    from .slide_speech import EXTRACTION_VERSION, deck_speech

    base = lesson.get('source_slide_speech') or lesson.get('slide_speech', {})
    if lesson.get('voice_only') or not base or base.get('extraction_version') == EXTRACTION_VERSION:
        return lesson
    if cancel and cancel.is_set():
        raise InterruptedError('Đã dừng chuẩn bị lời đọc.')
    before = copy.deepcopy(lesson)
    path = verified_presentation(lesson, library.directory)
    data = deck_speech(path, Presentation(path), lesson.get('conversion_format', ''))
    if data['sha256'] != lesson['source']['sha256']:
        raise ValueError('PowerPoint đã thay đổi trong lúc đọc. Mở lại bài giảng.')
    if cancel and cancel.is_set():
        raise InterruptedError('Đã dừng chuẩn bị lời đọc.')
    lesson['source_slide_speech'] = data
    refresh_speech(lesson)
    _retain_unchanged_audio(before, lesson)
    library.db.execute('BEGIN IMMEDIATE')
    try:
        if library.get(lesson['id']) != before:
            raise RevisionConflict('Bài vừa thay đổi. Mở lại để chuẩn bị lời đọc mới.')
        library._write(lesson)
    except BaseException:
        library.db.rollback()
        raise
    return lesson


def import_presentation(directory, path):
    path = Path(path)
    inspection = inspect_returned_deck(path)
    # Explicit narration notes may contain only one language. The bilingual
    # conversion inspector intentionally requires pairs; this workspace does not.
    from pptx import Presentation

    for slide, unit in zip(Presentation(path).slides, inspection['profile']['units'], strict=True):
        if not slide.has_notes_slide:
            continue
        parts, language = {'vi': [], 'en': []}, None
        for line in slide.notes_slide.notes_text_frame.text.splitlines():
            match = re.match(r'^\s*(VI|VN|EN|English|Tiếng Việt|Tiếng Anh)\s*[:：]\s*(.*)$', line, re.I)
            if match:
                language = 'vi' if match[1].casefold() in {'vi', 'vn', 'tiếng việt'} else 'en'
                parts[language].append(match[2])
            elif re.match(r'^\s*(CHECK|QUIZ|SPEECH|BILICLASS_NOTES)\s*[:：]', line, re.I):
                language = None
            elif language:
                parts[language].append(line)
        if any(parts.values()):
            vi, en = ('\n'.join(parts[lang]).strip() for lang in ('vi', 'en'))
            if len(vi) + len(en) > 100_000:
                raise ValueError('Mỗi slide tối đa 100.000 ký tự lời đọc.')
            unit.update(vi=vi, en=en)
    if sum(len(u['vi']) + len(u['en']) for u in inspection['profile']['units']) > 1_000_000:
        raise ValueError('Tổng lời đọc của bài tối đa 1.000.000 ký tự.')
    with closing(Library(directory)) as library:
        source = library.store_source(path)
        from .export_origin import bind_lesson, remember_source

        remember_source(library, path, source['sha256'])
        if source['sha256'] != inspection['slide_speech']['sha256']:
            raise ValueError('PowerPoint đã thay đổi trong lúc đọc. Chọn lại file trước khi tạo giọng.')
        # Reuse an earlier import of the exact file, including its voice variants.
        for row in library.db.execute('SELECT content FROM lessons'):
            lesson = json.loads(row[0])
            if lesson.get('voice_only') and (lesson.get('source') or {}).get('sha256') == source['sha256'] and not lesson.get('deleted_at'):
                from .narration_script import prepare

                lesson = prepare(library, prepare_script(library, lesson))
                bind_lesson(library, lesson, source['sha256'])
                return lesson, external_preview(lesson, directory)
        config = {'title': path.stem, 'subject': 'Bài giảng có sẵn', 'education_level': '', 'grade': '',
                  'level': 4, 'layout': 'split_view', 'style': 'source', 'preset': 'standard', 'mode': 'preserve'}
        lesson = library.create_external_lesson(config, source, inspection, voice_only=True)
        bind_lesson(library, lesson, source['sha256'])
        prepare_script(library, lesson)
        from .narration_script import prepare

        lesson = prepare(library, lesson)
        return lesson, external_preview(lesson, directory)
