"""Embed existing, verified narration without asking a model or rewriting artwork."""

import hashlib
import io
import wave
import zipfile
from pathlib import Path

from lxml import etree

from .lesson_narration import MODES, audio_status, valid_key
from .narration_script import segments as narration_segments
from .narration_timing import PARAGRAPH_GAP_MS
from .powerpoint import slide_for_locator
from .source_deck import NS, _bytes, _rels_path, _xml, source_slide_parts


def slide_audio(lesson, directory, *, scope='full_slide'):
    from . import speech
    from .slide_speech import speech_chunks

    if scope not in {'full_slide', 'visible'}:
        raise ValueError('Chọn phạm vi lời đọc hợp lệ.')
    status = audio_status(lesson, directory)
    if not status["ready"]:
        raise ValueError("Tạo đủ giọng đọc đã chọn trước khi tải PowerPoint kèm giọng.")
    record = lesson["narration"]
    pages = {}
    for page in lesson.get("slide_speech", {}).get("slides", []):
        segments = page["englishSpeechSegments"]
        safe = [
            s for s in segments if s["status"] == "ready" and (
                scope == 'full_slide' or s.get("visibility", {}).get("kind") == "always")
        ]
        languages = {}
        if "en" in MODES[record["mode"]]:
            languages["en"] = [part for s in safe for part in speech_chunks(s["english_text"])]
        # The explicit whole-slide button reads the saved script on demand.
        # The optional build-aware route excludes words hidden behind clicks.
        if "vi" in MODES[record["mode"]] and (scope == 'full_slide' or not any(
            s.get("visibility", {}).get("kind") != "always" for s in segments if s["english_text"]
        )):
            languages["vi"] = [
                part
                for unit in narration_segments(lesson)
                if slide_for_locator(unit.get("locator", "")) == page["slide"]
                for part in speech_chunks(unit["vi"])
            ]
        for language, texts in languages.items():
            keys = [speech.audio_key(text, speech.voice_for(record["voice"], language), record["rate"]) for text in texts]
            if not keys:
                continue
            if any(key not in record["keys"][language] or not valid_key(directory, key) for key in keys):
                raise ValueError("Giọng đọc không khớp bài đang xem. Tạo lại giọng trước khi tải.")
            pages.setdefault(page["slide"], {})[language] = keys
    if not pages:
        raise ValueError(
            "Chưa có lời đọc tĩnh phù hợp để nhúng; phần theo hiệu ứng vẫn dùng trong BiliClass."
        )
    return {number: {lang: languages[lang] for lang in ('vi', 'en') if lang in languages}
            for number, languages in pages.items()}


def static_vietnamese(slide, lesson, number):
    from .slide_speech import animation_rules, normalized, objects

    rules = animation_rules(slide)
    visible = normalized(' '.join(item['text'] for item in objects(slide)
                        if not item['hidden'] and not any(str(parent) in rules for parent in item['ancestors'])))
    texts = [normalized(unit['vi']) for unit in narration_segments(lesson)
             if slide_for_locator(unit.get('locator', '')) == number and unit['vi'].strip()]
    return bool(texts) and all(text in visible for text in texts)


def joined_wav(directory, keys):
    buffer = io.BytesIO()
    parameters = None
    total = 0
    with wave.open(buffer, "wb") as output:
        for key in keys:
            with wave.open(str(Path(directory) / "audio" / (key + ".wav")), "rb") as source:
                current = (source.getnchannels(), source.getsampwidth(), source.getframerate())
                if parameters is None:
                    parameters = current
                    output.setparams((*parameters, 0, "NONE", "not compressed"))
                if current != parameters or source.getcomptype() != "NONE":
                    raise ValueError("Các đoạn voice khác định dạng. Tạo lại giọng đọc của bài.")
                output.writeframes(source.readframes(source.getnframes()))
                silence = current[2] * PARAGRAPH_GAP_MS // 1000
                output.writeframes(b"\x00" * (silence * current[0] * current[1]))
                total += source.getnframes() + silence
    return buffer.getvalue(), round(total / parameters[2] * 1000)


def coverage_message(report):
    message = f"Đã nhúng {report.get('speakers', 0)} nút loa trên {report.get('slides', 0)} slide. Bấm F5 rồi bấm loa để nghe."
    if report.get('unplaced_pages'):
        pages = ', '.join(str(n) for n in report['unplaced_pages'])
        message += f' Slide {pages} không đủ chỗ đặt loa; nghe giọng đã lưu bằng Mascot trong BiliClass.'
    if report.get('restricted_pages'):
        message += ' Phần lời phụ thuộc hiệu ứng hoặc chưa xác định khả năng hiển thị chỉ phát trong BiliClass.'
    if report.get('missing_pages'):
        message += f" {len(report['missing_pages'])} slide chưa có lời đọc cho ngôn ngữ đã chọn."
    if report.get('scope') == 'full_slide':
        message += ' Loa đọc toàn bộ lời đã lưu của slide, kể cả phần xuất hiện sau; không tự phát theo hiệu ứng.'
    if report.get('overlap_pages'):
        message += ' Kiểm tra khoảng trống chân trang ở slide: ' + ', '.join(map(str, report['overlap_pages'])) + '.'
    return message


def _node(parent, prefix, tag, **attributes):
    return etree.SubElement(parent, f"{{{NS[prefix]}}}{tag}", **{k: str(v) for k, v in attributes.items()})


def media_shape(identity, language, box, audio_relation, media_relation, icon_relation):
    picture = etree.Element(f"{{{NS['p']}}}pic")
    nonvisual = _node(picture, "p", "nvPicPr")
    props = _node(
        nonvisual,
        "p",
        "cNvPr",
        id=identity,
        name="BiliClass audio " + language.upper(),
        descr="Play " + language.upper(),
        hidden="0",
    )
    link = _node(props, "a", "hlinkClick", action="ppaction://media",
                 tooltip="Bấm để đọc toàn bộ lời đã lưu của slide · " + language.upper())
    link.set(f"{{{NS['r']}}}id", "")
    _node(_node(nonvisual, "p", "cNvPicPr"), "a", "picLocks", noChangeAspect="1")
    application = _node(nonvisual, "p", "nvPr")
    audio = _node(application, "a", "audioFile")
    audio.set(f"{{{NS['r']}}}link", audio_relation)
    extension = _node(
        _node(application, "p", "extLst"), "p", "ext", uri="{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}"
    )
    media = etree.SubElement(
        extension,
        "{http://schemas.microsoft.com/office/powerpoint/2010/main}media",
        nsmap={"p14": "http://schemas.microsoft.com/office/powerpoint/2010/main"},
    )
    media.set(f"{{{NS['r']}}}embed", media_relation)
    fill = _node(picture, "p", "blipFill")
    _node(fill, "a", "blip").set(f"{{{NS['r']}}}embed", icon_relation)
    _node(_node(fill, "a", "stretch"), "a", "fillRect")
    geometry = _node(picture, "p", "spPr")
    transform = _node(geometry, "a", "xfrm")
    _node(transform, "a", "off", x=round(box[0] * 12700), y=round(box[1] * 12700))
    _node(transform, "a", "ext", cx=round(box[2] * 12700), cy=round(box[3] * 12700))
    _node(_node(geometry, "a", "prstGeom", prst="rect"), "a", "avLst")
    return picture


def embed_audio(source, destination, lesson, directory, *, scope='full_slide', managed=False):
    from pptx import Presentation

    from .layout_regions import intersection
    from .presentation_controls import add_controls, picture
    from .slide_speech import objects
    from .speaker_artwork import artwork
    from .speaker_layout import control_box, default_position, fixed_slots, position

    pages = slide_audio(lesson, directory, scope=scope)
    deck = Presentation(source)
    parts = source_slide_parts(source)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    original_parts = set(entries)
    types = _xml(entries["[Content_Types].xml"])
    for extension, kind in [("wav", "audio/wav"), ("png", "image/png")]:
        if not any(n.get("Extension") == extension for n in types):
            etree.SubElement(types, f"{{{NS['ct']}}}Default", Extension=extension, ContentType=kind)
    count, embedded_pages, unplaced_pages, overlap_pages = 0, [], [], []
    selected_position = position(lesson.get('speaker_position', default_position(directory)))
    deck_languages = [lang for lang in ('vi', 'en', 'both') if any(lang in languages for languages in pages.values())]
    bounds = (deck.slide_width / 12700, deck.slide_height / 12700)
    slots = fixed_slots(bounds, selected_position, languages=deck_languages)
    missing_pages = sorted(set(range(1, len(parts) + 1)) - set(pages))
    restricted_pages = list(missing_pages) if scope == 'visible' else []
    for number, languages in pages.items():
        if not 1 <= number <= len(parts):
            raise ValueError("Số trang giọng đọc không khớp PowerPoint.")
        part = parts[number - 1]
        tree = _xml(entries[part])
        if scope == 'visible' and tree.find("p:timing", NS) is not None and not static_vietnamese(deck.slides[number - 1], lesson, number):
            languages = {lang: keys for lang, keys in languages.items() if lang not in {"vi", "both"}}
        if not languages:
            restricted_pages.append(number)
            continue  # Vietnamese notes may include answers hidden in native builds.
        if set(languages) != set(pages[number]) or set(languages) - {'both'} != set(MODES[lesson['narration']['mode']]):
            restricted_pages.append(number)
        shape_tree = tree.find("p:cSld/p:spTree", NS)
        relationships = _xml(entries[_rels_path(part)])

        def relation(kind, target):
            identity = "rIdBiliClass" + str(len(relationships) + 1)
            while any(r.get("Id") == identity for r in relationships):
                identity += "x"
            etree.SubElement(
                relationships, f"{{{NS['rel']}}}Relationship", Id=identity, Type=kind, Target=target
            )
            return identity

        occupied = []
        for item in objects(deck.slides[number - 1]):
            box = tuple(item[k] / 12700 for k in ("x", "y", "w", "h"))
            if not item["hidden"] and (item['source_type'] != 'image' or
                    box[2] * box[3] < deck.slide_width * deck.slide_height / 12700**2 * 0.6):
                occupied.append(box)
        if any(intersection(control_box(bounds, slots[lang]), box) for lang in languages for box in occupied):
            overlap_pages.append(number)
        embedded_pages.append(number)
        identity = max(int(v) for v in tree.xpath(".//p:cNvPr/@id", namespaces=NS))
        controls = []
        for language, keys in languages.items():
            audio, duration = joined_wav(directory, keys)
            audio_name = "ppt/media/biliclass-" + hashlib.sha256(audio).hexdigest() + ".wav"
            entries[audio_name] = audio
            audio_id = relation(NS["r"] + "/audio", "../media/" + Path(audio_name).name)
            media_id = relation(
                "http://schemas.microsoft.com/office/2007/relationships/media",
                "../media/" + Path(audio_name).name,
            )
            identity += 1
            icon, crop = artwork(language)
            icon_name = 'ppt/media/biliclass-' + hashlib.sha256(icon).hexdigest() + '.png'
            entries[icon_name] = icon
            icon_id = relation(NS['r'] + '/image', '../media/' + Path(icon_name).name)
            box = control_box(bounds, slots[language])
            media = media_shape(identity, language, box, audio_id, media_id, icon_id)
            shape_tree.append(media)
            shape_tree.append(picture(media, identity + 1, language, icon_id, crop=crop))
            pause_icon, pause_crop = artwork(language, paused=True)
            pause_name = 'ppt/media/biliclass-' + hashlib.sha256(pause_icon).hexdigest() + '.png'
            entries[pause_name] = pause_icon
            pause_id = relation(NS['r'] + '/image', '../media/' + Path(pause_name).name)
            shape_tree.append(picture(media, identity + 2, language, pause_id, paused=True, crop=pause_crop))
            # Keep the native player enabled, but make its backing poster
            # transparent and tiny underneath the visible speaker control.
            # Selection-pane hiding prevents reliable playback in Office.
            _node(media.find('p:blipFill/a:blip', NS), 'a', 'alphaModFix', amt='0')
            extent = media.find('p:spPr/a:xfrm/a:ext', NS)
            extent.set('cx', '12700')
            extent.set('cy', '12700')
            controls.append((identity, identity + 1, identity + 2, duration))
            identity += 2
            count += 1
        add_controls(tree, controls, managed=managed)
        entries[part] = _bytes(tree)
        entries[_rels_path(part)] = _bytes(relationships)
    if not count:
        raise ValueError(
            "Không đủ chỗ đặt loa hoặc lời đọc phụ thuộc hiệu ứng. Nghe trong BiliClass hoặc tải bản không giọng."
        )
    entries["[Content_Types].xml"] = _bytes(types)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as archive:
        for name, data in entries.items():
            # Raster/video assets are already compressed. Re-deflating them
            # costs seconds on large teaching decks without useful savings.
            compressed_media = Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.mp3', '.mp4', '.wmv', '.avi'}
            archive.writestr(name, data, compress_type=zipfile.ZIP_STORED if compressed_media else zipfile.ZIP_DEFLATED)
    # Every original resource remains in the output byte-for-byte.
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination) as exported:
        for name in original_parts:
            if name.startswith(("ppt/media/", "ppt/charts/", "ppt/embeddings/")) and original.read(
                name
            ) != exported.read(name):
                raise ValueError("Tài nguyên nguồn không khớp sau khi nhúng voice.")
    return {"speakers": count, "slides": len(embedded_pages), "embedded_pages": embedded_pages, "scope": scope,
            "unplaced_pages": unplaced_pages, "restricted_pages": sorted(set(restricted_pages)),
            "missing_pages": missing_pages if scope == 'full_slide' else [],
            'speaker_position': selected_position, 'overlap_pages': overlap_pages}
