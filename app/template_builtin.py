"""Editable local classroom templates, shared by preview and conversion."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .local_template_style import THEMES
from .paths import RESOURCE_ROOT


def create(path, preset):
    theme = THEMES[preset]
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def panel(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        shape.name = 'BCL_DECORATION panel'
        return shape

    def text(slide, value, name, x, y, w, h, size, color, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        shape.text_frame.word_wrap = True
        shape.text_frame.margin_left = shape.text_frame.margin_right = 0
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
        shape.text = value
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.font.name = 'Arial'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        return shape

    for index, role in enumerate(('COVER', 'INTRO', 'CONTENT', 'CLOSING'), 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['paper'])
        accent = theme['accent']
        panel(slide, .32, .35, .1, 6.6, accent)
        text(slide, 'BiliClass', 'BCL_DECORATION brand', .7, .35, 4, .3, 13, accent, True)
        text(slide, f'{index:02}', 'BCL_DECORATION page', 11.9, 6.95, .6, .25, 10, accent)
        panel(slide, .7, 6.75, 11.9, .025, 'CDD9E8')
        title_y, title_h = (1.35, 2.05) if role in {'COVER', 'CLOSING'} else (.95, 1.15)
        title_w = 8.2 if role in {'COVER', 'CLOSING'} else 11.8
        text(slide, '[LESSON_TITLE]', 'BCL_TITLE', .9, title_y, title_w, title_h, 32, theme['ink'], True)
        body_y, body_h = (3.65, 1.9) if role in {'COVER', 'CLOSING'} else (2.6, 3.7)
        if role not in {'COVER', 'CLOSING'}:
            panel(slide, .7, body_y - .2, 11.9, body_h + .4, 'FFFFFF')
        text(slide, '[SOURCE_CONTENT]', 'BCL_CONTENT', 1, body_y, 11.2, body_h, 24, theme['ink'])
        if role in {'COVER', 'CLOSING'}:
            logo = slide.shapes.add_picture(str(RESOURCE_ROOT/'assets/brand/biliclass.png'),
                                            Inches(10.4), Inches(1.35), width=Inches(1.8))
            logo.name = 'BCL_DECORATION logo'
        transition = OxmlElement('p:transition')
        transition.set('spd', 'med')
        transition.set('advClick', '1')
        transition.append(OxmlElement('p:fade'))
        slide._element.append(transition)
    deck.save(path)
    return path
