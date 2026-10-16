"""Content-led frames inside the supplied Universal themes."""

from collections import OrderedDict

from pptx.util import Inches, Pt

COMMON = {'TWO_COLUMNS', 'THEORY', 'STACKED', 'ENGLISH', 'FORMULA', 'IMAGE_TEXT'}


def prepare(slide, layout, slots, mode):
    if layout not in COMMON or not slots:
        return slots
    key, pair = next(iter(slots.items()))
    for extra in list(slots.values())[1:]:
        for shape in extra.values():
            slide.shapes._spTree.remove(shape._element)
    for language, shape in pair.items():
        english = language == 'en'
        if layout == 'IMAGE_TEXT':
            box = (8.0, 1.85, 4.5, 4.65) if mode == 'english_only' else (8.0, 4.1 if english else 1.85, 4.5, 2.15)
        elif layout == 'FORMULA':
            box = (.8, 4.75, 11.7, 1.8) if mode == 'english_only' else (7.0 if english else .8, 4.75, 5.55, 1.8)
        elif mode == 'english_only':
            box = (.8, 1.8, 11.7, 4.9)
        elif mode == 'sentence_pairs' or layout == 'STACKED':
            box = (.8, 4.3 if english else 1.8, 11.7, 2.2)
        else:
            box = (7.0 if english else .8, 1.8, 5.55, 4.9)
        shape.left, shape.top, shape.width, shape.height = (Inches(n) for n in box)
        frame = shape.text_frame
        # Inline Office math and italic glyphs overhang the nominal line width.
        # Give them real padding rather than accepting a clipped right edge.
        frame.margin_left = frame.margin_right = Pt(6)
        frame.margin_top = frame.margin_bottom = Pt(6)
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(22 if english else 24)
            for run in paragraph.runs:
                run.font.size = paragraph.font.size
    return OrderedDict([(key, pair)])


def finish(slide, layout, native, assigned):
    if layout not in COMMON:
        return None
    # Drop the template's vacant cards and their captions, not actual content.
    # Original branding/header/footer and populated native text remain editable.
    for shape in list(slide.shapes):
        if (Inches(1.65) <= shape.top < Inches(6.95)
                and not shape.name.startswith(('BCL_FILLED ', 'BCL_SOURCE '))):
            slide.shapes._spTree.remove(shape._element)
    if not native:
        return None
    if layout == 'FORMULA':
        return tuple(Inches(n) for n in ((.9, 1.85, 11.5, 2.6) if assigned else (.9, 1.9, 11.5, 4.65)))
    if layout == 'IMAGE_TEXT':
        return tuple(Inches(n) for n in ((.8, 1.85, 6.75, 4.65) if assigned else (.8, 1.85, 11.7, 4.8)))
    return None


def equation_pages(objects):
    """Stack related standalone equations without stretching source coordinates."""
    pages, current, top = [], [], 0
    limit, gap = Inches(4.5), Pt(18)
    for obj in objects:
        height = obj['h']
        if current and (top + height > limit or len(current) >= 4):
            pages.append(current)
            current, top = [], 0
        current.append({**obj, 'x': 0, 'y': top})
        top += height + gap
    if current:
        pages.append(current)
    return pages
