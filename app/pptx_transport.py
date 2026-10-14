"""Keep large native media local; restore exact original bytes after web editing."""

import base64
import hashlib
import io
import json
import struct
import wave
import zipfile
from pathlib import Path

from .source_deck import NS

MP4_PLACEHOLDER = base64.b64decode('AAAAHGZ0eXBpc29tAAACAGlzb21pc28ybXA0MQAAAAhmcmVlAAAAGW1kYXQAAAGzABAHAAABthYHGCPbfgAAAzBtb292AAAAbG12aGQAAAAAAAAAAAAAAAAAAAPoAAAD6AABAAABAAAAAAAAAAAAAAAAAQAAAAAAAAAAAAAAAAAAAAEAAAAAAAAAAAAAAAAAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACAAACW3RyYWsAAABcdGtoZAAAAAMAAAAAAAAAAAAAAAEAAAAAAAAD6AAAAAAAAAAAAAAAAAAAAAAAAQAAAAAAAAAAAAAAAAAAAAEAAAAAAAAAAAAAAAAAAEAAAAAAEAAAABAAAAAAACRlZHRzAAAAHGVsc3QAAAAAAAAAAQAAA+gAAAAAAAEAAAAAAdNtZGlhAAAAIG1kaGQAAAAAAAAAAAAAAAAAAEAAAABAAFXEAAAAAAAtaGRscgAAAAAAAAAAdmlkZQAAAAAAAAAAAAAAAFZpZGVvSGFuZGxlcgAAAAF+bWluZgAAABR2bWhkAAAAAQAAAAAAAAAAAAAAJGRpbmYAAAAcZHJlZgAAAAAAAAABAAAADHVybCAAAAABAAABPnN0YmwAAADac3RzZAAAAAAAAAABAAAAym1wNHYAAAAAAAAAAQAAAAAAAAAAAAAAAAAAAAAAEAAQAEgAAABIAAAAAAAAAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAY//8AAABgZXNkcwAAAAADgICATwABAASAgIBBIBEAAAAAAAIAAAAAiAWAgIAvAAABsAEAAAG1iRMAAAEAAAABIADEjYgADQCEAhRjAAABskxhdmM2MS4xOS4xMDAGgICAAQIAAAAUYnRydAAAAAAAAAIAAAAAiAAAABhzdHRzAAAAAAAAAAEAAAABAABAAAAAABxzdHNjAAAAAAAAAAEAAAABAAAAAQAAAAEAAAAUc3RzegAAAAAAAAARAAAAAQAAABRzdGNvAAAAAAAAAAEAAAAsAAAAYXVkdGEAAABZbWV0YQAAAAAAAAAhaGRscgAAAAAAAAAAbWRpcmFwcGwAAAAAAAAAAAAAAAAsaWxzdAAAACSpdG9vAAAAHGRhdGEAAAABAAAAAExhdmY2MS43LjEwMA==')


def _placeholder(name, data):
    digest = hashlib.sha256(data).hexdigest()
    tag = ('BILICLASS_LOCAL_MEDIA:'+digest).encode('ascii')
    suffix = Path(name).suffix.casefold()
    if suffix in {'.mp4','.m4v','.mov'}:
        if not MP4_PLACEHOLDER:
            raise ValueError('Thiếu mẫu media nhẹ đã kiểm chứng.')
        return MP4_PLACEHOLDER + struct.pack('>I',len(tag)+8)+b'free'+tag
    if suffix == '.wav' and data.startswith(b'RIFF'):
        stream = io.BytesIO()
        with wave.open(stream,'wb') as audio:
            audio.setparams((1,2,16000,0,'NONE','not compressed'))
            audio.writeframes(b'\0'*320)
        result = stream.getvalue()+b'JUNK'+struct.pack('<I',len(tag))+tag
        if len(tag)%2:
            result+=b'\0'
        return result[:4]+struct.pack('<I',len(result)-8)+result[8:]
    return None


def media_references(path, restore=None):
    """Count media actually used by shapes, not orphaned files in the ZIP."""
    from pptx import Presentation

    from .native_bilingual import leaves

    restore = restore or {}
    rows = []
    for slide in Presentation(path).slides:
        row = []
        for shape in leaves(slide.shapes):
            found = set()
            nodes = shape._element.xpath('.//*[local-name()="videoFile" or local-name()="audioFile" or local-name()="media"]')
            for node in nodes:
                for name,value in node.attrib.items():
                    if not name.startswith('{'+NS['r']+'}') or value not in slide.part.rels:
                        continue
                    relation = slide.part.rels[value]
                    if relation.is_external:
                        continue
                    digest = hashlib.sha256(relation.target_part.blob).hexdigest()
                    found.add(restore.get(digest,digest))
            if found:
                row.append(sorted(found))
        rows.append(sorted(row))
    return rows


def verify_media(path, expected, restore):
    if media_references(path,restore) != expected:
        raise ValueError('Video/âm thanh của một nhóm slide đã thiếu hoặc đổi; giữ bản gốc để xử lý lại.')


def proxy(source, destination):
    source,destination = Path(source),Path(destination)
    if source.resolve()==destination.resolve():
        raise ValueError('Không ghi đè tài liệu đầy đủ khi tạo bản làm việc nhẹ.')
    mappings,removed = {},0
    pending = destination.with_suffix('.proxy-pending.pptx')
    destination.parent.mkdir(parents=True,exist_ok=True)
    try:
        with zipfile.ZipFile(source) as before,zipfile.ZipFile(pending,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=1) as after:
            for item in before.infolist():
                data = before.read(item)
                replacement = _placeholder(item.filename,data) if item.filename.startswith('ppt/media/') and len(data)>256*1024 else None
                if replacement is not None:
                    mappings[hashlib.sha256(replacement).hexdigest()] = hashlib.sha256(data).hexdigest()
                    removed += len(data)-len(replacement)
                    data = replacement
                after.writestr(item.filename,data)
        pending.replace(destination)
    finally:
        pending.unlink(missing_ok=True)
    result = {'version':1,'restore':mappings,'offloaded_bytes':removed,'media_references':media_references(source),
              'original_bytes':source.stat().st_size,'upload_bytes':destination.stat().st_size}
    destination.with_suffix('.transport.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result
