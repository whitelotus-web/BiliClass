"""Image translations stay visible, but are not automatic slide narration."""

import copy
import json
import re
from functools import lru_cache

PREFIX = 'BILICLASS_IMAGE_TEXT:'


def mark_template_text(shape, pairs):
    """Keep paragraph provenance while a known template frame is filled."""
    rows, offset = [], 0
    for locator, vi, en in pairs:
        lines = en.split('\n')
        if ' · chữ trong ảnh' in locator:
            rows.extend({'paragraph': offset + index, 'vi': '', 'en': line, 'in_notes': False}
                        for index, line in enumerate(lines) if line.strip())
        offset += len(lines)
    if rows:
        shape._element.xpath('./p:nvSpPr/p:cNvPr')[0].set('descr', PREFIX + json.dumps({'rows': rows}, ensure_ascii=False))


def image_text(items, notes):
    """Use our own shape provenance; never infer OCR from an ordinary caption."""
    pairs, current = [], {}
    language = None
    for line in notes.splitlines():
        match = re.match(r'^\s*(VI|EN)\s*[:：]\s*(.*)$', line, re.I)
        if match:
            language = match[1].lower()
            if language == 'vi' and current:
                pairs.append(current)
                current = {}
            current.setdefault(language, []).append(match[2])
        elif re.match(r'^\s*(CHECK|QUIZ|SPEECH|BILICLASS_NOTES)\s*[:：]', line, re.I):
            language = None
        elif language:
            current[language].append(line)
    if current:
        pairs.append(current)
    pairs = [{lang: '\n'.join(row.get(lang, [])).strip() for lang in ('vi', 'en')} for row in pairs]
    result, seen = [], set()
    for item in items:
        description = item['shape']._element.xpath('./p:nvSpPr/p:cNvPr/@descr')
        if description and description[0].startswith(PREFIX):
            try:
                declared = json.loads(description[0][len(PREFIX):])
                rows = declared.get('rows', []) if isinstance(declared, dict) else []
                for row in rows[:1000] if isinstance(rows, list) else []:
                    if (isinstance(row, dict) and row.get('paragraph') == item.get('paragraph')
                            and isinstance(row.get('en'), str) and row['en'].strip() == item['text'].strip()):
                        result.append({**row, 'shape_id': item['shape_id']})
            except (ValueError, TypeError):
                pass
        if item['name'] != 'BCL_EN image translation' or item['shape_id'] in seen:
            continue
        seen.add(item['shape_id'])
        text = item['shape'].text.strip()
        row = next((row for row in reversed(pairs) if row['en'] == text), {'vi': '', 'en': text})
        if description and description[0].startswith(PREFIX):
            try:
                declared = json.loads(description[0][len(PREFIX):])
                if (isinstance(declared, dict) and declared.get('en') == text
                        and isinstance(declared.get('vi'), str) and len(declared['vi']) <= 100_000):
                    row = declared
            except (ValueError, TypeError):
                pass
        result.append({'shape_id': item['shape_id'], 'vi': row['vi'], 'en': text})
    return result


def without_image_suffix(text, rows, language):
    """Remove only appended OCR note blocks; identical body text is preserved."""
    result = text.rstrip()
    seen = set()
    for row in reversed(rows):
        if row.get('in_notes') is False:
            continue
        pair = row.get('vi', ''), row.get('en', '')
        if pair in seen:
            continue
        seen.add(pair)
        suffix = row.get(language, '').strip()
        if suffix and (result == suffix or result.endswith('\n' + suffix)):
            result = result[:-len(suffix)].rstrip()
    return result


@lru_cache(maxsize=32)
def _legacy_pages(path, size, modified, digest):
    from pptx import Presentation

    from .slide_speech import objects

    return {str(index): rows for index, slide in enumerate(Presentation(path).slides, 1)
            if (rows := image_text(objects(slide), slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''))}


def prepare(lesson, library):
    """Upgrade old narration metadata without changing the PPTX or lesson text."""
    source_hash = (lesson.get('source') or {}).get('sha256')
    checked_hash = lesson.get('narration_image_source_sha256')
    if lesson.get('narration_image_policy') == 1 and (not checked_hash or checked_hash == source_hash):
        return False
    pages = {str(page['slide']): page['imageText'] for page in lesson.get('slide_speech', {}).get('slides', [])
             if page.get('imageText')}
    if not pages and (lesson.get('source') or {}).get('file', '').lower().endswith('.pptx'):
        from .powerpoint import verified_presentation

        try:
            path = verified_presentation(lesson, library.directory)
            stat = path.stat()
            pages = _legacy_pages(str(path), stat.st_size, stat.st_mtime_ns, lesson['source']['sha256'])
        except (OSError, ValueError):
            return False
    if not pages:
        # A successful scan with no image labels is a result too. Remember it
        # instead of reparsing the entire presentation on every subsequent open.
        lesson['narration_image_policy'] = 1
        lesson['narration_image_source_sha256'] = source_hash
        return True
    lesson['narration_image_text'] = copy.deepcopy(pages)
    lesson['narration_image_policy'] = 1
    lesson['narration_image_source_sha256'] = source_hash
    for key in ('slide_speech', 'source_slide_speech'):
        for page in lesson.get(key, {}).get('slides', []):
            rows = pages.get(str(page['slide']), [])
            if rows:
                page['imageText'] = copy.deepcopy(rows)
                page['englishSpeechSegments'] = [s for s in page.get('englishSpeechSegments', [])
                                                 if not any(s.get('source', {}).get('shape_id') == row['shape_id']
                                                            and (row.get('paragraph') is None or
                                                                 row['paragraph'] == s.get('source', {}).get('paragraph'))
                                                            for row in rows)]
    from .narration_script import refresh_speech

    scripted = [page['slide'] for page in lesson.get('slide_speech', {}).get('slides', [])
                if str(page['slide']) in pages and any(s.get('source', {}).get('source_type') == 'narration_script'
                                                      for s in page.get('englishSpeechSegments', []))]
    refresh_speech(lesson, edited=scripted)
    return True
