"""Extract reviewable text without rewriting the original document."""

import zipfile
from pathlib import Path

MAX_BYTES = 50 * 1024**2
MAX_TEXT = 1_000_000


def extraction_warnings(path):
    """Surface content that cannot be faithfully flattened into a bilingual text pair."""
    path = Path(path)
    if path.suffix.lower() not in (".pptx", ".docx"):
        return []
    inspect_zip(path)
    from lxml import etree
    formula = objects = False
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml") or not name.startswith(("word/", "ppt/slides/")):
                continue
            tree = etree.fromstring(archive.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
            for element in tree.iter():
                local = etree.QName(element).localname
                formula |= local in ("oMath", "oMathPara")
                objects |= local in ("oleObj", "OLEObject", "chart", "videoFile", "pic")
    result = []
    if formula:
        result.append("Nguồn có công thức dạng đối tượng. Chúng chưa được chuyển đầy đủ sang văn bản; mở bản nguồn để đối chiếu và nhập lại công thức cần dạy (có thể đặt trong dấu $…$ để khóa khi dịch).")
    if objects:
        result.append("Ảnh, biểu đồ, video và đối tượng nhúng vẫn có trong tệp nguồn. Chế độ giữ thiết kế PowerPoint sao chép các đối tượng này; chữ bên trong ảnh/biểu đồ và công thức nhúng chưa được dịch tự động. Hãy đối chiếu bản xuất trước khi dạy.")
    return result


def inspect_zip(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > 250 * 1024**2:
            raise ValueError("Tài liệu có dữ liệu giải nén quá lớn.")


def parse_document(path, language="vi", cancelled=None):
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError("Không đọc được tệp hoặc tệp vượt giới hạn 50 MB.")
    suffix = path.suffix.lower()
    blocks = []
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8-sig")
        blocks = [(f"Đoạn {i + 1}", paragraph.strip()) for i, paragraph in enumerate(text.split("\n\n"))]
    elif suffix == ".pptx":
        from pptx import Presentation

        inspect_zip(path)
        presentation = Presentation(path)
        from .source_deck import text_blocks

        blocks = [(unit["locator"], unit["text"]) for unit in text_blocks(presentation)]
        # Scan-style decks can still enter Template Mode. OCR only substantial
        # pictures on slides with no native text; never run it on every logo.
        occupied_slides = {int(locator.split()[1].split("·")[0]) for locator, _ in blocks}
        for index, slide in enumerate(presentation.slides, 1):
            if index in occupied_slides:
                continue
            pictures = [shape for shape in slide.shapes if int(shape.shape_type) == 13
                        and shape.width * shape.height >= presentation.slide_width * presentation.slide_height * .2]
            if not pictures:
                continue
            import io
            import tempfile

            from PIL import Image

            from .ocr import recognize

            with tempfile.TemporaryDirectory(prefix="biliclass-slide-ocr-") as temporary:
                for number, picture in enumerate(pictures, 1):
                    target = Path(temporary) / "scan.png"
                    with Image.open(io.BytesIO(picture.image.blob)) as image:
                        image.thumbnail((2400, 2400))
                        image.convert("RGB").save(target)
                    recognized = recognize(target, language, cancelled=cancelled)
                    text = "\n".join(value for _, value in recognized if value.strip())
                    if text.strip():
                        blocks.append((f"Slide {index} · ảnh {number} · OCR cần kiểm tra", text))
        blocks.sort(key=lambda item: int(item[0].split()[1]))
    elif suffix == ".docx":
        from docx import Document

        inspect_zip(path)
        document = Document(path)
        for index, item in enumerate(document.iter_inner_content(), 1):
            text = (
                item.text
                if hasattr(item, "text")
                else "\n".join(" | ".join(cell.text for cell in row.cells) for row in item.rows)
            )
            blocks.append((f"Mục {index}", text))
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("PDF được mã hóa; hãy xuất bản không có mật khẩu để nhập.")
        if len(reader.pages) > 300:
            raise ValueError("Bản thử hỗ trợ tối đa 300 trang PDF.")
        blocks = [(f"Trang {i + 1}", page.extract_text() or "") for i, page in enumerate(reader.pages)]
        missing = [i for i, (_, text) in enumerate(blocks) if not text.strip()]
        if missing:
            from .ocr import recognize
            recognized = dict(recognize(path, language, pages=missing, cancelled=cancelled))
            for index in missing:
                locator = f"Trang {index + 1} · OCR cần kiểm tra"
                blocks[index] = locator, recognized.get(locator, "")
    elif suffix in (".png", ".jpg", ".jpeg"):
        from .ocr import recognize
        blocks = recognize(path, language, cancelled=cancelled)
    else:
        raise ValueError("Hỗ trợ PPTX, DOCX, PDF, TXT, PNG và JPG.")
    blocks = [(locator, text.strip()) for locator, text in blocks if text.strip()]
    if not blocks:
        raise ValueError("Không tìm thấy văn bản. Tệp có thể là ảnh/scan và cần OCR.")
    if sum(len(text) for _, text in blocks) > MAX_TEXT:
        raise ValueError("Nội dung vượt giới hạn 1 triệu ký tự của bản thử.")
    return blocks
