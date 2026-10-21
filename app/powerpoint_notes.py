"""Read optional notes and restore only a missing notes-body placeholder."""

from pptx.enum.shapes import PP_PLACEHOLDER


def read_text(slide):
    if not slide.has_notes_slide:
        return ''
    frame = slide.notes_slide.notes_text_frame
    return frame.text if frame is not None else ''


def text_frame(slide):
    notes = slide.notes_slide
    if notes.notes_text_frame is None:
        body = next((shape for shape in notes.part.notes_master.placeholders
                     if shape.placeholder_format.type == PP_PLACEHOLDER.BODY), None)
        if body is not None:
            notes.shapes.clone_placeholder(body)
        else:
            # A valid custom notes master can omit its body too. Use the same
            # shape factory as python-pptx, with fresh IDs; retain every other
            # notes shape, slide thumbnail, footer and relationship.
            shape_id = max((shape.shape_id for shape in notes.shapes), default=0) + 1
            index = max((shape.placeholder_format.idx for shape in notes.placeholders), default=-1) + 1
            notes.shapes._spTree.add_placeholder(shape_id, 'BiliClass Notes', PP_PLACEHOLDER.BODY,
                                                'horz', 'full', index)
    return notes.notes_text_frame
