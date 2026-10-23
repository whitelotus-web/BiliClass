"""Apply validated translations to native OOXML; keep original package resources."""

import copy
import posixpath
import re
import zipfile
from pathlib import Path

from lxml import etree

from .native_layout import can_fit, fit_body
from .powerpoint_notes import text_frame as notes_text_frame
from .source_deck import NS, _bytes, _part_relative, _rels_path, _size, _xml, source_slide_parts
from .translation_batches import TOKEN, lock_numbers, restore


def leaves(shapes, *, modern=False):
    from .slide_speech import native_shapes

    for shape in native_shapes(shapes) if modern else shapes:
        if hasattr(shape, 'shapes'):
            yield from leaves(shape.shapes, modern=modern)
        else:
            yield shape


def picture_key(shape):
    try:
        return shape.image.sha1
    except (ValueError,KeyError):
        return None


def paragraph_text(paragraph):
    return ''.join(
        node.text or '' for node in paragraph.iter() if etree.QName(node).localname == 't')


def _title_height(shape, shapes, deck):
    size=_size(shape.text_frame,shape)
    required=int(((size+max(18,size*.8))*1.12)*12700+shape.text_frame.margin_top+shape.text_frame.margin_bottom)
    height=max(shape.height,required)
    height=min(height, deck.slide_height-shape.top)
    for other in shapes:
        if other.shape_id==shape.shape_id:
            continue
        if not (other.has_text_frame and other.text.strip()) and other.width*other.height>deck.slide_width*deck.slide_height*.6:
            continue  # A full-slide decorative background remains behind the title.
        if (shape.left<other.left+other.width and shape.left+shape.width>other.left
                and other.top>=shape.top+shape.height-12700 and other.top<shape.top+height):
            height=min(height, other.top-shape.top)
    return height if height >= shape.height and can_fit(shape, '', title_height=height) else None


def _fits_bilingual(shape,mode):
    return can_fit(shape,mode)


def _body_height(shape, shapes, deck, mode):
    """Use only genuinely empty space below a text box; never move a picture."""
    height = min(deck.slide_height-shape.top, shape.height+max(shape.height//2, 48*12700))
    for other in shapes:
        if other.shape_id == shape.shape_id:
            continue
        if not (other.has_text_frame and other.text.strip()) and other.width*other.height > deck.slide_width*deck.slide_height*.6:
            continue
        if (shape.left < other.left+other.width and shape.left+shape.width > other.left
                and other.top+other.height > shape.top+shape.height
                and other.top < shape.top+height):
            # An already overlapping decoration/figure forbids expansion too.
            height = min(height, max(shape.height, other.top-shape.top-12700))
    return height if height > shape.height and can_fit(shape, mode, height=height) else None


def _visible_insets(shape, box, deck, items):
    left,top,width,height = (box[k] for k in ('x','y','w','h'))
    right = min(deck.slide_width,left+width)
    for picture in items:
        if picture['source_type'] != 'image':
            continue
        overlap = min(top+height,picture['y']+picture['h'])-max(top,picture['y'])
        if (overlap > height*.5 and left+width*.5 < picture['x'] < right
                and picture['w']*picture['h'] < deck.slide_width*deck.slide_height*.6):
            right = picture['x']-12700
    return {'lIns':round(max(0,-left)*shape.width/max(1,width)),
            'rIns':round(max(0,left+width-right)*shape.width/max(1,width)),
            'tIns':round(max(0,-top)*shape.height/max(1,height)),
            'bIns':round(max(0,top+height-deck.slide_height)*shape.height/max(1,height))}


def extract(source, mode, *, preserve_design=False, same_slide=False):
    """Preflight before any send. Image/diagram-heavy decks retain the visual route."""
    from pptx import Presentation
    from pptx.text.text import TextFrame

    from .importers import inspect_zip
    from .input_analysis import language_of
    from .slide_semantics import equation_only, quantity_only, roman_label, title_frame

    preserve_design = preserve_design or same_slide
    inspect_zip(source)
    deck = Presentation(source)
    original_parts = source_slide_parts(source)
    units, unsupported, english = [], [], 0
    picture_slides = {}
    for index, slide in enumerate(deck.slides,1):
        for shape in leaves(slide.shapes, modern=preserve_design):
            if int(shape.shape_type) == 13:
                picture_slides.setdefault(picture_key(shape),set()).add(index)
    for index, slide in enumerate(deck.slides, 1):
        shapes = list(leaves(slide.shapes, modern=preserve_design))
        if preserve_design:
            from .slide_speech import objects
            visible_items = objects(slide)
            visible_boxes = {item['shape_id']:item for item in visible_items}
        else:
            visible_boxes = {}
        motion={int(value) for value in slide._element.xpath('.//p:spTgt/@spid') if value.isdigit()}
        moving={int(value) for value in slide._element.xpath('.//p:animMotion//p:spTgt/@spid') if value.isdigit()}
        if mode != 'english_only' and any(hasattr(shape,'shapes') and any(s.has_text_frame and s.text.strip() for s in leaves(shape.shapes)) for shape in slide.shapes):
            unsupported.append(f'Slide {index}: giữ bố cục nhóm đối tượng gốc')
        text_ids={shape.shape_id for shape in shapes if shape.has_text_frame and shape.text.strip()}
        if (mode in {'parallel_columns','sentence_pairs'} and slide._element.find('p:timing',NS) is not None
                and (not motion or motion & text_ids)):
            unsupported.append(f'Slide {index}: giữ đồng bộ hiệu ứng với bản song ngữ')
        if slide._element.xpath('.//*[local-name()="AlternateContent"]//*[local-name()="txBody" or local-name()="oMath"]'):
            unsupported.append(f'Slide {index}: đối tượng chữ/công thức có nhánh tương thích cần kiểm tra trực quan')
        text_count = sum(len(shape.text.strip()) for shape in shapes if shape.has_text_frame)
        pictures = [s for s in shapes if int(s.shape_type) == 13]
        if any(p.width*p.height > deck.slide_width*deck.slide_height*.08
               and (text_count < 60 or picture_key(p) is None or len(picture_slides[picture_key(p)]) <= max(2,len(deck.slides)*.15)) for p in pictures):
            unsupported.append(f'Slide {index}: nội dung phụ thuộc ảnh')
        for shape in shapes:
            box=visible_boxes.get(shape.shape_id)
            x,y,w,h=((box[k] for k in ('x','y','w','h')) if box else
                     (shape.left,shape.top,shape.width,shape.height))
            fully_outside=(x+w<=0 or y+h<=0 or x>=deck.slide_width or y>=deck.slide_height)
            animated=bool(motion.intersection(box.get('ancestors',[]))) if box else shape.shape_id in motion
            if fully_outside and not animated and mode != 'english_only':
                continue  # Keep unused bilingual source workspace objects intact.
            # English-only also translates off-canvas text: PowerPoint can
            # paint overflowing glyphs inside the slide despite the box bounds.
            transform = shape._element.find('p:spPr/a:xfrm',NS)
            if mode == 'parallel_columns' and shape.has_text_frame and (
                    transform is None or transform.get('rot') not in {None,'0'}):
                unsupported.append(f'Slide {index}: hộp chữ xoay hoặc thiếu tọa độ riêng')
            title_height=layout_height=None
            shape_title = title_frame(shape, deck.slide_width, deck.slide_height)
            if shape.has_text_frame and shape.text.strip():
                if (shape.left<0 or shape.top<0 or shape.left+shape.width>deck.slide_width
                        or shape.top+shape.height>deck.slide_height) and shape.shape_id not in motion:
                    unsupported.append(f'Slide {index}: dàn lại vùng chữ tràn khung có sẵn')
                if shape_title and mode!='english_only':
                    title_height=_title_height(shape,shapes,deck)
                    if title_height is None:
                        unsupported.append(f'Slide {index}: cần dành vùng cho tiêu đề song ngữ')
                elif len(shape.text)>12 and not _fits_bilingual(shape,mode):
                    layout_height=_body_height(shape,shapes,deck,mode)
                    if layout_height is None:
                        unsupported.append(f'Slide {index}: cần dàn lại vùng chữ chật')
            if shape._element.xpath('.//a:graphicData[not(a:tbl)]'):
                unsupported.append(f'Slide {index}: sơ đồ cần xử lý trực quan')
            if shape.has_chart:
                unsupported.append(f'Slide {index}: biểu đồ có nhãn riêng')
            if shape.has_table and mode == 'parallel_columns':
                unsupported.append(f'Slide {index}: bảng cần dàn hai cột trực quan')
            bodies = ([shape._element.find('p:txBody', NS)] if shape.has_text_frame else
                      shape._element.findall('.//a:tc/a:txBody', NS) if shape.has_table else [])
            for body_index, body in enumerate(bodies):
                if body is None:
                    continue
                for paragraph_index, paragraph in enumerate(body.findall('a:p', NS)):
                    raw = paragraph_text(paragraph).strip()
                    if not raw:
                        continue
                    if TOKEN.search(raw):
                        unsupported.append(f'Slide {index}: nhãn nguồn trùng định dạng chuyển đổi')
                        continue
                    parts, numbers, xml_locks = [], [], {}
                    math_words = False
                    for child in paragraph:
                        tag = etree.QName(child).localname
                        if tag in {'pPr','endParaRPr'}:
                            continue
                        if tag == 'br':
                            parts.append('\n')
                            continue
                        special = tag != 'r' or child.find('a:rPr', NS) is not None and (
                            child.find('a:rPr', NS).get('baseline', '0') != '0')
                        if preserve_design and etree.XPath('.//a:rPr/a:solidFill/*/a:alpha[@val="0"]|.//a:rPr/a:noFill',namespaces=NS)(child):
                            special = True  # Keep invisible source runs invisible, including private markers.
                        if special and preserve_design:
                            from .math_prose import fragments

                            exposed = fragments(child)
                            if exposed:
                                math_words = True
                                for kind, value in exposed:
                                    if kind == 'text':
                                        protected, locked = lock_numbers(value)
                                        parts.append(TOKEN.sub(lambda m: f'⟦BLC{int(m.group()[4:-1])+len(numbers)}⟧', protected))
                                        numbers.extend(locked)
                                    else:
                                        token = len(numbers)
                                        numbers.append(paragraph_text(value))
                                        xml_locks[str(token)] = _bytes(value).decode('utf-8')
                                        parts.append(f'⟦BLC{token}⟧')
                                continue
                        if special:
                            token = len(numbers)
                            numbers.append(paragraph_text(child))
                            xml_locks[str(token)] = _bytes(child).decode('utf-8')
                            parts.append(f'⟦BLC{token}⟧')
                        else:
                            value = paragraph_text(child)
                            protected, locked = lock_numbers(value)
                            protected = TOKEN.sub(lambda m: f'⟦BLC{int(m.group()[4:-1])+len(numbers)}⟧', protected)
                            numbers.extend(locked)
                            parts.append(protected)
                    text = ''.join(parts).strip()
                    if not text:
                        continue
                    if len(text)>14000:
                        unsupported.append(f'Slide {index}: vùng nội dung cần xử lý trực quan')
                    language = language_of(raw)
                    if (roman_label(raw) or not math_words and equation_only(paragraph) or quantity_only(raw)
                            or re.fullmatch(r'(?:https?://|www\.|(?:facebook|youtube|mailto):)\S+', raw.strip(), re.I)):
                        language='neutral'
                    if same_slide and language != 'vi' and not TOKEN.sub('',text).strip():
                        language='neutral'
                    remainder=TOKEN.sub('',text).strip()
                    if (same_slide and TOKEN.search(text) and language!='vi'
                            and re.fullmatch(r'[A-Za-zα-ωΑ-Ω=+−±\s-]*',remainder)
                            and all(len(word)==1 for word in re.findall(r'[A-Za-zα-ωΑ-Ω]+',remainder))):
                        language='neutral'
                    if re.fullmatch(r'[A-Za-zα-ωΑ-Ω]',raw) or len(raw)<=8 and re.search(r'[\d₀-₉/*^=]',raw) and re.fullmatch(r'[A-Za-zα-ωΑ-Ω\d₀-₉/*^=+().−-]+',raw):
                        language='neutral'
                    if re.fullmatch(r'https?://\S+', raw):
                        language='neutral'  # Source links have no translated counterpart.
                    if same_slide and re.match(r'^[A-D][.)]\s*',raw):
                        quantity=re.sub(r'^[A-D][.)]\s*','',raw)
                        quantity=re.sub(r'(?i)(?<![A-Za-z])(?:cm|mm|km|rad|kg|hz|m|s|ms)(?![A-Za-z])','',quantity)
                        if re.search(r'\d',quantity) and re.fullmatch(r'[\d₀-₉²³π√\s/,.;+−^=*()–-]+',quantity):
                            language='neutral'  # A numeric option has no language to duplicate.
                    english += language == 'en'
                    units.append({'id':f's{index}o{shape.shape_id}b{body_index}p{paragraph_index}',
                        'slide':index,'shape_id':shape.shape_id,'body':body_index,'paragraph':paragraph_index,
                        'text':text,'numbers':numbers,'xml_locks':xml_locks,'language':language,
                        'raw':raw,'title':bool(shape_title),'title_height':title_height,'part':original_parts[index-1],
                        'layout_height':layout_height, **({'math_prose':True} if math_words else {}),
                        **({'table_cell':True} if shape.has_table and not any(
                            n.find('p:txEl', NS) is not None for n in slide._element.xpath(
                                './/p:spTgt[@spid="' + str(shape.shape_id) + '"]')) else {}),
                        **({'clip_insets':_visible_insets(shape,visible_boxes[shape.shape_id],deck,visible_items)}
                           if shape.shape_id in visible_boxes and shape.shape_id not in moving and not shape.rotation else {}),
                        'font_size':round(_size(shape.text_frame if shape.has_text_frame else TextFrame(body, shape),shape)*100)})
    if mode != 'english_only' and english > max(5, len(units)*.35):
        unsupported.append('Nguồn có nhiều vùng English/song ngữ cần ghép cặp trực quan')
    if not units:
        unsupported.append('Nguồn chưa có văn bản chỉnh sửa được')
    visual = {int(match[1]) for reason in unsupported if (match:=re.match(r'Slide (\d+):',reason))}
    if any(not reason.startswith('Slide ') for reason in unsupported):
        visual = set(range(1,len(deck.slides)+1))
    native = units if preserve_design else [unit for unit in units if unit['slide'] not in visual]
    manifest = {'version':4 if same_slide else 3 if preserve_design else 2,'eligible':bool(units or visual),'reasons':list(dict.fromkeys(unsupported)),
            'slides':len(deck.slides),'units':native,'visual_slides':sorted(visual),
            **({'preserve_design':True,'paired_slides':sorted(visual) if mode != 'english_only' else []}
               if preserve_design else {}),
            'text_characters':sum(len(u['text']) for u in native)}
    if same_slide:
        from .same_slide import prepare
        from .table_prose import join_wrapped

        manifest.update(same_slide=True, paired_slides=[])
        manifest['units'] = join_wrapped(manifest['units'])
        prepare(source,manifest)
    elif preserve_design:
        from .native_fonts import missing_embedded_fonts

        manifest['font_fallbacks'] = missing_embedded_fonts(source)
    return manifest


def _element(parent, prefix, tag, **attributes):
    return etree.SubElement(parent, f'{{{NS[prefix]}}}{tag}', **{k:str(v) for k,v in attributes.items()})


def translated_paragraph(original, text, unit, *, english=False, smaller=False):
    """Translate natural text while keeping fields/equations/superscripts as XML."""
    restore(text, unit['numbers'])
    result = copy.deepcopy(original)
    for child in list(result):
        if etree.QName(child).localname != 'pPr':
            result.remove(child)
    template = next((props for run in original.findall('a:r', NS)
                     if (props := run.find('a:rPr', NS)) is not None
                     and props.get('baseline', '0') == '0'), None)
    if template is None and unit.get('math_prose'):
        template=next(iter(original.findall('.//a:rPr', NS)), None)
    if template is None:
        template=original.find('a:pPr/a:defRPr',NS)
    if template is None:
        template=next(iter(original.findall('.//a:rPr', NS)), None)
    template = copy.deepcopy(template) if template is not None else etree.Element(f'{{{NS["a"]}}}rPr')
    template.tag=f'{{{NS["a"]}}}rPr'
    template.attrib.pop('baseline', None)
    if template.get('sz') is None and unit.get('font_size'):
        template.set('sz',str(unit['font_size']))
    if english:
        template.set('lang','en-US')
    if smaller:
        template.set('i','1')
        size = int(template.get('sz', '2400'))
        template.set('sz',str(max(min(1400, size), round(size*.8))))
    for piece in re.split(r'(⟦BLC\d+⟧)', text):
        if not piece:
            continue
        if TOKEN.fullmatch(piece):
            key = piece[4:-1]
            if key in unit['xml_locks']:
                result.append(_xml(unit['xml_locks'][key].encode('utf-8')))
                continue
            piece = unit['numbers'][int(key)]
        for number, line in enumerate(piece.split('\n')):
            if number:
                _element(result,'a','br')
            if line:
                run = _element(result,'a','r')
                run.append(copy.deepcopy(template))
                _element(run,'a','t').text = line
    return result


def _fit_body(body):
    properties = body.find('a:bodyPr', NS)
    if properties is None:
        return
    properties.set('wrap','square')
    for child in list(properties):
        if etree.QName(child).localname in {'normAutofit','spAutoFit','noAutofit'}:
            properties.remove(child)
    _element(properties,'a','noAutofit')


def _pair_columns(tree, shape, english):
    """Static text uses sibling columns; animated/grouped slides use the visual lane."""
    properties = shape.find('p:spPr/a:xfrm', NS)
    if properties is None or properties.get('rot') not in {None,'0'}:
        raise ValueError('Đối tượng này cần dàn cột bằng luồng trực quan.')
    off, ext = properties.find('a:off', NS), properties.find('a:ext', NS)
    if off is None or ext is None:
        raise ValueError('Thiếu tọa độ hộp chữ để giữ PowerPoint gốc.')
    x,w = int(off.get('x')),int(ext.get('cx'))
    ids = [int(value) for value in tree.xpath('.//p:cNvPr/@id',namespaces=NS)]
    clone_id = english.find('p:nvSpPr/p:cNvPr', NS)
    clone_id.set('id',str(max(ids)+1))
    clone_id.set('name',clone_id.get('name','')+' English')
    gap = min(127000, int(w*.035))
    column = (w-gap)//2
    properties.find('a:ext', NS).set('cx',str(column))
    en_transform = english.find('p:spPr/a:xfrm', NS)
    en_transform.find('a:off', NS).set('x',str(x+column+gap))
    en_transform.find('a:ext', NS).set('cx',str(column))
    owner = shape.getparent()
    position = owner.index(shape)
    owner.insert(position+1,english)
    return int(clone_id.get('id'))


def assemble(source, destination, manifest, translations, mode):
    """Only the changed slide/notes XML is serialized; media and timing survive."""
    from pptx import Presentation

    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError('Không ghi đè PowerPoint nguồn.')
    if set(translations) != {u['id'] for u in manifest['units']}:
        raise ValueError('Chưa đủ bản dịch để dựng bài giảng.')
    from .translation_case import align_case

    # Also normalize cached/provider-independent translations before creating
    # paragraphs, notes and coverage hashes, keeping all three in agreement.
    translations = {u['id']: {**translations[u['id']], 'en': align_case(u['text'], translations[u['id']]['en'])}
                    if u['language'] == 'vi' else translations[u['id']] for u in manifest['units']}
    with zipfile.ZipFile(source) as archive:
        entries = {item.filename:archive.read(item) for item in archive.infolist()}
    grouped = {}
    for unit in manifest['units']:
        grouped.setdefault(unit['part'],{}).setdefault(unit['shape_id'],[]).append(unit)
    notes = {number: {'vi': [], 'en': []} for number in manifest.get('note_slides', [])}
    paired_trees, coverage, continuations = {}, {}, {}
    together = manifest.get('same_slide', False)
    if together:
        from . import same_slide
    for part, shapes in grouped.items():
        tree = _xml(entries[part])
        index = next(iter(shapes.values()))[0]['slide']
        paired = index in manifest.get('paired_slides', [])
        effective_mode = 'english_only' if paired else mode
        columns, paragraph_maps = {}, {}
        for shape_id, units in shapes.items():
            nodes = tree.xpath('.//p:cNvPr[@id=$id]',namespaces=NS,id=str(shape_id))
            if len(nodes) > 1 and manifest.get('preserve_design'):
                # Modern math/text and its compatibility bitmap share an ID.
                # Edit the Choice once; keep the original fallback asset intact.
                preferred = [node for node in nodes if not any(
                    etree.QName(parent).localname == 'Fallback' for parent in node.iterancestors())]
                if preferred:
                    nodes = preferred
            if len(nodes) != 1:
                raise ValueError('Đối tượng PowerPoint không còn khớp nguồn.')
            shape = nodes[0].getparent().getparent()
            effective_mode = 'english_only' if paired else mode
            transform=shape.find('p:spPr/a:xfrm',NS)
            if together and effective_mode == 'parallel_columns' and (
                    shape.tag != f'{{{NS["p"]}}}sp' or same_slide.prefer_inline(units, translations)
                    or transform is None or transform.get('rot') not in {None,'0'}):
                effective_mode = 'sentence_pairs'  # Cell-local pairs keep table row/column ownership.
            if together and units[0].get('suppressed'):
                for paragraph in shape.findall('p:txBody/a:p',NS):
                    for child in list(paragraph):
                        if etree.QName(child).localname != 'pPr':
                            paragraph.remove(child)
                continue
            if all(u['language'] == 'neutral' or u['language'] == 'en'
                   and translations[u['id']]['en'] == u['text'] for u in units):
                continue  # Keep exact XML only when the validated response agrees.
            from .native_fonts import normalize

            normalize(shape,manifest.get('font_fallbacks',{}))
            bodies = ([shape.find('p:txBody', NS)] if shape.tag == f'{{{NS["p"]}}}sp' else
                      shape.findall('.//a:tc/a:txBody', NS))
            if together:
                same_slide.apply_geometry(shape,units[0])
                if not any(u['title'] for u in units):
                    nodes[0].set('name',('BCL_LABEL ' if units[0].get('label') else 'BCL_BODY ')+nodes[0].get('name',''))
            if manifest.get('version', 1) >= 2 and not paired and not together:
                height=max((u.get('layout_height') or 0 for u in units),default=0)
                extent=shape.find('p:spPr/a:xfrm/a:ext',NS)
                if extent is not None and height:
                    extent.set('cy',str(height))
            clone = copy.deepcopy(shape)
            clone_bodies = ([clone.find('p:txBody', NS)] if shape.tag == f'{{{NS["p"]}}}sp' else
                            clone.findall('.//a:tc/a:txBody', NS))
            title = any(u['title'] for u in units)
            if title and effective_mode!='english_only' and not together:
                height=max((u.get('title_height') or 0 for u in units),default=0)
                extent=shape.find('p:spPr/a:xfrm/a:ext',NS)
                if extent is not None and height:
                    extent.set('cy',str(height))
            for body_index, body in enumerate(bodies):
                if body is None:
                    continue
                originals = list(body.findall('a:p', NS))
                selected = {u['paragraph']:u for u in units if u['body'] == body_index}
                if together and selected and shape.tag == f'{{{NS["p"]}}}graphicFrame':
                    margins = next(iter(selected.values())).get('cell_margins', {})
                    if margins:
                        cell = body.getparent()
                        props = cell.find('a:tcPr', NS)
                        if props is None:
                            props = etree.SubElement(cell, f'{{{NS["a"]}}}tcPr')
                        for name, value in margins.items():
                            props.set(name, str(value))
                joined = {p for u in selected.values() for p in u.get('merged_paragraphs', [])[1:]}
                for paragraph in originals:
                    body.remove(paragraph)
                en_body = clone_bodies[body_index]
                for paragraph in list(en_body.findall('a:p', NS)):
                    en_body.remove(paragraph)
                for paragraph_index, original in enumerate(originals):
                    if paragraph_index in joined:
                        continue
                    start = len(body.findall('a:p',NS))
                    unit = selected.get(paragraph_index)
                    if unit is None:
                        if together:
                            original = same_slide.style_paragraph(copy.deepcopy(original),units[0]['font_size'],title=title,
                                canvas_scale=units[0].get('canvas_scale', 1))
                        body.append(copy.deepcopy(original))
                        en_body.append(copy.deepcopy(original))
                        if together:
                            paragraph_maps.setdefault(shape_id,{})[paragraph_index]=[start]
                        continue
                    value = translations[unit['id']]
                    if unit.get('merged_paragraphs'):
                        original = copy.deepcopy(original)
                        for child in list(original):
                            if etree.QName(child).localname == 'endParaRPr':
                                original.remove(child)
                        for p in unit['merged_paragraphs'][1:]:
                            _element(_element(original, 'a', 'r'), 'a', 't').text = ' '
                            for child in originals[p]:
                                if etree.QName(child).localname not in {'pPr', 'endParaRPr'}:
                                    original.append(copy.deepcopy(child))
                    english = translated_paragraph(original,value['en'],unit,english=True,
                        smaller=not together and (effective_mode=='sentence_pairs' or title and
                        (effective_mode!='english_only' or not manifest.get('preserve_design'))))
                    if together:
                        if unit.get('math_prose') or unit.get('cell_heading'):
                            # Normal equation text is prose, not a formula. Let
                            # Office wrap/fit it as text while restoring every
                            # genuinely mathematical subtree from its XML lock.
                            original = translated_paragraph(original,unit['text'],unit)
                        original = same_slide.style_paragraph(copy.deepcopy(original),unit['font_size'],title=title,label=unit.get('label',False),
                            canvas_scale=unit.get('canvas_scale', 1))
                        english = same_slide.style_paragraph(english,unit['font_size'],english=True,title=title,label=unit.get('label',False),
                            canvas_scale=unit.get('canvas_scale', 1))
                        if unit.get('cell_heading'):
                            original.find('a:pPr', NS).set('algn', 'ctr')
                            english.find('a:pPr', NS).set('algn', 'ctr')
                        if effective_mode == 'sentence_pairs':
                            props = english.find('a:pPr', NS)
                            for child in list(props):
                                if etree.QName(child).localname in {'buAutoNum', 'buChar', 'buNone'}:
                                    props.remove(child)
                            etree.SubElement(props, f'{{{NS["a"]}}}buNone')
                    if effective_mode == 'english_only':
                        body.append(english)
                    elif effective_mode == 'sentence_pairs' or title:
                        body.append(copy.deepcopy(original))
                        if unit['language'] != 'neutral':
                            body.append(english)
                    elif effective_mode == 'integrated_keywords':
                        body.append(translated_paragraph(original,value['mixed'],unit))
                    else:
                        body.append(copy.deepcopy(original))
                        en_body.append(english)
                    if together:
                        paragraph_maps.setdefault(shape_id,{})[paragraph_index]=list(range(start,len(body.findall('a:p',NS))))
                        if effective_mode in {'parallel_columns','sentence_pairs'} or title and effective_mode!='english_only':
                            inline = effective_mode=='sentence_pairs' or title
                            shared = inline and unit['language']=='neutral'
                            english_index = (start if shared else start+1) if inline else len(en_body.findall('a:p',NS))-1
                            coverage.setdefault(index,[]).append({
                                'unit':unit['id'],'body':body_index,'source_shape':shape_id,
                                'vi_shape':shape_id,'en_shape':shape_id,
                                'vi_paragraph':start,
                                'en_paragraph':english_index,
                                'vi_hash':same_slide.fingerprint(original),'en_hash':same_slide.fingerprint(original if shared else english),
                                'inline':inline,'title':title,
                                'question':unit.get('question',False),'neutral':unit['language']=='neutral'})
                    pair = notes.setdefault(unit['slide'],{'vi':[],'en':[]})
                    pair['vi'].append(unit['raw'])
                    pair['en'].append(restore(value['en'],unit['numbers']))
                _fit_body(body)
                _fit_body(en_body)
                if paired:
                    props = body.find('a:bodyPr',NS)
                    if props is not None:
                        for key,value in units[0].get('clip_insets',{}).items():
                            if value:
                                props.set(key,str(int(props.get(key,'0'))+value))
            if effective_mode == 'parallel_columns' and not title:
                if shape.tag != f'{{{NS["p"]}}}sp':
                    raise ValueError('Bảng này cần dàn hai cột bằng luồng trực quan.')
                clone_id = _pair_columns(tree,shape,clone)
                if together:
                    columns[shape_id]=clone_id
                    for pair in coverage.get(index,[]):
                        if pair['source_shape']==shape_id:
                            pair['en_shape']=clone_id
                    same_slide.fit(clone,max(u['font_size'] for u in units))
                elif manifest.get('version', 1) >= 2:
                    fit_body(clone, max(u['font_size'] for u in units))
            if together:
                same_slide.fit(shape,max(u['font_size'] for u in units),title=title)
                if shape.tag == f'{{{NS["p"]}}}graphicFrame':
                    same_slide.fit_cells(shape, max(u['font_size'] for u in units))
            elif manifest.get('version', 1) >= 2:
                fit_body(shape, max(u['font_size'] for u in units), minimum=14 if paired else 18)
        from .slide_typography import protect_tree

        protect_tree(tree)
        if together:
            same_slide.synchronize(tree,columns,paragraph_maps)
            pages=same_slide.paginate(tree,coverage.get(index,[]))
            tree,coverage[index]=pages[0]
            continuations[index]=pages[1:]
        if paired:
            paired_trees[index] = (part, tree)
        else:
            entries[part] = _bytes(tree)
    # python-pptx creates only the notes plumbing in a temporary copy. Keep its
    # slide XML out of the output so unsupported source extensions stay intact.
    deck = Presentation(source)
    original_parts = source_slide_parts(source)
    normalized = {str(slide.part.partname)[1:]:part for slide,part in zip(deck.slides,original_parts,strict=True)}
    for index,pair in notes.items():
        frame = notes_text_frame(deck.slides[index-1])
        original = frame.text.split('BILICLASS_NOTES:')[0].rstrip()
        frame.text = original + '\nBILICLASS_NOTES:\nVI: ' + '\n'.join(pair['vi']) + '\nEN: ' + '\n'.join(pair['en'])
    notes_copy = destination.with_suffix('.notes.pptx')
    pending = destination.with_suffix('.pending.pptx')
    destination.parent.mkdir(parents=True,exist_ok=True)
    try:
        deck.save(notes_copy)
        with zipfile.ZipFile(notes_copy) as archive:
            incoming = {name:archive.read(name) for name in archive.namelist()}
        for name,data in incoming.items():
            if name.startswith(('ppt/notesSlides/','ppt/notesMasters/')):
                if name.endswith('.rels'):
                    owner=posixpath.join(posixpath.dirname(posixpath.dirname(name)),posixpath.basename(name)[:-5])
                    tree=_xml(data)
                    for relation in tree:
                        if relation.get('Type','').endswith('/slide'):
                            target=_part_relative(owner,relation.get('Target',''))
                            if target in normalized:
                                relation.set('Target',posixpath.relpath(normalized[target],posixpath.dirname(owner)))
                    data=_bytes(tree)
                entries[name] = data
            elif name.startswith('ppt/slides/_rels/') or name == 'ppt/_rels/presentation.xml.rels':
                destination_name=name
                if name.startswith('ppt/slides/_rels/'):
                    owner='ppt/slides/'+posixpath.basename(name)[:-5]
                    destination_name=_rels_path(normalized.get(owner,owner))
                tree = _xml(entries[destination_name]) if destination_name in entries else etree.Element(f'{{{NS["rel"]}}}Relationships')
                by_id = {node.get('Id'):node for node in tree}
                for node in _xml(data):
                    if not node.get('Type','').endswith(('/notesSlide','/notesMaster')):
                        continue
                    if node.get('Id') in by_id:
                        if by_id[node.get('Id')].get('Type') != node.get('Type'):
                            raise ValueError('Liên kết ghi chú trùng với tài nguyên nguồn.')
                    else:
                        tree.append(copy.deepcopy(node))
                entries[destination_name] = _bytes(tree)
        types = _xml(entries['[Content_Types].xml'])
        known = {(node.tag,node.get('PartName') or node.get('Extension')) for node in types}
        for node in _xml(incoming['[Content_Types].xml']):
            key = (node.tag,node.get('PartName') or node.get('Extension'))
            if key not in known and (node.tag.endswith('Default') or node.get('PartName','').startswith(('/ppt/notesSlides/','/ppt/notesMasters/'))):
                types.append(copy.deepcopy(node))
                known.add(key)
        entries['[Content_Types].xml'] = _bytes(types)
        presentation = _xml(entries['ppt/presentation.xml'])
        masters = _xml(incoming['ppt/presentation.xml']).find('p:notesMasterIdLst',NS)
        if masters is not None and presentation.find('p:notesMasterIdLst',NS) is None:
            following = presentation.find('p:sldIdLst',NS)
            presentation.insert(presentation.index(following) if following is not None else 0,copy.deepcopy(masters))
            entries['ppt/presentation.xml'] = _bytes(presentation)
        if manifest.get('preserve_design'):
            from .source_deck import _append_slide

            presentation = _xml(entries['ppt/presentation.xml'])
            relations = _xml(entries['ppt/_rels/presentation.xml.rels'])
            types = _xml(entries['[Content_Types].xml'])
            slide_list = presentation.find('p:sldIdLst', NS)
            originals = list(slide_list)
            highest_id = max(int(sid.get('id')) for sid in originals)
            plan = []
            for index, sid in enumerate(originals, 1):
                if index not in paired_trees:
                    role = 'en' if mode == 'english_only' else 'bilingual' if index in notes else 'original'
                    plan.append({'source_slide':index,'role':role,
                                 **({'pairs':coverage[index]} if index in coverage else {})})
                    previous=sid
                    for page,(extra,extra_pairs) in enumerate(continuations.get(index,[]),2):
                        highest_id+=1
                        previous=_append_slide(original_parts[index-1],extra,f'{index}-continuation-{page}',
                                               entries,types,relations,slide_list,previous,highest_id)
                        plan.append({'source_slide':index,'role':'bilingual','continuation':page,'pairs':extra_pairs})
                    continue
                part, tree = paired_trees[index]
                plan.append({'source_slide':index,'role':'vi'})
                highest_id += 1
                _append_slide(part, tree, index, entries, types, relations, slide_list, sid, highest_id)
                plan.append({'source_slide':index,'role':'en'})
            manifest['layout_plan'] = plan
            entries['ppt/presentation.xml'] = _bytes(presentation)
            entries['ppt/_rels/presentation.xml.rels'] = _bytes(relations)
            entries['[Content_Types].xml'] = _bytes(types)
            if 'docProps/app.xml' in entries:
                properties = _xml(entries['docProps/app.xml'])
                count = properties.find('{*}Slides')
                if count is not None:
                    count.text = str(len(plan))
                    entries['docProps/app.xml'] = _bytes(properties)
        with zipfile.ZipFile(pending,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=1) as archive:
            for name,data in entries.items():
                archive.writestr(name,data)
        with zipfile.ZipFile(pending) as archive:
            if archive.testzip() is not None:
                raise ValueError('Tệp PowerPoint xuất chưa hoàn chỉnh.')
        pending.replace(destination)
    finally:
        notes_copy.unlink(missing_ok=True)
        pending.unlink(missing_ok=True)
    return destination
