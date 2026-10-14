"""Make small source decks and merge returned slides without losing source parts."""

import copy
import hashlib
import posixpath
import zipfile
from pathlib import Path

from lxml import etree

from .source_deck import NS, _bytes, _part_relative, _rels_path, _xml, source_slide_parts


def subset(source, indices, destination):
    from pptx import Presentation

    if Path(source).resolve()==Path(destination).resolve():
        raise ValueError('Không ghi đè PowerPoint nguồn khi chia nhóm slide.')
    deck = Presentation(source)
    if not indices or sorted(set(indices)) != list(indices) or min(indices)<1 or max(indices)>len(deck.slides):
        raise ValueError('Nhóm slide nguồn không hợp lệ.')
    for index in range(len(deck.slides),0,-1):
        if index not in indices:
            sid = deck.slides._sldIdLst[index-1]
            deck.part.drop_rel(sid.rId)
            deck.slides._sldIdLst.remove(sid)
    for show in list(deck._element):
        if show.tag.rsplit('}',1)[-1] == 'custShowLst':
            deck._element.remove(show)
    for extension in list(deck._element.xpath('./p:extLst/p:ext')):
        if any(node.tag.rsplit('}',1)[-1]=='sectionLst' for node in extension.iter()):
            extension.getparent().remove(extension)
    deck.part.rename_slide_parts([sid.rId for sid in deck.slides._sldIdLst])
    destination = Path(destination)
    destination.parent.mkdir(parents=True,exist_ok=True)
    pending = destination.with_suffix('.pending.pptx')
    try:
        deck.save(pending)
        pending.replace(destination)
    finally:
        pending.unlink(missing_ok=True)
    return destination


def merge(base, replacements, destination):
    """Each replacement is (PPTX path, original one-based slide indices).

    Namespacing the full relationship closure prevents image/theme/chart name
    collisions. Original slide part names/IDs keep cross-slide links working.
    """
    from pptx import Presentation

    from .importers import inspect_zip
    from .translation_batches import signature

    base, destination = Path(base), Path(destination)
    if base.resolve()==destination.resolve():
        raise ValueError('Không ghi đè PowerPoint nguồn khi ghép bài giảng.')
    deck = Presentation(base)
    base_parts = source_slide_parts(base)
    with zipfile.ZipFile(base) as archive:
        entries = {name:archive.read(name) for name in archive.namelist()}
    types = _xml(entries['[Content_Types].xml'])
    media = {hashlib.sha256(data).hexdigest():name for name,data in entries.items() if name.startswith('ppt/media/')}
    used = set()
    for replacement in replacements:
        path,indices = replacement[:2]
        restore = replacement[2] if len(replacement)>2 else {}
        restored = set()
        if used.intersection(indices):
            raise ValueError('Hai kết quả cùng thay một slide nguồn.')
        used.update(indices)
        inspect_zip(path)
        candidate = Presentation(path)
        candidate_parts = source_slide_parts(path)
        if len(candidate.slides) != len(indices) or any(not 1<=i<=len(deck.slides) for i in indices):
            raise ValueError('Số slide kết quả không khớp nhóm nguồn; giữ các phần đã nhận.')
        if (candidate.slide_width,candidate.slide_height)!=(deck.slide_width,deck.slide_height):
            raise ValueError('Nhóm kết quả đổi kích thước PowerPoint nguồn.')
        with zipfile.ZipFile(path) as archive:
            incoming = {name:archive.read(name) for name in archive.namelist()}
        incoming_types = _xml(incoming['[Content_Types].xml'])
        prefix = 'biliclass-parts-'+signature({'slides':indices,'parts':sorted(incoming)})[:16]+'/'
        mapping = {part:base_parts[index-1] for part,index in zip(candidate_parts,indices,strict=True)}
        visited = set()
        canonical={}

        def same_data(part,target):
            if incoming.get(part)==entries.get(target):
                return True
            if not part.endswith('.xml') or not target.endswith('.xml'):
                return False
            for label,data in ((('incoming',part),incoming.get(part)),(('base',target),entries.get(target))):
                if label not in canonical:
                    canonical[label]=etree.tostring(_xml(data),method='c14n') if data else None
            return canonical[('incoming',part)]==canonical[('base',target)]

        def equivalent(part, target, seen=None):
            """Reuse complete template graphs, not just equal XML filenames."""
            seen=set() if seen is None else seen
            pair=(part,target)
            if pair in seen:
                return True
            seen.add(pair)
            if not same_data(part,target):
                return False
            first,second=_rels_path(part),_rels_path(target)
            if (first in incoming)!=(second in entries):
                return False
            if first not in incoming:
                return True
            left={n.get('Id'):n for n in _xml(incoming[first])}
            right={n.get('Id'):n for n in _xml(entries[second])}
            if left.keys()!=right.keys():
                return False
            for key,a in left.items():
                b=right[key]
                if {k:v for k,v in a.attrib.items() if k!='Target'}!={k:v for k,v in b.attrib.items() if k!='Target'}:
                    return False
                if a.get('TargetMode')=='External':
                    if a.get('Target')!=b.get('Target'):
                        return False
                elif not equivalent(_part_relative(part,a.get('Target','')),
                                    _part_relative(target,b.get('Target','')),seen):
                    return False
            return True

        def include(part):
            if part in visited:
                return mapping[part]
            if part not in incoming:
                raise ValueError('Một tài nguyên của nhóm kết quả bị thiếu.')
            visited.add(part)
            if part.startswith('ppt/notesMasters/') and not part.endswith('.rels'):
                # Keep the host's single notes master and import note text.
                for relation in _xml(entries['ppt/_rels/presentation.xml.rels']):
                    if relation.get('Type','').endswith('/notesMaster'):
                        mapping[part]=_part_relative('ppt/presentation.xml',relation.get('Target',''))
                        return mapping[part]
            if part.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/notesMasters/','ppt/theme/')):
                for possible in list(entries):
                    if (possible.startswith('ppt/'+part.split('/')[1]+'/') and same_data(part,possible)
                            and equivalent(part,possible)):
                        mapping[part]=possible
                        return possible
            if part.startswith('ppt/media/'):
                digest = hashlib.sha256(incoming[part]).hexdigest()
                if digest in restore:
                    original = restore[digest]
                    if original not in media:
                        raise ValueError('Không tìm thấy media gốc để ghép vào bài giảng.')
                    mapping[part] = media[original]
                    restored.add(digest)
                    return mapping[part]
                if digest in media:
                    restored.update(key for key,value in restore.items() if value==digest)
                    mapping[part] = media[digest]
                    return mapping[part]
            target = mapping.setdefault(part,prefix+part)
            entries[target] = incoming[part]
            if part.startswith('ppt/media/'):
                media[digest] = target
            relationships = _rels_path(part)
            if relationships in incoming:
                tree = _xml(incoming[relationships])
                for relationship in tree:
                    if relationship.get('TargetMode') == 'External':
                        continue
                    dependency = _part_relative(part,relationship.get('Target',''))
                    if dependency.startswith('../') or dependency.startswith('/'):
                        raise ValueError('Liên kết tài nguyên nằm ngoài gói PowerPoint.')
                    copied = include(dependency)
                    relationship.set('Target',posixpath.relpath(copied,posixpath.dirname(target)))
                entries[_rels_path(target)] = _bytes(tree)
            return target

        for part in list(mapping):
            include(part)
        if set(restore)-restored:
            raise ValueError('Một video/âm thanh đã bị bỏ khỏi nhóm kết quả; chưa thay bản gốc.')
        existing_overrides = {node.get('PartName') for node in types if node.tag.endswith('Override')}
        defaults = {node.get('Extension'):node.get('ContentType') for node in types if node.tag.endswith('Default')}
        for node in incoming_types:
            if node.tag.endswith('Override'):
                name = node.get('PartName','').lstrip('/')
                if name not in visited:
                    continue
                target = '/'+mapping[name]
                if target not in existing_overrides:
                    added = copy.deepcopy(node)
                    added.set('PartName',target)
                    types.append(added)
                    existing_overrides.add(target)
            elif node.tag.endswith('Default'):
                extension = node.get('Extension')
                if extension not in defaults:
                    types.append(copy.deepcopy(node))
                    defaults[extension] = node.get('ContentType')
                elif defaults[extension] != node.get('ContentType'):
                    raise ValueError('Kiểu tài nguyên của nhóm kết quả không tương thích.')
        # Office requires every referenced master to be registered at the
        # presentation root, including a genuinely changed incoming template.
        presentation=_xml(entries['ppt/presentation.xml'])
        relations=_xml(entries['ppt/_rels/presentation.xml.rels'])
        if presentation.find('p:notesMasterIdLst',NS) is None:
            existing=next((r for r in relations if r.get('Type','').endswith('/notesMaster')),None)
            if existing is not None:
                container=etree.Element('{'+NS['p']+'}notesMasterIdLst')
                node=etree.SubElement(container,'{'+NS['p']+'}notesMasterId')
                node.set('{'+NS['r']+'}id',existing.get('Id'))
                presentation.insert(1 if presentation.find('p:sldMasterIdLst',NS) is not None else 0,container)
        for kind,list_name,item_name in (('slideMaster','sldMasterIdLst','sldMasterId'),
                                         ('notesMaster','notesMasterIdLst','notesMasterId')):
            targets={_part_relative('ppt/presentation.xml',r.get('Target','')) for r in relations
                     if r.get('Type','').endswith('/'+kind)}
            for relation_name in [name for name in entries if name.endswith('.rels')]:
                owner=posixpath.join(posixpath.dirname(posixpath.dirname(relation_name)),posixpath.basename(relation_name)[:-5])
                for relation in _xml(entries[relation_name]):
                    if not relation.get('Type','').endswith('/'+kind) or relation.get('TargetMode')=='External':
                        continue
                    target=_part_relative(owner,relation.get('Target',''))
                    if target in targets:
                        continue
                    container=presentation.find('p:'+list_name,NS)
                    if container is None:
                        container=etree.Element('{'+NS['p']+'}'+list_name)
                        presentation.insert(0 if kind=='slideMaster' else 1,container)
                    taken={r.get('Id') for r in relations}
                    index=1
                    while 'rIdBiliClass'+str(index) in taken:
                        index+=1
                    rid='rIdBiliClass'+str(index)
                    added=copy.deepcopy(relation)
                    added.set('Id',rid)
                    added.set('Target',posixpath.relpath(target,'ppt'))
                    relations.append(added)
                    node=etree.SubElement(container,'{'+NS['p']+'}'+item_name)
                    node.set('{'+NS['r']+'}id',rid)
                    if kind=='slideMaster':
                        all_ids=[int(n.get('id')) for n in container if n.get('id')]
                        for name,data in entries.items():
                            if name.endswith('.xml') and '/slideMasters/' in name:
                                all_ids.extend(int(v) for v in _xml(data).xpath('.//p:sldLayoutId/@id',namespaces=NS))
                        next_id=max([2147483647]+all_ids)+1
                        node.set('id',str(next_id))
                        master=_xml(entries[target])
                        for layout in master.xpath('.//p:sldLayoutId',namespaces=NS):
                            next_id+=1
                            layout.set('id',str(next_id))
                        entries[target]=_bytes(master)
                    targets.add(target)
        entries['ppt/presentation.xml']=_bytes(presentation)
        entries['ppt/_rels/presentation.xml.rels']=_bytes(relations)
    entries['[Content_Types].xml'] = _bytes(types)
    destination.parent.mkdir(parents=True,exist_ok=True)
    pending = destination.with_suffix('.merge-pending.pptx')
    try:
        with zipfile.ZipFile(pending,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=1) as archive:
            for name,data in entries.items():
                archive.writestr(name,data)
        Presentation(pending)
        pending.replace(destination)
    finally:
        pending.unlink(missing_ok=True)
    return destination
