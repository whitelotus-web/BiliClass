"""Explicit portable-build test using its own temporary data directory."""

import json
import sys
import tempfile
import traceback
from pathlib import Path

from .importers import parse_document
from .library import Library
from .pack import export_pack, import_pack
from .speech import list_voices, synthesize, valid_audio
from .translation import translate_draft


def run(output):
    report = {
        "status": "running",
        "checks": [],
        "scope": ("packaged" if getattr(sys, "frozen", False) else "source") + " local pipeline; not a clean-Windows certification",
    }
    library = None
    try:
        with tempfile.TemporaryDirectory(prefix="biliclass-diagnostic-") as temporary:
            directory = Path(temporary)
            library = Library(directory / "library")
            try:
                from docx import Document
                from pptx import Presentation

                document = Document()
                document.add_paragraph("Hãy thảo luận theo nhóm.")
                document.save(directory / "Bài.docx")
                assert parse_document(directory / "Bài.docx")[0][1] == "Hãy thảo luận theo nhóm."
                deck = Presentation()
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                slide.shapes.title.text = "Thảo luận theo nhóm"
                deck.save(directory / "Bài.pptx")
                assert parse_document(directory / "Bài.pptx")[0][1] == "Thảo luận theo nhóm"
                report["checks"].append("DOCX and PPTX create/read with packaged dependencies")
                lesson = library.create(
                    "Kiểm tra", "Liên môn", "THPT", "11", [("Đoạn", "Thảo luận theo nhóm.")]
                )
                en = translate_draft("Thảo luận theo nhóm.", "vi")
                vi = translate_draft("Discuss in groups.", "en")
                assert en.strip() and vi.strip()
                library.edit_segment(
                    lesson["id"], lesson["segments"][0]["id"], "Thảo luận theo nhóm.", en, True
                )
                report["translations"] = {"vi_en": en, "en_vi": vi}
                report["checks"].append("both local model directions load and produce drafts")
                pack = export_pack(library, lesson["id"], directory / "Bài.biliclass")
                copied = import_pack(library, pack)
                assert copied["segments"][0]["en"] == en and not copied["segments"][0]["approved"]
                assert library.history(lesson["id"])
                restored = library.restore_revision(lesson["id"], 1)
                assert restored["segments"][0]["en"] == ""
                report["checks"].append("SQLite, history and draft pack roundtrip")
                voices = list_voices()
                report["voices"] = [{"name": v["name"], "language": v["language"]} for v in voices]
                english = next((v for v in voices if v["language"] == "en"), None)
                if english:
                    first = synthesize("Please explain your answer.", english["id"], 0, directory / "audio")
                    second = synthesize("Please explain your answer.", english["id"], 0, directory / "audio")
                    assert valid_audio(first["path"]) and second["cached"]
                    report["tts"] = {"voice": english["id"], "provider": english.get("provider", "sapi")}
                    report["checks"].append("offline English WAV synthesized and cache reused; no playback")
                vietnamese = next((v for v in voices if v.get("provider") == "vieneu"), None)
                if vietnamese:
                    first = synthesize("Chúng ta cùng tìm hiểu một ý tưởng mới.", vietnamese["id"], 0, directory / "audio")
                    second = synthesize("Chúng ta cùng tìm hiểu một ý tưởng mới.", vietnamese["id"], 0, directory / "audio")
                    adjusted = synthesize("Chào cả lớp.", vietnamese["id"], 1, directory / "audio")
                    assert valid_audio(first["path"]) and second["cached"] and valid_audio(adjusted["path"])
                    report["tts_vi"] = {"voice": vietnamese["id"], "provider": "vieneu", "cached": True}
                    report["checks"].append("offline Vietnamese VieNeu WAV, speed and cache; no playback")
                library.edit_segment(lesson["id"], lesson["segments"][0]["id"], "Thảo luận theo nhóm.", en, True)
                prepared = library.save_question(lesson["id"], {"kind": "single", "vi": "Chọn bằng chứng", "en": "Choose evidence", "approved": True,
                    "concept_id": lesson["segments"][0]["id"], "concept_label": "Lập luận", "correct": "A",
                    "options": [{"vi": "Dữ liệu", "en": "Data"}, {"vi": "Đoán", "en": "Guess"}]})
                prepared = library.mark_prepared(lesson["id"], prepared["revision"])
                import httpx
                from websockets.sync.client import connect

                from .analytics import report as classroom_report
                from .classroom_server import ClassroomRuntime
                runtime = ClassroomRuntime(library.directory)
                try:
                    info = runtime.start(prepared, {"title": "Bài thử", "mode": "anonymous", "size": 50, "host": "127.0.0.1"})
                    public = f"http://127.0.0.1:{info['port']}"
                    with httpx.Client(trust_env=False) as client:
                        assert client.get(public + "/").status_code == 200
                        student = client.post(public + "/api/join", json={"join": info["join"]}).json()
                        assert client.get(public + "/state").status_code == 404
                    state = runtime.request({"action": "open", "question_id": prepared["questions"][0]["id"], "duration": 60})
                    with connect(public.replace("http", "ws") + "/student", proxy=None) as socket:
                        socket.send(json.dumps({"token": student["token"]}))
                        student_state = json.loads(socket.recv(timeout=5))
                        assert "correct" not in student_state["round"]
                        socket.send(json.dumps({"submission_id": "diagnostic", "round_id": state["round"]["id"], "answer": "A"}))
                        for _ in range(10):
                            if json.loads(socket.recv(timeout=5))["type"] == "ack":
                                break
                        else:
                            raise AssertionError("No answer ACK")
                    runtime.request({"action": "close"})
                    runtime.request({"action": "end"})
                    assert classroom_report(library.directory, info["session_id"])["rounds"][0]["percent"] == 100
                    report["checks"].append("spawned classroom process, bundled student page, authenticated WebSocket ACK and saved report")
                finally:
                    runtime.stop()
                from .deck_export import export_deck
                from .ocr import languages
                from .storage import backup_library, restore_backup
                output_deck = export_deck(prepared, library.directory, directory / "Bilingual.pptx")
                assert len(Presentation(output_deck).slides) == 1
                restore_backup(backup_library(library.directory), directory / "restored")
                report["checks"].append("reviewed deck export and consistent library/classroom backup restore")
                report["ocr_languages"] = languages()
                if any(tag.startswith("en") for tag in report["ocr_languages"]):
                    from PIL import Image, ImageDraw, ImageFont
                    image = Image.new("RGB", (900, 140), "white")
                    ImageDraw.Draw(image).text((20, 30), "Discuss evidence in groups.", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 36), fill="black")
                    image.save(directory / "ocr.png")
                    assert "Discuss" in parse_document(directory / "ocr.png", "en")[0][1]
                    report["checks"].append("isolated Windows English OCR and installed native dependencies")
            finally:
                library.close()
        report["status"] = "passed"
    except Exception:
        report["status"] = "failed"
        report["error"] = traceback.format_exc()
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["status"] == "passed" else 1
