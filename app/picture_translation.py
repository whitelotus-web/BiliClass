"""Editable overlays on the original picture; never alter raster evidence."""

import hashlib
import json
import math
import zipfile
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image

from .layout_regions import empty_region, intersection
from .source_deck import NS, _bytes, _part_relative, _rels_path, _xml, source_slide_parts


def text_shape(identity, text, box, size, *, fill=None, color="112650", english=True):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(*(round(value * 12700) for value in box))
    shape.name = "BCL_EN image translation" if english else "BiliClass label"
    shape._element.find("p:nvSpPr/p:cNvPr", NS).set("id", str(identity))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(1.5)
    frame.margin_top = frame.margin_bottom = Pt(1)
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.italic = english
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.05
        for run in paragraph.runs:
            run._r.get_or_add_rPr().set("lang", "en-US" if english else "vi-VN")
    return shape._element


def background(image, picture_box, region, *, ring=False):
    x, y, w, h = picture_box
    left, top, width, height = region
    if left < x or top < y or left + width > x + w or top + height > y + h:
        return None
    crop = image.crop(
        (
            math.floor((left - x) / w * image.width),
            math.floor((top - y) / h * image.height),
            math.ceil((left + width - x) / w * image.width),
            math.ceil((top + height - y) / h * image.height),
        )
    )
    try:
        if crop.mode=='RGBA' and np.mean(np.asarray(crop.getchannel('A'))<250)>.01:
            return None
        data = np.asarray(crop.convert("RGB"))
        if not data.size:
            return None
        samples = (
            np.concatenate(
                (
                    data[:2].reshape(-1, 3),
                    data[-2:].reshape(-1, 3),
                    data[:, :2].reshape(-1, 3),
                    data[:, -2:].reshape(-1, 3),
                )
            )
            if ring
            else data.reshape(-1, 3)
        )
        color = np.median(samples, axis=0)
        # A busy/gradient border must not be treated as a plain covering color.
        if np.mean(np.max(abs(samples.astype(float) - color), axis=1) < 22) < 0.97:
            return None
        return "".join(f"{int(channel):02X}" for channel in color)
    finally:
        crop.close()


def apply_picture_translations(path, manifest, rows, translations, mode, cancel=None):
    from pptx import Presentation

    from .native_layout import text_height
    from .slide_speech import objects

    path = Path(path)
    deck = Presentation(path)
    parts = source_slide_parts(path)
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    plan = manifest["layout_plan"]
    applied = []
    warnings = []
    note_additions = {}
    images = {}
    canvases = {}
    prose_pages = [page for page, item in enumerate(plan, 1) if any(
        row.get('draft') and row['slide'] == item['source_slide'] for row in rows)]
    if prose_pages and mode != 'english_only':
        from .office_tasks import render_quality

        directory = path.parent/'picture-background'
        report = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                  'slides': [dict(slide=n, issues=[]) for n in prose_pages]}
        render_quality(path, directory, report, cancel)
        canvases = {row['slide']: directory/row['render_image'] for row in report['slides'] if row.get('rendered')}
    try:
        for page, slide in enumerate(deck.slides, 1):
            if cancel is not None and cancel.is_set():
                raise InterruptedError("Đã dừng sắp xếp bản dịch trong ảnh.")
            matching = [row for row in rows if row["slide"] == plan[page - 1]["source_slide"]]
            if not matching:
                continue
            tree = _xml(entries[parts[page - 1]])
            shape_tree = tree.find("p:cSld/p:spTree", NS)
            identity = max([int(v) for v in tree.xpath(".//p:cNvPr/@id", namespaces=NS)] + [1])
            items = objects(slide)
            placed = []
            for row in matching:
                if row.get('draft') and mode == 'english_only':
                    warnings.append(f"Slide {page}: đoạn đọc từ ảnh cần đối chiếu; giữ chữ gốc thay vì che bằng bản dịch nháp.")
                    continue
                text = translations[row["id"]]
                original = tuple(row["region"])
                x, y, w, h = original
                occupied = [
                    tuple(item[k] / 12700 for k in ("x", "y", "w", "h"))
                    for item in items
                    if item["shape_id"] != row["shape_id"] and not item["hidden"]
                ]
                # Background pictures are allowed only where the bitmap is blank.
                occupied = [
                    box
                    for box in occupied
                    if box[2] * box[3] < deck.slide_width * deck.slide_height / 12700**2 * 0.6
                ]
                # Do not translate raster duplicates behind editable source text.
                if any(intersection(original, box) > w * h * 0.25 for box in occupied):
                    warnings.append(f"Slide {page}: bỏ vùng ảnh trùng/chồng đối tượng chữ hoặc hình khác.")
                    continue
                image = images.get(row["image"])
                if image is None:
                    with Image.open(row['image']) as original_image:
                        image = original_image.convert('RGBA')
                    images[row["image"]] = image
                other_regions = [tuple(r["region"]) for r in matching if r["id"] != row["id"]]
                occupied += placed + other_regions
                box = None
                fill = None
                size = 16
                if mode == "english_only":
                    for scale in (1, 1.4, 1.8):
                        for grow in (1.3, 2, 3):
                            candidate = (x - 2, y - 2, w * scale + 4, h * grow + 4)
                            if any(intersection(candidate, other) > 0.5 for other in occupied):
                                continue
                            fill = background(image, row["box"], candidate, ring=True)
                            if not fill:
                                continue
                            # Only the original glyph region may contain dark pixels.
                            extra = (
                                candidate[0] + w + 4,
                                candidate[1],
                                max(0, candidate[2] - w - 4),
                                candidate[3],
                            )
                            below = (
                                candidate[0],
                                y + h + 2,
                                candidate[2],
                                max(0, candidate[1] + candidate[3] - y - h - 2),
                            )
                            if any(
                                a[2] > 2 and a[3] > 2 and not background(image, row["box"], a)
                                for a in (extra, below)
                            ):
                                continue
                            size = min(24, max(14, h * 0.85))
                            if text_height(text, candidate[2] - 3, size) <= candidate[3] - 2:
                                box = candidate
                                break
                        if box:
                            break
                else:
                    size = min(20, max(14, h * 0.75))
                    width = min(360, max(160, w * 1.3))
                    height = text_height(text, width - 4, size) + 6
                    box = empty_region(
                        original,
                        (width, height),
                        (deck.slide_width / 12700, deck.slide_height / 12700),
                        occupied + [original],
                        clear=lambda candidate: bool(background(image, row["box"], candidate)),
                    )
                    if box:
                        fill = background(image, row["box"], box)
                    if box is None and row.get('draft') and page in canvases:
                        canvas_key = 'canvas-' + str(page)
                        if canvas_key not in images:
                            with Image.open(canvases[page]) as rendered:
                                images[canvas_key] = rendered.convert('RGB')
                        canvas = images[canvas_key]
                        bounds = (deck.slide_width/12700, deck.slide_height/12700)
                        width = min(bounds[0]-24, max(w, bounds[0]*.92))
                        for size in (32, 28, 24, 20, 18):
                            height = text_height(text, width-6, size) + 8
                            box = empty_region(original, (width, height), bounds, occupied+[tuple(row['box'])],
                                clear=lambda candidate: bool(background(canvas, (0, 0, *bounds), candidate)))
                            if box:
                                fill = background(canvas, (0, 0, *bounds), box)
                                warnings.append(f"Slide {page}: đã đặt bản dịch nháp cho đoạn chữ trong ảnh; cần đối chiếu với ảnh gốc, chưa tự đọc đoạn này.")
                                break
                if box is None:
                    warnings.append(
                        f"Slide {page}: chưa đủ vùng trống/nền rõ để đặt bản dịch chữ trong ảnh; cần đối chiếu."
                    )
                    continue
                identity += 1
                rgb = [int(fill[i : i + 2], 16) for i in (0, 2, 4)]
                color = "FFFFFF" if sum(rgb) < 300 else "112650"
                overlay = text_shape(identity, text, box, size, fill=fill, color=color)
                from .picture_narration import PREFIX
                overlay.find('p:nvSpPr/p:cNvPr', NS).set('descr', PREFIX + json.dumps(
                    {'vi': row['text'], 'en': text}, ensure_ascii=False))
                shape_tree.append(overlay)
                if row.get('animation_source'):
                    from .same_slide import synchronize

                    synchronize(tree, {row['animation_source']: identity}, {})
                placed.append(box)
                applied.append(
                    {
                        "slide": page,
                        "source_shape": row["shape_id"],
                        "shape_id": identity,
                        "text": text,
                        "vi": row["text"],
                        "box": box,
                        "confidence": row["confidence"],
                    }
                )
                rels = _xml(entries[_rels_path(parts[page - 1])])
                note = next((r for r in rels if r.get("Type", "").endswith("/notesSlide")), None)
                if note is not None:
                    part = _part_relative(parts[page - 1], note.get("Target"))
                    pair = (row["text"], text)
                    if pair not in note_additions.setdefault(part, []):
                        note_additions[part].append(pair)
            entries[parts[page - 1]] = _bytes(tree)
        for part, pairs in note_additions.items():
            tree = _xml(entries[part])
            bodies = tree.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody', namespaces=NS)
            if not bodies:
                continue
            body = bodies[0]
            for vi, en in pairs:
                for text in ("VI: " + vi, "EN: " + en):
                    paragraph = etree.SubElement(body, f"{{{NS['a']}}}p")
                    run = etree.SubElement(paragraph, f"{{{NS['a']}}}r")
                    etree.SubElement(run, f"{{{NS['a']}}}t").text = text
            entries[part] = _bytes(tree)
        if applied:
            pending = path.with_suffix(".image-pending.pptx")
            try:
                with zipfile.ZipFile(pending, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as archive:
                    for name, data in entries.items():
                        archive.writestr(name, data)
                if cancel is not None and cancel.is_set():
                    raise InterruptedError("Đã dừng; bản trình chiếu trước vẫn được giữ.")
                pending.replace(path)
            finally:
                pending.unlink(missing_ok=True)
    finally:
        for image in images.values():
            image.close()
    return {"applied": applied, "warnings": list(dict.fromkeys(warnings))}
