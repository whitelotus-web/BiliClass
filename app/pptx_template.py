"""Populate editable PowerPoint frames; keep the template's native package.

Templates are local documents, never prompts or executable code. Page numbers
are explicit, and zero disables optional cover / intro / closing pages.
"""

import copy
import hashlib
import io
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.slide import Slide
from pptx.util import Pt

from .source_deck import NS, _append_slide, _bytes, _xml, source_slide_parts


def frames(slide):
    # Group-relative positions cannot safely be treated as slide coordinates.
    shapes = [s for s in slide.shapes if s.has_text_frame and s.width > 0 and s.height > 0]
    title = next((s for s in shapes if s.name.upper() == 'BCL_TITLE' or
                  re.search(r'\[(LESSON_TITLE|SOURCE_HEADING)\]', s.text)), None)
    title = title or slide.shapes.title
    if title is None:
        title = next(iter(sorted((s for s in shapes if s.text.strip()), key=lambda s: (s.top, -s.width))), None)
    bodies = [s for s in shapes if s is not title and s.width >= Pt(140) and s.height >= Pt(35)]
    body = next((s for s in bodies if s.name.upper() == 'BCL_CONTENT'), None)
    if body is None:
        for marker in ('SOURCE_CONTENT', 'SOURCE_POINT_1', 'SOURCE_QUESTION', 'SOURCE_TASK', 'SOURCE_SUBTITLE'):
            body = next((s for s in bodies if '[' + marker + ']' in s.text), None)
            if body is not None:
                break
    body = body or max(bodies, key=lambda s: s.width * s.height, default=None)
    return title, body


def inspect(path):
    from .document_limits import check_document_size

    path = Path(path)
    if path.suffix.lower() != '.pptx':
        raise ValueError('Mẫu cần là file .pptx có chữ chỉnh sửa được.')
    check_document_size(path)
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10_000 or sum(i.file_size for i in infos) > 512 * 1024**2:
            raise ValueError('Mẫu PowerPoint chứa quá nhiều tài nguyên.')
        if len({i.filename for i in infos}) != len(infos) or any('vbaProject' in i.filename for i in infos):
            raise ValueError('Mẫu PowerPoint không hợp lệ hoặc chứa macro.')
    deck = Presentation(path)
    if not 1 <= len(deck.slides) <= 30:
        raise ValueError('Chọn mẫu có từ 1 đến 30 trang.')
    pages = []
    for index, slide in enumerate(deck.slides, 1):
        title, body = frames(slide)
        pages.append({'number': index, 'title': (title.text[:90] if title is not None else ''),
                      'usable': bool(title is not None and body is not None and body.height >= Pt(60))})
    content = next((p['number'] for p in pages[2:] if p['usable']), None) if len(pages) >= 4 else None
    content = content or next((p['number'] for p in pages[1:] if p['usable']), None)
    content = content or next((p['number'] for p in pages if p['usable']), None)
    if content is None:
        raise ValueError('Mẫu chưa có khung tiêu đề và nội dung đủ rộng. Đặt tên ô chữ BCL_TITLE và BCL_CONTENT trong Selection Pane của PowerPoint.')
    return {'name': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'pages': pages,
            'roles': {'cover': 1 if content != 1 else 0, 'intro': 2 if len(pages) >= 4 and pages[1]['usable'] else 0,
                      'content': content, 'ending': len(pages) if len(pages) >= 4 else 0}}


def validate_roles(metadata, roles):
    count = len(metadata['pages'])
    normalized = {key: int(roles.get(key, 0)) for key in ('cover', 'intro', 'content', 'ending')}
    if any(not 0 <= n <= count for n in normalized.values()) or not normalized['content']:
        raise ValueError('Chọn trang nội dung; các trang tuỳ chọn dùng 0 để bỏ qua.')
    if any(not metadata['pages'][n - 1]['usable'] for key, n in normalized.items() if n and key in {'intro', 'content'}):
        raise ValueError('Trang intro/nội dung cần có ô tiêu đề và ô nội dung chỉnh sửa được.')
    return normalized


def store(directory, path):
    metadata = inspect(path)
    target = Path(directory) / 'templates' / (metadata['sha256'] + '.pptx')
    target.parent.mkdir(parents=True, exist_ok=True)
    data = Path(path).read_bytes()
    if hashlib.sha256(data).hexdigest() != metadata['sha256']:
        raise ValueError('File mẫu vừa thay đổi; chọn lại mẫu.')
    target.write_bytes(data)
    return metadata


def stored_path(directory, metadata):
    digest = metadata.get('sha256', '')
    if not re.fullmatch('[a-f0-9]{64}', digest):
        raise ValueError('Mẫu đã lưu không hợp lệ; chọn lại mẫu.')
    path = Path(directory) / 'templates' / (digest + '.pptx')
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError('File mẫu đã thay đổi hoặc bị xoá; chọn lại mẫu.')
    return path


def _size(shape, fallback):
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
    sizes += [p.font.size.pt for p in shape.text_frame.paragraphs if p.font.size]
    return min(32, max(18, max(sizes, default=fallback)))


def _height(shape, texts, size):
    from .native_layout import text_height

    frame = shape.text_frame
    width = (shape.width - frame.margin_left - frame.margin_right) / 12700
    return sum(text_height(text, width, size * (.82 if english else 1)) + size * .18
               for text, english in texts if text)


def _texts(vi, en, mode):
    return [(en, True)] if mode == 'english_only' else [(vi, False), (en, True)]


def _write(shape, texts, *, title=False):
    """Keep the frame, paint, alignment and representative run styling."""
    size = _size(shape, 28 if title else 24)
    frame = shape.text_frame
    available = (shape.height - frame.margin_top - frame.margin_bottom) / 12700
    minimum = 14 if title else 18
    while size > minimum and _height(shape, texts, size) > available:
        size -= 1
    if _height(shape, texts, size) > available:
        raise ValueError('Khung chữ mẫu quá hẹp cho tiêu đề/nội dung. Chọn khung lớn hơn; tool không cắt mất chữ.')
    original = copy.deepcopy(frame.paragraphs[0]._p)
    properties = original.find('a:pPr', NS)
    run_props = original.find('a:r/a:rPr', NS)
    if run_props is None:
        run_props = original.find('a:pPr/a:defRPr', NS)
    frame.clear()
    frame.word_wrap = True
    for index, (text, english) in enumerate(texts):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        if properties is not None:
            old = paragraph._p.find('a:pPr', NS)
            if old is not None:
                paragraph._p.remove(old)
            paragraph._p.insert(0, copy.deepcopy(properties))
        for run in paragraph.runs:
            if run_props is not None:
                old = run._r.find('a:rPr', NS)
                if old is not None:
                    run._r.remove(old)
                run._r.insert(0, copy.deepcopy(run_props))
            run.font.size = Pt(size * (.82 if english else 1))
            run.font.italic = english
            run.font.bold = title
            run._r.get_or_add_rPr().set('lang', 'en-US' if english else 'vi-VN')
        paragraph.line_spacing = 1.1
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(size * .18)


def _columns(body, mode):
    return mode == 'parallel_columns' and body.width >= Pt(450)


def _cost(body, vi, en, mode, size):
    if _columns(body, mode):
        from .native_layout import text_height
        width = (body.width / 2 - Pt(12)) / 12700
        return max(text_height(vi, width, size), text_height(en, width, size * .82)) + size * .25
    return _height(body, _texts(vi, en, mode), size)


def _body(slide, body, content, mode):
    if not _columns(body, mode):
        _write(body, [item for _, vi, en in content for item in _texts(vi, en, mode)])
        return
    # Keep the original full-width painted frame, with transparent text above.
    original = copy.deepcopy(body.text_frame._txBody)
    size = _size(body, 24)
    body.text_frame.clear()
    y = body.top + body.text_frame.margin_top
    x = body.left + body.text_frame.margin_left
    width = int((body.width - body.text_frame.margin_left - body.text_frame.margin_right - Pt(16)) / 2)
    for _, vi, en in content:
        height = Pt(_cost(body, vi, en, mode, size))
        for english, value in ((False, vi), (True, en)):
            shape = slide.shapes.add_textbox(x + (width + Pt(16) if english else 0), int(y), width, int(height))
            shape.name = 'BCL_BODY_EN' if english else 'BCL_BODY_VI'
            shape._element.replace(shape.text_frame._txBody, copy.deepcopy(original))
            _write(shape, [(value, english)])
        y += height


def _pages(pairs, body, mode):
    """Pack complete pairs, subdividing an oversized pair without dropping text."""
    size = _size(body, 24)
    available = (body.height - body.text_frame.margin_top - body.text_frame.margin_bottom) / 12700
    pending = list(pairs)
    page, used = [], 0
    while pending:
        locator, vi, en = pending.pop(0)
        cost = _cost(body, vi, en, mode, size)
        if cost > available:
            left, right = vi.split(), en.split()
            if max(len(left), len(right)) < 4:
                raise ValueError('Khung nội dung quá nhỏ; chọn trang mẫu có vùng nội dung lớn hơn.')
            halves = []
            for i in range(2):
                halves.append((locator, ' '.join(left[len(left)*i//2:len(left)*(i+1)//2]),
                               ' '.join(right[len(right)*i//2:len(right)*(i+1)//2])))
            pending[0:0] = halves
            continue
        if page and used + cost > available:
            yield page
            page, used = [], 0
        page.append((locator, vi, en))
        used += cost
    if page:
        yield page


def _visual_pages(deck, source, number):
    """Stage image pages before serializing, so all media relationships exist."""
    if not source or Path(source).suffix.lower() != '.pptx':
        return []
    from .slide_speech import objects

    original, template = Presentation(source), deck.slides[number - 1]
    pages = []
    for source_slide in original.slides:
        pictures = [o for o in objects(source_slide) if int(o['shape'].shape_type) == 13 and not o['hidden']]
        if not pictures:
            continue
        slide = deck.slides.add_slide(template.slide_layout)
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        mapping = {}
        for relation in template.part.rels.values():
            if relation.reltype.endswith(('/notesSlide', '/slide')):
                continue
            mapping[relation.rId] = (slide.part.rels.get_or_add_ext_rel(relation.reltype, relation.target_ref)
                                    if relation.is_external else
                                    slide.part.rels.get_or_add(relation.reltype, relation.target_part))
        for shape in template.shapes:
            element = copy.deepcopy(shape._element)
            for node in element.iter():
                for key, value in list(node.attrib.items()):
                    if key.startswith('{' + NS['r'] + '}') and value in mapping:
                        node.set(key, mapping[value])
            slide.shapes._spTree.insert_element_before(element, 'p:extLst')
        background = template._element.find('p:cSld/p:bg', NS)
        if background is not None:
            slide._element.find('p:cSld', NS).insert(0, copy.deepcopy(background))
        _, body = frames(slide)
        body.text_frame.clear()
        scale = min(body.width / original.slide_width, body.height / original.slide_height)
        for o in pictures:
            picture = slide.shapes.add_picture(io.BytesIO(o['shape'].image.blob),
                int(body.left + o['x'] * scale), int(body.top + o['y'] * scale),
                width=max(1, int(o['w'] * scale)), height=max(1, int(o['h'] * scale)))
            for attr in ('crop_left', 'crop_right', 'crop_top', 'crop_bottom'):
                setattr(picture, attr, getattr(o['shape'], attr))
        pages.append(len(deck.slides))
    return pages


def build(target, template, config, pairs, cancel, source=None, roles=None):
    """Clone native pages/resources and fill only the selected editable frames."""
    metadata = inspect(template)
    roles = validate_roles(metadata, roles or metadata['roles'])
    deck = Presentation(template)
    visuals = _visual_pages(deck, source, roles['content'])
    # Ensure each cloned page has independent narration notes before ZIP edits.
    for slide in deck.slides:
        slide.notes_slide.notes_text_frame.text = ''
    buffer = io.BytesIO()
    deck.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as archive:
        entries = {item.filename: archive.read(item) for item in archive.infolist()}
    parts = source_slide_parts(buffer)
    types = _xml(entries['[Content_Types].xml'])
    relations = _xml(entries['ppt/_rels/presentation.xml.rels'])
    presentation = _xml(entries['ppt/presentation.xml'])
    slide_list = presentation.find('p:sldIdLst', NS)
    originals = list(slide_list)
    added, notes = [], []
    mode = config['conversion_format']
    title_pair = _texts(config['title'], config['english_title'], mode)

    def page(role, content=(), closing=False, number=None):
        if cancel.is_set():
            raise InterruptedError('Đã dừng dựng bài theo mẫu.')
        number = number or roles[role]
        original = deck.slides[number - 1]
        slide = Slide(copy.deepcopy(original._element), original.part)
        title, body = frames(slide)
        protected = {id(s._element) for s in (title, body) if s is not None}
        for shape in slide.shapes:
            if shape.has_text_frame and id(shape._element) not in protected:
                text = shape.text
                if shape.name == 'BCL_DECORATION page':
                    runs = shape.text_frame.paragraphs[0].runs
                    if runs:
                        runs[0].text = f'{len(added) + 1:02}'
                # Keep static template artwork; neutral markers never leak.
                if re.search(r'\[[A-Z_0-9]+\]', text):
                    shape.text_frame.clear()
        if title is not None:
            _write(title, _texts('Cảm ơn!', 'Thank you!', mode) if closing else title_pair, title=True)
        vi, en = '\n'.join(p[1] for p in content), '\n'.join(p[2] for p in content)
        if body is not None:
            if content:
                _body(slide, body, content, mode)
            else:
                body.text_frame.clear()
        token = 'template-' + str(len(added) + 1)
        sid = _append_slide(parts[number - 1], copy.deepcopy(slide._element), token, entries, types,
                            relations, slide_list, originals[number - 1], 256 + len(originals) + len(added))
        added.append(sid)
        note_vi, note_en = ('Cảm ơn!', 'Thank you!') if closing else (config['title'], config['english_title'])
        notes.append((token, '\n'.join(filter(None, (note_vi, vi))), '\n'.join(filter(None, (note_en, en)))))

    if roles['cover']:
        page('cover')
    pending = list(pairs)
    if roles['intro'] and pending:
        intro_body = frames(deck.slides[roles['intro'] - 1])[1]
        intro = list(_pages([pending.pop(0)], intro_body, mode))
        for content in intro:
            page('intro', content)
    body = frames(deck.slides[roles['content'] - 1])[1]
    for content in _pages(pending, body, mode):
        page('content', content)
    for number in visuals:
        page('content', number=number)
    if roles['ending']:
        page('ending', closing=True)
    if not added:
        raise ValueError('Không có nội dung để điền vào mẫu.')
    for sid in list(slide_list):
        slide_list.remove(sid)
    for sid in added:
        slide_list.append(sid)
    active = {s.get('{' + NS['r'] + '}id') for s in added}
    for relation in list(relations):
        if relation.get('Type', '').endswith('/slide') and relation.get('Id') not in active:
            relations.remove(relation)
    for token, vi, en in notes:
        part = f'ppt/notesSlides/biliclass-en-{token}.xml'
        tree = _xml(entries[part])
        shape = tree.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS)[0]
        body_node = shape.find('p:txBody', NS)
        for p in list(body_node.findall('a:p', NS)):
            body_node.remove(p)
        for line in ('BILICLASS_NOTES:', 'VI: ' + vi, 'EN: ' + en):
            p = etree.SubElement(body_node, '{' + NS['a'] + '}p')
            r = etree.SubElement(p, '{' + NS['a'] + '}r')
            etree.SubElement(r, '{' + NS['a'] + '}t').text = line
        entries[part] = _bytes(tree)
    entries['ppt/presentation.xml'] = _bytes(presentation)
    entries['ppt/_rels/presentation.xml.rels'] = _bytes(relations)
    entries['[Content_Types].xml'] = _bytes(types)
    target = Path(target)
    pending_path = target.with_suffix('.pending.pptx')
    try:
        with zipfile.ZipFile(pending_path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        pending_path.replace(target)
    finally:
        pending_path.unlink(missing_ok=True)
    return {'pages': len(added), 'template_sha256': metadata['sha256'], 'roles': roles}
