"""Disposable installed Qt + real bilingual narration check, without playback."""

import json
import tempfile
import time
import traceback
from pathlib import Path


def run(output):
    from pptx import Presentation
    from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QTimer
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtQml import QQmlApplicationEngine
    from PySide6.QtQuickControls2 import QQuickStyle

    from . import speech
    from .chatgpt_handoff import external_preview, inspect_returned_deck
    from .library import Library
    from .paths import RESOURCE_ROOT
    from .ui import Bridge
    from .ui_watchdog import UiWatchdog

    app = QGuiApplication([])
    speech._ISOLATED_SYNTHESIS = True
    QQuickStyle.setStyle('Basic')
    output = Path(output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    result = dict(status='failed', scope='Disposable Qt conversion-result to real local Vietnamese-US pair narration; no API/playback')
    with tempfile.TemporaryDirectory(prefix='bcl-real-narration-', ignore_cleanup_errors=True) as temporary:
        root = Path(temporary)
        library = Library(root / 'data')
        bridge = engine = watcher = None
        ticks, warnings = [], []
        timer = QTimer()
        timer.setInterval(20)
        timer.timeout.connect(lambda: ticks.append(time.monotonic()))

        def pump():
            loop = QEventLoop()
            QTimer.singleShot(20, loop.quit)
            loop.exec()

        def until(condition, seconds=240):
            deadline = time.monotonic() + seconds
            while not condition():
                assert time.monotonic() < deadline, 'Real narration timed out'
                pump()

        try:
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text = 'VI: Năng lượng được bảo toàn.\nEN: Energy is conserved.'
            slide.shapes.add_textbox(914400, 914400, 6500000, 1200000).text = text
            slide.notes_slide.notes_text_frame.text = text
            source = root / 'fixture.pptx'
            deck.save(source)
            config = dict(title='Installed narration check', subject='Vật lý', education_level='THPT', grade='11',
                          conversion_format='parallel_columns', style='source', preset='standard')
            lesson = library.create_external_lesson(config, library.store_source(source), inspect_returned_deck(source))
            bridge = Bridge(library)
            engine = QQmlApplicationEngine()
            engine.warnings.connect(lambda values: warnings.extend(str(v) for v in values))
            engine.rootContext().setContextProperty('bridge', bridge)
            engine.load(str(RESOURCE_ROOT / 'qml/Main.qml'))
            assert engine.rootObjects(), warnings
            window = engine.rootObjects()[0]
            watcher = UiWatchdog(app, library.directory)
            bridge._browser_state = dict(phase='completed', lessonId=lesson['id'])
            bridge._show_lesson(lesson, external_preview(lesson, library.directory),
                                destination='chatgpt', auto_repair=False, load_preview=False)
            assert window.property('page') == 'chatgpt' and not bridge.narrationBusy
            bridge.openConvertedLesson()
            assert window.property('page') == 'voice'
            timer.start()
            started = time.monotonic()
            bridge.createLessonNarrationWithVoice('both', 'pair:north-female', 0)
            until(lambda: not bridge.busy and not bridge.narrationBusy)
            until(lambda: bridge.lessonNarration.get('status') != 'checking')
            status = bridge.lessonNarration
            assert status.get('ready') and status['complete'] == 2, bridge.narrationPreparation
            gaps = [b-a for a, b in zip(ticks, ticks[1:])]
            result.update(complete=status['complete'], language_mode=status['mode'],
                          elapsed_seconds=round(time.monotonic() - started, 2), heartbeat_ticks=len(ticks),
                          max_ui_gap_seconds=round(max(gaps), 4), qml_warnings=warnings,
                          watchdog_enabled=True, source_preserved=True)
            if watcher.path.is_file():
                output.with_suffix('.hang.log').write_bytes(watcher.path.read_bytes())
            assert len(ticks) >= 10 and max(gaps, default=0) < 5, result
            assert not warnings, warnings
            assert (library.directory / 'sources' / lesson['source']['file']).read_bytes() == source.read_bytes()
            result['status'] = 'passed'
        except Exception:
            result['error'] = traceback.format_exc()
        finally:
            timer.stop()
            if watcher is not None:
                watcher.stop()
                watcher.thread.join(2)
            if bridge is not None:
                # Wait for any requested operation before destroying its QObject.
                if bridge.worker is not None:
                    until(lambda: not bridge.busy)
                for service in (bridge.narration_preparation, bridge.english_reader, bridge.voice_controls,
                                bridge.classroom, bridge.libraryThumbnails, bridge.slidePreview, bridge._catalog,
                                bridge.browserAI, bridge.narrationPreferences):
                    service.shutdown()
                from .speech_process import shutdown

                shutdown()
            if engine is not None:
                for window in engine.rootObjects():
                    window.close()
                engine.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                app.processEvents()
            library.close()
            output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if result['status'] == 'passed' else 1
