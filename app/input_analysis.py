"""Conservative, local input assessment; language/pair guesses remain drafts."""

import re
import unicodedata
from collections import Counter

VI_WORDS = set("va la cua trong mot cac cho voi khi thi duoc hoc bai giang nghia vi du ham so gia tri lop em chung ta tinh dien tich chu ky phan tu te bao lich su dia ly cach doc viet hay nay neu bang tu nhung nguoi khong".split())
EN_WORDS = set("the a an and is are was were of to in for with when if then this that these those we you your each from by value values function equation calculate find explain example lesson learning class student students cell energy history geography area time what why how does can will has have not into about means change increases decreases english chapter sequences arithmetic geometric progressions grade mathematics standard curriculum objective".split())
LABELS = {"vi": "Tiếng Việt", "en": "Tiếng Anh", "bilingual": "Có cặp Việt–Anh",
          "mixed": "Hỗn hợp Việt/Anh", "unknown": "Chưa rõ ngôn ngữ", "neutral": "Số/ký hiệu"}


def language_of(text):
    words = re.findall(r"[^\W\d_]+", text.casefold(), re.UNICODE)
    if not words:
        return "neutral"
    accents = sum(any(ord(char) > 127 and char in "ăâđêôơưáàảãạắằẳẵặấầẩẫậéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ" for char in word) for word in words)
    flat = ["".join(char for char in unicodedata.normalize("NFD", word.replace("đ", "d"))
                    if not unicodedata.combining(char)) for word in words]
    vi, en = sum(w in VI_WORDS for w in flat), sum(w in EN_WORDS for w in words)
    if accents and accents / len(words) >= .12:
        return "vi"
    if en >= 2 and en > vi:
        return "en"
    if vi >= 2 and vi > en:
        return "vi"
    return "unknown"


def split_existing_pair(text):
    """Only split clear labelled/alternating lines; never invent a translation."""
    if "|" in text:
        # A single bilingual heading is common in teachers' decks. Multi-row
        # or multi-column grids remain intact for the native table exporter.
        if text.count("|") == 1 and len(text.splitlines()) == 1:
            left, right = (part.strip() for part in text.split("|"))
            languages = language_of(left), language_of(right)
            if set(languages) == {"vi", "en"}:
                return dict(zip(languages, (left, right), strict=True))
        return None
    parts = {"vi": [], "en": []}
    for line in text.splitlines():
        if not line.strip():
            continue
        match = re.match(r"^\s*(VI|VN|Tiếng Việt|EN|English|Tiếng Anh)\s*[:：]\s*(.+)$", line, re.I)
        if match:
            language = "vi" if match[1].casefold() in {"vi", "vn", "tiếng việt"} else "en"
            line = match[2]
        else:
            language = language_of(line)
        if language not in parts:
            return None  # Ambiguous formulas/other languages need teacher routing.
        parts[language].append(line.strip())
    if all(parts.values()):
        return {key: "\n".join(value) for key, value in parts.items()}
    return None


def assess_blocks(blocks, kind="text", fallback_language="vi", objects=None, table_locators=()):
    units = []
    for locator, text in blocks:
        pair = None if locator in table_locators else split_existing_pair(text)
        language = "bilingual" if pair else language_of(text)
        source_language = language if language in {"vi", "en"} else fallback_language
        draft = pair or {"vi": text if source_language == "vi" else "",
                         "en": text if source_language == "en" else ""}
        if language == "neutral":
            draft = {"vi": text, "en": text}
        units.append({"locator": locator, "language": language, "source_language": source_language,
                      "vi": draft["vi"], "en": draft["en"], "existing_pair": bool(pair), "paired_locator": ""})
    # A unique VI/EN block on one slide is a pairing suggestion only. Both sides
    # stay unapproved and the UI explicitly asks the teacher to verify meaning.
    slides = {}
    for unit in units:
        slide = re.match(r"^Slide (\d+)(?:\D|$)", unit["locator"])
        if slide:
            slides.setdefault(slide[1], []).append(unit)
    for slide_units in slides.values():
        vi = [u for u in slide_units if u["language"] == "vi"]
        en = [u for u in slide_units if u["language"] == "en"]
        if len(vi) == len(en) == 1:
            vi, en = vi[0], en[0]
            # A conflicting number is useful evidence that these are unrelated.
            if re.findall(r"\d+(?:[.,]\d+)?", vi["vi"]) == re.findall(r"\d+(?:[.,]\d+)?", en["en"]):
                vi.update(en=en["en"], paired_locator=en["locator"])
                en.update(vi=vi["vi"], paired_locator=vi["locator"])
    counts = Counter(unit["language"] for unit in units)
    identified = {key for key in ("vi", "en") if counts[key]}
    paired = sum(bool(u["existing_pair"] or u["paired_locator"]) for u in units)
    language = ("bilingual" if paired and paired == counts["vi"] + counts["en"] + counts["bilingual"]
                else "mixed" if len(identified) > 1 or counts["bilingual"]
                else next(iter(identified), "unknown"))
    primary = "en" if counts["en"] > counts["vi"] else "vi" if counts["vi"] > counts["en"] else fallback_language
    warnings = []
    if paired:
        warnings.append("Cặp có sẵn chỉ là bản nháp nhận diện; kiểm tra đúng nghĩa trước khi duyệt. Không tự dịch lại phần đã có.")
    if counts["unknown"]:
        warnings.append(f"{counts['unknown']} vùng chưa rõ ngôn ngữ; dùng ngôn ngữ dự phòng đã chọn và kiểm tra lại.")
    if any("OCR" in u["locator"] for u in units):
        warnings.append("Nội dung OCR cần kiểm tra dấu tiếng Việt, số, công thức và thứ tự đọc.")
    return {"version": 1, "kind": kind, "language": language, "label": LABELS[language],
            "primary_language": primary, "counts": dict(counts), "units": units, "objects": objects or {},
            "recommended_style": "source" if kind == "pptx" else "template",
            "recommended_mode": "preserve" if kind == "pptx" and language in {"bilingual", "mixed"} else "level",
            "warnings": warnings, "requires_review": True}


def analyze_document(path, language="vi", cancelled=None):
    from pathlib import Path

    from .importers import extraction_warnings, parse_document

    path = Path(path)
    blocks = parse_document(path, language, cancelled)
    objects = {}
    table_locators = set()
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation

        deck = Presentation(path)
        from .source_deck import text_blocks
        table_locators = {unit["locator"] for unit in text_blocks(deck) if unit["shape"].has_table}
        objects = {"slides": len(deck.slides), "images": 0, "charts": 0, "tables": 0, "image_only_slides": 0}
        for slide in deck.slides:
            def leaves(shapes):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        yield from leaves(shape.shapes)
                    else:
                        yield shape
            shapes = list(leaves(slide.shapes))
            objects["images"] += sum(int(shape.shape_type) == 13 for shape in shapes)
            objects["charts"] += sum(shape.has_chart for shape in shapes)
            objects["tables"] += sum(shape.has_table for shape in shapes)
            objects["image_only_slides"] += not any((shape.has_text_frame and shape.text.strip()) or shape.has_table for shape in shapes)
    profile = assess_blocks(blocks, path.suffix.lower().lstrip("."), language, objects, table_locators)
    if path.suffix.lower() == ".pptx" and any("OCR" in locator for locator, _ in blocks):
        profile.update(recommended_style="template", recommended_mode="level")
        profile["warnings"].append("PowerPoint có slide scan: nên tạo bài theo mẫu từ OCR; giữ nguồn để đối chiếu. Không thay trực tiếp chữ nằm trong ảnh.")
    profile["warnings"].extend(extraction_warnings(path))
    return {"blocks": blocks, "profile": profile}
