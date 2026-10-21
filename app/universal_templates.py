"""Manifest-driven, editable local layouts from the teacher's Universal pack.

The pack is design data, never executable code. Source slide boundaries and
native media travel together. Old jobs and teacher-uploaded templates keep
their existing renderer. Layout selection never invents knowledge or answers.
"""

import copy
import hashlib
import json
import re
import unicodedata
from collections import OrderedDict
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .native_layout import text_height
from .paths import RESOURCE_ROOT
from .pptx_template import _write
from .source_deck import NS

ROOT = RESOURCE_ROOT / 'assets/templates/universal2026'
THEMES = {'standard': 'CLASSROOM', 'visual': 'VISUAL', 'practice': 'PRACTICE'}
TOKEN = re.compile(r'\{\{([A-Z_0-9]+)\}\}')
SLIDE = re.compile(r'^Slide (\d+)(?:\D|$)')


def manifest():
    return json.loads((ROOT / 'manifest.json').read_text(encoding='utf-8'))


def asset(preset):
    item = manifest()['themes'][THEMES[preset]]
    path = (ROOT / item['file']).resolve()
    if not path.is_relative_to(ROOT.resolve()) or not path.is_file():
        raise ValueError('Thiếu bộ mẫu Universal 2026; cập nhật lại ứng dụng.')
    return path


def formula(text):
    """Share language-neutral equations once; do not translate section numbers."""
    if re.match(r'^\s*[A-D][.)]\s*', text):
        return False  # A formula-valued answer still belongs with the question.
    if '$' in text:
        prose = re.sub(r'\$[^$]+\$', '', text)
        prose = re.sub(r'\b(?:cm|mm|km|m|s|rad|Hz|kg|N)\b', '', prose)
        if re.search(r'[^\W\d_]', prose, re.UNICODE):
            return False
    return bool(re.search(r'[=∫∑√^]|[α-ωΑ-Ω]', text) and not
                re.search(r'[ăâđêôơưáàảãạắằẳẵặấầẩẫậéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ]', text, re.I))


def needs_translation(text):
    return bool(text.strip()) and not (formula(text) or re.fullmatch(r'[\d\sIVXLCDMivxlcdm.)\-–]+', text.strip()))


def source_blocks(deck):
    from .universal_content import source_blocks as visible_blocks

    return visible_blocks(deck)


def diagnose(report):
    """Offline packaging check; writes only disposable fixtures beside the report."""
    import time
    from threading import Event

    report=Path(report)
    report.parent.mkdir(parents=True,exist_ok=True)
    directory=report.parent/'template-fixtures'
    directory.mkdir(parents=True,exist_ok=True)
    rows=[]
    try:
        pairs=[('Slide 1','Cơ năng','Mechanical energy'),
               ('Slide 2 · ý 1','Bài học','Lesson'),
               ('Slide 2 · ý 2','Năng lượng được bảo toàn.','Energy is conserved.')]
        for preset in THEMES:
            for mode in ('parallel_columns','english_only'):
                start=time.monotonic()
                target=directory/(preset+'-'+mode+'.pptx')
                result=build(target,asset(preset),{'title':'Bài mẫu','english_title':'Sample lesson',
                             'preset':preset,'conversion_format':mode},pairs,Event())
                deck=Presentation(target)
                text='\n'.join(s.text for page in deck.slides for s in page.shapes if s.has_text_frame)
                assert len(deck.slides)==2 and 'Energy is conserved.' in text and '{{' not in text
                assert ('Năng lượng được bảo toàn.' in text)==(mode!='english_only')
                rows.append({'preset':preset,'mode':mode,'seconds':time.monotonic()-start,**result})
        result={'status':'passed','checks':rows,'rendering':'not run by this packaging check'}
    except Exception as exc:
        import traceback
        result={'status':'failed','error':type(exc).__name__,'message':str(exc),'traceback':traceback.format_exc()}
    report.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if result['status']=='passed' else 1


def _plain(text):
    return ''.join(c for c in unicodedata.normalize('NFD', text.lower().replace('đ', 'd'))
                   if unicodedata.category(c) != 'Mn')


def choose_layout(title, content, *, pictures=False, equations=False, table=False, mode='parallel_columns'):
    heading = _plain(title)
    if not content and not (pictures or equations or table):
        if re.match(r'^(?:chuong|phan|chapter|section)\b', heading):
            return 'SECTION'
        return 'TITLE_ONLY'
    if table:
        return 'TABLE'
    if pictures:
        return 'IMAGE_TEXT'
    if equations:
        return 'FORMULA'
    if re.search(r'\b(?:dap an|loi giai|answer|solution)\b', heading):
        # Do not infer which answer is correct or add a reveal page.
        return 'QUIZ_REVEAL'
    if re.search(r'\b(bai tap|vi du|exercise|example)\b', heading) and any(
            re.match(r'^(loi giai|giai|solution)\b', _plain(p[1])) for p in content):
        return 'EXERCISE'
    options = [re.match(r'^\s*([A-D])[.)]\s*', p[1]) for p in content]
    if {m[1] for m in options if m} == set('ABCD') and len(content) == 5:
        return 'QUIZ'
    for pattern, layout in (
        (r'\b(muc tieu|objectives?)\b', 'OBJECTIVES'),
        (r'\b(noi dung bai|muc luc|roadmap|outline)\b', 'ROADMAP'),
        (r'\b(quy trinh|cac buoc|process|steps)\b', 'PROCESS'),
        (r'\b(so sanh|comparison)\b', 'COMPARISON'),
        (r'\b(thao luan|hoat dong nhom|group work)\b', 'GROUP_WORK'),
        (r'\b(bai tap ve nha|homework)\b', 'HOMEWORK'),
        (r'\b(tom tat|cung co|ket luan|summary|conclusion)\b', 'SUMMARY'),
        (r'\b(mo dau|gioi thieu|introduction)\b', 'INTRO'),
        (r'\b(cam on|thank you)\b', 'THANK_YOU'),
    ):
        if re.search(pattern, heading):
            return layout
    if mode == 'english_only':
        return 'ENGLISH'
    if mode == 'sentence_pairs':
        return 'STACKED'
    return 'TWO_COLUMNS' if len(content) > 1 else 'THEORY'


def _clone(deck, original):
    slide = deck.slides.add_slide(original.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    mapping = {}
    for rel in original.part.rels.values():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide', '/slide')):
            continue
        mapping[rel.rId] = (slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                            if rel.is_external else slide.part.rels.get_or_add(rel.reltype, rel.target_part))
    for shape in original.shapes:
        node = copy.deepcopy(shape._element)
        for element in node.iter():
            if element.tag in {'{'+NS['a']+'}latin','{'+NS['a']+'}ea','{'+NS['a']+'}cs'} and element.get('typeface','').startswith('Aptos'):
                # Office 2019 can render the supplied Aptos italic glyphs incorrectly.
                element.set('typeface','Arial')
                element.attrib.pop('charset',None)
                element.attrib.pop('pitchFamily',None)
            for key, value in list(element.attrib.items()):
                if key.startswith('{' + NS['r'] + '}') and value in mapping:
                    element.set(key, mapping[value])
        slide.shapes._spTree.insert_element_before(node, 'p:extLst')
    background = original._element.find('p:cSld/p:bg', NS)
    if background is not None:
        slide._element.find('p:cSld', NS).insert(0, copy.deepcopy(background))
    return slide


def _fields(slide):
    return [(shape, TOKEN.search(shape.text)[1]) for shape in slide.shapes
            if shape.has_text_frame and TOKEN.search(shape.text)]


def _header(slide):
    """Stack title translations; the supplied pack puts most headers side by side."""
    fields = _fields(slide)
    headers = [(s, k) for s, k in fields if s.top < Inches(1.1)]
    if not headers:
        return
    # Reserve a real title band, then scale the existing native composition.
    for shape in slide.shapes:
        if Inches(1.04) <= shape.top < Inches(7):
            shape.top = int(Inches(1.72) + (shape.top - Inches(1.04)) * .84)
            shape.height = max(1, int(shape.height * .84))
    for shape, key in headers:
        english = key.endswith('_EN')
        shape.left, shape.width = Inches(.72), Inches(11.85)
        shape.top = Inches(1.1 if english else .44)
        shape.height = Inches(.51 if english else .61)
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment=PP_ALIGN.LEFT
    for shape, key in fields:
        if key.startswith('OPTION_'):
            shape.height = max(shape.height, Inches(.35))
        elif key.startswith('QUESTION_'):
            shape.height = max(shape.height, Inches(.8))
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text and not TOKEN.search(shape.text):
            sizes=[r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
            shape.height=max(shape.height,Pt(max(sizes,default=10)*1.4))


def _fit(shape, text, minimum=18):
    width = (shape.width - shape.text_frame.margin_left - shape.text_frame.margin_right) / 12700
    height = (shape.height - shape.text_frame.margin_top - shape.text_frame.margin_bottom) / 12700
    return width > 0 and text_height(text, width, minimum) + minimum * .18 <= height


def _fill(shape, value, english=False, title=False, *, maths=None):
    if not value:
        shape.text_frame.clear()
        return
    # English slots in the pack already have smaller type: do not shrink twice.
    _write(shape, [(value, False)], title=title)
    shape.name='BCL_FILLED '+shape.name
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run._r.get_or_add_rPr().set('lang', 'en-US' if english else 'vi-VN')
            run.font.italic = english
            if title:
                run.font.size = Pt(min(run.font.size.pt, 22 if english else 28))
    if maths:
        from .universal_content import insert_inline_math

        insert_inline_math(shape, maths, english)


def _slots(slide, mode):
    """Matched VI/EN native frames; never put two languages in separate slides."""
    groups = OrderedDict()
    for shape, key in _fields(slide):
        if shape.top < Inches(1.7) or key in {'IMAGE_SOURCE', 'FORMULA_SOURCE', 'TABLE_SOURCE',
                    'SECTION_NUMBER', 'SUBJECT', 'GRADE', 'TEACHER', 'TIME', 'FOOTER', 'CORRECT_ANSWER'}:
            continue
        if key.startswith(('TITLE_', 'LESSON_TITLE_', 'SECTION_TITLE_', 'THANKS_')):
            continue
        base = re.sub(r'_(VI|EN)$', '', key)
        groups.setdefault(base, {})['en' if key.endswith('_EN') or key == 'ENGLISH_BODY' else 'vi'] = shape
    if mode == 'english_only':
        for pair in groups.values():
            vi, en = pair.get('vi'), pair.get('en')
            if vi is not None and en is not None:
                x, y = min(vi.left, en.left), min(vi.top, en.top)
                w, h = max(vi.left+vi.width, en.left+en.width)-x, max(vi.top+vi.height, en.top+en.height)-y
                en.left, en.top, en.width, en.height = x, y, w, h
    return groups


def _pack(slots, pending, mode):
    """Fit complete pairs, using pagination rather than truncation."""
    assigned, taken = [], 0
    for key, pair in slots.items():
        values = []
        while taken < len(pending):
            candidate = values + [pending[taken]]
            vi, en = '\n'.join(p[1] for p in candidate), '\n'.join(p[2] for p in candidate)
            fits = (mode == 'english_only' or pair.get('vi') is not None and _fit(pair['vi'], vi))
            fits = fits and pair.get('en') is not None and _fit(pair['en'], en)
            if not fits:
                break
            values = candidate
            taken += 1
        if values:
            assigned.append((key, values))
    return assigned, taken


def _dense_columns(slide, slots):
    """Join the two short rows when a longer paragraph needs the full column."""
    first=slots.get('BODY_1')
    if not first:
        return slots
    for shape in first.values():
        shape.height=int(Inches(6.55)-shape.top)
    for shape in list(slide.shapes):
        if not shape.has_text_frame and shape.height < Pt(2) and Inches(3)<shape.top<Inches(6.7) and shape.width<Inches(6):
            slide.shapes._spTree.remove(shape._element)
    return OrderedDict([('BODY_1',first)])


def _dense_theory(slide, slots):
    pair=slots.get('TEXT')
    if not pair:
        return slots
    for language, shape in pair.items():
        shape.top=Inches(2.1 if language=='vi' else 4.4)
        shape.height=Inches(2.1)
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text in {'TIẾNG VIỆT','ENGLISH'}:
            shape.text_frame.clear()
        elif not shape.has_text_frame and shape.height<Pt(2) and Inches(2)<shape.top<Inches(6.6):
            slide.shapes._spTree.remove(shape._element)
    return slots


def _fallback(mode, pending):
    if mode=='english_only':
        return 'ENGLISH'
    return 'THEORY' if any(max(len(p[1]),len(p[2]))>250 for p in pending) else 'TWO_COLUMNS'


def _prune_unused(slide, slots, used):
    """Remove an unused native card and its number, not a panel with live content."""
    active=[s for key,pair in slots.items() if key in used for s in pair.values()]
    active += [s for s in slide.shapes if s.name.startswith('BCL_SOURCE ')]
    def contains(panel, shape):
        return (panel.left<=shape.left and panel.top<=shape.top and
                panel.left+panel.width>=shape.left+shape.width and
                panel.top+panel.height>=shape.top+shape.height)
    for key,pair in slots.items():
        if key in used:
            continue
        boxes=list(pair.values())
        for subset in [boxes]+[[b] for b in boxes]:
            panels=[s for s in slide.shapes if s.has_text_frame and not s.text.strip() and
                    s.width>Inches(1) and s.height>Inches(.6) and all(contains(s,b) for b in subset)
                    and not any(contains(s,b) for b in active)]
            if not panels:
                continue
            panel=min(panels,key=lambda s:s.width*s.height)
            for shape in list(slide.shapes):
                if shape is panel or contains(panel,shape):
                    slide.shapes._spTree.remove(shape._element)


def _media(source_slide, deck_size=None, inline_expressions=(), *, preserve_picture_layout=False):
    from .slide_speech import objects
    # Keep native equations and pictures, including group-relative pictures.
    result, seen = [], set()
    for obj in objects(source_slide):
        shape = obj['shape']
        if obj['hidden'] or shape._element in seen:
            continue
        if int(shape.shape_type) == 13 and not preserve_picture_layout:
            from .universal_content import decorative_picture

            if deck_size and decorative_picture(obj, deck_size):
                seen.add(shape._element)
                continue
        math = shape._element.xpath('.//*[local-name()="oMath" or local-name()="oMathPara"]')
        if math and shape.has_text_frame:
            # Prose surrounding native math is already translated as body
            # text. Carry only equations; copying the whole source textbox
            # duplicates untranslated paragraphs and crushes them into the
            # formula slot.
            from pptx.shapes.autoshape import Shape
            node = copy.deepcopy(shape._element)
            paragraphs = node.xpath('./p:txBody/a:p')
            for paragraph in paragraphs:
                from .universal_content import equation_only, math_text, visible_run

                standalone = equation_only(paragraph)

                expressions = [math_text(c) for c in paragraph if c.xpath('.//*[local-name()="oMath"]') or etree.QName(c).localname == 'oMath']
                if (not expressions or not standalone
                        and all(value in inline_expressions for value in expressions)):
                    paragraph.getparent().remove(paragraph)
                    continue
                for child in list(paragraph):
                    if standalone and etree.QName(child).localname in {'r','fld','br'} and visible_run(child):
                        continue  # Keep the native equation's unit/punctuation.
                    if child.tag != '{'+NS['a']+'}pPr' and not child.xpath('.//*[local-name()="oMath" or local-name()="oMathPara"]') and child not in math and child.tag.rsplit('}',1)[-1] not in {'oMath','oMathPara'}:
                        paragraph.remove(child)
            # Each equation is an independent native object. Source paragraph
            # spacing, bullets and decorated answer boxes do not belong to the
            # formula slot, and several equations must not share a tiny frame.
            for paragraph in node.xpath('./p:txBody/a:p'):
                equation = copy.deepcopy(node)
                body = equation.find('p:txBody', NS)
                for old in list(body.findall('a:p', NS)):
                    body.remove(old)
                clean = copy.deepcopy(paragraph)
                old = clean.find('a:pPr', NS)
                if old is not None:
                    clean.remove(old)
                props = etree.Element('{'+NS['a']+'}pPr', algn='ctr')
                etree.SubElement(props, '{'+NS['a']+'}buNone')
                etree.SubElement(props, '{'+NS['a']+'}defRPr', sz='2400')
                clean.insert(0, props)
                body.append(clean)
                body_props = body.find('a:bodyPr', NS)
                body_props.attrib.clear()
                body_props.attrib.update({'wrap':'none', 'lIns':'0', 'rIns':'0', 'tIns':'0', 'bIns':'0', 'anchor':'ctr'})
                for old in list(body_props):
                    body_props.remove(old)
                etree.SubElement(body_props, '{'+NS['a']+'}noAutofit')
                for element in clean.iter():
                    if element.get('sz'):
                        element.set('sz', '2400')
                cloned = Shape(equation, shape._parent)
                cloned.fill.background()
                cloned.line.fill.background()
                terms = ''.join(clean.xpath('.//*[local-name()="t"]/text()'))
                depth = max((sum(etree.QName(a).localname in {'f','rad','sSup','sSub','eqArr','m','nary'}
                                 for a in e.iterancestors()) for e in clean.iter()), default=0)
                cloned.width = max(shape.width, Pt(len(terms)*22+40))
                cloned.height = Pt(42+38*depth)
                result.append({**obj, 'shape': cloned, 'w': cloned.width, 'h': cloned.height})
            seen.add(shape._element)
        elif int(shape.shape_type) in {13, 16} or shape.has_chart or math:
            accent = (not preserve_picture_layout and deck_size is not None and int(shape.shape_type) == 13
                      and obj['w'] < deck_size[0]*.12 and obj['h'] < deck_size[1]*.2)
            result.append({**obj, 'accent': accent})
            seen.add(shape._element)
    return result


def _insert_media(slide, objects, region):
    if not objects:
        return
    x, y, w, h = region
    left, top = min(o['x'] for o in objects), min(o['y'] for o in objects)
    width = max(o['x']+o['w'] for o in objects)-left
    height = max(o['y']+o['h'] for o in objects)-top
    scale = min(w/max(1, width), h/max(1, height))
    for obj in objects:
        shape = obj['shape']
        node = copy.deepcopy(shape._element)
        # Source slide-jump links would import the original slide/master graph
        # into this new template, creating duplicate OPC part names. The new
        # layout has its own pagination, so those navigation targets no longer
        # identify a valid destination. Keep ordinary external hyperlinks.
        for link in list(node.iter()):
            rid = link.get('{' + NS['r'] + '}id')
            if rid in shape.part.rels:
                rel = shape.part.rels[rid]
                if not rel.is_external and rel.reltype.endswith('/slide'):
                    parent = link.getparent()
                    if parent is not None:
                        parent.remove(link)
        next_id = slide.shapes._next_shape_id
        for element in node.iter():
            if element.get('sz') and element.tag.rsplit('}',1)[-1] in {'rPr','defRPr','endParaRPr'}:
                element.set('sz', str(max(100, round(int(element.get('sz')) * scale))))
            if element.tag == '{'+NS['a']+'}spcPts':
                element.set('val', str(round(int(element.get('val','0')) * scale)))
            if element.tag == '{'+NS['a']+'}bodyPr':
                for inset in ('lIns','rIns','tIns','bIns'):
                    if element.get(inset):
                        element.set(inset, str(round(int(element.get(inset)) * scale)))
            if element.tag == '{' + NS['p'] + '}cNvPr':
                element.set('id', str(next_id))
                next_id += 1
                element.set('name', 'BCL_SOURCE ' + element.get('name', ''))
            for key, rid in list(element.attrib.items()):
                if key.startswith('{' + NS['r'] + '}') and rid in shape.part.rels:
                    rel = shape.part.rels[rid]
                    new = (slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external
                           else slide.part.rels.get_or_add(rel.reltype, rel.target_part))
                    element.set(key, new)
        slide.shapes._spTree.insert_element_before(node, 'p:extLst')
        placed = list(slide.shapes)[-1]
        placed.left = int(x+(w-width*scale)/2+(obj['x']-left)*scale)
        placed.top = int(y+(h-height*scale)/2+(obj['y']-top)*scale)
        placed.width, placed.height = max(1, int(obj['w']*scale)), max(1, int(obj['h']*scale))


def _table_pages(shape):
    """Size actual cell text and split rows before copying the native table."""
    width,height=Inches(11),Inches(3.9)
    table=shape.table
    original_width=sum(c.width for c in table.columns)
    for column in table.columns:
        column.width=int(column.width*width/original_width)
    heights=[]
    for row in table.rows:
        needed=Pt(32)
        for col,cell in zip(table.columns,row.cells):
            if cell.is_spanned:
                continue
            cell_width=max(1,(col.width-cell.margin_left-cell.margin_right)/12700)
            needed=max(needed,Pt(text_height(cell.text,cell_width,18)+12)+cell.margin_top+cell.margin_bottom)
        row.height=int(needed)
        heights.append(row.height)
    pending=list(range(len(heights)))
    while pending:
        indices=[]
        used=0
        if not (pending[0]==0) and table.first_row:
            indices=[0]
            used=heights[0]
        while pending and used+heights[pending[0]]<=height:
            index=pending.pop(0)
            indices.append(index)
            used+=heights[index]
        if not indices or indices==[0] and pending and pending[0]!=0:
            raise ValueError('Một ô bảng quá dài cho trang mẫu. Tách nội dung ô trước khi chuyển đổi; tool không bỏ chữ.')
        from pptx.shapes.graphfrm import GraphicFrame
        output=GraphicFrame(copy.deepcopy(shape._element),shape._parent)
        rows=list(output.table._tbl.tr_lst)
        for i,row in enumerate(rows):
            if i not in indices:
                output.table._tbl.remove(row)
        output.width=width
        output.height=sum(heights[i] for i in indices)
        yield {'shape':output,'x':0,'y':0,'w':output.width,'h':output.height,'hidden':False}


LABELS = {
    'VIỆT NAM / VI': '', 'TIẾNG VIỆT': '', 'VIỆT NAM': '',
    'GIẢI THÍCH TIẾNG VIỆT': '', 'GIỚI THIỆU / INTRO': 'INTRODUCTION',
    'NEW SECTION / CHỦ ĐỀ': 'NEW SECTION', 'GIẢI THÍCH / EXPLANATION': 'EXPLANATION',
    'CÔNG THỨC GỐC / ORIGINAL FORMULA': 'ORIGINAL FORMULA',
    'BẢNG DỮ LIỆU / DATA TABLE': 'DATA TABLE', 'THẢO LUẬN / GROUP ACTIVITY': 'GROUP ACTIVITY',
    'NHIỆM VỤ': '', 'BÀI TẬP / EXERCISE': 'EXERCISE', 'LỜI GIẢI / SOLUTION': 'SOLUTION',
    'ĐÁP ÁN / CORRECT ANSWER': 'CORRECT ANSWER', 'NHIỆM VỤ VỀ NHÀ / TAKE-HOME TASK': 'TAKE-HOME TASK',
}


def _clean(slide, mode, number):
    def label(shape, text):
        # These are small decorative captions, not 18pt teaching content.
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
        if runs:
            runs[0].text = text
            for run in runs[1:]:
                run.text = ''
        else:
            shape.text = text
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.name.startswith('BCL_FILLED '):
            continue
        value = shape.text
        if TOKEN.search(value) or value in {'+', '⏱', 'ANSWER IS NOT SHOWN ON THIS SLIDE'} or value.startswith('Ví dụ cấu trúc:'):
            shape.text_frame.clear()
        elif re.fullmatch(r'\d{2} / 23', value):
            label(shape, f'{number:02}')
        elif mode == 'english_only':
            if value in LABELS:
                label(shape, LABELS[value])
            elif value.startswith('CẶP '):
                label(shape, value.split(' / ')[-1])


def build(target, template, config, pairs, cancel, source=None):
    """Choose native layout pages, populate and save atomically with own notes."""
    catalog = config.get('template_layouts') or manifest()['layouts']
    deck = Presentation(template)
    originals = list(deck.slides)
    if len(originals) != 23 or any(not 1 <= int(v['slide_number']) <= 23 for v in catalog.values()):
        raise ValueError('Bộ mẫu Universal không khớp manifest 23 bố cục.')
    mode = config['conversion_format']
    groups = OrderedDict()
    for pair in pairs:
        match = SLIDE.match(pair[0])
        key = int(match[1]) if match else 0
        groups.setdefault(key, []).append(pair)
    source_deck = Presentation(source) if source and Path(source).suffix.lower() == '.pptx' else None
    from .universal_content import image_content_pages, inline_math, source_objects

    equations_xml = inline_math(source_deck) if source_deck else {}
    image_pages = set(image_content_pages(source_deck)) if source_deck else set()

    def fill(shape, value, english=False, title=False):
        return _fill(shape, value, english, title, maths=equations_xml)

    table_units = [{'slide': number, 'shape': obj['shape'], 'locator': f'Slide {number} · ý {index}'}
                   for number, page in enumerate(source_deck.slides, 1)
                   for index, obj in enumerate(source_objects(page), 1)
                   if obj['shape'].has_table and not obj['hidden']] if source_deck else []
    if source_deck:
        groups = OrderedDict((n, groups.get(n, [])) for n in range(1, len(source_deck.slides)+1))
    plan = []
    title_fallback = ('', config['title'], config['english_title'])

    def new(layout, title, body):
        if cancel.is_set():
            raise InterruptedError('Đã dừng dựng bài theo mẫu Universal.')
        slide = _clone(deck, originals[catalog[layout]['slide_number']-1])
        _header(slide)
        slots = _slots(slide, mode)
        from .universal_layout import prepare

        slots = prepare(slide, layout, slots, mode)
        assigned, taken = _pack(slots, body, mode)
        if body and not taken and layout=='TWO_COLUMNS':
            assigned,taken=_pack(_dense_columns(slide,slots),body,mode)
        if body and not taken and layout=='THEORY':
            assigned,taken=_pack(_dense_theory(slide,slots),body,mode)
        return slide, assigned, taken

    def discard(slide):
        sid = deck.slides._sldIdLst[-1]
        deck.part.drop_rel(sid.rId)
        deck.slides._sldIdLst.remove(sid)

    def heading_fits(layout, title):
        # A source's first paragraph is not necessarily a heading. Test the
        # actual template banner at readable size before taking it out of body
        # pagination; the full paragraph must remain visible somewhere.
        from pptx.slide import Slide
        original = originals[catalog[layout]['slide_number']-1]
        probe = Slide(copy.deepcopy(original._element), original.part)
        _header(probe)
        for shape, key in _fields(probe):
            if shape.top < Inches(1.7) or key.startswith(('TITLE_', 'LESSON_TITLE_', 'SECTION_TITLE_', 'THANKS_')):
                english = key.endswith('_EN')
                value = title[2] if english else ('' if mode == 'english_only' else title[1])
                if value and not _fit(shape, value):
                    return False
        return True

    def finish(slide, layout, title, assigned, native, source_number, *, answer='', extra=()):
        from .picture_narration import mark_template_text

        accents = [obj for obj in native if obj.get('accent')]
        native = [obj for obj in native if not obj.get('accent')]
        fields = _fields(slide)
        lookup = {base: values for base, values in assigned}
        # Fill header tokens by position; ROADMAP repeats SECTION_1 in its body.
        for shape, key in fields:
            if (shape.top < Inches(1.7) or key.startswith(('TITLE_', 'LESSON_TITLE_', 'SECTION_TITLE_', 'THANKS_'))):
                english = key.endswith('_EN')
                value = title[2] if english else ('' if mode == 'english_only' else title[1])
                fill(shape, value, english, True)
                if english:
                    mark_template_text(shape, [title])
            elif key in {'SUBJECT', 'GRADE', 'TEACHER'}:
                # Personal names/subject translation are not guessed for English-only.
                fill(shape, str(config.get({'SUBJECT':'subject','GRADE':'grade','TEACHER':'teacher'}[key], ''))
                      if mode != 'english_only' or key != 'SUBJECT' else '')
            elif key == 'SECTION_NUMBER':
                fill(shape, str(source_number))
            elif key == 'CORRECT_ANSWER' and answer:
                fill(shape, answer, True)
        slots = _slots(slide, mode)
        for base, values in lookup.items():
            pair = slots[base]
            if pair.get('vi') is not None:
                fill(pair['vi'], '\n'.join(p[1] for p in values) if mode != 'english_only' else '')
            if pair.get('en') is not None:
                fill(pair['en'], '\n'.join(p[2] for p in values), True)
                mark_template_text(pair['en'], values)
        _clean(slide, mode, len(plan)+1)
        from .universal_layout import finish as arrange_media

        media_region = arrange_media(slide, layout, native, assigned)
        if native:
            if layout == 'FORMULA':
                region = tuple(Inches(n) for n in (1.2, 2.55, 10.8, 1.2))
            elif layout == 'TABLE':
                region = tuple(Inches(n) for n in (1.1, 2.4, 11, 3.9))
            else:
                region = tuple(Inches(n) for n in (1.15, 2.45, 5.9, 2.5))
            if layout=='IMAGE_TEXT' and media_region is None:
                # The supplied image placeholder has a little circle/plus target;
                # it must not remain visible through transparent charts or PNGs.
                for shape in list(slide.shapes):
                    geometry=shape._element.find('p:spPr/a:prstGeom',NS)
                    if geometry is not None and geometry.get('prst')=='ellipse' and (
                            Inches(1)<shape.left<Inches(7) and Inches(2)<shape.top<Inches(5.2)):
                        slide.shapes._spTree.remove(shape._element)
                if not assigned:
                    for shape in slide.shapes:
                        if shape.has_text_frame and not shape.text.strip() and shape.left<Inches(2) and (
                                shape.width>Inches(5) and shape.height>Inches(2) and shape.top>Inches(1.7)):
                            shape.width=Inches(11.8 if shape.left<Inches(1) else 11.1)
                    region=tuple(Inches(n) for n in (1.15,2.4,11,3.1))
            _insert_media(slide, native, media_region or region)
        if accents:
            # Preserve small source illustrations at icon size, instead of
            # stretching a 20px ornament into the main teaching figure.
            _insert_media(slide, accents, tuple(Inches(n) for n in (11.8, 6.8, .6, .4)))
        if layout in {'OBJECTIVES','SUMMARY','ROADMAP','PROCESS','COMPARISON','FORMULA','IMAGE_TEXT'}:
            _prune_unused(slide, slots, set(lookup))
        spoken = [p for _, values in assigned for p in values] + list(extra)
        if answer:
            spoken.insert(0,('', 'Đáp án: '+answer, 'Answer: '+answer))
        narrated = [p for p in [title]+spoken if ' · chữ trong ảnh' not in p[0]]
        vi = '\n'.join(p[1] for p in narrated if p[1])
        en = '\n'.join(p[2] for p in narrated if p[2])
        slide.notes_slide.notes_text_frame.text = f'BILICLASS_NOTES:\nVI: {vi}\nEN: {en}'
        plan.append({'slide':len(plan)+1, 'source_slide':source_number, 'layout':layout})

    if config.get('template_bookends'):
        slide, assigned, _ = new('COVER', title_fallback, [])
        finish(slide, 'COVER', title_fallback, assigned, [], 0)
    for number, data in groups.items():
        data = list(data)
        # Cell translations stay in an editable table, outside body text packing.
        cells = [p for p in data if ' · ô ' in p[0]]
        data = [p for p in data if ' · ô ' not in p[0]]
        # Plain documents keep their first paragraph as content, not a new title.
        title = (data.pop(0) if number and data and ' · chữ trong ảnh' not in data[0][0]
                 and not formula(data[0][1]) else title_fallback)
        if data and re.fullmatch(r'[IVXLCDMivxlcdm\d.)\s]+',title[1]) and len(title[1])<12:
            heading=data.pop(0)
            title=(title[0],title[1]+' '+heading[1],title[1]+' '+heading[2])
        elif data and re.fullmatch(r'[IVX]{1,5}', data[0][1]) and len(title[1]) < 120:
            marker = data.pop(0)[1]
            title = (title[0], marker + ' · ' + title[1], marker + ' · ' + title[2])
        represented = set(re.findall(r'\$([^$]+)\$', '\n'.join(p[1] for p in [title]+data)))
        image_primary = number in image_pages and not any(' · chữ trong ảnh' not in pair[0] for pair in groups[number])
        native = _media(source_deck.slides[number-1], (source_deck.slide_width, source_deck.slide_height), represented,
                        preserve_picture_layout=image_primary) if source_deck and number else []
        # Native math is shared once. Plain equations also stay once in a formula slot.
        equations = [p for p in data if formula(p[1])]
        data = [p for p in data if not formula(p[1])]
        if equations:
            # Text equations retain Unicode and remain editable, without duplication.
            from pptx.slide import Slide
            carrier = Slide(copy.deepcopy(originals[12]._element), originals[12].part)
            for equation in equations:
                region = next(s for s,k in _fields(carrier) if k == 'FORMULA_SOURCE')
                fill(region, equation[1])
                native.append({'shape':region,'x':0,'y':0,'w':region.width,'h':region.height,'hidden':False})
                carrier = Slide(copy.deepcopy(originals[12]._element), originals[12].part)
        main_media = [o for o in native if not o.get('accent')]
        pictures = any(int(o['shape'].shape_type) in {13, 16} or o['shape'].has_chart for o in main_media)
        other_formulas=[]
        if pictures:
            # A graph and an expanded native equation have incompatible
            # source coordinates. Scaling their union into the image column
            # can reduce an equation below 10pt. Keep the figure prominent
            # and carry standalone equations once on readable formula pages.
            standalone = [o for o in main_media if o['shape'].has_text_frame]
            if standalone:
                from .universal_layout import equation_pages

                other_formulas = equation_pages(standalone)
                native = [o for o in native if o not in standalone]
                main_media = [o for o in native if not o.get('accent')]
        elif len(main_media)>1:
            from .universal_layout import equation_pages

            pages = equation_pages(main_media)
            native, other_formulas = pages[0] + [o for o in native if o.get('accent')], pages[1:]
        layout = choose_layout(title[1], data, pictures=pictures, equations=bool(main_media), mode=mode)
        if not heading_fits(layout, title):
            if title != title_fallback:
                data.insert(0, title)
            # A long question belongs in the body. Prefer its actual short
            # source caption over repeating the lesson name on every page.
            caption = next((p for p in data[1:] if len(p[1]) < 90 and re.match(
                r'^(cau hoi|hoat dong|phieu hoc tap|huong dan)\b', _plain(p[1]))), None)
            if caption:
                data.remove(caption)
                title = caption
            else:
                question = re.match(r'^\s*Câu\s*(\d+)\b', data[0][1], re.I) if data else None
                title = ('', 'Câu ' + question[1], 'Question ' + question[1]) if question else title_fallback
            if not heading_fits(layout, title):
                if title not in data:
                    data.insert(0, title)
                title = ('', 'Nội dung bài học', 'Lesson content')
            layout = choose_layout(title[1], data, pictures=pictures, equations=bool(main_media), mode=mode)
        if cells and source_deck:
            table_objects = []
            by_locator = {}
            for loc, vi, en in cells:
                by_locator.setdefault(loc, []).append((vi,en))
            # Copy table XML so the teacher's source is never edited in place.
            for unit in table_units:
                if unit['slide'] != number or not unit['shape'].has_table:
                    continue
                from pptx.shapes.graphfrm import GraphicFrame
                cloned = GraphicFrame(copy.deepcopy(unit['shape']._element), unit['shape']._parent)
                for r, row in enumerate(cloned.table.rows):
                    for c, cell in enumerate(row.cells):
                        translations = by_locator.get(unit['locator'] + f' · ô {r},{c}')
                        if translations:
                            vi, en = ('\n'.join(p[i] for p in translations) for i in (0,1))
                            cell.text = en if mode=='english_only' else vi+'\n'+en
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    run.font.size=Pt(18)
                                    run.font.name='Arial'
                table_objects.extend(_table_pages(cloned))
            if table_objects:
                for table_obj in table_objects:
                    slide, assigned, _ = new('TABLE',title,[])
                    # Notes follow the visible rows of this page, including a repeated header.
                    values=[c.text for row in table_obj['shape'].table.rows for c in row.cells if not c.is_spanned]
                    visible_cells=[p for p in cells if any(p[2] in text if mode=='english_only' else p[1] in text for text in values)]
                    finish(slide,'TABLE',title,assigned,[table_obj],number,extra=visible_cells)
                if not data and not native:
                    continue
        pending = list(data)
        first = True
        while pending or first:
            slide, assigned, taken = new(layout, title, pending)
            # Tiny specialised slots must not force missing text or unreadable fonts.
            answer = ''
            if layout == 'EXERCISE':
                split = next((i for i,p in enumerate(pending) if re.match(r'^(loi giai|giai|solution)\b',_plain(p[1]))),0)
                exact = [('QUESTION',pending[:split]),('SOLUTION',pending[split:])]
                fields = _slots(slide,mode)
                if split and all(_pack(OrderedDict([(k,fields[k])]),v,mode)[1]==len(v) for k,v in exact):
                    assigned,taken=exact,len(pending)
                else:
                    taken=0
            if layout == 'QUIZ_REVEAL':
                found = re.fullmatch(r'\s*(?:Đáp án|Answer)\s*[:=]\s*([A-D])\s*[.!]?\s*',pending[0][1],re.I) if pending else None
                if found:
                    answer=found[1]
                    assigned,taken=_pack(_slots(slide,mode),pending[1:],mode)
                    taken+=1
                else:
                    taken=0
            if pending and (not taken or layout in {'QUIZ','THANK_YOU'}):
                # Quiz fields need exact question/option mapping; fall back if the source
                # cannot fit all five fields in this single, honest question layout.
                if layout == 'QUIZ':
                    options = {re.match(r'^\s*([A-D])[.)]',p[1])[1]:p for p in pending if re.match(r'^\s*([A-D])[.)]',p[1])}
                    question = [p for p in pending if not re.match(r'^\s*([A-D])[.)]',p[1])]
                    if len(options)==4 and len(question)==1:
                        fields = _slots(slide,mode)
                        exact = [('QUESTION',question)]+[('OPTION_'+k,[options[k]]) for k in 'ABCD']
                        if all(_pack(OrderedDict([(k,fields[k])]),v,mode)[1]==1 for k,v in exact):
                            assigned,taken = exact,5
                        else:
                            taken=0
                else:
                    taken=0
                if not taken:
                    discard(slide)
                    if main_media and first:
                        slide,assigned,_=new(layout,title,[])
                        finish(slide,layout,title,[],native,number,extra=equations)
                        first=False
                        layout=_fallback(mode,pending)
                        continue
                    layout = _fallback(mode,pending)
                    slide,assigned,taken = new(layout,title,pending)
            if pending and not taken:
                discard(slide)
                # Retain every word. Translation pairs are already short API units.
                loc,vi,en=pending.pop(0)
                # Keep editable inline math atomic across continuation pages.
                a,b=(re.findall(r'\$[^$]+\$|\S+',value) for value in (vi,en))
                if max(len(a),len(b))<4:
                    raise ValueError('Nội dung không vừa khung mẫu; chọn khung chữ lớn hơn.')
                pending[:0]=[(loc,' '.join(a[len(a)*i//2:len(a)*(i+1)//2]),
                                  ' '.join(b[len(b)*i//2:len(b)*(i+1)//2])) for i in range(2)]
                continue
            finish(slide,layout,title,assigned,native if first else [],number,answer=answer,extra=equations if first else ())
            del pending[:taken]
            first=False
            if pending:
                layout=_fallback(mode,pending)
        for formula_group in other_formulas:
            slide,assigned,_=new('FORMULA',title,[])
            finish(slide,'FORMULA',title,assigned,formula_group,number)
    last_source = next(reversed(groups.values()), [])
    already_closed = any(re.match(r'^(?:cam on|thank(?:s| you))\b', _plain(p[1])) for p in last_source)
    if config.get('template_bookends') and not already_closed:
        ending=('', 'Cảm ơn!', 'Thank you!')
        slide,assigned,_=new('THANK_YOU',ending,[])
        finish(slide,'THANK_YOU',ending,assigned,[],0)
    if not plan:
        raise ValueError('Chưa có nội dung để tạo bài giảng theo mẫu.')
    for sid in list(deck.slides._sldIdLst)[:23]:
        deck.part.drop_rel(sid.rId)
        deck.slides._sldIdLst.remove(sid)
    target=Path(target)
    pending_path=target.with_suffix('.pending.pptx')
    try:
        deck.save(pending_path)
        if cancel.is_set():
            raise InterruptedError('Đã dừng dựng bài theo mẫu.')
        pending_path.replace(target)
    finally:
        pending_path.unlink(missing_ok=True)
    result={'version':3,'template_sha256':hashlib.sha256(Path(template).read_bytes()).hexdigest(),'pages':len(plan),'layout_plan':plan}
    target.with_suffix('.layout.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result
