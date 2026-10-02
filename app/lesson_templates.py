"""Source-owned slide templates shared by QML preview and PowerPoint export.

Templates arrange supplied text. They never add lesson facts, answers or approval.
Coordinates and font sizes use a 1280 x 720 canvas in CSS pixels.
"""

import json
import textwrap
from functools import lru_cache
from pathlib import Path

from .lesson_design import paired_pages
from .presentation_policy import presentation_content


@lru_cache(maxsize=1)
def catalog():
    data = json.loads((Path(__file__).parent / "assets/templates/lesson_templates.json").read_text(encoding="utf-8"))
    if data["schema_version"] != 1:
        raise ValueError("Phiên bản mẫu bài giảng không được hỗ trợ.")
    return data


def block_type(kind):
    aliases = {"heading": "title", "unknown": "explanation"}
    key = aliases.get(kind, kind) if isinstance(kind, str) else "explanation"
    return next((item for item in catalog()["blocks"] if item["id"] == key), catalog()["blocks"][5])


def _fits(text, width, height, size):
    # Conservative fallback shared across renderers. Keep pairs together; ask
    # teachers to split a long indivisible idea rather than silently shrink it.
    columns = max(1, int(width / (size * .6)))
    lines = sum(max(1, len(textwrap.wrap(line, columns))) for line in text.splitlines())
    return lines * size * 1.35 <= height


def slide_plan(lesson, segment, *, terms=(), rescue=False, image="", sample=False, profile=None):
    preset_id = lesson.get("teaching_preset", "standard")
    preset = next((item for item in catalog()["presets"] if item["id"] == preset_id), None)
    if preset is None:
        raise ValueError("Mẫu bài giảng không hợp lệ.")
    block = block_type(segment.get("kind", "unknown"))
    layout = lesson.get("layout", "line_pair")
    content = presentation_content(segment, lesson.get("level", 2), layout, rescue, terms)
    result = {"width": 1280, "height": 720, "font": "Arial", "preset": preset_id,
              "kind": block["id"], "background": preset["background"], "accent": preset["accent"],
              "ink": preset["ink"], "muted": preset["muted"], "elements": [], "image": image,
              "image_box": {}, "overflow": False, "sample": sample}

    def element(text, x, y, width, height, size, color, *, bold=False, italic=False, align="left", body=False):
        if not text:
            return
        result["elements"].append({"text": text, "x": x, "y": y, "width": width, "height": height,
                                   "size": size, "color": color, "bold": bold, "italic": italic,
                                   "align": align, "body": body})
        if body and not _fits(text, width, height, size):
            result["overflow"] = True

    element("BiliClass   /   " + lesson.get("subject", "") + "   /   " + str(lesson.get("grade", "")),
            64, 28, 1152, 32, 22, preset["accent"], bold=True)
    element(lesson.get("title", ""), 64, 76, 1152, 64, 40, preset["ink"], bold=True)
    byline = " · ".join(filter(None, ((profile or {}).get("teacher", ""), (profile or {}).get("school", ""))))
    show_profile = bool(byline and (profile or {}).get("show_profile", True))
    block_label = block["label"] + "  /  " + block["en"] if content["show_vi"] else block["en"]
    element(block_label, 64, 176 if show_profile else 152, 1152, 30, 20, preset["muted"])
    element(segment.get("locator", ""), 64, 670, 1152, 28, 17, preset["muted"])
    if show_profile:
        element(byline[:150], 64, 145, 1152, 23, 17, preset["muted"])
    x, y, width, height = 64, 214, 1152, 410
    composition = block["composition"]
    if composition == "hero":
        x, y, width, height = 96, 228, 1088, 402
    if composition == "visual" and (image or sample):
        width = 548
        result["image_box"] = {"x": 696, "y": 214, "width": 520, "height": 410}
        if sample and not image:
            element("[Hình từ tài liệu bài học]\n[Image from the lesson source]", 728, 320, 456, 180,
                    28, preset["muted"], align="center")
    size = block["font_size"]
    if preset_id == "practice" and composition == "prompt":
        size = 42
    if content["show_vi"] and content["show_en"]:
        if layout == "split_view":
            gap = 36
            regions = [(x, y, (width-gap)/2, height), (x+(width+gap)/2, y, (width-gap)/2, height)]
            size = min(size, 34)
        else:
            regions = [(x, y, width, (height-34)/2), (x, y+(height+34)/2, width, (height-34)/2)]
        for text, rect, color, italic in ((content["vi"], regions[0], preset["ink"], False),
                                         (content["en"], regions[1], preset["accent"], layout == "line_pair")):
            element(text, *rect, size, color, bold=composition in {"hero", "focus", "prompt"},
                    italic=italic, align=block["align"], body=True)
    else:
        text = content["vi"] if content["show_vi"] else content["en"]
        element(text, x, y, width, height, size, preset["ink"] if content["show_vi"] else preset["accent"],
                bold=composition in {"hero", "focus", "prompt"}, align=block["align"], body=True)
    return result


def slide_pages(lesson, segment, *, terms=(), image="", profile=None, rescue=False):
    content = presentation_content(segment, lesson.get("level", 2), lesson.get("layout", "line_pair"), rescue, terms)
    vi = content["vi"] if content["show_vi"] else ""
    en = content["en"] if content["show_en"] else ""
    # Start at corresponding lines, then pack those pairs into the actual regions.
    pairs = paired_pages(vi, en, limit=340 if lesson.get("layout") == "split_view" else 430)
    line_pairs = []
    for left, right in pairs:
        left_lines, right_lines = left.splitlines(), right.splitlines()
        if left_lines and right_lines and len(left_lines) == len(right_lines):
            line_pairs.extend(zip(left_lines, right_lines, strict=True))
        elif not left:
            line_pairs.extend(("", line) for line in right_lines)
        elif not right:
            line_pairs.extend((line, "") for line in left_lines)
        else:
            line_pairs.append((left, right))
    packed, current = [], ("", "")
    for left, right in line_pairs:
        candidate = ("\n".join(filter(None, (current[0], left))), "\n".join(filter(None, (current[1], right))))
        plan = slide_plan(lesson, {**segment, "vi": candidate[0], "en": candidate[1], "support": []},
                          image=image, profile=profile, rescue=rescue)
        if plan["overflow"] and any(current):
            packed.append(current)
            current = (left, right)
        else:
            current = candidate
    if any(current):
        packed.append(current)
    plans = []
    for index, (left, right) in enumerate(packed):
        # Already resolved keywords/easy English: avoid applying transformations twice.
        page_segment = {**segment, "vi": left, "en": right, "support": []}
        plan = slide_plan(lesson, page_segment, terms=(), image=image, profile=profile, rescue=rescue)
        if plan["overflow"]:
            raise ValueError(f"{segment.get('locator', '')}: nội dung vượt khung mẫu. Hãy tách ý và duyệt lại cặp Việt–Anh.")
        if len(packed) > 1:
            for item in plan["elements"]:
                if item["y"] == 670:
                    item["text"] = segment.get("locator", "") + f" · {index+1}/{len(packed)}"
        plans.append(plan)
    return plans


def source_image(lesson, directory, segment):
    """Extract the first ordinary source image for an explicitly visual block."""
    if block_type(segment.get("kind"))["id"] != "visual":
        return None
    from .powerpoint import slide_for_locator, verified_presentation

    source = lesson.get("source") or {}
    slide_index = slide_for_locator(segment.get("locator", ""))
    if not source.get("file", "").lower().endswith(".pptx") or not slide_index:
        return None
    folder = Path(directory) / "temp/template-images"
    prefix = source["sha256"] + "-" + str(slide_index)
    cached = next(iter(folder.glob(prefix + ".*")), None)
    if cached:
        return cached
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(verified_presentation(lesson, directory))
    if slide_index > len(deck.slides):
        return None

    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                yield shape

    picture = next(pictures(deck.slides[slide_index - 1].shapes), None)
    if picture is None:
        return None
    folder.mkdir(parents=True, exist_ok=True)
    destination = folder / (prefix + "." + picture.image.ext)
    destination.write_bytes(picture.image.blob)
    return destination


def example_plan(preset, kind, level=2, layout="line_pair"):
    block = block_type(kind)
    return slide_plan({"title": "Mẫu bài giảng song ngữ", "subject": "Môn học của thầy cô", "grade": "10–12",
                       "teaching_preset": preset, "level": level, "layout": layout},
                      {"kind": kind, "vi": block["vi"], "en": block["sample_en"],
                       "locator": "Mẫu cấu trúc · Thay bằng nội dung bài học"}, sample=True)
