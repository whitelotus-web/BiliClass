import io
import json
import urllib.error
import zipfile
from contextlib import closing
from threading import Event

import pytest

from app import azure_translation
from app.azure_translation import AzureKeyStore, AzureTranslator
from app.conversion_input import creation_config
from app.conversion_progress import ConversionProgress
from app.google_conversion import convert, load, prepare
from app.google_settings_tasks import change, read_state
from app.google_translation import KeyStore, TranslationProblem, Translator
from app.library import Library
from app.storage import backup_library
from app.translation_provider import build_translator, require_credentials, snapshot

AZURE_KEY = 'a' * 32
GOOGLE_KEY = 'AIza' + 'g' * 35


def test_google_default_and_old_checkpoint_stay_google_after_azure_selected(tmp_path):
    assert snapshot(tmp_path)['translation_provider'] == 'google'
    with closing(Library(tmp_path)):
        pass
    config = creation_config('Thí nghiệm', 'Vật lý', 'THPT', '11', 'parallel_columns', 'standard', 'template')
    old = prepare(tmp_path, config, text='Vật chuyển động.')
    raw = json.loads((tmp_path / 'google_jobs' / old['config']['request_id'] / 'config.json').read_text())
    raw.pop('translation_provider')
    raw.pop('translation_budget')
    (tmp_path / 'google_jobs' / raw['request_id'] / 'config.json').write_text(json.dumps(raw))
    with closing(Library(tmp_path)) as library:
        library.set_setting('translation_provider', 'azure')  # An older RC168 installation.
    assert snapshot(tmp_path) == {'translation_provider': 'google', 'translation_budget': 500_000}
    called = []
    result = build_translator(tmp_path, load(tmp_path, old['folder'])['config'], 123,
                              lambda *args: called.append(args) or 'google')
    assert result == 'google' and called == [(tmp_path, 123)]
    new = prepare(tmp_path, {**config, 'translation_provider': 'azure'}, text='Vật chuyển động.')
    assert new['config']['translation_provider'] == 'azure'
    assert not change(tmp_path, 'provider', ('google',))['error']
    assert load(tmp_path, new['folder'])['config']['translation_provider'] == 'azure'
    assert isinstance(build_translator(tmp_path, new['config'], 500_000, None), AzureTranslator)
    assert change(tmp_path, 'provider', ('unknown',))['error']
    assert snapshot(tmp_path)['translation_provider'] == 'google'


def test_azure_secret_rotation_and_remove_readd_keep_resource_budget(tmp_path, monkeypatch):
    store = AzureKeyStore(tmp_path)
    store.configure(AZURE_KEY, 'southeastasia', 'Trường học')
    assert AZURE_KEY.encode() not in store.path.read_bytes()
    assert AZURE_KEY not in str(store.public())
    store.reserve(1_500_000, 2_000_000)
    identity = store.public()[0]['id']
    store.configure('b' * 32, 'southeastasia', 'Tên mới')
    assert store.public()[0]['used'] == 1_500_000
    assert store.public()[0]['id'] == identity
    store.remove(identity)
    store.configure(AZURE_KEY, 'global', '')
    assert store.public()[0]['used'] == 1_500_000
    with pytest.raises(TranslationProblem):
        store.reserve(500_001, 2_000_000)
    monkeypatch.setattr(KeyStore, 'month', staticmethod(lambda: '2099-01'))
    assert store.public()[0]['used'] == 0


def test_provider_credentials_cache_and_budget_are_isolated(tmp_path):
    KeyStore(tmp_path).add(GOOGLE_KEY)
    AzureKeyStore(tmp_path).configure(AZURE_KEY, 'global', '')
    google_calls, azure_calls = [], []
    google = Translator(tmp_path, transport=lambda rows, *_: google_calls.extend(rows) or ['Google'] * len(rows))
    azure = AzureTranslator(tmp_path, transport=lambda rows, *_: azure_calls.extend(rows) or ['Azure'] * len(rows))
    assert google.translate(['Xin chào'], Event()) == ['Google']
    assert azure.translate(['Xin chào'], Event()) == ['Azure']
    assert len(google_calls) == len(azure_calls) == 1
    assert google.translate(['Xin chào'], Event()) == ['Google']
    assert azure.translate(['Xin chào'], Event()) == ['Azure']
    assert len(google_calls) == len(azure_calls) == 1
    assert KeyStore(tmp_path).public()[0]['used'] == len('Xin chào')
    assert AzureKeyStore(tmp_path).public()[0]['used'] == len('Xin chào<div></div>')
    AzureKeyStore(tmp_path).remove(AzureKeyStore(tmp_path).public()[0]['id'])
    with pytest.raises(TranslationProblem, match='Azure'):
        require_credentials(tmp_path, 'azure')
    require_credentials(tmp_path, 'google')


@pytest.mark.parametrize('region', ['global', 'southeastasia'])
def test_azure_http_contract_and_no_redirect(monkeypatch, region):
    seen = []

    class Response:
        def __enter__(self): return self
        def __exit__(self, *_): pass
        def read(self, limit):
            assert limit == 1_000_001
            return b'[{"translations":[{"text":"Hello","to":"en"}]}]'

    class Opener:
        def open(self, request, timeout):
            seen.append(request)
            assert timeout == 20
            assert request.full_url.startswith(azure_translation.ENDPOINT + '?')
            assert request.get_header('Ocp-apim-subscription-key') == AZURE_KEY
            assert request.get_header('Ocp-apim-subscription-region') == (None if region == 'global' else region)
            assert AZURE_KEY not in request.full_url
            assert 'textType=html' in request.full_url
            assert json.loads(request.data) == [{'Text': '<div>Xin chào</div>'}]
            return Response()

    def opener(handler):
        assert handler().redirect_request(None, None, 302, '', {}, 'https://example.org') is None
        return Opener()

    monkeypatch.setattr(azure_translation.urllib.request, 'build_opener', opener)
    assert azure_translation.request(['Xin chào'], AZURE_KEY, 'vi', 'en', region) == ['Hello']
    assert len(seen) == 1


@pytest.mark.parametrize('status,detail,code', [(401, 401000, 'api_key'), (403, 403001, 'api_quota'),
    (403, 403000, 'api_key'), (429, 429000, 'api_rate'), (503, 503000, 'api_network'), (400, 400000, 'api_request')])
def test_azure_http_failure_is_bounded_and_does_not_echo_secret(monkeypatch, status, detail, code):
    class Opener:
        def open(self, *_args, **_kwargs):
            body = json.dumps({'error': {'code': detail, 'message': AZURE_KEY + 'Private source'}}).encode()
            raise urllib.error.HTTPError(azure_translation.ENDPOINT, status, AZURE_KEY, {}, io.BytesIO(body))

    monkeypatch.setattr(azure_translation.urllib.request, 'build_opener', lambda *_: Opener())
    with pytest.raises(TranslationProblem) as caught:
        azure_translation.request(['Private source'], AZURE_KEY, 'vi', 'en', 'global')
    assert caught.value.code == code
    assert AZURE_KEY not in str(caught.value) and 'Private source' not in str(caught.value)


@pytest.mark.parametrize('payload', [[], {}, [{'translations': []}],
    [{'translations': [{'text': 'Hello', 'to': 'fr'}]}], [{'translations': [{'text': '', 'to': 'en'}]}]])
def test_azure_never_accepts_missing_or_wrong_language_result(monkeypatch, payload):
    class Response:
        def __enter__(self): return self
        def __exit__(self, *_): pass
        def read(self, _limit): return json.dumps(payload).encode()
    monkeypatch.setattr(azure_translation.urllib.request, 'build_opener',
                        lambda *_: type('Opener', (), {'open': lambda *_args, **_kwargs: Response()})())
    with pytest.raises(TranslationProblem) as caught:
        azure_translation.request(['Xin chào'], AZURE_KEY, 'vi', 'en', 'global')
    assert caught.value.code == 'api_result'


def test_azure_uses_existing_number_guards_cancel_and_never_falls_back(tmp_path):
    AzureKeyStore(tmp_path).configure(AZURE_KEY, 'global', '')
    KeyStore(tmp_path).add(GOOGLE_KEY)
    calls = []

    def bad(rows, key, *_):
        assert key == AZURE_KEY
        calls.append(rows)
        return ['Velocity 999']

    with pytest.raises(TranslationProblem):
        AzureTranslator(tmp_path, transport=bad).translate(['Vận tốc 12'], Event(), fragment_repair=False)
    assert 1 <= len(calls) <= 2
    assert KeyStore(tmp_path).public()[0]['used'] == 0
    cancel = Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        AzureTranslator(tmp_path, transport=lambda *_: pytest.fail('sent after cancel')).translate(['Xin chào'], cancel)


def test_azure_bad_key_keeps_authentication_error_without_using_google(tmp_path):
    AzureKeyStore(tmp_path).configure(AZURE_KEY, 'global', '')
    KeyStore(tmp_path).add(GOOGLE_KEY)
    calls = []

    def reject(*args):
        calls.append(args)
        raise TranslationProblem('api_key', 'Kiểm tra key và Region Azure.')

    with pytest.raises(TranslationProblem) as caught:
        AzureTranslator(tmp_path, transport=reject).translate(['Xin chào'], Event())
    assert caught.value.code == 'api_key' and len(calls) == 1
    assert AzureKeyStore(tmp_path).public()[0]['used'] == 0
    assert KeyStore(tmp_path).public()[0]['used'] == 0


def test_settings_mask_keys_and_inactive_corruption_does_not_break_google(tmp_path):
    with closing(Library(tmp_path)):
        pass
    KeyStore(tmp_path).add(GOOGLE_KEY)
    assert not change(tmp_path, 'azure_save', (AZURE_KEY, 'southeastasia', 'BiliClass'))['error']
    state = read_state(tmp_path)
    assert state['provider'] == 'google' and state['activeConfigured'] and state['azure']['configured']
    assert AZURE_KEY not in json.dumps(state) and GOOGLE_KEY not in json.dumps(state)
    AzureKeyStore(tmp_path).path.write_bytes(b'broken test credential')
    state = read_state(tmp_path)
    assert state['activeConfigured'] and state['azure']['message']
    assert change(tmp_path, 'provider', ('azure',))['error']
    assert not read_state(tmp_path, provider='azure')['activeConfigured']
    assert not change(tmp_path, 'azure_save', (AZURE_KEY, 'global', 'Đã sửa'))['error']
    assert read_state(tmp_path)['activeConfigured']


@pytest.mark.parametrize('style', ['source', 'template'])
@pytest.mark.parametrize('mode', ['parallel_columns', 'english_only'])
def test_azure_pptx_conversion_routes_pinned_engine_and_resumes_without_network(tmp_path, monkeypatch, style, mode):
    from pptx import Presentation
    from pptx.util import Inches

    from app import google_conversion

    with closing(Library(tmp_path)) as library:
        library.set_setting('translation_provider', 'azure')
    AzureKeyStore(tmp_path).configure(AZURE_KEY, 'global', '')
    calls = []

    def fixture(rows, key, source, target, region):
        assert (key, source, target, region) == (AZURE_KEY, 'vi', 'en', 'global')
        calls.extend(rows)
        return [text.replace('Vật có ', 'The object has a ') for text in rows]

    monkeypatch.setattr(azure_translation.AzureTranslator, '_send',
        lambda self, rows, credential, source, target: fixture(rows, credential['key'], source, target, credential['region']))
    monkeypatch.setattr(google_conversion, 'Translator', lambda *_: pytest.fail('Google was invoked'))
    config = creation_config('Cơ năng', 'Vật lý', 'THPT', '11', mode, 'standard', style)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = 'Vật có vận tốc 12 m/s.'
    source = tmp_path / 'Thí nghiệm gốc.pptx'
    deck.save(source)
    original = source.read_bytes()
    job = prepare(tmp_path, {**config, 'translation_provider': 'azure'}, source)
    change(tmp_path, 'provider', ('google',))
    progress = ConversionProgress(job['folder'], lambda *_: None)
    lesson, preview = convert(tmp_path, job, Event(), progress)
    assert calls and preview['path']
    assert lesson['external_deck']['translation_provider'] == 'azure'
    assert lesson['narration']['mode'] == 'none'
    assert source.read_bytes() == original
    result = Presentation(preview['path'])
    text = '\n'.join(shape.text for page in result.slides for shape in page.shapes if shape.has_text_frame)
    assert 'The object has a velocity' in text and '12' in text
    assert ('Vật có vận tốc' in text) == (mode == 'parallel_columns')
    used = AzureKeyStore(tmp_path).public()[0]['used']
    again, _ = convert(tmp_path, job, Event(), progress)
    assert again['id'] == lesson['id']
    assert AzureKeyStore(tmp_path).public()[0]['used'] == used
    assert not KeyStore(tmp_path).keys()


def test_selection_is_not_persisted_if_destination_settings_cannot_be_read(tmp_path, monkeypatch):
    from app import google_settings_tasks

    with closing(Library(tmp_path)):
        pass

    def fail(*_args, **_kwargs):
        raise OSError('Read failed')

    monkeypatch.setattr(google_settings_tasks, 'read_state', fail)
    assert change(tmp_path, 'provider', ('azure',))['error']
    assert snapshot(tmp_path)['translation_provider'] == 'google'


def test_backup_excludes_both_provider_secrets_and_caches(tmp_path):
    with closing(Library(tmp_path)):
        pass
    KeyStore(tmp_path).add(GOOGLE_KEY)
    AzureKeyStore(tmp_path).configure(AZURE_KEY, 'global', '')
    archive = backup_library(tmp_path)
    with zipfile.ZipFile(archive) as package:
        assert all(not name.startswith(('translation_api/', 'translation_azure/')) for name in package.namelist())
        assert all(AZURE_KEY.encode() not in package.read(name) for name in package.namelist())
