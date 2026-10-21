"""Opening another deck is a transaction; choosing or cancelling is not one."""

import pytest
from PySide6.QtCore import QCoreApplication, QObject, Signal

from app.library import Library
from app.ui import Bridge


class Office(QObject):
    stateChanged = Signal(object)
    failed = Signal(str)
    finished = Signal()

    def __init__(self):
        super().__init__()
        self.commands = []
        self.stops = 0

    def navigate(self, *args):
        self.commands.append(args)

    def stop(self):
        self.stops += 1


@pytest.fixture
def session(tmp_path, monkeypatch):
    app = QCoreApplication.instance() or QCoreApplication([])
    library = Library(tmp_path / 'data')
    lesson = library.create('First', 'Vật lý', 'THPT', '11', [('Energy', 'Năng lượng')])
    bridge = Bridge(library)
    worker = Office()
    bridge.powerpoint_worker = worker
    bridge._lesson = lesson
    bridge._powerpoint_lesson_id = lesson['id']
    bridge._powerpoint_state = dict(active=True, opened=True, slide=3, total=8)
    bridge.teaching_session.bound = dict(lesson=lesson, preview={}, media={})
    monkeypatch.setattr(bridge, 'focusTeachingPresentation', lambda: True)
    monkeypatch.setattr(bridge, 'stopSpeech', lambda: None)
    yield bridge, worker, lesson
    bridge.powerpoint_worker = None
    bridge.teaching_session.bound = {}
    for item in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                 bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                 bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
        item.shutdown()
    app.processEvents()
    library.close()


def test_open_same_lesson_focuses_without_reset_and_mascot_returns_to_library(session, monkeypatch):
    bridge, worker, lesson = session
    pages = []
    bridge.navigate.connect(pages.append)
    monkeypatch.setattr(bridge, 'launch', lambda *_: pytest.fail('Same deck must not be prepared again'))
    token = bridge.teaching_session.token
    bridge.openCompanionLibraryLesson(lesson['id'])
    assert pages == ['library'] and bridge.powerpoint_worker is worker
    assert bridge.powerpointState['slide'] == 3 and bridge.powerpointState['active']
    assert bridge.teaching_session.token == token and worker.stops == 0
    bridge._powerpoint_state['active'] = False
    bridge.openTeachingLibraryLesson(lesson['id'])
    assert pages == ['library'] and bridge.powerpoint_worker is worker
    bridge.openTeachingPresentation('')  # Cancelled file selection.
    assert worker.stops == 0


@pytest.mark.parametrize('state', [{'starting': True}, {'closing': True}, {'show_pending': 'start'}, {'show_pending': 'end'}])
def test_all_open_routes_reject_concurrent_office_transitions(session, monkeypatch, state):
    bridge, worker, lesson = session
    bridge._powerpoint_state.update(state)
    monkeypatch.setattr(bridge, 'launch', lambda *_: pytest.fail('Cannot queue another open'))
    bridge.openTeachingLibraryLesson(lesson['id'])
    bridge.openCompanionLibraryLesson(lesson['id'])
    bridge.openTeachingPresentation('file:///not-used.pptx')
    bridge.openInputPowerPoint('file:///not-used.pptx')
    bridge.openNarratedPowerPoint('not-used')
    bridge.useConvertedLesson()
    bridge.openConvertedDeck()
    bridge.startPowerPoint()
    bridge.startBilingualPowerPoint()
    bridge.previewSourcePowerPoint()
    bridge._lesson['source'] = {'file': 'not-used.pptx'}
    bridge.openSourceCopy()
    bridge.openTemplateSample('standard')
    bridge.openLocalTemplateSample('standard', 'parallel_columns')
    assert worker.stops == 0 and bridge.powerpoint_worker is worker


def test_original_and_converted_products_are_distinct_bindings(session):
    bridge, worker, lesson = session
    assert not bridge._focus_open_teaching(lesson['id'], product='original')
    bridge.teaching_session.bound['preview'] = {'product': 'original'}
    assert bridge._focus_open_teaching(lesson['id'], product='original')
    assert not bridge._focus_open_teaching(lesson['id'])


def test_changed_speaker_position_requires_fresh_managed_product(session):
    bridge, worker, lesson = session
    bridge._powerpoint_state['active'] = False
    saved = bridge.library.get(lesson['id'])
    saved['speaker_position'] = 'top_left'
    bridge.library._write(saved)
    assert not bridge._focus_open_teaching(lesson['id'])
    assert worker.stops == 0, 'Checking freshness must not close the bound editor'
    bridge.teaching_session.bound['lesson']['speaker_position'] = 'top_left'
    assert bridge._focus_open_teaching(lesson['id'])


@pytest.mark.parametrize('opener,destination', [('openInputPowerPoint', None), ('openTeachingPresentation', 'library')])
def test_local_file_open_uses_managed_copy_and_keeps_its_navigation_contract(session, monkeypatch, tmp_path, opener, destination):
    from pptx import Presentation
    from PySide6.QtCore import QUrl

    bridge, _, _ = session
    path = tmp_path / 'Bài 6 – Dao động tắt dần.pptx'
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    original = path.read_bytes()
    bridge.powerpoint_worker = None
    bridge.teaching_session.bound = {}
    pages, prepared = [], []
    bridge.navigate.connect(pages.append)
    monkeypatch.setattr(bridge, 'launch', lambda action, callback, **_: callback(action()))
    monkeypatch.setattr(bridge, '_activate_teaching', prepared.append)
    monkeypatch.setattr('app.ui.QDesktopServices.openUrl', lambda *_: pytest.fail('PowerPoint must be managed'))
    getattr(bridge, opener)(QUrl.fromLocalFile(str(path)).toString())
    assert pages == ([destination] if destination else [])
    assert len(prepared) == 1 and prepared[0]['lesson']['title'] == path.stem
    assert prepared[0]['owned_path'] != str(path)
    assert path.read_bytes() == original and not bridge.narrationBusy
    # Cancellation is inert; reopening the same bound product keeps its session.
    bridge.powerpoint_worker = Office()
    bridge._powerpoint_state = dict(opened=True, active=True, slide=2)
    bridge.teaching_session.bound = prepared[0]
    getattr(bridge, opener)('')
    getattr(bridge, opener)(QUrl.fromLocalFile(str(path)).toString())
    assert len(prepared) == 1 and bridge.powerpointState['slide'] == 2


def test_replacement_waits_for_old_office_close_and_ignores_late_old_signals(session, monkeypatch):
    bridge, old, _ = session
    other = bridge.library.create('Second', 'Vật lý', 'THPT', '11', [('Another', 'Khác')])
    prepared = dict(lesson=other, preview={'lesson_id': other['id'], 'revision': other['revision']})
    activated = []

    def activate(value):
        activated.append(value)
        bridge.powerpoint_worker = current
        bridge.teaching_session.bound = value
        bridge._powerpoint_state = dict(active=False, opened=True, slide=1)

    current = Office()
    monkeypatch.setattr(bridge, '_activate_teaching', activate)
    old.finished.connect(bridge._powerpoint_finished)
    old.stateChanged.connect(bridge._powerpoint_changed)
    old.failed.connect(bridge._powerpoint_failed)
    bridge._open_library_teaching(prepared, destination='library')
    assert old.stops == 1 and not activated and bridge.teachingTransitioning
    assert bridge._pending_teaching is prepared
    old.stateChanged.emit(dict(active=True, opened=True, slide=7))
    assert bridge.powerpointState.get('closing')
    old.finished.emit()
    assert activated == [prepared] and bridge.powerpoint_worker is current
    old.stateChanged.emit(dict(active=True, opened=True, slide=7))
    old.failed.emit('Stale failure')
    old.finished.emit()
    assert bridge.powerpoint_worker is current and bridge.powerpointState['slide'] == 1
    assert not bridge.error and not bridge._pending_teaching


def test_start_end_are_idempotent_and_ending_retains_ready_binding(session):
    bridge, worker, lesson = session
    bridge.startTeachingShow()
    assert not worker.commands, 'Starting an active show must preserve it'
    token = bridge.teaching_session.token
    bridge.navigatePowerPoint('end_show')
    bridge.navigatePowerPoint('end_show')
    assert worker.commands == [('end_show',)]
    assert not bridge.companionCanRead and bridge.teaching_session.token > token
    bridge._powerpoint_changed(dict(active=False, opened=True, slide=1, total=8))
    assert bridge.teaching_session.state['lesson_id'] == lesson['id'] and not bridge.teachingTransitioning
    bridge.startTeachingShow()
    bridge.startTeachingShow()
    assert worker.commands == [('end_show',), ('start',)]
    bridge._powerpoint_changed(dict(active=True, opened=True, slide=1, total=8))
    assert not bridge.teachingTransitioning and bridge.companionCanRead


def test_failed_preparation_does_not_close_current_show(session, monkeypatch):
    bridge, worker, lesson = session
    actions = []
    monkeypatch.setattr(bridge, 'launch', lambda action, _callback: actions.append(action))
    bridge.openCompanionLibraryLesson('missing-id')
    with pytest.raises(ValueError, match='Không tìm thấy'):
        actions[0]()
    assert bridge.teaching_session.state['lesson_id'] == lesson['id']
    assert worker.stops == 0 and bridge.powerpointState['active']


def test_explicit_close_is_idempotent_and_clears_only_the_bound_session(session):
    bridge, worker, lesson = session
    before = bridge.library.get(lesson['id'])
    worker.finished.connect(bridge._powerpoint_finished)
    bridge.stopPowerPoint()
    bridge.stopPowerPoint()
    assert worker.stops == 1 and bridge.teachingTransitioning
    worker.finished.emit()
    assert not bridge.teaching_session.bound and bridge.powerpoint_worker is None
    assert not bridge.teachingTransitioning
    assert bridge.library.get(lesson['id']) == before
