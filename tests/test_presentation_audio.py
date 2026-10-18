import wave
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app import speech
from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.lesson_files import save_presentation
from app.lesson_narration import make_plan
from app.library import Library
from app.presentation_audio import control_slots, slide_audio
from app.source_deck import NS, source_slide_parts


def test_crowded_slide_can_place_controls_in_separate_clear_gaps():
    from app.layout_regions import intersection

    occupied = [(0, 50, 300, 100), (70, 0, 32, 50), (166, 0, 32, 50), (262, 0, 38, 50)]
    slots = control_slots(['vi', 'en', 'both'], (300, 150), occupied, lambda _: True)
    assert slots and len(slots) == 3
    boxes = [(x, y, 40 if compact else 62, 48 if compact else 42) for x, y, compact in slots.values()]
    assert all(not intersection(box, obstacle) for box in boxes for obstacle in occupied)
    assert not any(intersection(boxes[i], boxes[j]) for i in range(3) for j in range(i))


def test_controls_never_cover_evidence_when_no_clear_slot_exists():
    assert control_slots(['en', 'vi'], (300, 150), [(0, 0, 300, 150)], lambda _: True) is None


def test_labelled_speaker_can_use_a_clear_narrow_header_without_covering_text():
    slots = control_slots(['en'], (960, 540), [(0, 28, 960, 512)], lambda _: True)
    assert slots and slots['en'][2] == 'mini'
    assert slots['en'][1] + 20 <= 28


def fixture(tmp_path, mode="both", crowded=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(7), Inches(2)
    ).text = "Năng lượng được bảo toàn.\nEnergy is conserved."
    slide.notes_slide.notes_text_frame.text = "VI: Năng lượng được bảo toàn.\nEN: Energy is conserved."
    if crowded:
        second = deck.slides.add_slide(deck.slide_layouts[6])
        second.shapes.add_textbox(0, 0, deck.slide_width, deck.slide_height).text = slide.shapes[0].text
        second.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    path = tmp_path / "source.pptx"
    deck.save(path)
    library = Library(tmp_path / "library")
    config = {
        "title": "Test",
        "subject": "Vật lý",
        "education_level": "THPT",
        "grade": "10",
        "conversion_format": "parallel_columns",
        "style": "source",
        "preset": "standard",
    }
    lesson = library.create_external_lesson(config, library.store_source(path), inspect_returned_deck(path))
    plan, _ = make_plan(lesson, mode, {"voice": "vieneu:Mai Anh", "rate": 0})
    for keys in plan["keys"].values():
        for key in keys:
            target = library.directory / "audio" / (key + ".wav")
            target.parent.mkdir(exist_ok=True)
            with wave.open(str(target), "wb") as writer:
                writer.setparams((1, 2, 24000, 0, "NONE", "not compressed"))
                writer.writeframes(bytes([int(key[:2],16),0]) * 2400)
    lesson = library.update_narration(lesson["id"], plan)
    return library, lesson, path


def test_crowded_page_does_not_discard_other_speakers_and_is_reported(tmp_path):
    library, lesson, source = fixture(tmp_path, crowded=True)
    original = source.read_bytes()
    try:
        report = {}
        target = tmp_path / 'partial.pptx'
        save_presentation(lesson, external_preview(lesson, library.directory), library.directory,
                          target, with_voice=True, report=report)
        assert report['embedded_pages'] == [1]
        assert report['unplaced_pages'] == [2]
        assert report['speakers'] == 2
        assert source.read_bytes() == original
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(target) as after:
            part = source_slide_parts(source)[1]
            assert before.read(part) == after.read(part), 'Unplaceable artwork and animations stay intact'
    finally:
        library.close()


@pytest.mark.parametrize("mode,count", [("en", 1), ("vi", 1), ("both", 2)])
def test_selected_speakers_are_embedded_internal_click_only_and_source_is_unchanged(
    tmp_path, mode, count, monkeypatch
):
    library, lesson, source = fixture(tmp_path, mode)
    monkeypatch.setattr(speech, "synthesize", lambda *_: pytest.fail("No model at export"))
    original = source.read_bytes()
    target = tmp_path / "exported.pptx"
    try:
        save_presentation(
            lesson, external_preview(lesson, library.directory), library.directory, target, with_voice=True
        )
        with zipfile.ZipFile(target) as archive:
            wavs = [n for n in archive.namelist() if n.startswith("ppt/media/") and n.endswith(".wav")]
            assert len(wavs) == count
            part = source_slide_parts(target)[0]
            tree = etree.fromstring(archive.read(part))
            assert len(tree.findall(".//p:audio", NS)) == count
            clocks = tree.xpath(".//p:cTn/@id", namespaces=NS)
            assert len(clocks) == len(set(clocks))
            assert tree.xpath('.//p:cond[@evt="onClick"]/p:tgtEl/p:spTgt', namespaces=NS)
            assert (
                len(
                    tree.xpath('.//p:cMediaNode/p:cTn/p:stCondLst/p:cond[@delay="indefinite"]', namespaces=NS)
                )
                == count
            )
            for name in archive.namelist():
                if name.endswith(".rels"):
                    assert all(
                        n.get("TargetMode") != "External"
                        for n in etree.fromstring(archive.read(name))
                        if n.get("Type", "").endswith(("/audio", "/media"))
                    )
        assert source.read_bytes() == original
    finally:
        library.close()


def test_missing_audio_fails_atomically_and_does_not_overwrite_download(tmp_path):
    library, lesson, _ = fixture(tmp_path)
    key = lesson["narration"]["keys"]["en"][0]
    (library.directory / "audio" / (key + ".wav")).unlink()
    target = tmp_path / "output.pptx"
    target.write_bytes(b"previous file")
    try:
        with pytest.raises(ValueError, match="Tạo đủ"):
            save_presentation(
                lesson,
                external_preview(lesson, library.directory),
                library.directory,
                target,
                with_voice=True,
            )
        assert target.read_bytes() == b"previous file"
    finally:
        library.close()


def test_slide_recording_excludes_answers_hidden_behind_click(tmp_path):
    library, lesson, _ = fixture(tmp_path)
    try:
        page = lesson["slide_speech"]["slides"][0]
        page["englishSpeechSegments"][0]["visibility"] = {
            "kind": "click",
            "show_after": 1,
            "hide_after": None,
            "settle_seconds": 0,
        }
        # Keep the signature/voice manifest in sync just as an edited lesson must.
        plan, _ = make_plan(lesson, "both", {"voice": "vieneu:Mai Anh", "rate": 0})
        lesson["narration"] = plan
        with pytest.raises(ValueError, match="Chưa có lời đọc tĩnh"):
            slide_audio(lesson, library.directory, scope='visible')
        assert set(slide_audio(lesson, library.directory)[1]) == {'vi', 'en'}, 'An explicit whole-slide speaker retains saved words behind builds'
    finally:
        library.close()


def test_bilingual_product_has_individual_and_combined_speakers_with_no_new_inference(tmp_path, monkeypatch):
    import io

    from app.presentation_audio import joined_wav

    library, lesson, source = fixture(tmp_path)
    try:
        plan, _ = make_plan(lesson, 'both', {'voice': 'vieneu:Mai Anh', 'rate': 0}, playback='both_controls')
        lesson = library.update_narration(lesson['id'], plan)
        pages = slide_audio(lesson, library.directory)
        assert set(pages[1]) == {'vi', 'en', 'both'}
        assert pages[1]['both'] == pages[1]['vi'] + pages[1]['en']
        audio, _ = joined_wav(library.directory, pages[1]['both'])
        with wave.open(io.BytesIO(audio)) as reader:
            samples = reader.readframes(reader.getnframes())
        assert samples[:2] == bytes([int(plan['keys']['vi'][0][:2], 16), 0])
        assert samples[2*(2400+4320):][:2] == bytes([int(plan['keys']['en'][0][:2], 16), 0])
        monkeypatch.setattr(speech, 'synthesize', lambda *_: pytest.fail('Reuse prepared chunks'))
        target = tmp_path / 'three-speakers.pptx'
        save_presentation(lesson, external_preview(lesson, library.directory), library.directory, target, with_voice=True)
        with zipfile.ZipFile(target) as archive:
            tree = etree.fromstring(archive.read(source_slide_parts(target)[0]))
            assert len(tree.findall('.//p:audio', NS)) == 3
            assert set(tree.xpath('.//p:cNvPr/@name', namespaces=NS)) >= {'BiliClass audio VI', 'BiliClass audio EN', 'BiliClass audio BOTH'}
        en, _ = make_plan(lesson, 'en', {'voice': 'vieneu:Mai Anh', 'rate': 0})
        library.update_narration(lesson['id'], en)
        from app.lesson_narration import variant_id
        lesson = library.select_narration_variant(lesson['id'], variant_id(en))
        assert set(slide_audio(lesson, library.directory)[1]) == {'en'}
    finally:
        library.close()


def test_static_vietnamese_is_allowed_despite_unrelated_timing_but_not_if_its_shape_is_animated(tmp_path):
    from app.presentation_audio import static_vietnamese
    library, lesson, source = fixture(tmp_path)
    try:
        deck = Presentation(source)
        slide = deck.slides[0]
        timing = etree.SubElement(slide._element, '{'+NS['p']+'}timing')
        assert static_vietnamese(slide, lesson, 1)
        target = etree.SubElement(timing, '{'+NS['p']+'}spTgt')
        target.set('spid', str(slide.shapes[0].shape_id))
        assert not static_vietnamese(slide, lesson, 1)
    finally:
        library.close()
