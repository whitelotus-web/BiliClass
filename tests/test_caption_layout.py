from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from app.caption_layout import reserve_band
from app.source_deck import NS


def caption_page(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    art = tmp_path/'picture.png'
    Image.new('RGB', (400, 250), 'blue').save(art)
    for x in (50, 490):
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(str(art), Pt(x), Pt(220), Pt(400), Pt(250))
        for offset in (0, 210):
            caption = group.shapes.add_textbox(Pt(x+offset), Pt(480), Pt(190), Pt(50))
            caption.name = 'BCL_BODY Caption'
            caption.text = 'A vehicle traveling 108 km is considered a point mass in this example.'
            caption.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    return deck, slide


def test_grouped_picture_band_preserves_aspect_media_and_words(tmp_path):
    deck, slide = caption_page(tmp_path)
    tree = etree.fromstring(etree.tostring(slide._element))
    before_text = tree.xpath('.//a:t/text()', namespaces=NS)
    before_media = tree.xpath('.//a:blip/@r:embed', namespaces={**NS, 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})
    pictures = tree.findall('.//p:pic/p:spPr/a:xfrm/a:ext', NS)
    ratios = [int(ext.get('cx'))/int(ext.get('cy')) for ext in pictures]
    changes = reserve_band(tree, deck.slide_width, deck.slide_height, slide)
    assert len([row for row in changes if row['operation'] == 'reserve_caption_band']) == 2
    assert len([row for row in changes if row['operation'] == 'align_caption_band']) == 4
    assert [int(ext.get('cx'))/int(ext.get('cy')) for ext in pictures] == ratios
    assert tree.xpath('.//a:t/text()', namespaces=NS) == before_text
    assert tree.xpath('.//a:blip/@r:embed', namespaces={**NS, 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}) == before_media


def test_caption_band_never_covers_unrelated_footer(tmp_path):
    deck, slide = caption_page(tmp_path)
    footer = slide.shapes.add_textbox(Pt(20), Pt(510), Pt(600), Pt(20))
    footer.text = 'Teacher footer'
    tree = etree.fromstring(etree.tostring(slide._element))
    before = etree.tostring(tree)
    assert not reserve_band(tree, deck.slide_width, deck.slide_height, slide)
    assert etree.tostring(tree) == before
