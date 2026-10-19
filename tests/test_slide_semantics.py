import pytest
from pptx import Presentation
from pptx.util import Inches

from app.native_bilingual import assemble, extract
from app.slide_semantics import quantity_only


@pytest.mark.parametrize('value', ['12 kg', '24 m/s', '12,5 m/s²', '3.2e-3 V', '10 MPa', '20 °C'])
def test_spaced_scientific_quantities_are_language_independent(value):
    assert quantity_only(value)


@pytest.mark.parametrize('value', ['12 lớp', 'The mass is 12 kg', 'Khối lượng 12 kg', '3 in class'])
def test_prose_is_never_treated_as_a_shared_quantity(value):
    assert not quantity_only(value)


def test_bilingual_table_keeps_one_copy_of_spaced_quantity(tmp_path):
    source, target = tmp_path / 'source.pptx', tmp_path / 'result.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(8), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = 'Khối lượng', '12 kg'
    deck.save(source)
    original = source.read_bytes()
    manifest = extract(source, 'parallel_columns', preserve_design=True, same_slide=True)
    values = {unit['id']: {'en': unit['text'].replace('Khối lượng', 'Mass')} for unit in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    table = next(shape.table for shape in Presentation(target).slides[0].shapes if shape.has_table)
    assert table.cell(0, 1).text == '12\u00a0kg'
    assert table.cell(0, 0).text == 'Khối lượng\nMass'
    assert source.read_bytes() == original
