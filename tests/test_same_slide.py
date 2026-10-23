"""Same-slide delivery, visible coverage and native build synchronization."""

import hashlib
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.chatgpt_handoff import inspect_returned_deck
from app.native_bilingual import assemble, extract
from app.source_deck import NS, source_slide_parts


@pytest.mark.parametrize('content', [
    'Phải quay kim phút mấy phần của một vòng tròn để nó chỉ đúng số 12?',
    'Dòng nội dung thứ nhất\nDòng nội dung thứ hai\nDòng nội dung thứ ba',
    'Một đoạn nội dung giải thích chuyển động của vật. ' * 6,
    'Cho tam giác ABC có điểm M là trung điểm của cạnh BC.',
])
def test_title_placeholder_used_for_content_is_not_a_heading(tmp_path, content):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    title = slide.shapes.title
    title.text = content
    title.text_frame.paragraphs[0].font.size = Pt(32)
    path = tmp_path/'misused-title.pptx'
    deck.save(path)
    manifest = extract(path, 'parallel_columns', same_slide=True)
    assert all(not u['title'] for u in manifest['units'] if u['shape_id'] == title.shape_id)


def test_native_table_uses_its_cell_font_and_fits_inside_the_source_grid(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(26.667), Inches(15)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(.4), Inches(1.5), Inches(25.5), Inches(10))
    table = shape.table
    table.rows[0].height = Inches(1.5)
    table.rows[1].height = Inches(8.5)
    table.cell(0, 0).text = 'CÂU HỎI'
    table.cell(1, 0).text = 'Phải quay kim phút mấy phần của một vòng tròn để nó chỉ đúng số 12?'
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(44)
    source = tmp_path/'table.pptx'
    deck.save(source)
    before = source.read_bytes()
    manifest = extract(source, 'parallel_columns', same_slide=True)
    assert all(u['font_size'] == 4400 for u in manifest['units'])
    assert all(not u['title'] and 'same_slide_box' not in u for u in manifest['units'])
    values = {u['id']: {'en': 'QUESTION' if u['raw'] == 'CÂU HỎI' else
                         'What fraction of a full turn must the minute hand rotate to point to ⟦BLC0⟧?'}
              for u in manifest['units']}
    output = assemble(source, tmp_path/'paired.pptx', manifest, values, 'parallel_columns')
    result = Presentation(output).slides[0].shapes[0]
    assert (result.left, result.top, result.width, result.height) == (shape.left, shape.top, shape.width, shape.height)
    assert 'What fraction' in result.table.cell(1, 0).text
    assert max((p.runs[0].font.size or p.font.size).pt for p in result.table.cell(1, 0).text_frame.paragraphs if p.runs) <= 44
    assert source.read_bytes() == before


def test_manual_wrapped_table_question_is_translated_as_one_sentence(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(7), Inches(4))
    shape.table.cell(0, 0).text = 'Phải quay kim phút mấy\nphần của một vòng tròn để nó chỉ đúng\nsố 12?'
    source = tmp_path/'wrapped.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    assert len(manifest['units']) == 1
    unit = manifest['units'][0]
    assert unit['merged_paragraphs'] == [0, 1, 2]
    assert 'kim phút mấy phần' in unit['raw']
    output = assemble(source, tmp_path/'paired.pptx', manifest,
        {unit['id']: {'en': 'What fraction of a turn must the hand rotate to point to ⟦BLC0⟧?'}}, 'parallel_columns')
    paragraphs = Presentation(output).slides[0].shapes[0].table.cell(0, 0).text_frame.paragraphs
    assert len([p for p in paragraphs if p.text.strip()]) == 2
    assert 'kim phút mấy phần' in paragraphs[0].text
    assert 'What fraction' in paragraphs[1].text


def test_small_source_label_does_not_double_with_a_large_canvas(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(26.667), Inches(15)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(.7))
    label.text = 'Giải'
    label.text_frame.paragraphs[0].font.size = Pt(24)
    source = tmp_path/'small-label.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    output = assemble(source, tmp_path/'paired.pptx', manifest,
        {manifest['units'][0]['id']: {'en': 'Solution'}}, 'parallel_columns')
    paragraph = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0]
    assert (paragraph.runs[0].font.size or paragraph.font.size).pt <= 24


def test_small_running_footer_keeps_its_band_and_source_font_hierarchy(tmp_path):
    deck=Presentation()
    deck.slide_width,deck.slide_height=Inches(13.333),Inches(7.5)
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    label=slide.shapes.add_textbox(Inches(.6),Inches(7.2),Inches(6),Inches(.22))
    label.text='BILICLASS / BÀI GIẢNG MẪU TIẾNG VIỆT'
    label.text_frame.paragraphs[0].font.size=Pt(8.2)
    source=tmp_path/'running-footer.pptx'
    deck.save(source)
    before=source.read_bytes()
    manifest=extract(source,'parallel_columns',same_slide=True)
    unit=manifest['units'][0]
    assert unit['label'] and not unit['title'] and unit['inline_pair']
    output=assemble(source,tmp_path/'paired.pptx',manifest,
                    {unit['id']:{'en':'BILICLASS / VIETNAMESE SAMPLE LESSON'}},'parallel_columns')
    result=Presentation(output).slides[0].shapes[0]
    assert (result.left,result.top,result.width,result.height)==(label.left,label.top,label.width,label.height)
    assert 'BÀI GIẢNG MẪU' in result.text and 'VIETNAMESE SAMPLE LESSON' in result.text
    sizes=[int(n.get('sz'))/100 for n in result._element.findall('.//a:defRPr',NS)]
    assert max(sizes)<=8.2
    assert source.read_bytes()==before


def test_table_with_only_office_math_still_has_cell_geometry_and_reserves_a_solution_label(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(26.667), Inches(15)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_table(1, 1, Inches(1), Inches(5), Inches(12), Inches(2))
    cell = shape.table.cell(0, 0)
    cell.text_frame.paragraphs[0]._p.append(etree.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<m:r><m:rPr><m:nor/></m:rPr><a:rPr sz="4400"/><m:t>Ta có công thức sau đây.</m:t></m:r></m:oMath>'
    ))
    label = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(.8))
    label.text = 'Lời giải'
    path = tmp_path/'math-table.pptx'
    deck.save(path)
    manifest = extract(path, 'parallel_columns', same_slide=True)
    unit = next(u for u in manifest['units'] if u['shape_id'] == shape.shape_id)
    assert unit['canvas_scale'] == 2 and not unit['title']
    assert unit['cell_margins']['marL'] > Inches(2)


def test_short_solution_body_reserves_room_for_a_stacked_solution_heading(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label = slide.shapes.add_textbox(Inches(.5), Inches(3), Inches(2), Inches(.6))
    label.text = 'Lời giải'
    body = slide.shapes.add_textbox(Inches(.5), Inches(3), Inches(7), Inches(2))
    body.text = 'Bán kính quỹ đạo của trạm vũ trụ quốc tế là:'
    source = tmp_path/'solution-label.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    unit = next(u for u in manifest['units'] if u['shape_id'] == body.shape_id)
    assert unit['same_slide_box']['y'] >= label.top+label.height


@pytest.mark.parametrize('label', ['❶', '❷', '❸', '❹', '❺', '①', '⓿', '½', 'Ⅲ'])
def test_shared_numeric_badges_keep_one_source_copy_and_its_size(tmp_path, label):
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1), Inches(1), Inches(.4), Inches(.5))
    badge.text = label
    badge.text_frame.paragraphs[0].font.size = Pt(18)
    caption = slide.shapes.add_textbox(Inches(2), Inches(1), Inches(6), Inches(1))
    caption.text = 'Khái niệm mở đầu'
    source = tmp_path/'badges.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    unit = next(u for u in manifest['units'] if u['raw'] == label)
    assert unit['language'] == 'neutral'
    values = {u['id']: {'en': u['text'] if u['language'] == 'neutral' else 'Introduction'}
              for u in manifest['units']}
    output = assemble(source, tmp_path/'paired.pptx', manifest, values, 'parallel_columns')
    converted = Presentation(output)
    result = next(shape for shape in converted.slides[0].shapes if shape.shape_id == badge.shape_id)
    assert result.text == label
    assert (result.left, result.top, result.width, result.height) == (badge.left, badge.top, badge.width, badge.height)
    assert result.text_frame.paragraphs[0].font.size == Pt(18)


def test_bilingual_title_retains_top_anchor_in_oversized_source_textbox(tmp_path):
    from pptx.enum.text import MSO_ANCHOR

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(.5), Inches(7), Inches(2))
    title.text = 'ĐƯỜNG THẲNG VÀ MẶT PHẲNG TRONG KHÔNG GIAN'
    title.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    title.text_frame.paragraphs[0].font.size = Pt(32)
    source = tmp_path/'banner.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    assert manifest['units'][0]['title']
    values = {u['id']: {'en': 'LINES AND PLANES IN SPACE'} for u in manifest['units']}
    result = assemble(source, tmp_path/'result.pptx', manifest, values, 'parallel_columns')
    converted = Presentation(result).slides[0].shapes[0]
    assert converted.text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert 'LINES AND PLANES' in converted.text


def test_typography_scales_with_canvas_instead_of_halving_large_deck_text(tmp_path):
    results = []
    for scale in (1, 2):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333 * scale), Inches(7.5 * scale)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(scale), Inches(scale), Inches(10 * scale), Inches(scale))
        title.text = 'Cơ năng'
        title.text_frame.paragraphs[0].font.size = Pt(32 * scale)
        source = tmp_path/f'canvas-{scale}.pptx'
        deck.save(source)
        manifest = extract(source, 'parallel_columns', same_slide=True)
        values = {u['id']: {'en': 'Mechanical energy'} for u in manifest['units']}
        output = assemble(source, tmp_path/f'result-{scale}.pptx', manifest, values, 'parallel_columns')
        results.append(Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0].font.size)
    assert results[1] == results[0] * 2


def test_native_measurement_visits_group_children_but_does_not_fit_rotated_groups():
    from types import SimpleNamespace as Shape

    from app.native_layout import office_text_shapes

    text = Shape(Type=17, Rotation=0, Id=11, HasTextFrame=True)
    nested = Shape(Type=6, Rotation=0, GroupItems=[text])
    rotated = Shape(Type=6, Rotation=15, GroupItems=[Shape(Type=17, Rotation=0, Id=12)])
    assert [
        s.Id
        for s in office_text_shapes([nested, rotated, Shape(Type=13, Rotation=0, Id=13, HasTextFrame=False)])
    ] == [11]


def fixture(path, *, animated=False):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    title.text = "Cơ năng"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(7), Inches(4))
    body.text = (
        "Năng lượng của vật được bảo toàn khi không có lực ma sát.\nVận tốc của vật tăng khi thế năng giảm."
    )
    for shape in (title, body):
        for p in shape.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(28)
    picture = path.with_suffix(".png")
    Image.new("RGB", (200, 200), "#46A3CF").save(picture)
    slide.shapes.add_picture(str(picture), Inches(8), Inches(2), width=Inches(4))
    if animated:
        timing = etree.SubElement(slide._element, f"{{{NS['p']}}}timing")
        sequence = etree.SubElement(timing, f"{{{NS['p']}}}tnLst")
        parent = etree.SubElement(sequence, f"{{{NS['p']}}}par")
        clock = etree.SubElement(parent, f"{{{NS['p']}}}cTn", id="1", nodeType="clickEffect")
        children = etree.SubElement(clock, f"{{{NS['p']}}}childTnLst")
        behavior = etree.SubElement(children, f"{{{NS['p']}}}animEffect", transition="in", filter="fade")
        common = etree.SubElement(behavior, f"{{{NS['p']}}}cBhvr")
        etree.SubElement(common, f"{{{NS['p']}}}cTn", id="2", dur="500")
        target = etree.SubElement(
            etree.SubElement(common, f"{{{NS['p']}}}tgtEl"), f"{{{NS['p']}}}spTgt", spid=str(body.shape_id)
        )
        text = etree.SubElement(target, f"{{{NS['p']}}}txEl")
        etree.SubElement(text, f"{{{NS['p']}}}pRg", st="1", end="1")
        builds = etree.SubElement(timing, f"{{{NS['p']}}}bldLst")
        etree.SubElement(builds, f"{{{NS['p']}}}bldP", spid=str(body.shape_id), grpId="0")
    deck.save(path)
    return path


def translations(manifest):
    values = [
        "Mechanical energy",
        "The energy of the object is conserved in the absence of friction.",
        "The velocity increases when the potential energy decreases.",
    ]
    return {unit["id"]: {"en": values[i]} for i, unit in enumerate(manifest["units"])}


def test_english_only_keeps_source_art_without_visible_vietnamese(tmp_path):
    source = fixture(tmp_path / "original.pptx")
    manifest = extract(source, "english_only", preserve_design=True)
    output = assemble(source, tmp_path / "english.pptx", manifest, translations(manifest), "english_only")
    deck = Presentation(output)
    text = "\n".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert "The energy" in text and "The velocity" in text
    assert "Năng lượng" not in text and "Cơ năng" not in text
    assert any(int(s.shape_type) == 13 for s in deck.slides[0].shapes)
    original = Presentation(source)
    assert [(s.shape_id, s.left, s.top, s.width, s.height) for s in deck.slides[0].shapes] == [
        (s.shape_id, s.left, s.top, s.width, s.height) for s in original.slides[0].shapes
    ]


@pytest.mark.parametrize("mode", ["parallel_columns", "sentence_pairs"])
def test_both_languages_visible_on_each_source_page_and_resources_preserved(tmp_path, mode):
    source = fixture(tmp_path / "original.pptx")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    manifest = extract(source, mode, same_slide=True)
    assert manifest["version"] == 4 and not manifest["paired_slides"]
    output = assemble(source, tmp_path / "output.pptx", manifest, translations(manifest), mode)
    deck = Presentation(output)
    assert len(deck.slides) == 1
    text = "\n".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert "Năng lượng" in text and "The energy" in text and "The velocity" in text
    assert deck.slides[0].shapes[0].text == "Cơ năng\nMechanical energy"
    assert len(manifest["layout_plan"][0]["pairs"]) == 3
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        for name in before.namelist():
            if name.startswith(("ppt/media/", "ppt/theme/", "ppt/slideLayouts/", "ppt/slideMasters/")):
                assert before.read(name) == after.read(name)
    report = inspect_returned_deck(output, source, mode, layout_plan=manifest["layout_plan"])
    assert not any(
        i["code"] in {"translation_pair_missing", "english_body_missing", "parallel_columns_missing"}
        for row in report["slide_quality"]["slides"]
        for i in row["issues"]
    )
    paragraph = deck.slides[0].shapes[0].text_frame.paragraphs[1]
    assert paragraph.runs[0].font.italic
    assert paragraph.runs[0].font.name == "Arial"
    vi = deck.slides[0].shapes[0].text_frame.paragraphs[0]
    assert paragraph.runs[0].font.size < (vi.runs[0].font.size or vi.font.size)


def test_notes_and_declared_pairs_cannot_hide_deleted_english(tmp_path):
    source = fixture(tmp_path / "source.pptx")
    manifest = extract(source, "parallel_columns", same_slide=True)
    output = assemble(source, tmp_path / "output.pptx", manifest, translations(manifest), "parallel_columns")
    deck = Presentation(output)
    pair = manifest["layout_plan"][0]["pairs"][1]
    next(s for s in deck.slides[0].shapes if s.shape_id == pair["en_shape"]).text = ""
    deck.save(output)
    report = inspect_returned_deck(output, source, "parallel_columns", layout_plan=manifest["layout_plan"])
    assert any(
        i["code"] == "translation_pair_missing" and i["severity"] == "FAIL"
        for row in report["slide_quality"]["slides"]
        for i in row["issues"]
    )


def test_transparent_english_is_not_visible_coverage(tmp_path):
    from app.same_slide import visible_pairs

    source = fixture(tmp_path / "source.pptx")
    manifest = extract(source, "parallel_columns", same_slide=True)
    output = assemble(source, tmp_path / "result.pptx", manifest, translations(manifest), "parallel_columns")
    with zipfile.ZipFile(output) as archive:
        tree = etree.fromstring(archive.read(source_slide_parts(output)[0]))
    pairs = manifest["layout_plan"][0]["pairs"]
    pair = pairs[1]
    node = tree.xpath(".//p:sp[p:nvSpPr/p:cNvPr[@id=$id]]", namespaces=NS, id=str(pair["en_shape"]))[0]
    for props in node.findall(".//a:rPr", NS):
        fill = etree.SubElement(props, f"{{{NS['a']}}}solidFill")
        color = etree.SubElement(fill, f"{{{NS['a']}}}srgbClr", val="000000")
        etree.SubElement(color, f"{{{NS['a']}}}alpha", val="0")
    assert not visible_pairs(tree, pairs)


def test_numeric_options_remain_single_native_objects(tmp_path):
    source = fixture(tmp_path / "source.pptx")
    deck = Presentation(source)
    option = deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(2), Inches(0.5))
    option.text = "A. 2,5 m/s²"
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    unit = next(u for u in manifest["units"] if u["shape_id"] == option.shape_id)
    assert unit["language"] == "neutral"
    values = {u["id"]: {"en": "Mechanical energy" if u["title"] else u["text"]} for u in manifest["units"]}
    output = assemble(source, tmp_path / "result.pptx", manifest, values, "parallel_columns")
    assert (
        sum(s.text.replace('\u00a0', ' ').count("A. 2,5 m/s²") for s in Presentation(output).slides[0].shapes if s.has_text_frame)
        == 1
    )


def test_formula_variable_and_protected_xml_are_language_neutral(tmp_path):
    source = fixture(tmp_path / "source.pptx")
    deck = Presentation(source)
    paragraph = deck.slides[0].shapes[1].text_frame.add_paragraph()
    paragraph.add_run().text = "a"
    math = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    formula = etree.SubElement(paragraph._p, f"{{{math}}}oMath")
    run = etree.SubElement(formula, f"{{{math}}}r")
    etree.SubElement(run, f"{{{math}}}t").text = "=±ω²A"
    deck.save(source)
    manifest = extract(source, "sentence_pairs", same_slide=True)
    unit = next(u for u in manifest["units"] if u["paragraph"] == 2)
    assert unit["text"] == "a⟦BLC0⟧" and unit["language"] == "neutral"


def test_roman_badges_and_standalone_cosine_equations_are_not_translated(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(0.5), Inches(0.5)).text = "II"
    shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
    shape.text = "Phương trình của gia tốc là:\na = -ω²A cos(ωt + φ)"
    for p in shape.text_frame.paragraphs:
        p.font.name = "Arial"
        p.font.size = Pt(24)
    source = tmp_path / "equations.pptx"
    deck.save(source)
    manifest = extract(source, "sentence_pairs", same_slide=True)
    assert [u["language"] for u in manifest["units"]] == ["neutral", "vi", "neutral"]
    values = {
        u["id"]: {"en": u["text"] if u["language"] == "neutral" else "The equation of acceleration is:"}
        for u in manifest["units"]
    }
    output = assemble(source, tmp_path / "paired.pptx", manifest, values, "sentence_pairs")
    text = "\n".join(s.text for s in Presentation(output).slides[0].shapes if s.has_text_frame)
    assert text.count("II") == 1 and text.count("a = -ω²A cos(ωt + φ)") == 1
    assert "The equation of acceleration is:" in text
    assert not any(
        i["code"] == "translation_pair_missing"
        for row in inspect_returned_deck(
            output, source, "sentence_pairs", layout_plan=manifest["layout_plan"]
        )["slide_quality"]["slides"]
        for i in row["issues"]
    )


def test_short_bold_section_headers_keep_same_stacked_style_at_different_heights(tmp_path):
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for text, y in [("Vận tốc của vật dao động điều hòa", 0.5), ("Gia tốc của vật dao động điều hòa", 4.5)]:
        shape = slide.shapes.add_textbox(Inches(4), Inches(y), Inches(5), Inches(0.6))
        shape.text = text
        p = shape.text_frame.paragraphs[0]
        p.font.name = "Arial"
        p.font.size = Pt(24)
        for run in p.runs:
            run.font.bold = True
    source = tmp_path / "headers.pptx"
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    assert all(u["title"] for u in manifest["units"])
    _output = assemble(
        source,
        tmp_path / "headers-en.pptx",
        manifest,
        {
            u["id"]: {"en": "The acceleration of an object undergoing simple harmonic motion."}
            for u in manifest["units"]
        },
        "parallel_columns",
    )
    assert all(p["inline"] for p in manifest["layout_plan"][0]["pairs"])


def test_short_caption_keeps_original_height_over_a_raster_backdrop(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = tmp_path / "backdrop.png"
    Image.new("RGB", (400, 300), "#ffe066").save(background)
    slide.shapes.add_picture(str(background), 0, 0, width=deck.slide_width, height=deck.slide_height)
    caption = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4.5), Inches(0.78))
    caption.text = "Nội dung bài học"
    caption.text_frame.paragraphs[0].font.size = Pt(32)
    source = tmp_path / "decorated.pptx"
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    unit = next(u for u in manifest["units"] if u["shape_id"] == caption.shape_id)
    assert unit["same_slide_box"]["h"] == caption.height
    output = assemble(
        source,
        tmp_path / "bilingual.pptx",
        manifest,
        {unit["id"]: {"en": "Lesson content"}},
        "parallel_columns",
    )
    result = next(s for s in Presentation(output).slides[0].shapes if s.shape_id == caption.shape_id)
    assert result.height == caption.height and "Lesson content" in result.text


def test_group_child_outside_in_local_coordinates_is_translated_when_visible(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(50), Inches(50), Inches(7), Inches(3))
    child.text = "Kéo quả nặng lệch khỏi vị trí cân bằng rồi buông tay."
    child.text_frame.paragraphs[0].font.size = Pt(24)
    transform = group._element.find("p:grpSpPr/a:xfrm", NS)
    transform.find("a:off", NS).set("x", str(Inches(1)))
    transform.find("a:off", NS).set("y", str(Inches(2)))
    source = tmp_path / "group.pptx"
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    assert len(manifest["units"]) == 1 and manifest["units"][0]["shape_id"] == child.shape_id
    output = assemble(
        source,
        tmp_path / "output.pptx",
        manifest,
        {manifest["units"][0]["id"]: {"en": "Displace the weight from equilibrium and release it."}},
        "parallel_columns",
    )
    from app.slide_speech import objects

    text = " ".join(item["text"] for item in objects(Presentation(output).slides[0]))
    assert "Kéo quả nặng" in text and "Displace the weight" in text


@pytest.mark.parametrize("mode", ["parallel_columns", "sentence_pairs"])
def test_answer_english_shares_source_click_and_paragraph_build(tmp_path, mode):
    source = fixture(tmp_path / "source.pptx", animated=True)
    manifest = extract(source, mode, same_slide=True)
    output = assemble(source, tmp_path / "output.pptx", manifest, translations(manifest), mode)
    with zipfile.ZipFile(output) as archive:
        tree = etree.fromstring(archive.read(source_slide_parts(output)[0]))
    timing = tree.find("p:timing", NS)
    ids = timing.xpath(".//p:cTn/@id", namespaces=NS)
    assert len(set(ids)) == len(ids)
    pair = manifest["layout_plan"][0]["pairs"][2]
    targets = timing.findall(".//p:spTgt", NS)
    if mode == "parallel_columns":
        assert {int(n.get("spid")) for n in targets} == {pair["vi_shape"], pair["en_shape"]}
        effects = timing.findall(".//p:animEffect", NS)
        assert len(effects) == 2 and effects[0].getparent() is effects[1].getparent()
        assert len(timing.findall(".//p:bldP", NS)) == 2
    else:
        region = targets[0].find("p:txEl/p:pRg", NS)
        assert region.attrib == {"st": "2", "end": "3"}


def test_long_math_or_single_body_is_not_misclassified_as_title(tmp_path):
    source = fixture(tmp_path / "source.pptx")
    deck = Presentation(source)
    body = deck.slides[0].shapes[1]
    body.top = Inches(0.1)
    body.text = "Nhập hàm số và kiểm tra vận tốc của vật. " * 12
    body.text_frame.paragraphs[0].font.size = Pt(32)
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    assert not next(u for u in manifest["units"] if u["shape_id"] == body.shape_id)["title"]


def test_top_question_block_with_three_paragraphs_uses_body_columns(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    body = slide.shapes.add_textbox(Inches(1), Inches(.7), Inches(11), Inches(1.4))
    body.text = ('Quan sát các vật trong hình vẽ.\n'
                 '1. Giải thích tại sao quả cầu đứng yên khi được treo.\n'
                 '2. Tại sao người trượt ván có thể giữ nguyên vận tốc của mình?')
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.size = Pt(25)
    source = tmp_path / 'question.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    assert all(not unit['title'] for unit in manifest['units'])
    from app.translation_batches import TOKEN

    values = {unit['id']: {'en': ' '.join(TOKEN.findall(unit['text'])) +
                          ' Observe and explain the motion of the object.'} for unit in manifest['units']}
    assemble(source, tmp_path / 'bilingual.pptx', manifest, values, 'parallel_columns')
    assert all(not pair['inline'] for pair in manifest['layout_plan'][0]['pairs'])


def test_table_cells_keep_both_languages(tmp_path):
    source = fixture(tmp_path / "source.pptx")
    deck = Presentation(source)
    table = deck.slides[0].shapes.add_table(1, 1, Inches(0.5), Inches(6.4), Inches(10), Inches(0.7)).table
    table.cell(0, 0).text = "Năng lượng được bảo toàn."
    deck.save(source)
    manifest = extract(source, "parallel_columns", same_slide=True)
    values = {u["id"]: {"en": "The energy is conserved."} for u in manifest["units"]}
    output = assemble(source, tmp_path / "result.pptx", manifest, values, "parallel_columns")
    cell = next(s.table.cell(0, 0) for s in Presentation(output).slides[0].shapes if s.has_table)
    assert cell.text == "Năng lượng được bảo toàn.\nThe energy is conserved."


@pytest.mark.parametrize("mode", ["parallel_columns", "sentence_pairs"])
def test_dense_frames_continue_with_whole_bilingual_pairs(tmp_path, mode):
    source = fixture(tmp_path / "source.pptx", animated=True)
    deck = Presentation(source)
    body = deck.slides[0].shapes[1]
    body.height = Inches(1.4)
    body.text = "\n".join("Năng lượng của vật được bảo toàn khi không có lực ma sát." for _ in range(8))
    for p in body.text_frame.paragraphs:
        p.font.size = Pt(26)
    deck.save(source)
    manifest = extract(source, mode, same_slide=True)
    values = {
        u["id"]: {
            "en": "Mechanical energy"
            if u["title"]
            else "The energy of the object is conserved in the absence of friction."
        }
        for u in manifest["units"]
    }
    output = assemble(source, tmp_path / "result.pptx", manifest, values, mode)
    plan = manifest["layout_plan"]
    assert len(plan) > 1
    assert all(row["role"] == "bilingual" and row["source_slide"] == 1 for row in plan)
    assert [row["continuation"] for row in plan[1:]] == list(range(2, len(plan) + 1))
    body_units = [p["unit"] for row in plan for p in row["pairs"] if not p["title"]]
    assert len(body_units) == 8 and len(set(body_units)) == 8
    for slide in Presentation(output).slides:
        text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert "Năng lượng" in text and "The energy" in text
    report = inspect_returned_deck(output, source, mode, layout_plan=plan)
    assert not any(
        i["code"] in {"translation_pair_missing", "slide_count"}
        for row in report["slide_quality"]["slides"]
        for i in row["issues"]
    )


def test_one_long_paragraph_must_not_disable_pagination_of_its_whole_frame(tmp_path):
    source = fixture(tmp_path / 'source.pptx')
    deck = Presentation(source)
    body = deck.slides[0].shapes[1]
    body.height = Inches(1.3)
    body.text = '\n'.join('Năng lượng của vật được bảo toàn khi không có lực ma sát.' for _ in range(3))
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.size = Pt(26)
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    body_units = [u for u in manifest['units'] if not u['title']]
    values = {u['id']: {'en': 'Mechanical energy' if u['title'] else 'A short translated paragraph.'}
              for u in manifest['units']}
    values[body_units[0]['id']]['en'] = 'The object moves under gravity and keeps its mechanical energy. ' * 4
    output = assemble(source, tmp_path / 'result.pptx', manifest, values, 'parallel_columns')
    plan = manifest['layout_plan']
    assert len(plan) > 1
    scheduled = [pair['unit'] for row in plan for pair in row['pairs'] if not pair['title']]
    assert sorted(scheduled) == sorted(u['id'] for u in body_units)
    english = '\n'.join(s.text for slide in Presentation(output).slides for s in slide.shapes
                        if s.has_text_frame and s.name.endswith('English'))
    assert english.count('The object moves under gravity') == 4
