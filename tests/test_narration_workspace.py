import json
import wave
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.chatgpt_handoff import external_preview
from app.lesson_narration import audio_status, make_plan
from app.library import Library, RevisionConflict
from app.narration_audio_export import export_audio
from app.narration_script import segments as narration_segments
from app.readiness import text_readiness
from app.voice_import import import_presentation, prepare_library_lesson, save_script


def source(path, text='', notes=''):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if text:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = text
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    return path


def test_idle_bound_mascot_allows_voice_configuration_without_rebinding(tmp_path, monkeypatch):
    import copy
    from PySide6.QtCore import QCoreApplication
    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    library = Library(tmp_path / 'data')
    first = library.create('Bound lesson', 'Vật lý', 'THPT', '11', [('Slide 1', 'Năng lượng.')])
    second = library.create('Voice draft', 'Vật lý', 'THPT', '11', [('Slide 1', 'Dao động.')])
    bridge = Bridge(library)
    worker = object()
    bound = dict(lesson=copy.deepcopy(first))
    stopped, prepared = [], []
    try:
        bridge.powerpoint_worker = worker
        bridge._powerpoint_state = dict(opened=True, active=False)
        bridge.teaching_session.bound = copy.deepcopy(bound)
        monkeypatch.setattr(bridge, 'stopSpeech', lambda: stopped.append(True))
        monkeypatch.setattr(bridge, 'launch', lambda action, callback, **_: callback(action()))
        monkeypatch.setattr('app.voice_import.prepare_library_lesson',
                            lambda *_: (second, dict(lesson_id=second['id'], revision=second['revision'], path='fixture.pptx')))
        workspace = bridge.narration_workspace
        assert workspace.configurationEnabled and not workspace.blockReason
        workspace.setSpeakerPosition('top_left')
        assert workspace.speakerPosition == 'top_left'
        workspace.selectLibrary(second['id'])
        assert bridge._lesson['id'] == second['id'] and stopped
        assert bridge.teaching_session.bound == bound and bridge.powerpoint_worker is worker
        assert not bridge.error
        monkeypatch.setattr(bridge, '_available_voice', lambda _: True)
        def prepare(*args):
            prepared.append(args[1])
            return dict(already_created=True)
        monkeypatch.setattr('app.narration_preparation.prepare_request', prepare)
        bridge._create_narration('vi', bridge.voiceSettings['voice'], 0)
        assert prepared == [second['id']], 'Idle binding must not silently discard Create voice'
        assert bridge.teaching_session.bound == bound and bridge.powerpoint_worker is worker
        monkeypatch.setattr(bridge, 'launch', lambda *_args, **_kwargs: pytest.fail('Blocked configuration must not start a job'))
        for state in (dict(active=True), dict(starting=True), dict(closing=True), dict(show_pending='start')):
            bridge._powerpoint_state = dict(opened=True, **state)
            assert not workspace.configurationEnabled and workspace.blockReason
            workspace.selectLibrary(first['id'])
            assert bridge._lesson['id'] == second['id']
            assert bridge.error and bridge.message == workspace.blockReason
    finally:
        bridge.powerpoint_worker = None
        bridge._powerpoint_state = {}
        bridge.teaching_session.bound = {}
        for service in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                        bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                        bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
            service.shutdown()
        app.processEvents()
        library.close()


def test_intentionally_silent_slide_preserves_text_and_can_be_restored(tmp_path):
    path = source(tmp_path / 'Original.pptx', 'VI: Năng lượng\nEN: Energy')
    original = path.read_bytes()
    lesson, _ = import_presentation(tmp_path / 'library', path)
    library = Library(tmp_path / 'library')
    try:
        row = lesson['segments'][0]
        original_plan, _ = make_plan(lesson, 'both', {'voice': 'vieneu:Mai Anh', 'rate': 0})
        assert original_plan['total']
        quiet = save_script(library, lesson['id'], row['id'], lesson['revision'],
                             row['vi'], row['en'], silent=True)
        assert quiet['segments'][0]['vi'] == row['vi'] and quiet['segments'][0]['en'] == row['en']
        assert narration_segments(quiet)[0]['silent']
        assert not narration_segments(quiet)[0]['vi'] and not narration_segments(quiet)[0]['en']
        plan, _ = make_plan(quiet, 'both', {'voice': 'vieneu:Mai Anh', 'rate': 0})
        assert plan['total'] == 0
        restored = save_script(library, quiet['id'], row['id'], quiet['revision'],
                                row['vi'], row['en'], silent=False)
        assert not narration_segments(restored)[0]['silent']
        restored_plan, _ = make_plan(restored, 'both', {'voice': 'vieneu:Mai Anh', 'rate': 0})
        assert restored_plan['keys'] == original_plan['keys']
        assert not restored['segments'][0]['approved']
        assert path.read_bytes() == original
    finally:
        library.close()


def test_uncertain_language_is_explained_and_manual_narration_wins():
    from app.narration_script import classification

    assert classification('VI: Năng lượng', 'EN: Energy')['confidence'] == 1
    assert classification('x² = y²', '')['needs_review']
    assert not classification('x² = y²', '', manual=True)['needs_review']


def complete(library, lesson, mode='both'):
    plan, _ = make_plan(lesson, mode, {'voice': 'vieneu:Mai Anh', 'rate': 0}, playback='both_controls' if mode == 'both' else 'separate')
    for keys in plan['keys'].values():
        for key in keys:
            target = library.directory / 'audio' / (key + '.wav')
            target.parent.mkdir(exist_ok=True)
            with wave.open(str(target), 'wb') as writer:
                writer.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                writer.writeframes(b'\1\0' * 100)
    return library.update_narration(lesson['id'], plan)


def test_single_english_notes_are_the_actual_spoken_text_and_blank_slide_is_not_fabricated(tmp_path):
    path = source(tmp_path / 'English.pptx', 'Computer', 'EN: A computer processes information.\nCHECK: Do not read this check.')
    lesson, preview = import_presentation(tmp_path / 'data', path)
    assert lesson['segments'][0]['en'] == 'A computer processes information.'
    assert lesson['segments'][1]['vi'] == lesson['segments'][1]['en'] == ''
    plan, _ = make_plan(lesson, 'en', {'voice': 'vieneu:Mai Anh', 'rate': 0})
    assert plan['total'] == 1 and not plan['skipped']
    library = Library(tmp_path / 'data')
    try:
        lesson = complete(library, lesson, 'en')
        assert audio_status(lesson, library.directory)['ready']
        reviewed = library.review_external_deck(lesson['id'], lesson['revision'])
        assert text_readiness(reviewed, library.directory)['text_ready']
        assert not reviewed['segments'][1]['approved']
        assert external_preview(reviewed, library.directory)['path'] == preview['path']
    finally:
        library.close()


def test_edit_script_preserves_pptx_invalidates_audio_requires_new_review_and_reimports_saved_script(tmp_path):
    path = source(tmp_path / 'Lesson.pptx', 'Máy tính xử lý thông tin.')
    original = path.read_bytes()
    lesson, _ = import_presentation(tmp_path / 'data', path)
    library = Library(tmp_path / 'data')
    try:
        lesson = complete(library, lesson, 'vi')
        lesson = library.review_external_deck(lesson['id'], lesson['revision'])
        changed = save_script(library, lesson['id'], lesson['segments'][0]['id'], lesson['revision'], 'Lời đọc mới.', 'New narration.')
        assert changed['revision'] == lesson['revision'] + 1
        assert not changed['segments'][0]['approved'] and not text_readiness(changed, library.directory)['text_ready']
        assert audio_status(changed, library.directory)['status'] == 'stale'
        assert path.read_bytes() == original
        assert external_preview(changed, library.directory)['revision'] == changed['revision']
        with pytest.raises(RevisionConflict):
            save_script(library, lesson['id'], lesson['segments'][0]['id'], lesson['revision'], 'Old edit.', '')
        again, _ = import_presentation(tmp_path / 'data', path)
        assert again['id'] == changed['id'] and again['segments'][0]['en'] == 'New narration.'
        assert library.get(changed['id'])['narration_variants']
    finally:
        library.close()


def test_blank_presentation_can_receive_script_and_be_prepared_without_design_repair(tmp_path):
    lesson, _ = import_presentation(tmp_path / 'data', source(tmp_path / 'Images.pptx'))
    library = Library(tmp_path / 'data')
    try:
        assert not text_readiness(lesson, library.directory)['text_ready']
        changed = save_script(library, lesson['id'], lesson['segments'][0]['id'], lesson['revision'], 'Xin chào.', '')
        changed['slide_quality'] = {'counts': {'FAIL': 2}}
        library._write(changed)
        reviewed = library.review_external_deck(changed['id'], changed['revision'])
        assert text_readiness(reviewed, library.directory)['text_ready']
        assert reviewed['prepared']['revision'] == reviewed['revision']
        row = library.list_lessons()[0]
        assert row['status'] == 'READY_TO_TEACH' and row['prepared'] and not row['quality_failures']
    finally:
        library.close()


def test_audio_zip_contains_selected_slide_languages_and_matching_audio_without_touching_source(tmp_path):
    path = source(tmp_path / 'Both.pptx', 'Energy', 'VI: Năng lượng được bảo toàn.\nEN: Energy is conserved.')
    before = path.read_bytes()
    lesson, _ = import_presentation(tmp_path / 'data', path)
    library = Library(tmp_path / 'data')
    try:
        lesson = complete(library, lesson)
        target = export_audio(lesson, library.directory, tmp_path / 'Audio.zip')
        with zipfile.ZipFile(target) as archive:
            assert set(archive.namelist()) == {'Slide-001-vi.wav', 'Slide-001-en.wav', 'loi-doc.txt', 'manifest.json'}
            assert json.loads(archive.read('manifest.json'))['mode'] == 'both'
            assert 'Năng lượng' in archive.read('loi-doc.txt').decode('utf-8')
        assert path.read_bytes() == before
        with pytest.raises(ValueError, match='ngoài thư viện'):
            export_audio(lesson, library.directory, library.directory / 'Bad.zip')
        key = lesson['narration']['keys']['en'][0]
        (library.directory / 'audio' / (key + '.wav')).unlink()
        target.write_bytes(b'existing')
        with pytest.raises(ValueError, match='Tạo đủ'):
            export_audio(lesson, library.directory, target)
        assert target.read_bytes() == b'existing'
    finally:
        library.close()


def test_select_regular_library_lesson_prepares_a_matching_preview(tmp_path, monkeypatch):
    library = Library(tmp_path / 'data')
    lesson = library.create('Bài cũ', 'Tin học', 'THPT', '10', [('Slide 1', 'Máy tính.')])
    lesson = library.edit_segment(lesson['id'], lesson['segments'][0]['id'], 'Máy tính.', 'A computer.')
    library.close()
    monkeypatch.setattr('app.quick_conversion.build_preview', lambda selected, directory, profile: {'lesson_id': selected['id'], 'path': 'prepared.pptx', 'revision': selected['revision']})
    selected, preview = prepare_library_lesson(tmp_path / 'data', lesson['id'])
    assert selected['id'] == preview['lesson_id'] == lesson['id']
    assert preview['revision'] == lesson['revision'] and not selected.get('voice_only')


def test_failed_narration_exposes_cause_and_stops_repeated_failures(tmp_path, monkeypatch):
    from app.narration_preparation import NarrationJob, NarrationPreparation

    job = NarrationJob({'profile': {'units': [{'vi': f'Câu {i}.', 'en': ''} for i in range(20)]}},
                       {'vi': 'vieneu:Mai Anh', 'rate': 0}, tmp_path, None)
    monkeypatch.setattr('app.speech.synthesize_background', lambda *_: (_ for _ in ()).throw(RuntimeError('engine missing')))
    job.run()
    assert job.result['failed'] == 3 and job.result['last_error'] == 'engine missing'
    assert 'engine missing' in NarrationPreparation._completion_message(job.result)


def test_labelled_multiline_heading_and_formula_are_split_without_losing_continuations(tmp_path):
    from app.input_analysis import split_existing_pair

    text = 'VI: I. ĐỘNG NĂNG\nCÔNG THỨC VÀ ĐỒ THỊ\nW = 1/2 mv²\nEN: I. KINETIC ENERGY\nFORMULAS AND GRAPHS\nW = 1/2 mv²'
    pair = split_existing_pair(text)
    assert pair['vi'] == 'I. ĐỘNG NĂNG\nCÔNG THỨC VÀ ĐỒ THỊ\nW = 1/2 mv²'
    assert pair['en'] == 'I. KINETIC ENERGY\nFORMULAS AND GRAPHS\nW = 1/2 mv²'
    lesson, _ = import_presentation(tmp_path / 'data', source(tmp_path / 'Bilingual.pptx', text))
    assert lesson['segments'][0]['vi'] == pair['vi']
    assert lesson['segments'][0]['en'] == pair['en']


def test_old_converted_lesson_is_repaired_and_editable_without_changing_slide_text_or_source(tmp_path):
    from app.chatgpt_handoff import text_signature

    text = 'VI: Động năng của vật.\nEN: The kinetic energy of the object.'
    path = source(tmp_path / 'Old.pptx', text)
    before = path.read_bytes()
    lesson, _ = import_presentation(tmp_path / 'data', path)
    library = Library(tmp_path / 'data')
    try:
        lesson.pop('narration_classification_version', None)
        lesson['voice_only'] = False
        lesson['segments'][0].update(vi=text, en='')
        lesson['external_deck']['text_signature'] = text_signature(lesson)
        library._write(lesson)
        signature_before = text_signature(lesson)
        restored, preview = prepare_library_lesson(library.directory, lesson['id'])
        assert restored['id'] == lesson['id'] and not restored['voice_only']
        assert restored['segments'][0]['vi'] == text
        assert narration_segments(restored)[0]['en'] == 'The kinetic energy of the object.'
        assert text_signature(restored) == signature_before
        assert preview['revision'] == restored['revision']
        complete(library, restored)
        restored = library.get(lesson['id'])
        changed = save_script(library, restored['id'], restored['segments'][0]['id'],
                              restored['revision'], 'Động năng đã sửa.', 'Edited kinetic energy.')
        assert changed['segments'][0]['vi'] == text
        assert narration_segments(changed)[0]['en'] == 'Edited kinetic energy.'
        assert text_signature(changed) == signature_before
        assert audio_status(changed, library.directory)['status'] == 'stale'
        changed = complete(library, changed)
        target = export_audio(changed, library.directory, tmp_path / 'Edited.zip')
        with zipfile.ZipFile(target) as archive:
            assert 'Edited kinetic energy.' in archive.read('loi-doc.txt').decode('utf-8')
        assert path.read_bytes() == before
        assert external_preview(changed, library.directory)['path'] == preview['path']
        again, _ = prepare_library_lesson(library.directory, changed['id'])
        assert narration_segments(again)[0]['en'] == 'Edited kinetic energy.'
    finally:
        library.close()
