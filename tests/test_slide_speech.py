import hashlib
import json

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from app.chatgpt_handoff import inspect_returned_deck, speaker_pair
from app.library import Library
from app.slide_quality import inspect_quality
from app.slide_speech import deck_speech, inspect_slide, speech_chunks, visible_segments


def textbox(slide, text, x=1, y=1, name=""):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(4), Inches(.6))
    shape.text = text
    if name:
        shape.name = name
    return shape


def animation(slide, shape, kind="entr", node_type="clickEffect", paragraph=False):
    timing = OxmlElement("p:timing")
    tree = OxmlElement("p:tnLst")
    main = OxmlElement("p:cTn")
    main.set("nodeType", "mainSeq")
    effect = OxmlElement("p:cTn")
    effect.set("nodeType", node_type)
    effect.set("presetClass", kind)
    behavior = OxmlElement("p:cBhvr")
    duration = OxmlElement("p:cTn")
    duration.set("dur", "500")
    target = OxmlElement("p:tgtEl")
    shape_target = OxmlElement("p:spTgt")
    shape_target.set("spid", str(shape.shape_id))
    if paragraph:
        shape_target.append(OxmlElement("p:txEl"))
    target.append(shape_target)
    behavior.extend([duration, target])
    effect.append(behavior)
    main.append(effect)
    tree.append(main)
    timing.append(tree)
    slide._element.append(timing)


@pytest.mark.parametrize("mode", ["parallel_columns", "sentence_pairs", "integrated_keywords", "english_only"])
def test_reads_objects_not_title_only_notes_and_never_vietnamese(mode):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "Động năng", y=.5)
    textbox(slide, "Kinetic energy", y=.5, name="BCL_EN_title")
    textbox(slide, "The energy of an object depends on its mass and speed.", y=1.8)
    textbox(slide, "Năng lượng phụ thuộc khối lượng và vận tốc.", y=2.8)
    slide.notes_slide.notes_text_frame.text = "VI: Động năng\nEN: Kinetic energy"
    result = inspect_slide(slide, 1, mode, "Kinetic energy")
    assert [s["english_text"] for s in result["englishSpeechSegments"]] == [
        "Kinetic energy", "The energy of an object depends on its mass and speed."]
    assert "notes_mismatch" in result["issues"]
    assert result["speech_status"] == "CHECK"


def test_inline_keywords_table_and_group_native_sources_have_stable_ids():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "Vận tốc (velocity) và gia tốc (acceleration).")
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(.5))
    shape.text = "Discuss the result in groups."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Khối lượng"
    table.cell(0, 1).text = "Mass"
    first = inspect_slide(slide, 1, "integrated_keywords")
    assert {s["english_text"] for s in first["englishSpeechSegments"]} == {
        "velocity", "acceleration", "Discuss the result in groups.", "Mass"}
    assert first == inspect_slide(slide, 1, "integrated_keywords")
    assert any(s["source"]["cell"] == [0, 1] for s in first["englishSpeechSegments"])


def test_options_read_a_to_d_even_when_shape_creation_and_coordinates_differ():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i, text in enumerate(["D. Potential energy", "B. Kinetic energy", "A. Mechanical energy", "C. Thermal energy"]):
        textbox(slide, text, y=1+i)
    values = inspect_slide(slide, 1)
    assert [s["english_text"][0] for s in visible_segments(values)] == list("ABCD")


def test_hidden_animation_never_reads_until_click_has_settled():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "What is kinetic energy?")
    answer = textbox(slide, "The answer is kinetic energy.", y=2)
    animation(slide, answer)
    result = inspect_slide(slide, 1)
    assert len(visible_segments(result, 0, 2)) == 1
    assert len(visible_segments(result, 1, .1)) == 1
    assert len(visible_segments(result, 1, 2)) == 2
    assert len(visible_segments(result, None, 2)) == 1


@pytest.mark.parametrize("node,paragraph", [("afterEffect", False), ("clickEffect", True), ("withEffect", False)])
def test_unsupported_animation_is_not_assumed_visible(node, paragraph):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    answer = textbox(slide, "The answer is potential energy.")
    animation(slide, answer, node_type=node, paragraph=paragraph)
    result = inspect_slide(slide, 1)
    assert not visible_segments(result, 100, 100)
    assert "animation_needs_review" in result["issues"]


def test_repeated_text_is_kept_at_distinct_meaningful_sources_not_metadata_duplicates():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = textbox(slide, "The energy is conserved.")
    textbox(slide, "The energy is conserved.", y=2)
    metadata = {"version": 1, "segments": [{"shape_id": first.shape_id, "paragraph": 0,
                                         "text": first.text}] * 2}
    slide.notes_slide.notes_text_frame.text = "SPEECH: " + json.dumps(metadata)
    result = inspect_slide(slide, 1)
    assert len(result["englishSpeechSegments"]) == 2
    assert len({s["id"] for s in result["englishSpeechSegments"]}) == 2


def test_notes_cannot_invent_speech_and_metadata_never_leaks_to_vi_en():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = textbox(slide, "Kinetic energy")
    metadata = {"version": 1, "segments": [{"shape_id": shape.shape_id, "paragraph": 0,
                                         "text": "This explanation does not appear on the slide."}]}
    notes = "VI: Động năng\nEN: Kinetic energy\nSPEECH: " + json.dumps(metadata)
    slide.notes_slide.notes_text_frame.text = notes
    result = inspect_slide(slide, 1)
    assert [s["english_text"] for s in result["englishSpeechSegments"]] == ["Kinetic energy"]
    assert "speech_metadata_mismatch" in result["issues"]
    assert speaker_pair(notes) == {"vi": "Động năng", "en": "Kinetic energy"}


@pytest.mark.parametrize("text", ["CHAPTER s 1", "19wx568", "CHECK: Kinetic energy", "https://example.com/technical"])
def test_garbage_metadata_and_url_never_spoken(text):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, text, name="BCL_EN")
    assert not visible_segments(inspect_slide(slide, 1))


def test_image_content_needs_review_and_foreground_image_detects_overlap(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "The energy is conserved.")
    image = tmp_path / "diagram.png"
    Image.new("RGB", (300, 60), "white").save(image)
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(4), Inches(.6))
    path = tmp_path / "deck.pptx"
    deck.save(path)
    data = deck_speech(path, deck)
    assert "image_or_formula_needs_review" in data["slides"][0]["issues"]
    quality = inspect_quality(deck, data)
    assert "possible_overlap" in [i["code"] for i in quality["slides"][0]["issues"]]
    assert quality["status"] == "CHECK"


def test_garbage_and_quiz_contradiction_block_classroom_confirmation(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "A. Kinetic energy ✓\nB. Potential energy")
    textbox(slide, "19wx568", y=3)
    slide.notes_slide.notes_text_frame.text = 'VI: Câu hỏi\nEN: What is the answer?\nQUIZ: {"correct":"B"}'
    path = tmp_path / "deck.pptx"
    deck.save(path)
    inspection = inspect_returned_deck(path)
    codes = [issue["code"] for issue in inspection["slide_quality"]["slides"][0]["issues"]]
    assert "quiz_mismatch" in codes and "technical_text" in codes
    library = Library(tmp_path / "library")
    try:
        source = library.store_source(path)
        config = {"title": "Test", "subject": "Vật lý", "education_level": "THPT", "grade": "11",
                  "conversion_format": "parallel_columns", "style": "source", "preset": "standard"}
        lesson = library.create_external_lesson(config, source, inspection)
        assert library.get(lesson["id"])["slide_speech"] == inspection["slide_speech"]
        with pytest.raises(ValueError, match="FAIL"):
            library.review_external_deck(lesson["id"], lesson["revision"])
    finally:
        library.close()


def test_powerpoint_soft_line_break_is_not_technical_garbage(tmp_path):
    from app.slide_speech import speech_problem

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "Bước 2: Nhập hàm số\vQuan sát đồ thị vận tốc và gia tốc.")
    path = tmp_path / "soft-break.pptx"
    deck.save(path)
    inspection = inspect_returned_deck(path, conversion_format="parallel_columns")
    assert not any(i["code"] == "technical_text" for i in inspection["slide_quality"]["slides"][0]["issues"])
    assert speech_problem("19wx568") == "technical_text"
    assert speech_problem("Energy\x0cwrong") == "technical_text"


def test_notes_and_a_long_english_title_cannot_replace_visible_body_translation(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "Velocity and acceleration in simple harmonic motion", y=.1, name="BCL_EN_title")
    textbox(slide, "Vận tốc thay đổi theo thời gian và bằng đạo hàm của li độ. " * 5, y=2)
    slide.notes_slide.notes_text_frame.text = "VI: Quan sát vận tốc.\nEN: " + "Velocity is the time derivative of displacement. " * 5
    path = tmp_path / "title-only-translation.pptx"
    deck.save(path)
    report = inspect_returned_deck(path, conversion_format="parallel_columns")["slide_quality"]
    assert any(i["code"] == "english_body_missing" and i["severity"] == "FAIL" for i in report["slides"][0]["issues"])


def test_parallel_columns_reject_stacked_body_but_accept_separate_left_and_right(tmp_path):
    vi = "Vận tốc là đạo hàm của li độ theo thời gian. Vận tốc thay đổi khi vật dao động."
    en = "Velocity is the time derivative of displacement. It changes as the object oscillates."
    for stacked in (True, False):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        first = textbox(slide, vi, x=.3, y=2)
        if stacked:
            first.text += "\n" + en
        else:
            textbox(slide, en, x=5.1, y=2)
        path = tmp_path / ("stacked.pptx" if stacked else "columns.pptx")
        deck.save(path)
        report = inspect_returned_deck(path, conversion_format="parallel_columns")["slide_quality"]
        assert any(i["code"] == "parallel_columns_missing" for i in report["slides"][0]["issues"]) == stacked


def test_language_columns_can_share_half_the_slide_beside_a_large_diagram(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    vi = textbox(slide, "Vận tốc bằng đạo hàm của li độ theo thời gian, thay đổi khi vật dao động.", x=.2, y=2)
    vi.width = Inches(2)
    en = textbox(slide, "Velocity is the time derivative of displacement and changes during oscillation.", x=2.3, y=2)
    en.width = Inches(2)
    path = tmp_path / "columns-beside-diagram.pptx"
    deck.save(path)
    report = inspect_returned_deck(path, conversion_format="parallel_columns")["slide_quality"]
    assert not any(i["code"] == "parallel_columns_missing" for i in report["slides"][0]["issues"])


def test_english_only_remaining_vietnamese_is_not_pass_and_hash_changes_with_edits(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = textbox(slide, "Năng lượng")
    path = tmp_path / "deck.pptx"
    deck.save(path)
    data = deck_speech(path, deck)
    assert not visible_segments(data["slides"][0])
    assert any(i["code"] == "vietnamese_remaining" for i in inspect_quality(deck, data, "english_only")["slides"][0]["issues"])
    shape.text = "Kinetic energy"
    deck.save(path)
    assert hashlib.sha256(path.read_bytes()).hexdigest() != data["sha256"]


def test_chunking_keeps_every_word_without_cutting_normal_sentence():
    text = "The energy is conserved. " * 300
    chunks = speech_chunks(text)
    assert all(len(c) <= 1500 for c in chunks)
    assert " ".join(chunks) == text.strip()
    assert all(c.endswith(".") for c in chunks)
    with pytest.raises(ValueError):
        speech_chunks("a" * 1600)


def test_alternate_content_reads_modern_shape_once_and_does_not_skip_math():
    import copy

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = textbox(slide, "The energy is conserved.")
    element = shape._element
    tree = element.getparent()
    tree.remove(element)
    ns = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alternate = etree.Element("{"+ns+"}AlternateContent")
    choice = etree.SubElement(alternate, "{"+ns+"}Choice")
    choice.set("Requires", "a14")
    choice.append(element)
    fallback = etree.SubElement(alternate, "{"+ns+"}Fallback")
    fallback.append(copy.deepcopy(element))
    tree.append(alternate)
    math = etree.SubElement(element, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
    text = etree.SubElement(math, "{http://schemas.openxmlformats.org/officeDocument/2006/math}t")
    text.text = "Với vận tốc"
    data = inspect_slide(slide, 1)
    assert [s["english_text"] for s in data["englishSpeechSegments"] if s["status"] == "ready"] == ["The energy is conserved."]
    assert any(s["status"] == "needs_review" and not s["english_text"] for s in data["englishSpeechSegments"])
    assert "image_or_formula_needs_review" in data["issues"]


def test_hidden_group_cannot_leak_answer_through_visible_child_shape():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, "What is kinetic energy?")
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(.5))
    child.text = "The correct answer is potential energy."
    group._element.xpath('./p:nvGrpSpPr/p:cNvPr')[0].set("hidden", "true")
    result = inspect_slide(slide, 1)
    assert [s["english_text"] for s in visible_segments(result)] == ["What is kinetic energy?"]

def test_tagged_english_keeps_isolated_source_identifier_without_accepting_vietnamese():
    from app.slide_speech import english_text
    assert english_text('Predict what happens when pendulum Đ oscillates.',tagged=True)
    assert not english_text('Con lắc Đ dao động.',tagged=True)
    assert not english_text('The answer by Đức is shown below.',tagged=True)
    assert not english_text('Predict what happens when pendulum Đ oscillates.')
