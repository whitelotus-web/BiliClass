import hashlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.contrast_repair import repair


def test_measured_solid_fill_contrast_changes_only_color_and_keeps_math(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    label.text = 'Force is measured in newtons.'
    run = label.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.italic = True
    run.font.color.rgb = RGBColor(230, 230, 230)
    math_shape = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    math_shape.text_frame.paragraphs[0]._p.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'))
    source = tmp_path / 'source.pptx'
    deck.save(source)
    original = source.read_bytes()
    report = {'sha256': hashlib.sha256(original).hexdigest(), 'slides': [{'slide': 1, 'issues': [
        {'code': 'low_contrast', 'shape_id': shape.shape_id, 'background_rgb': 0xffffff, 'threshold': 4.5}
        for shape in (label, math_shape)]}]}
    result = repair(source, tmp_path, report)
    output = Presentation(result['path'])
    corrected = output.slides[0].shapes[0]
    assert corrected.text == label.text and corrected.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(0, 0, 0)
    assert corrected.text_frame.paragraphs[0].runs[0].font.size.pt == 24
    assert corrected.text_frame.paragraphs[0].runs[0].font.italic
    assert etree.tostring(output.slides[0].shapes[1]._element) == etree.tostring(math_shape._element)
    assert source.read_bytes() == original and len(result['changes']) == 1
    report['slides'][0]['issues'][0].pop('background_rgb')
    unchanged = repair(source, tmp_path, report)
    assert Path(unchanged['path']) == source and not unchanged['changes']
