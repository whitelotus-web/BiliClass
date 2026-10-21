import copy
import time
import wave
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QCoreApplication

from app import speech
from app.chatgpt_handoff import inspect_returned_deck
from app.lesson_narration import audio_status, make_plan
from app.library import Library, RevisionConflict
from app.ui import Bridge

VOICE = "vieneu:Mai Anh"
VOICES = {"voice": VOICE, "vi": VOICE, "en": VOICE, "rate": 0}


def settle_narration(app, bridge):
    deadline = time.monotonic() + 8
    while bridge.busy or bridge.narrationBusy or bridge._catalog.selected_narration.worker is not None:
        app.processEvents()
        assert time.monotonic() < deadline, bridge.message
        time.sleep(.005)
    app.processEvents()


def wav(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with wave.open(str(path), "wb") as writer:
        writer.setparams((1, 2, 24000, 0, "NONE", "not compressed"))
        writer.writeframes(b"\x01\x00" * 100)


@pytest.fixture
def saved(tmp_path):
    path = tmp_path / "lesson.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(7), Inches(2)
    ).text = "Năng lượng được bảo toàn.\nEnergy is conserved."
    slide.notes_slide.notes_text_frame.text = "VI: Năng lượng được bảo toàn.\nEN: Energy is conserved."
    deck.save(path)
    library = Library(tmp_path / "library")
    config = {
        "title": "Test",
        "subject": "Vật lý",
        "education_level": "THPT",
        "grade": "10",
        "conversion_format": "parallel_columns",
        "style": "source",
        "preset": "standard",
    }
    lesson = library.create_external_lesson(config, library.store_source(path), inspect_returned_deck(path))
    yield library, lesson, path
    library.close()


@pytest.mark.parametrize("mode,languages", [("en", {"en"}), ("vi", {"vi"}), ("both", {"vi", "en"})])
def test_selected_languages_only_and_metadata_does_not_approve_or_change_pptx(saved, mode, languages):
    library, lesson, path = saved
    original = path.read_bytes()
    plan, selected = make_plan(lesson, mode, VOICES)
    assert set(plan["keys"]) == languages
    assert {lang for lang in ("vi", "en") if selected[lang]} == languages
    for keys in plan["keys"].values():
        for key in keys:
            wav(library.directory / "audio" / (key + ".wav"))
    changed = library.update_narration(lesson["id"], plan)
    assert changed["revision"] == lesson["revision"] and changed["created_at"] == lesson["created_at"]
    assert not any(s["approved"] for s in changed["segments"])
    status = audio_status(changed, library.directory, VOICES)
    assert status["ready"] and set(status["languages"]) == languages
    assert library.list_lessons()[0]["narration"]["mode"] == mode
    assert len(library.list_lessons()) == 1
    assert path.read_bytes() == original


def test_new_conversion_has_no_voice_even_if_other_lessons_cached_identical_text(saved):
    library, lesson, _ = saved
    plan, _ = make_plan(lesson, "both", VOICES)
    for keys in plan["keys"].values():
        for key in keys:
            wav(library.directory / "audio" / (key + ".wav"))
    assert audio_status(lesson, library.directory, VOICES)["status"] == "none"
    assert library.list_lessons()[0]["narration"]["status"] == "none"


def test_status_checks_each_distinct_wav_once_and_can_cancel(saved, monkeypatch):
    from threading import Event

    from app import lesson_narration

    library, lesson, _ = saved
    lesson["narration"], _ = make_plan(lesson, "both", VOICES)
    expected = {key for keys in lesson["narration"]["keys"].values() for key in keys}
    for key in expected:
        wav(library.directory / "audio" / (key + ".wav"))
    original = lesson_narration.valid_key
    calls = []

    def counted(directory, key):
        calls.append(key)
        return original(directory, key)

    monkeypatch.setattr(lesson_narration, "valid_key", counted)
    assert audio_status(lesson, library.directory)["ready"]
    assert set(calls) == expected and len(calls) == len(expected)
    cancel = Event()
    cancel.set()
    calls.clear()
    with pytest.raises(InterruptedError):
        audio_status(lesson, library.directory, cancel=cancel)
    assert calls == []


def test_legacy_cached_voice_is_recognized_without_loading_or_starting_model(saved, monkeypatch):
    library, lesson, _ = saved
    lesson.pop("narration")
    plan, _ = make_plan(lesson, "en", VOICES)
    for key in plan["keys"]["en"]:
        wav(library.directory / "audio" / (key + ".wav"))
    monkeypatch.setattr(
        speech, "_vieneu_engine", lambda *_: (_ for _ in ()).throw(AssertionError("No model"))
    )
    status = audio_status(lesson, library.directory, VOICES)
    assert status["ready"] and status["mode"] == "en" and status["complete"] == status["total"]


def test_missing_or_corrupt_wavs_and_unrelated_keys_cannot_claim_ready(saved):
    library, lesson, _ = saved
    plan, _ = make_plan(lesson, "en", VOICES)
    lesson["narration"] = plan
    path = library.directory / "audio" / (plan["keys"]["en"][0] + ".wav")
    wav(path)
    assert audio_status(lesson, library.directory)["ready"]
    path.write_bytes(b"broken")
    assert not audio_status(lesson, library.directory)["ready"]
    forged = "f" * 64
    wav(path.parent / (forged + ".wav"))
    lesson["narration"]["keys"]["en"] = [forged]
    assert audio_status(lesson, library.directory)["status"] == "stale"


def test_shared_variants_reuse_wav_checks_without_hiding_later_file_loss(saved, monkeypatch):
    from app import lesson_narration

    library, lesson, _ = saved
    for mode in ("en", "vi", "both"):
        plan, _ = make_plan(lesson, mode, VOICES)
        for keys in plan["keys"].values():
            for key in keys:
                wav(library.directory / "audio" / (key + ".wav"))
        library.update_narration(lesson["id"], plan)
    lesson = library.get(lesson["id"])
    checks = []
    original = lesson_narration.valid_key

    def counted(directory, key):
        checks.append(key)
        return original(directory, key)

    monkeypatch.setattr(lesson_narration, "valid_key", counted)
    status = lesson_narration.library_audio_status(lesson, library.directory)
    assert all(row["ready"] for row in status["variants"])
    assert len(checks) == len(set(checks))
    (library.directory / "audio" / (checks[0] + ".wav")).unlink()
    changed = lesson_narration.library_audio_status(lesson, library.directory)
    assert any(not row["ready"] for row in changed["variants"])


def test_late_audio_metadata_cannot_overwrite_edited_lesson(saved):
    library, lesson, _ = saved
    plan, _ = make_plan(lesson, "en", VOICES)
    changed = copy.deepcopy(lesson)
    changed["segments"][0]["en"] = "Changed content."
    library._write(changed)
    with pytest.raises(RevisionConflict):
        library.update_narration(lesson["id"], plan)
    assert library.get(lesson["id"])["narration"]["mode"] == "none"


def test_multiple_voice_variants_survive_selection_reopen_and_no_voice_choice(saved):
    from app.lesson_narration import records, variant_id

    library, lesson, path = saved
    original = path.read_bytes()
    identities = []
    for mode, voice, rate in [("en", VOICE, 0), ("vi", "vieneu:Hải Đăng", 1), ("both", VOICE, 0)]:
        plan, _ = make_plan(
            lesson,
            mode,
            {"voice": voice, "rate": rate},
            playback="both_controls" if mode == "both" else "separate",
        )
        for keys in plan["keys"].values():
            for key in keys:
                wav(library.directory / "audio" / (key + ".wav"))
        library.update_narration(lesson["id"], plan)
        identities.append(variant_id(plan))
    current = library.get(lesson["id"])
    assert len(records(current)) == 3
    for identity, mode in zip(identities, ["en", "vi", "both"]):
        current = library.select_narration_variant(lesson["id"], identity)
        status = audio_status(current, library.directory)
        assert status["ready"] and status["mode"] == mode and status["id"] == identity
        assert current["revision"] == lesson["revision"]
    current = library.select_narration_variant(lesson["id"], "")
    assert audio_status(current, library.directory)["status"] == "none"
    assert len(records(current)) == 3
    badge = library.list_lessons()[0]["narration"]
    assert len(badge["variants"]) == 3 and all(v["ready"] for v in badge["variants"])
    assert badge["label"] == "Đã có giọng đọc: Anh · Việt · Song ngữ"
    assert library.setting("voice_options", None) is None
    assert not any(s["approved"] for s in current["segments"])
    assert current["created_at"] == lesson["created_at"] and path.read_bytes() == original


def test_old_single_voice_is_kept_when_another_voice_is_created(saved):
    from app.lesson_narration import records, variant_id

    library, lesson, _ = saved
    old, _ = make_plan(lesson, "en", VOICES)
    lesson["narration"] = old  # Simulate a genuine pre-variant library record.
    lesson.pop("narration_variants", None)
    library._write(lesson)
    new, _ = make_plan(lesson, "vi", {"voice": "vieneu:Thùy Dung", "rate": 0})
    current = library.update_narration(lesson["id"], new)
    assert set(records(current)) == {variant_id(old), variant_id(new)}
    assert library.select_narration_variant(lesson["id"], variant_id(old))["narration"] == old


def test_different_voice_generation_and_teaching_use_selected_variant_not_global_setting(saved, monkeypatch):
    from app.lesson_narration import variant_id

    app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = saved
    other = "vieneu:Hải Đăng"
    monkeypatch.setattr(
        speech,
        "list_voices",
        lambda: [dict(v, provider="vieneu", languages=["vi", "en"]) for v in speech.VIENEU_VOICES],
    )
    calls = []

    def render(text, voice, rate, directory, cancel):
        calls.append((text, voice, rate))
        target = Path(directory) / (speech.audio_key(text, voice, rate) + ".wav")
        wav(target)
        return {"path": str(target), "cached": False}

    monkeypatch.setattr(speech, "synthesize_background", render)
    bridge = Bridge(library)
    try:
        bridge._lesson = lesson
        for mode, voice in [("en", VOICE), ("vi", other), ("both", VOICE)]:
            bridge.createLessonNarrationWithVoice(mode, voice, 0)
            settle_narration(app, bridge)
            assert not bridge.narrationBusy and bridge.lessonNarration["ready"]
        assert len(calls) == 3  # Both reuses the previously generated English.
        assert len(library.get(lesson["id"])["narration_variants"]) == 3
        assert library.setting("voice_options", None) is None
        queues = []
        bridge.english_reader.start = lambda texts, voice, rate: queues.append((texts, voice, rate))
        bridge.speakSegment("both")
        assert queues == [] and "loa Anh hoặc loa Việt" in bridge.message
        vi, _ = make_plan(lesson, "vi", {"voice": other, "rate": 0})
        bridge.selectNarrationVariant(variant_id(vi))
        settle_narration(app, bridge)
        assert bridge._lesson_voice_settings()["vi"] == other and not bridge._lesson_voice_settings()["en"]
        bridge.speakSegment("both")
        assert queues == [] and "loa Anh hoặc loa Việt" in bridge.message
    finally:
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_incomplete_variant_cannot_silently_start_teaching(saved):
    library, lesson, _ = saved
    _app = QCoreApplication.instance() or QCoreApplication([])
    bridge = Bridge(library)
    try:
        plan, _ = make_plan(lesson, "en", VOICES)
        bridge._lesson = library.update_narration(lesson["id"], plan)
        bridge.useConvertedLesson()
        assert "chưa hoàn tất" in bridge.message and bridge.powerpoint_worker is None
    finally:
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_accessory_subject_is_independent_of_the_lesson_and_visibility_is_always_enabled(saved):
    import time

    _app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = saved
    bridge = Bridge(library)
    def saved_preferences():
        deadline = time.monotonic() + 5
        while bridge.mascotSaving and time.monotonic() < deadline:
            _app.processEvents()
            time.sleep(.01)
        assert not bridge.mascotSaving
    try:
        bridge._lesson = lesson
        library.set_setting('mascot_options', {'subject_mode': 'fixed', 'subject': 'Mĩ thuật',
                                              'visible': False, 'show_explanation': False, 'show_quiz': False, 'reduced_motion': True})
        assert bridge.mascotSubject == 'Vật lý'
        assert bridge.mascotAccessorySubject == 'Mĩ thuật'
        assert bridge.mascotSettings['reduced_motion'] is False
        assert all(bridge.mascotSettings[name] for name in ('visible', 'show_explanation', 'show_quiz'))
        bridge.saveMascotPreferences("Lumi", False, 110, "right", True, True, False, False, "fixed", "Mĩ thuật")
        saved_preferences()
        assert bridge.mascotSubject == "Vật lý"
        assert bridge.mascotAccessorySubject == 'Mĩ thuật'
        assert all(library.setting('mascot_options', {})[name] for name in ('visible', 'show_explanation', 'show_quiz'))
        assert library.setting('mascot_options', {})['reduced_motion'] is False
        bridge._lesson = {**lesson, "subject": "Tin học"}
        assert bridge.mascotSubject == "Tin học"
        assert bridge.mascotAccessorySubject == 'Mĩ thuật'
        bridge.saveMascotPreferences("Lumi", True, 110, "right", False, True, True, True, "auto", "Mĩ thuật")
        saved_preferences()
        assert bridge.mascotSubject == "Tin học"
        assert bridge.mascotSettings['subject_mode'] == 'auto'
        assert bridge.mascotAccessorySubject == 'Tin học'
        assert library.get(lesson["id"])["subject"] == "Vật lý"
    finally:
        bridge._companion_preferences.shutdown()
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_companion_cannot_read_without_live_matching_presentation_and_stops_on_close(saved, monkeypatch):
    from types import SimpleNamespace

    _app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, path = saved
    bridge = Bridge(library)
    calls, stopped = [], []
    monkeypatch.setattr(bridge, "speakSegment", calls.append)
    monkeypatch.setattr(bridge, "stopSpeech", lambda: stopped.append(True))
    try:
        bridge._lesson = lesson
        bridge.companionRead("en")
        assert calls == [] and not bridge.companionCanRead
        bridge.powerpoint_worker = SimpleNamespace(path=str(path))
        bridge._powerpoint_lesson_id = lesson["id"]
        bridge._powerpoint_state = {"active": True, "slide": 1, "total": 1}
        assert bridge.companionCanRead
        bridge.companionRead("en")
        bridge.replayCompanion()
        assert calls == ["en", "en"]
        bridge._powerpoint_changed({"active": False, "slide": 1, "total": 1})
        bridge.companionRead("vi")
        assert calls == ["en", "en"] and stopped and not bridge.companionCanRead
    finally:
        bridge.powerpoint_worker = None
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_companion_vietnamese_skips_hidden_partner_and_changed_file(saved, monkeypatch):
    from types import SimpleNamespace

    _app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = saved
    bridge = Bridge(library)
    path = library.directory / "sources" / lesson["source"]["file"]
    bridge._lesson = lesson
    bridge.powerpoint_worker = SimpleNamespace(path=str(path))
    bridge._powerpoint_lesson_id = lesson["id"]
    bridge._powerpoint_state = {"active": True, "slide": 1, "total": 1, "click_index": 0}
    bridge._presentation_stamp = (path.stat().st_size, path.stat().st_mtime_ns)
    page = lesson["slide_speech"]["slides"][0]
    for item in page["englishSpeechSegments"]:
        item["visibility"] = {"kind": "click", "show_after": 1, "hide_after": None, "settle_seconds": 0}
    calls = []
    monkeypatch.setattr(bridge, "_lesson_voice_settings", lambda: VOICES)
    monkeypatch.setattr(bridge.english_reader, "start", lambda *args: calls.append(args))
    try:
        bridge.companionRead("vi")
        assert not calls
        bridge._powerpoint_state["click_index"] = 1
        bridge.companionRead("vi")
        assert calls and calls[0][0] == ["Năng lượng được bảo toàn."]
        bridge._presentation_stamp = (0, 0)
        bridge.companionRead("vi")
        assert len(calls) == 1
    finally:
        bridge.powerpoint_worker = None
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_companion_assistance_refuses_missing_or_wrong_presentation(saved, monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setenv('BILICLASS_EXPERIMENTAL_TEACHING', '1')
    _app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, path = saved
    bridge = Bridge(library)
    bridge._lesson = lesson
    calls = []
    monkeypatch.setattr(bridge.teaching, "ask", lambda *args: calls.append(args))
    try:
        bridge.companionAsk("explanation", "en")
        assert not calls
        bridge.powerpoint_worker = SimpleNamespace(path=str(path))
        bridge._powerpoint_lesson_id = "another-lesson"
        bridge._powerpoint_state = {"active": True, "slide": 1, "total": 2}
        bridge.companionAsk("explanation", "en")
        assert not calls
        bridge._powerpoint_lesson_id = lesson["id"]
        bridge._powerpoint_state["slide"] = 2
        bridge.companionAsk("explanation", "en")
        assert not calls
        bridge._powerpoint_state["slide"] = 1
        bridge.companionAsk("explanation", "en")
        assert calls == [("explanation", "en")]
    finally:
        bridge.powerpoint_worker = None
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_creation_choices_lock_only_the_exact_completed_product(saved):
    from app.lesson_narration import creation_choices, variant_statuses

    library, lesson, _ = saved
    plan, _ = make_plan(lesson, "en", VOICES)
    for key in plan["keys"]["en"]:
        wav(library.directory / "audio" / (key + ".wav"))
    lesson = library.update_narration(lesson["id"], plan)
    checked = variant_statuses(lesson, library.directory)
    choices = creation_choices(lesson, VOICE, 0, checked)
    assert choices["en"]["ready"] and not choices["en"]["canCreate"]
    assert choices["vi"]["canCreate"] and choices["both"]["canCreate"]
    assert creation_choices(lesson, VOICE, 1, checked)["en"]["canCreate"]
    assert creation_choices(lesson, "vieneu:Hải Đăng", 0, checked)["en"]["canCreate"]
    damaged = [{**row, "ready": False, "status": "partial"} for row in checked]
    assert creation_choices(lesson, VOICE, 0, damaged)["en"]["phase"] == "partial"
    assert creation_choices(lesson, VOICE, 0)["en"]["phase"] == "checking"


def test_duplicate_creation_never_restarts_worker_rewrites_record_or_resets_clock(saved, monkeypatch):
    _app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = saved
    monkeypatch.setattr(
        speech,
        "list_voices",
        lambda: [dict(v, provider="vieneu", languages=["vi", "en"]) for v in speech.VIENEU_VOICES],
    )
    plan, _ = make_plan(lesson, "en", VOICES)
    for key in plan["keys"]["en"]:
        wav(library.directory / "audio" / (key + ".wav"))
    lesson = library.update_narration(lesson["id"], plan)
    bridge = Bridge(library)
    try:
        bridge._lesson = lesson
        bridge.narration_preparation._elapsed = 123
        before = library.get(lesson["id"])
        changes = library.db.total_changes
        monkeypatch.setattr(
            bridge.narration_preparation, "start", lambda *_: pytest.fail("Already completed")
        )
        bridge.createLessonNarrationWithVoice("en", VOICE, 0)
        settle_narration(_app, bridge)
        assert library.get(lesson["id"]) == before and library.db.total_changes == changes
        assert bridge.narrationClock.elapsedSeconds == 123 and not bridge.narrationBusy
        assert "đã sẵn sàng" in bridge.message
    finally:
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_missing_wav_unlocks_existing_variant_and_only_missing_work_is_recreated(saved, monkeypatch):
    app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = saved
    monkeypatch.setattr(
        speech,
        "list_voices",
        lambda: [dict(v, provider="vieneu", languages=["vi", "en"]) for v in speech.VIENEU_VOICES],
    )
    plan, _ = make_plan(lesson, "en", VOICES)
    lesson = library.update_narration(lesson["id"], plan)
    calls = []

    def render(text, voice, rate, directory, cancel):
        calls.append(text)
        path = Path(directory) / (speech.audio_key(text, voice, rate) + ".wav")
        wav(path)
        return {"path": str(path), "cached": False}

    monkeypatch.setattr(speech, "synthesize_background", render)
    bridge = Bridge(library)
    try:
        bridge._lesson = lesson
        bridge.changed.emit()
        settle_narration(app, bridge)
        choice = bridge.narrationCreationChoices(bridge.narrationVariants, VOICE, 0)["en"]
        assert choice["canCreate"] and not choice["ready"]
        bridge.createLessonNarrationWithVoice("en", VOICE, 0)
        settle_narration(app, bridge)
        assert not bridge.narrationBusy and bridge.lessonNarration["ready"]
        assert calls == ["Energy is conserved."]
        choice = bridge.narrationCreationChoices(bridge.narrationVariants, VOICE, 0)["en"]
        assert choice["ready"] and not choice["canCreate"]
    finally:
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()


def test_explicit_local_generation_cache_reuse_and_reopen_do_not_send_web_request(saved, monkeypatch):
    app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, path = saved
    calls = []
    monkeypatch.setattr(
        speech, "list_voices", lambda: [{"id": VOICE, "provider": "vieneu", "languages": ["vi", "en"]}]
    )
    monkeypatch.setattr(
        "app.browser_automation.convert",
        lambda *a, **kw: (_ for _ in ()).throw(AssertionError("No web request")),
    )

    def render(text, voice, rate, directory, cancel):
        calls.append(text)
        target = Path(directory) / (speech.audio_key(text, voice, rate) + ".wav")
        wav(target)
        return {"path": str(target), "cached": False}

    monkeypatch.setattr(speech, "synthesize_background", render)
    bridge = Bridge(library)
    try:
        bridge._lesson = lesson
        bridge.changed.emit()
        settle_narration(app, bridge)
        assert bridge.lessonNarration["status"] == "none" and not calls
        bridge.createLessonNarration("en")
        settle_narration(app, bridge)
        assert not bridge.narrationBusy and bridge.lessonNarration["ready"]
        assert calls == ["Energy is conserved."]
        assert bridge._lesson_voice_settings()["vi"] == ""
        bridge.createLessonNarration("en")
        settle_narration(app, bridge)
        assert not bridge.narrationBusy and len(calls) == 1
        assert library.get(lesson["id"])["revision"] == lesson["revision"]
        assert path.read_bytes() == (library.directory / "sources" / lesson["source"]["file"]).read_bytes()
        bridge.speakSegment("vi")
        assert len(calls) == 1 and "chưa chọn giọng" in bridge.message
    finally:
        bridge.narration_preparation.shutdown()
        bridge.english_reader.shutdown()
        bridge.browserAI.shutdown()
