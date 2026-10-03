import copy
import json
import zipfile
from threading import Event

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.ai_lesson import convert_request, lesson_from_result, normalize_source, validate_result
from app.chatgpt_auth import PlanError
from app.chatgpt_handoff import prepare_request
from app.content import assistant_response
from app.level_conversion import narration_text
from app.library import Library
from app.pack import export_pack, import_pack
from app.presentation_policy import presentation_content
from app.quick_conversion import build_preview, review_snapshot, verify_preview


class EchoProvider:
    """Controlled AI fixture; no network, credentials or factual generation."""
    def __init__(self):
        self.calls, self.fail_at = [], 0

    def models(self, _account):
        return [{"slug": "test-model", "display_name": "Fixture", "input_modalities": ["text", "image"]}]

    def generate(self, _account, _model, _instructions, content, _schema, _cancel, _progress):
        context = json.loads(content[0]["text"])
        self.calls.append(content)
        if self.fail_at == len(self.calls):
            raise PlanError("Interrupted fixture", "stream_interrupted")
        blocks = []
        for unit in context["sources"]:
            text = unit["source_text"] or "Có một giá trị."
            vi, en = unit["vi"], unit["en"]
            if not unit["existing_pair"]:
                vi, en = (text, "There is one value.") if unit["language"] != "en" else ("Có một giá trị.", text)
            blocks.append({"source_ref": unit["source_ref"], "source_text": text,
                           "language": unit["language"] if unit["language"] in {"en", "vi"} else "vi",
                           "vi": vi, "en": en, "kind": "concept", "issues": [],
                           "support": [{"kind": "easy_en", "vi": "", "en": "One value.", "source_refs": [unit["source_ref"]]},
                                       {"kind": "vocabulary", "vi": "giá trị", "en": "value", "source_refs": [unit["source_ref"]]}]
                           if "giá trị" in vi else []})
        return {"schema_version": 1, "source_sha256": context["source_sha256"], "level": context["level"], "blocks": blocks}


@pytest.fixture
def library(tmp_path):
    value = Library(tmp_path / "library")
    yield value
    value.close()


def request_for(library, source=None, text="Có một giá trị.", **options):
    config = {"title": "Bài thử", "subject": "Toán", "education_level": "THPT", "grade": "10", "level": 2,
              "layout": "line_pair", "preset": "standard", "style": "source" if source and source.suffix == ".pptx" else "template",
              "mode": "level", "provider": "chatgpt_plan", **options}
    return prepare_request(library.directory, config, source, "" if source else text)


def execute(library, request, provider=None, **options):
    return convert_request(request, library.directory, provider or EchoProvider(), "account-one", Event(), lambda _: None, **options)


def test_draft_preview_whole_review_and_level_support(library):
    request = request_for(library)
    lesson = lesson_from_result(library, execute(library, request))
    assert not lesson["segments"][0]["approved"]
    assert not lesson["segments"][0]["support"][0]["approved"]
    preview = build_preview(lesson, library.directory)
    assert preview["segment_map"] and verify_preview(lesson, preview, library.directory).is_file()
    assert not library.get(lesson["id"])["segments"][0]["approved"]
    checked = library.review_lesson(lesson["id"], lesson["revision"])
    segment = checked["segments"][0]
    assert segment["approved"] and all(item["approved"] for item in segment["support"])
    assert assistant_response(segment, "easy_en", "en", 2)["available"]
    assert presentation_content(segment, 0, "split_view")["en"] == "giá trị — value"
    assert presentation_content(segment, 2, "split_view")["en"] == "One value."
    assert presentation_content(segment, 4, "english_rescue")["en"] == "There is one value."
    assert narration_text(segment, "en", 0) == "value"
    assert narration_text(segment, "en", 2) == "One value."
    assert narration_text(segment, "en", 4) == "There is one value."
    assert narration_text(segment, "vi", 4) == "Có một giá trị."


def test_native_pptx_keeps_media_and_source_objects(library, tmp_path, monkeypatch):
    import app.powerpoint_review
    monkeypatch.setattr(app.powerpoint_review, "render_slide", lambda *_: (_ for _ in ()).throw(ValueError("no Office")))
    picture = tmp_path / "figure.png"
    Image.new("RGB", (300, 200), "#93b9e4").save(picture)
    source = tmp_path / "original.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(7), Inches(.8)).text = "Có một giá trị."
    slide.shapes.add_picture(str(picture), Inches(.5), Inches(1.5), width=Inches(3))
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(7), Inches(.8)).text = "VI: Có một giá trị.\nEN: There is one value."
    deck.save(source)
    original = source.read_bytes()
    request = request_for(library, source)
    lesson = lesson_from_result(library, execute(library, request))
    assert lesson["segments"][1]["existing_pair"]
    preview = build_preview(lesson, library.directory)
    assert source.read_bytes() == original
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(preview["path"]) as after:
        for name in before.namelist():
            if name.startswith("ppt/media/") or name.startswith("ppt/slideMasters/") or name.startswith("ppt/theme/"):
                assert before.read(name) == after.read(name)
        # A slide already bilingual at L2 remains byte-for-byte unchanged.
        assert before.read("ppt/slides/slide2.xml") == after.read("ppt/slides/slide2.xml")
    assert "Có một giá trị." in "\n".join(shape.text for slide in Presentation(preview["path"]).slides for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize("mode,level", [("level", 0), ("level", 2), ("level", 4), ("paired", 2)])
def test_image_only_pptx_retains_picture(library, tmp_path, monkeypatch, mode, level):
    import app.powerpoint_review
    monkeypatch.setattr(app.powerpoint_review, "render_slide", lambda *_: (_ for _ in ()).throw(ValueError("no Office")))
    image = tmp_path / "book.png"
    Image.new("RGB", (300, 200), "white").save(image)
    source = tmp_path / "scan-slide.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(source)
    lesson = lesson_from_result(library, execute(library, request_for(library, source, mode=mode, level=level)))
    assert lesson["segments"][0]["source_image_only"]
    preview = build_preview(lesson, library.directory)
    assert preview["total"] >= 2 and all(index == 1 for index in preview["slide_map"])
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(preview["path"]) as after:
        assert before.read("ppt/media/image1.png") == after.read("ppt/media/image1.png")
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")


def test_long_native_object_stays_mapped(library, tmp_path):
    source = tmp_path / "long.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = "Có một giá trị. " * 900
    slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(5)).text = text
    deck.save(source)
    lesson = lesson_from_result(library, execute(library, request_for(library, source)))
    assert len(lesson["segments"]) == 2
    assert all(segment["source_text"] == text.strip() for segment in lesson["segments"])
    assert build_preview(lesson, library.directory)["path"]


def test_scanned_image_pack_remains_draft_and_hash_checked(library, tmp_path):
    source = tmp_path / "book.png"
    Image.new("RGB", (300, 200), "white").save(source)
    provider = EchoProvider()
    lesson = lesson_from_result(library, execute(library, request_for(library, source), provider))
    assert any(item["type"] == "input_image" for item in provider.calls[0])
    assert lesson["segments"][0]["kind"] == "visual"
    assert build_preview(lesson, library.directory)["path"]
    packed = export_pack(library, lesson["id"], tmp_path / "share.biliclass")
    recipient = Library(tmp_path / "recipient")
    try:
        imported = import_pack(recipient, packed)
        assert not imported["segments"][0]["approved"]
        assert not imported["segments"][0]["support"][0]["approved"]
        assert build_preview(imported, recipient.directory)["path"]
        assert recipient.review_lesson(imported["id"], imported["revision"])["segments"][0]["support"][0]["approved"]
        image = recipient.directory / "assets" / imported["segments"][0]["source_image"]
        image.write_bytes(b"changed")
        with pytest.raises(ValueError, match="Hình nguồn"):
            build_preview(imported, recipient.directory)
    finally:
        recipient.close()


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_word_and_scanned_pdf_keep_visual_sources(library, tmp_path, suffix):
    image_path = tmp_path / "page.png"
    Image.new("RGB", (500, 700), "white").save(image_path)
    source = tmp_path / ("document" + suffix)
    if suffix == ".pdf":
        with Image.open(image_path) as image:
            image.save(source, "PDF")
    else:
        from docx import Document

        document = Document()
        document.add_paragraph("Có một giá trị.")
        document.add_picture(str(image_path))
        document.save(source)
    original = source.read_bytes()
    provider = EchoProvider()
    result = execute(library, request_for(library, source), provider)
    assert result["manifest"]["images"]
    assert any(item["type"] == "input_image" for content in provider.calls for item in content)
    lesson = lesson_from_result(library, result)
    assert any(segment.get("source_image") for segment in lesson["segments"])
    assert build_preview(lesson, library.directory)["segment_map"]
    assert source.read_bytes() == original


def test_unsupported_source_requires_teacher_correction(library):
    result = execute(library, request_for(library))
    result["blocks"][0]["issues"] = ["SOURCE_MISSING"]
    lesson = lesson_from_result(library, result)
    with pytest.raises(ValueError, match="đọc rõ|kiểm tra"):
        review_snapshot(lesson, library.directory)
    changed = library.edit_segment(lesson["id"], lesson["segments"][0]["id"], "Có hai giá trị.", "There are two values.", False, False)
    snapshot = review_snapshot(changed, library.directory)
    assert not any(item["approved"] for item in snapshot["segments"][0]["support"])


def test_interrupted_request_pins_account_and_reuses_completed_parts(library):
    request = request_for(library, text="\n\n".join("Có một giá trị." for _ in range(17)))
    provider = EchoProvider()
    provider.fail_at = 2
    with pytest.raises(PlanError, match="Interrupted"):
        execute(library, request, provider)
    with pytest.raises(PlanError) as failure:
        execute(library, request, provider)
    assert failure.value.code == "confirmation_required" and len(provider.calls) == 2
    with pytest.raises(PlanError, match="tài khoản"):
        convert_request(request, library.directory, provider, "account-two", Event(), lambda _: None, retry_unconfirmed=True)
    result = execute(library, request, provider, retry_unconfirmed=True)
    assert len(result["blocks"]) == 17 and len(provider.calls) == 4
    again = execute(library, request, provider)
    assert again["blocks"] == result["blocks"] and len(provider.calls) == 4
    assert not library.list_lessons()  # no SQLite mutation in workers


@pytest.mark.parametrize("mutation", ["source", "ref", "level", "approval", "pair", "support"])
def test_invalid_ai_data_rejected_before_lesson_creation(library, mutation):
    request = request_for(library, text="VI: Có một giá trị.\nEN: There is one value.")
    manifest = normalize_source(request, library.directory, Event(), lambda _: None)
    result = execute(library, request)
    value = {"schema_version": 1, "source_sha256": request["config"]["source_sha256"], "level": 2,
             "blocks": copy.deepcopy(result["blocks"])}
    if mutation == "source":
        value["blocks"][0]["source_text"] += "invented"
    elif mutation == "ref":
        value["blocks"][0]["source_ref"] = "foreign"
    elif mutation == "level":
        value["level"] = 4
    elif mutation == "approval":
        value["blocks"][0]["approved"] = True
    elif mutation == "pair":
        value["blocks"][0]["en"] = "rewritten"
    else:
        value["blocks"][0]["support"][0]["source_refs"] = ["another-object"]
    with pytest.raises(PlanError):
        validate_result(value, manifest["units"], request["config"])
    assert not library.list_lessons()


def test_source_changed_or_cancelled_never_calls_ai(library):
    request = request_for(library)
    provider = EchoProvider()
    source = library.directory / "chatgpt" / request["config"]["request_id"] / request["config"]["source_file"]
    source.write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="thay đổi"):
        execute(library, request, provider)
    request = request_for(library)
    cancelled = Event()
    cancelled.set()
    with pytest.raises(PlanError):
        convert_request(request, library.directory, provider, "account-one", cancelled, lambda _: None)
    assert not provider.calls
