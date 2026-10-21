from contextlib import closing
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.conversion_input import creation_config
from app.conversion_progress import ConversionProgress
from app.google_conversion import convert, load, prepare
from app.google_translation import KeyStore, Translator
from app.library import Library


def transport(texts, key, source, target):
    pairs = {'Cơ năng': 'Mechanical energy', 'Động năng': 'Kinetic energy',
             'Vật có vận tốc': 'The object has a velocity of',
             'Năng lượng được bảo toàn.': 'Energy is conserved.'}
    results = []
    for text in texts:
        for vi, en in pairs.items():
            text = text.replace(vi, en)
        results.append(text)
    return results


def config(style='template', mode='parallel_columns'):
    return creation_config('Cơ năng', 'Vật lý', 'THPT', '11', mode, 'standard', style)


@pytest.mark.parametrize('mode', ['parallel_columns', 'sentence_pairs', 'english_only'])
def test_text_to_local_pptx_paired_notes_and_idempotent_save(tmp_path, mode):
    KeyStore(tmp_path).add('AIza' + 'a' * 35)
    translator = Translator(tmp_path, transport=transport)
    job = prepare(tmp_path, config(mode=mode), text='Năng lượng được bảo toàn.\nVật có vận tốc 12 m/s.')
    events = []
    progress = ConversionProgress(job['folder'], events.append)
    lesson, preview = convert(tmp_path, job, Event(), progress, translator=translator)
    assert lesson['external_deck']['provider'] == 'google_translate'
    assert lesson['narration']['mode'] == 'none'
    assert lesson['slide_quality']['counts'].get('FAIL', 0) == 0
    assert all(not segment['approved'] for segment in lesson['segments'])
    deck = Presentation(preview['path'])
    text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert 'Energy is conserved.' in text and '12\u00a0m/s' in text
    assert ('Năng lượng được bảo toàn.' in text) == (mode != 'english_only')
    assert all('BILICLASS_NOTES:' in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert events[-1]['stage'] == 'save'
    again, _ = convert(tmp_path, job, Event(), progress, translator=translator)
    assert again['id'] == lesson['id']
    with closing(Library(tmp_path)) as library:
        assert len(library.list_lessons()) == 1


def test_source_mode_keeps_native_editable_slides(tmp_path):
    source = tmp_path / 'original.pptx'
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for text, y, height, size in [('Cơ năng', .4, 1, 32), ('Vật có vận tốc 12 m/s.', 2, 4, 24)]:
        box = slide.shapes.add_textbox(Inches(.6), Inches(y), Inches(12), Inches(height))
        box.text = text
        box.text_frame.paragraphs[0].font.size = Pt(size)
        box.text_frame.paragraphs[0].font.name = 'Arial'
    deck.save(source)
    original = source.read_bytes()
    KeyStore(tmp_path).add('AIza' + 'a' * 35)
    job = prepare(tmp_path, config('source'), source)
    lesson, preview = convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None),
                              translator=Translator(tmp_path, transport=transport))
    assert source.read_bytes() == original
    assert len(Presentation(preview['path']).slides) == 1
    assert lesson['external_deck']['requested_style'] == 'source'
    assert lesson['segments'][0]['en']


def test_changed_source_cannot_resume(tmp_path):
    job = prepare(tmp_path, config(), text='Cơ năng')
    from pathlib import Path
    (Path(job['folder']) / 'source.txt').write_text('changed')
    with pytest.raises(ValueError, match='thay đổi'):
        load(tmp_path, job['folder'])


def test_only_final_checked_artifact_is_published_without_retranslation(tmp_path, monkeypatch):
    import hashlib
    from pathlib import Path

    from app.chatgpt_handoff import inspect_returned_deck

    KeyStore(tmp_path).add('AIza' + 'a' * 35)
    checks, requests = [], []

    def checked(path, directory, **kwargs):
        checks.append(Path(path))
        candidate = Path(path).with_name('checked.pptx')
        deck = Presentation(path)
        deck.core_properties.comments = 'Local QA final artifact'
        deck.save(candidate)
        inspection = inspect_returned_deck(candidate)
        inspection['review_notes'].append('Review retained after automatic QA')
        return {'path': str(candidate), 'inspection': inspection, 'changes': [{'operation': 'fixture-fit'}]}

    def translated(*args):
        requests.append(1)
        return transport(*args)

    monkeypatch.setattr('app.slide_repair.repair_and_inspect', checked)
    job = prepare(tmp_path, config(), text='Năng lượng được bảo toàn.')
    lesson, preview = convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None),
                              translator=Translator(tmp_path, transport=translated))
    assert len(checks) == len(requests) == 1
    assert lesson['source']['sha256'] == hashlib.sha256(checks[0].with_name('checked.pptx').read_bytes()).hexdigest()
    assert Presentation(preview['path']).core_properties.comments == 'Local QA final artifact'
    assert 'Review retained after automatic QA' in lesson['external_deck']['review_notes']


def test_google_job_cannot_load_outside_current_library(tmp_path):
    with pytest.raises(ValueError):
        load(tmp_path, tmp_path.parent)


def test_low_disk_is_reported_before_creating_a_google_job(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from app.storage_space import MIB, StorageProblem

    monkeypatch.setattr('app.storage_space.shutil.disk_usage', lambda _: SimpleNamespace(free=64 * MIB))
    with pytest.raises(StorageProblem):
        prepare(tmp_path, config(), text='Cơ năng')
    assert not (tmp_path / 'google_jobs').exists()


def test_storage_failure_at_publish_preserves_code_and_resume_reuses_checked_deck(tmp_path, monkeypatch):
    from pathlib import Path

    from app.chatgpt_handoff import inspect_returned_deck
    from app.storage_space import StorageProblem

    checks, calls = [], []
    def checked(path, directory, **kwargs):
        checks.append(1)
        return {'path': str(path), 'inspection': inspect_returned_deck(path), 'changes': []}
    def translated(*args):
        calls.append(1)
        return transport(*args)
    monkeypatch.setattr('app.slide_repair.repair_and_inspect', checked)
    job = prepare(tmp_path, config(), text='Năng lượng được bảo toàn.')
    KeyStore(tmp_path).add('AIza' + 'a' * 35)
    store = Library.store_source
    monkeypatch.setattr(Library, 'store_source', lambda *_: (_ for _ in ()).throw(StorageProblem('disk full')))
    with pytest.raises(StorageProblem) as caught:
        convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None),
                translator=Translator(tmp_path, transport=translated))
    assert caught.value.code == 'storage'
    assert (Path(job['folder']) / 'publish-ready.json').is_file()
    monkeypatch.setattr(Library, 'store_source', store)
    lesson, preview = convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None),
                              translator=Translator(tmp_path, transport=translated))
    assert preview['lesson_id'] == lesson['id']
    assert len(calls) == len(checks) == 1


def test_document_parser_failure_is_not_reported_as_missing_key(tmp_path, monkeypatch):
    import json
    from pathlib import Path

    from app import importers
    from app.google_translation import TranslationProblem

    job = prepare(tmp_path, config(), text='Cơ năng')
    def fail(*args, **kwargs):
        raise AttributeError('document has no inherited font')
    monkeypatch.setattr(importers, 'parse_document', fail)
    with pytest.raises(TranslationProblem) as caught:
        convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None), translator=object())
    assert caught.value.code == 'source_document'
    diagnostic = json.loads((Path(job['folder']) / 'conversion-error.json').read_text())
    assert diagnostic['type'] == 'AttributeError'
    assert diagnostic['code'] == 'source_document'
    assert 'document has no inherited font' not in str(diagnostic)


@pytest.mark.parametrize('already_saved', [False, True])
def test_confirmed_discard_is_durable_scoped_and_keeps_source_keys_and_usage(tmp_path, monkeypatch, already_saved):
    from pathlib import Path

    from app.chatgpt_handoff import inspect_returned_deck
    from app.google_conversion import ConversionDiscarded, discard, mark_discarded

    monkeypatch.setattr('app.slide_repair.repair_and_inspect', lambda path, *_args, **_kwargs:
                        {'path': str(path), 'inspection': inspect_returned_deck(path), 'changes': []})
    store = KeyStore(tmp_path)
    store.add('AIza' + 'a' * 35)
    job = prepare(tmp_path, config(), text='Năng lượng được bảo toàn.')
    source = Path(job['folder']) / 'source.txt'
    original = source.read_bytes()
    with closing(Library(tmp_path)) as library:
        other = library.create('Keep this lesson', 'Vật lý', 'THPT', '11', [('Giữ nguyên', 'Keep')])
    saved = None
    if already_saved:
        saved, _ = convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None),
                           translator=Translator(tmp_path, transport=transport))
    usage = store.public()
    encrypted = store.path.read_bytes()
    mark_discarded(tmp_path, job['folder'])
    with pytest.raises(ConversionDiscarded):
        convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None), translator=object())
    discarded = discard(tmp_path, job['folder'])
    assert discarded == ([saved['id']] if saved else [])
    assert discard(tmp_path, job['folder']) == [], 'Crash recovery must be idempotent'
    with closing(Library(tmp_path)) as library:
        assert library.get(other['id'])['title'] == 'Keep this lesson'
        assert len(library.list_lessons(verify_audio=False)) == 1
        assert not library.setting('google_request_folder', '')
        if saved:
            assert library.get(saved['id'], include_trash=True)['deleted_at']
    assert source.read_bytes() == original
    assert store.path.read_bytes() == encrypted and store.public() == usage
    with pytest.raises(ValueError):
        mark_discarded(tmp_path, tmp_path.parent)


@pytest.mark.parametrize('already_saved', [False, True])
def test_dismissed_input_forgets_job_but_never_deletes_a_saved_lesson(tmp_path, monkeypatch, already_saved):
    from pathlib import Path
    from app.chatgpt_handoff import inspect_returned_deck
    from app.google_conversion import ConversionDiscarded, dismiss, load

    monkeypatch.setattr('app.slide_repair.repair_and_inspect', lambda path, *_args, **_kwargs:
                        {'path': str(path), 'inspection': inspect_returned_deck(path), 'changes': []})
    store=KeyStore(tmp_path); store.add('AIza'+'a'*35)
    job=prepare(tmp_path,config(),text='Năng lượng được bảo toàn.')
    source=Path(job['folder'])/'source.txt'; original=source.read_bytes()
    lesson=None
    if already_saved:
        lesson,_=convert(tmp_path,job,Event(),ConversionProgress(job['folder'],lambda _:None),
                         translator=Translator(tmp_path,transport=transport))
    usage=store.public(); encrypted=store.path.read_bytes()
    dismiss(tmp_path,job['folder'])
    dismiss(tmp_path,job['folder'])
    with closing(Library(tmp_path)) as library:
        assert library.setting('google_request_folder','')==''
        assert len(library.list_lessons(verify_audio=False))==int(already_saved)
        if lesson:assert not library.get(lesson['id']).get('deleted_at')
    if not already_saved:
        with pytest.raises(ConversionDiscarded):load(tmp_path,job['folder'])
    else:
        assert load(tmp_path,job['folder'])['folder']==job['folder']
    assert source.read_bytes()==original and store.public()==usage and store.path.read_bytes()==encrypted


@pytest.mark.skipif(__import__('os').name != 'nt', reason='Windows job lock')
def test_same_job_locked_by_another_window_does_not_send(tmp_path):
    import msvcrt
    from pathlib import Path
    job = prepare(tmp_path, config(), text='Cơ năng')
    with (Path(job['folder']) / 'convert.lock').open('w+b') as lock:
        lock.write(b'0')
        lock.flush()
        lock.seek(0)
        msvcrt.locking(lock.fileno(), msvcrt.LK_NBLCK, 1)
        try:
            with pytest.raises(ValueError, match='cửa sổ'):
                convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None))
        finally:
            lock.seek(0)
            msvcrt.locking(lock.fileno(), msvcrt.LK_UNLCK, 1)
