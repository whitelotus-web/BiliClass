import copy
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.lesson_completion import complete_existing, prepare_library_teaching
from app.lesson_narration import audio_status, make_plan, variant_id
from app.library import Library, RevisionConflict
from app.voice_import import import_presentation


def converted(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'VI: Xin chào.\nEN: Hello.'
    slide.notes_slide.notes_text_frame.text = 'VI: Xin chào.\nEN: Hello.'
    path = tmp_path / 'original.pptx'
    deck.save(path)
    lesson, _ = import_presentation(tmp_path / 'library', path)
    library = Library(tmp_path / 'library')
    lesson['voice_only'] = False
    lesson['slide_quality'] = {'counts': {'FAIL': 1}, 'slides': []}
    library._write(lesson)
    return library, lesson, path


def checked_copy(library, path):
    from app.chatgpt_handoff import inspect_returned_deck

    deck = Presentation(path)
    deck.core_properties.subject = 'Locally verified layout copy'
    target = library.directory / 'temp' / 'completed.pptx'
    target.parent.mkdir(exist_ok=True)
    deck.save(target)
    inspection = inspect_returned_deck(target)
    inspection['slide_quality']['counts'] = {'FAIL': 0, 'CHECK': 1, 'PASS': 0}
    return {'path': str(target), 'inspection': inspection}


def test_completion_keeps_source_history_and_identical_voice_tasks(tmp_path, monkeypatch):
    import wave

    library, lesson, path = converted(tmp_path)
    plan, _ = make_plan(lesson, 'en', {'voice': 'vieneu:Mai Anh', 'rate': 0})
    for key in plan['keys']['en']:
        target = library.directory / 'audio' / (key + '.wav')
        target.parent.mkdir(exist_ok=True)
        with wave.open(str(target), 'wb') as wav:
            wav.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
            wav.writeframes(b'\0\0' * 24000)
    plan.update(lesson_id=lesson['id'], revision=lesson['revision'], complete=plan['total'], running=False)
    lesson = library.update_narration(lesson['id'], plan)
    assert audio_status(lesson, library.directory)['ready']
    before = path.read_bytes()
    candidate = checked_copy(library, path)
    monkeypatch.setattr('app.slide_repair.repair_and_inspect', lambda *a, **k: candidate)
    updated = complete_existing(library, lesson)
    assert updated['revision'] == lesson['revision'] + 1
    assert updated['slide_quality']['counts']['FAIL'] == 0
    assert path.read_bytes() == before
    assert not any(segment['approved'] for segment in updated['segments'])
    assert audio_status(updated, library.directory)['ready']
    assert variant_id(updated['narration']) != variant_id(plan)
    assert updated['narration']['keys'] == plan['keys']
    assert library.db.execute('SELECT count(*) FROM lesson_revisions WHERE lesson_id=?', (lesson['id'],)).fetchone()[0]
    library.close()


def test_completion_rejects_concurrent_rename_and_keeps_it(tmp_path, monkeypatch):
    library, lesson, path = converted(tmp_path)
    candidate = checked_copy(library, path)

    def repair(*args, **kwargs):
        current = library.get(lesson['id'])
        current['title'] = 'Renamed while Office was working'
        library._write(current)
        return candidate

    monkeypatch.setattr('app.slide_repair.repair_and_inspect', repair)
    with pytest.raises(RevisionConflict):
        complete_existing(library, lesson)
    saved = library.get(lesson['id'])
    assert saved['title'].startswith('Renamed') and saved['source'] == lesson['source']
    library.close()


def test_no_improvement_is_recorded_once_and_never_called_ready(tmp_path, monkeypatch):
    from app.chatgpt_handoff import inspect_returned_deck

    library, lesson, _ = converted(tmp_path)
    source = library.directory / 'sources' / lesson['source']['file']
    inspection = inspect_returned_deck(source)
    inspection['slide_quality']['counts']['FAIL'] = 1
    calls = []
    monkeypatch.setattr('app.slide_repair.repair_and_inspect',
                        lambda *a, **k: calls.append(1) or {'path': str(source), 'inspection': inspection})
    updated = complete_existing(library, lesson)
    assert updated['revision'] == lesson['revision'] and updated['slide_quality']['counts']['FAIL'] == 1
    assert complete_existing(library, updated) == updated and len(calls) == 1
    library.close()


def test_cancel_during_completion_never_publishes_candidate(tmp_path, monkeypatch):
    library, lesson, path = converted(tmp_path)
    candidate, cancel = checked_copy(library, path), Event()

    def repair(*args, **kwargs):
        cancel.set()
        return candidate

    monkeypatch.setattr('app.slide_repair.repair_and_inspect', repair)
    with pytest.raises(InterruptedError):
        complete_existing(library, lesson, cancel)
    assert library.get(lesson['id']) == lesson
    library.close()


def test_open_library_reuses_ready_voice_when_active_variant_is_incomplete(tmp_path, monkeypatch):
    from test_teaching_session import prepared_lesson

    lesson, _, directory, _ = prepared_lesson(tmp_path)
    library = Library(directory)
    old = copy.deepcopy(lesson)
    lesson['narration'] = {'version': 2, 'mode': 'none'}
    library._write(lesson)
    library.close()
    monkeypatch.setattr('app.speech.synthesize', lambda *a, **k: pytest.fail('Opening must not generate voice'))
    result = prepare_library_teaching(directory, lesson['id'])
    assert result['lesson']['narration']['keys'] == old['narration']['keys']
    assert result['media'][1] == ['vi', 'en', 'both']
    assert not any(segment['approved'] for segment in result['lesson']['segments'])


def test_open_saved_powerpoint_does_not_run_full_office_repair_or_change_the_saved_product(tmp_path, monkeypatch):
    library, lesson, _ = converted(tmp_path)
    before = library.get(lesson['id'])
    monkeypatch.setattr('app.slide_repair.repair_and_inspect',
                        lambda *args, **kwargs: pytest.fail('Open must not run a full Office repair'))
    monkeypatch.setattr('app.speech.synthesize',
                        lambda *args, **kwargs: pytest.fail('Open must not create narration'))
    try:
        result = prepare_library_teaching(library.directory, lesson['id'])
        assert result['path'] and result['lesson']['id'] == lesson['id']
        assert result['lesson']['slide_quality']['counts']['FAIL'] == 1
        assert '1 slide' in result['warning']
        assert library.get(lesson['id'])['source'] == before['source']
        assert library.get(lesson['id'])['revision'] == before['revision']
    finally:
        library.close()


def test_open_and_narration_are_separate_entry_points(tmp_path, monkeypatch):
    from PySide6.QtCore import QCoreApplication

    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    library, lesson, _ = converted(tmp_path)
    prepared = prepare_library_teaching(library.directory, lesson['id'])
    bridge = Bridge(library)
    destinations, opened = [], []
    bridge.navigate.connect(destinations.append)
    monkeypatch.setattr(bridge, '_activate_teaching', opened.append)
    try:
        bridge._open_library_teaching(prepared)
        assert opened == [prepared] and destinations == []
        assert not bridge.narrationBusy and bridge.powerpoint_worker is None
        bridge.openLessonNarration(lesson['id'])
        assert destinations == ['voice'] and opened == [prepared]
        app.processEvents()
    finally:
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                       bridge.voice_controls, bridge._companion_preferences, bridge.classroom):
            worker.shutdown()
        library.close()


def test_new_spoken_tasks_never_reuse_old_voice_signature(tmp_path):
    from app.lesson_completion import _retain_unchanged_audio

    library, lesson, _ = converted(tmp_path)
    plan, _ = make_plan(lesson, 'en', {'voice': 'vieneu:Mai Anh', 'rate': 0})
    lesson['narration'] = plan
    updated = copy.deepcopy(lesson)
    updated['segments'][0]['en'] = 'A changed sentence.'
    updated['source']['sha256'] = 'f' * 64
    # Explicit script has precedence over the native text.
    updated['narration_script'] = {updated['segments'][0]['id']: {'vi': 'Xin chào.', 'en': 'A changed sentence.', 'manual': True}}
    _retain_unchanged_audio(lesson, updated)
    assert updated['narration']['signature'] == plan['signature']
    assert not audio_status(updated, library.directory)['ready']
    library.close()
