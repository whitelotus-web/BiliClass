import io
import zipfile
from threading import Event

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.google_conversion import load, prepare
from app.pptx_template import build, inspect, store, stored_path, validate_roles
from app.source_deck import _part_relative, _rels_path, _xml, source_slide_parts
from app.template_builtin import create


def settings():
    return {'title': 'Cơ năng', 'english_title': 'Mechanical energy', 'preset': 'standard',
            'style': 'template', 'conversion_format': 'parallel_columns'}


def test_template_snapshot_survives_original_removal_and_detects_tampering(tmp_path):
    path = create(tmp_path/'own.pptx', 'standard')
    metadata = store(tmp_path, path)
    job = prepare(tmp_path, {**settings(), 'custom_template': metadata}, text='Năng lượng.')
    path.unlink()
    assert stored_path(tmp_path, metadata).is_file()
    assert load(tmp_path, job['folder'])['config']['template_roles'] == metadata['roles']
    from pathlib import Path
    (Path(job['folder'])/'template.pptx').write_bytes(b'changed')
    with pytest.raises(ValueError, match='Mẫu'):
        load(tmp_path, job['folder'])


@pytest.mark.parametrize('mode', ['parallel_columns', 'sentence_pairs', 'english_only'])
def test_populates_native_template_and_paginates_without_losing_text(tmp_path, mode):
    template = create(tmp_path/'template.pptx', 'visual')
    source = Presentation(template)
    image = io.BytesIO()
    Image.new('RGB', (50, 40), '#eabd7e').save(image, 'PNG')
    source.slides[2].shapes.add_picture(io.BytesIO(image.getvalue()), Inches(12), Inches(.4), width=Inches(.4))
    source.save(template)
    data = [('Ý '+str(i), f'Năng lượng được bảo toàn thứ {i}.', f'Energy is conserved at step {i}.') for i in range(25)]
    target = tmp_path/'output.pptx'
    build(target, template, {**settings(), 'conversion_format': mode}, data, Event(),
          roles={'cover': 1, 'intro': 2, 'content': 3, 'ending': 4})
    deck = Presentation(target)
    assert len(deck.slides) > 4
    text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert '[SOURCE_' not in text and '[LESSON_' not in text
    for i in range(25):
        assert f'Energy is conserved at step {i}.' in text
        assert (f'Năng lượng được bảo toàn thứ {i}.' in text) == (mode != 'english_only')
    assert 'BILICLASS_NOTES:' in deck.slides[0].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(template) as before, zipfile.ZipFile(target) as after:
        for name in before.namelist():
            if name.startswith('ppt/media/'):
                assert before.read(name) == after.read(name)
        note_parts = set()
        for part in source_slide_parts(target):
            rels = _xml(after.read(_rels_path(part)))
            note = next(r for r in rels if r.get('Type', '').endswith('/notesSlide'))
            resolved = _part_relative(part, note.get('Target'))
            assert resolved not in note_parts
            note_parts.add(resolved)
            back = _xml(after.read(_rels_path(resolved)))
            link = next(r for r in back if r.get('Type', '').endswith('/slide'))
            assert _part_relative(resolved, link.get('Target')) == part


def test_source_picture_remains_embedded_with_own_relationship(tmp_path):
    template = create(tmp_path/'template.pptx', 'practice')
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    image = io.BytesIO()
    Image.new('RGB', (120, 70), '#bdbd71').save(image, 'PNG')
    slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(1), width=Inches(3))
    source_path = tmp_path/'source.pptx'
    source.save(source_path)
    output = tmp_path/'output.pptx'
    build(output, template, settings(), [('Ý', 'Năng lượng.', 'Energy.')], Event(), source_path,
          {'cover':0, 'intro':0, 'content':3, 'ending':0})
    deck = Presentation(output)
    images = [s.image.blob for page in deck.slides for s in page.shapes if int(s.shape_type)==13]
    assert images == [image.getvalue()]
    assert len(deck.slides) == 2


def test_invalid_role_or_noneditable_template_is_rejected(tmp_path):
    template = create(tmp_path/'template.pptx', 'standard')
    metadata = inspect(template)
    with pytest.raises(ValueError):
        validate_roles(metadata, {'content':99})
    blank = Presentation()
    blank.slides.add_slide(blank.slide_layouts[6])
    blank.save(tmp_path/'blank.pptx')
    with pytest.raises(ValueError, match='khung'):
        inspect(tmp_path/'blank.pptx')


def test_ocr_model_paths_do_not_import_legacy_translation(tmp_path, monkeypatch):
    import builtins
    import sys

    from app.local_ocr import model_paths
    from app.vietnamese_ocr import pack_directory
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name == 'translation' or name.startswith('biliclass_m0'):
            raise AssertionError('OCR imported the obsolete translation provider')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', guarded)
    monkeypatch.setattr(sys, 'frozen', True, raising=False)
    monkeypatch.setenv('BILICLASS_DATA', str(tmp_path))
    monkeypatch.setenv('BILICLASS_MODEL_DIR', str(tmp_path/'models'))
    assert pack_directory() is None
    assert isinstance(model_paths(), dict)
