import copy
import hashlib
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from app.deck_export import export_deck
from app.importers import parse_document
from app.library import Library
from app.pack import export_pack, import_pack
from app.readiness import preparation_current
from app.source_deck import NS, export_source_deck, prepare_source_deck


def fixture(tmp_path):
    source = tmp_path / "teacher.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("182C3D")
    heading = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(9), Inches(1))
    heading.text = "Hàm số"
    run = heading.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.name, run.font.bold = Pt(32), "Arial", True
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    body = slide.shapes.add_textbox(Inches(.5), Inches(1.6), Inches(6), Inches(1.2))
    body.text = "y = 2x + 1"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    picture = tmp_path / "image.png"
    Image.new("RGB", (80, 80), "#f38c46").save(picture)
    slide.shapes.add_picture(str(picture), Inches(8), Inches(2), width=Inches(2))
    chart = CategoryChartData()
    chart.categories = ["A", "B"]
    chart.add_series("Điểm", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(4), Inches(4), Inches(2), chart)
    slide.notes_slide.notes_text_frame.text = "Ghi chú của thầy cô"
    slide._element.append(etree.Element(f"{{{NS['p']}}}transition", advClick="1"))
    # Preserve unknown/unsupported extension payloads without re-saving them.
    ext = etree.SubElement(slide._element, f"{{{NS['p']}}}extLst")
    etree.SubElement(ext, f"{{{NS['p']}}}ext", uri="test-native-payload").text = "native"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    table = second.shapes.add_table(2, 2, Inches(.5), Inches(.5), Inches(10), Inches(2)).table
    for row, values in zip(table.rows, (("Đại lượng", "Giá trị"), ("x", "2")), strict=True):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(picture), Inches(4), Inches(1))
    deck.save(source)
    library = Library(tmp_path / "library")
    lesson = library.create("Bài hàm số", "Toán", "THPT", "10", parse_document(source), library.store_source(source))
    translations = ["Functions", "y = 2x + 1", "Quantity | Value\nx | 2"]
    for segment, english in zip(lesson["segments"], translations, strict=True):
        lesson = library.edit_segment(lesson["id"], segment["id"], segment["vi"], english, True)
    return source, library, lesson


def test_native_copy_preserves_design_media_chart_notes_and_image_only_slide(tmp_path):
    source, library, lesson = fixture(tmp_path)
    try:
        original_bytes = source.read_bytes()
        before = copy.deepcopy(lesson)
        result = export_source_deck(lesson, library.directory, tmp_path / "bilingual.pptx")
        assert result["slide_map"] == [1, 1, 2, 2, 3]
        output = Presentation(result["path"])
        assert (output.slide_width, output.slide_height) == (Inches(12), Inches(6.75))
        assert output.slides[0].shapes[0].text == "Hàm số"
        assert output.slides[1].shapes[0].text == "Functions"
        assert output.slides[1].shapes[0].left == output.slides[0].shapes[0].left
        assert output.slides[1].shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor.from_string("FFFFFF")
        assert output.slides[1].notes_slide.notes_text_frame.text == output.slides[0].notes_slide.notes_text_frame.text
        original_chart = output.slides[0].shapes[3].chart
        english_chart = output.slides[1].shapes[3].chart
        # Office cannot open a deck if duplicated slides share one chart part.
        assert original_chart.part.partname != english_chart.part.partname
        assert original_chart.part.blob == english_chart.part.blob
        assert english_chart.series[0].values == original_chart.series[0].values
        assert original_chart.part.chart_workbook.xlsx_part.partname != english_chart.part.chart_workbook.xlsx_part.partname
        assert original_chart.part.chart_workbook.xlsx_part.blob == english_chart.part.chart_workbook.xlsx_part.blob
        assert output.slides[3].shapes[0].table.cell(0, 0).text == "Quantity"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(result["path"]) as converted:
            unchanged = [name for name in original.namelist() if name.startswith(("ppt/media/", "ppt/charts/", "ppt/slideMasters/", "ppt/theme/"))]
            assert unchanged
            assert all(original.read(name) == converted.read(name) for name in unchanged)
            assert original.read("ppt/slides/slide3.xml") == converted.read("ppt/slides/slide3.xml")
            english = etree.fromstring(converted.read("ppt/slides/biliclass-en-1.xml"))
            assert english.find("p:transition", NS).get("advClick") == "1"
            assert english.find("p:extLst/p:ext", NS).text == "native"
        assert source.read_bytes() == original_bytes and lesson == before
        assert hashlib.sha256((library.directory / "sources" / lesson["source"]["file"]).read_bytes()).hexdigest() == lesson["source"]["sha256"]
        assert export_deck(lesson, library.directory, tmp_path / "ordinary-export.pptx").is_file()
    finally:
        library.close()


def test_cached_native_deck_is_reused_and_corrupt_output_is_regenerated(tmp_path):
    _, library, lesson = fixture(tmp_path)
    try:
        result = prepare_source_deck(lesson, library.directory)
        path = Path(result["path"])
        modified = path.stat().st_mtime_ns
        assert prepare_source_deck(lesson, library.directory) == result
        assert path.stat().st_mtime_ns == modified
        path.write_bytes(b"damaged output")
        assert prepare_source_deck(lesson, library.directory) == result
        assert len(Presentation(path).slides) == 5
        updated = copy.deepcopy(lesson)
        updated["segments"][0]["en"] = "A function"
        assert prepare_source_deck(updated, library.directory)["path"] != str(path)
    finally:
        library.close()


def test_layout_language_rules_and_reviewed_split_keep_source_mapping(tmp_path):
    _, library, lesson = fixture(tmp_path)
    try:
        lesson = library.set_presentation(lesson["id"], 4, "english_rescue")
        result = export_source_deck(lesson, library.directory, tmp_path / "english.pptx")
        assert result["slide_map"] == [1, 2, 3]
        assert Presentation(result["path"]).slides[0].shapes[0].text == "Functions"
        lesson = library.set_presentation(lesson["id"], 0, "keyword_overlay")
        result = export_source_deck(lesson, library.directory, tmp_path / "keywords.pptx", [{"vi": "Hàm số", "en": "function"}])
        assert len(Presentation(result["path"]).slides) == 3
        assert "function" in Presentation(result["path"]).slides[0].shapes[0].text
        lesson = library.set_presentation(lesson["id"], 2, "line_pair")
        segment = lesson["segments"][0]
        lesson = library.split_segment(lesson["id"], segment["id"], 3, 4)
        for segment in lesson["segments"][:2]:
            lesson = library.edit_segment(lesson["id"], segment["id"], segment["vi"], segment["en"], True)
        result = export_source_deck(lesson, library.directory, tmp_path / "split.pptx")
        assert "Func\ntions" == Presentation(result["path"]).slides[1].shapes[0].text
    finally:
        library.close()


def test_source_style_rejects_unreviewed_mismatched_overflow_and_source_overwrite(tmp_path):
    _, library, lesson = fixture(tmp_path)
    try:
        target = tmp_path / "existing.pptx"
        target.write_bytes(b"previous output")
        invalid = copy.deepcopy(lesson)
        invalid["segments"][0]["approved"] = False
        with pytest.raises(ValueError, match="Duyệt"):
            export_source_deck(invalid, library.directory, target)
        invalid = copy.deepcopy(lesson)
        invalid["segments"][0]["locator"] = "Slide 999"
        with pytest.raises(ValueError, match="khớp"):
            export_source_deck(invalid, library.directory, target)
        invalid = copy.deepcopy(lesson)
        invalid["segments"][0]["en"] = "Long English translation " * 150
        with pytest.raises(ValueError, match="quá dài"):
            export_source_deck(invalid, library.directory, target)
        with pytest.raises(ValueError, match="giữ nguyên"):
            export_source_deck(lesson, library.directory, library.directory / "sources" / lesson["source"]["file"])
        assert target.read_bytes() == b"previous output"
    finally:
        library.close()


def test_style_choice_is_revisioned_and_survives_shared_pack(tmp_path):
    _, library, lesson = fixture(tmp_path)
    try:
        assert lesson["presentation_style"] == "source"
        lesson = library.mark_prepared(lesson["id"], lesson["revision"])
        changed = library.set_presentation_style(lesson["id"], "template")
        assert not preparation_current(changed)
        assert all(segment["approved"] for segment in changed["segments"])
        changed = library.set_presentation_style(lesson["id"], "source")
        imported = import_pack(library, export_pack(library, changed["id"], tmp_path / "lesson.biliclass"))
        assert imported["presentation_style"] == "source"
        assert all(not segment["approved"] for segment in imported["segments"])
        plain = library.create("Bài chữ", "Toán", "THPT", "10", [("Ý 1", "Văn bản")])
        with pytest.raises(ValueError, match="PowerPoint"):
            library.set_presentation_style(plain["id"], "source")
    finally:
        library.close()
