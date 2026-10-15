import hashlib
import json
import wave
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from threading import Event
from types import SimpleNamespace

import numpy as np
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app import speech
from app.browser_audio import narration_tasks
from app.chatgpt_handoff import inspect_returned_deck
from app.conversion_progress import ConversionProgress, measured
from app.lesson_import import save_conversion
from app.library import Library
from app.native_bilingual import assemble, extract
from app.slide_speech import english_text
from app.translation_batches import parse_response
from app.storage_space import StorageProblem


def pcm(path):
    with wave.open(str(path), 'wb') as writer:
        writer.setparams((1, 2, 48000, 0, 'NONE', 'not compressed'))
        time = np.arange(48000) / 48000
        writer.writeframes((np.sin(2*np.pi*220*time)*8000).astype('<i2').tobytes())


def test_all_rates_share_one_inference_even_when_first_request_is_fast(tmp_path, monkeypatch):
    calls = []

    def infer(text, voice, rate, target):
        calls.append((text, voice, rate))
        pcm(target)

    monkeypatch.setattr(speech, '_synthesize_vieneu', infer)
    for rate in [3, -2, 0, 1, 3]:
        result = speech.synthesize('Xin chào. Hello class.', 'vieneu:Mai Anh', rate, tmp_path)
        with wave.open(result['path']) as audio:
            assert audio.getnframes() == round(48000 / speech._SPEEDS[rate+3])
    assert calls == [('Xin chào. Hello class.', 'vieneu:Mai Anh', 0)]
    assert result['cached']
    assert len(list(tmp_path.glob('*.wav'))) == 4


def test_two_readers_do_not_infer_twice_and_a_failed_variant_keeps_base(tmp_path, monkeypatch):
    calls = []

    def infer(text, voice, rate, target):
        calls.append(rate)
        pcm(target)

    monkeypatch.setattr(speech, '_synthesize_vieneu', infer)
    with ThreadPoolExecutor(max_workers=2) as workers:
        results = list(workers.map(lambda rate: speech.synthesize('Hello.', 'vieneu:Mai Anh', rate, tmp_path), [1, -1]))
    assert calls == [0] and all(speech.valid_audio(r['path']) for r in results)
    base = tmp_path / (speech.audio_key('Hello.', 'vieneu:Mai Anh', 0) + '.wav')
    digest = hashlib.sha256(base.read_bytes()).hexdigest()

    def fail(*_):
        raise RuntimeError('speed failed')

    monkeypatch.setattr(speech, '_speed_variant', fail)
    with pytest.raises(RuntimeError, match='speed failed'):
        speech.synthesize('Hello.', 'vieneu:Mai Anh', 2, tmp_path)
    assert hashlib.sha256(base.read_bytes()).hexdigest() == digest
    assert calls == [0]
    assert not (tmp_path / (speech.audio_key('Hello.', 'vieneu:Mai Anh', 2) + '.wav')).exists()
    assert len(list(tmp_path.glob('*.wav'))) == 3


def test_existing_portable_rate_variant_is_reused_without_neutral_audio(tmp_path, monkeypatch):
    destination = tmp_path / (speech.audio_key('Hello.', 'vieneu:Mai Anh', 2) + '.wav')
    pcm(destination)
    monkeypatch.setattr(speech, '_synthesize_vieneu', lambda *_: pytest.fail('must reuse existing audio'))
    assert speech.synthesize('Hello.', 'vieneu:Mai Anh', 2, tmp_path)['cached']
    assert len(list(tmp_path.glob('*.wav'))) == 1


def deck_file(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.4), Inches(.3), Inches(8), Inches(1.2))
    title.text = 'Cơ năng'
    title.text_frame.paragraphs[0].font.size = Pt(28)
    body = slide.shapes.add_textbox(Inches(.4), Inches(2), Inches(8), Inches(3))
    body.text = 'Vận tốc của vật thay đổi theo thời gian.'
    body.text_frame.paragraphs[0].font.size = Pt(24)
    slide.notes_slide.notes_text_frame.text = 'VI: Vận tốc thay đổi.\nEN: The velocity changes.'
    deck.save(path)
    return path


def config():
    return {'title': 'Bài thử', 'subject': 'Vật lý', 'education_level': 'THPT', 'grade': '11',
            'conversion_format': 'parallel_columns', 'style': 'source', 'preset': 'standard',
            'request_id': 'request-for-import-test'}


def test_failed_external_record_is_not_published_half_complete(tmp_path):
    source_path = deck_file(tmp_path / 'source.pptx')
    with closing(Library(tmp_path / 'library')) as library:
        source = library.store_source(source_path)
        inspection = inspect_returned_deck(source_path)
        inspection['slide_quality'] = object()  # A serialization failure after assembly.
        with pytest.raises(TypeError):
            library.create_external_lesson(config(), source, inspection)
        assert library.db.execute('SELECT COUNT(*) FROM lessons').fetchone()[0] == 0


def test_worker_import_is_idempotent_across_connections_and_never_starts_voice(tmp_path):
    source_path = deck_file(tmp_path / 'source.pptx')
    root = tmp_path / 'library'
    with closing(Library(root)) as library:
        source = library.store_source(source_path)
    inspection = inspect_returned_deck(source_path)
    with ThreadPoolExecutor(max_workers=2) as workers:
        results = list(workers.map(lambda _: save_conversion(root, config(), source, inspection), range(2)))
    assert results[0][0]['id'] == results[1][0]['id']
    with closing(Library(root)) as library:
        assert len(library.list_lessons()) == 1
        assert library.db.execute('SELECT COUNT(*) FROM lesson_revisions').fetchone()[0] == 0
    assert results[0][0]['narration']['mode'] == 'none'
    assert not (root / 'audio').exists()
    cancel = Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        save_conversion(root, {**config(), 'request_id': 'cancelled'}, source, inspection, cancel)


def test_source_artifact_requires_matching_paragraph_and_geometry(tmp_path):
    source = deck_file(tmp_path / 'source.pptx')
    deck = Presentation(source)
    box = deck.slides[0].shapes[1]
    box.text = 'Nguồn tham khảo https://example.test/lesson'
    deck.save(source)
    result = inspect_returned_deck(source, source, 'parallel_columns', 'source')
    issues = result['slide_quality']['slides'][0]['issues']
    assert any(i['code'] == 'source_text_needs_review' and i['evidence'] == 'same_native_text_and_box' for i in issues)
    box.text += ' 19ab568'
    changed = tmp_path / 'changed.pptx'
    deck.save(changed)
    issues = inspect_returned_deck(changed, source, 'parallel_columns', 'source')['slide_quality']['slides'][0]['issues']
    assert any(i['code'] == 'technical_text' and i['severity'] == 'FAIL' for i in issues)


def test_source_urls_and_garbage_are_not_spoken_after_provenance_review():
    inspection = {'profile': {'units': [{'vi': 'https://example.test/lesson', 'en': ''},
                                       {'vi': 'Chuỗi 19ab568', 'en': ''},
                                       {'vi': 'Nội dung hợp lệ.', 'en': 'Valid lesson content.'}]}}
    tasks, skipped, _ = narration_tasks(inspection, {'vi': 'speaker', 'en': 'speaker'})
    assert skipped == 2
    assert tasks == [('Nội dung hợp lệ.', 'speaker'), ('Valid lesson content.', 'speaker')]


def test_english_identifier_does_not_create_false_vietnamese_column(tmp_path):
    assert english_text('Predict what happens when pendulum Đ is pulled to one side and released.')
    assert not english_text('Con lắc Đ được kéo sang một bên rồi thả cho dao động.')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.3), Inches(2), Inches(4), Inches(2)).text = (
        'Dự đoán hiện tượng xảy ra khi con lắc Đ được kéo sang bên rồi thả cho dao động.')
    slide.shapes.add_textbox(Inches(5), Inches(2), Inches(4), Inches(2)).text = (
        'Predict what happens when pendulum Đ is pulled to one side and released.')
    path = tmp_path / 'columns.pptx'
    deck.save(path)
    result = inspect_returned_deck(path, conversion_format='parallel_columns')
    assert not any(i['code'] in {'parallel_columns_missing', 'english_body_missing'}
                   for i in result['slide_quality']['slides'][0]['issues'])
    assert any('Đ' in s['english_text'] for s in result['slide_speech']['slides'][0]['englishSpeechSegments'])


def test_large_source_font_does_not_pass_a_tiny_box_preflight(tmp_path):
    path = deck_file(tmp_path / 'source.pptx')
    deck = Presentation(path)
    shape = deck.slides[0].shapes[1]
    shape.width, shape.height = Inches(4), Inches(.55)
    shape.text = 'Vận tốc của vật thay đổi liên tục theo thời gian trong quá trình dao động.'
    shape.text_frame.paragraphs[0].font.size = Pt(40)
    deck.save(path)
    assert extract(path, 'parallel_columns')['visual_slides'] == [1]


def test_keyword_mode_keeps_a_stacked_bilingual_title(tmp_path):
    source = deck_file(tmp_path / 'source.pptx')
    manifest = extract(source, 'integrated_keywords')
    assert not manifest['visual_slides']
    values = {u['id']: {'en': 'Mechanical energy' if u['title'] else 'The velocity changes over time.',
                        'mixed': 'Cơ năng (mechanical energy)' if u['title'] else 'Vận tốc (velocity) thay đổi.'}
              for u in manifest['units']}
    result = assemble(source, tmp_path / 'out.pptx', manifest, values, 'integrated_keywords')
    title = Presentation(result).slides[0].shapes[0]
    assert title.text.splitlines() == ['Cơ năng', 'Mechanical energy']
    assert title.text_frame.paragraphs[1].runs[0].font.italic


def test_collapsing_three_source_lines_is_not_complete_translation():
    text = 'Hãy quan sát các hiện tượng xảy ra trong thí nghiệm.\nGiải thích nguyên nhân làm chuyển động của vật thay đổi.\nNêu kết luận và các điều kiện áp dụng trong thực tế.'
    source = {'id': 's1p1', 'slide': 1, 'text': text, 'raw': text, 'numbers': []}
    answer = {'batch': 'test', 'items': [{'id': 's1p1', 'en': 'Observe the experiment and explain why the motion changes. Give a conclusion and describe the conditions for practical application.'}]}
    with pytest.raises(ValueError, match='dòng'):
        parse_response(json.dumps(answer), [source], 'test', 'sentence_pairs')


def test_private_operation_timings_do_not_leak_text_into_metrics(tmp_path, monkeypatch):
    clock = [10.0]
    monkeypatch.setattr('app.conversion_progress.time.monotonic', lambda: clock[0])
    sink = ConversionProgress(tmp_path, lambda _: None)
    with measured(sink, 'local.inspect'):
        clock[0] += 2
    with pytest.raises(ValueError):
        with measured(sink, 'local.inspect'):
            clock[0] += 3
            raise ValueError('private lesson or browser details')
    saved = json.loads(sink.path.read_text(encoding='utf-8'))
    assert saved['operations']['local.inspect'] == {'calls': 2, 'seconds': 5, 'failures': 1}
    assert 'private lesson' not in sink.path.read_text(encoding='utf-8')


def test_full_destination_stops_before_upload_or_creating_an_outgoing_job(tmp_path, monkeypatch):
    from app.browser_dispatch import convert_available
    from app.chatgpt_handoff import prepare_request

    source = deck_file(tmp_path / 'source.pptx')
    monkeypatch.setattr('app.storage_space.shutil.disk_usage', lambda _: SimpleNamespace(free=0))
    with pytest.raises(StorageProblem, match='dung lượng'):
        prepare_request(tmp_path / 'library', config(), source)
    assert not (tmp_path / 'library' / 'chatgpt').exists()
    with pytest.raises(StorageProblem):
        convert_available([], tmp_path, tmp_path, Event(), lambda _: None,
                          converter=lambda *_: pytest.fail('must not open browser'))


def test_full_disk_still_allows_cached_voice_but_never_starts_inference(tmp_path, monkeypatch):
    path = tmp_path / (speech.audio_key('Hello.', 'vieneu:Mai Anh', 0) + '.wav')
    pcm(path)
    monkeypatch.setattr('app.storage_space.shutil.disk_usage', lambda _: SimpleNamespace(free=0))
    monkeypatch.setattr(speech, '_synthesize_vieneu', lambda *_: pytest.fail('disk is full'))
    assert speech.synthesize('Hello.', 'vieneu:Mai Anh', 0, tmp_path)['cached']
    with pytest.raises(StorageProblem):
        speech.synthesize('Different text.', 'vieneu:Mai Anh', 0, tmp_path)


def test_partial_stored_source_is_repaired_from_intact_input(tmp_path):
    source_path = deck_file(tmp_path / 'source.pptx')
    original = source_path.read_bytes()
    with closing(Library(tmp_path / 'library')) as library:
        stored = library.store_source(source_path)
        target = library.directory / 'sources' / stored['file']
        target.write_bytes(b'partial previous write')
        assert library.store_source(source_path) == stored
        assert target.read_bytes() == source_path.read_bytes() == original
        assert not list(target.parent.glob('*.pending'))


def test_failed_source_publish_never_leaves_a_partial_hash_named_file(tmp_path, monkeypatch):
    from pathlib import Path

    source_path = deck_file(tmp_path / 'source.pptx')
    original = source_path.read_bytes()
    replace = Path.replace

    def fail(source, destination):
        if source.suffix == '.pending':
            raise OSError(28, 'disk full')
        return replace(source, destination)

    with closing(Library(tmp_path / 'library')) as library:
        monkeypatch.setattr(Path, 'replace', fail)
        with pytest.raises(OSError):
            library.store_source(source_path)
        assert source_path.read_bytes() == original
        assert not list((library.directory / 'sources').iterdir())
