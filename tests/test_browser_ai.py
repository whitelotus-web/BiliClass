import json
from pathlib import Path
from threading import Event

import pytest

from app.browser_accounts import BrowserAccounts, ProfileBusy, profile_lock
from app.browser_audio import prepare_narration
from app.browser_automation import BrowserProblem, conversation_url, read_record, write_record


def test_profiles_are_separate_persisted_and_deleted_with_session_data(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    first = accounts.add("Cô A", "chrome")
    second = accounts.add("Thầy B", "chrome")
    assert accounts.get()["id"] == first["id"]
    accounts.select(second["id"])
    accounts.options(False, True)
    resumed = BrowserAccounts(tmp_path)
    assert resumed.get()["id"] == second["id"] and not resumed.data["automatic"]
    assert resumed.profile(first["id"]) != resumed.profile(second["id"])
    (resumed.profile(first["id"]) / "session-marker.txt").write_text("fixture only")
    with profile_lock(resumed.root, first["id"]):
        with pytest.raises(ProfileBusy):
            resumed.remove(first["id"])
    resumed.remove(first["id"])
    assert not resumed.profile(first["id"]).exists()
    assert resumed.get()["id"] == second["id"]
    resumed.remove(second["id"])
    with pytest.raises(ValueError, match="Đăng nhập ChatGPT"):
        resumed.get()
    assert resumed.data["active"] == ""
    with pytest.raises(ValueError):
        resumed.profile("../../elsewhere")


def test_plus_priority_changes_after_account_downgrades_to_free(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    plus = accounts.add("Plus")
    free = accounts.add("Free")
    accounts.observe(plus["id"], "plus")
    accounts.observe(free["id"], "free")
    assert accounts.preferred()["id"] == plus["id"]
    accounts.observe(plus["id"], "free", "GPT-6 Luna")
    assert accounts.get(plus["id"])["plan"] == "free"
    assert accounts.get(plus["id"])["model"] == "GPT-6 Luna"
    assert accounts.preferred()["channel"] == "chrome"


def test_account_focus_does_not_override_automatic_paid_priority(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    free, plus = accounts.add("Free"), accounts.add("Plus")
    with pytest.raises(ValueError, match="Đăng nhập"):
        accounts.preferred()
    accounts.observe(free["id"], "free")
    accounts.observe(plus["id"], "plus")
    assert accounts.preferred()["id"] == plus["id"]
    accounts.select(free["id"])
    accounts.observe(plus["id"], "plus", "Thinking")
    assert BrowserAccounts(tmp_path).preferred()["id"] == plus["id"]
    accounts.data["selection"] = "manual"  # Old persisted UI setting is migrated.
    accounts.save()
    assert BrowserAccounts(tmp_path).data["selection"] == "auto"
    accounts.prefer_paid()
    assert accounts.preferred()["id"] == plus["id"]
    accounts.remove(plus["id"])
    assert accounts.preferred()["id"] == free["id"]


def test_unknown_plan_is_not_free_and_downgrade_prefers_remaining_plus(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    first, second, unknown = accounts.add("One"), accounts.add("Two"), accounts.add("Unknown")
    accounts.observe(first["id"], "plus")
    accounts.observe(second["id"], "plus")
    accounts.observe(unknown["id"], "unknown", identity={"name": "Teacher", "email": ""})
    assert accounts.get(unknown["id"])["plan"] == "unknown"
    assert accounts.get(unknown["id"])["name"] == "Teacher"
    assert accounts.get(unknown["id"])["saved_at"]
    accounts.observe(first["id"], "free")
    assert accounts.preferred()["id"] == second["id"]


def test_failed_added_account_is_retryable_without_changing_conversion_preference(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    free, plus = accounts.add("Free"), accounts.add("Plus")
    accounts.observe(free["id"], "free")
    accounts.observe(plus["id"], "plus")
    accounts.select(free["id"])
    new = accounts.add("New account")
    accounts.begin_login(new["id"])
    accounts.login_error(new["id"], "verification", "Fixture verification")
    resumed = BrowserAccounts(tmp_path)
    assert resumed.data["pending_login"] == new["id"]
    assert resumed.preferred()["id"] == plus["id"]
    resumed.observe(new["id"], "plus", identity={"name": "New teacher"})
    assert not resumed.data["pending_login"] and resumed.preferred()["plan"] == "plus"
    resumed.prefer_paid()
    assert resumed.preferred()["plan"] == "plus"


@pytest.mark.parametrize("code", ["login", "verification", "auth_response"])
def test_expired_profile_is_not_used_for_new_jobs_and_recovers_after_login(tmp_path, code):
    accounts = BrowserAccounts(tmp_path)
    plus, free = accounts.add("Plus"), accounts.add("Free")
    accounts.observe(plus["id"], "plus")
    accounts.observe(free["id"], "free")
    accounts.login_error(plus["id"], code, "Session unavailable")
    assert not accounts.get(plus["id"])["ready"] and accounts.preferred()["id"] == free["id"]
    assert not BrowserAccounts(tmp_path).get(plus["id"])["ready"]
    accounts.observe(plus["id"], "plus", identity={"name": "Teacher"})
    assert accounts.preferred()["id"] == plus["id"]
    accounts.login_error(plus["id"], "limit", "Web usage limit")
    assert accounts.get(plus["id"])["ready"]  # Quota is distinct from authentication.


def test_legacy_edge_profile_is_preserved_but_chrome_requires_new_login(tmp_path):
    accounts = BrowserAccounts(tmp_path)
    item = accounts.add("Old profile")
    marker = accounts.profile(item["id"]) / "old-session-marker"
    marker.write_text("Fixture")
    accounts.data["accounts"][0].update(channel="msedge", ready=True, plan="plus")
    accounts.data["accounts"][0].pop("profile_layout")
    accounts.save()
    resumed = BrowserAccounts(tmp_path)
    assert resumed.get()["channel"] == "chrome"
    assert Path(resumed.get()["profile"]) == marker.parent / "chrome"
    assert not resumed.get()["ready"] and resumed.get()["plan"] == "unknown"
    assert resumed.get()["last_error"]["code"] == "browser_changed"
    assert marker.read_text() == "Fixture"
    assert BrowserAccounts(tmp_path).get()["profile"] == resumed.get()["profile"]
    with pytest.raises(ValueError, match="Google Chrome"):
        resumed.add("New profile", "msedge")


def test_job_keeps_conversation_and_refuses_another_account_or_foreign_url(tmp_path):
    job = read_record(tmp_path, "account-a")
    job.update(state="waiting", url="https://chatgpt.com/c/fixture-123")
    write_record(tmp_path, job)
    assert read_record(tmp_path, "account-a") == job
    with pytest.raises(BrowserProblem, match="tài khoản khác"):
        read_record(tmp_path, "account-b")
    job["url"] = "https://chatgpt.com.evil.invalid/c/fixture-123"
    write_record(tmp_path, job)
    with pytest.raises(BrowserProblem, match="không hợp lệ"):
        read_record(tmp_path, "account-a")
    assert conversation_url("https://chatgpt.com/c/fixture-123")
    assert not conversation_url("https://chatgpt.com/c/fixture-123?token=secret")


def test_draft_audio_never_approves_or_plays_and_does_not_discard_other_voice(monkeypatch, tmp_path):
    from app import speech

    calls = []
    units = [{"vi": "Một câu tiếng Việt.", "en": "One English sentence.", "approved": False},
             {"vi": "", "en": "x" * 5001, "approved": False}]

    def synthesize(text, voice, rate, directory):
        calls.append((text, voice, rate, directory))
        if voice == "unavailable":
            raise ValueError("Missing model")

    monkeypatch.setattr(speech, "synthesize", synthesize)
    monkeypatch.setattr(speech, "play", lambda *_: pytest.fail("Draft audio must not play"))
    result = prepare_narration({"profile": {"units": units}}, {"vi": "unavailable", "en": "good", "rate": 0},
                               tmp_path, Event(), lambda _: None)
    assert result == {"complete": 1, "skipped": 1, "failed": 1, "total": 3}
    assert len(calls) == 2 and not any(u["approved"] for u in units)
    cancelled = Event()
    cancelled.set()
    result = prepare_narration({"profile": {"units": units}}, {"vi": "good", "en": "good", "rate": 0},
                               tmp_path, cancelled, lambda _: None)
    assert result["complete"] == 0 and len(calls) == 2


def test_profiles_do_not_enter_library_backup(tmp_path):
    from app.library import Library
    from app.storage import backup_library

    library = Library(tmp_path / "library")
    try:
        accounts = BrowserAccounts(library.directory)
        item = accounts.add("Cô A", "chrome")
        (accounts.profile(item["id"]) / "session-marker").write_text("Not a credential")
        path = backup_library(library.directory, tmp_path / "backup.bcbackup")
        import zipfile
        with zipfile.ZipFile(path) as archive:
            assert not any("browser_ai" in name or "session-marker" in name for name in archive.namelist())
    finally:
        library.close()


# A real Chrome browser against a routed fixture: no ChatGPT account or uploads to the Internet.
FIXTURE = """<!DOCTYPE html><html><head><title>Browser AI fixture</title></head><body>
<button data-testid="accounts-profile-button">Fixture account</button>
<textarea id="prompt-textarea"></textarea><input type="file" multiple id="files"><div id="attached"></div>
<button data-testid="send-button" id="send">Send</button><div id="messages"></div>
<script>
function result() { document.querySelector('#messages').innerHTML = '<div data-message-author-role="assistant"><a href="/backend-api/files/result.pptx">bai-giang-song-ngu.pptx</a></div>'; }
if (location.pathname.startsWith('/c/')) result();
document.querySelector('#files').onchange = (e) => {document.querySelector('#attached').replaceChildren(...[...e.target.files].map(f => {const n=document.createElement('span');n.textContent=f.name;return n;}));};
document.querySelector('#send').onclick = () => {console.log('FIXTURE-SEND:' + JSON.stringify({prompt:document.querySelector('#prompt-textarea').value,files:[...document.querySelector('#files').files].map(f=>f.name)}));history.pushState({},'', '/c/fixture-123');result();};
</script></body></html>"""


@pytest.fixture
def browser_fixture(monkeypatch, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from app import browser_automation as adapter
    from app.chatgpt_handoff import prepare_request

    source = tmp_path / "source.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2)).text = "Một slide để kiểm tra"
    slide.notes_slide.notes_text_frame.text = "VI: Một câu tiếng Việt.\nEN: One English sentence.\nCHECK: Kiểm tra số liệu."
    deck.save(source)
    accounts = BrowserAccounts(tmp_path / "library")
    accounts.add("Fixture account", "chrome")
    config = {"title": "Fixture", "subject": "Toán", "education_level": "THPT", "grade": "11",
              "level": 2, "layout": "split_view", "preset": "standard", "style": "source", "mode": "level"}
    request = prepare_request(tmp_path / "library", config, source)
    original = adapter.open_context
    captured = {"sends": [], "contexts": 0, "challenge": False, "download_failure": False, "html": ""}

    def opening(playwright, account, *, background, cancel=None):
        captured["contexts"] += 1
        context = original(playwright, account, background=background, cancel=cancel)

        def route(req):
            from urllib.parse import urlsplit

            if urlsplit(req.request.url).hostname == "127.0.0.1":
                req.continue_()  # Optional local download server for a large fixture.
                return
            if "/backend-api/files/result-" in req.request.url:
                if captured["download_failure"]:
                    req.fulfill(status=500, content_type="text/plain", body="Fixture download failed")
                elif captured.get("download_url"):
                    req.fulfill(status=302, headers={"Location": captured["download_url"]})
                else:
                    req.fulfill(body=source.read_bytes(), headers={"Content-Type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                                                                 "Content-Disposition": f'attachment; filename="fixture-{captured["contexts"]}.pptx"'})
            else:
                content = '<html><head><title>Just a moment...</title></head><body id="challenge-stage">Verify</body></html>' if captured["challenge"] else captured["html"] or FIXTURE
                content = content.replace("result.pptx", f'result-{captured["contexts"]}.pptx')
                req.fulfill(status=403 if captured["challenge"] else 200, content_type="text/html", body=content)

        context.route("**/*", route)
        def observe_page(page):
            page.on("console", lambda msg: captured["sends"].append(json.loads(msg.text.split(":", 1)[1]))
                    if msg.text.startswith("FIXTURE-SEND:") else None)
        context.on("page", observe_page)
        for page in context.pages:
            observe_page(page)
        return context

    monkeypatch.setattr(adapter, "open_context", opening)
    return adapter, accounts, request, captured, source


def test_real_browser_upload_download_and_repeat_does_not_send_twice(browser_fixture):
    adapter, accounts, request, captured, source = browser_fixture
    result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert Path(result["path"]).read_bytes() == source.read_bytes()
    assert len(captured["sends"]) == 1
    assert captured["sends"][0] == {"files": ["tai-lieu-goc.pptx"], "prompt": request["prompt"]}
    assert "Hai cột" in captured["sends"][0]["prompt"]
    assert read_record(request["folder"], accounts.get()["id"])["state"] == "completed"
    cached = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None)
    assert cached == dict(result, cached=True) and not result["cached"]
    assert len(captured["sends"]) == 1 and captured["contexts"] == 1
    inspection = adapter.inspect_returned_deck(result["path"])
    assert inspection["profile"]["units"][0]["en"] == "One English sentence."
    Path(result["path"]).write_bytes(b"altered")
    with pytest.raises(BrowserProblem, match="bị thay đổi"):
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None)
    assert len(captured["sends"]) == 1


def test_powerpoint_above_old_limit_uploads_downloads_stores_and_reopens(browser_fixture):
    import hashlib
    import shutil
    import zipfile
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Thread

    from app.chatgpt_handoff import external_preview, prepare_request
    from app.importers import MAX_BYTES
    from app.library import Library

    adapter, accounts, request, captured, source = browser_fixture
    # An unused stored media part increases file size without changing native
    # slide text/notes; the fixture server returns these exact bytes.
    with zipfile.ZipFile(source, "a", compression=zipfile.ZIP_STORED) as archive:
        with archive.open("ppt/media/large-fixture.bin", "w") as stream:
            for _ in range(51):
                stream.write(b"\0" * 1024**2)
    assert source.stat().st_size > MAX_BYTES
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    request = prepare_request(accounts.root, request["config"], source)

    class Download(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
            self.send_header("Content-Disposition", 'attachment; filename="large-fixture.pptx"')
            self.send_header("Content-Length", str(source.stat().st_size))
            self.end_headers()
            with source.open("rb") as stream:
                shutil.copyfileobj(stream, self.wfile, 1024**2)

        def log_message(self, *_):
            pass

    # Stream actual download bytes instead of a huge base64 route.fulfill body.
    server = ThreadingHTTPServer(("127.0.0.1", 0), Download)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    captured["download_url"] = f"http://127.0.0.1:{server.server_port}/large-fixture.pptx"
    try:
        result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=30)
        assert hashlib.sha256(Path(result["path"]).read_bytes()).hexdigest() == digest
        assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
        assert len(captured["sends"]) == 1
        cached = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None)
        assert cached["cached"] and captured["contexts"] == 1
        library = Library(accounts.root)
        try:
            stored = library.store_source(result["path"])
            assert stored["sha256"] == digest
            lesson = library.create_external_lesson(request["config"], stored, adapter.inspect_returned_deck(result["path"]))
            preview = external_preview(library.get(lesson["id"]), library.directory)
            assert preview["sha256"] == digest and preview["draft"]
        finally:
            library.close()
        assert captured["sends"][0] == {"files": ["tai-lieu-goc.pptx"], "prompt": request["prompt"]}
        assert read_record(request["folder"], accounts.get()["id"])["state"] == "completed"
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)


def test_filename_button_downloads_real_pptx_and_never_matches_user_source(browser_fixture):
    adapter, accounts, request, captured, source = browser_fixture
    captured["html"] = FIXTURE.replace(
        '<a href="/backend-api/files/result.pptx">bai-giang-song-ngu.pptx</a>',
        '<button aria-label="bai-giang-song-ngu.pptx" onclick="downloadFile()">PPTX</button>')
    captured["html"] = captured["html"].replace('<script>', "<script>function downloadFile() { location.href='/backend-api/files/result.pptx'; }")
    result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert Path(result["path"]).read_bytes() == source.read_bytes() and len(captured["sends"]) == 1
    from playwright.sync_api import sync_playwright
    with sync_playwright() as playwright:
        context = adapter.open_context(playwright, accounts.get(), background=True)
        try:
            page = context.pages[0]
            page.goto(adapter.CHATGPT)
            page.locator('#messages').evaluate("e => { e.innerHTML = '<div data-message-author-role=\"user\"><button aria-label=\"source.pptx\">PPTX</button></div><div data-message-author-role=\"assistant\">No result yet</div>'; }")
            assert adapter.result_link(page) is None
        finally:
            context.close()


def test_file_preview_download_uses_matching_artifact_and_sends_once(browser_fixture):
    adapter, accounts, request, captured, source = browser_fixture
    captured["html"] = FIXTURE.replace(
        '<a href="/backend-api/files/result.pptx">bai-giang-song-ngu.pptx</a>',
        '<button aria-label="bai-giang-song-ngu.pptx" onclick="showFile()">PPTX</button>')
    captured["html"] = captured["html"].replace('<script>', '''<script>
function showFile() {
    const preview = document.createElement('div');
    preview.innerHTML = `<h3>bai-giang-song-ngu.pptx</h3><button aria-label="Download file" onclick="location.href='/backend-api/files/result.pptx'">Download</button>`;
    document.body.appendChild(preview);
    const other = document.createElement('div');
    other.innerHTML = `<h3>other.pptx</h3><button aria-label="Download file" onclick="throw Error('Do not download unrelated artifact')">Download</button>`;
    document.body.appendChild(other);
}
''', 1)
    result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert Path(result["path"]).read_bytes() == source.read_bytes() and len(captured["sends"]) == 1


def test_download_transport_failure_keeps_login_and_resumes_without_another_send(browser_fixture, monkeypatch):
    from playwright.sync_api import TimeoutError

    from app.browser_dispatch import convert_available

    adapter, accounts, request, captured, _ = browser_fixture
    accounts.observe(accounts.get()["id"], "free")
    original = adapter.receive_download
    monkeypatch.setattr(adapter, "receive_download", lambda *_a, **_k: (_ for _ in ()).throw(TimeoutError("Download fixture interruption")))
    with pytest.raises(BrowserProblem) as exc:
        convert_available(accounts.candidates(), accounts.root, request["folder"], Event(), lambda _: None,
                          failed=accounts.login_error, retry_delays=())
    assert exc.value.code == "download" and accounts.get()["ready"]
    assert read_record(request["folder"], accounts.get()["id"])["state"] == "waiting"
    monkeypatch.setattr(adapter, "receive_download", original)
    result = convert_available(accounts.candidates(), accounts.root, request["folder"], Event(), lambda _: None)
    assert Path(result["path"]).is_file() and len(captured["sends"]) == 1


def test_challenge_stops_before_upload_and_can_retry_same_request(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["challenge"] = True
    with pytest.raises(BrowserProblem, match="xác minh"):
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    record = read_record(request["folder"], accounts.get()["id"])
    assert record["state"] == "prepared" and record["error"] == "verification" and not captured["sends"]
    captured["challenge"] = False
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert len(captured["sends"]) == 1


def test_resume_download_uses_conversation_without_resending_input(browser_fixture):
    adapter, accounts, request, captured, source = browser_fixture
    record = read_record(request["folder"], accounts.get()["id"])
    record.update(state="waiting", url="https://chatgpt.com/c/fixture-123")
    write_record(request["folder"], record)
    result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert Path(result["path"]).read_bytes() == source.read_bytes() and not captured["sends"]


def test_uncertain_send_is_not_repeated(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    record = read_record(request["folder"], accounts.get()["id"])
    record["state"] = "submitting"
    write_record(request["folder"], record)
    with pytest.raises(BrowserProblem, match="không gửi trùng"):
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert not captured["sends"]


def test_latest_answer_only_and_lesson_limits_do_not_become_quota_errors(browser_fixture):
    from playwright.sync_api import sync_playwright

    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = FIXTURE.replace('<div id="messages"></div>',
        '<div id="messages"><div data-message-author-role="assistant"><a href="/backend-api/files/old.pptx">old.pptx</a></div>'
        '<div data-message-author-role="assistant">Bài toán giới hạn; hạn mức là một khái niệm trong bài.</div></div>')
    with sync_playwright() as playwright:
        context = adapter.open_context(playwright, accounts.get(), background=True)
        try:
            page = context.pages[0]
            page.goto(adapter.CHATGPT)
            adapter.page_problem(page)
            assert adapter.result_link(page) is None  # Ignore the older PPTX.
            page.locator('#messages').evaluate("node => node.insertAdjacentHTML('beforeend', "
                + json.dumps('<div data-message-author-role="assistant"><a href="/backend-api/files/new.pptx">new.pptx</a></div>') + ")")
            assert adapter.result_link(page, 2).inner_text() == "new.pptx"
            assert adapter.result_link(page, 3) is None
            page.locator('body').evaluate("node => node.insertAdjacentHTML('beforeend', "
                + json.dumps('<div role="alert">You have reached your usage limit. Try again later.</div>') + ")")
            with pytest.raises(BrowserProblem) as failure:
                adapter.page_problem(page)
            assert failure.value.code == "limit"
        finally:
            context.close()


def test_followup_waits_for_new_answer_without_resending_source(browser_fixture):
    adapter, accounts, request, captured, source = browser_fixture
    captured["html"] = FIXTURE[:FIXTURE.index('<script>')] + '''<script>
let sent = 0;
document.querySelector('#files').onchange = (e) => {
    document.querySelector('#attached').innerHTML = [...e.target.files].map(f => `<span>${f.name}</span>`).join('');
};
document.querySelector('#send').onclick = () => {
    console.log('FIXTURE-SEND:' + JSON.stringify({prompt:document.querySelector('#prompt-textarea').value,
                                               files:[...document.querySelector('#files').files].map(f=>f.name)}));
    sent++; history.pushState({},'', '/c/fixture-123');
    document.querySelector('#messages').insertAdjacentHTML('beforeend', '<div data-message-author-role="user">Sent</div>');
    document.querySelector('#files').value = '';
    if (sent === 1) document.querySelector('#messages').insertAdjacentHTML('beforeend',
        '<div data-message-author-role="assistant">Bài giảng về giới hạn và hạn mức. Đây là dàn ý.</div>');
    else setTimeout(() => document.querySelector('#messages').insertAdjacentHTML('beforeend',
        '<div data-message-author-role="assistant"><a href="/backend-api/files/result.pptx">new.pptx</a></div>'), 700);
};</script></body></html>'''
    result = adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=15)
    assert Path(result["path"]).read_bytes() == source.read_bytes()
    assert len(captured["sends"]) == 2 and captured["sends"][1]["prompt"] == adapter.EXPORT_PROMPT
    assert captured["sends"][1]["files"] == []
    record = read_record(request["folder"], accounts.get()["id"])
    assert record["followups"] == 1 and record["assistant_before"] == 1 and record["user_before"] == 1
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None)
    assert len(captured["sends"]) == 2


def test_prompt_removed_during_upload_is_not_sent(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["html"] = FIXTURE.replace("document.querySelector('#files').onchange = (e) => {",
        "document.querySelector('#files').onchange = (e) => {document.querySelector('#prompt-textarea').value='';")
    with pytest.raises(BrowserProblem, match="Ô prompt đã thay đổi"):
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert not captured["sends"]
    assert read_record(request["folder"], accounts.get()["id"])["state"] == "prepared"


def test_upload_event_timeout_still_waits_for_real_attachment_and_sends_once(browser_fixture, monkeypatch):
    from playwright.sync_api import Locator, TimeoutError

    adapter, accounts, request, captured, _ = browser_fixture
    original = Locator.set_input_files
    calls = []

    def delayed_event(self, *args, **kwargs):
        calls.append(1)
        original(self, *args, **kwargs)
        raise TimeoutError("Fixture input event timed out after files were set")

    monkeypatch.setattr(Locator, "set_input_files", delayed_event)
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert len(calls) == len(captured["sends"]) == 1
    assert captured["sends"][0]["files"] == ["tai-lieu-goc.pptx"]


def test_upload_timeout_without_rendered_attachment_never_sends(browser_fixture, monkeypatch):
    from playwright.sync_api import Locator, TimeoutError

    adapter, accounts, request, captured, _ = browser_fixture
    monkeypatch.setattr(Locator, "set_input_files", lambda *_a, **_k: (_ for _ in ()).throw(TimeoutError("No attachment")))
    original = adapter.wait_upload
    monkeypatch.setattr(adapter, "wait_upload", lambda *args: original(*args, timeout=.2))
    with pytest.raises(BrowserProblem) as error:
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert error.value.code == "upload" and not captured["sends"]
    assert read_record(request["folder"], accounts.get()["id"])["state"] == "prepared"


def test_file_upload_skips_image_only_input_and_accepts_editor_paragraph_spacing(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["html"] = FIXTURE.replace('<textarea id="prompt-textarea"></textarea>',
        '<input type="file" id="camera" accept="image/*" multiple>'
        '<div id="prompt-textarea" contenteditable="true"></div>')
    captured["html"] = captured["html"].replace("document.querySelector('#prompt-textarea').value",
                                               "document.querySelector('#prompt-textarea').innerText")
    captured["html"] = captured["html"].replace('<script>', '''<script>
document.querySelector('#camera').onchange = () => { throw Error('Must not upload PPTX as photo'); };
document.querySelector('#prompt-textarea').oninput = (e) => {
    const lines = e.target.innerText.split('\\n');
    e.target.replaceChildren(...lines.map(text => { const p = document.createElement('p'); p.textContent = text || ' '; return p; }));
};
''', 1)
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    import re
    assert len(captured["sends"]) == 1 and captured["sends"][0]["files"] == ["tai-lieu-goc.pptx"]
    assert re.sub(r"\s+", " ", captured["sends"][0]["prompt"]).strip() == re.sub(r"\s+", " ", request["prompt"]).strip()


def test_web_renamed_attachment_is_recognized_after_app_draft_cleanup(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["html"] = FIXTURE.replace('<div id="attached"></div>',
        '<div id="attached"><div id="stale" aria-label="tai-lieu-goc(7).pptx">'
        '<button aria-label="Remove file 1: tai-lieu-goc(7).pptx" '
        'onclick="this.parentElement.remove()">Remove old draft</button></div></div>')
    captured["html"] = captured["html"].replace('n.textContent=f.name;',
        '''n.setAttribute('aria-label', f.name.replace('.pptx','(8).pptx'));
        n.textContent=f.name.replace('.pptx','(8)');
        if (document.querySelector('#stale')) throw Error('Old draft was not cleared');''')
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert len(captured["sends"]) == 1 and captured["sends"][0]["files"] == ["tai-lieu-goc.pptx"]


def test_active_profile_cannot_overwrite_another_running_job(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    record = read_record(request["folder"], accounts.get()["id"])
    record.update(state="waiting", url="https://chatgpt.com/c/fixture-123")
    write_record(request["folder"], record)
    with profile_lock(accounts.root, accounts.get()["id"]):
        with pytest.raises(ProfileBusy):
            adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None)
    assert read_record(request["folder"], accounts.get()["id"]) == record and not captured["contexts"]


def models_fixture(plan):
    profile = f'<button data-testid="accounts-profile-button"><div>Fixture account</div><div>{plan}</div></button>'
    choices = '''<button data-testid="model-switcher-dropdown-button" id="model-picker" onclick="document.querySelector('#model-menu').hidden=false">GPT-6.1 Sol</button>
    <div role="menu" id="model-menu" hidden>
    <button role="menuitemradio" onclick="choose(this)">GPT-6 Luna</button>
    <button role="menuitemradio" onclick="choose(this)">GPT-6.1 Sol</button>
    <button role="menuitemradio" %s onclick="choose(this)">Astra</button>
    <button role="menuitemradio" onclick="throw Error('Must not buy a plan')">Astra · Upgrade to Plus</button></div>
    <script>function choose(item) {item.setAttribute('aria-checked','true');document.querySelector('#model-picker').textContent=item.textContent;document.querySelector('#model-menu').hidden=true;}</script>''' % ("disabled" if plan == "Free" else "")
    if plan == "Free":
        choices = choices.replace('onclick="choose(this)">GPT-6.1 Sol', 'disabled onclick="choose(this)">GPT-6.1 Sol')
    return FIXTURE.replace('<button data-testid="accounts-profile-button">Fixture account</button>', profile + choices).replace(
        "files:[...document.querySelector('#files').files]", "model:document.querySelector('#model-picker').textContent,files:[...document.querySelector('#files').files]")


def test_downgrade_selects_free_model_before_upload_and_keeps_cancelled_request(browser_fixture):
    from app.chatgpt_handoff import prepare_request

    adapter, accounts, request, captured, source = browser_fixture
    captured["html"] = models_fixture("Plus")
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10,
                    observed=accounts.observe)
    assert captured["sends"][0]["model"] == "Astra"
    assert accounts.get()["plan"] == "plus"
    captured["html"] = models_fixture("Free")
    next_request = prepare_request(accounts.root.parent, request["config"], source)
    cancel = Event()

    def stop_before_upload(account_id, plan, model):
        accounts.observe(account_id, plan, model)
        cancel.set()

    with pytest.raises(BrowserProblem, match="Đã dừng"):
        adapter.convert(accounts.get(), accounts.root, next_request["folder"], cancel, lambda _: None, timeout=10,
                        observed=stop_before_upload)
    assert accounts.get()["plan"] == "free" and accounts.get()["model"] == "GPT-6 Luna"
    record = read_record(next_request["folder"], accounts.get()["id"])
    assert record["plan"] == "free" and record["model"] == "GPT-6 Luna" and record["state"] == "prepared"
    assert record["error"] == "cancelled" and len(captured["sends"]) == 1


def test_login_saves_automatically_without_finish_button(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = models_fixture("Plus")
    opening = adapter.open_context

    def login_browser(playwright, account, *, background, cancel=None):
        return opening(playwright, account, background=True, cancel=cancel)  # The test remains unobtrusive.

    monkeypatch.setattr(adapter, "open_context", login_browser)
    result = adapter.login_observed(accounts.get(), accounts.root, Event(), lambda _: None)
    assert result["ready"] and result["plan"] == "plus" and not captured["sends"]
    assert captured["contexts"] == 2  # Confirm the saved session in a new Chrome process.
    with profile_lock(accounts.root, accounts.get()["id"]):
        pass  # The automatic login closed its browser and released the profile.


def test_guest_composer_with_profile_control_is_not_a_saved_login(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = FIXTURE.replace("<textarea", '<button data-testid="login-button">Log in</button><textarea')
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    with pytest.raises(BrowserProblem, match="Hết thời gian"):
        adapter.login_observed(accounts.get(), accounts.root, Event(), lambda _: None, timeout=2)
    assert not accounts.get()["ready"] and not captured["sends"]


@pytest.fixture
def plain_login(browser_fixture, monkeypatch):
    from app import chrome_session

    adapter, accounts, _, captured, _ = browser_fixture
    native = {"opened": [], "closed": 0, "running": False}

    class Plain:
        def __init__(self, account, url):
            native["opened"].append((account["id"], url))

        def running(self):
            return native["running"]

        def finish(self):
            native["running"] = False

        def close(self):
            native["closed"] += 1
            native["running"] = False

    monkeypatch.setattr(chrome_session, "PlainChromeLogin", Plain)
    return adapter, accounts, captured, native


def test_plain_login_attaches_only_after_user_confirmation_and_saves_once(plain_login):
    adapter, accounts, captured, native = plain_login
    native["running"] = True
    finish = Event()
    phases = []

    def confirmation(waiting):
        phases.append(waiting)
        if waiting:
            assert captured["contexts"] == 0 and not captured["sends"]
            finish.set()  # Simulated human confirmation, not web form automation.

    result = adapter.login(accounts.get(), accounts.root, Event(), lambda _: None,
        finish=finish, awaiting_confirmation=confirmation, timeout=10)
    assert result["ready"] and phases == [True, False] and captured["contexts"] == 1
    assert native["opened"] == [(accounts.get()["id"], adapter.CHATGPT)] and native["closed"] == 1
    assert not captured["sends"]


def test_plain_login_cancel_does_not_attach_or_mark_account_ready(plain_login):
    adapter, accounts, captured, native = plain_login
    native["running"] = True
    cancel = Event()
    with pytest.raises(BrowserProblem) as error:
        adapter.login(accounts.get(), accounts.root, cancel, lambda _: None,
            awaiting_confirmation=lambda waiting: cancel.set() if waiting else None)
    assert error.value.code == "cancelled" and captured["contexts"] == 0
    assert native["closed"] == 1 and not accounts.get()["ready"]
    with profile_lock(accounts.root, accounts.get()["id"]):
        pass


def test_plain_login_closed_before_sign_in_is_not_connected(plain_login):
    adapter, accounts, captured, _ = plain_login
    captured["html"] = FIXTURE.replace('<textarea', '<button data-testid="login-button">Log in</button><textarea')
    with pytest.raises(BrowserProblem) as error:
        adapter.login(accounts.get(), accounts.root, Event(), lambda _: None)
    assert error.value.code == "login" and not accounts.get()["ready"] and not captured["sends"]


def test_authentication_content_type_error_is_not_session_timeout(plain_login):
    adapter, accounts, captured, _ = plain_login
    captured["html"] = '<html><h1>Oops, an error occurred!</h1><p>Route Error (400 Invalid content type: text/html; charset=UTF-8)</p></html>'
    with pytest.raises(BrowserProblem) as error:
        adapter.login(accounts.get(), accounts.root, Event(), lambda _: None)
    assert error.value.code == "auth_response" and "400" in str(error.value)
    assert not accounts.get()["ready"] and not captured["sends"]


def test_readonly_health_probe_reads_current_plan_without_upgrade_or_sending(browser_fixture):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = '''<html><title>Fixture</title>
    <button data-testid="accounts-profile-button" onclick="document.querySelector('#menu').hidden=false">TO</button>
    <textarea id="prompt-textarea"></textarea><button data-testid="send-button">Send</button>
    <div id="menu" role="menu" hidden><span>Upgrade to Plus</span><span data-testid="account-name">Fixture teacher</span>
    <button role="menuitem" onclick="document.querySelector('#dialog').hidden=false">Settings</button></div>
    <div id="dialog" role="dialog" hidden><button role="tab">Account</button>
    <div data-testid="current-plan">Current plan: Free</div><button>Upgrade to Plus</button></div></html>'''
    result = adapter.check_session(accounts.get(), accounts.root, Event(), lambda _: None)
    assert result["ready"] and result["plan"] == "free" and result["name"] == "Fixture teacher"
    assert not result["quota_limited"] and not captured["sends"]


def test_health_probe_limit_keeps_authentication_and_does_not_upload(browser_fixture):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = FIXTURE + '<div role="alert">You have reached your usage limit. Try again later.</div>'
    result = adapter.check_session(accounts.get(), accounts.root, Event(), lambda _: None)
    assert result["ready"] and result["quota_limited"] and not captured["sends"]


def test_health_probe_reports_lost_session_without_opening_human_login(browser_fixture):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["html"] = '<html><title>Fixture</title><button data-testid="login-button">Log in</button></html>'
    with pytest.raises(BrowserProblem) as error:
        adapter.check_session(accounts.get(), accounts.root, Event(), lambda _: None)
    assert error.value.code == "login" and captured["contexts"] == 1 and not captured["sends"]


def test_slow_first_document_keeps_login_alive_and_can_finish(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    opening, navigating = adapter.open_context, adapter.navigate
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    navigations = []

    def slow_navigation(page, url, cancel):
        navigating(page, url, cancel)
        navigations.append(url)
        if len(navigations) == 1:
            raise BrowserProblem("network", "Fixture slow first document")

    monkeypatch.setattr(adapter, "navigate", slow_navigation)
    messages = []
    result = adapter.login_observed(accounts.get(), accounts.root, Event(), messages.append, timeout=10)
    assert result["ready"] and captured["contexts"] == 2 and not captured["sends"]
    assert any("tải chậm" in message for message in messages)


def test_slow_first_document_waits_for_user_cancel_without_false_success(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    monkeypatch.setattr(adapter, "navigate", lambda *_: (_ for _ in ()).throw(BrowserProblem("network", "Fixture timeout")))
    cancel = Event()

    def progress(message):
        if "tải chậm" in message:
            cancel.set()

    result = adapter.login_observed(accounts.get(), accounts.root, cancel, progress, timeout=10)
    assert not result["ready"] and cancel.is_set() and captured["contexts"] == 1
    assert not accounts.get()["ready"] and not captured["sends"]


def test_visible_login_is_not_ready_if_saved_chrome_session_cannot_reopen(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))

    def progress(message):
        if "kiểm tra phiên" in message:
            captured["html"] = FIXTURE.replace('<textarea', '<button data-testid="login-button">Log in</button><textarea')

    with pytest.raises(BrowserProblem) as error:
        adapter.login_observed(accounts.get(), accounts.root, Event(), progress, timeout=10)
    assert error.value.code == "login"
    assert captured["contexts"] == 2 and not accounts.get()["ready"] and not captured["sends"]


def test_login_verification_loop_stops_with_recoverable_reason_and_releases_profile(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["challenge"] = True
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    messages = []
    with pytest.raises(BrowserProblem) as failure:
        adapter.login_observed(accounts.get(), accounts.root, Event(), messages.append, timeout=10, verification_timeout=.3)
    assert failure.value.code == "verification"
    assert "Cloudflare" in str(failure.value)
    assert any("xác minh" in message for message in messages)
    assert not accounts.get()["ready"] and not captured["sends"] and captured["contexts"] == 1
    accounts.login_error(accounts.get()["id"], failure.value.code, str(failure.value))
    assert BrowserAccounts(accounts.root.parent).get()["last_error"]["code"] == "verification"
    with profile_lock(accounts.root, accounts.get()["id"]):
        pass


def test_login_can_finish_after_user_resolves_fixture_challenge(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    # The fixture simulates a human completing the challenge, not an automated
    # click or a request to any real challenge endpoint.
    captured["html"] = '<html><head><title>Just a moment...</title></head><body>' + (
        '<script>setTimeout(() => { document.title="ChatGPT"; document.body.innerHTML='
        + json.dumps('<button data-testid="accounts-profile-button">Teacher<br>Free</button><textarea id="prompt-textarea"></textarea>')
        + '; }, 700)</script></body></html>')
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    accounts.login_error(accounts.get()["id"], "verification", "Fixture")
    def progress(message):
        if "kiểm tra phiên" in message:
            # The mock verification has completed; the persisted fixture
            # session returns directly to signed-in controls on next launch.
            captured["html"] = models_fixture("Free")

    result = adapter.login_observed(accounts.get(), accounts.root, Event(), progress, timeout=10, verification_timeout=5)
    assert result["ready"] and result["plan"] == "free"
    accounts.observe(accounts.get()["id"], result["plan"], identity=result)
    assert "last_error" not in accounts.get() and not captured["sends"]


def test_cancel_during_verification_is_not_success_or_unnecessary_timeout(browser_fixture, monkeypatch):
    adapter, accounts, _, captured, _ = browser_fixture
    captured["challenge"] = True
    opening = adapter.open_context
    monkeypatch.setattr(adapter, "open_context", lambda playwright, account, **_: opening(playwright, account, background=True))
    cancel = Event()

    def on_progress(message):
        if "xác minh" in message:
            cancel.set()

    result = adapter.login_observed(accounts.get(), accounts.root, cancel, on_progress, timeout=10, verification_timeout=5)
    assert not result["ready"] and cancel.is_set() and not captured["sends"]


def test_upgrade_labels_are_not_plan_or_model_entitlements():
    from app.browser_capabilities import model_priority, plan_from_text

    assert plan_from_text("Upgrade to Plus") == "unknown"
    assert plan_from_text("ChatGPT Free\nUpgrade to Plus") == "free"
    assert model_priority("Astra · Upgrade to Plus") is None
    assert model_priority("GPT-6.1 Sol") > model_priority("GPT-6 Sol") > model_priority("GPT-6 Luna")


@pytest.mark.parametrize("composer_label", ["Ask ChatGPT", "Hỏi ChatGPT"])
def test_current_web_composer_saves_real_profile_identity_and_sends_full_prompt(browser_fixture, composer_label):
    adapter, accounts, request, captured, _ = browser_fixture
    content = models_fixture("Free").replace("<head>", '<head><meta charset="utf-8">').replace(
        "<div>Fixture account</div>", "<div>Loading profile</div><div>Fixture account</div>")
    content = content.replace('<textarea id="prompt-textarea"></textarea>',
        f'<form><div contenteditable="true" role="textbox" aria-label="{composer_label}" id="modern-composer"></div>'
        '<button type="button" aria-pressed="false" onclick="this.setAttribute(\'aria-pressed\',\'true\')">Think</button>')
    content = content.replace('<button data-testid="send-button"', '<button type="button" data-testid="send-button"')
    content = content.replace('<div id="messages">', '</form><div id="messages">')
    content = content.replace("document.querySelector('#prompt-textarea').value",
                              "document.querySelector('#modern-composer').innerText")
    captured["html"] = content
    result = adapter.check_session(accounts.get(), accounts.root, Event(), lambda _: None)
    assert result["ready"] and result["plan"] == "free" and result["name"] == "Fixture account"
    assert not captured["sends"], "Checking the saved session does not send a prompt"
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert len(captured["sends"]) == 1
    assert " ".join(captured["sends"][0]["prompt"].split()) == " ".join(request["prompt"].split())
    record = read_record(request["folder"], accounts.get()["id"])
    assert record["model"] == "GPT-6 Luna" and record["reasoning"] == "Đã bật suy luận trên web"


def test_highest_model_must_be_confirmed_before_any_upload_or_send(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["html"] = models_fixture("Plus").replace(
        'onclick="choose(this)">Astra',
        'onclick="document.querySelector(\'#model-menu\').hidden=true">Astra', 1)
    with pytest.raises(BrowserProblem) as error:
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert error.value.code == "model" and not captured["sends"]
    record = read_record(request["folder"], accounts.get()["id"])
    assert record["error"] == "model" and record["state"] == "prepared"


def test_model_with_generic_heading_is_verified_in_reopened_visible_menu(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    captured["html"] = models_fixture("Plus").replace(
        "document.querySelector('#model-picker').textContent=item.textContent;", "")
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    record = read_record(request["folder"], accounts.get()["id"])
    assert record["model"] == "Astra" and len(captured["sends"]) == 1


def test_free_think_toggle_and_account_priority_do_not_change_a_sent_job(browser_fixture):
    adapter, accounts, request, captured, _ = browser_fixture
    free_id = accounts.get()["id"]
    accounts.observe(free_id, "free")
    captured["html"] = models_fixture("Free").replace('<textarea id="prompt-textarea">',
        '<button aria-pressed="false" onclick="this.setAttribute(\'aria-pressed\',\'true\')">Think</button>'
        '<textarea id="prompt-textarea">')
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    record = read_record(request["folder"], free_id)
    assert record["reasoning"] == "Đã bật suy luận trên web"
    plus = accounts.add("New Plus")
    accounts.observe(plus["id"], "plus")
    assert accounts.preferred()["id"] == plus["id"]
    with pytest.raises(BrowserProblem, match="tài khoản khác"):
        adapter.convert(accounts.preferred(), accounts.root, request["folder"], Event(), lambda _: None)
    before = captured["contexts"]
    adapter.convert(accounts.get(free_id), accounts.root, request["folder"], Event(), lambda _: None)
    assert len(captured["sends"]) == 1 and captured["contexts"] == before


@pytest.mark.parametrize("state", ["aria-pressed", "data-state"])
def test_reasoning_toggle_is_idempotent_and_verified_from_live_controls(browser_fixture, state):
    from playwright.sync_api import sync_playwright

    from app.browser_capabilities import select_reasoning

    adapter, accounts, _request, captured, _source = browser_fixture
    off, on = ("false", "true") if state == "aria-pressed" else ("off", "on")
    captured["html"] = FIXTURE.replace('<textarea id="prompt-textarea">',
        f'<button id="think" {state}="{off}" onclick="window.clicks=(window.clicks||0)+1;this.setAttribute(\'{state}\',\'{on}\')">Think</button>'
        '<textarea id="prompt-textarea">')
    with sync_playwright() as playwright:
        context = adapter.open_context(playwright, accounts.get(), background=True)
        try:
            page = context.pages[0]
            page.goto(adapter.CHATGPT)
            assert select_reasoning(page) == "Đã bật suy luận trên web"
            assert select_reasoning(page) == "Đã bật suy luận trên web"
            assert page.evaluate("window.clicks") == 1 and not captured["sends"]
        finally:
            context.close()


def test_strongest_enabled_effort_is_selected_before_sending(browser_fixture):
    adapter, accounts, request, captured, _source = browser_fixture
    controls = '''<button id="effort" onclick="document.querySelector('#efforts').hidden=false">Standard</button>
    <div id="efforts" role="menu" hidden>
      <button role="menuitemradio" aria-checked="false" onclick="document.querySelector('#effort').textContent='High';this.setAttribute('aria-checked','true');this.parentNode.hidden=true">High</button>
      <button role="menuitemradio" aria-checked="false" onclick="document.querySelector('#effort').textContent='Extended';this.setAttribute('aria-checked','true');this.parentNode.hidden=true">Extended</button>
      <button role="menuitemradio" disabled>Extra high</button>
      <button role="menuitemradio" onclick="throw Error('Must not upgrade')">Max · Upgrade to Plus</button>
    </div>'''
    captured["html"] = FIXTURE.replace('<textarea id="prompt-textarea">', controls + '<textarea id="prompt-textarea">')
    adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert read_record(request["folder"], accounts.get()["id"])["reasoning"] == "Suy luận trên web: Extended"
    assert len(captured["sends"]) == 1


def test_free_tools_menu_can_enable_think_without_sending_a_chat(browser_fixture):
    from playwright.sync_api import sync_playwright

    from app.browser_capabilities import select_reasoning

    adapter, accounts, _request, captured, _source = browser_fixture
    controls = '''<button onclick="document.querySelector('#tools').hidden=false">Tools</button>
    <div id="tools" role="menu" hidden><button role="menuitemcheckbox" onclick="document.querySelector('#think').hidden=false;this.parentNode.hidden=true">Think</button></div>
    <button id="think" aria-pressed="true" hidden onclick="throw Error('Must not disable Think')">Think</button>'''
    captured["html"] = FIXTURE.replace('<textarea id="prompt-textarea">', controls + '<textarea id="prompt-textarea">')
    with sync_playwright() as playwright:
        context = adapter.open_context(playwright, accounts.get(), background=True)
        try:
            page = context.pages[0]
            page.goto(adapter.CHATGPT)
            assert select_reasoning(page) == "Đã bật suy luận trên web" and not captured["sends"]
        finally:
            context.close()


def test_unconfirmed_reasoning_selection_stops_before_upload_or_send(browser_fixture):
    adapter, accounts, request, captured, _source = browser_fixture
    captured["html"] = FIXTURE.replace('<textarea id="prompt-textarea">',
        '<button aria-pressed="false">Think</button><textarea id="prompt-textarea">')
    with pytest.raises(BrowserProblem) as exc:
        adapter.convert(accounts.get(), accounts.root, request["folder"], Event(), lambda _: None, timeout=10)
    assert exc.value.code == "model" and not captured["sends"]
    assert read_record(request["folder"], accounts.get()["id"])["state"] == "prepared"


def test_lesson_buttons_and_generic_retry_notice_are_not_effort_or_quota(browser_fixture):
    from playwright.sync_api import sync_playwright

    from app.browser_capabilities import select_reasoning

    adapter, accounts, _request, captured, _source = browser_fixture
    captured["html"] = FIXTURE.replace('<div id="messages">',
        '<div role="alert">Something went wrong. Try again later.</div><div id="messages">'
        '<div data-message-author-role="assistant"><button aria-pressed="false" onclick="throw Error(\'Must not click lesson text\')">Think</button></div>')
    with sync_playwright() as playwright:
        context = adapter.open_context(playwright, accounts.get(), background=True)
        try:
            page = context.pages[0]
            page.goto(adapter.CHATGPT)
            adapter.page_problem(page)
            assert select_reasoning(page) == "Theo tùy chọn web đang có" and not captured["sends"]
        finally:
            context.close()
