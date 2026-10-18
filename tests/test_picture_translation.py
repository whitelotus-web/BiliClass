import zipfile
from threading import Event

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from app.native_bilingual import assemble, extract
from app.picture_ocr import extract_pictures
from app.picture_translation import apply_picture_translations, background


def source(tmp_path, *, busy=False):
    image = Image.new("RGB", (1000, 600), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (100, 100), "Vận tốc của vật", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 40), fill="black"
    )
    if busy:
        # A graph-like background must never be painted over as a plain label.
        for x in range(0, 1000, 6):
            draw.line((x, 0, x, 600), fill="navy", width=2)
    picture = tmp_path / "picture.png"
    image.save(picture)
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(0.5), Inches(0.5), width=Inches(12), height=Inches(6))
    path = tmp_path / "source.pptx"
    deck.save(path)
    return path


def recognizer(files, cancel, checkpoint):
    rows = [
        {
            "text": "Vận tốc của vật",
            "confidence": 0.99,
            "verified": True,
            "points": [[0.10, 0.16], [0.40, 0.16], [0.40, 0.24], [0.10, 0.24]],
        },
        {
            "text": "Chữ không rõ",
            "confidence": 0.60,
            "verified": False,
            "points": [[0.6, 0.3], [0.8, 0.3], [0.8, 0.35], [0.6, 0.35]],
        },
        {
            "text": "a = -ω²x",
            "confidence": 0.99,
            "verified": True,
            "points": [[0.1, 0.7], [0.4, 0.7], [0.4, 0.75], [0.1, 0.75]],
        },
    ]
    result = {key: rows for key, _ in files}
    for key, rows in result.items():
        checkpoint(key, rows)
    return result


def test_ocr_cache_keeps_boxes_and_rejects_uncertain_math_and_rotated_pictures(tmp_path):
    path = source(tmp_path)
    result = extract_pictures(path, tmp_path / "data", Event(), recognizer=recognizer)
    assert len(result["rows"]) == 1 and result["warnings"]
    again = extract_pictures(
        path,
        tmp_path / "data",
        Event(),
        recognizer=lambda *_: (_ for _ in ()).throw(AssertionError("cache miss")),
    )
    assert again == result
    deck = Presentation(path)
    deck.slides[0].shapes[0].rotation = 15
    deck.save(path)
    assert not extract_pictures(path, tmp_path / "data", Event(), recognizer=recognizer)["rows"]


def output(tmp_path, mode, busy=False):
    path = source(tmp_path, busy=busy)
    data = extract_pictures(path, tmp_path / "data", Event(), recognizer=recognizer)
    manifest = extract(path, mode, preserve_design=True, same_slide=mode != "english_only")
    manifest["note_slides"] = [1]
    target = tmp_path / "translated.pptx"
    assemble(path, target, manifest, {}, mode)
    result = apply_picture_translations(
        target, manifest, data["rows"], {data["rows"][0]["id"]: "The velocity of the object"}, mode, Event()
    )
    return path, target, result


def test_english_overlay_covers_source_text_preserves_image_and_adds_notes(tmp_path):
    original, target, result = output(tmp_path, "english_only")
    assert len(result["applied"]) == 1
    deck = Presentation(target)
    text = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert text.fill.type is not None and text.text == "The velocity of the object"
    assert text.left / 12700 < 0.5 * 72 + 0.1 * 12 * 72
    assert "VI: Vận tốc của vật" in deck.slides[0].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(target) as after:
        for name in before.namelist():
            if name.startswith("ppt/media/"):
                assert before.read(name) == after.read(name)


def test_bilingual_caption_does_not_hide_the_original_label(tmp_path):
    _, target, result = output(tmp_path, "parallel_columns")
    assert result["applied"]
    original = (0.5 * 72 + 0.1 * 12 * 72, 0.5 * 72 + 0.16 * 6 * 72, 0.3 * 12 * 72, 0.08 * 6 * 72)
    from app.layout_regions import intersection

    assert intersection(original, result["applied"][0]["box"]) == 0
    assert "The velocity" in next(s.text for s in Presentation(target).slides[0].shapes if s.has_text_frame)


def test_busy_image_is_not_covered_and_missing_model_does_not_claim_translation(tmp_path, monkeypatch):
    _, _, result = output(tmp_path, "english_only", busy=True)
    assert not result["applied"] and result["warnings"]
    monkeypatch.setattr("app.local_ocr.model_paths", lambda: {})
    result = extract_pictures(tmp_path / "source.pptx", tmp_path / "no-model", Event())
    assert not result["rows"] and result["warnings"]


def test_background_sampling_handles_crop_without_modifying_pixels():
    image = Image.new("RGB", (200, 100), "#fbf6eb")
    before = np.array(image)
    assert background(image, (0, 0, 200, 100), (20, 20, 100, 40)) == "FBF6EB"
    assert background(image, (0, 0, 200, 100), (-1, 0, 100, 40)) is None
    assert np.array_equal(before, np.array(image))


def test_two_readers_must_agree_on_letters_before_accent_correction():
    from app.picture_ocr import crosscheck

    assert crosscheck("Vân tóc cúa vât", 0.95, "Vận tốc của vật", 0.90) == ("Vận tốc của vật", True)
    assert crosscheck("Vận tốc của vật", 0.99, "Vận tốc của đất", 0.99)[1] is False
    assert crosscheck("Vận tốc của vật", 0.99, "Vận tốc của vật", 0.50)[1] is False


def test_google_pipeline_sends_image_text_once_and_saves_editable_overlay(tmp_path, monkeypatch):
    from app import picture_ocr
    from app.conversion_input import creation_config
    from app.conversion_progress import ConversionProgress
    from app.google_conversion import convert, prepare
    from app.google_translation import KeyStore, Translator

    path = source(tmp_path)
    data = tmp_path / "library"
    recognized = extract_pictures(path, data, Event(), recognizer=recognizer)
    monkeypatch.setattr(picture_ocr, "extract_pictures", lambda *_args, **_kwargs: recognized)
    KeyStore(data).add("AIza" + "a" * 35)
    sent = []

    def transport(texts, *_):
        sent.extend(texts)
        return ["The velocity of the object" for _ in texts]

    config = creation_config("Chữ trong ảnh", "Vật lý", "THPT", "11", "english_only", "standard", "source")
    job = prepare(data, config, path)
    lesson, preview = convert(
        data,
        job,
        Event(),
        ConversionProgress(job["folder"], lambda _: None),
        translator=Translator(data, transport=transport),
    )
    assert len(sent) == 1 and "Vận tốc" in sent[0]
    assert any(
        "The velocity" in s.text for s in Presentation(preview["path"]).slides[0].shapes if s.has_text_frame
    )
    assert any("Vận tốc" in unit["vi"] for unit in lesson["segments"])
    from app.lesson_narration import make_plan
    plan, _ = make_plan(lesson, 'both', {'voice': 'vieneu:Mai Anh', 'rate': 0})
    assert plan['total'] == 0, 'A slide containing only OCR labels must not synthesize those labels'


def test_image_labels_are_silent_without_losing_identical_body_text_or_manual_script(tmp_path):
    import copy
    from contextlib import closing

    from app.chatgpt_handoff import inspect_returned_deck
    from app.conversion_input import creation_config
    from app.lesson_narration import signature
    from app.library import Library
    from app.narration_script import prepare, refresh_speech, segments

    _, target, _ = output(tmp_path, 'english_only')
    deck = Presentation(target)
    slide = deck.slides[0]
    body = slide.shapes.add_textbox(Inches(.5), Inches(6.8), Inches(12), Inches(.4))
    body.text = 'The velocity of the object'
    notes = slide.notes_slide.notes_text_frame
    notes.text = ('BILICLASS_NOTES:\nVI: Vận tốc của vật\nEN: The velocity of the object\n'
                  'VI: Vận tốc của vật\nEN: The velocity of the object')
    # Legacy files lack the explicit description added by the new exporter.
    overlay = next(shape for shape in slide.shapes if shape.name == 'BCL_EN image translation')
    overlay._element.xpath('./p:nvSpPr/p:cNvPr')[0].attrib.pop('descr', None)
    overlay_id = overlay.shape_id
    deck.save(target)
    original = target.read_bytes()
    inspection = inspect_returned_deck(target)
    page = inspection['slide_speech']['slides'][0]
    assert [s['english_text'] for s in page['englishSpeechSegments'] if s['status'] == 'ready'] == [body.text]
    with closing(Library(tmp_path / 'library')) as library:
        stored = library.store_source(target)
        lesson = library.create_external_lesson(creation_config('Fixture', 'Vật lý', 'THPT', '11',
                                                                'english_only', 'standard', 'source'), stored, inspection)
        assert segments(lesson)[0]['vi'] == 'Vận tốc của vật'
        assert segments(lesson)[0]['en'] == 'The velocity of the object'
        assert lesson['segments'][0]['vi'].count('Vận tốc') == 2, 'Lesson content must stay intact'
        # Recover old persisted speech even when classification had already run.
        lesson.pop('narration_image_policy')
        lesson.pop('narration_image_text')
        page = lesson['slide_speech']['slides'][0]
        page.pop('imageText')
        page['englishSpeechSegments'].append({'id': 'old-ocr', 'english_text': body.text, 'status': 'ready',
                                              'source': {'shape_id': overlay_id}, 'visibility': {'kind': 'always'}})
        lesson['source_slide_speech'] = copy.deepcopy(lesson['slide_speech'])
        old_signature = signature(lesson)
        recovered = prepare(library, lesson)
        assert signature(recovered) != old_signature, 'Old recordings containing OCR must become stale'
        assert not any(s['id'] == 'old-ocr' for s in recovered['slide_speech']['slides'][0]['englishSpeechSegments'])
        assert segments(recovered)[0]['vi'] == 'Vận tốc của vật'
        identity = recovered['segments'][0]['id']
        recovered['narration_script'] = {identity: {'vi': 'Lời giáo viên nhập.', 'en': body.text, 'manual': True}}
        refresh_speech(recovered)
        assert segments(recovered)[0]['en'] == body.text
        assert segments(recovered)[0]['vi'] == 'Lời giáo viên nhập.'
    from app.voice_import import import_presentation
    imported, _ = import_presentation(tmp_path / 'voice-import', target)
    assert segments(imported)[0]['vi'] == 'Vận tốc của vật'
    assert segments(imported)[0]['en'] == 'The velocity of the object'
    assert target.read_bytes() == original


def test_builtin_template_keeps_ocr_text_visible_but_excludes_it_from_voice(tmp_path):
    from app.chatgpt_handoff import inspect_returned_deck
    from app.universal_templates import asset, build

    target = tmp_path / 'template.pptx'
    build(target, asset('standard'), {'title': 'Bài học', 'english_title': 'Lesson',
          'subject': 'Vật lý', 'grade': '11', 'conversion_format': 'parallel_columns', 'preset': 'standard'},
          [('Slide 1', 'Bài học', 'Lesson'), ('Slide 1', 'Năng lượng được bảo toàn.', 'Energy is conserved.'),
           ('Slide 1 · chữ trong ảnh', 'Nhãn nằm trong hình', 'Label inside the picture')], Event())
    deck = Presentation(target)
    assert any('Label inside the picture' in s.text for page in deck.slides for s in page.shapes if s.has_text_frame)
    inspection = inspect_returned_deck(target)
    spoken = '\n'.join(s['english_text'] for page in inspection['slide_speech']['slides']
                       for s in page['englishSpeechSegments'] if s['status'] == 'ready')
    assert 'Energy is conserved.' in spoken and 'Label inside the picture' not in spoken
    assert all('Nhãn nằm trong hình' not in u['vi'] for u in inspection['profile']['units'])
