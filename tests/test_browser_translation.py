import io
import json
from pathlib import Path
from threading import Event
from urllib.parse import urlsplit

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import browser_automation as web
from app.browser_accounts import BrowserAccounts
from app.chatgpt_handoff import prepare_request
from app.conversion_progress import ConversionProgress

HTML = r"""<html><head><meta charset="utf-8"></head><body><button data-testid="accounts-profile-button">Fixture</button>
<textarea id="prompt-textarea"></textarea><input type="file" multiple><div id="attached"></div>
<button data-testid="send-button" id="send">Send</button>
<button data-testid="stop-button" id="stop" hidden>Stop</button><div id="messages"></div>
<script>
const initial=INITIAL, chat=CHAT, bad=BAD; let count=0;
function finish(prompt){
 const native=prompt.includes('NỘI DUNG:');
 const answer=document.createElement('div');answer.dataset.messageAuthorRole='assistant';
 answer.dataset.messageId='answer-'+count;
 if(prompt.includes('BLCIMG02')){const nonce=prompt.match(/bytes.fromhex\("([a-f0-9]{32})"\)/)[1];
  const holder=document.createElement('div');holder.style.cssText='overflow:hidden;max-width:300px;transform:translateX(20px)';
  const image=document.createElement('img');image.alt='BiliClass transfer '+nonce;image.src='/transfer/'+nonce+'.png';holder.append(image);answer.append(holder);
 }else if(native){const data=JSON.parse(prompt.split('NỘI DUNG:').at(-1));
  const batch=prompt.match(/"batch": *"([^"]+)"/)[1];
  const result={batch,items:data.map(i=>({id:i.id,en:bad&&count===1?'Missing number':'The object has a velocity of ⟦BLC0⟧ m/s.'}))};
  const pre=document.createElement('pre'),code=document.createElement('code');code.textContent=JSON.stringify(result);pre.append(code);answer.append(pre);
 }else{answer.innerHTML='<a href="/backend-api/files/result.pptx">bai-giang-song-ngu.pptx</a>';}
 document.querySelector('#messages').append(answer);document.querySelector('#stop').hidden=true;
 console.log('DONE:'+JSON.stringify({native,url:location.pathname}));
}
function addUser(prompt){const user=document.createElement('div');user.dataset.messageAuthorRole='user';user.textContent=prompt;document.querySelector('#messages').append(user);}
if(initial){count=initial.count;addUser(initial.prompt);finish(initial.prompt);}
document.querySelector('input').onchange=e=>{document.querySelector('#attached').replaceChildren();for(const f of e.target.files){const n=document.createElement('span');n.textContent=f.name;document.querySelector('#attached').append(n);}};
document.querySelector('#send').onclick=()=>{count++;const prompt=document.querySelector('textarea').value;
 addUser(prompt);history.pushState({},'',chat);document.querySelector('#stop').hidden=false;
 console.log('SENT:'+JSON.stringify({prompt,count,url:location.pathname,files:[...document.querySelector('input').files].map(f=>f.name)}));
 document.querySelector('input').value='';document.querySelector('#attached').replaceChildren();
 setTimeout(()=>finish(prompt),prompt.includes('NỘI DUNG:')?2500:1200);
};</script></body></html>"""


@pytest.fixture
def translation_site(tmp_path, monkeypatch):
    from app import browser_capabilities

    source = tmp_path / "source.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    first = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_textbox(Inches(0.5), Inches(2), Inches(12), Inches(3)).text = "Vật có vận tốc 12 m/s."
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1)).text = "Năng lượng của vật"
    image = io.BytesIO()
    Image.new("RGB", (800, 400), "red").save(image, format="PNG")
    second.shapes.add_picture(
        io.BytesIO(image.getvalue()), Inches(1), Inches(2), width=Inches(10), height=Inches(4)
    )
    deck.save(source)
    returned = tmp_path / "returned.pptx"
    from app.pptx_parts import subset

    subset(source, [2], returned)
    candidate = Presentation(returned)
    candidate.slides[0].shapes[0].text = "Năng lượng của vật\nEnergy of the object"
    candidate.slides[0].notes_slide.notes_text_frame.text = "VI: Năng lượng của vật\nEN: Energy of the object"
    candidate.save(returned)
    store = BrowserAccounts(tmp_path / "library")
    account = store.add("Fixture")
    store.observe(account["id"], "free")
    request = prepare_request(
        tmp_path / "library",
        {
            "title": "Vận tốc",
            "subject": "Vật lý",
            "education_level": "THPT",
            "grade": "10",
            "conversion_format": "english_only",
            "preset": "standard",
            "style": "source",
            "provider": "browser_web",
            "require_highest_model": True,
            "speed_pipeline": 1,
            "prompt_speed_version": 1,
        },
        source,
    )
    capture = {
        "sent": [],
        "done": [],
        "chats": {},
        "pages": 0,
        "bad": False,
        "in_flight": 0,
        "maximum": 0,
        "contexts": 0,
        "errors": [],
    }
    original = web.open_context

    def opening(playwright, account, **kwargs):
        context = original(playwright, account, **kwargs)
        capture["contexts"] += 1

        def watch(page):
            page.on("pageerror", lambda error: capture["errors"].append(str(error)))

            def console(message):
                if message.text.startswith("SENT:"):
                    value = json.loads(message.text[5:])
                    capture["sent"].append(value)
                    capture["chats"][value["url"]] = value
                    capture["in_flight"] += 1
                    capture["maximum"] = max(capture["maximum"], capture["in_flight"])
                elif message.text.startswith("DONE:"):
                    capture["done"].append(json.loads(message.text[5:]))
                    capture["in_flight"] = max(0, capture["in_flight"] - 1)

            page.on("console", console)

        context.on("page", watch)
        for page in context.pages:
            watch(page)

        def route(req):
            path = urlsplit(req.request.url).path
            if path.startswith("/transfer/"):
                import base64
                import hashlib
                import lzma
                import zipfile

                from test_pptx_text_transport import data_image

                original = next(Path(request["folder"]).glob("visual-batches/*/tai-lieu-goc.pptx"))
                with zipfile.ZipFile(original) as before, zipfile.ZipFile(returned) as after:
                    parts = [
                        {
                            "name": name,
                            "sha256": hashlib.sha256(after.read(name)).hexdigest(),
                            "xml": after.read(name).decode("utf-8"),
                        }
                        for name in after.namelist()
                        if name not in before.namelist() or before.read(name) != after.read(name)
                    ]
                    payload = {
                        "format": "biliclass-ooxml-text-v1",
                        "source_sha256": hashlib.sha256(original.read_bytes()).hexdigest(),
                        "target_part_count": len(after.namelist()),
                        "removed": [],
                        "parts": parts,
                    }
                raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
                envelope = json.dumps(
                    {
                        "format": "biliclass-ooxml-xz-v1",
                        "sha256": hashlib.sha256(raw).hexdigest(),
                        "unpacked_bytes": len(raw),
                        "data_b64": base64.b64encode(lzma.compress(raw)).decode(),
                    }
                ).encode()
                image = data_image(envelope, Path(path).stem)
                if capture.get("corrupt_image"):
                    image.putpixel((186, 16), (255, 255, 255))
                output = io.BytesIO()
                image.save(output, format="PNG")
                req.fulfill(body=output.getvalue(), content_type="image/png")
            elif path == "/backend-api/files/result.pptx":
                req.fulfill(
                    body=returned.read_bytes(),
                    headers={
                        "Content-Type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        "Content-Disposition": 'attachment; filename="result.pptx"',
                    },
                )
            else:
                capture["pages"] += 1
                chat = "/c/batch-" + str(capture["pages"])
                body = (
                    HTML.replace("INITIAL", json.dumps(capture["chats"].get(path)))
                    .replace("CHAT", json.dumps(chat))
                    .replace("BAD", json.dumps(capture["bad"]))
                )
                req.fulfill(content_type="text/html; charset=utf-8", body=body)

        context.route("**/*", route)
        return context

    monkeypatch.setattr(web, "open_context", opening)
    monkeypatch.setattr(browser_capabilities, "detect_plan", lambda *_: "free")
    monkeypatch.setattr(browser_capabilities, "select_best_model", lambda *_: "Thinking")
    monkeypatch.setattr(browser_capabilities, "select_reasoning", lambda *_: "Suy luận trên web: High")
    # Geometry is separately checked against real Office by the Qt/native smoke.
    # Keep browser protocol fixtures independent of Office process availability.
    monkeypatch.setattr("app.slide_repair.repair_and_inspect", lambda path, *a, **kw: {"path": str(path)})
    return store, request, capture, source


def test_parallel_text_and_visual_receive_out_of_order_and_merge_correctly(translation_site, monkeypatch):
    import sys

    monkeypatch.setattr(sys.modules[__name__], "HTML", HTML.replace("?2500:1200", "?8000:300"))
    store, request, capture, source = translation_site
    progress = ConversionProgress(request["folder"], lambda _: None)
    try:
        result = web.convert(store.get(), store.root, request["folder"], Event(), progress, timeout=20)
    except Exception:
        assert not capture["errors"], capture["errors"]
        raise
    assert capture["maximum"] == 2
    assert len(capture["sent"]) == 2 and len(capture["sent"][0]["files"]) == 0
    assert capture["done"][0]["native"] is False
    deck = Presentation(result["path"])
    assert len(deck.slides) == 2
    assert "velocity" in deck.slides[0].shapes[0].text
    assert "Energy of the object" in deck.slides[1].shapes[0].text
    assert len(Presentation(source).slides) == 2
    for journal in Path(request["folder"]).glob("*-batches/*/browser-job.json"):
        record = json.loads(journal.read_text(encoding="utf-8"))
        assert record["highest_model_verified"] and record["reasoning_verified"]
    again = web.convert(store.get(), store.root, request["folder"], Event(), progress, timeout=20)
    assert again["cached"] and len(capture["sent"]) == 2


def test_lost_receipt_resumes_old_chat_without_duplicate_send(translation_site, monkeypatch):
    store, request, capture, _ = translation_site
    original = web.submit
    failed = []

    def lose_receipt(*args, **kwargs):
        original(*args, **kwargs)
        if not failed:
            failed.append(True)
            raise web.BrowserProblem("receipt", "Fixture lost acknowledgement")

    monkeypatch.setattr(web, "submit", lose_receipt)
    with pytest.raises(web.BrowserProblem):
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert len(capture["sent"]) == 1
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert Path(result["path"]).is_file() and len(capture["sent"]) == 2


def test_invalid_numeric_translation_gets_only_one_focused_correction(translation_site):
    store, request, capture, _ = translation_site
    capture["bad"] = True
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    native = [row for row in capture["sent"] if not row["files"]]
    assert len(native) == 2 and "Sửa JSON" in native[1]["prompt"]
    assert "12 m/s." in Presentation(result["path"]).slides[0].shapes[0].text


def test_received_batches_resume_local_merge_without_new_browser_or_send(translation_site, monkeypatch):
    store, request, capture, _ = translation_site
    from app import pptx_parts

    original = pptx_parts.merge

    def fail(*args, **kwargs):
        raise OSError("Fixture local merge failure after all batches received")

    monkeypatch.setattr(pptx_parts, "merge", fail)
    with pytest.raises(OSError):
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    contexts = capture["contexts"]
    # The per-job completed result also restores a missing language cache;
    # resuming must not require opening the web or translating again.
    for cached in (Path(request["folder"]).parent.parent / "translation-cache").glob("*.json"):
        cached.unlink()
    monkeypatch.setattr(pptx_parts, "merge", original)
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert Path(result["path"]).is_file()
    assert len([row for row in capture["sent"] if not row["files"]]) == 1
    assert len([row for row in capture["sent"] if row["files"]]) == 1
    assert capture["contexts"] == contexts


def test_corrupt_completed_language_result_is_not_reused_or_resent(translation_site, monkeypatch):
    store, request, capture, _ = translation_site
    from app import pptx_parts

    original = pptx_parts.merge
    monkeypatch.setattr(
        pptx_parts, "merge", lambda *a, **kw: (_ for _ in ()).throw(OSError("Local fixture failure"))
    )
    with pytest.raises(OSError):
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    contexts, sends = capture["contexts"], len(capture["sent"])
    for cached in (Path(request["folder"]).parent.parent / "translation-cache").glob("*.json"):
        cached.unlink()
    result = next(Path(request["folder"]).glob("translation-batches/*/result.json"))
    result.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(pptx_parts, "merge", original)
    with pytest.raises(web.BrowserProblem) as error:
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert error.value.code == "record" and capture["contexts"] == contexts and len(capture["sent"]) == sends


@pytest.mark.parametrize("structured_body", [False, True])
def test_json_rendered_as_prose_with_role_heading_needs_no_correction(
    translation_site, monkeypatch, structured_body
):
    import sys

    store, request, capture, _ = translation_site
    old = "const pre=document.createElement('pre'),code=document.createElement('code');code.textContent=JSON.stringify(result);pre.append(code);answer.append(pre);"
    replacement = """const heading=document.createElement('h4');heading.textContent='ChatGPT said:';answer.append(heading);
const body=document.createElement('div');body.textContent=JSON.stringify(result);BODY_ATTRIBUTE answer.append(body);"""
    replacement = replacement.replace(
        "BODY_ATTRIBUTE", "body.dataset.markdownTextStyle='assistant-message';" if structured_body else ""
    )
    monkeypatch.setattr(sys.modules[__name__], "HTML", HTML.replace(old, replacement))
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert Path(result["path"]).is_file()
    assert len([row for row in capture["sent"] if not row["files"]]) == 1
    assert len(capture["sent"]) == 2 and not capture["errors"]


def test_visual_transfer_failure_keeps_valid_language_reply_without_translating_again(
    translation_site, monkeypatch
):
    import sys

    store, request, capture, _ = translation_site
    monkeypatch.setattr(sys.modules[__name__], "HTML", HTML.replace("?2500:1200", "?100:100"))
    original = web.receive_download
    monkeypatch.setattr(
        "app.browser_text_transport.receive",
        lambda *a, **kw: (_ for _ in ()).throw(
            web.BrowserProblem("file_transfer", "Fixture interrupted result stream")
        ),
    )
    monkeypatch.setattr(
        web,
        "receive_download",
        lambda *a, **kw: (_ for _ in ()).throw(
            web.BrowserProblem("file_transfer", "Fixture interrupted result stream")
        ),
    )
    with pytest.raises(web.BrowserProblem, match="interrupted result stream"):
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    results = list(Path(request["folder"]).glob("translation-batches/*/result.json"))
    assert results and len([row for row in capture["sent"] if not row["files"]]) == 1
    monkeypatch.setattr(web, "receive_download", original)
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert Path(result["path"]).is_file()
    assert len([row for row in capture["sent"] if not row["files"]]) == 1
    assert len(capture["sent"]) == 2


def test_resume_reuses_exact_restored_task_tab_and_leaves_other_chats_alone(monkeypatch, tmp_path):
    from types import SimpleNamespace

    from app import browser_translation

    target = SimpleNamespace(
        url="https://chatgpt.com/c/task", is_closed=lambda: False, set_default_timeout=lambda _: None
    )
    unrelated = SimpleNamespace(url="https://chatgpt.com/c/other", is_closed=lambda: False)
    context = SimpleNamespace(pages=[unrelated, target], new_page=lambda: pytest.fail("Duplicate task tab"))
    navigated = []
    monkeypatch.setattr(web, "navigate", lambda page, url, cancel: navigated.append((page, url)))
    monkeypatch.setattr(web, "signed_in", lambda page: page is target)
    monkeypatch.setattr(web, "require_account", lambda *a: None)
    task = {"folder": tmp_path, "request": {}}
    record = {"state": "waiting", "url": target.url}
    assert browser_translation._prepare_page(context, task, record, Event()) is target
    assert navigated == [] and unrelated.url.endswith("/other")


def test_failed_file_transfer_recovers_xml_image_without_source_reupload(translation_site, monkeypatch):
    from app import browser_text_transport

    store, request, capture, _ = translation_site
    original = browser_text_transport._capture
    restored = []

    def capture_and_check(image, path):
        before = image.evaluate("e=>({style:e.style.cssText,parent:e.parentElement.outerHTML.slice(0,80)})")
        original(image, path)
        assert (
            image.evaluate("e=>({style:e.style.cssText,parent:e.parentElement.outerHTML.slice(0,80)})")
            == before
        )
        restored.append(True)

    monkeypatch.setattr(browser_text_transport, "_capture", capture_and_check)

    def fail(*a, **kw):
        raise web.BrowserProblem("file_transfer", "Fixture broken download")

    monkeypatch.setattr(web, "receive_download", fail)
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert len(capture["sent"]) == 3
    assert sum(bool(r["files"]) for r in capture["sent"]) == 1
    exports = [r for r in capture["sent"] if "BLCIMG02" in r["prompt"]]
    assert len(exports) == 1 and not exports[0]["files"]
    assert restored == [True]
    assert "Energy of the object" in Presentation(result["path"]).slides[1].shapes[0].text
    journal = next(Path(request["folder"]).glob("visual-batches/*/browser-job.json"))
    child = json.loads(journal.read_text(encoding="utf-8"))
    assert child["state"] == "completed" and child["candidate_send_sha256"] == child["send_sha256"]
    receipt = next(journal.parent.glob("receive-data/*/browser-job.json"))
    assert json.loads(receipt.read_text(encoding="utf-8"))["state"] == "completed"


def test_lost_export_receipt_resumes_same_export_without_sending_again(translation_site, monkeypatch):
    store, request, capture, _ = translation_site
    original = web.submit

    def fail(*a, **kw):
        raise web.BrowserProblem("file_transfer", "Fixture broken download")

    def lose(page, folder, record, prompt, *a, **kw):
        original(page, folder, record, prompt, *a, **kw)
        if "BLCIMG02" in prompt and not capture.get("lost"):
            capture["lost"] = True
            raise web.BrowserProblem("receipt", "Fixture lost export receipt")

    monkeypatch.setattr(web, "receive_download", fail)
    monkeypatch.setattr(web, "submit", lose)
    with pytest.raises(web.BrowserProblem):
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert Path(result["path"]).is_file()
    assert len([r for r in capture["sent"] if "BLCIMG02" in r["prompt"]]) == 1
    assert sum(bool(r["files"]) for r in capture["sent"]) == 1


def test_corrupt_recovery_image_never_completes_or_reuploads(translation_site, monkeypatch):
    store, request, capture, _ = translation_site
    capture["corrupt_image"] = True

    def fail(*a, **kw):
        raise web.BrowserProblem("file_transfer", "Fixture broken download")

    monkeypatch.setattr(web, "receive_download", fail)
    with pytest.raises(web.BrowserProblem) as error:
        web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=20)
    assert error.value.code == "file_transfer"
    child = next(Path(request["folder"]).glob("visual-batches/*/browser-job.json"))
    assert json.loads(child.read_text(encoding="utf-8"))["state"] != "completed"
    assert not (child.parent / "result.pptx").exists()
    assert len([r for r in capture["sent"] if "BLCIMG02" in r["prompt"]]) == 1


def test_focused_correction_includes_late_errors_instead_of_first_25_warnings():
    from app.browser_translation import _correction_errors

    quality = {
        "slides": [
            {"slide": n, "issues": [{"code": "image_or_formula_needs_review", "severity": "CHECK"}]}
            for n in range(1, 65)
        ]
    }
    quality["slides"][42]["issues"].append({"code": "english_body_missing", "severity": "FAIL"})
    quality["slides"][63]["issues"].append({"code": "parallel_columns_missing", "severity": "FAIL"})
    message = _correction_errors(quality)
    assert "43" in message and "64" in message
    assert "Cần rà soát" not in message and message.count("Cần sửa") == 2


def test_quiet_active_turn_after_two_refreshes_is_not_false_network_failure(translation_site, monkeypatch):
    from app import browser_translation, browser_wait

    store, request, capture, _ = translation_site
    read_text = browser_translation._read_text
    result_link = web.result_link
    polls = []

    def quiet(page, record):
        polls.append(1)
        return "" if len(polls) < 20 else read_text(page, record)

    monkeypatch.setattr(browser_translation, "_read_text", quiet)
    monkeypatch.setattr(
        web, "result_link", lambda *a, **kw: None if len(polls) < 20 else result_link(*a, **kw)
    )
    monkeypatch.setattr(browser_wait, "stale_refresh", lambda *a, **kw: True)
    result = web.convert(store.get(), store.root, request["folder"], Event(), lambda _: None, timeout=45)
    assert Path(result["path"]).is_file() and len(polls) >= 20
    assert len(capture["sent"]) == 2
