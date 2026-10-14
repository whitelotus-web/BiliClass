import copy

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.chatgpt_handoff import inspect_returned_deck
from app.source_deck import NS


def source_motion(path):
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    shape=slide.shapes.add_textbox(-Inches(3),Inches(2),Inches(2),Inches(1))
    shape.text='Nội dung nguồn'
    timing=etree.SubElement(slide._element,f'{{{NS["p"]}}}timing')
    motion=etree.SubElement(timing,f'{{{NS["p"]}}}animMotion')
    target=etree.SubElement(etree.SubElement(motion,f'{{{NS["p"]}}}cBhvr'),f'{{{NS["p"]}}}tgtEl')
    etree.SubElement(target,f'{{{NS["p"]}}}spTgt',spid=str(shape.shape_id))
    deck.save(path)
    return path


def outside(report):
    return [issue for issue in report['slide_quality']['slides'][0]['issues'] if issue['code']=='outside_slide']


def test_source_motion_position_is_reviewed_without_automatic_relocation(tmp_path):
    source=source_motion(tmp_path/'source.pptx')
    candidate=tmp_path/'candidate.pptx'
    candidate.write_bytes(source.read_bytes())
    report=inspect_returned_deck(candidate,source,'english_only','source')
    assert outside(report)[0]['severity']=='CHECK'
    assert 'source_motion_position' in {i['code'] for i in report['slide_quality']['slides'][0]['issues']}


def test_lost_motion_or_new_off_canvas_geometry_still_fails(tmp_path):
    source=source_motion(tmp_path/'source.pptx')
    for change in ('timing','geometry'):
        deck=Presentation(source)
        if change=='timing':
            node=deck.slides[0]._element.find('p:timing',NS)
            deck.slides[0]._element.remove(node)
        else:
            deck.slides[0].shapes[0].left=-Inches(4)
        candidate=tmp_path/(change+'.pptx')
        deck.save(candidate)
        assert outside(inspect_returned_deck(candidate,source,'english_only','source'))[0]['severity']=='FAIL'


def test_new_translation_overflow_is_not_excused_by_source_motion(tmp_path):
    from app.slide_quality import update_status
    row={'slide':1,'source_motion_shapes':[2],'source_motion_unchanged_text':[],
         'issues':[{'code':'text_overflow','severity':'FAIL','shape_id':2}]}
    report=update_status({'slides':[copy.deepcopy(row)]})
    assert report['counts']['FAIL']==1


def test_unchanged_unused_workspace_is_not_relocated_or_read(tmp_path):
    from app.slide_speech import visible_segments
    source=source_motion(tmp_path/'source.pptx')
    deck=Presentation(source)
    slide=deck.slides[0]
    slide._element.remove(slide._element.find('p:timing',NS))
    slide.shapes[0].text='The object has a velocity of twelve metres per second.'
    deck.save(source)
    report=inspect_returned_deck(source,source,'parallel_columns','source')
    row=report['slide_quality']['slides'][0]
    assert row['source_workspace_shapes']
    assert not any(i['severity']=='FAIL' for i in row['issues'])
    assert not visible_segments(report['slide_speech']['slides'][0])
