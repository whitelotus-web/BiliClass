import copy
import wave
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from PySide6.QtCore import QObject, Signal

from app import speech
from app.lesson_narration import make_plan, variant_id
from app.library import Library
from app.teaching_session import TeachingSession, prepare_teaching
from app.voice_import import import_presentation, save_script


def prepared_lesson(tmp_path, mode='both'):
    from app.chatgpt_handoff import external_preview

    path = tmp_path / 'Independent.pptx'
    deck = Presentation()
    for _ in range(3):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    directory = tmp_path / 'library'
    lesson, preview = import_presentation(directory, path)
    library = Library(directory)
    for index, (vi, en) in enumerate([('Xin chào.', 'Hello.'), ('', ''), ('Tạm biệt.', '')]):
        lesson = save_script(library, lesson['id'], lesson['segments'][index]['id'], lesson['revision'], vi, en)
    plan, _ = make_plan(lesson, mode, {'voice': 'vieneu:Mai Anh', 'rate': 0}, playback='both_controls')
    for keys in plan['keys'].values():
        for key in keys:
            target = directory / 'audio' / (key + '.wav')
            target.parent.mkdir(parents=True, exist_ok=True)
            with wave.open(str(target), 'wb') as audio:
                audio.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
                audio.writeframes(b'\0\0' * 24000)
    plan.update(lesson_id=lesson['id'], revision=lesson['revision'], complete=plan['total'], running=False)
    lesson = library.update_narration(lesson['id'], plan)
    lesson = library.select_narration_variant(lesson['id'], variant_id(plan))
    library.close()
    return lesson, external_preview(lesson, directory), directory, path


def test_original_source_is_managed_without_borrowing_converted_scripts_or_voice(tmp_path):
    from app.teaching_session import prepare_original_teaching

    lesson, _, directory, source = prepared_lesson(tmp_path)
    before = copy.deepcopy(lesson)
    original = source.read_bytes()
    prepared = prepare_original_teaching(lesson, directory)
    assert prepared['preview']['product'] == 'original'
    assert prepared['path'] == prepared['owned_path'] and Path(prepared['path']).is_file()
    assert Path(prepared['path']).read_bytes() == original
    assert not prepared['lesson'].get('narration') and not prepared['lesson'].get('narration_variants')
    assert not prepared['media'][1], 'A generated voice must not be attached to the original'
    assert not prepared['speech']['slides'][0]['englishSpeechSegments'], 'Original is blank; edited scripts belong to the saved product'
    assert lesson == before and source.read_bytes() == original
    with_source = copy.deepcopy(lesson)
    with_source['source']['sha256'] = 'invalid'
    with pytest.raises(ValueError, match='thay đổi'):
        prepare_original_teaching(with_source, directory)


def test_ready_voice_opens_powerpoint_without_approving_or_blocking_on_layout_warning(tmp_path, monkeypatch):
    import time

    from PySide6.QtCore import QCoreApplication

    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    lesson, preview, directory, source = prepared_lesson(tmp_path)
    lesson['voice_only'] = False
    lesson['slide_quality'] = {'counts': {'FAIL': 1, 'CHECK': 2}, 'slides': []}
    library = Library(directory)
    library._write(lesson)
    original, opened = source.read_bytes(), []
    bridge = Bridge(library)
    bridge._lesson, bridge._quick_result = lesson, preview
    completions = []
    monkeypatch.setattr('app.lesson_completion.complete_existing',
                        lambda library, item, cancel=None: completions.append(item['id']) or item)
    monkeypatch.setattr(bridge, '_activate_teaching', opened.append)
    monkeypatch.setattr(speech, 'synthesize', lambda *a, **k: pytest.fail('Opening must reuse existing WAVs'))
    try:
        bridge.openNarratedPowerPoint(variant_id(lesson['narration']))
        deadline = time.monotonic() + 8
        while bridge.busy and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(.01)
        assert not bridge.busy and len(opened) == 1, bridge.message
        assert set(opened[0]['media'][1]) == {'vi', 'en'}
        saved = library.get(lesson['id'])
        assert saved['revision'] == lesson['revision']
        assert not any(row['approved'] for row in saved['segments'])
        assert saved['slide_quality']['counts']['FAIL'] == 1
        assert not completions, 'Opening a saved deck does not request another layout repair'
        assert source.read_bytes() == original
    finally:
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader):
            worker.shutdown()
        library.close()


def test_completed_preview_uses_the_saved_voice_without_synthesis(tmp_path, monkeypatch):
    from PySide6.QtCore import QCoreApplication

    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    library = Library(directory)
    bridge = Bridge(library)
    bridge._lesson, bridge._quick_result = lesson, preview
    calls = []
    monkeypatch.setattr(bridge.english_reader, 'start', lambda *a, **k: calls.append((a, k)))
    try:
        bridge.narrationWorkspace.listenCreated(1, 'en', variant_id(lesson['narration']))
        assert calls == [((['Hello.'], lesson['narration']['voice'], 0), {'cache_only': True})]
        bridge.narrationWorkspace.listenCreated(2, 'en', variant_id(lesson['narration']))
        assert len(calls) == 1 and 'không có lời' in bridge.message
        lesson['segments'][0]['en'] = 'Changed since generation.'
        bridge.narrationWorkspace.listenCreated(1, 'en', variant_id(lesson['narration']))
        assert len(calls) == 1 and bridge.error
        app.processEvents()
    finally:
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader):
            worker.shutdown()
        library.close()


class FakeBridge(QObject):
    changed = Signal()
    speechStateChanged = Signal()

    def __init__(self, directory):
        super().__init__()
        self.library = SimpleNamespace(directory=directory)
        self._powerpoint_state = {'active': True, 'slide': 1, 'slide_id': 256}
        self.calls = []
        self.powerpoint_worker = SimpleNamespace(navigate=lambda *args: self.calls.append(args))
        self.english_reader = SimpleNamespace(start=lambda *args, **kwargs: self.calls.append((args, kwargs)))

    def stopSpeech(self):
        self.calls.append('stop')

    def pauseEnglish(self):
        self.calls.append('pause_cached')

    def inform(self, *args, **kwargs):
        self.calls.append(('message', args))


def test_unbinding_disables_controls_and_drops_old_office_updates(tmp_path):
    from PySide6.QtCore import QCoreApplication

    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    library = Library(tmp_path / 'library')
    bridge = Bridge(library)
    stopped = []
    try:
        bridge.powerpoint_worker = SimpleNamespace(stop=lambda: stopped.append(True))
        previous = {'active': True, 'opened': True, 'slide': 2, 'total': 3}
        bridge._powerpoint_state = previous
        bridge.stopPowerPoint()
        assert stopped == [True] and not bridge.companionCanRead
        assert not bridge.powerpointState['opened'] and bridge.powerpointState['closing']
        bridge._powerpoint_changed(previous)
        assert bridge.powerpointState['closing'] and not bridge.powerpointState['active']
        app.processEvents()
    finally:
        bridge.powerpoint_worker = None
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader,
                       bridge.voice_controls, bridge._companion_preferences):
            worker.shutdown()
        bridge.classroom.shutdown()
        library.close()


def test_no_voice_bindings_use_separate_session_copies_and_cleanup_preserves_source(tmp_path):
    lesson, preview, directory, source = prepared_lesson(tmp_path)
    lesson['narration'] = {'version': 2, 'mode': 'none'}
    original, product = source.read_bytes(), Path(preview['path'])
    product_bytes = product.read_bytes()
    first = prepare_teaching(lesson, preview, directory)
    second = prepare_teaching(lesson, preview, directory)
    first_path, second_path = Path(first['path']), Path(second['path'])
    assert first_path != second_path and first_path not in (source, product)
    assert first_path.read_bytes() == second_path.read_bytes() == product_bytes
    assert first['owned_path'] == str(first_path) and not any(first['media'].values())
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(first)
    session.clear()
    assert not first_path.exists() and second_path.is_file()
    assert source.read_bytes() == original and product.read_bytes() == product_bytes
    second_path.unlink()


def test_bound_snapshot_follows_slide_not_browsed_lesson(tmp_path, monkeypatch):
    lesson, preview, directory, source = prepared_lesson(tmp_path)
    original = source.read_bytes()
    monkeypatch.setattr(speech, 'synthesize', lambda *_: (_ for _ in ()).throw(AssertionError('No synthesis')))
    prepared = prepare_teaching(lesson, preview, directory)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepared)
    assert session.audio == {'vi': True, 'en': True, 'both': False}
    bridge._lesson = {'id': 'another-lesson', 'segments': []}
    lesson['segments'][0]['vi'] = 'Changed after binding'
    assert session.texts('vi') == ['Xin chào.']
    assert session.state['lesson_id'] == prepared['lesson']['id']
    bridge._powerpoint_state['slide'] = 2
    assert session.audio == {'vi': False, 'en': False, 'both': False}
    bridge._powerpoint_state['slide'] = 3
    assert session.audio == {'vi': True, 'en': False, 'both': False}
    assert source.read_bytes() == original
    import zipfile

    with zipfile.ZipFile(prepared['path']) as archive:
        assert any(name.startswith('ppt/media/') and name.endswith('.wav') for name in archive.namelist())


def test_cached_fallback_waits_for_powerpoint_stop_and_rejects_late_slide(tmp_path, monkeypatch):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    from app import presentation_audio

    monkeypatch.setattr(presentation_audio, 'embed_audio', lambda *_, **kwargs: (_ for _ in ()).throw(ValueError('No room')))
    prepared = prepare_teaching(lesson, preview, directory)
    assert prepared['warning'] and not any(prepared['media'].values())
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepared)
    session.read('vi')
    token = session.token
    assert bridge.calls[-1] == ('stop_media', {'token': token, 'slide_id': 256, 'language': 'vi'})
    bridge._powerpoint_state['slide'] = 2
    session.office_event({'event': 'media_stopped', 'token': token})
    assert not any(isinstance(call, tuple) and call[-1] == {'cache_only': True} for call in bridge.calls)
    bridge._powerpoint_state['slide'] = 1
    session.read('vi')
    session.office_event({'event': 'media_stopped', 'token': session.token})
    assert bridge.calls[-1][1] == {'cache_only': True}
    session.read('vi')
    token = session.token
    session.stop()
    calls = len(bridge.calls)
    session.office_event({'event': 'media_stopped', 'token': token})
    assert len(bridge.calls) == calls, 'A stop must invalidate a queued read acknowledgement'
    changed = copy.deepcopy(prepared)
    changed['sha256'] = 'changed'
    session.bind(changed)
    session.read('vi')
    assert bridge.calls[-1][0] == 'message'


def test_explicit_no_voice_and_independent_embedded_powerpoint_are_teachable(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    with_voice = prepare_teaching(lesson, preview, directory)
    lesson['narration'] = {'version': 2, 'mode': 'none'}
    without = prepare_teaching(lesson, preview, directory)
    assert without['path'] != preview['path'] and not any(without['media'].values())
    assert Path(without['path']).read_bytes() == Path(preview['path']).read_bytes()
    imported, preview = import_presentation(directory, with_voice['path'])
    assert imported['narration']['mode'] == 'none'
    independent = prepare_teaching(imported, preview, directory)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(independent)
    assert session.audio == {'vi': True, 'en': True, 'both': False}


def test_key_and_mouse_toggle_pause_resume_switch_and_missing_language(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    bridge._powerpoint_state['slide_id'] = 256
    session.office_event({'event': 'read_requested', 'slide_id': 256, 'language': 'en'})
    assert bridge.calls[-1] == ('play_media', session._media_request())
    session.office_event({'event': 'media_started', 'language': 'en'})
    session.toggleRead('en')
    assert bridge.calls[-1] == ('pause_media', session._media_request())
    session.office_event({'event': 'media_paused'})
    session.toggleRead('en')
    assert bridge.calls[-1] == ('resume_media', session._media_request())
    session.office_event({'event': 'media_resumed'})
    session.toggleRead('vi')
    assert bridge.calls[-1] == ('play_media', session._media_request())
    assert session.language == 'vi'
    before = list(bridge.calls)
    session.office_event({'event': 'read_requested', 'slide_id': 999, 'language': 'en'})
    assert bridge.calls == before, 'Drop key-downs belonging to a previous slide'
    bridge._powerpoint_state['slide'] = 3
    session.toggleRead('en')
    assert bridge.calls[-1][0] == 'message'
    assert not any(call[0] == 'play_media' for call in bridge.calls[len(before):] if isinstance(call, tuple))


def test_auto_reveal_only_new_english_and_matched_vietnamese_paragraphs(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    prepared = prepare_teaching({**lesson, 'narration': {'mode':'none'}}, preview, directory)
    lesson['segments'][0].update(vi='Năng lượng\nĐoạn thứ nhất.\nĐoạn thứ hai.',
                                  en='Energy\nFirst paragraph.\nSecond paragraph.')
    lesson['narration_script'] = {}
    prepared['lesson'] = lesson
    page = prepared['speech']['slides'][0]
    page['englishSpeechSegments'] = [
        {'id': 'title', 'kind':'title', 'status':'ready', 'english_text':'Energy', 'visibility':{'kind':'always'}},
        *[{'id': f'body-{n}', 'kind':'body', 'status':'ready', 'english_text':text,
           'visibility':{'kind':'click', 'show_after':n, 'hide_after':None, 'settle_seconds':0}}
          for n,text in ((1,'First paragraph.'),(2,'Second paragraph.'))]]
    bridge = FakeBridge(directory)
    bridge._powerpoint_state.update(click_index=1)
    session = TeachingSession(bridge)
    session.bind(prepared)
    assert session.texts('en') == ['Energy', 'First paragraph.']
    assert session.texts('vi') == ['Năng lượng', 'Đoạn thứ nhất.']
    for lang in ('vi','en'):
        visible = session._selected_items(lang)
        session._progress.finish(256, lang, session._selected_items(lang))
        assert session.texts(lang) == []
        assert session._selected_items(lang, replay=True) == visible
    bridge._powerpoint_state['click_index'] = 2
    session._watch_build()
    assert session.texts('en') == ['Second paragraph.']
    assert session.texts('vi') == ['Đoạn thứ hai.']
    assert [row[1] for row in session._selected_items('en', replay=True)] == ['Second paragraph.'], 'Auto still prefers the newly revealed paragraph'


def test_repeated_press_while_native_player_prepares_queues_pause_without_restarting(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.toggleRead('en')
    assert session.native_state == 'preparing'
    before = list(bridge.calls)
    session.toggleRead('en')
    assert bridge.calls == before
    session.office_event({'event': 'media_started', 'language': 'en'})
    assert bridge.calls[-1] == ('pause_media', session._media_request())
    assert sum(call[0] == 'play_media' for call in bridge.calls if isinstance(call, tuple)) == 1


def test_reading_can_pause_after_auto_marks_all_current_words_complete(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.toggleRead('en')
    session.office_event({'event': 'media_started', 'language': 'en'})
    session._progress.finish(256, 'en', session._read_batch)
    assert session.audio['en'], 'Completed audio remains available for deliberate replay'
    session.toggleRead('en')
    assert bridge.calls[-1] == ('pause_media', session._media_request())
    session.office_event({'event': 'media_finished', 'slide_id': 256})
    session.toggleRead('en')
    assert bridge.calls[-1] == ('play_media', session._media_request())


def test_build_aware_reading_bypasses_whole_slide_media_and_only_uses_visible_cached_words(tmp_path):
    import time

    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    prepared = prepare_teaching(lesson, preview, directory)
    page = prepared['speech']['slides'][0]
    for row in page['englishSpeechSegments']:
        row['visibility'] = {'kind': 'click', 'show_after': 1, 'hide_after': None, 'settle_seconds': 0}
    bridge = FakeBridge(directory)
    bridge._powerpoint_state.update(click_index=0, build_started=time.monotonic() - 2)
    session = TeachingSession(bridge)
    session.bind(prepared)
    assert session.state['reading_mode'] == 'auto' and session.reading_scope == 'visible'
    assert not session.audio['en'], 'Auto never reads unrevealed animated content'
    assert session.metaObject().indexOfMethod('setReadingScope(QString)') == -1, 'No manual override is exposed to QML'
    assert not session.audio['en']
    bridge._powerpoint_state['click_index'] = 1
    assert session.audio['en']
    session.read('en')
    assert bridge.calls[-1] == ('stop_media', session._media_request()), 'Wait for native stop before cached playback'
    pending = session.pending_read
    session.office_event({'event': 'media_stopped', 'token': None})
    assert session.pending_read == pending and session.native_state == 'preparing'
    session.office_event({'event': 'media_stopped', 'token': session.token})
    assert bridge.calls[-1][1] == {'cache_only': True}


@pytest.mark.parametrize('action', ['read', 'cancel', 'next_build', 'new_slide', 'stop', 'language'])
def test_speakers_stay_available_during_entrance_and_queue_only_the_current_build(tmp_path, action):
    import time

    from PySide6.QtCore import QCoreApplication

    app = QCoreApplication.instance() or QCoreApplication([])
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    prepared = prepare_teaching(lesson, preview, directory)
    for row in prepared['speech']['slides'][0]['englishSpeechSegments']:
        row['visibility'] = {'kind': 'click', 'show_after': 1, 'hide_after': None, 'settle_seconds': .12}
    bridge = FakeBridge(directory)
    bridge._powerpoint_state.update(click_index=1, build_started=time.monotonic())
    session = TeachingSession(bridge)
    session.bind(prepared)
    try:
        assert session.state['settling'] and session.audio['en'] and session.audio['vi']
        assert session.texts('en') == [], 'Unfinished entrances cannot be read early'
        session.toggleRead('en')
        assert session._deferred_read and not session.pending_read
        if action == 'cancel':
            session.toggleRead('en')
        elif action == 'language':
            session.toggleRead('vi')
        elif action in {'next_build', 'new_slide'}:
            bridge._powerpoint_state.update(click_index=2, build_started=time.monotonic())
            if action == 'new_slide':
                bridge._powerpoint_state.update(slide=3, slide_id=258)
            session.reset_reading()
            session._watch_build()
        elif action == 'stop':
            session.stop()
        deadline = time.monotonic() + .6
        while time.monotonic() < deadline:
            app.processEvents()
            time.sleep(.005)
        if action in {'read', 'language'}:
            expected = 'vi' if action == 'language' else 'en'
            assert session.language == expected and session.pending_read
            assert session.pending_read[3] == (['Xin chào.'] if expected == 'vi' else ['Hello.'])
            assert bridge.calls[-1][0] == 'stop_media'
        else:
            assert not session.pending_read and not session._deferred_read
            assert not any(isinstance(call, tuple) and call[0] == 'play_media' for call in bridge.calls)
    finally:
        session.clear()


def test_player_failure_reuses_verified_cache_and_preserves_pause_intent(tmp_path, monkeypatch):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    monkeypatch.setattr(speech, 'synthesize', lambda *a: pytest.fail('Never synthesize during teaching'))
    session.toggleRead('en')
    request = session._media_request()
    session.toggleRead('en')  # Pause while native Office prepares.
    session.office_event(dict(request, event='media_error', operation='play_media', message='Native player unavailable'))
    assert session.pending_read and bridge.calls[-1][0] == 'stop_media'
    session.office_event(dict(request, event='media_stopped'))
    assert bridge.calls[-2][1] == {'cache_only': True} and bridge.calls[-1] == 'pause_cached'
    assert not session.pending_read and session.native_state == 'stopped'


def test_only_separate_languages_can_play_even_with_a_legacy_combined_recording(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    prepared = prepare_teaching(lesson, preview, directory)
    session.bind(prepared)
    assert not session.audio['both']
    before = list(bridge.calls)
    session.read('both')
    session.toggleRead('both')
    assert bridge.calls == before and not session.pending_read
    for language in ('en', 'vi'):
        session.read(language)
        assert session.language == language and bridge.calls[-1][0] == 'play_media'


def test_switching_language_while_cache_fallback_waits_keeps_only_latest_intent(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.read('en')
    session.office_event(dict(session._media_request(), event='media_error', operation='play_media', message='Unavailable'))
    assert session.pending_read
    old_request = session._media_request()
    session.toggleRead('vi')
    assert session.language == 'vi' and session.native_state == 'preparing' and not session.pending_read
    assert session.token != old_request['token'] and bridge.calls[-1][0] == 'play_media'
    calls = list(bridge.calls)
    session.office_event(dict(old_request, event='media_stopped'))
    assert bridge.calls == calls and session.language == 'vi'


def test_late_playback_events_cannot_revive_stopped_or_new_slide_audio(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.read('en')
    request = session._media_request()
    session.stop()
    calls = list(bridge.calls)
    session.office_event(dict(request, event='media_started'))
    assert session.native_state == 'stopped' and bridge.calls == calls
    session.read('en')
    request = session._media_request()
    bridge._powerpoint_state.update(slide=3, slide_id=258)
    session.reset_reading()
    calls = list(bridge.calls)
    session.office_event(dict(request, event='media_error', operation='play_media', message='Old slide'))
    assert not session.pending_read and bridge.calls == calls


def test_missing_native_player_and_missing_cache_finishes_with_clear_error(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.read('en')
    for key in lesson['narration']['keys']['en']:
        (directory / 'audio' / (key + '.wav')).unlink(missing_ok=True)
    session.office_event(dict(session._media_request(), event='media_error', operation='play_media', message='Native player unavailable'))
    assert session.native_state == 'stopped' and not session.pending_read
    assert bridge.calls[-1] == ('message', ('Native player unavailable', True))


def test_office_rejecting_queued_request_releases_preparing_without_playing_cache(tmp_path):
    lesson, preview, directory, _ = prepared_lesson(tmp_path)
    bridge = FakeBridge(directory)
    session = TeachingSession(bridge)
    session.bind(prepare_teaching(lesson, preview, directory))
    session.read('en')
    request = session._media_request()
    session.office_event(dict(request, slide_id=258, event='media_rejected'))
    assert session.native_state == 'stopped' and not session.pending_read
    assert bridge.calls[-1][0] == 'play_media'


@pytest.mark.parametrize('scope,expected', [('full_slide', 'reading'), ('visible', 'stopped')])
def test_animation_build_keeps_full_slide_audio_but_invalidates_visible_scope(tmp_path, scope, expected):
    from PySide6.QtCore import QCoreApplication

    from app.ui import Bridge

    app = QCoreApplication.instance() or QCoreApplication([])
    library = Library(tmp_path / 'library')
    bridge = Bridge(library)
    try:
        bridge.teaching_session.bound = {'lesson': {'id': 'bound'}}
        bridge.teaching_session.reading_scope = scope
        bridge.teaching_session.native_state = 'reading'
        bridge._powerpoint_state = {'active': True, 'slide': 1, 'slide_id': 256, 'click_index': 0}
        bridge._powerpoint_changed({'active': True, 'slide': 1, 'slide_id': 256, 'click_index': 1})
        assert bridge.teaching_session.native_state == expected
        app.processEvents()
    finally:
        for worker in (bridge._catalog, bridge.browserAI, bridge.slidePreview, bridge.libraryThumbnails,
                       bridge.narrationPreferences, bridge.narration_preparation, bridge.english_reader):
            worker.shutdown()
        library.close()
