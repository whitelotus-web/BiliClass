import hashlib
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from app.chatgpt_handoff import load_request, make_prompt, prepare_request
from app.conversion_formats import FORMATS, METHODS
from app.template_rules import (
    STYLE_RULES,
    audience_guidance,
    reference_asset,
    reference_catalog,
    subject_guidance,
)


def config(**updates):
    return {"title": "Bài của giáo viên", "subject": "Ngữ văn", "education_level": "THCS", "grade": "7",
            "conversion_format": "sentence_pairs", "style": "template", "preset": "standard", **updates}


@pytest.mark.parametrize("preset", ["standard", "visual", "practice"])
@pytest.mark.parametrize("spec", FORMATS, ids=lambda item: item["id"])
def test_sent_template_never_contains_examples_or_previews(tmp_path, preset, spec):
    request = prepare_request(tmp_path, config(preset=preset, conversion_format=spec["id"]), text="NGỮ LIỆU CỦA GIÁO VIÊN")
    assert load_request(request["folder"]) == request
    with zipfile.ZipFile(request["bundle"]) as archive:
        assert request["attachments"] == ["tai-lieu-goc.txt", "mau-biliclass.pptx"]
        assert not any(name.endswith(".png") for name in archive.namelist())
        assert archive.read("mau-biliclass.pptx") == reference_asset(preset).read_bytes()
        assert archive.read("tai-lieu-goc.txt").decode() == "NGỮ LIỆU CỦA GIÁO VIÊN"
    deck = Presentation(Path(request["folder"]) / "mau-biliclass.pptx")
    assert len(deck.slides) == 4
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "[SUBJECT]" in text and "[SOURCE_IMAGE]" in text
    for demo in ("hình chữ nhật", "cấp số cộng", "vòng tuần hoàn", "6 cm", "u₁", "sin(", "Vận tốc"):
        assert demo.casefold() not in text.casefold()
    assert all(not slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)
    with zipfile.ZipFile(reference_asset(preset)) as archive:
        assert not any(name.startswith("ppt/media/") for name in archive.namelist())
    prompt = request["prompt"]
    assert METHODS[spec["id"]] in prompt
    assert STYLE_RULES[preset] in prompt
    assert all(rule not in prompt for key, rule in STYLE_RULES.items() if key != preset)
    assert "THCS, lớp 6–9" in prompt and "ngữ liệu" in prompt
    assert "không áp đặt hai cột" in prompt
    assert "ba lựa chọn A/B/C" in prompt


@pytest.mark.parametrize("preset", ["standard", "visual", "practice"])
def test_source_design_has_no_template_attachment_or_style_module(tmp_path, preset):
    source = tmp_path / "Nguồn.pptx"
    # Preparation preserves bytes and does not rebuild the PPTX source.
    source.write_bytes(b"native source fixture")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    request = prepare_request(tmp_path / "library", config(style="source", preset=preset), source_path=source)
    assert request["attachments"] == ["tai-lieu-goc.pptx"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert not any(rule in request["prompt"] for rule in STYLE_RULES.values())
    assert "mau-biliclass.pptx" not in request["prompt"]
    assert "CHỈNH BẢN SAO POWERPOINT GỐC" in request["prompt"]


@pytest.mark.parametrize("grade, expected", [("1", "28–36"), ("2", "28–36"), ("5", "24–32"), ("6", "22–30"), ("12", "20–28")])
def test_age_guidance_changes_without_changing_language_or_source(grade, expected):
    prompt = make_prompt(config(grade=grade, conversion_format="parallel_columns"), "tai-lieu-goc.pptx")
    assert expected in prompt and METHODS["parallel_columns"] in prompt
    assert "giữ đủ kiến thức/câu Việt" in prompt and "không phải giới hạn cứng" in prompt


def test_unknown_grade_does_not_assume_high_school():
    assert "không mặc định THPT" in audience_guidance({"grade": "unknown"})


@pytest.mark.parametrize("subject, expected", [("Toán", "chứng minh"), ("Hóa học", "hóa trị"), ("Vật lý", "vectơ"),
    ("Lịch sử và Địa lí", "mốc thời gian"), ("Ngữ văn", "ngữ liệu"), ("Tin học và Công nghệ", "cú pháp"),
    ("Âm nhạc", "nhịp/phách"), ("Tiếng Anh", "phiên âm"), ("Giáo dục thể chất", "an toàn"),
    ("Giáo dục kinh tế và pháp luật", "quy định"), ("Mĩ thuật", "tác phẩm"), ("Khác: Thực hành", "ngữ cảnh")])
def test_subject_guidance_preserves_specialised_material(subject, expected):
    assert expected in subject_guidance(subject)


def test_all_previews_exist_and_indices_are_local_to_the_four_slide_deck():
    for preset, item in reference_catalog().items():
        assert [layout["slide"] for layout in item["layouts"]] == [1, 2, 3, 4]
        for role in ("title", "concept", "visual", "practice"):
            assert reference_asset(preset, role=role).read_bytes().startswith(b"\x89PNG")
        for spec in FORMATS:
            assert reference_asset(preset, conversion_format=spec["id"]).read_bytes().startswith(b"\x89PNG")
    for invalid in ("../", "unknown"):
        with pytest.raises(ValueError):
            reference_asset(invalid)
    with pytest.raises(ValueError):
        reference_asset("visual", conversion_format="bad")


def test_resuming_keeps_the_sent_prompt_after_application_upgrade(tmp_path, monkeypatch):
    request = prepare_request(tmp_path, config(), text="Nội dung riêng.")
    monkeypatch.setattr("app.chatgpt_handoff.make_prompt", lambda *_: "New prompt after upgrade")
    assert load_request(request["folder"])["prompt"] == request["prompt"]
    file = Path(request["folder"]) / "mau-biliclass.pptx"
    file.write_bytes(b"changed")
    with pytest.raises(ValueError, match="Mẫu.*đã thay đổi"):
        load_request(request["folder"])


def test_missing_template_fails_without_sending_an_incomplete_package(tmp_path, monkeypatch):
    monkeypatch.setattr("app.template_rules.reference_asset", lambda *_: tmp_path / "missing.pptx")
    with pytest.raises(FileNotFoundError):
        prepare_request(tmp_path, config(), text="Nội dung riêng.")
    assert not list((tmp_path / "chatgpt").glob("*/config.json"))
