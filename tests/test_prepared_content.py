import wave

import pytest

from app.content import assistant_response
from app.deck_export import export_deck
from app.library import Library
from app.pack import export_pack, import_pack
from app.portable_audio import available_audio
from app.speech import audio_key
from app.storage import backup_library, restore_backup
from app.text_quality import protected_parts, review_warnings


@pytest.fixture
def prepared(tmp_path):
    library = Library(tmp_path / "library")
    lesson = library.create("Môn tự tạo", "Liên môn", "THPT", "12", [("Đoạn 1", "Giải thích. Ví dụ.")])
    identity, segment = lesson["id"], lesson["segments"][0]["id"]
    library.edit_segment(identity, segment, "Giải thích. Ví dụ.", "Explain. Example.", True)
    library.save_support(identity, segment, {"kind": "example", "vi": "Ví dụ", "en": "Example", "approved": True})
    library.save_question(identity, {"kind": "single", "vi": "Chọn ví dụ", "en": "Choose an example", "concept_id": segment,
                                   "correct": "A", "approved": True, "options": [{"vi": "Một", "en": "One"}, {"vi": "Hai", "en": "Two"}]})
    yield library, library.get(identity)
    library.close()


def test_pack_resets_all_reviews_and_remaps_concepts(prepared, tmp_path):
    library, lesson = prepared
    imported = import_pack(library, export_pack(library, lesson["id"], tmp_path / "pack.biliclass"))
    assert not imported["questions"][0]["approved"]
    assert not imported["segments"][0]["support"][0]["approved"]
    assert imported["questions"][0]["concept_id"] == imported["segments"][0]["id"]
    assert imported["questions"][0]["id"] != lesson["questions"][0]["id"]


@pytest.mark.parametrize("change", ["edit", "split", "subject", "restore"])
def test_source_changes_invalidate_dependent_reviews(prepared, change):
    library, lesson = prepared
    identity, segment = lesson["id"], lesson["segments"][0]["id"]
    if change == "edit":
        result = library.edit_segment(identity, segment, "Mới", "New", True)
    elif change == "split":
        result = library.split_segment(identity, segment, 10, 8)
    elif change == "subject":
        result = library.update_metadata(identity, "Bài", "Môn khác", "THPT", "11")
    else:
        library.set_presentation(identity, 3, "split_view")
        result = library.restore_revision(identity, lesson["revision"])
    assert not result["questions"][0]["approved"]
    assert all(not item["approved"] for seg in result["segments"] for item in seg.get("support", []))
    assert result["questions"][0]["concept_id"] in {s["id"] for s in result["segments"]}


def test_assistant_only_serves_approved_prepared_items(prepared):
    _, lesson = prepared
    segment = lesson["segments"][0]
    assert assistant_response(segment, "example", "en")["text"] == "Example"
    assert not assistant_response(segment, "explanation", "en")["available"]
    assert assistant_response(segment, "rescue", "vi")["text"] == segment["vi"]
    segment["approved"] = False
    assert not assistant_response(segment, "example", "en")["available"]


def test_portable_audio_without_installed_voice_and_exact_content(prepared, tmp_path):
    library, lesson = prepared
    library.set_setting("voice_en", "test-voice")
    text = lesson["segments"][0]["en"]
    folder = library.directory / "audio"
    folder.mkdir()
    wav = folder / (audio_key(text, "test-voice", 0) + ".wav")
    with wave.open(str(wav), "wb") as stream:
        stream.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        stream.writeframes(b"\0\0" * 1600)
    pack = export_pack(library, lesson["id"], tmp_path / "audio.biliclass")
    with_other = Library(tmp_path / "other")
    try:
        imported = import_pack(with_other, pack)
        assert not imported["segments"][0]["approved"]
        assert available_audio(with_other.directory, text, "en")
        assert not available_audio(with_other.directory, text + " Changed", "en")
        assert not available_audio(with_other.directory, text, "vi")
        assert not available_audio(with_other.directory, text, "en", "another-voice", 0)
    finally:
        with_other.close()


def test_consistent_backup_restore_into_new_library(prepared, tmp_path):
    library, lesson = prepared
    backup = backup_library(library.directory, daily=True)
    assert backup_library(library.directory, daily=True) == backup
    restored_path = restore_backup(backup, tmp_path / "restored")
    restored = Library(restored_path)
    try:
        assert restored.get(lesson["id"]) == lesson
    finally:
        restored.close()
    with pytest.raises(ValueError, match="trống"):
        restore_backup(backup, library.directory)
    assert library.get(lesson["id"]) == lesson


def test_editable_pptx_preserves_text_and_requires_review(prepared, tmp_path):
    from pptx import Presentation
    library, lesson = prepared
    target = export_deck(lesson, library.directory, tmp_path / "Bài mới.pptx")
    deck = Presentation(target)
    text = "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert lesson["segments"][0]["vi"] in text and lesson["segments"][0]["en"] in text
    with_profile = export_deck(lesson, library.directory, tmp_path / "Có hồ sơ.pptx", {
        "teacher": "Cô Minh", "school": "THPT Bình Minh", "show_profile": True,
    })
    profile_text = "\n".join(shape.text for shape in Presentation(with_profile).slides[0].shapes if shape.has_text_frame)
    assert "Cô Minh · THPT Bình Minh" in profile_text
    hidden_profile = export_deck(lesson, library.directory, tmp_path / "Ẩn hồ sơ.pptx", {
        "teacher": "Cô Minh", "show_profile": False,
    })
    hidden_text = "\n".join(shape.text for shape in Presentation(hidden_profile).slides[0].shapes if shape.has_text_frame)
    assert "Cô Minh" not in hidden_text
    lesson["segments"][0]["approved"] = False
    with pytest.raises(ValueError, match="Duyệt"):
        export_deck(lesson, library.directory, target)


@pytest.mark.parametrize("layout", ["keyword_overlay", "line_pair", "split_view", "english_rescue"])
def test_pptx_follows_per_lesson_display_mode(prepared, tmp_path, layout):
    from pptx import Presentation

    library, lesson = prepared
    lesson["layout"] = layout
    terms = [{"vi": "Giải thích", "en": "Explain"}]
    path = export_deck(lesson, library.directory, tmp_path / f"{layout}.pptx", terms=terms)
    slide = Presentation(path).slides[0]
    text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert ("Giải thích" in text) == (layout != "english_rescue")
    assert ("Explain." in text) == (layout in {"line_pair", "split_view", "english_rescue"})
    if layout == "keyword_overlay":
        assert "Giải thích (Explain)" in text
    if layout == "line_pair":
        assert any(paragraph.font.italic for shape in slide.shapes if shape.has_text_frame
                   for paragraph in shape.text_frame.paragraphs if "Explain." in paragraph.text)


def test_school_logo_is_in_deck_and_library_backup(prepared, tmp_path):
    from PIL import Image
    from pptx import Presentation

    library, lesson = prepared
    assets = library.directory / "assets"
    assets.mkdir()
    logo = assets / "school-logo.png"
    Image.new("RGBA", (120, 80), "#0869f9").save(logo)
    deck = export_deck(lesson, library.directory, tmp_path / "branded.pptx", {"show_profile": True})
    assert any(shape.shape_type == 13 for shape in Presentation(deck).slides[0].shapes)
    backup = backup_library(library.directory)
    restored_path = restore_backup(backup, tmp_path / "logo-restored")
    assert (restored_path / "assets" / "school-logo.png").read_bytes() == logo.read_bytes()


def test_protected_scientific_and_curated_spans():
    parts = protected_parts("Tính $x^2 + 2x = 3$ với 20 kg DNA và tế bào.", [{"vi": "tế bào", "en": "cell", "locked": True}])
    protected = [p["text"] for p in parts if p["protected"]]
    assert "$x^2 + 2x = 3$" in protected and "20 kg" in protected and "DNA" in protected and "cell" in protected
    assert review_warnings("Có 20 học sinh", "There are 30 students")
    assert not review_warnings("Có 20 học sinh", "There are 20 students")
