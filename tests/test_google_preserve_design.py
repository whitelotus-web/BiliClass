import hashlib
import copy
import zipfile
from threading import Event

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.chatgpt_handoff import inspect_returned_deck
from app.native_bilingual import assemble, extract
from app.source_deck import NS, source_slide_parts


def source(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    simple = deck.slides.add_slide(deck.slide_layouts[6])
    title = simple.shapes.add_textbox(Inches(.5), Inches(.4), Inches(12), Inches(1.2))
    title.text = 'Cơ năng'
    body = simple.shapes.add_textbox(Inches(.5), Inches(2), Inches(12), Inches(4))
    body.text = 'Năng lượng của vật được bảo toàn khi không có lực ma sát.'
    busy = deck.slides.add_slide(deck.slide_layouts[6])
    box = busy.shapes.add_textbox(Inches(.5), Inches(2), Inches(6), Inches(2.5))
    box.text = body.text
    picture = path.with_suffix('.png')
    Image.new('RGB', (200, 200), '#46A3CF').save(picture)
    busy.shapes.add_picture(str(picture), Inches(7), Inches(1), width=Inches(5))
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = 'Arial'
                    paragraph.font.size = Pt(26)
    timing = etree.SubElement(busy._element, f'{{{NS["p"]}}}timing')
    target = etree.SubElement(timing, f'{{{NS["p"]}}}spTgt', spid=str(box.shape_id))
    etree.SubElement(target, f'{{{NS["p"]}}}txEl')
    deck.save(path)
    return path


def translated(manifest):
    return {unit['id']: {'en': 'Mechanical energy' if unit['title'] else
                        'The energy of the object is conserved in the absence of friction.'}
            for unit in manifest['units']}


def test_mixed_deck_preserves_complex_source_and_uses_adjacent_language_pages(tmp_path):
    original = source(tmp_path/'source.pptx')
    digest = hashlib.sha256(original.read_bytes()).hexdigest()
    manifest = extract(original, 'parallel_columns', preserve_design=True)
    assert manifest['paired_slides'] == [2]
    assert {unit['slide'] for unit in manifest['units']} == {1, 2}
    output = assemble(original, tmp_path/'result.pptx', manifest, translated(manifest), 'parallel_columns')
    assert manifest['layout_plan'] == [{'source_slide':1,'role':'bilingual'},
                                      {'source_slide':2,'role':'vi'}, {'source_slide':2,'role':'en'}]
    assert hashlib.sha256(original.read_bytes()).hexdigest() == digest
    parts = source_slide_parts(output)
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(output) as after:
        assert before.read('ppt/slides/slide2.xml') == after.read(parts[1])
        for name in before.namelist():
            if name.startswith(('ppt/media/', 'ppt/theme/', 'ppt/slideMasters/', 'ppt/slideLayouts/')):
                assert before.read(name) == after.read(name)
        vi, en = (etree.fromstring(after.read(part)) for part in parts[1:])
        assert etree.tostring(vi.find('p:timing', NS)) == etree.tostring(en.find('p:timing', NS))
        assert vi.xpath('.//p:cNvPr/@id', namespaces=NS) == en.xpath('.//p:cNvPr/@id', namespaces=NS)
        assert vi.xpath('.//a:xfrm/a:off/@x', namespaces=NS) == en.xpath('.//a:xfrm/a:off/@x', namespaces=NS)
    deck = Presentation(output)
    assert len(deck.slides) == 3
    assert deck.slides[1].shapes[0].text.startswith('Năng lượng')
    assert deck.slides[2].shapes[0].text.startswith('The energy')
    assert deck.slides[1].shapes[1].image.sha1 == deck.slides[2].shapes[1].image.sha1
    report = inspect_returned_deck(output, original, 'parallel_columns', layout_plan=manifest['layout_plan'])
    assert not any(issue['code'] in {'slide_count', 'english_body_missing', 'parallel_columns_missing'}
                   for row in report['slide_quality']['slides'] for issue in row['issues'])


def test_paired_layout_still_detects_missing_visible_english(tmp_path):
    original = source(tmp_path/'source.pptx')
    manifest = extract(original, 'parallel_columns', preserve_design=True)
    output = assemble(original, tmp_path/'result.pptx', manifest, translated(manifest), 'parallel_columns')
    deck = Presentation(output)
    deck.slides[2].shapes[0].text = ''
    deck.save(output)
    report = inspect_returned_deck(output, original, 'parallel_columns', layout_plan=manifest['layout_plan'])
    assert any(issue['code'] == 'english_body_missing'
               for issue in report['slide_quality']['slides'][1]['issues'])


def test_invalid_plan_cannot_hide_missing_source_slides(tmp_path):
    original = source(tmp_path/'source.pptx')
    with pytest.raises(ValueError, match='Thiếu trang Anh'):
        inspect_returned_deck(original, original, 'parallel_columns',
                              layout_plan=[{'source_slide':1,'role':'vi'}, {'source_slide':2,'role':'en'}])


def test_google_conversion_keeps_layout_mapping_after_library_save(tmp_path):
    from contextlib import closing
    from app.conversion_input import creation_config
    from app.conversion_progress import ConversionProgress
    from app.google_conversion import convert, prepare, original_source
    from app.google_translation import KeyStore, Translator
    from app.library import Library

    original = source(tmp_path/'source.pptx')
    KeyStore(tmp_path).add('AIza' + 'a'*35)
    config = creation_config('Cơ năng', 'Vật lý', 'THPT', '11', 'parallel_columns', 'standard', 'source')
    job = prepare(tmp_path, config, original)
    translator = Translator(tmp_path, transport=lambda texts, *_: [
        'Mechanical energy' if text == 'Cơ năng' else
        'The energy of the object is conserved in the absence of friction.' for text in texts])
    lesson, preview = convert(tmp_path, job, Event(), ConversionProgress(job['folder'], lambda _: None), translator=translator)
    assert lesson['external_deck']['requested_style'] == 'source'
    assert lesson['external_deck']['total'] == 3
    plan = lesson['external_deck']['layout_plan']
    assert original_source(tmp_path,lesson).read_bytes()==original.read_bytes()
    wrong = copy.deepcopy(lesson)
    wrong['external_deck']['original_source_sha256']='0'*64
    assert original_source(tmp_path,wrong) is None
    again = inspect_returned_deck(preview['path'], conversion_format='parallel_columns', layout_plan=plan)
    assert not any(issue['code'] == 'english_body_missing'
                   for row in again['slide_quality']['slides'] for issue in row['issues'])
    with closing(Library(tmp_path)) as library:
        assert library.get(lesson['id'])['external_deck']['layout_plan'] == plan


def test_modern_math_body_is_translated_once_and_formula_and_fallback_survive(tmp_path):
    path = source(tmp_path/'modern.pptx')
    deck = Presentation(path)
    slide = deck.slides[1]
    shape = slide.shapes[0]
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = 'Vận tốc của vật là '
    math_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    formula = etree.SubElement(paragraph._p, f'{{{math_ns}}}oMath')
    run = etree.SubElement(formula, f'{{{math_ns}}}r')
    etree.SubElement(run, f'{{{math_ns}}}t').text = 'v=ωA'
    original = shape._element
    compatibility = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    wrapper = etree.Element(f'{{{compatibility}}}AlternateContent', nsmap={
        'mc':compatibility,'a14':'http://schemas.microsoft.com/office/drawing/2010/main'})
    choice = etree.SubElement(wrapper, f'{{{compatibility}}}Choice', Requires='a14')
    fallback = etree.SubElement(wrapper, f'{{{compatibility}}}Fallback')
    fallback.append(copy.deepcopy(original))
    position = original.getparent().index(original)
    original.getparent().insert(position,wrapper)
    choice.append(original)
    deck.save(path)
    manifest = extract(path, 'parallel_columns', preserve_design=True)
    unit = next(u for u in manifest['units'] if u['slide'] == 2)
    assert unit['text'].endswith('⟦BLC0⟧')
    values = translated(manifest)
    values[unit['id']] = {'en':'The velocity of the object is ⟦BLC0⟧'}
    output = assemble(path, tmp_path/'translated.pptx', manifest, values, 'parallel_columns')
    with zipfile.ZipFile(path) as before, zipfile.ZipFile(output) as after:
        source_tree = etree.fromstring(before.read(source_slide_parts(path)[1]))
        english = etree.fromstring(after.read(source_slide_parts(output)[2]))
        assert english.xpath('.//*[local-name()="Choice"]//a:t/text()',namespaces=NS)[0].startswith('The velocity')
        assert english.xpath('.//*[local-name()="oMath"]//*[local-name()="t"]/text()') == ['v=ωA','v=ωA']
        assert etree.tostring(source_tree.find(f'.//{{{compatibility}}}Fallback')) == etree.tostring(
            english.find(f'.//{{{compatibility}}}Fallback'))


def test_group_coordinates_are_measured_on_slide_not_in_child_coordinate_space(tmp_path):
    from app.slide_speech import objects

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(10), Inches(10), Inches(4), Inches(2))
    box.text = 'Nội dung trong nhóm'
    group.left,group.top,group.width,group.height = Inches(1),Inches(2),Inches(2),Inches(1)
    row = objects(slide)[0]
    assert (row['x'],row['y'],row['w'],row['h']) == (Inches(1),Inches(2),Inches(2),Inches(1))


def test_image_only_source_keeps_exact_slide_without_claiming_a_translation(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    image = tmp_path/'image.png'
    Image.new('RGB',(100,100),'#46A3CF').save(image)
    slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    original = tmp_path/'source.pptx'
    deck.save(original)
    manifest = extract(original,'parallel_columns',preserve_design=True)
    assert not manifest['units']
    output = assemble(original,tmp_path/'result.pptx',manifest,{},'parallel_columns')
    assert manifest['layout_plan'] == [{'source_slide':1,'role':'original'}]
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(output) as after:
        assert before.read('ppt/slides/slide1.xml') == after.read('ppt/slides/slide1.xml')


def test_measured_fitting_is_bounded_and_cannot_modify_an_original_vi_page(tmp_path):
    from app.native_layout import fit_preserved_frames

    original = source(tmp_path/'source.pptx')
    manifest = extract(original,'parallel_columns',preserve_design=True)
    output = assemble(original,tmp_path/'result.pptx',manifest,translated(manifest),'parallel_columns')
    with zipfile.ZipFile(output) as archive:
        before = {name:archive.read(name) for name in archive.namelist()}
    calls = []
    def measure(path,pages,cancel):
        calls.append(pages)
        return [(2,2,.1),(3,2,.1)]
    result = fit_preserved_frames(output,manifest,measure=measure)
    assert len(calls)==2 and all(2 not in pages for pages in calls)
    assert result['available'] and all(row['page']==3 for row in result['changes'])
    with zipfile.ZipFile(output) as archive:
        for name,data in before.items():
            if name != source_slide_parts(output)[2]:
                assert archive.read(name)==data
        tree = etree.fromstring(archive.read(source_slide_parts(output)[2]))
        scale = int(tree.find('.//a:normAutofit',NS).get('fontScale'))
        assert scale*26/100000 >= 12-.001


def test_missing_office_keeps_the_portable_preserved_deck(tmp_path):
    from app.native_layout import fit_preserved_frames

    original = source(tmp_path/'source.pptx')
    manifest = extract(original,'parallel_columns',preserve_design=True)
    output = assemble(original,tmp_path/'result.pptx',manifest,translated(manifest),'parallel_columns')
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    def unavailable(*args):
        raise RuntimeError('Office not installed')
    result = fit_preserved_frames(output,manifest,measure=unavailable)
    assert not result['available']
    assert hashlib.sha256(output.read_bytes()).hexdigest()==digest


def test_measured_fit_compacts_empty_paragraphs_without_changing_build_indices(tmp_path):
    from app.native_layout import fit_preserved_frames

    original = source(tmp_path/'source.pptx')
    deck = Presentation(original)
    deck.slides[1].shapes[0].text_frame.add_paragraph().text = '\v'
    deck.save(original)
    manifest = extract(original,'parallel_columns',preserve_design=True)
    output = assemble(original,tmp_path/'result.pptx',manifest,translated(manifest),'parallel_columns')
    before = Presentation(output).slides[2].shapes[0].text
    fit_preserved_frames(output,manifest,measure=lambda *_:[(3,2,.7)])
    after = Presentation(output).slides[2].shapes[0]
    assert after.text == before
    assert len(after.text_frame.paragraphs)==2
    assert after.text_frame.paragraphs[1]._p.find('a:pPr/a:defRPr',NS).get('sz')=='100'
    assert all(node.get('sz')=='100' for node in after.text_frame.paragraphs[1]._p.findall('.//a:rPr',NS))


def test_transparent_source_marker_keeps_its_visibility_and_is_not_spoken(tmp_path):
    from app.slide_speech import objects

    original = source(tmp_path/'source.pptx')
    deck = Presentation(original)
    paragraph = deck.slides[1].shapes[0].text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = ' 19vm568'
    color = etree.SubElement(etree.SubElement(run._r.get_or_add_rPr(),f'{{{NS["a"]}}}solidFill'),
                             f'{{{NS["a"]}}}srgbClr',val='FFFFFF')
    etree.SubElement(color,f'{{{NS["a"]}}}alpha',val='0')
    deck.save(original)
    manifest = extract(original,'parallel_columns',preserve_design=True)
    values = translated(manifest)
    unit = next(u for u in manifest['units'] if u['slide']==2)
    assert unit['numbers']==[' 19vm568']
    values[unit['id']]['en'] += ' ⟦BLC0⟧'
    output = assemble(original,tmp_path/'out.pptx',manifest,values,'parallel_columns')
    english = Presentation(output).slides[2]
    marker = english.shapes[0].text_frame.paragraphs[0].runs[-1]
    assert marker.text==' 19vm568'
    assert marker._r.find('.//a:alpha',NS).get('val')=='0'
    assert all('19vm568' not in item['text'] for item in objects(english))


def test_english_insets_reserve_space_beside_a_source_figure(tmp_path):
    original = source(tmp_path/'source.pptx')
    deck = Presentation(original)
    slide = deck.slides[1]
    box,picture = slide.shapes[0],slide.shapes[1]
    picture.left = box.left+box.width-Inches(.3)
    picture.top = box.top
    picture.height = box.height
    deck.save(original)
    manifest = extract(original,'parallel_columns',preserve_design=True)
    unit = next(u for u in manifest['units'] if u['slide']==2)
    assert unit['clip_insets']['rIns']>=Inches(.3)
    output = assemble(original,tmp_path/'out.pptx',manifest,translated(manifest),'parallel_columns')
    result = Presentation(output)
    assert result.slides[2].shapes[0].text_frame.margin_right>=Inches(.3)
    assert result.slides[2].shapes[1].left==picture.left
