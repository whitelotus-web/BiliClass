import hashlib
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.library import Library
from app.voice_import import import_presentation


def deck(path, text, notes=''):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if text:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = text
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    presentation.save(path)
    return path


def test_import_keeps_original_bytes_and_language_and_reuses_exact_file(tmp_path):
    source = deck(tmp_path / 'Bài có sẵn.pptx', 'Máy tính giúp xử lý thông tin và lưu trữ dữ liệu.')
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    directory = tmp_path / 'library'
    lesson, preview = import_presentation(directory, source)
    assert lesson['voice_only']
    assert lesson['external_deck']['provider'] == 'voice_import'
    assert lesson['segments'][0]['vi'] and not lesson['segments'][0]['en']
    assert not lesson['segments'][0]['approved']
    assert lesson['source']['sha256'] == before
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert hashlib.sha256(Path(preview['path']).read_bytes()).hexdigest() == before
    again, _ = import_presentation(directory, source)
    assert again['id'] == lesson['id']
    library = Library(directory)
    assert len(library.list_lessons()) == 1
    from app.export_origin import destination

    assert destination(library, again, directory) == (source.parent, source.stem)
    reviewed = library.review_external_deck(lesson['id'], lesson['revision'])
    assert reviewed['segments'][0]['approved']
    from app.readiness import text_readiness

    assert text_readiness(reviewed, directory)['text_ready']
    library.close()


def test_bilingual_notes_available_without_translation(tmp_path):
    source = deck(tmp_path / 'Bilingual.pptx', 'Computer', 'VI: Máy tính xử lý thông tin.\nEN: A computer processes information.')
    lesson, _ = import_presentation(tmp_path / 'library', source)
    assert lesson['segments'][0]['vi'] == 'Máy tính xử lý thông tin.'
    assert lesson['segments'][0]['en'] == 'A computer processes information.'


def test_no_text_can_be_completed_with_explicit_narration_and_wrong_type_rejected(tmp_path):
    source = deck(tmp_path / 'Empty.pptx', '')
    lesson, _ = import_presentation(tmp_path / 'library', source)
    assert not lesson['segments'][0]['vi'] and not lesson['segments'][0]['en']
    assert lesson['voice_only'] and not lesson['segments'][0]['approved']
    wrong = tmp_path / 'image.png'
    wrong.write_bytes(b'not a deck')
    with pytest.raises(ValueError):
        import_presentation(tmp_path / 'library', wrong)


def test_reimport_reuses_script_and_audio_without_parsing_and_changed_bytes_are_reinspected(tmp_path, monkeypatch):
    from contextlib import closing

    from app.voice_import import save_script

    source = deck(tmp_path/'Known.pptx', 'Năng lượng được bảo toàn.')
    directory = tmp_path/'library'
    lesson, _ = import_presentation(directory, source)
    with closing(Library(directory)) as library:
        lesson = save_script(library, lesson['id'], lesson['segments'][0]['id'], lesson['revision'],
                             'Lời thầy cô đã sửa.', 'An edited script.')
        audio = directory/'audio/keep.wav'
        audio.parent.mkdir(exist_ok=True)
        audio.write_bytes(b'cached audio')
        lesson['narration'] = {'mode': 'both', 'keys': {'vi': ['keep']}}
        library._write(lesson)
    def unexpected_parse(*_, **__):
        raise AssertionError('The cached import must not parse again')
    with monkeypatch.context() as patch:
        patch.setattr('app.voice_import.inspect_returned_deck', unexpected_parse)
        again, _ = import_presentation(directory, source)
    assert again['id'] == lesson['id'] and again['revision'] == lesson['revision']
    assert again['segments'] == lesson['segments'] and again['narration'] == lesson['narration']
    assert audio.read_bytes() == b'cached audio'
    deck(source, 'Nội dung khác hoàn toàn.')
    changed, _ = import_presentation(directory, source)
    assert changed['id'] != lesson['id'] and changed['source']['sha256'] != lesson['source']['sha256']


@pytest.mark.parametrize(('notes', 'vi', 'en'), [
    ('VI: Lời Việt\ntiếp theo\nCHECK: Không đọc dòng này', 'Lời Việt\ntiếp theo', ''),
    ('EN: An English script.\nSPEECH: do not read', '', 'An English script.'),
    ('BILICLASS_NOTES:\nVI: Lời Việt\nEN: An English script.', 'Lời Việt', 'An English script.'),
])
def test_narration_notes_read_once_and_new_import_skips_conversion_layout_review(tmp_path, monkeypatch, notes, vi, en):
    import pptx

    calls = []
    presentation = pptx.Presentation
    def parsed(*args, **kwargs):
        calls.append(1)
        return presentation(*args, **kwargs)
    def unexpected_scan(*_, **__):
        raise AssertionError('Current extraction already checked image narration')
    source = deck(tmp_path/'Notes.pptx', 'Nội dung hiển thị.', notes)
    monkeypatch.setattr(pptx, 'Presentation', parsed)
    monkeypatch.setattr('app.picture_narration._legacy_pages', unexpected_scan)
    monkeypatch.setattr('app.slide_quality.inspect_quality', unexpected_scan)
    lesson, _ = import_presentation(tmp_path/'library', source)
    assert lesson['segments'][0]['vi'] == vi and lesson['segments'][0]['en'] == en
    assert calls == [1] and not lesson['slide_quality']
    assert lesson['source_slide_speech']['extraction_version']
