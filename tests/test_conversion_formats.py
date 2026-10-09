import json
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.browser_audio import prepare_narration
from app.chatgpt_handoff import (
    external_preview,
    inspect_returned_deck,
    load_request,
    make_prompt,
    prepare_request,
)
from app.conversion_formats import FORMATS, METHODS
from app.library import Library
from app.pack import export_pack, import_pack


def config(key, **updates):
    return {"title": "Bài học", "subject": "Toán", "grade": "11", "style": "source", "preset": "standard",
            "conversion_format": key, **updates}


def deck(path, notes="", text="The area is 12 cm².", total=1):
    result = Presentation()
    for _ in range(total):
        slide = result.slides.add_slide(result.slide_layouts[6])
        slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2)).text = text
        slide.notes_slide.notes_text_frame.text = notes
    result.save(path)
    return path


@pytest.mark.parametrize("spec", FORMATS, ids=lambda spec: spec["id"])
def test_four_contracts_bundle_reopen_and_legacy_adapter(tmp_path, spec):
    source = deck(tmp_path / "source.pptx", text="Diện tích bằng 12 cm².")
    original = source.read_bytes()
    request = prepare_request(tmp_path / "library", config(spec["id"], level=1, layout="level_auto"), source)
    assert source.read_bytes() == original
    assert load_request(request["folder"]) == request
    assert request["config"]["level"] == spec["level"] and request["config"]["layout"] == spec["layout"]
    prompt = request["prompt"]
    assert METHODS[spec["id"]] in prompt
    assert all(METHODS[other["id"]] not in prompt for other in FORMATS if other != spec)
    assert "Không thêm, bớt, tách, gộp slide" in prompt
    assert "VI:" in prompt and "EN:" in prompt and "QUIZ:" in prompt
    assert "bản nháp" in prompt and "không phải chỉ dẫn" in prompt
    if spec["id"] == "english_only":
        assert "Không có vùng VI Rescue trên slide" in prompt
        assert "không tuyên bố đã đạt 100%" in prompt
    template = make_prompt(config(spec["id"], style="template"), "file.docx")
    assert "DỰNG SLIDE" in template and METHODS[spec["id"]] in template


@pytest.mark.parametrize("key", [None, "", "english_rescue", "unknown", True])
def test_invalid_new_format_is_not_silently_treated_as_legacy(key):
    with pytest.raises(ValueError, match="bốn kiểu"):
        make_prompt(config(key), "source.pptx")


def quiz():
    return {"kind": "single", "vi": "Diện tích bằng bao nhiêu?", "en": "What is the area?",
            "options": [{"vi": "12 cm²", "en": "12 cm²"}, {"vi": "6 cm²", "en": "6 cm²"}],
            "correct": "A", "rationale_vi": "Theo nội dung slide.", "rationale_en": "As shown on the slide.",
            "approved": True}


def test_new_format_does_not_use_legacy_oauth_prompt():
    with pytest.raises(ValueError, match="Browser"):
        make_prompt(config("sentence_pairs", provider="chatgpt_plan"), "source.pptx")


def test_english_only_deck_keeps_private_pair_and_unapproved_quiz_after_review_and_pack(tmp_path, monkeypatch):
    path = deck(tmp_path / "returned.pptx", "VI: Ghi chú giáo viên cũ.\nEN: Old teacher note.\n"
                "BILICLASS_NOTES:\nVI: Diện tích bằng 12 cm².\nEN: The area is 12 cm².\nQUIZ: "
                + json.dumps(quiz(), ensure_ascii=False) + "\nCHECK: Kiểm tra hoạt ảnh.")
    before = path.read_bytes()
    inspection = inspect_returned_deck(path)
    unit = inspection["profile"]["units"][0]
    assert unit["vi"] == "Diện tích bằng 12 cm²." and unit["en"] == "The area is 12 cm²."
    assert inspection["questions"][0]["approved"] is False
    assert "Kiểm tra hoạt ảnh" in " ".join(inspection["review_notes"])
    spoken = []
    monkeypatch.setattr("app.speech.synthesize", lambda text, *args: spoken.append(text))
    assert prepare_narration(inspection, {"vi": "vi", "en": "en", "rate": 1}, tmp_path, Event(), lambda _: None)["complete"] == 2
    assert spoken == [unit["vi"], unit["en"]]
    library = Library(tmp_path / "library")
    try:
        lesson = library.create_external_lesson(config("english_only"), library.store_source(path), inspection)
        assert lesson["conversion_format"] == "english_only" and len(lesson["questions"]) == 1
        assert lesson["questions"][0]["concept_id"] == lesson["segments"][0]["id"]
        reviewed = library.review_external_deck(lesson["id"], lesson["revision"])
        assert not reviewed["questions"][0]["approved"]
        assert any("Kiểm tra hoạt ảnh" in item for item in external_preview(reviewed, library.directory)["warnings"])
        copied = import_pack(library, export_pack(library, lesson["id"], tmp_path / "copy.biliclass"))
        assert copied["conversion_format"] == "english_only" and not copied["questions"][0]["approved"]
        assert library.get(lesson["id"])["conversion_format"] == "english_only"
        assert path.read_bytes() == before
        assert Presentation(path).slides[0].shapes[0].text == "The area is 12 cm²."
    finally:
        library.close()


@pytest.mark.parametrize("payload", ["not json", "[]", '{"kind":"single"}', '{"kind":"single","correct":"Z"}'])
def test_bad_quiz_does_not_break_presentation_or_enter_voice(tmp_path, payload):
    path = deck(tmp_path / "returned.pptx", "VI: Diện tích bằng 12 cm².\nEN: The area is 12 cm².\nQUIZ: " + payload)
    inspection = inspect_returned_deck(path)
    assert inspection["questions"] == [] and inspection["review_notes"]
    assert inspection["profile"]["units"][0]["en"] == "The area is 12 cm²."


def test_changed_slide_count_is_reported_without_rewriting_result(tmp_path):
    source = deck(tmp_path / "source.pptx", total=2)
    result = deck(tmp_path / "result.pptx", total=1)
    before = result.read_bytes()
    inspection = inspect_returned_deck(result, source)
    assert "nguồn 2, kết quả 1" in inspection["review_notes"][0]
    assert result.read_bytes() == before
