import json
import sqlite3

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.chatgpt_handoff import external_preview, inspect_returned_deck
from app.lesson_files import save_presentation
from app.library import Library


@pytest.fixture
def library(tmp_path):
    lib = Library(tmp_path / "library")
    yield lib
    lib.close()


def create(lib, title="Bài thứ nhất"):
    return lib.create(title, "Vật lý", "THPT", "11", [("Slide 1", "Năng lượng được bảo toàn.")])


def test_creation_date_is_immutable_across_edit_rename_restore_and_restart(library, monkeypatch):
    monkeypatch.setattr("app.library.now", lambda: "2026-10-01T03:00:00+00:00")
    first = create(library)
    monkeypatch.setattr("app.library.now", lambda: "2026-10-12T09:00:00+00:00")
    renamed = library.update_metadata(first["id"], "Tên mới", "Vật lý", "THPT", "11")
    assert renamed["created_at"] == first["created_at"] and renamed["updated_at"] != first["updated_at"]
    restored = library.restore_revision(first["id"], 1)
    assert restored["created_at"] == first["created_at"]
    reopened = Library(library.directory)
    try:
        assert reopened.get(first["id"])["created_at"] == first["created_at"]
    finally:
        reopened.close()


def test_chronological_creation_sort_does_not_change_when_old_lesson_is_edited(library, monkeypatch):
    monkeypatch.setattr("app.library.now", lambda: "2026-10-01T09:00:00+07:00")
    first = create(library)
    monkeypatch.setattr("app.library.now", lambda: "2026-10-01T03:00:00+00:00")
    second = create(library, "Bài thứ hai")
    monkeypatch.setattr("app.library.now", lambda: "2026-10-12T09:00:00+00:00")
    library.update_metadata(first["id"], "Tên vừa sửa", "Vật lý", "THPT", "11")
    assert [r["id"] for r in library.list_lessons("created_desc")] == [second["id"], first["id"]]
    assert [r["id"] for r in library.list_lessons("created_asc")] == [first["id"], second["id"]]
    assert library.list_lessons("updated_desc")[0]["id"] == first["id"]
    with pytest.raises(ValueError):
        library.list_lessons("untrusted SQL")


def test_legacy_creation_is_recovered_from_revision_content_not_edit_save_time(library, monkeypatch):
    monkeypatch.setattr("app.library.now", lambda: "2026-10-01T03:00:00+00:00")
    original = create(library)
    monkeypatch.setattr("app.library.now", lambda: "2026-10-12T09:00:00+00:00")
    library.update_metadata(original["id"], "Đã sửa", "Vật lý", "THPT", "11")
    # Emulate a pre-RC15 library: JSON never had a creation field.
    for table in ("lessons", "lesson_revisions"):
        key = "id" if table == "lessons" else "lesson_id"
        for row in library.db.execute(f"SELECT rowid,content FROM {table} WHERE {key}=?", (original["id"],)).fetchall():
            value = json.loads(row["content"])
            value.pop("created_at", None)
            library.db.execute(f"UPDATE {table} SET content=? WHERE rowid=?", (json.dumps(value), row["rowid"]))
    library.db.commit()
    reopened = Library(library.directory)
    try:
        assert reopened.get(original["id"])["created_at"] == "2026-10-01T03:00:00+00:00"
        assert not reopened.get(original["id"])["created_at_estimated"]
    finally:
        reopened.close()


def test_trash_restore_keeps_sources_audio_history_and_approved_memory(library, tmp_path):
    source = tmp_path / "original.txt"
    source.write_text("Năng lượng được bảo toàn.", encoding="utf-8")
    lesson = library.create("Bài cần xóa", "Vật lý", "THPT", "11", [("Đoạn", source.read_text(encoding="utf-8"))], library.store_source(source))
    lesson = library.edit_segment(lesson["id"], lesson["segments"][0]["id"], "Năng lượng được bảo toàn.", "The energy is conserved.", True)
    audio = library.directory / "audio" / "shared.wav"
    audio.parent.mkdir()
    audio.write_bytes(b"shared audio fixture")
    assert library.approved_translations("Vật lý", "Năng lượng được bảo toàn.")
    history = library.history(lesson["id"])
    library.trash_lesson(lesson["id"])
    assert not library.list_lessons()
    assert library.list_lessons(trash=True)[0]["id"] == lesson["id"]
    with pytest.raises(ValueError, match="thùng rác"):
        library.get(lesson["id"])
    assert not library.approved_translations("Vật lý", "Năng lượng được bảo toàn.")
    assert audio.read_bytes() == b"shared audio fixture"
    assert (library.directory / "sources" / lesson["source"]["file"]).read_bytes() == source.read_bytes()
    restored = library.restore_lesson(lesson["id"])
    assert restored["created_at"] == lesson["created_at"] and restored["segments"] == lesson["segments"]
    assert library.history(lesson["id"]) == history
    assert library.approved_translations("Vật lý", "Năng lượng được bảo toàn.")


def test_trash_cannot_remove_unfinished_classroom_lesson_and_keeps_ended_reports(library):
    lesson = create(library)
    path = library.directory / "classroom.db"
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE sessions (lesson TEXT, status TEXT)")
        db.execute("INSERT INTO sessions VALUES(?, 'active')", (json.dumps({"id": lesson["id"], "title": lesson["title"]}),))
    with pytest.raises(ValueError, match="Kết thúc tiết"):
        library.trash_lesson(lesson["id"])
    with sqlite3.connect(path) as db:
        db.execute("UPDATE sessions SET status='ended'")
    library.trash_lesson(lesson["id"])
    with sqlite3.connect(path) as db:
        assert json.loads(db.execute("SELECT lesson FROM sessions").fetchone()[0])["id"] == lesson["id"]


def external_lesson(library, tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = "The energy is conserved."
    slide.notes_slide.notes_text_frame.text = "VI: Năng lượng được bảo toàn.\nEN: The energy is conserved."
    path = tmp_path / "returned.pptx"
    deck.save(path)
    config = {"title": "Bài đã nhận", "subject": "Vật lý", "education_level": "THPT", "grade": "11",
              "conversion_format": "parallel_columns", "preset": "standard", "style": "source"}
    return path, library.create_external_lesson(config, library.store_source(path), inspect_returned_deck(path))


def test_download_saves_exact_viewed_file_even_as_draft_and_rename_keeps_review(library, tmp_path):
    source, lesson = external_lesson(library, tmp_path)
    preview = external_preview(lesson, library.directory)
    output = tmp_path / "download.pptx"
    save_presentation(lesson, preview, library.directory, output)
    assert output.read_bytes() == source.read_bytes()
    lesson = library.review_external_deck(lesson["id"], lesson["revision"])
    renamed = library.update_metadata(lesson["id"], "Tên mới", "Vật lý", "THPT", "11")
    assert not external_preview(renamed, library.directory)["draft"]
    assert renamed["created_at"] == lesson["created_at"]


def test_download_rejects_library_target_stale_revision_and_tampered_pptx(library, tmp_path):
    source, lesson = external_lesson(library, tmp_path)
    preview = external_preview(lesson, library.directory)
    with pytest.raises(ValueError, match="ngoài dữ liệu"):
        save_presentation(lesson, preview, library.directory, library.directory / "sources/overwrite.pptx")
    assert source.read_bytes() == (library.directory / "sources" / lesson["source"]["file"]).read_bytes()
    lesson["revision"] += 1
    with pytest.raises(ValueError, match="vừa được sửa"):
        save_presentation(lesson, preview, library.directory, tmp_path / "output.pptx")
    lesson["revision"] -= 1
    (library.directory / "sources" / lesson["source"]["file"]).write_bytes(b"tampered")
    with pytest.raises(ValueError, match="thay đổi"):
        save_presentation(lesson, preview, library.directory, tmp_path / "output.pptx")
    assert not (tmp_path / "output.pptx").exists()
