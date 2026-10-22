import threading
import time

from PySide6.QtCore import QCoreApplication

from app.input_documents import InputDocuments


def pump_until(predicate):
    app = QCoreApplication.instance() or QCoreApplication([])
    deadline = time.monotonic() + 5
    while not predicate():
        assert time.monotonic() < deadline
        app.processEvents()
        time.sleep(.002)
    app.processEvents()


def test_slow_file_selection_runs_off_gui_and_only_latest_choice_is_applied(monkeypatch):
    app = QCoreApplication.instance() or QCoreApplication([])
    entered, release = threading.Event(), threading.Event()
    calls, results = [], []
    gui_thread = threading.get_ident()

    def validate(urls):
        assert threading.get_ident() != gui_thread
        calls.append(urls)
        if urls == ['first']:
            entered.set()
            assert release.wait(3)
        return {'valid': True, 'name': urls[0]}

    monkeypatch.setattr('app.input_documents.document_selection', validate)
    picker = InputDocuments()
    picker.selected.connect(lambda result, error: results.append((result, error)))
    try:
        started = time.monotonic()
        picker.select(['first'])
        assert time.monotonic() - started < .1 and picker.busy
        pump_until(entered.is_set)
        picker.select(['obsolete'])
        picker.select(['latest'])
        app.processEvents()
        assert not results
        release.set()
        pump_until(lambda: not picker.busy)
        assert calls == [['first'], ['latest']]
        assert results == [({'valid': True, 'name': 'latest'}, '')]
    finally:
        release.set()
        picker.shutdown()


def test_clearing_selection_discards_late_result_and_reports_next_error(monkeypatch):
    app = QCoreApplication.instance() or QCoreApplication([])
    entered, release = threading.Event(), threading.Event()
    results = []

    def validate(urls):
        if urls == ['removed']:
            entered.set()
            assert release.wait(3)
            return {'valid': True}
        raise ValueError('Không đọc được tài liệu.')

    monkeypatch.setattr('app.input_documents.document_selection', validate)
    picker = InputDocuments()
    picker.selected.connect(lambda result, error: results.append((result, error)))
    try:
        picker.select(['removed'])
        pump_until(entered.is_set)
        picker.clear()
        assert not picker.busy
        release.set()
        pump_until(lambda: picker.worker is None)
        assert not results
        picker.select(['missing'])
        pump_until(lambda: not picker.busy)
        assert results == [({'valid': False}, 'Không đọc được tài liệu.')]
        app.processEvents()
    finally:
        release.set()
        picker.shutdown()


def test_source_preview_cancels_replaced_file_and_ignores_its_late_image(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from PySide6.QtCore import QObject

    app = QCoreApplication.instance() or QCoreApplication([])
    parent = QObject()
    parent.library = SimpleNamespace(directory=tmp_path)
    entered, cancelled = threading.Event(), threading.Event()
    gui_thread = threading.get_ident()

    def validate(urls):
        return {'valid': True, 'name': urls[0].split('/')[-1], 'url': urls[0]}

    def preview(url, _directory, cancel, number=1):
        assert threading.get_ident() != gui_thread
        if url.endswith('first.pptx'):
            entered.set()
            assert cancel.wait(3)
            cancelled.set()
        return {'image': url + '.png', 'total': 2}

    monkeypatch.setattr('app.input_documents.document_selection', validate)
    monkeypatch.setattr('app.input_documents.input_preview', preview)
    picker = InputDocuments(parent)
    try:
        picker.select(['file:///first.pptx'])
        pump_until(entered.is_set)
        assert not picker.busy and picker.preview['busy']
        picker.select(['file:///second.pptx'])
        pump_until(lambda: picker.preview.get('image') == 'file:///second.pptx.png')
        assert cancelled.is_set() and not picker.preview['busy']
        picker.clear()
        app.processEvents()
        assert not picker.preview
    finally:
        picker.shutdown()


def test_page_requests_cancel_the_old_page_and_keep_only_latest(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from PySide6.QtCore import QObject

    _app = QCoreApplication.instance() or QCoreApplication([])
    parent = QObject()
    parent.library = SimpleNamespace(directory=tmp_path)
    entered, cancelled = threading.Event(), threading.Event()

    def preview(url, directory, cancel, number=1):
        if number == 2:
            entered.set()
            assert cancel.wait(3)
            cancelled.set()
        return {'image': f'file:///slide-{number}.png', 'slide': number, 'total': 3}

    monkeypatch.setattr('app.input_documents.input_preview', preview)
    picker = InputDocuments(parent)
    picker.selected_url = 'file:///source.pptx'
    try:
        picker.retryPreview()
        pump_until(lambda: picker.preview.get('slide') == 1 and not picker.preview.get('busy'))
        picker.showSlide(2)
        pump_until(entered.is_set)
        picker.showSlide(3)
        pump_until(lambda: picker.preview.get('image') == 'file:///slide-3.png')
        assert cancelled.is_set() and not picker.busy
        picker.showSlide(4)
        assert picker.preview['slide'] == 3
        picker.clear()
        assert not picker.preview
    finally:
        picker.shutdown()


def test_source_preview_reuses_only_verified_pngs_and_prefetches_neighbours(tmp_path, monkeypatch):
    import hashlib
    import json
    from pathlib import Path

    from pptx import Presentation

    from app.input_documents import input_preview

    deck = Presentation()
    for _ in range(3):
        deck.slides.add_slide(deck.slide_layouts[6])
    path = tmp_path/'source.pptx'
    deck.save(path)
    calls = []
    monkeypatch.setattr('app.slide_fingerprint.environment_key', lambda: 'test-environment')

    def render(source, pages, directory, digest, cancel):
        calls.append(pages)
        folder = Path(directory)/'temp/lesson-preview-images'/digest[:24]
        folder.mkdir(parents=True, exist_ok=True)
        rows = {}
        for number in pages:
            content = b'\x89PNG\r\n\x1a\n' + str(number).encode()
            (folder/f'slide-{number}.png').write_bytes(content)
            rows[str(number)] = dict(image_sha256=hashlib.sha256(content).hexdigest(),
                                    environment='test-environment', date_sensitive=False)
        index = dict(version=3, sha256=digest, slides=rows)
        (folder/'preview-cache.json').write_text(json.dumps(index))
        return folder, index

    monkeypatch.setattr('app.powerpoint_review.render_preview_pages', render)
    cancel = threading.Event()
    first = input_preview(path.as_uri(), tmp_path, cancel)
    assert first['slide'] == 1 and first['total'] == 3 and calls == [[1, 2]]
    second = input_preview(path.as_uri(), tmp_path, cancel, 2)
    assert calls == [[1, 2]] and second['slide'] == 2
    from PySide6.QtCore import QUrl
    Path(QUrl(second['image']).toLocalFile()).write_bytes(b'damaged')
    input_preview(path.as_uri(), tmp_path, cancel, 2)
    assert calls == [[1, 2], [2, 3, 1]]
