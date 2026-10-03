import copy
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.bulk_translation import plan_batch
from app.input_analysis import analyze_document, assess_blocks
from app.library import Library, RevisionConflict
from app.quick_conversion import build_preview, verify_preview
from app.readiness import preparation_current, text_readiness
from app.translation import translate_document_text


def lesson_with_pair(library):
    blocks = [("Ý 1", "VI: Hàm số có 2 giá trị.\nEN: The function has 2 values.")]
    return library.create("Bài thử", "Toán", "THPT", "10", blocks, analysis=assess_blocks(blocks))


def test_preview_and_whole_lesson_review_are_separate_and_keep_other_approvals(tmp_path):
    library = Library(tmp_path)
    try:
        lesson = lesson_with_pair(library)
        support = {"id": "support", "approved": False, "kind": "prompt", "vi": "Gợi ý", "en": "Prompt"}
        lesson["segments"][0]["support"] = [support]
        lesson["questions"] = [{"id": "quiz", "approved": False}]
        library._write(lesson)
        before = copy.deepcopy(lesson)
        preview = build_preview(lesson, tmp_path)
        assert preview["total"] > 0 and preview["draft"]
        assert len(preview["segment_map"]) == preview["total"]
        assert verify_preview(lesson, preview, tmp_path).is_file()
        assert library.get(lesson["id"]) == before and lesson == before
        approved = library.review_lesson(lesson["id"], preview["revision"])
        assert text_readiness(approved, tmp_path)["text_ready"] and preparation_current(approved)
        assert approved["segments"][0]["support"] == [support]
        assert approved["questions"] == before["questions"]
        assert approved["segments"][0]["source_text"] == before["segments"][0]["source_text"]
    finally:
        library.close()


def test_review_rejects_incomplete_stale_or_changed_source_without_saving(tmp_path):
    library = Library(tmp_path)
    try:
        lesson = lesson_with_pair(library)
        preview = build_preview(lesson, tmp_path)
        segment = lesson["segments"][0]
        changed = library.edit_segment(lesson["id"], segment["id"], segment["vi"], "")
        with pytest.raises(RevisionConflict):
            library.review_lesson(lesson["id"], lesson["revision"])
        with pytest.raises(ValueError, match="Chưa đủ"):
            library.review_lesson(lesson["id"], changed["revision"])
        with pytest.raises(ValueError, match="vừa được sửa"):
            verify_preview(changed, preview, tmp_path)
        assert library.get(lesson["id"]) == changed
        library.directory.joinpath("sources").mkdir(exist_ok=True)
        (library.directory / "sources/source.txt").write_text("changed", encoding="utf-8")
        lesson["source"] = {"file": "source.txt", "sha256": "different"}
        library._write(lesson)
        with pytest.raises(ValueError, match="Bản nguồn"):
            library.review_lesson(lesson["id"], lesson["revision"])
        assert library.get(lesson["id"]) == lesson
    finally:
        library.close()


def test_modified_generated_file_cannot_be_approved_as_the_preview(tmp_path):
    library = Library(tmp_path)
    try:
        lesson = lesson_with_pair(library)
        result = build_preview(lesson, tmp_path)
        verify_preview(lesson, result, tmp_path).write_bytes(b"modified")
        with pytest.raises(ValueError, match="trình chiếu đã thay đổi"):
            verify_preview(lesson, result, tmp_path)
        assert not library.get(lesson["id"])["segments"][0]["approved"]
    finally:
        library.close()


def test_inline_heading_is_paired_but_a_native_one_row_table_remains_a_grid(tmp_path):
    heading = "CHƯƠNG II: DÃY SỐ, CẤP SỐ CỘNG VÀ CẤP SỐ NHÂN | CHAPTER II: SEQUENCES, ARITHMETIC & GEOMETRIC PROGRESSIONS"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(9), Inches(1)).text = heading
    table = slide.shapes.add_table(1, 2, Inches(.5), Inches(2), Inches(9), Inches(1)).table
    table.cell(0, 0).text = "Hàm số có một giá trị."
    table.cell(0, 1).text = "The function has one value."
    path = tmp_path / "source.pptx"
    deck.save(path)
    units = analyze_document(path)["profile"]["units"]
    assert units[0]["existing_pair"] and "CHAPTER II" in units[0]["en"]
    assert not units[1]["existing_pair"] and "|" in units[1]["vi"]


def test_native_preview_preserves_the_original_and_keeps_source_warnings(tmp_path):
    source = tmp_path / "teacher.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(9), Inches(1)).text = "VI: Hàm số có 2 giá trị.\nEN: The function has 2 values."
    deck.save(source)
    original = source.read_bytes()
    data = analyze_document(source)
    library = Library(tmp_path / "library")
    try:
        reference = library.store_source(source)
        reference["warnings"] = ["Cần kiểm tra chữ trong ảnh."]
        lesson = library.create("Bài thử", "Toán", "THPT", "10", data["blocks"], reference,
                                data["profile"]["primary_language"], data["profile"])
        lesson = library.set_presentation(lesson["id"], 2, "level_auto")
        lesson = library.set_conversion_mode(lesson["id"], "level")
        result = build_preview(lesson, library.directory)
        assert result["slide_map"] and result["draft"]
        assert "Cần kiểm tra chữ trong ảnh." in result["warnings"]
        assert source.read_bytes() == original
        assert (library.directory / "sources" / reference["file"]).read_bytes() == original
        assert library.get(lesson["id"]) == lesson
    finally:
        library.close()


def test_whole_document_plans_more_than_fifty_missing_units(tmp_path):
    library = Library(tmp_path)
    try:
        blocks = [(f"Ý {i}", f"Hàm số có {i} giá trị.") for i in range(65)]
        lesson = library.create("Bài dài", "Toán", "THPT", "10", blocks)
        assert len(plan_batch(library, lesson)) == 50
        assert len(plan_batch(library, lesson, limit=None)) == 65
    finally:
        library.close()


def test_long_translation_preserves_content_tables_and_formula_boundaries(monkeypatch):
    calls = []

    class Tokenizer:
        def encode(self, value, out_type=str):
            return value.split()

    def translate(span, language, terms, tokenizer, translator):
        assert len(span) <= 2000 and len(tokenizer.encode(span)) <= 350
        calls.append(span)
        return span

    monkeypatch.setattr("app.translation.translate_with_resources", translate)
    text = "Một giá trị " * 500 + "`f(x) = x + 25`\nHàm số | Giá trị\n3 | 5"
    result = translate_document_text(text, "vi", [], Tokenizer(), object())
    assert result == text and len(calls) > 5
    assert sum("`f(x) = x + 25`" in part for part in calls) == 1
    cancel = Event()
    cancel.set()
    with pytest.raises(ValueError, match="Đã dừng"):
        translate_document_text(text, "vi", [], Tokenizer(), object(), cancel)
