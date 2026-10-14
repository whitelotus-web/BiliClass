import io
import os
import struct
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.pptx_parts import merge
from app.pptx_transport import MP4_PLACEHOLDER, media_references, proxy, verify_media


def movie_deck(path):
    deck=Presentation()
    poster=io.BytesIO()
    Image.new('RGB',(80,80),'blue').save(poster,format='PNG')
    for i in range(2):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        payload=os.urandom(300*1024)
        video=MP4_PLACEHOLDER+struct.pack('>I',len(payload)+8)+b'free'+payload
        slide.shapes.add_movie(io.BytesIO(video),Inches(1),Inches(1),Inches(8),Inches(4),
            poster_frame_image=io.BytesIO(poster.getvalue()),mime_type='video/mp4')
    deck.save(path)
    return path


def test_light_upload_restores_exact_original_videos_and_relations(tmp_path):
    original=movie_deck(tmp_path/'original.pptx')
    light=tmp_path/'proxy.pptx'
    plan=proxy(original,light)
    assert plan['offloaded_bytes']>590*1024 and len(plan['restore'])==2
    assert plan['upload_bytes']<plan['original_bytes']*.1
    assert len(set(plan['restore']))==2
    verify_media(light,plan['media_references'],plan['restore'])
    assert media_references(original)==plan['media_references']
    target=merge(original,[(light,[1,2],plan['restore'])],tmp_path/'restored.pptx')
    assert media_references(target)==media_references(original)
    with zipfile.ZipFile(original) as old,zipfile.ZipFile(target) as new:
        for name in old.namelist():
            if name.startswith('ppt/media/'):
                assert old.read(name)==new.read(name)
        assert not any(b'BILICLASS_LOCAL_MEDIA:' in new.read(name) for name in new.namelist() if name.startswith('ppt/media/'))


def test_dropped_video_shape_is_not_hidden_by_orphaned_media_parts(tmp_path):
    original=movie_deck(tmp_path/'original.pptx')
    light=tmp_path/'proxy.pptx'
    plan=proxy(original,light)
    deck=Presentation(light)
    shape=deck.slides[0].shapes[0]
    shape._element.getparent().remove(shape._element)
    deck.save(light)
    with pytest.raises(ValueError):
        verify_media(light,plan['media_references'],plan['restore'])


def test_proxy_never_overwrites_teacher_source(tmp_path):
    original=movie_deck(tmp_path/'original.pptx')
    before=original.read_bytes()
    with pytest.raises(ValueError):
        proxy(original,original)
    assert original.read_bytes()==before
