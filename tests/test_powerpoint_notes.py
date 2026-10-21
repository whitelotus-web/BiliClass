import hashlib

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

from app.chatgpt_handoff import inspect_returned_deck
from app.native_bilingual import assemble, extract
from app.powerpoint_notes import read_text, text_frame
from app.slide_speech import inspect_slide


def without_body(slide, *, remove_master=False):
    notes = slide.notes_slide
    shape = notes.notes_placeholder.element
    shape.getparent().remove(shape)
    if remove_master:
        for placeholder in list(notes.part.notes_master.placeholders):
            if placeholder.placeholder_format.type == PP_PLACEHOLDER.BODY:
                shape = placeholder.element
                shape.getparent().remove(shape)
    assert notes.notes_text_frame is None


def test_reading_notes_does_not_create_missing_notes_pages():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    assert not slide.has_notes_slide and read_text(slide) == '' and not slide.has_notes_slide


@pytest.mark.parametrize('remove_master', [False, True])
def test_repair_missing_body_retains_other_notes_shapes_and_never_duplicates(remove_master):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    without_body(slide, remove_master=remove_master)
    before = {shape.shape_id: shape.element.xml for shape in slide.notes_slide.shapes}
    assert read_text(slide) == ''
    frame = text_frame(slide)
    frame.text = 'Ghi chú của giáo viên.\nVI: Lực.\nEN: Force.'
    assert read_text(slide).startswith('Ghi chú của giáo viên.')
    assert all(shape.element.xml == before[shape.shape_id] for shape in slide.notes_slide.shapes
               if shape.shape_id in before)
    count = len(slide.notes_slide.shapes)
    assert text_frame(slide).text == frame.text and len(slide.notes_slide.shapes) == count


@pytest.mark.parametrize('mode', ['parallel_columns', 'english_only'])
def test_both_conversion_and_narration_accept_a_deck_with_deleted_notes_body(tmp_path, mode):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = 'Đường thẳng và mặt phẳng'
    without_body(slide)
    source = tmp_path/'teacher-notes.pptx'
    deck.save(source)
    digest = hashlib.sha256(source.read_bytes()).digest()
    imported = inspect_returned_deck(source)
    assert imported['profile']['units']
    inspect_slide(Presentation(source).slides[0], 1)
    manifest = extract(source, mode, preserve_design=True)
    translations = {unit['id']: {'en': 'Lines and planes'} for unit in manifest['units']}
    result = assemble(source, tmp_path/'converted.pptx', manifest, translations, mode)
    converted = Presentation(result)
    assert all(slide.notes_slide.notes_text_frame is not None for slide in converted.slides)
    assert all('BILICLASS_NOTES:' in read_text(slide) for slide in converted.slides)
    assert hashlib.sha256(source.read_bytes()).digest() == digest
