import base64
import hashlib
import io
import json
import lzma
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import pptx_text_transport


def data_image(data,nonce):
    import math
    import numpy as np

    packet=b'BLCIMG02'+bytes.fromhex(nonce)+len(data).to_bytes(4,'big')+hashlib.sha256(data).digest()+data
    bits=np.unpackbits(np.frombuffer(packet,dtype=np.uint8),bitorder='big')
    height=math.ceil(len(bits)/(1024*3))
    pixels=np.pad(bits,(0,height*1024*3-len(bits))).reshape(height,1024,3)*255
    image=Image.new('RGB',(1056,height+32),'white')
    image.paste(Image.fromarray(pixels.astype(np.uint8)),(16,16))
    return image


@pytest.mark.parametrize('fault',['none','nonce','hash','scale','truncated','oversized'])
def test_data_image_requires_exact_bytes_and_current_request(tmp_path,fault):
    nonce='1234567890abcdef'*2
    data='{"text":"Năng lượng"}'.encode('utf-8')
    if fault=='oversized':data=b'x'*240001
    image=data_image(data,nonce)
    if fault=='nonce':nonce='0'*32
    elif fault=='hash':image.putpixel((16+170,16),(255,255,255))
    elif fault=='scale':image=image.resize((528,image.height//2))
    elif fault=='truncated':image=image.crop((0,0,1056,32))
    path=tmp_path/'data.png'
    image.save(path)
    if fault=='none':assert pptx_text_transport.decode_image(path,nonce)==data.decode('utf-8')
    else:
        with pytest.raises(ValueError):pptx_text_transport.decode_image(path,nonce)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def envelope(payload):
    raw=json.dumps(payload,ensure_ascii=False,separators=(',',':')).encode()
    return json.dumps({'format':'biliclass-ooxml-patch-xz-v1','sha256':digest(raw),'unpacked_bytes':len(raw),
                       'data_b64':base64.b64encode(lzma.compress(raw)).decode()})


@pytest.fixture
def patch_fixture(tmp_path):
    source=tmp_path/'source.pptx'
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2)).text='Vận tốc'
    image=io.BytesIO()
    Image.new('RGB',(80,60),'blue').save(image,format='PNG')
    slide.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(1),Inches(3))
    slide.notes_slide.notes_text_frame.text='VI: Vận tốc'
    deck.save(source)
    with zipfile.ZipFile(source) as archive:
        old=archive.read('ppt/slides/slide1.xml')
        count=len(archive.namelist())
    target=old.replace('Vận tốc'.encode(),'Velocity'.encode())
    index=old.index('Vận tốc'.encode())
    part={'name':'ppt/slides/slide1.xml','source_sha256':digest(old),'sha256':digest(target),
          'edits':[[index,index+len('Vận tốc'.encode()),base64.b64encode(b'Velocity').decode()]]}
    payload={'format':'biliclass-ooxml-patch-v1','source_sha256':digest(source.read_bytes()),
             'target_part_count':count,'removed':[],'parts':[part]}
    return source,payload,target


def test_reconstructs_exact_xml_and_keeps_every_unchanged_part_byte_for_byte(patch_fixture,tmp_path):
    source,payload,target=patch_fixture
    original=source.read_bytes()
    path=pptx_text_transport.apply(source,tmp_path/'received.pptx',envelope(payload))
    assert source.read_bytes()==original
    assert Presentation(path).slides[0].shapes[0].text=='Velocity'
    with zipfile.ZipFile(source) as old,zipfile.ZipFile(path) as new:
        assert set(old.namelist())==set(new.namelist())
        for name in old.namelist():
            assert new.read(name)==(target if name=='ppt/slides/slide1.xml' else old.read(name))


@pytest.mark.parametrize('change',['wrong-source','wrong-part','wrong-target','overlap','outside','traversal',
    'binary','duplicate','count','remove-required','remove-binary','invalid-xml','external','executable'])
def test_invalid_or_unsafe_change_never_creates_a_deck(patch_fixture,tmp_path,change):
    source,payload,_=patch_fixture
    part=payload['parts'][0]
    with zipfile.ZipFile(source) as archive:
        old=archive.read(part['name'])
        if change=='wrong-source':payload['source_sha256']='0'*64
        elif change=='wrong-part':part['source_sha256']='0'*64
        elif change=='wrong-target':part['sha256']='0'*64
        elif change=='overlap':part['edits']*=2
        elif change=='outside':part['edits'][0][1]=len(old)+1
        elif change=='traversal':part['name']='../outside.xml'
        elif change=='binary':part['name']='ppt/media/image1.png'
        elif change=='duplicate':payload['parts']*=2
        elif change=='count':payload['target_part_count']+=1
        elif change=='remove-required':payload['removed']=['ppt/presentation.xml']
        elif change=='remove-binary':payload['removed']=['ppt/media/image1.png']
        else:
            if change=='invalid-xml':new=b'<!DOCTYPE x [<!ENTITY a "bad">]><x>&a;</x>'
            elif change=='external':
                part['name']='ppt/slides/_rels/slide1.xml.rels'
                old=archive.read(part['name'])
                new=old.replace(b'</Relationships>',b'<Relationship Id="rId999" Type="unsafe" TargetMode="External" Target="file:///private"/></Relationships>')
            else:
                part['name']='[Content_Types].xml'
                old=archive.read(part['name'])
                new=old.replace(b'</Types>',b'<Override PartName="/ppt/evil.xml" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            part.update(source_sha256=digest(old),sha256=digest(new),edits=[[0,len(old),base64.b64encode(new).decode()]])
    output=tmp_path/'received.pptx'
    with pytest.raises(ValueError):pptx_text_transport.apply(source,output,envelope(payload))
    assert not output.exists() and not list(tmp_path.glob('*.pending.pptx'))


def test_corrupt_truncated_oversized_and_concatenated_envelopes_are_rejected(patch_fixture):
    _,payload,_=patch_fixture
    data=json.loads(envelope(payload))
    for changed in [dict(data,sha256='0'*64),dict(data,unpacked_bytes=1),dict(data,data_b64='invalid'),
        dict(data,data_b64=data['data_b64'][:-8]),dict(data,data_b64=data['data_b64']+'AAAA')]:
        with pytest.raises((ValueError,lzma.LZMAError)):
            pptx_text_transport.decode(json.dumps(changed))


def test_source_is_never_overwritten(patch_fixture):
    source,payload,_=patch_fixture
    with pytest.raises(ValueError,match='ghi đè'):
        pptx_text_transport.apply(source,source,envelope(payload))


def test_compact_whole_xml_transfer_has_the_same_validation_and_preserves_media(patch_fixture,tmp_path):
    source,payload,target=patch_fixture
    payload['format']='biliclass-ooxml-text-v1'
    payload['parts']=[{'name':'ppt/slides/slide1.xml','sha256':digest(target),'xml':target.decode('utf-8')}]
    encoded=json.loads(envelope(payload))
    encoded['format']='biliclass-ooxml-xz-v1'
    output=pptx_text_transport.apply(source,tmp_path/'from-text.pptx',json.dumps(encoded))
    with zipfile.ZipFile(source) as old,zipfile.ZipFile(output) as new:
        assert new.read('ppt/slides/slide1.xml')==target
        for name in old.namelist():
            if name.startswith('ppt/media/'):
                assert old.read(name)==new.read(name)
