import hashlib
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.native_bilingual import assemble, extract


@pytest.mark.parametrize('kind,rotation,text,expected', [(5,0,True,True), (5,90,True,False), (5,0,False,False), (13,0,True,False)])
def test_measure_and_repair_accept_native_freeform_text_only(kind, rotation, text, expected):
    from types import SimpleNamespace

    from app.native_layout import fittable_text

    shape = SimpleNamespace(Type=kind, Rotation=rotation, HasTextFrame=True, TextFrame2=SimpleNamespace(HasText=text))
    assert bool(fittable_text(shape)) == expected


def source(path, *, picture=False, cramped=False):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(12), Inches(.9))
    title.text = 'Cơ năng'
    title.text_frame.paragraphs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].font.name = 'Arial'
    body = slide.shapes.add_textbox(Inches(.5), Inches(2), Inches(9), Inches(1.8 if not cramped else .35))
    body.text = 'Vận tốc của vật thay đổi theo thời gian trong quá trình dao động.'
    body.text_frame.paragraphs[0].font.size = Pt(40)
    body.text_frame.paragraphs[0].font.name = 'Arial'
    if picture:
        # An opaque foreground shape blocks expansion exactly like a picture.
        from pptx.enum.shapes import MSO_SHAPE
        block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.5), body.top+body.height,
                                       Inches(9), Inches(2))
        block.fill.solid()
    deck.save(path)
    return path


@pytest.mark.parametrize('mode', ['parallel_columns', 'sentence_pairs', 'integrated_keywords', 'english_only'])
def test_native_fit_uses_readable_text_and_preserves_package(tmp_path, mode):
    path = source(tmp_path/'source.pptx')
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = extract(path, mode)
    assert manifest['version'] == 2 and not manifest['visual_slides']
    translations = {u['id']: {'en': 'Mechanical energy' if u['title'] else
                             'The velocity changes over time during oscillation.',
                             'mixed': 'Cơ năng (mechanical energy)' if u['title'] else
                             'Vận tốc (velocity) của vật thay đổi theo thời gian.'} for u in manifest['units']}
    result = assemble(path, tmp_path/'result.pptx', manifest, translations, mode)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    deck = Presentation(result)
    shapes = [s for s in deck.slides[0].shapes if s.has_text_frame]
    sizes = [p.font.size.pt for s in shapes for p in s.text_frame.paragraphs if p.font.size]
    sizes += [r.font.size.pt for s in shapes for p in s.text_frame.paragraphs for r in p.runs if r.font.size]
    assert min(sizes) >= 18 and min(sizes) < 40
    if mode == 'parallel_columns':
        vi, en = shapes[1:]
        assert vi.left+vi.width < en.left
        assert vi.top == en.top and vi.height == en.height
    with zipfile.ZipFile(path) as before, zipfile.ZipFile(result) as after:
        assert all(before.read(n) == after.read(n) for n in before.namelist()
                   if n.startswith(('ppt/media/', 'ppt/theme/', 'ppt/slideLayouts/', 'ppt/slideMasters/')))


def test_a_blocked_tiny_box_stays_visual_instead_of_hiding_overlap(tmp_path):
    path = source(tmp_path/'source.pptx', picture=True, cramped=True)
    manifest = extract(path, 'parallel_columns')
    assert manifest['visual_slides'] == [1] and not manifest['units']


def test_old_saved_manifest_does_not_silently_change_layout_algorithm(tmp_path):
    path = source(tmp_path/'source.pptx')
    manifest = extract(path, 'english_only')
    manifest['version'] = 1
    values = {u['id']: {'en': 'An extended description of the experiment and its observations.'} for u in manifest['units']}
    result = assemble(path, tmp_path/'old.pptx', manifest, values, 'english_only')
    shape = Presentation(result).slides[0].shapes[1]
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 40


def test_text_fit_leaves_rotated_text_unchanged(tmp_path):
    from app.native_layout import fit_body

    path = source(tmp_path/'rotated.pptx', cramped=True)
    shape = Presentation(path).slides[0].shapes[1]
    shape.rotation = 45
    before = shape.element.xml
    fit_body(shape.element, 4000)
    assert shape.element.xml == before


def test_measured_math_fit_changes_font_properties_not_only_ignored_autofit(tmp_path):
    from lxml import etree

    from app.native_layout import fit_preserved_frames
    from app.source_deck import NS, source_slide_parts

    path = source(tmp_path/'formula.pptx')
    deck = Presentation(path)
    shape = deck.slides[0].shapes[1]
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = ''
    paragraph._p.append(etree.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<m:r><m:rPr><m:sty m:val="b"/></m:rPr><a:rPr sz="4400"/><m:t>x²</m:t></m:r></m:oMath>'
    ))
    deck.save(path)
    calls = []

    def measure(*_):
        calls.append(1)
        return [(1, shape.shape_id, .7)] if len(calls) == 1 else []

    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = fit_preserved_frames(path, {'layout_plan': [{'role': 'bilingual'}], 'same_slide': True}, measure=measure)
    assert result['changes'] and len(calls) == 2
    assert hashlib.sha256(path.read_bytes()).hexdigest() != before
    with zipfile.ZipFile(path) as archive:
        tree = etree.fromstring(archive.read(source_slide_parts(path)[0]))
    assert tree.xpath('.//*[local-name()="oMath"]//*[local-name()="t"]/text()') == ['x²']
    size = int(tree.xpath('.//*[local-name()="oMath"]//a:rPr/@sz', namespaces=NS)[0])
    assert 1400 <= size < 4400
