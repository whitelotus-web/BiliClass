from pptx import Presentation
from pptx.util import Inches

from app.slideshow_controls import prepare_controls, speaker_index


def test_dense_slides_without_speakers_never_traverse_office_shapes(tmp_path):
    deck = Presentation()
    for _ in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for _ in range(50):
            slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).name = 'Repeated source name'
    path = tmp_path / 'Many shapes.pptx'
    deck.save(path)

    class NoOfficeCalls:
        def __getattr__(self, name):
            raise AssertionError('Navigation must not scan source shapes: ' + name)

    media, badges, other = prepare_controls(NoOfficeCalls(), path)
    assert media == {slide.slide_id: ([], []) for slide in deck.slides}
    assert not badges and not other


def test_index_keeps_all_speakers_and_separate_control_positions(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for name in ['Source', 'BiliClass audio EN', 'BiliClass control EN', 'BiliClass audio VI']:
        slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).name = name
    path = tmp_path / 'Speakers.pptx'
    deck.save(path)
    index, other = speaker_index(path)
    assert index == {slide.slide_id: [(2, slide.shapes[1].shape_id, 'en', 3, None),
                                     (4, slide.shapes[3].shape_id, 'vi', 4, None)]}
    assert not other


def test_badge_failure_does_not_disable_indexed_audio(monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setattr('app.slideshow_controls.speaker_index', lambda path: ({256: [(1, 9, 'en', 1, None)]}, False))
    shape = SimpleNamespace(Id=9, Left=10, Top=20, Width=60, Height=40)

    class Shapes:
        def Item(self, position):
            assert position == 1
            return shape

        def AddPicture(self, *args):
            raise RuntimeError('Office cannot add an overlay')

    slides = SimpleNamespace(FindBySlideID=lambda identity: SimpleNamespace(Shapes=Shapes()))
    warnings = []
    media, badges, other = prepare_controls(SimpleNamespace(Slides=slides), 'fixture', on_warning=lambda: warnings.append(1))
    assert media == {256: ([(9, 'en', shape)], [('en', 10., 20., 60., 40.)])}
    assert not badges and not other and warnings == [1]


def test_named_controls_survive_office_shape_order_differences(monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setattr('app.slideshow_controls.speaker_index', lambda path: ({256: [(5, 9, 'en', 6, 7)]}, False))
    media = SimpleNamespace(Id=9)
    ready = SimpleNamespace(Left=10, Top=20, Width=50, Height=20, Visible=True)
    badge = SimpleNamespace(Visible=True)

    class Shapes:
        def Item(self, name):
            assert isinstance(name, str), 'Stale positions must not choose a source shape'
            return {'BiliClass audio EN':media, 'BiliClass control EN':ready, 'BiliClass pause EN':badge}[name]

    slides = SimpleNamespace(FindBySlideID=lambda _:SimpleNamespace(Shapes=Shapes()))
    players, badges, _ = prepare_controls(SimpleNamespace(Slides=slides), 'fixture')
    assert players[256][0] == [(9, 'en', media)]
    assert badges[(256,9)] == (ready,badge) and badge.Visible is False


def test_source_audio_keeps_native_media_controls(tmp_path):
    import wave

    from app.paths import RESOURCE_ROOT

    sound = tmp_path / 'source.wav'
    with wave.open(str(sound), 'wb') as audio:
        audio.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
        audio.writeframes(b'\0\0' * 2400)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_movie(str(sound), 0, 0, Inches(1), Inches(1),
                           poster_frame_image=str(RESOURCE_ROOT / 'assets/actions/speaker-en.png'),
                           mime_type='audio/wav')
    path = tmp_path / 'Source media.pptx'
    deck.save(path)
    index, other = speaker_index(path)
    assert index == {slide.slide_id: []} and other


def test_index_never_reads_binary_media_and_follows_presentation_order(tmp_path, monkeypatch):
    from zipfile import ZipFile

    from app.paths import RESOURCE_ROOT

    deck = Presentation()
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(RESOURCE_ROOT / 'assets/actions/speaker-en.png'), 0, 0).name = 'Source image'
    # Slide part filenames are unchanged after reordering; slide IDs are stable.
    ids = deck.slides._sldIdLst
    ids.insert(0, ids[-1])
    path = tmp_path / 'Reordered images.pptx'
    deck.save(path)
    read = ZipFile.read
    opened = []

    def metadata_only(archive, name, *args, **kwargs):
        opened.append(name)
        assert '/media/' not in name
        return read(archive, name, *args, **kwargs)

    monkeypatch.setattr(ZipFile, 'read', metadata_only)
    index, other = speaker_index(path)
    assert list(index) == [slide.slide_id for slide in deck.slides]
    assert len(opened) == 4 and not other
