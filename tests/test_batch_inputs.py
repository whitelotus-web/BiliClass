from threading import Event

from pptx import Presentation

from app.batch_inputs import input_details


def test_metadata_reads_slide_count_without_changing_input(tmp_path):
    source = tmp_path / 'Bài có dấu, #1.pptx'
    deck = Presentation()
    for _ in range(4):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    original = source.read_bytes()
    result = input_details([dict(id='first', url=source.as_uri())], Event())
    assert result == {'first': {'total': 4, 'bytes': len(original)}}
    assert source.read_bytes() == original


def test_bad_input_does_not_hide_other_files_and_cancel_stops_scan(tmp_path):
    source = tmp_path / 'bad.pptx'
    source.write_bytes(b'not a presentation')
    result = input_details([dict(id='missing', url=(tmp_path/'missing.pptx').as_uri()),
                            dict(id='bad', url=source.as_uri())], Event())
    assert result['missing']['error'] and result['bad']['error']
    cancel = Event()
    cancel.set()
    assert input_details([dict(id='bad', url=source.as_uri())], cancel) == {}
