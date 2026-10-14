import io
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.pptx_parts import merge, subset
from app.source_deck import NS


def source_deck(path):
    deck=Presentation()
    image=io.BytesIO()
    Image.new('RGB',(80,80),'red').save(image,format='PNG')
    for i in range(4):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text=f'Source {i+1}'
        slide.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(1),Inches(3))
        slide.notes_slide.notes_text_frame.text=f'Original note {i+1}'
    deck.save(path)
    return path


def test_subsets_and_merge_keep_order_notes_and_deduplicate_original_media(tmp_path):
    source=source_deck(tmp_path/'source.pptx')
    part=subset(source,[2,4],tmp_path/'part.pptx')
    deck=Presentation(part)
    assert len(deck.slides)==2
    for index,slide in enumerate(deck.slides):
        slide.shapes[0].text=f'Translated {index+1}'
    deck.save(part)
    target=merge(source,[(part,[2,4])],tmp_path/'out.pptx')
    result=Presentation(target)
    assert [s.shapes[0].text for s in result.slides]==['Source 1','Translated 1','Source 3','Translated 2']
    assert [s.notes_slide.notes_text_frame.text for s in result.slides]==[
        'Original note 1','Original note 2','Original note 3','Original note 4']
    with zipfile.ZipFile(source) as old,zipfile.ZipFile(target) as new:
        media=[old.read(n) for n in old.namelist() if n.startswith('ppt/media/')]
        assert all(sum(new.read(n)==blob for n in new.namelist() if not n.endswith('.rels'))==1 for blob in media)
        assert old.read('ppt/slides/slide1.xml')==new.read('ppt/slides/slide1.xml')
        assert old.read('ppt/slides/slide3.xml')==new.read('ppt/slides/slide3.xml')


def test_mismatched_or_duplicate_slide_results_are_rejected(tmp_path):
    source=source_deck(tmp_path/'source.pptx')
    part=subset(source,[2],tmp_path/'part.pptx')
    with pytest.raises(ValueError):
        merge(source,[(part,[1,2])],tmp_path/'out.pptx')
    with pytest.raises(ValueError):
        merge(source,[(part,[2]),(part,[2])],tmp_path/'out.pptx')


def test_original_part_names_can_differ_from_slide_order(tmp_path):
    source=source_deck(tmp_path/'source.pptx')
    with zipfile.ZipFile(source) as archive:
        entries={name:archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(source,'w') as archive:
        for name,data in entries.items():
            archive.writestr(name.replace('slide1.xml','slide91.xml'),
                            data.replace(b'slide1.xml',b'slide91.xml') if name.endswith(('.xml','.rels')) else data)
    part=subset(source,[1,3],tmp_path/'part.pptx')
    deck=Presentation(part)
    deck.slides[0].shapes[0].text='Replacement first'
    deck.save(part)
    target=merge(source,[(part,[1,3])],tmp_path/'merged.pptx')
    result=Presentation(target)
    assert [s.shapes[0].text for s in result.slides]==['Replacement first','Source 2','Source 3','Source 4']
    assert result.slides[0].notes_slide.notes_text_frame.text=='Original note 1'


def test_changed_master_registers_unique_master_and_layout_ids_and_one_notes_master(tmp_path):
    source=source_deck(tmp_path/'source.pptx')
    part=subset(source,[2],tmp_path/'part.pptx')
    deck=Presentation(part)
    deck.slide_master._element.set('name','Changed incoming master')
    deck.save(part)
    target=merge(source,[(part,[2])],tmp_path/'out.pptx')
    with zipfile.ZipFile(target) as archive:
        root=etree.fromstring(archive.read('ppt/presentation.xml'))
        masters=root.xpath('.//p:sldMasterId/@id',namespaces=NS)
        assert len(masters)==2
        assert len(root.xpath('.//p:notesMasterId',namespaces=NS))==1
        ids=list(masters)
        for name in archive.namelist():
            if '/slideMasters/' in name and name.endswith('.xml'):
                ids.extend(etree.fromstring(archive.read(name)).xpath('.//p:sldLayoutId/@id',namespaces=NS))
        assert len(ids)==len(set(ids))
