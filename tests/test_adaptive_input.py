import copy
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.deck_export import export_deck
from app.input_analysis import analyze_document, assess_blocks
from app.lesson_templates import slide_pages
from app.library import Library
from app.pack import export_pack, import_pack
from app.presentation_policy import presentation_content
from app.source_deck import NS, export_source_deck


def make_lesson(tmp_path, texts, *, dense=False, animated=False, rotated=False):
    source = tmp_path / "teacher.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
    for text in texts:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(11), Inches(5.7 if dense else .9))
        box.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        if rotated:
            box.rotation = 10
        if animated:
            slide._element.append(etree.Element(f"{{{NS['p']}}}timing"))
    deck.save(source)
    data = analyze_document(source)
    library = Library(tmp_path / "library")
    lesson = library.create("Bài thử", "Toán", "THPT", "10", data["blocks"], library.store_source(source),
                            data["profile"]["primary_language"], data["profile"])
    return source, library, lesson


@pytest.mark.parametrize(("blocks", "expected"), [
    ([("Slide 1", "Hàm số là một quy tắc cho giá trị.")], "vi"),
    ([("Slide 1", "A function gives each input a value.")], "en"),
    ([("Slide 1", "VI: Hàm số là một quy tắc.\nEN: A function is a rule.")], "bilingual"),
    ([("Slide 1 · ý 1", "Hàm số là một quy tắc."),
      ("Slide 1 · ý 2", "A function is a rule.")], "bilingual"),
    ([("Slide 1", "Có 2 giá trị."), ("Slide 2", "There are 3 values.")], "mixed"),
    ([("Ý 1", "Bonjour tout le monde")], "unknown"),
])
def test_language_assessment_routes_inputs_without_approving_guesses(blocks, expected):
    profile = assess_blocks(blocks, "pptx")
    assert profile["language"] == expected
    assert profile["recommended_mode"] == ("preserve" if expected in {"mixed", "bilingual"} else "level")
    assert profile["requires_review"]
    if expected == "bilingual":
        assert all(unit["vi"] and unit["en"] for unit in profile["units"])
        assert profile["warnings"]


def test_conflicting_numbers_and_table_grid_are_not_paired():
    data = assess_blocks([("Slide 1 · ý 1", "Hàm số có 2 giá trị."),
                          ("Slide 1 · ý 2", "The function has 3 values.")], "pptx")
    assert data["language"] == "mixed"
    assert not any(unit["paired_locator"] for unit in data["units"])
    table = "Hàm số | Function\nGiá trị | Value"
    assert not assess_blocks([("Slide 1", table)])["units"][0]["existing_pair"]
    formula = assess_blocks([("Slide 2", "25% + 3 = 28")])["units"][0]
    assert formula["vi"] == formula["en"] == "25% + 3 = 28"


def test_mixed_source_review_and_project_roundtrip_preserve_languages(tmp_path):
    source, library, lesson = make_lesson(tmp_path, ["Hàm số là một quy tắc.", "Each input has one value."])
    try:
        vi, en = lesson["segments"]
        assert en["source_language"] == "en" and not en["vi"]
        with pytest.raises(ValueError, match="nguồn"):
            library.edit_segment(lesson["id"], en["id"], "Giá trị", "", False)
        library.save_term("Toán", "Hàm số", "function")
        library.set_setting("voice_en", "test-voice")
        lesson = library.edit_segment(lesson["id"], vi["id"], vi["vi"], "A function is a rule.", True)
        lesson = library.edit_segment(lesson["id"], en["id"], "Mỗi đầu vào có một giá trị.", en["en"], True)
        lesson = library.set_presentation(lesson["id"], 3, "level_auto")
        lesson = library.set_conversion_mode(lesson["id"], "level")
        pack = export_pack(library, lesson["id"], tmp_path / "shared.biliclass")
        recipient = Library(tmp_path / "recipient")
        try:
            imported = import_pack(recipient, pack)
            assert [s["source_language"] for s in imported["segments"]] == ["vi", "en"]
            assert imported["layout"] == "level_auto" and imported["conversion_mode"] == "level"
            assert imported["project_terms"][0]["en"] == "function"
            assert imported["project_preferences"]["voice_en"] == "test-voice"
            assert not any(s["approved"] for s in imported["segments"])
            assert not recipient.glossary() and recipient.setting("voice_en", "unchanged") == "unchanged"
            assert source.read_bytes() == (recipient.directory / "sources" / imported["source"]["file"]).read_bytes()
        finally:
            recipient.close()
    finally:
        library.close()


def reviewed(lesson, library, english="A function is a rule. Each input has one value."):
    for segment in lesson["segments"]:
        lesson = library.edit_segment(lesson["id"], segment["id"], segment["vi"], english, True)
    return lesson


def test_sparse_slide_adds_only_level_support_without_covering_original(tmp_path):
    source, library, lesson = make_lesson(tmp_path, ["Hàm số là một quy tắc."])
    try:
        original = source.read_bytes()
        lesson = reviewed(lesson, library)
        lesson.update(conversion_mode="level", layout="level_auto", level=0)
        terms = [{"vi": "Hàm số", "en": "function"}]
        result = export_source_deck(lesson, library.directory, tmp_path / "L0.pptx", terms)
        deck = Presentation(result["path"])
        assert result["slide_map"] == [1] and result["report"][0]["action"] == "added"
        original_box, added = deck.slides[0].shapes
        assert original_box.text == "Hàm số là một quy tắc."
        assert "Hàm số — function" in added.text and "A function is a rule" not in added.text
        assert added.top >= original_box.top + original_box.height or added.left >= original_box.left + original_box.width
        assert source.read_bytes() == original
        lesson["level"] = 2
        result = export_source_deck(lesson, library.directory, tmp_path / "L2.pptx")
        assert "A function is a rule." in Presentation(result["path"]).slides[0].shapes[-1].text
        assert "Each input" not in Presentation(result["path"]).slides[0].shapes[-1].text
        assert any("chưa có câu Anh dễ" in warning for warning in result["report"][0]["warnings"])
        lesson["segments"][0]["support"] = [{"kind": "easy_en", "en": "An input has one value.", "approved": True}]
        result = export_source_deck(lesson, library.directory, tmp_path / "L2-easy.pptx")
        assert "An input has one value." in Presentation(result["path"]).slides[0].shapes[-1].text
        assert not result["report"][0]["warnings"]
    finally:
        library.close()


@pytest.mark.parametrize("reason", ["dense", "animated", "rotated"])
def test_complex_slide_keeps_source_xml_and_paginates_all_support(tmp_path, reason):
    source, library, lesson = make_lesson(tmp_path, ["Hàm số là một quy tắc."], **{reason: True})
    try:
        lesson = reviewed(lesson, library, "A function gives each input exactly one value. " * 45)
        lesson.update(conversion_mode="level", layout="level_auto", level=3)
        result = export_source_deck(lesson, library.directory, tmp_path / "L3.pptx")
        deck = Presentation(result["path"])
        assert len(deck.slides) > 2 and all(value == 1 for value in result["slide_map"])
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(result["path"]) as exported:
            assert original.read("ppt/slides/slide1.xml") == exported.read("ppt/slides/slide1.xml")
        text = " ".join(slide.shapes[0].text.split("\n", 1)[1] for slide in list(deck.slides)[1:])
        assert " ".join(text.split()) == " ".join(lesson["segments"][0]["en"].split())
        assert all(slide._element.find("p:timing", NS) is None for slide in list(deck.slides)[1:])
    finally:
        library.close()


def test_existing_bilingual_has_no_duplicate_english_and_l4_has_rescue_in_project(tmp_path):
    source, library, lesson = make_lesson(tmp_path, ["VI: Hàm số là một quy tắc.\nEN: A function is a rule."])
    try:
        assert lesson["conversion_mode"] == "preserve"
        assert not lesson["segments"][0]["approved"]
        preserve = export_deck(lesson, library.directory, tmp_path / "unchanged.pptx")
        assert source.read_bytes() == preserve.read_bytes()
        segment = lesson["segments"][0]
        lesson = library.edit_segment(lesson["id"], segment["id"], segment["vi"], segment["en"], True)
        lesson.update(conversion_mode="level", layout="level_auto", level=3)
        result = export_source_deck(lesson, library.directory, tmp_path / "L3.pptx")
        assert len(Presentation(result["path"]).slides) == 1
        assert result["report"][0]["action"] == "existing"
        original = copy.deepcopy(lesson)
        lesson["level"] = 4
        result = export_source_deck(lesson, library.directory, tmp_path / "L4.pptx")
        assert Presentation(result["path"]).slides[0].shapes[0].text == "A function is a rule."
        assert lesson["segments"] == original["segments"]
        assert "VI Rescue" in result["report"][0]["warnings"][0]
    finally:
        library.close()


def test_paired_choice_remains_two_languages_at_l0(tmp_path):
    _, library, lesson = make_lesson(tmp_path, ["Hàm số là một quy tắc."])
    try:
        lesson = reviewed(lesson, library, "A function is a rule.")
        lesson.update(conversion_mode="paired", layout="level_auto", level=0)
        result = export_source_deck(lesson, library.directory, tmp_path / "paired.pptx")
        assert result["slide_map"] == [1, 1]
        assert Presentation(result["path"]).slides[1].shapes[0].text == "A function is a rule."
    finally:
        library.close()


@pytest.mark.parametrize("level", range(5))
def test_template_auto_level_uses_reviewed_support_in_preview_and_export(tmp_path, level):
    segment = {"id": "test-segment", "vi": "Hàm số là một quy tắc.", "en": "A function is a rule. Each input has one value.",
               "approved": True, "locator": "Ý 1", "kind": "concept", "support": [
                   {"kind": "vocabulary", "vi": "Hàm số", "en": "function", "approved": True},
                   {"kind": "prompt", "vi": "Quan sát ví dụ.", "en": "Look at the example.", "approved": True},
                   {"kind": "easy_en", "en": "One input has one value.", "approved": True},
                   {"kind": "prompt", "en": "UNREVIEWED", "approved": False},
               ]}
    lesson = {"title": "Bài hàm số", "subject": "Toán", "grade": "10", "level": level,
              "layout": "level_auto", "teaching_preset": "standard", "segments": [segment]}
    content = presentation_content(segment, level, "level_auto")
    expected = ("Hàm số (function)", "Look at the example.", "One input has one value.",
                "A function is a rule.", "A function is a rule.")[level]
    plans = slide_pages(lesson, segment)
    preview = "\n".join(item["text"] for plan in plans for item in plan["elements"] if item.get("body"))
    assert expected in preview and "UNREVIEWED" not in preview
    exported = Presentation(export_deck(lesson, tmp_path, tmp_path / f"L{level}.pptx"))
    output = "\n".join(shape.text for slide in exported.slides for shape in slide.shapes if shape.has_text_frame)
    assert expected in output and "UNREVIEWED" not in output
    assert content["show_vi"] == (level < 4)


def test_image_only_powerpoint_ocr_routes_to_template_and_preserves_source(tmp_path, monkeypatch):
    image = tmp_path / "scan.png"
    Image.new("RGB", (1200, 675), "white").save(image)
    source = tmp_path / "scan.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(image), 0, 0,
                                                                 width=deck.slide_width, height=deck.slide_height)
    deck.save(source)
    before = source.read_bytes()
    calls = []

    def ocr(path, language, cancelled=None):
        calls.append(language)
        assert path.is_file()
        return [("Ảnh 1", "Hàm số là một quy tắc.")]

    monkeypatch.setattr("app.ocr.recognize", ocr)
    data = analyze_document(source, "vi")
    assert calls == ["vi"] and data["profile"]["recommended_style"] == "template"
    assert "OCR" in data["blocks"][0][0] and data["profile"]["objects"]["image_only_slides"] == 1
    assert source.read_bytes() == before


def test_vietnamese_ocr_uses_local_backend_when_windows_only_has_english(tmp_path, monkeypatch):
    from app.ocr import _recognize

    image = tmp_path / "vietnamese.png"
    Image.new("RGB", (100, 100), "white").save(image)
    selected = []
    monkeypatch.setattr("app.ocr.languages", lambda: ["en-US"])
    monkeypatch.setattr("app.local_ocr.engine", lambda language: selected.append(language) or object())
    monkeypatch.setattr("app.local_ocr.image_text", lambda engine, path: "Hàm số là một quy tắc.")
    result = _recognize(image, "vi")
    assert selected == ["vi"] and result[0][1] == "Hàm số là một quy tắc."
    assert "OCR cần kiểm tra" in result[0][0]


def test_local_ocr_setup_rejects_wrong_checksum_without_replacing_existing_model(tmp_path, monkeypatch):
    from app.local_ocr import LATIN_MODEL, prepare_models

    monkeypatch.setattr("app.translation.model_directory", lambda: tmp_path)
    folder = tmp_path / "ocr-vi-en-v1"
    folder.mkdir()
    target = folder / LATIN_MODEL
    target.write_bytes(b"previous model")

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def raise_for_status(self):
            pass

        def iter_bytes(self):
            yield b"broken download"

    monkeypatch.setattr("httpx.stream", lambda *args, **kwargs: Response())
    with pytest.raises(ValueError, match="Checksum"):
        prepare_models()
    assert target.read_bytes() == b"previous model"
    assert list(folder.iterdir()) == [target]
