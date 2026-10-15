import hashlib
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from app.native_bilingual import assemble, extract, leaves
from app.source_deck import NS


def fixture(path):
    deck=Presentation()
    deck.slide_width,deck.slide_height=Inches(13.333),Inches(7.5)
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(.6),Inches(.4),Inches(12),Inches(1))
    title.text='Cơ năng'
    title.text_frame.paragraphs[0].font.name='Arial'
    title.text_frame.paragraphs[0].font.size=Pt(32)
    box=slide.shapes.add_textbox(Inches(.6),Inches(2),Inches(12),Inches(4))
    box.text='Vật có vận tốc 12 m/s.'
    box.text_frame.paragraphs[0].font.name='Arial'
    box.text_frame.paragraphs[0].font.size=Pt(24)
    deck.save(path)
    return path


@pytest.mark.parametrize('mode',['parallel_columns','sentence_pairs','integrated_keywords','english_only'])
def test_four_modes_are_editable_keep_source_and_have_paired_notes(tmp_path,mode):
    source=fixture(tmp_path/'source.pptx')
    original=hashlib.sha256(source.read_bytes()).hexdigest()
    manifest=extract(source,mode)
    assert not manifest['visual_slides']
    translations={row['id']:{'en':'Mechanical energy' if row['title'] else 'The object has a velocity of ⟦BLC0⟧ m/s.',
        **({'mixed':'Cơ năng (mechanical energy)' if row['title'] else 'Vật có vận tốc (velocity) ⟦BLC0⟧ m/s.'}
           if mode=='integrated_keywords' else {})} for row in manifest['units']}
    target=assemble(source,tmp_path/'result.pptx',manifest,translations,mode)
    deck=Presentation(target)
    assert len(deck.slides)==1
    shapes=[s for s in leaves(deck.slides[0].shapes) if s.has_text_frame]
    text='\n'.join(s.text for s in shapes)
    assert '12 m/s' in text
    assert ('Vật có' in text)==(mode!='english_only')
    assert ('velocity' in text)
    notes=deck.slides[0].notes_slide.notes_text_frame.text
    assert 'BILICLASS_NOTES:' in notes and 'VI:' in notes and 'EN:' in notes
    assert hashlib.sha256(source.read_bytes()).hexdigest()==original
    if mode=='parallel_columns':
        vi=next(s for s in shapes if 'Vật có' in s.text)
        en=next(s for s in shapes if 'The object' in s.text)
        assert en.left >= vi.left+vi.width
    with zipfile.ZipFile(source) as old,zipfile.ZipFile(target) as new:
        for name in old.namelist():
            if name.startswith(('ppt/theme/','ppt/slideMasters/','ppt/slideLayouts/','ppt/media/')):
                assert old.read(name)==new.read(name)


def test_animation_slides_requiring_extra_text_are_routed_to_visual_tasks(tmp_path):
    source=fixture(tmp_path/'source.pptx')
    deck=Presentation(source)
    etree.SubElement(deck.slides[0]._element,f'{{{NS["p"]}}}timing')
    deck.save(source)
    paired=extract(source,'parallel_columns')
    assert paired['visual_slides']==[1] and not paired['units']
    english=extract(source,'english_only')
    assert english['units'] and not english['visual_slides']
    translations={u['id']:{'en':'Mechanical energy' if u['title'] else 'Velocity is ⟦BLC0⟧ m/s.'} for u in english['units']}
    target=assemble(source,tmp_path/'result.pptx',english,translations,'english_only')
    with zipfile.ZipFile(source) as old,zipfile.ZipFile(target) as new:
        before=etree.fromstring(old.read('ppt/slides/slide1.xml')).find('p:timing',NS)
        after=etree.fromstring(new.read('ppt/slides/slide1.xml')).find('p:timing',NS)
        assert etree.tostring(before)==etree.tostring(after)


def test_missing_translation_or_overwriting_source_is_rejected(tmp_path):
    source=fixture(tmp_path/'source.pptx')
    manifest=extract(source,'english_only')
    with pytest.raises(ValueError):
        assemble(source,tmp_path/'out.pptx',manifest,{},'english_only')
    with pytest.raises(ValueError):
        assemble(source,source,manifest,{},'english_only')


def test_already_english_source_keeps_exact_text_runs_without_a_visual_rewrite(tmp_path):
    source = fixture(tmp_path / 'english-source.pptx')
    deck = Presentation(source)
    for shape in deck.slides[0].shapes:
        shape.text = '\n'.join('The energy is conserved in the system.' for _ in range(4))
    deck.save(source)
    manifest = extract(source, 'english_only')
    # Use generous frames so this test isolates language routing, not overflow.
    for shape in deck.slides[0].shapes:
        shape.height = Inches(3)
        shape.text_frame.paragraphs[0].font.size = Pt(18)
    deck.save(source)
    manifest = extract(source, 'english_only')
    assert not manifest['visual_slides'] and len(manifest['units']) == 8
    values = {u['id']: {'en': u['text']} for u in manifest['units']}
    result = assemble(source, tmp_path / 'english-copy.pptx', manifest, values, 'english_only')
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result) as after:
        assert before.read('ppt/slides/slide1.xml') == after.read('ppt/slides/slide1.xml')


def test_leading_superscript_is_kept_but_not_applied_to_all_translated_words(tmp_path):
    from app.native_bilingual import translated_paragraph
    from app.source_deck import _bytes

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    paragraph = shape.text_frame.paragraphs[0]
    exponent = paragraph.add_run()
    exponent.text = '2'
    exponent._r.get_or_add_rPr().set('baseline', '30000')
    ordinary = paragraph.add_run()
    ordinary.text = ' Gia tốc'
    ordinary.font.name = 'Arial'
    unit = {'numbers': ['2'], 'xml_locks': {'0': _bytes(exponent._r).decode('utf-8')}}
    translated = translated_paragraph(paragraph._p, '⟦BLC0⟧ Acceleration', unit, english=True)
    runs = translated.findall('a:r', NS)
    assert runs[0].find('a:rPr', NS).get('baseline') == '30000'
    assert runs[-1].find('a:rPr', NS).get('baseline', '0') == '0'
    assert runs[-1].find('a:rPr/a:latin', NS).get('typeface') == 'Arial'


def test_native_assembly_keeps_nonsequential_part_names_and_formula_runs(tmp_path):
    source=fixture(tmp_path/'source.pptx')
    deck=Presentation(source)
    p=deck.slides[0].shapes[1].text_frame.paragraphs[0]
    p.text='Đại lượng x'
    run=p.add_run()
    run.text='2'
    run._r.get_or_add_rPr().set('baseline','30000')
    deck.save(source)
    with zipfile.ZipFile(source) as archive:
        entries={name:archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(source,'w') as archive:
        for name,data in entries.items():
            archive.writestr(name.replace('slide1.xml','slide91.xml'),
                            data.replace(b'slide1.xml',b'slide91.xml') if name.endswith(('.xml','.rels')) else data)
    manifest=extract(source,'english_only')
    assert manifest['units'] and {u['part'] for u in manifest['units']}=={'ppt/slides/slide91.xml'}
    values={u['id']:{'en':'Mechanical energy' if u['title'] else 'Quantity x⟦BLC0⟧'} for u in manifest['units']}
    target=assemble(source,tmp_path/'result.pptx',manifest,values,'english_only')
    result=Presentation(target)
    assert result.slides[0].shapes[1].text=='Quantity x2'
    assert result.slides[0].shapes[1].text_frame.paragraphs[0].runs[-1]._r.rPr.get('baseline')=='30000'
    assert 'Quantity x2' in result.slides[0].notes_slide.notes_text_frame.text
