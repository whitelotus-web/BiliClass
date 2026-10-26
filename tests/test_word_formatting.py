from contextlib import closing
from pathlib import Path
from threading import Event

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from app import word_jobs as jobs
from app.library import Library
from app.word_document import NS, Package, keyword_annotations


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    monkeypatch.setattr('app.translation_knowledge.lookup', lambda *_: None)
    source = tmp_path / 'Bài học.docx'
    document = Document()
    document.styles['Normal'].font.name = 'Arial'
    run = document.add_paragraph().add_run('Vận tốc 12 m/s.')
    run.font.name = 'Cambria'
    run.bold = True
    run.font.color.rgb = RGBColor.from_string('9C1111')
    document.add_paragraph('Giữ nguyên phần này.')
    document.save(source)
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    package = Package(source)
    region = package.selection([dict(paragraph='p0', start=0, end=len(package.rows[0]['text']))])[0]
    jobs.selections(directory, job['id'], [region])
    return directory, source, job, region


@pytest.mark.parametrize('mode', ['below', 'columns', 'keywords', 'english'])
@pytest.mark.parametrize('italic,bold', [(False, False), (True, False), (False, True), (True, True)])
@pytest.mark.parametrize('font,color', [('', 'original'), ('Arial', '1667B1')])
def test_saved_english_uses_selected_style_and_original_or_custom_font_color(prepared, tmp_path, mode, italic, bold, font, color):
    _, source, _, region = prepared
    original = source.read_bytes()
    region['en'] = 'Velocity 12 m/s.'
    if mode == 'keywords':
        region['en'] = region['vi']
        region['annotations'] = keyword_annotations(region, {'vận tốc': 'velocity'})
    output = tmp_path / 'formatted.docx'
    Package(source).save(output, [region], dict(mode=mode, italic=italic, bold=bold, font=font, color=color))
    tree = Package(output).tree
    translated = [run for run in tree.findall('.//w:r', NS)
                  if 'velocity' in ''.join(run.xpath('./w:t/text()', namespaces=NS)).lower()]
    assert translated
    for run in translated:
        assert run.find('w:rPr/w:i', NS).get('{' + NS['w'] + '}val') == str(int(italic))
        assert run.find('w:rPr/w:b', NS).get('{' + NS['w'] + '}val') == str(int(bold))
        fonts = run.find('w:rPr/w:rFonts', NS)
        assert fonts.get('{' + NS['w'] + '}ascii') == (font or 'Cambria')
        assert fonts.get('{' + NS['w'] + '}hAnsi') == (font or 'Cambria')
        assert run.find('w:rPr/w:color', NS).get('{' + NS['w'] + '}val') == ('9C1111' if color == 'original' else color)
        if font:
            assert not any('Theme' in key for key in fonts.attrib)
    assert 'Giữ nguyên phần này.' in '\n'.join(row['text'] for row in Package(output).rows)
    assert source.read_bytes() == original


def test_formatting_defaults_survive_reopen_and_apply_only_to_new_documents(prepared):
    directory, source, first, _ = prepared
    assert jobs.formatting_defaults(directory) == dict(color='original', font='', italic=True, bold=False, underline=False, strike=False, script='normal', font_size=0)
    recent = dict(color='7030A0', font='Arial', italic=False, bold=True, underline=True, strike=True, script='superscript', font_size=26)
    jobs.formatting_defaults(directory, recent)
    assert jobs.formatting_defaults(directory) == recent
    other = source.with_name('Bài tiếp theo.docx')
    document = Document(source)
    document.add_paragraph('Nội dung mới.')
    document.save(other)
    new = jobs.create(directory, other)
    assert all(new['options'][key] == value for key, value in recent.items())
    assert jobs.read(directory, first['id'])['options']['font'] == ''
    assert jobs.create(directory, source)['options']['color'] == 'original'
    (jobs.root(directory) / 'formatting.json').write_text('broken', encoding='utf-8')
    assert jobs.formatting_defaults(directory)['color'] == 'original'


@pytest.mark.parametrize('mode,script', [('below', 'subscript'), ('columns', 'superscript'), ('english', 'normal'), ('keywords', 'superscript')])
def test_word_toolbar_styles_and_palette_color_reach_export_without_changing_source(tmp_path, mode, script):
    document = Document()
    document.add_paragraph('Vận tốc 12 m/s.')
    source = tmp_path / 'source.docx'
    document.save(source)
    original = source.read_bytes()
    package = Package(source)
    region = package.selection([dict(paragraph='p0', start=0, end=len(package.rows[0]['text']))])[0]
    region['en'] = 'Velocity 12 m/s.'
    if mode == 'keywords':
        region['annotations'] = keyword_annotations(region, {'vận tốc': 'velocity'})
    output = tmp_path / 'output.docx'
    package.save(output, [region], dict(mode=mode, color='7030A0', underline=True, strike=True, script=script))
    exported = Package(output)
    runs = [run for run in exported.tree.findall('.//w:r', NS)
            if 'velocity' in ''.join(run.xpath('./w:t/text()', namespaces=NS)).lower()]
    assert runs
    for run in runs:
        assert run.find('w:rPr/w:color', NS).get('{' + NS['w'] + '}val') == '7030A0'
        assert run.find('w:rPr/w:u', NS).get('{' + NS['w'] + '}val') == 'single'
        assert run.find('w:rPr/w:strike', NS).get('{' + NS['w'] + '}val') == '1'
        assert run.find('w:rPr/w:vertAlign', NS).get('{' + NS['w'] + '}val') == ('baseline' if script == 'normal' else script)
    assert source.read_bytes() == original


@pytest.mark.parametrize('mode', ['below', 'columns', 'english', 'keywords'])
@pytest.mark.parametrize('size', [0, 12.5, 26])
def test_size_keeps_source_by_default_and_applies_explicit_points_in_all_modes(tmp_path, mode, size):
    document = Document()
    document.styles['Normal'].font.size = Pt(14)
    document.add_paragraph().add_run('Vận tốc 12 m/s.').font.size = Pt(18)
    source = tmp_path / 'source.docx'
    document.save(source)
    original = source.read_bytes()
    package = Package(source)
    region = package.selection([dict(paragraph='p0', start=0, end=len(package.rows[0]['text']))])[0]
    region['en'] = 'Velocity 12 m/s.'
    if mode == 'keywords':
        region['annotations'] = keyword_annotations(region, {'vận tốc': 'velocity'})
    output = tmp_path / 'output.docx'
    package.save(output, [region], dict(mode=mode, font_size=size))
    translated = [run for run in Package(output).tree.findall('.//w:r', NS)
                  if 'velocity' in ''.join(run.xpath('./w:t/text()', namespaces=NS)).lower()]
    assert translated
    assert all(run.find('w:rPr/w:sz', NS).get('{' + NS['w'] + '}val') == str(int((size or 18) * 2)) for run in translated)
    assert source.read_bytes() == original


class Translator:
    calls = 0

    def configure_terms(self, *_):
        pass

    def configure_knowledge(self, *_):
        pass

    def translate(self, rows, *_):
        self.calls += 1
        return [row.replace('Vận tốc', 'Velocity') for row in rows]


def test_conversion_autosaves_suffix_and_renaming_or_formatting_reuses_translation(prepared):
    directory, source, job, _ = prepared
    original = source.read_bytes()
    translator = Translator()
    converted = jobs.convert(directory, job['id'], dict(mode='below'), 'Vật lý', '11', Event(), lambda _: None,
                             translator=translator, title='Chuyển động đều')
    assert converted['phase'] == 'completed'
    assert Path(converted['export_path']) == source.with_name('Bài học (Song ngu VI-EN).docx')
    assert Path(converted['export_path']).is_file() and converted['title'] == 'Chuyển động đều'
    assert jobs.catalog(directory)[0]['title'] == 'Chuyển động đều'
    again = jobs.convert(directory, job['id'], dict(mode='below', bold=True, font='Arial'), 'Vật lý', '11', Event(), lambda _: None,
                         translator=translator, title='Tên đã sửa')
    assert translator.calls == 1 and again['export_path'] == converted['export_path']
    assert source.read_bytes() == original
    assert jobs.restore_previous(directory, job['id'])['title'] == 'Chuyển động đều'


def test_failed_autosave_keeps_translation_and_can_retry_without_another_request(prepared, monkeypatch):
    directory, _, job, _ = prepared
    translator = Translator()
    real_export = jobs.export
    monkeypatch.setattr(jobs, 'export', lambda *_: (_ for _ in ()).throw(OSError('fixture disk unavailable')))
    failed = jobs.convert(directory, job['id'], {}, 'Vật lý', '11', Event(), lambda _: None, translator=translator)
    assert failed['phase'] == 'review' and failed['regions'][0]['en']
    assert 'chưa lưu được' in failed['message']
    monkeypatch.setattr(jobs, 'export', real_export)
    retried = jobs.convert(directory, job['id'], {}, 'Vật lý', '11', Event(), lambda _: None, translator=translator)
    assert retried['phase'] == 'completed' and translator.calls == 1


def test_stop_before_autosave_does_not_publish_a_result(prepared):
    directory, _, job, _ = prepared
    cancel = Event()
    stopped = jobs.convert(directory, job['id'], {}, 'Vật lý', '11', cancel,
                           lambda value: cancel.set() if value.get('percent') == 98 else None, translator=Translator())
    assert stopped['phase'] == 'paused' and stopped['regions'][0]['en']
    assert not (jobs.folder_for(directory, job['id']) / 'result.docx').exists()


@pytest.mark.parametrize('options', [dict(font='Arial\n'), dict(font='a' * 121), dict(bold='yes'), dict(color='invalid'),
    *[dict(font_size=value) for value in (-1, .5, 12.2, 1639, float('nan'), float('inf'), True, '26')]])
def test_invalid_typography_cannot_be_saved(options):
    with pytest.raises(ValueError):
        jobs.options_for(options)


@pytest.mark.parametrize('stop', [False, True])
def test_windows_output_lock_is_bounded_and_cancellation_preserves_previous_file(tmp_path, monkeypatch, stop):
    pending, target = tmp_path / 'candidate.docx', tmp_path / 'result.docx'
    pending.write_bytes(b'new result')
    target.write_bytes(b'previous result')
    original_replace = Path.replace
    cancel, attempts = Event(), []

    def sharing_lock(path, destination):
        attempts.append(path)
        if len(attempts) <= 2:
            if stop:
                cancel.set()
            error = PermissionError('fixture Windows sharing lock')
            error.winerror = 32
            raise error
        return original_replace(path, destination)

    monkeypatch.setattr(Path, 'replace', sharing_lock)
    if stop:
        with pytest.raises(InterruptedError):
            jobs.replace_output(pending, target, cancel)
        assert target.read_bytes() == b'previous result' and pending.is_file()
    else:
        jobs.replace_output(pending, target, cancel)
        assert len(attempts) == 3 and target.read_bytes() == b'new result'
