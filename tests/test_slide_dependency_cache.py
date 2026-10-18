import io
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from test_slide_render_cache import office_fixture, report

from app.slide_quality import render_quality


def source(path):
    deck=Presentation()
    image=io.BytesIO()
    Image.new('RGB',(30,30),'green').save(image,format='PNG')
    for text in ('First','Second'):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2)).text=text
        slide.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(1),Inches(4))
        slide.notes_slide.notes_text_frame.text='Original narration'
    deck.save(path)
    return path


def changed(original,target,name,replacement):
    with zipfile.ZipFile(original) as before,zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED) as after:
        for item in before.infolist():
            data=before.read(item)
            after.writestr(item,replacement(data) if item.filename==name else data)
    return target


def test_local_fit_renders_only_changed_slide_without_losing_native_failures(tmp_path,monkeypatch):
    path=source(tmp_path/'source.pptx')
    calls=office_fixture(monkeypatch)
    render_quality(path,tmp_path,report(path))
    new=changed(path,tmp_path/'changed.pptx','ppt/slides/slide2.xml',lambda b:b.replace(b'Second',b'Corrected'))
    result=render_quality(new,tmp_path,report(new))
    assert len(calls)==3
    assert result['render_cache']=={'reused':1,'rendered':1,'total':2}
    assert result['counts']['FAIL']==2  # Actual overflow measurements survive reuse.


def test_notes_only_change_reuses_render_but_not_content_checks(tmp_path,monkeypatch):
    path=source(tmp_path/'source.pptx')
    calls=office_fixture(monkeypatch)
    render_quality(path,tmp_path,report(path))
    new=changed(path,tmp_path/'changed.pptx','ppt/notesSlides/notesSlide1.xml',lambda b:b.replace(b'Original narration',b'New narration'))
    quality=report(new)
    quality['slides'][0]['issues'].append({'code':'notes_mismatch','severity':'CHECK'})
    result=render_quality(new,tmp_path,quality)
    assert len(calls)==2 and result['render_cache']['reused']==2
    assert 'notes_mismatch' in {i['code'] for i in result['slides'][0]['issues']}


def test_changed_theme_or_picture_invalidates_every_affected_slide(tmp_path,monkeypatch):
    for part in ('ppt/theme/theme1.xml','ppt/media/image1.png'):
        directory=tmp_path/Path(part).name
        directory.mkdir()
        path=source(directory/'source.pptx')
        calls=office_fixture(monkeypatch)
        render_quality(path,directory,report(path))
        def mutate(data):
            if part.endswith('.xml'):
                return data.replace(b'Office Theme',b'Changed Theme')
            image=io.BytesIO()
            Image.new('RGB',(30,30),'blue').save(image,format='PNG')
            return image.getvalue()
        new=changed(path,directory/'changed.pptx',part,mutate)
        result=render_quality(new,directory,report(new))
        assert len(calls)==4 and result['render_cache']['reused']==0


def test_changed_font_environment_rejects_old_geometry(tmp_path,monkeypatch):
    path=source(tmp_path/'source.pptx')
    calls=office_fixture(monkeypatch)
    render_quality(path,tmp_path,report(path))
    monkeypatch.setattr('app.slide_fingerprint.environment_key',lambda:'changed-environment')
    result=render_quality(path,tmp_path,report(path))
    assert len(calls)==4 and result['render_cache']['reused']==0


def test_reused_external_images_are_available_to_the_next_local_repair(tmp_path, monkeypatch):
    path = source(tmp_path / 'source.pptx')
    calls = office_fixture(monkeypatch)
    request, library = tmp_path / 'request', tmp_path / 'library'
    render_quality(path, request, report(path))
    imported = render_quality(path, library, report(path), reuse_directory=request)
    assert imported['render_cache']['reused'] == 2 and len(calls) == 2
    new = changed(path, tmp_path / 'changed.pptx', 'ppt/slides/slide2.xml', lambda b: b.replace(b'Second', b'Corrected'))
    repaired = render_quality(new, library, report(new))
    assert repaired['render_cache'] == {'reused': 1, 'rendered': 1, 'total': 2}
    assert len(calls) == 3


def test_static_slide_cache_survives_midnight(tmp_path, monkeypatch):
    from datetime import date

    from app import slide_fingerprint as fingerprints

    path = source(tmp_path / 'source.pptx')
    first_environment = fingerprints.environment_key()
    first_keys = fingerprints.SlideFingerprints(path).keys()
    import app.slide_render_cache as cache_module
    monkeypatch.setattr(cache_module, 'date', type('NextDay', (), {'today': staticmethod(lambda: date(2026, 10, 19))}))
    assert fingerprints.environment_key() == first_environment
    assert fingerprints.SlideFingerprints(path).keys() == first_keys


def test_non_placeholder_dynamic_field_is_never_cached(tmp_path):
    from lxml import etree

    from app import slide_fingerprint as fingerprints

    path = source(tmp_path / 'source.pptx')
    def add_date(data):
        tree = etree.fromstring(data)
        ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        paragraph = tree.find('.//{' + ns + '}p')
        field = etree.SubElement(paragraph, '{' + ns + '}fld',
                                 id='{11111111-1111-1111-1111-111111111111}', type='datetime1')
        etree.SubElement(field, '{' + ns + '}t').text = '18/10/2026'
        return etree.tostring(tree)
    dynamic = changed(path, tmp_path / 'date.pptx', 'ppt/slideMasters/slideMaster1.xml',
                      add_date)
    before = fingerprints.SlideFingerprints(dynamic).keys()
    assert before and all(value is None for value in before.values())


def test_only_visible_date_footer_renders_again_after_midnight(tmp_path, monkeypatch):
    from datetime import date
    import app.slide_render_cache as cache_module

    path = source(tmp_path / 'source.pptx')
    calls = office_fixture(monkeypatch, date_slides=(2,))
    monkeypatch.setattr(cache_module, 'date', type('DayOne', (), {'today': staticmethod(lambda: date(2026, 10, 18))}))
    render_quality(path, tmp_path, report(path))
    monkeypatch.setattr(cache_module, 'date', type('DayTwo', (), {'today': staticmethod(lambda: date(2026, 10, 19))}))
    result = render_quality(path, tmp_path, report(path))
    assert result['render_cache'] == {'reused': 1, 'rendered': 1, 'total': 2}
    assert len(calls) == 3
