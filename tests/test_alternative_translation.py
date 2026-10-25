import io
import json
import shutil
import subprocess
import urllib.error
import zipfile
from contextlib import closing
from pathlib import Path
from threading import Event

import pytest

from app import alternative_translation as api
from app.google_settings_tasks import change, read_state
from app.google_translation import KeyStore, TranslationProblem
from app.library import Library
from app.translation_provider import SELECTABLE, build_translator, require_credentials, snapshot

URL = 'https://script.google.com/macros/s/' + 'A' * 40 + '/exec'
OPTIONS = {'languageapp': URL, 'gemini': api.DEFAULT_MODEL, 'deepl': 'developer'}
SECRET = 'test_secret_do_not_use_1234567890'


@pytest.mark.parametrize('provider', SELECTABLE[1:])
def test_isolated_encrypted_configuration_and_durable_ledger(tmp_path, provider):
    with closing(Library(tmp_path)):
        pass
    KeyStore(tmp_path).add('AIza' + 'g' * 35)
    original = KeyStore(tmp_path).path.read_bytes()
    assert not change(tmp_path, 'alternative_save', (provider, SECRET, OPTIONS[provider]))['error']
    assert not change(tmp_path, 'provider', (provider,))['error']
    state = read_state(tmp_path)
    assert state['provider'] == provider and state['activeConfigured']
    assert SECRET not in json.dumps(state)
    store = api.AlternativeStore(tmp_path, provider)
    assert SECRET.encode() not in store.path.read_bytes()
    store.reserve(10, api.LIMITS[provider])
    store.clear()
    with pytest.raises(TranslationProblem):
        require_credentials(tmp_path, provider)
    store.configure(SECRET + '_new', OPTIONS[provider])
    assert store.public()[0]['used'] == 10
    assert KeyStore(tmp_path).path.read_bytes() == original
    assert KeyStore(tmp_path).public()[0]['used'] == 0


@pytest.mark.parametrize('provider', SELECTABLE[1:])
def test_routing_cache_guard_and_profile_resume(tmp_path, provider):
    with closing(Library(tmp_path)):
        pass
    store = api.AlternativeStore(tmp_path, provider)
    store.configure(SECRET, OPTIONS[provider])
    change(tmp_path, 'provider', (provider,))
    config = snapshot(tmp_path)
    translator = build_translator(tmp_path, config, 500_000, lambda *_: pytest.fail('Google invoked'))
    calls = []
    translator.transport = lambda texts, *_: calls.extend(texts) or [text.replace('Vật có', 'The object has') for text in texts]
    assert translator.translate(['Vật có 12 m/s.'], Event()) == ['The object has 12 m/s.']
    before = store.public()[0]['used']
    assert translator.translate(['Vật có 12 m/s.'], Event()) == ['The object has 12 m/s.']
    assert len(calls) == 1 and store.public()[0]['used'] == before
    assert before == (1 if provider == 'languageapp' else len(calls[0]))
    change(tmp_path, 'provider', ('google',))
    assert build_translator(tmp_path, config, 500_000, None).provider == provider
    if provider != 'deepl':
        store.configure('', URL.replace('AAAA', 'BBBB') if provider == 'languageapp' else 'gemini-3.1-flash-lite')
        with pytest.raises(TranslationProblem, match='Cấu hình'):
            build_translator(tmp_path, config, 500_000, None)


def test_daily_and_lifetime_counters_have_distinct_periods(tmp_path, monkeypatch):
    assert api.AlternativeStore(tmp_path, 'deepl').month() == 'lifetime'
    assert len(api.AlternativeStore(tmp_path, 'languageapp').month()) == 10
    store = api.AlternativeStore(tmp_path, 'deepl')
    store.configure(SECRET, 'developer')
    store.reserve(999_999, 1_000_000)
    monkeypatch.setattr(KeyStore, 'month', staticmethod(lambda: '2099-12'))
    assert store.public()[0]['used'] == 999_999
    with pytest.raises(TranslationProblem):
        store.reserve(2, 1_000_000)


@pytest.mark.parametrize('url', ['http://localhost/exec', 'https://script.google.com.evil/macros/s/' + 'a'*40 + '/exec',
    URL + '?key=secret', URL.replace('/exec', '/dev'), URL.replace('script.google.com', 'accounts.google.com')])
def test_languageapp_endpoint_rejects_arbitrary_hosts(tmp_path, url):
    with pytest.raises(TranslationProblem):
        api.AlternativeStore(tmp_path, 'languageapp').configure(SECRET, url)


class Response:
    def __init__(self, payload):
        self.raw = json.dumps(payload).encode()
    def __enter__(self): return self
    def __exit__(self, *_): pass
    def read(self, limit): return self.raw[:limit]


def opener(monkeypatch, handler):
    class Opener:
        def open(self, request, timeout):
            assert timeout == 45
            return handler(request)
    monkeypatch.setattr(api.urllib.request, 'build_opener', lambda *_: Opener())


def test_script_follows_only_google_result_redirect_without_post_or_code(monkeypatch):
    seen = []
    def handler(request):
        seen.append(request)
        if len(seen) == 1:
            raise urllib.error.HTTPError(URL, 302, '', {'Location': 'https://script.googleusercontent.com/macros/echo?x=1'}, io.BytesIO())
        assert request.method == 'GET' and request.data is None
        assert SECRET not in str(request.headers)
        return Response({'translations': ['Hello']})
    opener(monkeypatch, handler)
    assert api.languageapp_request(['Xin chào'], {'key': SECRET, 'option': URL}, 'vi', 'en') == ['Hello']
    assert json.loads(seen[0].data)['accessCode'] == SECRET


@pytest.mark.parametrize('location', ['https://evil.example/macros/x', 'https://accounts.google.com/login',
    'http://script.googleusercontent.com/macros/x', 'https://script.googleusercontent.com:443/macros/x'])
def test_script_never_forwards_credentials_to_other_redirects(monkeypatch, location):
    calls = []
    def handler(request):
        calls.append(request)
        raise urllib.error.HTTPError(URL, 302, '', {'Location': location}, io.BytesIO())
    opener(monkeypatch, handler)
    with pytest.raises(TranslationProblem):
        api.languageapp_request(['x'], {'key': SECRET, 'option': URL}, 'vi', 'en')
    assert len(calls) == 1


def gemini_result(rows, reason='STOP'):
    return {'candidates': [{'finishReason': reason, 'content': {'parts': [{'text': json.dumps({'translations': rows})}]}}]}


def test_gemini_fixed_host_header_and_strict_indexed_output(monkeypatch):
    def handler(request):
        assert request.full_url == 'https://generativelanguage.googleapis.com/v1beta/models/' + api.DEFAULT_MODEL + ':generateContent'
        assert request.headers['X-goog-api-key'] == SECRET and SECRET not in request.full_url
        payload = json.loads(request.data)
        assert payload['generationConfig']['responseMimeType'] == 'application/json'
        return Response(gemini_result([{'index': 0, 'text': 'Hello'}]))
    opener(monkeypatch, handler)
    assert api.gemini_request(['Xin chào'], {'key': SECRET, 'option': api.DEFAULT_MODEL}, 'vi', 'en') == ['Hello']


@pytest.mark.parametrize('payload', [gemini_result([{'index': 1, 'text': 'wrong order'}]),
    gemini_result([{'index': 0, 'text': 'cut'}], 'MAX_TOKENS'), gemini_result([]),
    gemini_result([{'index': 0, 'text': ''}]), {'promptFeedback': {'blockReason': 'SAFETY'}}])
def test_gemini_missing_reordered_blocked_truncated_is_not_accepted(monkeypatch, payload):
    opener(monkeypatch, lambda _: Response(payload))
    with pytest.raises(TranslationProblem):
        api.gemini_request(['x'], {'key': SECRET, 'option': api.DEFAULT_MODEL}, 'vi', 'en')


@pytest.mark.parametrize('suffix,host', [('', 'api.deepl.com'), (':fx', 'api-free.deepl.com')])
def test_deepl_batch_limit_html_auth_and_exact_count(monkeypatch, suffix, host):
    sizes = []
    def handler(request):
        assert request.full_url == 'https://' + host + '/v2/translate'
        assert request.headers['Authorization'] == 'DeepL-Auth-Key ' + SECRET + suffix
        payload = json.loads(request.data)
        assert payload['source_lang'] == 'VI' and payload['target_lang'] == 'EN'
        assert payload['tag_handling'] == 'html' and len(payload['text']) <= 50
        sizes.append(len(payload['text']))
        return Response({'translations': [{'text': 'Hello'} for _ in payload['text']]})
    opener(monkeypatch, handler)
    assert len(api.deepl_request(['x']*51, {'key': SECRET + suffix}, 'vi', 'en')) == 51
    assert sizes == [50, 1]


@pytest.mark.parametrize('status,code', [(403, 'api_key'), (429, 'api_rate'), (456, 'api_quota'), (500, 'api_network')])
def test_provider_errors_do_not_echo_secrets_or_lecture(monkeypatch, status, code):
    def handler(request):
        raise urllib.error.HTTPError(request.full_url, status, SECRET, {}, io.BytesIO(SECRET.encode()))
    opener(monkeypatch, handler)
    with pytest.raises(TranslationProblem) as caught:
        api.deepl_request(['private lesson'], {'key': SECRET}, 'vi', 'en')
    assert caught.value.code == code
    assert SECRET not in str(caught.value) and 'private lesson' not in str(caught.value)


@pytest.mark.parametrize('body,marker', [(b'<html>Script function not found: doPost</html>', 'doPost'),
    ('Lỗi Không tìm thấy hàm tập lệnh: doPost'.encode(), 'doPost'), (b'<html>Sign in to Google</html>', '')])
def test_languageapp_html_errors_are_deployment_errors_not_bad_translation(monkeypatch, body, marker):
    response = Response({})
    response.raw = body
    opener(monkeypatch, lambda _: response)
    with pytest.raises(TranslationProblem) as caught:
        api.languageapp_request(['hello'], {'key': SECRET, 'option': URL}, 'vi', 'en')
    assert caught.value.code == 'api_deployment'
    assert SECRET not in str(caught.value)
    if marker:
        assert marker in str(caught.value)


def test_languageapp_login_redirect_is_actionable_without_following_login(monkeypatch):
    def handler(request):
        raise urllib.error.HTTPError(URL, 302, '', {'Location': 'https://accounts.google.com/login'}, io.BytesIO())
    opener(monkeypatch, handler)
    with pytest.raises(TranslationProblem) as caught:
        api.languageapp_request(['hello'], {'key': SECRET, 'option': URL}, 'vi', 'en')
    assert caught.value.code == 'api_deployment' and 'Mọi người' in str(caught.value)


@pytest.mark.parametrize('status,code,marker', [(402,'api_billing','Free'), (404,'api_model','model')])
def test_gemini_depleted_prepayment_and_unavailable_model_are_distinct(monkeypatch, status, code, marker):
    def handler(request):
        raise urllib.error.HTTPError(request.full_url, status, SECRET, {}, io.BytesIO(SECRET.encode()))
    opener(monkeypatch, handler)
    with pytest.raises(TranslationProblem) as caught:
        api.gemini_request(['hello'], {'key':SECRET, 'option':api.DEFAULT_MODEL}, 'vi', 'en')
    assert caught.value.code == code and marker in str(caught.value)
    assert SECRET not in str(caught.value)


@pytest.mark.parametrize('provider', SELECTABLE[1:])
def test_connection_probe_is_fresh_counted_and_does_not_change_selection(tmp_path, monkeypatch, provider):
    with closing(Library(tmp_path)):
        pass
    store = api.AlternativeStore(tmp_path, provider)
    store.configure(SECRET, OPTIONS[provider])
    calls = []
    monkeypatch.setitem(api.TRANSPORTS, provider, lambda texts,*_: calls.append(texts) or ['Hello.'])
    api.check_connection(tmp_path, provider)
    api.check_connection(tmp_path, provider)
    assert len(calls) == 2  # A cached old success must not hide a new account/quota error.
    assert store.public()[0]['used'] == (2 if provider == 'languageapp' else 2*len('Xin chào.'))
    assert snapshot(tmp_path)['translation_provider'] == 'google'
    def fail(*_):
        raise TranslationProblem('api_billing','Test billing failure')
    monkeypatch.setitem(api.TRANSPORTS, provider, fail)
    result = change(tmp_path, 'alternative_check', (provider,))
    assert 'Test billing failure' == result['error']
    assert store.public()[0]['used'] == (2 if provider == 'languageapp' else 2*len('Xin chào.'))


@pytest.mark.parametrize('provider', SELECTABLE[1:])
def test_rejection_does_not_fallback_or_consume_budget_and_cancel_is_immediate(tmp_path, provider):
    store = api.AlternativeStore(tmp_path, provider)
    store.configure(SECRET, OPTIONS[provider])
    calls = []
    def reject(*_):
        calls.append(True)
        raise TranslationProblem('api_key', 'Rejected')
    translator = api.AlternativeTranslator(tmp_path, provider, transport=reject)
    with pytest.raises(TranslationProblem, match='Rejected'):
        translator.translate(['Xin chào'], Event())
    assert calls == [True] and store.public()[0]['used'] == 0
    cancel = Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        translator.translate(['Xin chào'], cancel)
    assert calls == [True]


def test_corrupt_inactive_service_and_failed_save_do_not_change_google(tmp_path):
    with closing(Library(tmp_path)):
        pass
    KeyStore(tmp_path).add('AIza' + 'g'*35)
    api.AlternativeStore(tmp_path, 'gemini').path.write_bytes(b'corrupt')
    assert read_state(tmp_path)['activeConfigured']
    assert read_state(tmp_path)['alternatives']['gemini']['message']
    assert not change(tmp_path, 'provider', ('gemini',))['error']
    assert not read_state(tmp_path)['activeConfigured']  # A broken config can still be selected and repaired in UI.
    change(tmp_path, 'provider', ('google',))
    assert change(tmp_path, 'alternative_save', ('languageapp', SECRET, 'wrong URL'))['error']
    assert snapshot(tmp_path)['translation_provider'] == 'google'
    # Explicit re-entry can repair a corrupt configuration.
    assert not change(tmp_path, 'alternative_save', ('gemini', SECRET, api.DEFAULT_MODEL))['error']


@pytest.mark.parametrize('provider', SELECTABLE[1:])
@pytest.mark.parametrize('style,mode', [('source', 'parallel_columns'), ('source', 'english_only'), ('template', 'parallel_columns')])
def test_real_pptx_conversion_uses_selected_adapter_without_modifying_source(tmp_path, monkeypatch, provider, style, mode):
    from pptx import Presentation
    from pptx.util import Inches

    from app import google_conversion
    from app.conversion_input import creation_config
    from app.conversion_progress import ConversionProgress

    with closing(Library(tmp_path)):
        pass
    api.AlternativeStore(tmp_path, provider).configure(SECRET, OPTIONS[provider])
    change(tmp_path, 'provider', (provider,))
    calls = []
    def fixture(self, texts, credential, source, target):
        assert self.provider == provider and credential['key'] == SECRET
        calls.extend(texts)
        return [text.replace('Vật có ', 'The object has a ') for text in texts]
    monkeypatch.setattr(api.AlternativeTranslator, '_send', fixture)
    monkeypatch.setattr(google_conversion, 'Translator', lambda *_: pytest.fail('Unselected Google invoked'))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = 'Vật có vận tốc 12 m/s.'
    source = tmp_path / 'Thí nghiệm gốc.pptx'
    deck.save(source)
    original = source.read_bytes()
    config = creation_config('Cơ năng', 'Vật lý', 'THPT', '11', mode, 'standard', style)
    job = google_conversion.prepare(tmp_path, config, source)
    change(tmp_path, 'provider', ('google',))
    progress = ConversionProgress(job['folder'], lambda *_: None)
    lesson, preview = google_conversion.convert(tmp_path, job, Event(), progress)
    assert calls and source.read_bytes() == original
    assert lesson['external_deck']['translation_provider'] == provider
    text = '\n'.join(shape.text for page in Presentation(preview['path']).slides for shape in page.shapes if shape.has_text_frame)
    assert 'The object has a velocity' in text and '12' in text
    assert ('Vật có vận tốc' in text) == (mode == 'parallel_columns')
    used = api.AlternativeStore(tmp_path, provider).public()[0]['used']
    again, _ = google_conversion.convert(tmp_path, job, Event(), progress)
    assert again['id'] == lesson['id']
    assert api.AlternativeStore(tmp_path, provider).public()[0]['used'] == used


def test_backup_excludes_every_new_provider_credential_and_cache(tmp_path):
    from app.storage import backup_library

    with closing(Library(tmp_path)):
        pass
    for provider in SELECTABLE[1:]:
        store = api.AlternativeStore(tmp_path, provider)
        store.configure(SECRET, OPTIONS[provider])
        store.cache('id', 'private lecture cache')
    output = backup_library(tmp_path)
    with zipfile.ZipFile(output) as archive:
        assert not any('translation_' in name for name in archive.namelist())


def test_corrupt_inactive_usage_database_does_not_block_other_engines(tmp_path):
    with closing(Library(tmp_path)):
        pass
    KeyStore(tmp_path).add('AIza' + 'g'*35)
    folder = tmp_path / 'translation_deepl'
    folder.mkdir()
    (folder / 'usage.db').write_bytes(b'not a SQLite database')
    state = read_state(tmp_path)
    assert state['provider'] == 'google' and state['activeConfigured']
    assert state['alternatives']['deepl']['message']
    assert not change(tmp_path, 'provider', ('deepl',))['error']
    assert not read_state(tmp_path)['activeConfigured']


def test_apps_script_protocol_and_shared_rolling_quota():
    node = shutil.which('node')
    if not node:
        pytest.skip('Node required to execute deployed Apps Script protocol fixture')
    script = (Path(api.__file__).parent / 'resources/languageapp.gs').read_text(encoding='utf-8')
    harness = r'''
const assert = require('node:assert/strict');
let now = 1700000000000, usedCalls = 0;
const RealDate = Date;
Date = class extends RealDate { static now() { return now; } };
const values = {BILICLASS_ACCESS_CODE: 'test_secret_1234567890'};
const PropertiesService = {getScriptProperties: () => ({getProperty: k => values[k], setProperty: (k,v) => {values[k] = v;}})};
const LockService = {getScriptLock: () => ({tryLock: () => true, releaseLock: () => {}})};
const ContentService = {MimeType: {JSON: 'json'}, createTextOutput: text => ({setMimeType: () => JSON.parse(text)})};
const LanguageApp = {translate: (text, source, target, options) => {usedCalls++; assert.equal(options.contentType, 'html'); return text.replace('Xin chào', 'Hello');}};
function post(body) {return doPost({postData: {contents: JSON.stringify(body)}});}
const body = {accessCode: values.BILICLASS_ACCESS_CODE, source: 'vi', target: 'en', texts: ['Xin chào']};
assert.deepEqual(post({...body, accessCode: 'wrong'}), {error: 'auth'});
assert.equal(usedCalls, 0);
assert.deepEqual(post(body), {translations: ['Hello']});
for(let i=1; i<4500; i++) assert.ok(post(body).translations);
assert.deepEqual(post(body), {error: 'quota'});
assert.equal(usedCalls, 4500);
assert.ok(values.BILICLASS_USAGE.length < 9000);
now += 26*3600000;
assert.deepEqual(post(body), {translations: ['Hello']});
assert.deepEqual(post({...body, texts: Array(26).fill('x')}), {error: 'request'});
console.log('passed');
'''
    result = subprocess.run([node], input=script + '\n' + harness, text=True, capture_output=True, timeout=20)
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == 'passed'
