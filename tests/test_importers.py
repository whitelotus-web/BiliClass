import hashlib

import pytest

from app.importers import parse_document


def test_txt_unicode_and_symbols(tmp_path):
    path = tmp_path / "Bài giảng.txt"
    path.write_text("Hãy so sánh x² = 4.\n\nGiữ 25% và NaCl.", encoding="utf-8-sig")
    blocks = parse_document(path)
    assert blocks == [("Đoạn 1", "Hãy so sánh x² = 4."), ("Đoạn 2", "Giữ 25% và NaCl.")]


def test_docx_preserves_paragraph_table_order_and_source(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Mở đầu")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Tiếng Việt"
    table.cell(0, 1).text = "English"
    document.add_paragraph("Kết thúc")
    path = tmp_path / "Bài.docx"
    document.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    assert [text for _, text in parse_document(path)] == ["Mở đầu", "Tiếng Việt | English", "Kết thúc"]
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_pptx_slide_locator_and_source(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    for text in ("Phần 1: A = 2", "Phần 2: B = 3"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = text
    path = tmp_path / "Bài.pptx"
    presentation.save(path)
    before = path.read_bytes()
    assert parse_document(path) == [("Slide 1", "Phần 1: A = 2"), ("Slide 2", "Phần 2: B = 3")]
    assert path.read_bytes() == before


def test_pdf_text_and_blank_scan_message(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=500, height=500)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 40 400 Td (Compare two ideas.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "Bai.pdf"
    writer.write(path)
    assert parse_document(path) == [("Trang 1", "Compare two ideas.")]
    blank = PdfWriter()
    blank.add_blank_page(width=500, height=500)
    blank.write(path)
    with pytest.raises(ValueError, match="OCR"):
        parse_document(path)


def test_unsupported_format_is_explicit(tmp_path):
    path = tmp_path / "scan.png"
    path.write_bytes(b"not a supported document")
    with pytest.raises(ValueError, match="OCR"):
        parse_document(path)
