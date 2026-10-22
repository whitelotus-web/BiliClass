import copy
import hashlib
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.chatgpt_handoff import inspect_returned_deck
from app.library import Library, RevisionConflict
from app.slide_refinement import improved, source_media_preserved
from app.slide_repair import _compact_fixed_spacing, normalize_titles, repair_layout


def test_fitting_font_also_releases_excessive_fixed_line_spacing():
    from lxml import etree

    from app.slide_repair import NS

    body = etree.fromstring(('''<p:txBody xmlns:p="''' + NS['p'] + '" xmlns:a="' + NS['a'] + '''">
      <a:p><a:pPr><a:lnSpc><a:spcPts val="5755"/></a:lnSpc><a:spcAft><a:spcPts val="3000"/></a:spcAft></a:pPr>
      <a:r><a:rPr sz="2800"/><a:t>Original text</a:t></a:r></a:p>
      <a:p><a:pPr><a:lnSpc><a:spcPts val="2800"/></a:lnSpc></a:pPr><a:r><a:rPr sz="2800"/><a:t>Second line</a:t></a:r></a:p>
    </p:txBody>''').encode())
    _compact_fixed_spacing(body)
    assert body.xpath('.//a:lnSpc/a:spcPct/@val', namespaces=NS) == ['110000', '100000']
    assert body.xpath('.//a:lnSpc/a:spcPts/@val', namespaces=NS) == []
    assert body.xpath('.//a:spcAft/a:spcPts/@val', namespaces=NS) == ['420']
    assert body.xpath('.//a:t/text()', namespaces=NS) == ['Original text', 'Second line']


def test_separate_stacked_title_boxes_use_same_font_without_moving_content(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    vi = slide.shapes.add_textbox(Inches(.4), Inches(.1), Inches(8), Inches(.5))
    vi.text = 'Năng lượng cơ học'
    vi.text_frame.paragraphs[0].font.name = 'Arial'
    vi.text_frame.paragraphs[0].font.size = Pt(32)
    en = slide.shapes.add_textbox(Inches(.4), Inches(.65), Inches(8), Inches(.5))
    en.text = 'Mechanical energy'
    en.text_frame.paragraphs[0].runs[0].font.name = 'Courier New'
    en.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    original_box = en.left, en.top, en.width, en.height
    path = tmp_path / 'headings.pptx'
    deck.save(path)
    styled = normalize_titles(path, tmp_path, 'parallel_columns')
    result = Presentation(styled['path']).slides[0].shapes[1]
    run = result.text_frame.paragraphs[0].runs[0]
    assert run.font.name == 'Arial' and run.font.italic and run.font.size.pt < 28
    assert result.text == en.text
    assert (result.left, result.top, result.width, result.height) == original_box


def presentation(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(.4), Inches(6), Inches(1))
    shape.text = "Năng lượng cơ học\nMechanical energy"
    shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    shape.text_frame.paragraphs[1].runs[0].font.name = "Courier New"
    shape.text_frame.paragraphs[1].runs[0].font.size = Pt(24)
    slide.notes_slide.notes_text_frame.text = "VI: Năng lượng cơ học.\nEN: Mechanical energy."
    image = path.with_suffix(".png")
    Image.new("RGB", (40, 40), "blue").save(image)
    slide.shapes.add_picture(str(image), Inches(1), Inches(2))
    deck.save(path)
    return shape.shape_id


def test_title_style_preserves_words_media_notes_and_unrelated_parts(tmp_path):
    source = tmp_path / "original.pptx"
    shape_id = presentation(source)
    before = source.read_bytes()
    result = normalize_titles(source, tmp_path, "parallel_columns")
    assert result["changes"] and source.read_bytes() == before
    fixed = Presentation(result["path"])
    shape = next(s for s in fixed.slides[0].shapes if s.shape_id == shape_id)
    vi, en = shape.text_frame.paragraphs
    assert shape.text == "Năng lượng cơ học\nMechanical energy"
    assert en.runs[0].font.name == vi.runs[0].font.name == "Arial"
    assert en.runs[0].font.italic and en.runs[0].font.size < vi.runs[0].font.size
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(result["path"]) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "ppt/slides/slide1.xml":
                assert a.read(name) == b.read(name), name
    assert source_media_preserved(source, result["path"], "source")


def test_layout_rejects_stale_report_and_never_touches_original(tmp_path):
    source = tmp_path / "original.pptx"
    presentation(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="thay đổi"):
        repair_layout(source, tmp_path, {"sha256": "f" * 64, "slides": []})
    assert source.read_bytes() == before


@pytest.mark.parametrize('blocked,generated', [(False, True), (True, True), (False, False)])
def test_formula_frame_grows_only_into_free_space_and_preserves_math(tmp_path, monkeypatch, blocked, generated):
    import hashlib

    from lxml import etree

    from app.slide_repair import NS

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    formula = group.shapes.add_textbox(Inches(4), Inches(1), Inches(4), Inches(.5))
    formula.name = 'BCL_BODY English' if generated else 'Original formula'
    formula.text = 'Speed = '
    formula.text_frame.paragraphs[0].font.size = Pt(20)
    math_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    math = etree.SubElement(formula.text_frame.paragraphs[0]._p, '{' + math_ns + '}oMath')
    fraction = etree.SubElement(math, '{' + math_ns + '}f')
    for tag, text in [('num', 'v'), ('den', 'ω')]:
        run = etree.SubElement(etree.SubElement(fraction, '{' + math_ns + '}' + tag), '{' + math_ns + '}r')
        etree.SubElement(run, '{' + math_ns + '}t').text = text
    original_math = etree.tostring(math, method='c14n')
    if blocked:
        obstacle = slide.shapes.add_textbox(Inches(4), Inches(1.6), Inches(4), Inches(1))
        obstacle.text = 'Do not cover this caption'
    source = tmp_path / 'formula.pptx'
    deck.save(source)
    original = source.read_bytes()
    monkeypatch.setattr('app.office_tasks.fit_text', lambda *_: pytest.fail('Never serialize math through COM text fitting'))
    report = {'sha256': hashlib.sha256(original).hexdigest(), 'slides': [
        {'slide': 1, 'issues': [{'code': 'text_overflow', 'shape_id': formula.shape_id}]}]}
    result = repair_layout(source, tmp_path, report)
    assert source.read_bytes() == original
    if blocked or not generated:
        assert result['path'] == str(source) and not result['changes']
        return
    assert result['changes'][0]['operation'] == 'grow_formula_frame'
    repaired = Presentation(result['path']).slides[0].shapes[0].shapes[0]
    assert repaired.height > formula.height
    assert (repaired.left, repaired.top, repaired.width) == (formula.left, formula.top, formula.width)
    assert etree.tostring(repaired._element.find('.//{' + math_ns + '}oMath'), method='c14n') == original_math
    assert repaired._element.xpath('.//a:rPr/@sz') == formula._element.xpath('.//a:rPr/@sz')
    assert repaired._element.find('p:txBody/a:bodyPr', NS).get('anchor') == 't'


def test_repair_handles_source_and_office_using_different_slide_part_names(tmp_path, monkeypatch):
    import hashlib

    from lxml import etree

    from app.slide_repair import NS

    source = tmp_path / 'original.pptx'
    shape_id = presentation(source)

    def rename_slide(before, after, old, new, fit=False):
        with zipfile.ZipFile(before) as a, zipfile.ZipFile(after, 'w') as b:
            for item in a.infolist():
                data = a.read(item)
                name = item.filename.replace(old, new)
                if name.endswith(('.xml', '.rels')):
                    data = data.replace(old.encode(), new.encode())
                if fit and name == 'ppt/slides/slide1.xml':
                    root = etree.fromstring(data)
                    for props in root.xpath('.//a:rPr', namespaces=NS):
                        props.set('sz', '1400')
                    data = etree.tostring(root)
                b.writestr(name, data)

    reordered = tmp_path / 'reordered.pptx'
    rename_slide(source, reordered, 'slide1.xml', 'slide37.xml')
    original = reordered.read_bytes()
    report = {'sha256': hashlib.sha256(original).hexdigest(), 'slides': [
        {'slide': 1, 'issues': [{'code': 'text_overflow', 'shape_id': shape_id}]}]}
    monkeypatch.setattr('app.office_tasks.fit_text',
                        lambda before, after, *_: rename_slide(before, after, 'slide37.xml', 'slide1.xml', True))
    repaired = repair_layout(reordered, tmp_path, report)
    assert repaired['changes'] and reordered.read_bytes() == original
    with zipfile.ZipFile(repaired['path']) as archive:
        assert 'ppt/slides/slide37.xml' in archive.namelist()
        root = etree.fromstring(archive.read('ppt/slides/slide37.xml'))
        assert set(root.xpath('.//a:rPr/@sz', namespaces=NS)) == {'1400'}
    # Styling also follows the presentation relationship, not assumed filenames.
    styled = normalize_titles(reordered, tmp_path, 'parallel_columns')
    assert styled['changes']


def test_regression_cannot_hide_an_incorrect_quiz_answer():
    def q(code): return {"slides": [{"slide": 1, "issues": [{"severity": "FAIL", "code": code}]}]}
    assert not improved(q("text_overflow"), q("quiz_mismatch"))


def test_partial_table_improvement_keeps_distinct_cell_failures():
    from app.slide_repair import failure_keys

    first = {'slide': 1, 'issues': [{'code': 'text_overflow', 'severity': 'FAIL', 'shape_id': 3, 'cell': [n, 0]}
                                  for n in (0, 1)]}
    before = {'slides': [first]}
    after = {'slides': [{**first, 'issues': first['issues'][1:]}]}
    assert len(failure_keys(before)) == 2 and failure_keys(after) < failure_keys(before)


@pytest.mark.parametrize('rewrite_words', [False, True])
def test_cell_fit_retains_grid_other_cells_and_math_and_rejects_changed_words(tmp_path, monkeypatch, rewrite_words):
    from lxml import etree

    from app.slide_repair import NS
    from app.source_deck import source_slide_parts

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    table_shape.table.cell(0, 0).text = 'A triangle has three vertices.'
    table_shape.table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    table_shape.table.cell(1, 1).text = 'Unrelated cell'
    math = etree.fromstring(b'<a14:m xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>x+1</m:t></m:r></m:oMath></a14:m>')
    table_shape.table.cell(0, 1).text_frame.paragraphs[0]._p.append(math)
    source = tmp_path / 'table.pptx'
    deck.save(source)
    original = source.read_bytes()
    part = source_slide_parts(source)[0]
    identity = table_shape.shape_id

    def fitted(before, after, targets, *_):
        assert targets == {1: {(identity, 0, 0)}}
        with zipfile.ZipFile(before) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        tree = etree.fromstring(entries[part])
        cell = tree.find('.//a:tbl/a:tr/a:tc', NS)
        cell.find('.//a:rPr', NS).set('sz', '1400')
        if rewrite_words:
            cell.find('.//a:t', NS).text = 'Incorrect rewritten content'
        # An Office rewrite must not leak any edits to an unrelated cell.
        tree.findall('.//a:tbl/a:tr', NS)[1].findall('a:tc', NS)[1].find('.//a:t', NS).text = 'Changed by Office'
        frame = tree.find('p:cSld/p:spTree/p:graphicFrame', NS)
        alternate = etree.Element('{' + NS['mc'] + '}AlternateContent')
        frame.getparent().replace(frame, alternate)
        etree.SubElement(alternate, '{' + NS['mc'] + '}Choice', Requires='a14').append(frame)
        etree.SubElement(alternate, '{' + NS['mc'] + '}Fallback').append(copy.deepcopy(frame))
        entries[part] = etree.tostring(tree)
        with zipfile.ZipFile(after, 'w') as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

    monkeypatch.setattr('app.office_tasks.fit_text', fitted)
    report = {'sha256': hashlib.sha256(original).hexdigest(), 'slides': [{'slide': 1, 'issues': [
        {'code': 'text_overflow', 'severity': 'FAIL', 'shape_id': identity, 'cell': [0, 0]}]}]}
    result = repair_layout(source, tmp_path, report)
    assert source.read_bytes() == original
    if rewrite_words:
        assert not result['changes'] and Path(result['path']) == source
    else:
        repaired = Presentation(result['path']).slides[0].shapes[0]
        assert repaired.table.cell(0, 0).text == table_shape.table.cell(0, 0).text
        assert repaired.table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size.pt == 14
        assert repaired.table.cell(1, 1).text == 'Unrelated cell'
        assert etree.tostring(repaired.table.cell(0, 1)._tc) == etree.tostring(table_shape.table.cell(0, 1)._tc)
        assert (repaired.left, repaired.top, repaired.width, repaired.height) == (table_shape.left, table_shape.top, table_shape.width, table_shape.height)
        assert result['changes'][0]['cell'] == [0, 0]


def test_reducing_overflows_cannot_hide_new_missing_body_translation():
    before = {"slides": [{"slide": 1, "issues": [
        {"severity": "FAIL", "code": "text_overflow", "shape_id": n} for n in (1, 2, 3)]}]}
    after = {"slides": [{"slide": 1, "issues": [{"severity": "FAIL", "code": "english_body_missing"}]}]}
    assert not improved(before, after)


@pytest.mark.parametrize('translated,blocked', [(False,False),(True,False),(True,True)])
def test_only_declared_translation_can_grow_an_original_named_frame(tmp_path, translated, blocked):
    from app.text_frame_space import grow

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(.35))
    shape.text = 'Ví dụ thực tế:\nReal-world example:'
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(18)
    before = shape.left, shape.top, shape.width, shape.height
    if blocked:
        obstacle = slide.shapes.add_textbox(Inches(.2), Inches(.2), Inches(8), Inches(4))
        obstacle.text = 'Keep this figure caption clear'
    changed = grow(shape._element, slide, deck.slide_width, deck.slide_height, translated=translated)
    assert changed == (translated and not blocked)
    assert shape.text == 'Ví dụ thực tế:\nReal-world example:'
    if changed:
        assert shape.height > before[3] and shape.width >= before[2]
    else:
        assert (shape.left, shape.top, shape.width, shape.height) == before


def test_independent_page_fixes_are_kept_without_accepting_regression(tmp_path):
    from app.slide_repair import retain_improved_pages
    from app.source_deck import source_slide_parts

    original = tmp_path/'original.pptx'
    deck=Presentation()
    for text in ['Original A', 'Original B']:
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text=text
    deck.save(original)
    for slide in deck.slides:
        slide.shapes[0].text='Candidate'
    candidate=tmp_path/'candidate.pptx'
    deck.save(candidate)
    def report(first, second):
        return {'slides':[{'slide':n,'issues':[{'code':code,'severity':'FAIL','shape_id':1} for code in codes]}
                          for n,codes in enumerate([first,second],1)]}
    output,pages=retain_improved_pages(original,candidate,report(['text_overflow'],['text_overflow']),
                                      report([],['english_body_missing']),tmp_path)
    assert pages=={1}
    assert [s.shapes[0].text for s in Presentation(output).slides]==['Candidate','Original B']
    with zipfile.ZipFile(original) as a,zipfile.ZipFile(output) as b:
        assert a.read(source_slide_parts(original)[1])==b.read(source_slide_parts(output)[1])


def test_saved_correction_keeps_id_history_and_requires_new_teacher_confirmation(tmp_path):
    path = tmp_path / "original.pptx"
    presentation(path)
    library = Library(tmp_path / "library")
    try:
        inspection = inspect_returned_deck(path, conversion_format="parallel_columns")
        source = library.store_source(path)
        config = {"title": "Test", "subject": "Vật lý", "education_level": "THPT", "grade": "10",
                  "conversion_format": "parallel_columns", "style": "source", "preset": "standard"}
        lesson = library.create_external_lesson(config, source, inspection)
        result = normalize_titles(path, tmp_path, "parallel_columns")
        corrected = inspect_returned_deck(result["path"], conversion_format="parallel_columns")
        corrected['layout_plan'] = [{'source_slide': n, 'role': 'bilingual'} for n in range(1, corrected['total']+1)]
        corrected['content_review'] = {'coverage': 'unverified', 'image_pages': [1]}
        new_source = library.store_source(Path(result["path"]))
        changed = library.replace_external_presentation(lesson["id"], lesson["revision"], new_source, corrected)
        assert changed["id"] == lesson["id"] and changed["revision"] == lesson["revision"] + 1
        assert changed["segments"][0]["id"] == lesson["segments"][0]["id"]
        assert not any(s["approved"] for s in changed["segments"])
        assert changed["external_deck"]["reviewed_revision"] == 0
        assert changed['external_deck']['layout_plan'] == corrected['layout_plan']
        assert changed['external_deck']['content_review'] == corrected['content_review']
        corrected['content_review']['image_pages'].append(99)
        assert changed['external_deck']['content_review']['image_pages'] == [1]
        assert (library.directory / "sources" / source["file"]).read_bytes() == path.read_bytes()
        with pytest.raises(RevisionConflict):
            library.replace_external_presentation(lesson["id"], lesson["revision"], new_source, corrected)
    finally:
        library.close()
