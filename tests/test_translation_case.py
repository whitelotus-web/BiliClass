from threading import Event

import pytest

from app.translation_case import align_case


@pytest.mark.parametrize('source,translated,expected', [
    ('Hàm số lượng giác', 'trigonometric function', 'Trigonometric function'),
    ('Phương trình lượng giác cơ bản', 'basic trigonometric equation', 'Basic trigonometric equation'),
    ('ĐỒ THỊ CỦA HÀM SỐ ⟦BLC0⟧', 'Graph of the function ⟦BLC0⟧', 'GRAPH OF THE FUNCTION ⟦BLC0⟧'),
    ('BÀI ⟦BLC12⟧:', 'Lesson ⟦BLC12⟧:', 'LESSON ⟦BLC12⟧:'),
    ('• Hàm số', '• function', '• Function'),
    ('A. Hàm số', 'A. function', 'A. Function'),
    ('⟦BLC0⟧ là hàm số', '⟦BLC0⟧ is a function', '⟦BLC0⟧ is a function'),
    ('hàm số', 'trigonometric function', 'trigonometric function'),
    ('Định luật Newton và DNA', "Newton's law and DNA", "Newton's law and DNA"),
    ('Chỉ số pH', 'pH value', 'pH value'),
    ('Điện thoại', 'iPhone', 'iPhone'),
    ('Công thức $sin(x)$', '$sin(x)$ is a formula', '$sin(x)$ is a formula'),
    ('https://example.com Chú ý', 'https://example.com note', 'https://example.com Note'),
])
def test_heading_case_preserves_math_names_and_other_letters(source, translated, expected):
    assert align_case(source, translated) == expected
    assert align_case(source, expected) == expected, 'Reading a cached result twice is idempotent'


def test_uppercase_chemistry_heading_does_not_change_element_symbols():
    assert align_case('PHẢN ỨNG CỦA HCl VÀ NaOH', 'Reaction of HCl and NaOH', chemistry=True) == 'REACTION OF HCl AND NaOH'


def test_local_and_cached_translations_follow_the_source_case_without_network(tmp_path):
    import hashlib

    from app.google_translation import Translator

    translator = Translator(tmp_path, transport=lambda *_: pytest.fail('A case correction must reuse local/cache results'))
    translator.configure_terms('Toán', [])
    source = 'Phương trình lượng giác cơ bản'
    cache = hashlib.sha256(('google-nmt-context-v3|' + translator.term_policy.fingerprint(translator.term_policy.matches(source))
        + '|vi|en|' + source).encode()).hexdigest()
    # The phrase contains the protected default term "phương trình".
    translator.store.cache(cache, 'basic trigonometric equation')
    assert translator.translate([source, 'HÀM SỐ', 'Hàm số'], Event()) == [
        'Basic trigonometric equation', 'FUNCTION', 'Function']


def test_existing_verified_pairs_keep_math_and_styles_and_refuse_stale_edits():
    import copy

    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    from app.same_slide import NS, _shape, fingerprint, visible_pairs
    from app.translation_case import align_pairs

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    shape.text = 'HÀM SỐ\ntrigonometric function '
    en = shape.text_frame.paragraphs[1]._p
    equation = etree.fromstring(b'<a14:m xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>sin(x)+a</m:t></m:r></m:oMath></a14:m>')
    en.append(equation)
    vi = shape.text_frame.paragraphs[0]._p
    pair = dict(vi_shape=shape.shape_id, en_shape=shape.shape_id, body=0, vi_paragraph=0, en_paragraph=1,
                vi_hash=fingerprint(vi), en_hash=fingerprint(en), neutral=False)
    tree = etree.fromstring(etree.tostring(slide._element))
    native_shape = _shape(tree, shape.shape_id)
    en = native_shape.findall('p:txBody/a:p', NS)[1]
    equation = en.find('{http://schemas.microsoft.com/office/drawing/2010/main}m')
    math_before = etree.tostring(equation)
    geometry = etree.tostring(native_shape.find('p:spPr', NS))
    stale = copy.deepcopy(pair)
    assert align_pairs(tree, [pair]) == [{'shape_id': shape.shape_id, 'operation': 'align_translation_case'}]
    assert ''.join(en.xpath('./a:r/a:t/text()', namespaces=NS)) == 'TRIGONOMETRIC FUNCTION '
    assert etree.tostring(equation) == math_before
    assert etree.tostring(native_shape.find('p:spPr', NS)) == geometry
    assert visible_pairs(tree, [pair])
    assert not align_pairs(tree, [pair]), 'A second explicit repair is a no-op'
    en.find('a:r/a:t', NS).text = 'teacher edited this text'
    assert not align_pairs(tree, [stale])
    assert ''.join(en.xpath('./a:r/a:t/text()', namespaces=NS)) == 'teacher edited this text'


def test_native_assembly_normalizes_old_cached_case_in_both_slide_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from app.native_bilingual import assemble, extract

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = 'Hàm số lượng giác'
    source, target = tmp_path/'source.pptx', tmp_path/'bilingual.pptx'
    deck.save(source)
    original = source.read_bytes()
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': 'trigonometric function'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    final = Presentation(target)
    assert any('Trigonometric function' in s.text for s in final.slides[0].shapes if s.has_text_frame)
    assert 'EN: Trigonometric function' in final.slides[0].notes_slide.notes_text_frame.text
    assert source.read_bytes() == original
    assert all(row['en'] == 'trigonometric function' for row in values.values()), 'Input cache is not rewritten'
