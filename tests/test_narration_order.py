import copy
import hashlib
from contextlib import closing

from pptx import Presentation
from pptx.util import Inches

from app.chatgpt_handoff import inspect_returned_deck
from app.library import Library
from app.narration_order import visual_order
from app.slide_speech import inspect_slide
from app.source_deck import text_blocks
from app.voice_import import import_presentation, prepare_library_lesson, save_script


def box(slide, text, x, y, w, h=.6, *, name=''):
    shape = slide.shapes.add_textbox(*(Inches(value) for value in (x, y, w, h)))
    shape.text = text
    if name:
        shape.name = name
    return shape


def fixture(path):
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    box(first, 'Hai đường thẳng song song', 2, .8, 6)
    box(first, 'Bài 11:', .4, .95, 1.4)
    first.shapes.add_shape(1, Inches(0), Inches(.78), Inches(9), Inches(.1))
    second = deck.slides.add_slide(deck.slide_layouts[6])
    box(second, 'A-B-C-D', 2.4, .8, 6, 2)
    box(second, 'Nội dung:', .4, 1.15, 1.8)
    deck.save(path)
    return deck


def test_voice_reads_prefix_and_contents_before_value_without_changing_export_locators(tmp_path):
    path = tmp_path/'Hai duong thang.pptx'
    deck = fixture(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    assert [u['text'] for u in text_blocks(deck)][:2] == ['Hai đường thẳng song song', 'Bài 11:']
    inspection = inspect_returned_deck(path, narration_only=True)
    assert inspection['profile']['units'][0]['vi'] == 'Bài 11:\nHai đường thẳng song song'
    assert inspection['profile']['units'][1]['vi'] == 'Nội dung:\nA-B-C-D'
    lesson, _ = import_presentation(tmp_path/'library', path)
    assert lesson['segments'][0]['vi'].startswith('Bài 11:')
    assert lesson['segments'][1]['vi'].startswith('Nội dung:')
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_visual_order_handles_group_offsets_and_centered_contents_list():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box(group, 'Hai đường thẳng song song', 2, .8, 5)
    box(group, 'Bài 11:', .3, .95, 1.4)
    assert [u['text'] for u in text_blocks(deck, narration_order=True)] == ['Bài 11:', 'Hai đường thẳng song song']
    items = [dict(text='Tên bài', x=500, y=80, w=700, h=100),
             *[dict(text=f'Mục {i}', x=700, y=240+i*150, w=1000, h=100) for i in range(4)],
             dict(text='Nội dung', x=100, y=510, w=300, h=80)]
    # Use native EMUs, as in a real slide with the contents label centered
    # vertically beside several separate list entries.
    for row in items:
        row.update({key:row[key]*12700 for key in ('x','y','w','h')})
    assert [row['text'] for row in visual_order(items)] == ['Tên bài', 'Nội dung', 'Mục 0', 'Mục 1', 'Mục 2', 'Mục 3']


def test_native_english_prefix_precedes_title_and_paragraphs_stay_in_order():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box(slide, 'Parallel lines', 2, .8, 5, name='BCL_EN_title')
    box(slide, 'Lesson 11:', .2, .95, 1.5, name='BCL_EN_label')
    box(slide, 'The first paragraph.\nThe second paragraph.', 1, 2, 7, name='BCL_EN_body')
    assert [s['english_text'] for s in inspect_slide(slide, 1)['englishSpeechSegments']] == [
        'Lesson 11:', 'Parallel lines', 'The first paragraph.', 'The second paragraph.']


def test_old_import_is_upgraded_once_but_teacher_script_and_wavs_are_preserved(tmp_path):
    path = tmp_path/'Old.pptx'
    fixture(path)
    directory = tmp_path/'library'
    lesson, _ = import_presentation(directory, path)
    with closing(Library(directory)) as library:
        lesson = save_script(library, lesson['id'], lesson['segments'][1]['id'], lesson['revision'],
                             'Lời đọc riêng của thầy cô.', '')
        lesson['segments'][0]['vi'] = 'Hai đường thẳng song song\nBài 11:'
        lesson['source_slide_speech']['extraction_version'] = 2
        library._write(lesson)
        untouched = copy.deepcopy(lesson['narration_script'])
        audio = directory/'audio/teacher-recording.wav'
        audio.parent.mkdir(exist_ok=True)
        audio.write_bytes(b'preserve unrelated cached recording')
    upgraded, _ = prepare_library_lesson(directory, lesson['id'])
    assert upgraded['segments'][0]['vi'] == 'Bài 11:\nHai đường thẳng song song'
    assert upgraded['segments'][1]['vi'] == 'Lời đọc riêng của thầy cô.'
    assert upgraded['narration_script'] == untouched
    assert upgraded['revision'] == lesson['revision'] + 1
    again, _ = prepare_library_lesson(directory, lesson['id'])
    assert again == upgraded
    assert audio.read_bytes() == b'preserve unrelated cached recording'
