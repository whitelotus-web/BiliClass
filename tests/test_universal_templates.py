import io
import json
import zipfile
from pathlib import Path
from threading import Event

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.google_conversion import load, prepare
from app.universal_templates import asset, build, choose_layout, manifest, source_blocks


def config(preset='standard', mode='parallel_columns'):
    return dict(title='Năng lượng', english_title='Energy', preset=preset,
                conversion_format=mode, style='template')


def text(deck):
    return '\n'.join(s.text for page in deck.slides for s in page.shapes if s.has_text_frame)


def examples():
    return [('Slide 1','Cơ năng','Mechanical energy'),
            ('Slide 2 · ý 1','Mục tiêu','Objectives'),
            ('Slide 2 · ý 2','Nhận biết năng lượng.','Identify energy.'),
            ('Slide 2 · ý 3','Nêu ví dụ.','Give examples.'),
            ('Slide 3 · ý 1','Công thức','Formula'),
            ('Slide 3 · ý 2','E = W + K','E = W + K'),
            ('Slide 3 · ý 3','Cơ năng bằng tổng động năng và thế năng.',
             'Mechanical energy is the sum of kinetic and potential energy.'),
            ('Slide 4 · ý 1','Câu hỏi','Question'),
            ('Slide 4 · ý 2','Đơn vị năng lượng là gì?','What is the unit of energy?'),
            ('Slide 4 · ý 3','A. Joule','A. Joule'),
            ('Slide 4 · ý 4','B. Mét','B. Metre'),
            ('Slide 4 · ý 5','C. Giây','C. Second'),
            ('Slide 4 · ý 6','D. Kilôgam','D. Kilogram')]


@pytest.mark.parametrize('preset',['standard','visual','practice'])
@pytest.mark.parametrize('mode',['parallel_columns','sentence_pairs','english_only'])
def test_real_pack_populates_without_tokens_or_unrequested_pages(tmp_path,preset,mode):
    target=tmp_path/'result.pptx'
    info=build(target,asset(preset),config(preset,mode),examples(),Event())
    deck=Presentation(target)
    combined=text(deck)
    assert '{{' not in combined and 'SOURCE_CONTENT' not in combined
    assert info['pages']==4 and len(deck.slides)==4
    assert {p['layout'] for p in info['layout_plan']}=={'TITLE_ONLY','OBJECTIVES','FORMULA','QUIZ'}
    assert combined.count('E = W + K')==1
    assert 'What is the unit of energy?' in combined
    assert ('Đơn vị năng lượng là gì?' in combined)==(mode!='english_only')
    if mode=='english_only':
        assert not any(c in combined for c in 'ăâđêôơưáàảãạắằẳẵặấầẩẫậ')
    for page in deck.slides:
        notes=page.notes_slide.notes_text_frame.text
        assert 'VI:' in notes and 'EN:' in notes and '{{' not in notes
    ids=[id(p.notes_slide.part) for p in deck.slides]
    assert len(ids)==len(set(ids))


def test_snapshot_is_self_contained_for_other_machine_and_resume(tmp_path):
    job=prepare(tmp_path,config(),text='Năng lượng được bảo toàn.')
    folder=Path(job['folder'])
    assert job['config']['template_version']==3
    assert len(job['config']['template_layouts'])==23
    assert (folder/'template.pptx').read_bytes()==asset('standard').read_bytes()
    assert load(tmp_path,folder)['config']['template_sha256']==job['config']['template_sha256']
    (folder/'template.pptx').write_bytes(b'bad')
    with pytest.raises(ValueError,match='Mẫu'):
        load(tmp_path,folder)


@pytest.mark.parametrize('mode',['parallel_columns','english_only'])
def test_image_and_editable_table_follow_their_own_source_slide(tmp_path,mode):
    source=Presentation()
    image=io.BytesIO();Image.new('RGB',(140,60),'#b8d3a1').save(image,'PNG')
    page=source.slides.add_slide(source.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(1)).text='Đồ thị'
    page.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(1), Inches(2), width=Inches(4))
    page=source.slides.add_slide(source.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(1)).text='Bảng số liệu'
    table=page.shapes.add_table(2,2,Inches(.5),Inches(2),Inches(8),Inches(2)).table
    for cell,value in zip((c for r in table.rows for c in r.cells),['Thời gian','Vận tốc','Một giây','Hai mét']):
        cell.text=value
    path=tmp_path/'source.pptx';source.save(path)
    translations={'Đồ thị':'Graph','Bảng số liệu':'Data table','Thời gian':'Time','Vận tốc':'Velocity',
                  'Một giây':'One second','Hai mét':'Two metres'}
    pairs=[(loc,vi,translations[vi]) for loc,vi in source_blocks(source)]
    target=tmp_path/'out.pptx'
    info=build(target,asset('visual'),config('visual',mode),pairs,Event(),path)
    deck=Presentation(target)
    assert len(deck.slides)==2
    assert [p['layout'] for p in info['layout_plan']]==['IMAGE_TEXT','TABLE']
    pictures=[s for s in deck.slides[0].shapes if int(s.shape_type)==13]
    assert len(pictures)==1 and pictures[0].image.blob==image.getvalue()
    assert pictures[0].width/pictures[0].height==pytest.approx(140/60,abs=.001)
    tables=[s.table for s in deck.slides[1].shapes if s.has_table]
    assert len(tables)==1
    assert tables[0].cell(0,0).text==('Time' if mode=='english_only' else 'Thời gian\nTime')
    assert '{{' not in text(deck)


def test_long_units_paginate_without_losing_words_or_moving_between_source_slides(tmp_path):
    pairs=[('Slide 1 · ý 1','Nội dung dài','Long content')]
    pairs += [(f'Slide 1 · ý {n}',f'Ý {n} '+('năng lượng '*40),f'Unit {n} '+('energy '*45)) for n in range(2,8)]
    pairs += [('Slide 2','Kết thúc','End')]
    result=build(tmp_path/'result.pptx',asset('practice'),config('practice'),pairs,Event())
    deck=Presentation(tmp_path/'result.pptx')
    combined=text(deck)
    assert result['pages']>2
    assert combined.count('năng lượng')==6*40
    assert combined.count('energy')==6*45
    assert result['layout_plan'][-1]['source_slide']==2
    assert all(p['source_slide']==1 for p in result['layout_plan'][:-1])


def test_cancel_does_not_publish_partial_output(tmp_path):
    cancel=Event();cancel.set()
    with pytest.raises(InterruptedError):
        build(tmp_path/'result.pptx',asset('standard'),config(),examples(),cancel)
    assert not (tmp_path/'result.pptx').exists()


def test_manifest_pages_and_semantic_selection():
    assert len(manifest()['layouts'])==23
    for preset in ('standard','visual','practice'):
        assert len(Presentation(asset(preset)).slides)==23
    assert choose_layout('Quy trình', [('','Bước một','First step')])=='PROCESS'
    assert choose_layout('Thảo luận nhóm', [('','Trao đổi','Discuss')])=='GROUP_WORK'
    assert choose_layout('Bài tập về nhà', [('','Hoàn thành bài','Complete the exercise')])=='HOMEWORK'
    assert choose_layout('Câu hỏi', [('','A. Một','A. One')])!='QUIZ'


def test_bookends_are_explicit_and_names_do_not_enter_english_slide(tmp_path):
    settings={**config(mode='english_only'),'template_bookends':True,'subject':'Vật lý','grade':'11'}
    info=build(tmp_path/'out.pptx',asset('standard'),settings,examples(),Event())
    deck=Presentation(tmp_path/'out.pptx')
    assert len(deck.slides)==6
    assert info['layout_plan'][0]['layout']=='COVER'
    assert info['layout_plan'][-1]['layout']=='THANK_YOU'
    assert 'Vật lý' not in text(deck) and 'Thank you!' in text(deck)


def test_native_equation_is_not_duplicated_for_its_paragraphs(tmp_path):
    from lxml import etree
    source=Presentation();page=source.slides.add_slide(source.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5),Inches(.3),Inches(8),Inches(1)).text='Công thức'
    shape=page.shapes.add_textbox(Inches(1),Inches(2),Inches(4),Inches(1))
    shape.text='a = 2\nF = ma'
    paragraph=shape.text_frame.paragraphs[0]._p
    paragraph.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'))
    from app.universal_templates import _media
    assert len(_media(page))==1


def test_mixed_prose_and_native_math_share_only_equation_not_untranslated_source_box(tmp_path):
    from lxml import etree
    source=Presentation();page=source.slides.add_slide(source.slide_layouts[6])
    shape=page.shapes.add_textbox(Inches(1),Inches(2),Inches(8),Inches(3))
    vi='Phương trình của vận tốc là:'
    shape.text=vi+'\nVận tốc biến thiên tuần hoàn theo thời gian.'
    paragraph=shape.text_frame.paragraphs[0]._p
    paragraph.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>v = 2</m:t></m:r></m:oMath>'))
    source_path=tmp_path/'source.pptx';source.save(source_path)
    from app.universal_templates import _media
    media=_media(page)
    assert len(media)==1 and not media[0]['shape'].text.strip()
    assert media[0]['shape']._element.xpath('.//*[local-name()="oMath"]')
    target=tmp_path/'out.pptx'
    build(target,asset('standard'),config(),[('Slide 1 · ý 1','Vận tốc','Velocity'),
        ('Slide 1 · ý 2',vi,'The equation for velocity is:')],Event(),source_path)
    deck=Presentation(target)
    assert text(deck).count(vi)==1 and text(deck).count('The equation for velocity is:')==1
    assert sum(len(s._element.xpath('.//*[local-name()="oMath"]')) for page in deck.slides for s in page.shapes)==1


@pytest.mark.parametrize('preset', ['standard','visual','practice'])
@pytest.mark.parametrize('mode', ['parallel_columns','english_only'])
def test_long_first_paragraph_goes_to_paginated_body_instead_of_title_banner(tmp_path, preset, mode):
    vi = 'Nội dung về vận tốc biến thiên tuần hoàn theo thời gian. ' * 5
    en = 'The velocity varies periodically with time during simple harmonic motion. ' * 5
    target = tmp_path/'out.pptx'
    build(target,asset(preset),config(preset,mode),[('Slide 1 · dòng 1',vi,en)],Event())
    combined = text(Presentation(target))
    assert combined.count('The velocity varies periodically') == 5
    assert combined.count('Nội dung về vận tốc') == (0 if mode=='english_only' else 5)
    assert '{{' not in combined


def test_picture_slide_jump_does_not_import_original_slides_or_duplicate_parts(tmp_path):
    import zipfile
    source=Presentation()
    first=source.slides.add_slide(source.slide_layouts[6])
    second=source.slides.add_slide(source.slide_layouts[6])
    picture=io.BytesIO();Image.new('RGB',(60,40),'blue').save(picture,'PNG')
    shape=first.shapes.add_picture(io.BytesIO(picture.getvalue()), Inches(1), Inches(2))
    shape.click_action.target_slide=second
    second.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='Original navigation target'
    original=tmp_path/'source.pptx';source.save(original)
    target=tmp_path/'out.pptx'
    build(target,asset('standard'),config(),[('Slide 1','Đồ thị','Graph')],Event(),original)
    with zipfile.ZipFile(target) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    deck=Presentation(target)
    assert 'Original navigation target' not in text(deck)
    images=[s for page in deck.slides for s in page.shapes if int(s.shape_type)==13]
    assert len(images)==1 and images[0].image.blob==picture.getvalue()


def test_cell_lines_and_question_options_have_stable_translation_units(tmp_path):
    source=Presentation();page=source.slides.add_slide(source.slide_layouts[6])
    shape=page.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(4))
    shape.text='Câu hỏi?\nA. Một\nB. Hai\nC. Ba\nD. Bốn'
    blocks=source_blocks(source)
    assert len(blocks)==5 and len({loc for loc,_ in blocks})==5


def test_plain_document_first_paragraph_is_content_and_roman_header_is_shared(tmp_path):
    pairs=[('Nội dung','Đoạn văn đầu tiên.','The first paragraph.')]
    build(tmp_path/'plain.pptx',asset('standard'),config(),pairs,Event())
    combined=text(Presentation(tmp_path/'plain.pptx'))
    assert 'Energy' in combined and 'The first paragraph.' in combined
    pairs=[('Slide 1 · ý 1','II','II'),('Slide 1 · ý 2','Vận tốc','Velocity')]
    build(tmp_path/'roman.pptx',asset('standard'),config(),pairs,Event())
    combined=text(Presentation(tmp_path/'roman.pptx'))
    assert 'II Vận tốc' in combined and 'II Velocity' in combined


def test_long_table_splits_by_rows_and_preserves_native_chart(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    source=Presentation();page=source.slides.add_slide(source.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5),Inches(.2),Inches(8),Inches(1)).text='Bảng'
    table=page.shapes.add_table(12,2,Inches(1),Inches(2),Inches(8),Inches(3)).table
    for r,row in enumerate(table.rows):
        row.cells[0].text=f'Mốc {r}';row.cells[1].text=f'Giá trị {r}'
    page=source.slides.add_slide(source.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5),Inches(.2),Inches(8),Inches(1)).text='Biểu đồ'
    data=CategoryChartData();data.categories=['A','B'];data.add_series('Energy',[2,4])
    page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(2),Inches(5),Inches(3),data)
    source.save(tmp_path/'source.pptx')
    pairs=[(loc,vi,vi.replace('Mốc','Point').replace('Giá trị','Value').replace('Biểu đồ','Chart').replace('Bảng','Table'))
           for loc,vi in source_blocks(source)]
    result=build(tmp_path/'out.pptx',asset('visual'),config('visual','english_only'),pairs,Event(),tmp_path/'source.pptx')
    deck=Presentation(tmp_path/'out.pptx')
    values=[cell.text for page in deck.slides for shape in page.shapes if shape.has_table
            for row in shape.table.rows for cell in row.cells]
    assert all(f'Point {r}' in values and f'Value {r}' in values for r in range(12))
    assert len(deck.slides)>2
    charts=[shape for page in deck.slides for shape in page.shapes if shape.has_chart]
    assert len(charts)==1 and list(charts[0].chart.series[0].values)==[2,4]
    assert result['layout_plan'][-1]['layout']=='IMAGE_TEXT'


def test_multiple_native_math_paragraphs_paginate_independently_without_source_decoration(tmp_path):
    from lxml import etree
    from app.universal_templates import _media
    from pptx.util import Pt
    source=Presentation();page=source.slides.add_slide(source.slide_layouts[6])
    shape=page.shapes.add_textbox(Inches(1),Inches(2),Inches(6),Inches(3))
    shape.text='\n\n'
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        paragraph.space_after=Pt(40)
        paragraph._p.append(etree.fromstring(('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x = '+str(index)+'</m:t></m:r></m:oMath>').encode()))
    media=_media(page)
    assert len(media)==3
    assert all(len(obj['shape']._element.xpath('.//*[local-name()="oMath"]'))==1 for obj in media)
    assert all(obj['shape'].text_frame.margin_top==0 for obj in media)
    assert all(not obj['shape']._element.xpath('.//a:spcAft') for obj in media)
    source.save(tmp_path/'source.pptx')
    build(tmp_path/'out.pptx',asset('standard'),config(),[('Slide 1','Công thức','Formula')],Event(),tmp_path/'source.pptx')
    deck=Presentation(tmp_path/'out.pptx')
    assert len(deck.slides)==3
    assert all(len(s._element.xpath('.//*[local-name()="oMath"]'))==1 for s in deck.slides)
