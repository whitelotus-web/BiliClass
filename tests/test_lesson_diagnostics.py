import json
from pathlib import Path
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.conversion_input import creation_config
from app.conversion_presentation import display_state
from app.conversion_progress import ConversionProgress
from app.google_conversion import convert, prepare
from app.google_translation import KeyStore, TranslationProblem, Translator
from app.lesson_diagnostics import failed_speech, layout


def test_layout_explains_specific_pages_without_turning_generic_review_into_errors():
    report = layout({'slide_quality': {'slides': [
        {'slide': 1, 'issues': [{'code': 'visual_review', 'severity': 'CHECK'}]},
        {'slide': 2, 'source_slide': 1, 'issues': [{'code': 'possible_overlap', 'severity': 'CHECK', 'excerpt': 'Định nghĩa'}]},
        {'slide': 3, 'issues': [{'code': 'text_overflow', 'severity': 'FAIL'}]},
    ], 'local_repairs': [{'slide': 2, 'operation': 'grow_text_frame'}]}})
    assert [row['slide'] for row in report['pages']] == [3, 2]
    assert report['failures'] == 1 and report['remaining'] == 2 and report['repaired'] == 1
    assert report['pages'][1]['source_slide'] == 1
    assert report['pages'][1]['issues'][0]['excerpt'] == 'Định nghĩa'
    assert report['title'] == 'Cần sửa bố cục trước khi dạy'
    assert layout({})['has_report'] is False


def test_image_review_is_plain_and_retained_without_a_layout_failure():
    report = layout({'external_deck': {'content_review': {'image_pages': [1], 'message': 'Legacy technical message'}},
                     'slide_quality': {'slides': [{'slide': 1, 'issues': [{'code': 'notes_mismatch', 'severity': 'CHECK'}]}]}})
    assert report['title'] == 'Cần xem lại trước khi dạy'
    assert report['failures'] == 0 and report['remaining'] == 1
    assert report['summary'] == 'Có 1 slide cần xem lại.'
    assert report['image_note'] == 'Chữ trong ảnh ở slide nguồn 1 chưa được kiểm tra hết. Hãy đối chiếu chữ, bảng và công thức với bản gốc.'
    assert layout({'external_deck': {'content_review': {'message': 'Cần xem ảnh.'}}})['image_note'] == 'Cần xem ảnh.'


def test_source_review_keeps_exact_reason_and_does_not_invent_an_answer():
    report = layout({'slide_quality': {'slides': [{'slide': 6, 'issues': [
        {'code': 'source_review_note', 'severity': 'CHECK', 'excerpt': 'Kiểm tra giả thiết trong đề gốc.'}]}]}})
    issue = report['pages'][0]['issues'][0]
    assert report['failures'] == 0 and report['pages'][0]['slide'] == 6
    assert issue['excerpt'] == 'Kiểm tra giả thiết trong đề gốc.'
    assert 'giữ nguyên giả thiết và đáp án' in issue['message']


def test_distinct_table_cells_keep_their_own_location_and_excerpt():
    report = layout({'slide_quality': {'slides': [{'slide': 4, 'issues': [
        {'code': 'text_overflow', 'severity': 'FAIL', 'shape_id': 2, 'cell': [1, 0], 'excerpt': 'First question'},
        {'code': 'text_overflow', 'severity': 'FAIL', 'shape_id': 2, 'cell': [2, 0], 'excerpt': 'Second question'},
    ]}]}})
    issues = report['pages'][0]['issues']
    assert len(issues) == 2 and report['failures'] == 1
    assert issues[0]['location'] == 'Ô bảng hàng 2, cột 1.'
    assert issues[1]['location'] == 'Ô bảng hàng 3, cột 1.'
    assert issues[1]['excerpt'] == 'Second question'


def test_failed_shared_speech_lists_every_source_slide_and_language():
    inspection = {'profile': {'units': [
        {'locator': 'Slide 2', 'vi': 'Số một.', 'en': ''},
        {'locator': 'Slide 5', 'vi': 'Số một.', 'en': ''}]}}
    result = failed_speech(inspection, {'vi': 'fixture'}, [('Số một.', 'fixture')])
    assert result == [{'slides': [2, 5], 'languages': ['vi'], 'excerpt': 'Số một.'}]
    assert failed_speech({}, {}, [('Unknown.', 'fixture')])[0]['slides'] == []


def test_failed_native_translation_identifies_slide_and_keeps_original(tmp_path):
    source = tmp_path/'input.pptx'
    deck = Presentation()
    for text in ('Khái niệm mở đầu', 'Một đường thẳng có 2 điểm.'):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = text
    deck.save(source)
    before = source.read_bytes()
    KeyStore(tmp_path/'data').add('AIza'+'x'*35)
    calls = []

    def invalid(rows, *_):
        calls.append(len(rows))
        return [row+' 999' for row in rows]

    translator = Translator(tmp_path/'data', transport=invalid)
    config = creation_config('Bài Toán', 'Toán', 'THPT', '11', 'parallel_columns', 'standard', 'source')
    job = prepare(tmp_path/'data', config, source)
    with pytest.raises(TranslationProblem, match='Slide nguồn 2'):
        convert(tmp_path/'data', job, Event(), ConversionProgress(job['folder'], lambda _: None), translator=translator)
    assert len(calls) == 2, 'One bounded repair; never loop over the entire lesson'
    context = json.loads((Path(job['folder'])/'conversion-issue.json').read_text(encoding='utf-8'))
    assert context['source_slides'] == [2] and context['excerpt'] == 'Một đường thẳng có 2 điểm.'
    shown = display_state({'running': False, 'phase': 'error', 'errorCode': 'api_result', 'issueContext': context})
    assert 'Slide nguồn: 2' in shown['displayMessage'] and context['excerpt'] in shown['displayMessage']
    assert source.read_bytes() == before
