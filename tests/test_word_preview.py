import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from PySide6.QtGui import QGuiApplication, QTextDocument

from app.word_document import Package
from app.word_preview import Preview, length


@pytest.fixture(scope='module', autouse=True)
def gui():
    application = QGuiApplication.instance() or QGuiApplication([])
    yield application


@pytest.fixture
def preview(tmp_path):
    source = tmp_path / 'preview.docx'
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Bài học 😀 chuyển động')
    run.bold, run.font.size, run.font.color.rgb = True, Pt(20), RGBColor.from_string('C00000')
    line = doc.add_paragraph('Vận tốc ')
    line.add_run('12 m/s.').italic = True
    line.add_run('\nDòng hai\tgiữ khoảng trắng.')
    image = io.BytesIO()
    Image.new('RGB', (40, 30), '#0088cc').save(image, format='PNG')
    image.seek(0)
    doc.add_paragraph('Hình: ').add_run().add_picture(image, width=Inches(.8))
    table = doc.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = 'Nội dung giống nhau'
    doc.add_paragraph('Đoạn cuối.')
    doc.save(source)
    signature = source.read_bytes()
    package = Package(source)
    view = Preview(package)
    document = QTextDocument()
    document.setHtml(view.html)
    assert source.read_bytes() == signature
    return package, view, document


def blocks(document):
    block = document.begin()
    while block.isValid():
        yield block
        block = block.next()


def test_preview_shows_original_style_embedded_image_and_table(preview):
    _, view, document = preview
    assert 'font-size:20.0pt' in view.html and 'text-align:center' in view.html
    assert '#C00000' in view.html and 'font-weight:bold' in view.html and 'font-style:italic' in view.html
    assert 'data:image/png;base64,' in view.html and '<table ' in view.html
    assert '\ufffc' in document.toPlainText()
    assert '[Hình minh họa]' not in document.toPlainText()
    assert view.fonts == ['Arial']


def test_source_font_display_reports_multiple_fonts_without_forcing_one_on_export(tmp_path):
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.add_paragraph('Đây là phần nội dung chính sử dụng phông mặc định của tài liệu.')
    doc.add_paragraph().add_run('Đổi phông riêng.').font.name = 'Times New Roman'
    source = tmp_path / 'mixed-fonts.docx'
    doc.save(source)
    view = Preview(Package(source))
    assert view.fonts == ['Arial', 'Times New Roman']


def test_original_size_follows_selected_run_and_reports_mixed_sizes(tmp_path):
    doc = Document()
    doc.styles['Normal'].font.size = Pt(14)
    doc.add_paragraph('Nội dung chính dùng cỡ mười bốn có độ dài lớn nhất.')
    paragraph = doc.add_paragraph()
    paragraph.add_run('Tiêu đề').font.size = Pt(26)
    paragraph.add_run(' và nội dung').font.size = Pt(12)
    source = tmp_path / 'sizes.docx'
    doc.save(source)
    view = Preview(Package(source))
    assert view.selection_size([]) == '14'
    assert view.selection_size([dict(paragraph='p1', start=0, end=7)]) == '26'
    assert view.selection_size([dict(paragraph='p1', start=7, end=19)]) == '12'
    assert view.selection_size([dict(paragraph='p1', start=0, end=19)]) == 'Nhiều cỡ'


def test_selection_maps_repeated_table_cells_to_distinct_original_paragraphs(preview):
    package, view, document = preview
    selected, skipped = view.ranges(document, 0, document.characterCount() - 1)
    assert skipped == 0
    assert [row['paragraph'] for row in selected] == [row['id'] for row in package.rows]
    assert all(region['start'] == 0 and region['end'] == len(package.paragraphs[region['paragraph']][2]['text'])
               for region in selected)


def test_confirmed_highlights_preserve_text_styles_and_distinguish_identical_table_cells(preview):
    from app.word_highlights import SourceHighlights

    package, view, document = preview
    before = document.toHtml()
    first = dict(paragraph='p0', start=0, end=len(package.rows[0]['text']))
    positions = view.selected_positions(document, [first])
    assert positions == [(0, length(package.rows[0]['text']))]
    highlighter = SourceHighlights(document)
    highlighter.set_ranges(positions)
    assert document.firstBlock().layout().formats()[0].format.background().color().name() == '#dbeafe'
    table_region = dict(paragraph='p4', start=0, end=len(package.rows[4]['text']))
    cell_positions = view.selected_positions(document, [table_region])
    selected, _ = view.ranges(document, *cell_positions[0])
    assert selected == [table_region]
    highlighter.set_ranges(cell_positions)
    assert not document.firstBlock().layout().formats()
    assert document.toHtml() == before


def test_selection_after_emoji_and_across_linebreaks_is_not_offset_by_rich_formatting(preview):
    package, view, document = preview
    first = document.firstBlock()
    text = first.text()
    begin = text.index('chuyển')
    selected, _ = view.ranges(document, length(text[:begin]), length(text))
    assert package.rows[0]['text'][selected[0]['start']:selected[0]['end']] == 'chuyển động'
    second = first.next()
    start = second.position() + length('Vận tốc ')
    selected, _ = view.ranges(document, start, second.position() + length(second.text()))
    assert selected == [dict(paragraph='p1', start=8, end=len(package.rows[1]['text']))]


def test_image_object_counts_as_the_full_original_protected_object(preview):
    package, view, document = preview
    picture = next(block for block in blocks(document) if '\ufffc' in block.text())
    selected, _ = view.ranges(document, picture.position(), picture.position() + length(picture.text()))
    assert selected[0]['end'] == len(package.rows[2]['text'])
    assert package.rows[2]['text'][selected[0]['start']:selected[0]['end']] == 'Hình: [Hình minh họa]'


def test_changed_or_unmapped_text_is_rejected_instead_of_translating_wrong_ranges(preview):
    _, view, _ = preview
    document = QTextDocument()
    document.setHtml('<p><a name="blc_p0"></a>Different content</p>')
    with pytest.raises(ValueError, match='Bố cục xem trước'):
        view.ranges(document, 0, 5)


def test_unanchored_blank_line_between_source_paragraphs_is_only_spacing(preview):
    package, view, _ = preview
    document = QTextDocument()
    document.setHtml(view.html.replace('</p>', '</p><p><br/></p>', 1))
    selected, skipped = view.ranges(document, 0, document.characterCount() - 1)
    assert skipped == 0
    assert [region['paragraph'] for region in selected] == [row['id'] for row in package.rows]
    blank = next(block for block in blocks(document) if block.text() and not block.text().strip())
    with pytest.raises(ValueError, match='chưa có phần chữ dịch được'):
        view.ranges(document, blank.position(), blank.position() + length(blank.text()))


def test_sentence_excerpt_keeps_original_offsets_and_style(preview):
    package, view, _ = preview
    excerpt = view.excerpt('p1', 8, len(package.rows[1]['text']))
    document = QTextDocument()
    document.setHtml(excerpt.html)
    selected, _ = excerpt.ranges(document, 0, document.characterCount() - 1)
    assert selected == [dict(paragraph='p1', start=8, end=len(package.rows[1]['text']))]
    assert 'font-style:italic' in excerpt.html


def test_external_images_and_links_are_not_loaded(tmp_path):
    import zipfile

    from lxml import etree

    source = tmp_path / 'external.docx'
    doc = Document()
    doc.add_paragraph('Không truy cập mạng từ tài liệu.')
    doc.save(source)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    rels = etree.fromstring(entries['word/_rels/document.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='rId99', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',
        Target='https://example.invalid/private.png', TargetMode='External')
    entries['word/_rels/document.xml.rels'] = etree.tostring(rels)
    with zipfile.ZipFile(source, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    assert 'https://' not in Preview(Package(source)).html
