import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from app.same_slide import fingerprint
from app.slide_speech import speech_chunks
from app.slide_typography import protect_paragraph
from app.text_frame_space import grow


def paragraph(text):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    shape.text = text
    return shape.text_frame.paragraphs[0]


@pytest.mark.parametrize('text,expected', [
    ('Travel 108 km in 10 s.', 'Travel 108\u00a0km in 10\u00a0s.'),
    ('Speed 108 km/h and 9.81 m/s².', 'Speed 108\u00a0km/h and 9.81\u00a0m/s².'),
    ('A 2,5 kg mass at −10 °C.', 'A 2,5\u00a0kg mass at −10\u00a0°C.'),
    ('10s', '10s'),
    ('10 students, page 108, example 10.', '10 students, page 108, example 10.'),
])
def test_quantities_stay_together_without_changing_numbers_or_normal_words(text, expected):
    p = paragraph(text)
    before = fingerprint(p._p)
    protect_paragraph(p._p)
    assert p.text == expected
    assert fingerprint(p._p) == before
    assert speech_chunks(p.text) == speech_chunks(text)


def test_keep_number_and_unit_across_runs_and_soft_break_with_style_and_math_unchanged():
    p = paragraph('')
    first = p.add_run()
    first.text = 'Travel 108'
    first.font.bold = True
    p.add_line_break()
    unit = p.add_run()
    unit.text = 'km in '
    unit.font.italic = True
    number = p.add_run()
    number.text = '10 '
    last = p.add_run()
    last.text = 's.'
    math = etree.SubElement(p._p, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
    etree.SubElement(math, '{http://schemas.openxmlformats.org/officeDocument/2006/math}t').text = '10 m'
    before = etree.tostring(math)
    assert protect_paragraph(p._p)
    assert p.text == 'Travel 108\u00a0km in 10\u00a0s.'
    assert first.font.bold and unit.font.italic
    assert etree.tostring(math) == before
    assert not protect_paragraph(p._p), 'Typography must be idempotent'


def test_generated_frame_uses_free_space_and_does_not_cover_neighbor_text():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label = slide.shapes.add_textbox(Inches(2), Inches(1), Inches(.45), Inches(.5))
    label.name = 'BCL_BODY Label'
    label.text = 'Bắc\nNorth'
    for p in label.text_frame.paragraphs:
        p.font.size = Pt(18)
    neighbor = slide.shapes.add_textbox(Inches(3.3), Inches(.7), Inches(1), Inches(2))
    neighbor.text = 'Other content'
    before = (label.left, label.top, label.width, label.height)
    assert grow(label._element, slide, deck.slide_width, deck.slide_height)
    assert label.width > before[2] or label.height > before[3]
    assert label.left + label.width < neighbor.left
    assert label.text == 'Bắc\nNorth'
    label.name = 'Original teacher label'
    assert not grow(label._element, slide, deck.slide_width, deck.slide_height)
