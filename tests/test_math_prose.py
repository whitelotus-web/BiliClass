import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.math_prose import MATH, fragments
from app.native_bilingual import assemble, extract
from app.source_deck import NS


def equation(*parts):
    root = etree.Element('{http://schemas.microsoft.com/office/drawing/2010/main}m')
    math = etree.SubElement(root, '{' + MATH + '}oMath')
    for text, normal in parts:
        run = etree.SubElement(math, '{' + MATH + '}r')
        if normal:
            etree.SubElement(etree.SubElement(run, '{' + MATH + '}rPr'), '{' + MATH + '}nor')
        etree.SubElement(run, '{' + MATH + '}t').text = text
    return root


def test_words_split_across_normal_math_runs_leave_symbols_intact():
    source = equation(('𝒄', False), (' ', True), ('v', True), ('ớ', True), ('i ', True), ('𝒂', False))
    before = etree.tostring(source)
    parts = fragments(source)
    assert [(kind, value if kind == 'text' else ''.join(value.itertext())) for kind, value in parts] == [
        ('xml', '𝒄'), ('text', ' với '), ('xml', '𝒂')]
    assert etree.tostring(source) == before
    fraction = equation()
    node = etree.SubElement(fraction[0], '{' + MATH + '}f')
    node.append(equation(('nếu', True))[0])
    assert fragments(fraction) is None, 'Never split inside fractions or other structural math'
    hidden = equation(('Có một mặt phẳng', True))
    etree.SubElement(etree.SubElement(hidden[0][0], '{' + NS['a'] + '}rPr'), '{' + NS['a'] + '}noFill')
    assert fragments(hidden) is None


@pytest.mark.parametrize('display_math', [False, True])
def test_entire_option_typed_as_math_is_translated_without_altering_vietnamese(tmp_path, display_math):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]._p
    original = equation(('Có một và chỉ một mặt phẳng đi qua ba điểm không thẳng hàng.', True))
    props = etree.SubElement(original[0][0], '{' + NS['a'] + '}rPr', sz='4800')
    etree.SubElement(etree.SubElement(props, '{' + NS['a'] + '}solidFill'), '{' + NS['a'] + '}schemeClr', val='bg1')
    etree.SubElement(etree.SubElement(paragraph, '{' + NS['a'] + '}pPr'), '{' + NS['a'] + '}defRPr')
    if display_math:
        body = original[0]
        original.remove(body)
        wrapper = etree.SubElement(original, '{' + MATH + '}oMathPara')
        etree.SubElement(wrapper, '{' + MATH + '}oMathParaPr')
        wrapper.append(body)
    paragraph.append(original)
    source, target = tmp_path/'source.pptx', tmp_path/'bilingual.pptx'
    deck.save(source)
    before = source.read_bytes()
    manifest = extract(source, 'parallel_columns', same_slide=True)
    assert len(manifest['units']) == 1 and manifest['units'][0]['language'] == 'vi'
    assert 'mặt phẳng' in manifest['units'][0]['text']
    assemble(source, target, manifest, {manifest['units'][0]['id']: {
        'en': 'There is exactly one plane through three noncollinear points.'}}, 'parallel_columns')
    converted = etree.fromstring(etree.tostring(Presentation(target).slides[0]._element))
    assert 'There is exactly one plane' in ''.join(converted.xpath('.//a:t/text()', namespaces=NS))
    assert 'Có một và chỉ một mặt phẳng đi qua ba điểm không thẳng hàng.' in ''.join(converted.xpath('.//a:t/text()', namespaces=NS))
    assert not converted.xpath('.//*[local-name()="oMath"]'), 'Pure prose must wrap and fit as ordinary editable text'
    assert converted.xpath('.//a:rPr/a:solidFill/a:schemeClr[@val="bg1"]', namespaces=NS), 'Retain light ink over a dark answer band'
    assert source.read_bytes() == before


def test_native_math_size_and_symbols_survive_prose_fitting(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    math = equation(('𝑺𝑶', False), (' với ', True), ('𝑶', False))
    for run in math[0]:
        etree.SubElement(run, '{' + NS['a'] + '}rPr', sz='4800')
    shape.text_frame.paragraphs[0]._p.append(math)
    source, target = tmp_path/'source.pptx', tmp_path/'output.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    unit = manifest['units'][0]
    assert unit['font_size'] == 4800, 'Read equation fonts instead of the ordinary-run default'
    assemble(source, target, manifest, {unit['id']: {'en': unit['text'].replace(' với ', ', where ')}}, 'parallel_columns')
    tree = etree.fromstring(etree.tostring(Presentation(target).slides[0]._element))
    symbols = tree.xpath('.//*[local-name()="oMath"]//*[local-name()="t"]/text()')
    assert symbols.count('𝑺𝑶') == symbols.count('𝑶') == 2
    assert not any(value in (' với ', ', where ') for value in symbols)
