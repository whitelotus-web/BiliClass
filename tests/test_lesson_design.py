import json

import pytest

from app.analytics import list_sessions, report
from app.classroom_store import ClassroomStore
from app.deck_export import export_deck
from app.importers import parse_document
from app.lesson_design import paired_pages
from app.library import Library
from app.readiness import assess


def test_paired_pages_never_splits_one_language_away_from_the_other():
    pages = paired_pages("Ý một.\nÝ hai.", "First idea.\nSecond idea.", limit=19)
    assert pages == [("Ý một.", "First idea."), ("Ý hai.", "Second idea.")]
    with pytest.raises(ValueError, match="tách đoạn"):
        paired_pages("Một ý rất dài mà không có cặp tương ứng được tách ra", "Short", limit=25)


def test_powerpoint_keeps_multiple_ideas_and_visual_in_export(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "source.pptx"
    image = tmp_path / "diagram.png"
    Image.new("RGB", (120, 80), "#0869f9").save(image)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(image), Inches(2), Inches(2))
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(.5)).text = "Khái niệm"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(.5)).text = "Ví dụ: lực"
    slide.shapes.add_picture(str(image), Inches(8), Inches(2), width=Inches(2))
    deck.save(source)

    blocks = parse_document(source)
    assert blocks == [("Slide 2 · ý 1", "Khái niệm"), ("Slide 2 · ý 2", "Ví dụ: lực")]
    library = Library(tmp_path / "library")
    try:
        saved = library.store_source(source)
        lesson = library.create("Bài lực", "Vật lý", "THPT", "10", blocks, saved)
        assert [segment["kind"] for segment in lesson["segments"]] == ["concept", "example"]
        for segment, english in zip(lesson["segments"], ("Concept", "Example: force"), strict=True):
            lesson = library.edit_segment(lesson["id"], segment["id"], segment["vi"], english, True)
        lesson = library.set_teaching_preset(lesson["id"], "visual")
        lesson = library.set_presentation_style(lesson["id"], "template")
        lesson = library.mark_prepared(lesson["id"], lesson["revision"])
        exported = Presentation(export_deck(lesson, library.directory, tmp_path / "bilingual.pptx"))
        assert len(exported.slides) == 4
        assert any(shape.shape_type == 13 for shape in exported.slides[0].shapes)
        assert any(shape.shape_type == 13 for shape in exported.slides[1].shapes)
        for index, expected in enumerate(("Concept", "Example: force"), 2):
            assert expected in "\n".join(shape.text for shape in exported.slides[index].shapes if shape.has_text_frame)
    finally:
        library.close()


def test_preparation_is_required_for_real_lesson_and_old_session_is_stable(tmp_path):
    library = Library(tmp_path)
    classroom = ClassroomStore(tmp_path / "classroom.db")
    try:
        lesson = library.create("Bài test", "Toán", "THPT", "10", [("Ý 1", "Một"), ("Ý 2", "Hai")])
        first, second = lesson["segments"]
        lesson = library.edit_segment(lesson["id"], first["id"], "Một", "One", True)
        with pytest.raises(ValueError, match="Duyệt đủ"):
            library.mark_prepared(lesson["id"], lesson["revision"])
        lesson = library.edit_segment(lesson["id"], second["id"], "Hai", "Two", True)
        with pytest.raises(ValueError, match="chốt bản chuẩn bị"):
            classroom.create(lesson, "Lớp A")
        lesson = library.mark_prepared(lesson["id"], lesson["revision"])
        assert assess(lesson, library.directory, [], {}, 0)["prepared_ready"]
        session, _ = classroom.create(lesson, "Lớp A")
        original = json.loads(classroom.session(session)["lesson"])
        listed = list_sessions(tmp_path)[0]
        assert (listed["lesson_id"], listed["grade"], listed["revision"]) == (lesson["id"], "10", lesson["revision"])
        assert report(tmp_path, session)["lesson_title"] == "Bài test"
        edited = library.edit_segment(lesson["id"], first["id"], "Một mới", "New one", True)
        assert not assess(edited, library.directory, [], {}, 0)["prepared_ready"]
        with pytest.raises(ValueError, match="chốt bản chuẩn bị"):
            classroom.create(edited, "Lớp B")
        assert json.loads(classroom.session(session)["lesson"]) == original
    finally:
        classroom.close()
        library.close()
