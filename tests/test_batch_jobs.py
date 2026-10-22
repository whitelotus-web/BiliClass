import copy
import errno
import hashlib
import json
import shutil
import wave
from contextlib import closing
from pathlib import Path
from threading import Event
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtCore import QUrl

from app import batch_jobs, speech
from app.completed_export import save_converted
from app.google_translation import TranslationProblem
from app.library import Library
from app.storage_space import StorageProblem


def batch(tmp_path, kind='convert', count=3):
    rows = batch_jobs.selection(kind, [QUrl.fromLocalFile(str(tmp_path/f'Bài {n}.pptx')).toString() for n in range(count)])
    rows['options'] = {'subject': 'Toán'}
    return rows


@pytest.mark.parametrize('kind', ['convert', 'voice'])
def test_stop_arriving_after_last_saved_result_does_not_leave_a_stuck_paused_queue(tmp_path, kind):
    state = batch(tmp_path, kind=kind, count=2)
    cancel = Event()
    seen = []

    def complete(directory, row, options, event, update):
        seen.append(row['id'])
        if len(seen) == 2:
            event.set()
        return {'lesson_id': 'saved-' + row['id']}

    result = batch_jobs.run(tmp_path, state, cancel, lambda _: None, process=complete)
    assert cancel.is_set() and len(seen) == 2
    assert result['phase'] == 'finished' and all(row['status'] == 'done' for row in result['rows'])
    assert batch_jobs.load(tmp_path, state['id'])['phase'] == 'finished'


def test_elapsed_clock_keeps_per_file_time_and_excludes_pause_between_runs(tmp_path, monkeypatch):
    now = [100.0]
    monkeypatch.setattr(batch_jobs, 'monotonic', lambda: now[0])
    cancel, updates = Event(), []

    def interrupted(directory, row, options, event, update):
        now[0] += 3.25
        update({'percent': 40}, True)
        event.set()
        raise InterruptedError()

    paused = batch_jobs.run(tmp_path, batch(tmp_path, count=2), cancel, updates.append, process=interrupted)
    assert paused['phase'] == 'paused'
    assert paused['elapsed_seconds'] == paused['rows'][0]['elapsed_seconds'] == 3.25
    assert paused['rows'][1].get('elapsed_seconds', 0) == 0
    assert batch_jobs.load(tmp_path, paused['id'])['elapsed_seconds'] == 3.25
    now[0] += 900  # Closing/reopening or pausing must not add 15 minutes of work.
    cancel.clear()

    def finished(directory, row, options, event, update):
        now[0] += 2
        update({'percent': 80})
        return {}

    result = batch_jobs.run(tmp_path, paused, cancel, updates.append, process=finished)
    assert result['elapsed_seconds'] == 7.25
    assert [row['elapsed_seconds'] for row in result['rows']] == [5.25, 2]
    assert result['current_row_id'] == result['rows'][1]['id']
    assert updates[-1]['phase'] == 'finished'


def test_selection_deduplicates_paths_and_never_reads_or_imports_files(tmp_path, monkeypatch):
    monkeypatch.setattr(Path, 'is_file', lambda *_: pytest.fail('No file IO while choosing'))
    url = QUrl.fromLocalFile(str(tmp_path/'Bài mới.pptx')).toString()
    state = batch_jobs.selection('voice', [url, url])
    assert len(state['rows']) == 1 and state['rows'][0]['name'] == 'Bài mới.pptx'
    with pytest.raises(ValueError):
        batch_jobs.selection('voice', ['https://example.com/a.pptx'])
    with pytest.raises(ValueError):
        batch_jobs.selection('voice', [url]*101)


def test_edited_original_is_detected_but_missing_original_can_resume_from_managed_copy(tmp_path):
    path = tmp_path/'Source.pptx'
    path.write_bytes(b'original')
    row = {'url': QUrl.fromLocalFile(str(path)).toString()}
    expected = hashlib.sha256(path.read_bytes()).hexdigest()
    batch_jobs.verify_input_version(row, expected)
    path.write_bytes(b'teacher edited')
    with pytest.raises(ValueError, match='đã sửa'):
        batch_jobs.verify_input_version(row, expected)
    path.unlink()
    batch_jobs.verify_input_version(row, expected)


def test_invalid_saved_queue_is_reported_as_a_read_error(tmp_path):
    state = batch(tmp_path)
    path = batch_jobs.journal(tmp_path, state['id'])
    path.parent.mkdir(parents=True)
    path.write_text('[]', encoding='utf-8')
    with pytest.raises(ValueError, match='đọc được'):
        batch_jobs.load(tmp_path, state['id'])


def test_each_lesson_finishes_before_next_and_local_failure_does_not_block_the_rest(tmp_path):
    state, calls, events = batch(tmp_path), [], []

    def process(directory, row, options, cancel, update):
        calls.append(row['id'])
        assert options == {'subject': 'Toán'}
        if row['name'] == 'Bài 1.pptx':
            raise ValueError('Slide 3: chưa đọc được một vùng chữ.')
        return {'output': row['name']+'.result'}

    finished = batch_jobs.run(tmp_path, state, Event(), events.append, process=process)
    assert calls == [row['id'] for row in state['rows']]
    assert [row['status'] for row in finished['rows']] == ['done', 'failed', 'done']
    assert 'Slide 3' in finished['rows'][1]['message']
    assert all(sum(row['status'] == 'running' for row in event['rows']) <= 1 for event in events)
    calls.clear()
    retried = batch_jobs.run(tmp_path, finished, Event(), events.append,
                             process=lambda _, row, *_args: calls.append(row['id']) or {}, retry=True)
    assert calls == [state['rows'][1]['id']]
    assert all(row['status'] == 'done' for row in retried['rows'])
    assert retried['rows'][0]['output'] == finished['rows'][0]['output']


def test_stop_preserves_current_checkpoint_and_resumes_without_repeating_completed_files(tmp_path):
    state, cancel = batch(tmp_path), Event()

    def process(_, row, options, event, update):
        if row['name'] == 'Bài 1.pptx':
            update({'folder': 'saved-per-file-checkpoint'}, True)
            event.set()
            raise InterruptedError()
        return {'output': 'first result'}

    stopped = batch_jobs.run(tmp_path, state, cancel, lambda _: None, process=process)
    assert stopped['phase'] == 'paused'
    restored = batch_jobs.load(tmp_path, stopped['id'])
    assert [row['status'] for row in restored['rows']] == ['done', 'waiting', 'waiting']
    assert restored['rows'][1]['folder'] == 'saved-per-file-checkpoint'
    calls = []
    resumed = batch_jobs.run(tmp_path, restored, Event(), lambda _: None,
        process=lambda _, row, *_args: calls.append(row['name']) or {})
    assert resumed['phase'] == 'finished' and calls == ['Bài 1.pptx', 'Bài 2.pptx']


def test_stop_during_final_save_keeps_that_success_and_does_not_start_another_lesson(tmp_path):
    state, cancel, calls = batch(tmp_path), Event(), []

    def finishing(_, row, options, event, update):
        calls.append(row['id'])
        event.set()
        return {'output': 'completed-file.pptx'}

    result = batch_jobs.run(tmp_path, state, cancel, lambda _: None, process=finishing)
    assert result['phase'] == 'paused' and len(calls) == 1
    assert [row['status'] for row in result['rows']] == ['done', 'waiting', 'waiting']
    assert result['rows'][0]['output'] == 'completed-file.pptx'


@pytest.mark.parametrize('error', [TranslationProblem('api_key', 'Kiểm tra quyền dịch.'),
                                 StorageProblem('Ổ đĩa chưa đủ chỗ.')])
def test_global_failure_pauses_without_failing_every_input_and_retry_starts_at_same_file(tmp_path, error):
    calls = []

    def process(_, row, *_args):
        calls.append(row['id'])
        raise error

    state = batch(tmp_path)
    stopped = batch_jobs.run(tmp_path, state, Event(), lambda _: None, process=process)
    assert stopped['phase'] == 'paused' and len(calls) == 1
    assert all(row['status'] == 'waiting' for row in stopped['rows'])
    calls.clear()
    result = batch_jobs.run(tmp_path, stopped, Event(), lambda _: None,
        process=lambda _, row, *_args: calls.append(row['id']) or {})
    assert calls == [row['id'] for row in state['rows']] and result['phase'] == 'finished'


def test_failed_checkpoint_cannot_start_costly_work_and_can_be_retried_after_disk_recovers(tmp_path, monkeypatch):
    state, calls = batch(tmp_path), []
    replace = Path.replace
    with monkeypatch.context() as patch:
        patch.setattr(Path, 'replace', lambda *_: (_ for _ in ()).throw(OSError(errno.ENOSPC, 'full')))
        stopped = batch_jobs.run(tmp_path, state, Event(), lambda _: None,
            process=lambda *_: calls.append(True))
    assert stopped['phase'] == 'paused' and stopped['revision'] == 0 and not calls
    assert Path.replace is replace
    result = batch_jobs.run(tmp_path, stopped, Event(), lambda _: None, process=lambda *_: {})
    assert result['phase'] == 'finished'


def test_two_windows_cannot_run_or_overwrite_same_queue_and_stale_state_is_rejected(tmp_path):
    state = batch(tmp_path)
    with batch_jobs.queue_lock(tmp_path):
        blocked = batch_jobs.run(tmp_path, state, Event(), lambda _: None,
                                process=lambda *_: pytest.fail('Second queue must not run'))
    assert blocked['phase'] == 'paused' and 'cửa sổ' in blocked['message']
    completed = batch_jobs.run(tmp_path, state, Event(), lambda _: None, process=lambda *_: {})
    stale = copy.deepcopy(completed)
    batch_jobs.persist(tmp_path, completed)
    rejected = batch_jobs.run(tmp_path, stale, Event(), lambda _: None,
                             process=lambda *_: pytest.fail('Stale queue must not run'))
    assert rejected['phase'] == 'paused'
    assert batch_jobs.load(tmp_path, state['id'])['revision'] == completed['revision']
    with pytest.raises(ValueError):
        batch_jobs.journal(tmp_path, '../outside')
    batch_jobs.dismiss(tmp_path, completed)
    assert batch_jobs.load(tmp_path, state['id'])['phase'] == 'dismissed'


def deck(path, text='Năng lượng được bảo toàn.'):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = text
    slide.notes_slide.notes_text_frame.text = f'VI: {text}\nEN: Energy is conserved.'
    presentation.save(path)
    return path


def test_real_narration_batch_creates_voiced_products_next_to_each_origin_reuses_wavs_and_preserves_originals(tmp_path, monkeypatch):
    paths = [deck(tmp_path/'Bài năng lượng.pptx'), deck(tmp_path/'Bài vận tốc.pptx', 'Vật có vận tốc.')]
    before = {path: path.read_bytes() for path in paths}
    calls = []

    def synthesize(text, voice, rate, directory, cancel):
        calls.append((text, voice))
        path = Path(directory)/(speech.audio_key(text, voice, rate)+'.wav')
        path.parent.mkdir(parents=True, exist_ok=True)
        with wave.open(str(path), 'wb') as wav:
            wav.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
            wav.writeframes(b'\1\0'*2400)
        return {'path': str(path), 'cached': False}

    monkeypatch.setattr(speech, 'synthesize_background', synthesize)
    monkeypatch.setattr(speech, 'narration_workers', lambda *_: 1)
    state = batch_jobs.selection('voice', [QUrl.fromLocalFile(str(path)).toString() for path in paths])
    state['options'] = {'mode': 'both', 'voice': 'pair:truc-ly', 'rate': 0, 'position': 'top_left'}
    directory = tmp_path/'data'
    result = batch_jobs.run(directory, state, Event(), lambda _: None)
    assert [row['status'] for row in result['rows']] == ['done', 'done'], result
    assert calls
    for row in result['rows']:
        output = Path(row['output'])
        assert output.parent == tmp_path and output not in paths
        with ZipFile(output) as archive:
            assert any(name.endswith('.wav') for name in archive.namelist())
        with closing(Library(directory)) as library:
            lesson = library.get(row['lesson_id'])
            assert lesson['voice_published'] and lesson['speaker_position'] == 'top_left'
    assert all(path.read_bytes() == before[path] for path in paths)
    # Same exact presentation in another original folder must export beside THAT file too.
    another = tmp_path/'different folder'
    another.mkdir()
    copied = another/'Same bytes new name.pptx'
    shutil.copyfile(paths[0], copied)
    wavs = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (directory/'audio').glob('*.wav')}
    calls.clear()
    repeated = batch_jobs.selection('voice', [QUrl.fromLocalFile(str(copied)).toString()])
    repeated['options'] = state['options']
    repeated = batch_jobs.run(directory, repeated, Event(), lambda _: None)
    assert repeated['rows'][0]['status'] == 'done', repeated
    assert Path(repeated['rows'][0]['output']).parent == another and not calls
    assert wavs == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (directory/'audio').glob('*.wav')}
    # A checkpoint from the first origin must not inherit the later import's destination.
    resumed = copy.deepcopy(result)
    resumed['rows'][0]['status'] = 'waiting'
    resumed['phase'] = 'paused'
    batch_jobs.persist(directory, resumed)
    resumed = batch_jobs.run(directory, resumed, Event(), lambda _: None)
    assert resumed['rows'][0]['output'] == result['rows'][0]['output'] and not calls
    first_origin = batch_jobs.selection('voice', [QUrl.fromLocalFile(str(paths[0])).toString()])
    first_origin['options'] = state['options']
    first_origin = batch_jobs.run(directory, first_origin, Event(), lambda _: None)
    assert first_origin['rows'][0]['output'] == result['rows'][0]['output'] and not calls


def test_real_narration_stop_keeps_first_wav_and_resume_reuses_it_before_next_file(tmp_path, monkeypatch):
    paths = [deck(tmp_path/'Bài đầu.pptx'), deck(tmp_path/'Bài sau.pptx', 'Vật chuyển động.')]
    calls, cancel = [], Event()

    def synthesize(text, voice, rate, directory, event):
        path = Path(directory)/(speech.audio_key(text, voice, rate)+'.wav')
        path.parent.mkdir(parents=True, exist_ok=True)
        with wave.open(str(path), 'wb') as wav:
            wav.setparams((1, 2, 24000, 0, 'NONE', 'not compressed'))
            wav.writeframes(b'\1\0'*2400)
        calls.append(path)
        if len(calls) == 1:
            event.set()
        return {'path': str(path), 'cached': False}

    monkeypatch.setattr(speech, 'synthesize_background', synthesize)
    monkeypatch.setattr(speech, 'narration_workers', lambda *_: 1)
    directory = tmp_path/'data'
    state = batch_jobs.selection('voice', [QUrl.fromLocalFile(str(path)).toString() for path in paths])
    state['options'] = {'mode': 'both', 'voice': 'pair:truc-ly', 'rate': 0, 'position': 'bottom_right'}
    stopped = batch_jobs.run(directory, state, cancel, lambda _: None)
    assert stopped['phase'] == 'paused' and all(row['status'] == 'waiting' for row in stopped['rows'])
    assert len(calls) == 1 and stopped['rows'][0].get('lesson_id') and not stopped['rows'][1].get('lesson_id')
    original = calls[0].read_bytes()
    cancel.clear()
    finished = batch_jobs.run(directory, stopped, cancel, lambda _: None)
    assert [row['status'] for row in finished['rows']] == ['done', 'done'], finished
    assert calls.count(calls[0]) == 1 and calls[0].read_bytes() == original


def test_real_conversion_batch_uses_each_title_and_exports_without_voice_without_touching_sources(tmp_path, monkeypatch):
    from test_google_conversion import transport

    from app.google_conversion import convert
    from app.google_translation import KeyStore, Translator

    paths = [deck(tmp_path/'Bài Một.pptx'), deck(tmp_path/'Bài Hai.pptx')]
    directory = tmp_path/'data'
    KeyStore(directory).add('AIza'+'a'*35)
    monkeypatch.setattr('app.google_conversion.convert', lambda *args, **kwargs:
        convert(*args, **kwargs, translator=Translator(directory, transport=transport)))
    state = batch_jobs.selection('convert', [QUrl.fromLocalFile(str(path)).toString() for path in paths])
    state['options'] = {'subject': 'Vật lý', 'education': 'THPT', 'grade': '11', 'format': 'parallel_columns',
                        'preset': 'standard', 'style': 'source'}
    before = [path.read_bytes() for path in paths]
    result = batch_jobs.run(directory, state, Event(), lambda _: None)
    assert [row['status'] for row in result['rows']] == ['done', 'done'], result
    for path, row in zip(paths, result['rows'], strict=True):
        with closing(Library(directory)) as library:
            lesson = library.get(row['lesson_id'])
            assert lesson['title'] == path.stem and lesson['subject'] == 'Vật lý'
            assert not lesson['narration'].get('keys')
            from app.chatgpt_handoff import external_preview
            exported = save_converted(directory, lesson, external_preview(lesson, directory))
            assert exported['reused'] and exported['path'] == row['output']
        with ZipFile(row['output']) as archive:
            assert not any(name.endswith('.wav') for name in archive.namelist())
    assert [path.read_bytes() for path in paths] == before


def test_library_selection_keeps_existing_ids_skips_ready_voice_and_detects_later_edits(tmp_path, monkeypatch):
    from app.voice_import import import_presentation

    directory = tmp_path/'data'
    lessons = [import_presentation(directory, deck(tmp_path/f'Bài {index}.pptx', f'Nội dung {index}.'))[0]
               for index in range(2)]
    monkeypatch.setattr('app.lesson_narration.library_audio_status',
                        lambda lesson, _: {'ready': lesson['id'] == lessons[1]['id']})
    state = batch_jobs.library_selection(directory, [lesson['id'] for lesson in lessons])
    assert state['skipped_ready'] == 1 and len(state['rows']) == 1
    row = state['rows'][0]
    assert row['library_lesson'] and row['lesson_id'] == lessons[0]['id']
    assert row['name'] == lessons[0]['title']
    with closing(Library(directory)) as library:
        assert library.db.execute('SELECT COUNT(*) FROM lessons').fetchone()[0] == 2
        changed = library.get(lessons[0]['id'])
        changed['revision'] += 1
        library.db.execute('UPDATE lessons SET content=? WHERE id=?', (json.dumps(changed), changed['id']))
        library.db.commit()
    state['options'] = {'mode': 'en', 'voice': 'pair:truc-ly', 'rate': 0, 'position': 'bottom_right'}
    result = batch_jobs.run(directory, state, Event(), lambda _: None)
    assert result['rows'][0]['status'] == 'failed' and 'đã sửa' in result['rows'][0]['message']
    with pytest.raises(ValueError, match='đều có giọng'):
        batch_jobs.library_selection(directory, [lessons[1]['id']])
