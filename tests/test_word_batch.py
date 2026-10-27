import copy
import json
from contextlib import closing
from pathlib import Path
from threading import Event

import pytest
from docx import Document
from PySide6.QtCore import QUrl

from app import word_batch_jobs as batches
from app import word_jobs as jobs
from app.google_translation import TranslationProblem
from app.library import Library
from app.storage_space import StorageProblem
from app.word_document import Package


class Translator:
    def configure_terms(self, *_):
        pass

    def configure_knowledge(self, *_):
        pass

    def translate(self, rows, cancel, progress):
        if cancel.is_set():
            raise InterruptedError()
        progress(len(rows), len(rows))
        return [row.replace('Vận tốc', 'Velocity').replace('Bài học', 'Lesson') for row in rows]


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    monkeypatch.setattr('app.translation_knowledge.lookup', lambda *_: None)
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    paths = []
    for number in range(3):
        path = tmp_path / f'Thư mục {number + 1}' / 'Bài học.docx'
        path.parent.mkdir()
        doc = Document()
        doc.add_paragraph('Bài học')
        doc.add_paragraph('Vận tốc 12 m/s.')
        doc.save(path)
        paths.append(path)
    # Identical files in different locations must save beside each origin.
    paths[1].write_bytes(paths[0].read_bytes())
    original_convert = jobs.convert
    calls = []

    def convert(*args, **kwargs):
        calls.append(args[1])
        return original_convert(*args, translator=Translator(), **kwargs)

    monkeypatch.setattr(jobs, 'convert', convert)
    return directory, paths, calls


def selected(directory, paths, mode='below'):
    from app.translation_knowledge import MODE_VERSION

    data = batches.selection([path.as_uri() for path in paths])
    data['options'] = dict(formatting=jobs.options_for(dict(mode=mode)), subject='Vật lý', grade='11',
        route=dict(translation_provider='google', translation_knowledge_online=False, translation_knowledge_version=MODE_VERSION))
    batches.persist(directory, data)
    return data


@pytest.mark.parametrize('mode', ['below', 'columns', 'english', 'keywords'])
def test_batch_always_selects_all_and_keeps_each_origin_and_manual_job(prepared, mode):
    directory, paths, calls = prepared
    original_bytes = [path.read_bytes() for path in paths]
    manual = jobs.create(directory, paths[0])
    manual = jobs.selections(directory, manual['id'], [dict(paragraph='p0', start=0, end=3)])
    manual_checkpoint = jobs.read(directory, manual['id'])
    result = batches.run(directory, selected(directory, paths[:2], mode), Event(), lambda _: None)
    assert result['phase'] == 'finished'
    assert [row['status'] for row in result['rows']] == ['done', 'done']
    assert len(set(calls)) == 2 and manual['id'] not in calls
    assert jobs.read(directory, manual['id']) == manual_checkpoint
    for source, row in zip(paths, result['rows']):
        job = jobs.read(directory, row['word_id'])
        assert job['whole_document'] and len(job['regions']) == 2
        assert job['options']['mode'] == mode
        assert job['title'] == source.stem and job['source_original'] == str(source.resolve())
        output = source.with_name(source.stem + ' (Song ngu VI-EN).docx')
        assert Path(row['output']) == output and output.is_file()
        assert Package(output).rows
    assert [path.read_bytes() for path in paths] == original_bytes


def test_bad_file_does_not_block_later_files_and_retry_skips_completed(prepared):
    directory, paths, calls = prepared
    valid = paths[1].read_bytes()
    paths[1].write_text('Not a Word package', encoding='utf-8')
    data = selected(directory, paths)
    result = batches.run(directory, data, Event(), lambda _: None)
    assert [row['status'] for row in result['rows']] == ['done', 'failed', 'done']
    assert len(calls) == 2
    first_output = Path(result['rows'][0]['output']).read_bytes()
    paths[1].write_bytes(valid)
    result = batches.run(directory, result, Event(), lambda _: None, retry=True)
    assert all(row['status'] == 'done' for row in result['rows'])
    assert len(calls) == 3 and len(set(calls)) == 3
    assert Path(result['rows'][0]['output']).read_bytes() == first_output


def test_stop_retains_completed_files_and_reloaded_queue_resumes(prepared):
    directory, paths, calls = prepared
    data = selected(directory, paths)
    cancel = Event()

    def stop_after_first(value):
        if value['rows'][0]['status'] == 'done':
            cancel.set()

    result = batches.run(directory, data, cancel, stop_after_first)
    assert result['phase'] == 'paused'
    assert [row['status'] for row in result['rows']] == ['done', 'waiting', 'waiting']
    loaded = batches.load(directory, data['id'])
    first_identity = loaded['rows'][0]['word_id']
    result = batches.run(directory, loaded, Event(), lambda _: None)
    assert result['phase'] == 'finished' and all(row['status'] == 'done' for row in result['rows'])
    assert result['rows'][0]['word_id'] == first_identity and calls.count(first_identity) == 1


@pytest.mark.parametrize('failure', [TranslationProblem('api_quota', 'Hết hạn mức.'), StorageProblem('Ổ lưu trữ đầy.')])
def test_shared_failure_pauses_remaining_files_and_can_retry(prepared, monkeypatch, failure):
    directory, paths, calls = prepared
    convert = jobs.convert
    fail = [True]

    def sometimes(*args, **kwargs):
        if fail[0] and len(calls) == 1:
            raise failure
        return convert(*args, **kwargs)

    monkeypatch.setattr(jobs, 'convert', sometimes)
    result = batches.run(directory, selected(directory, paths), Event(), lambda _: None)
    assert result['phase'] == 'paused'
    assert [row['status'] for row in result['rows']] == ['done', 'failed', 'waiting']
    fail[0] = False
    result = batches.run(directory, result, Event(), lambda _: None, retry=True)
    assert all(row['status'] == 'done' for row in result['rows']) and len(calls) == 3


def test_edited_source_after_stop_requires_reselection(prepared):
    directory, paths, _ = prepared
    data = selected(directory, paths[:2])
    row = data['rows'][0]
    job = jobs.create(directory, paths[0], reuse=False)
    job['word_batch_id'] = data['id']
    jobs.write(jobs.folder_for(directory, job['id']), job)
    jobs.selections(directory, job['id'], [], whole_document=True)
    row['word_id'] = job['id']
    batches.persist(directory, data)
    doc = Document(paths[0])
    doc.add_paragraph('Nội dung mới.')
    doc.save(paths[0])
    original = paths[0].read_bytes()
    result = batches.run(directory, data, Event(), lambda _: None)
    assert result['rows'][0]['status'] == 'failed' and 'đã được sửa' in result['rows'][0]['message']
    assert result['rows'][1]['status'] == 'done' and paths[0].read_bytes() == original


def test_full_disk_during_export_pauses_without_losing_translation(prepared, monkeypatch):
    directory, paths, calls = prepared
    export = jobs.export
    fail = [True]

    def sometimes(*args, **kwargs):
        if fail[0] and len(calls) == 2:
            raise StorageProblem('Chưa đủ chỗ để lưu Word.')
        return export(*args, **kwargs)

    monkeypatch.setattr(jobs, 'export', sometimes)
    result = batches.run(directory, selected(directory, paths), Event(), lambda _: None)
    assert result['phase'] == 'paused'
    assert [row['status'] for row in result['rows']] == ['done', 'failed', 'waiting']
    second = jobs.read(directory, result['rows'][1]['word_id'])
    assert second['error_code'] == 'storage' and all(row.get('en') for row in second['regions'])
    fail[0] = False
    result = batches.run(directory, result, Event(), lambda _: None, retry=True)
    assert all(row['status'] == 'done' for row in result['rows'])


def test_recovery_between_import_and_whole_document_selection(prepared):
    directory, paths, _ = prepared
    data = selected(directory, paths[:2])
    job = jobs.create(directory, paths[0], reuse=False)
    job['word_batch_id'] = data['id']
    jobs.write(jobs.folder_for(directory, job['id']), job)
    data['rows'][0].update(word_id=job['id'], status='running')
    data['phase'] = 'running'
    batches.persist(directory, data)
    result = batches.run(directory, batches.load(directory, data['id']), Event(), lambda _: None)
    assert all(row['status'] == 'done' for row in result['rows'])
    assert result['rows'][0]['word_id'] == job['id']


def test_stale_window_and_dismissed_queue_cannot_restart(prepared):
    directory, paths, calls = prepared
    data = selected(directory, paths[:2])
    stale = copy.deepcopy(data)
    batches.persist(directory, data)
    with pytest.raises(ValueError, match='cửa sổ khác'):
        batches.run(directory, stale, Event(), lambda _: None)
    batches.discard(directory, data)
    with pytest.raises(ValueError, match='đã được bỏ'):
        batches.run(directory, batches.load(directory, data['id']), Event(), lambda _: None)
    assert not calls


def test_selection_is_bounded_local_and_deduplicates_paths(prepared):
    _, paths, _ = prepared
    result = batches.selection([paths[0].as_uri(), paths[0].as_uri(), paths[1].as_uri()])
    assert len(result['rows']) == 2
    for urls in ([paths[0].as_uri()], [paths[0].as_uri()] * 101,
                 [paths[0].as_uri()] * 2, ['https://example.org/a.docx', paths[1].as_uri()]):
        with pytest.raises(ValueError):
            batches.selection(urls)


@pytest.mark.parametrize('invalid', ['rows', 'elapsed', 'identity'])
def test_corrupt_queue_does_not_load(prepared, invalid):
    directory, paths, _ = prepared
    data = selected(directory, paths[:2])
    if invalid == 'rows':
        data['rows'][0]['status'] = 'unknown'
    elif invalid == 'elapsed':
        data['elapsed_seconds'] = 'yesterday'
    else:
        data['rows'][0]['word_id'] = '../source'
    batches.journal(directory, data['id']).write_text(json.dumps(data), encoding='utf-8')
    with pytest.raises(ValueError):
        batches.load(directory, data['id'])
    assert all(QUrl(row['url']).isLocalFile() for row in data['rows'])
