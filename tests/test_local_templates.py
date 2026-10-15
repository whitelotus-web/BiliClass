from threading import Event

import pytest
from pptx import Presentation

from app.google_conversion import _template
from app.local_template_style import THEMES


@pytest.mark.parametrize('preset', list(THEMES))
@pytest.mark.parametrize('mode', ['parallel_columns', 'sentence_pairs', 'english_only'])
def test_local_theme_output_keeps_paired_content_and_click_transition(tmp_path, preset, mode):
    target = tmp_path/'sample.pptx'
    _template(target, {'title': 'Cơ năng', 'english_title': 'Mechanical energy',
                      'conversion_format': mode, 'preset': preset},
              [('Ý 1', 'Năng lượng được bảo toàn.', 'Energy is conserved.')], Event())
    deck = Presentation(target)
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    assert 'Energy is conserved.' in text
    assert ('Năng lượng được bảo toàn.' in text) == (mode != 'english_only')
    if mode == 'english_only':
        from app.slide_speech import VI_MARKS
        assert not any(c in VI_MARKS for c in text.casefold())
    assert str(slide.background.fill.fore_color.rgb) == THEMES[preset]['paper']
    assert len(slide._element.xpath('./p:transition/p:fade')) == 1
    assert slide._element.xpath('./p:transition/@advClick') == ['1']
    # Every content box remains within the page, with space reserved for the footer.
    assert all(s.left >= 0 and s.top >= 0 and s.left+s.width <= deck.slide_width
               and s.top+s.height <= deck.slide_height for s in slide.shapes)
    assert 'VI: Năng lượng được bảo toàn.' in slide.notes_slide.notes_text_frame.text
    from app.slide_speech import inspect_slide
    spoken = ' '.join(s['english_text'] for s in inspect_slide(slide,1)['englishSpeechSegments'])
    assert 'BiliClass' not in spoken and 'Page 1' not in spoken
