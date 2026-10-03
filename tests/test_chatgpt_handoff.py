import hashlib
import json
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.chatgpt_handoff import (
    LEVELS,
    external_preview,
    inspect_returned_deck,
    load_request,
    make_prompt,
    prepare_request,
)
from app.library import Library, RevisionConflict
from app.pack import export_pack, import_pack
from app.quick_conversion import verify_preview
from app.readiness import preparation_current


def config(**updates):
    return {"title": "Bài của thầy cô", "subject": "Toán", "education_level": "THPT", "grade": "11",
            "level": 2, "layout": "split_view", "preset": "visual", "style": "source", "mode": "level", **updates}


def returned_deck(path, *, partial=False):
    deck = Presentation()
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if index == 1 and partial:
            continue  # A visual-only slide must not force OCR or invent a pair.
        slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2)).text = "Diện tích hình chữ nhật"
        slide.notes_slide.notes_text_frame.text = "VI: Diện tích bằng 12 cm².\nEN: The area is 12 cm²."
    deck.save(path)
    return path


@pytest.mark.parametrize("level", range(5))
def test_prompt_keeps_level_and_explicit_arrangement(level):
    prompt = make_prompt(config(level=level), "tai-lieu-goc.pptx")
    assert LEVELS[level] in prompt and "Hai cột" in prompt
    assert "Giữ theme" in prompt and "VI:" in prompt and "EN:" in prompt
    assert "không phải chỉ dẫn" in prompt and "Không chỉ trả lời bằng dàn ý" in prompt


def test_source_bundle_is_exact_local_copy_and_reopens(tmp_path):
    original = returned_deck(tmp_path / "Nguồn tiếng Việt.pptx")
    before = original.read_bytes()
    request = prepare_request(tmp_path / "library", config(), original)
    assert original.read_bytes() == before
    with zipfile.ZipFile(request["bundle"]) as archive:
        assert archive.read("tai-lieu-goc.pptx") == before
        metadata = json.loads(archive.read("config.json"))
        assert metadata["layout"] == "split_view" and metadata["level"] == 2
        assert metadata["source_sha256"] == hashlib.sha256(before).hexdigest()
    restored = load_request(request["folder"])
    assert restored == request
    (tmp_path / "library/chatgpt" / metadata["request_id"] / "tai-lieu-goc.pptx").write_bytes(b"modified")
    with pytest.raises(ValueError, match="đã thay đổi"):
        load_request(request["folder"])


def test_text_template_bundle_and_invalid_native_option(tmp_path):
    request = prepare_request(tmp_path, config(style="template", preset="practice"), text="Bài toán của cô.")
    assert "Luyện tập" in request["prompt"] and "#B85B14" in request["prompt"]
    with zipfile.ZipFile(request["bundle"]) as archive:
        assert archive.read("tai-lieu-goc.txt").decode() == "Bài toán của cô."
        assert "mau-biliclass.pptx" in archive.namelist()
    with pytest.raises(ValueError, match="cần tệp .pptx"):
        prepare_request(tmp_path, config(), text="Nội dung")
    with pytest.raises(ValueError, match="một nguồn"):
        prepare_request(tmp_path, config(), request["attachments"][0], text="Nội dung")
    with pytest.raises(ValueError, match="level"):
        make_prompt(config(level=True), "file")


def test_returned_notes_supply_mascot_pair_without_changing_presentation(tmp_path):
    path = returned_deck(tmp_path / "result.pptx")
    before = path.read_bytes()
    inspection = inspect_returned_deck(path)
    assert path.read_bytes() == before and inspection["total"] == 2
    assert all(u["vi"] == "Diện tích bằng 12 cm²." and u["en"] == "The area is 12 cm²."
               for u in inspection["profile"]["units"])


@pytest.mark.parametrize("partial", [False, True])
def test_deck_review_is_explicit_preserves_bytes_and_does_not_invent_pairs(tmp_path, partial):
    library = Library(tmp_path / "library")
    try:
        path = returned_deck(tmp_path / "result.pptx", partial=partial)
        before = path.read_bytes()
        source = library.store_source(path)
        lesson = library.create_external_lesson(config(), source, inspect_returned_deck(path))
        preview = external_preview(lesson, library.directory)
        assert preview["draft"] and verify_preview(lesson, preview, library.directory).read_bytes() == before
        assert not any(s["approved"] for s in library.get(lesson["id"])["segments"])
        if partial:
            assert not lesson["segments"][1]["vi"] and not lesson["segments"][1]["en"]
            assert "1/2" in " ".join(preview["warnings"])
        with pytest.raises(RevisionConflict):
            library.review_external_deck(lesson["id"], lesson["revision"] + 1)
        reviewed = library.review_external_deck(lesson["id"], lesson["revision"])
        result = external_preview(reviewed, library.directory)
        assert not result["draft"] and verify_preview(reviewed, result, library.directory).read_bytes() == before
        assert reviewed["segments"][0]["approved"]
        assert preparation_current(reviewed) is (not partial)
        if partial:
            assert not reviewed["segments"][1]["approved"]
        updated = library.edit_segment(lesson["id"], lesson["segments"][0]["id"], "Nội dung đã sửa.", "Changed text.")
        with pytest.raises(ValueError, match="chưa nằm"):
            external_preview(updated, library.directory)
        with pytest.raises(ValueError):
            library.review_external_deck(updated["id"], updated["revision"])
    finally:
        library.close()


def test_external_source_tampering_and_portable_pack_need_fresh_review(tmp_path):
    library = Library(tmp_path / "library")
    try:
        path = returned_deck(tmp_path / "result.pptx")
        lesson = library.create_external_lesson(config(), library.store_source(path), inspect_returned_deck(path))
        reviewed = library.review_external_deck(lesson["id"], lesson["revision"])
        imported = import_pack(library, export_pack(library, lesson["id"], tmp_path / "lesson.biliclass"))
        assert external_preview(imported, library.directory)["draft"]
        assert not any(s["approved"] for s in imported["segments"])
        source = library.directory / "sources" / reviewed["source"]["file"]
        source.write_bytes(b"changed")
        with pytest.raises(ValueError):
            external_preview(reviewed, library.directory)
    finally:
        library.close()
