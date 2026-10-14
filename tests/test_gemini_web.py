import hashlib
import json
import sys
from pathlib import Path
from threading import Event
from urllib.parse import urlsplit

import pytest
from pptx import Presentation
from pptx.util import Inches

from app import gemini_automation as web
from app.browser_accounts import BrowserAccounts
from app.chatgpt_handoff import prepare_request

HTML = '''<!doctype html><html><body>
<a aria-label="Google Account: Fixture Teacher (fixture@example.test)" href="#">Account</a>
<div id="history"></div><div id="messages"></div>
<form class="input-area-container" onsubmit="return false">
<button type="button" data-test-id="bard-mode-menu-button" id="model">Thinking</button>
<button type="button" id="effort">Standard</button><div id="menu" role="menu" hidden></div>
<rich-textarea><div class="ql-editor" contenteditable="true" role="textbox"></div></rich-textarea>
<input type="file" multiple><div id="attached"></div><button type="button" class="send-button" aria-label="Send message">Send</button>
</form><script>
const init=INITIAL; const messages=document.querySelector('#messages');
function addMessage(prompt,i) {
 const u=document.createElement('user-query'); const t=document.createElement('div');t.className='query-text';t.textContent=prompt;u.append(t);messages.append(u);
 const r=document.createElement('model-response');r.id='response-'+i;
 r.innerHTML=SLIDES ? '<a href="https://docs.google.com/presentation/d/fixture/edit" target="_blank">Open in Slides</a>' : '<a href="/generated/result.pptx">bai-giang-song-ngu.pptx</a>';messages.append(r);
}
for(let i=0;i<init.length;i++)addMessage(init[i].prompt,i);
if(HAS_HISTORY)document.querySelector('#history').innerHTML='<a href="/app/fixtureGemini">Recent chat</a>';
const menu=document.querySelector('#menu'),model=document.querySelector('#model'),effort=document.querySelector('#effort');
function choices(names,button){menu.replaceChildren();menu.hidden=false;for(const name of names){const e=document.createElement('button');e.type='button';e.setAttribute('role','menuitemradio');e.textContent=name;e.onclick=()=>{if(!BROKEN)button.textContent=name;menu.hidden=true;};menu.append(e);}}
model.onclick=()=>choices(['Fast','Thinking','Gemini 2.5 Pro','Gemini 3.1 Pro'],model);
effort.onclick=()=>choices(['Light','Standard','High'],effort);
document.querySelector('input').onchange=e=>{document.querySelector('#attached').replaceChildren();for(const f of e.target.files){const chip=document.createElement('file-preview');chip.textContent=f.name;document.querySelector('#attached').append(chip);}};
document.querySelector('.send-button').onclick=()=>{const q=document.querySelector('.ql-editor');const prompt=q.innerText;console.log('FIXTURE-SEND:'+JSON.stringify({prompt,files:[...document.querySelector('input').files].map(f=>f.name),model:model.textContent,effort:effort.textContent}));history.pushState({},'','/app/fixtureGemini');addMessage(prompt,messages.querySelectorAll('user-query').length);q.textContent='';};
</script></body></html>'''

SLIDES = '''<html><body><button role="menuitem" id="docs-file-menu">File</button><div role="menu" id="menu"></div><script>
const menu=document.querySelector('#menu');document.querySelector('#docs-file-menu').onclick=()=>{menu.innerHTML='<button role="menuitem" id="download">Download</button>';document.querySelector('#download').onmouseover=()=>{if(!document.querySelector('#pptx')){const a=document.createElement('a');a.role='menuitem';a.id='pptx';a.textContent='Microsoft PowerPoint (.pptx)';a.href='/generated/result.pptx';menu.append(a);}};};
</script></body></html>'''


@pytest.fixture
def site(tmp_path, monkeypatch):
    path = tmp_path / "source.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = "Năng lượng được bảo toàn.\nEnergy is conserved."
    slide.notes_slide.notes_text_frame.text = "VI: Năng lượng được bảo toàn.\nEN: Energy is conserved."
    deck.save(path)
    store = BrowserAccounts(tmp_path / "library")
    store.choose_service("gemini")
    item = store.add("Fixture")
    store.observe(item["id"], "pro")
    config = {"title":"Bài thử", "subject":"Vật lý", "grade":"10", "education_level":"THPT", "style":"source",
              "preset":"standard", "conversion_format":"parallel_columns", "provider":"browser_web",
              "browser_service":"gemini", "require_highest_model":True}
    request = prepare_request(store.root.parent, config, path)
    capture = {"sends":[],"slides":False,"broken":False}
    real_open = web.open_context
    def opening(playwright, account, *, background, cancel=None):
        context = real_open(playwright, account, background=background, cancel=cancel)
        def watch(page):
            page.on('console', lambda msg: capture['sends'].append(json.loads(msg.text[len('FIXTURE-SEND:'):])) if msg.text.startswith('FIXTURE-SEND:') else None)
        context.on('page', watch)
        for page in context.pages:
            watch(page)
        def route(req):
            parsed = urlsplit(req.request.url)
            if parsed.path == '/generated/result.pptx':
                req.fulfill(status=200, content_type='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                            headers={'Content-Disposition':'attachment; filename="result.pptx"'}, body=path.read_bytes())
            elif parsed.hostname == 'docs.google.com':
                req.fulfill(content_type='text/html', body=SLIDES)
            else:
                initial = capture['sends'] if parsed.path.startswith('/app/fixture') else []
                req.fulfill(content_type='text/html', body=HTML.replace('INITIAL',json.dumps(initial)).replace('SLIDES',json.dumps(capture['slides'])).replace('BROKEN',json.dumps(capture['broken'])).replace('HAS_HISTORY',json.dumps(bool(capture['sends']))))
        context.route('**/*', route)
        return context
    monkeypatch.setattr(web,'open_context',opening)
    return store, item, request, path, capture


@pytest.mark.parametrize("slides", [False, True])
def test_real_chrome_upload_confirm_highest_model_receive_pptx_and_reuse_without_resend(site, slides):
    store, item, request, source, capture = site
    capture['slides'] = slides
    result = web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=15)
    assert Path(result['path']).read_bytes() == source.read_bytes()
    assert len(capture['sends']) == 1
    sent = capture['sends'][0]
    assert sent['model'] == 'Gemini 3.1 Pro' and sent['effort'] == 'High'
    assert ' '.join(sent['prompt'].split()) == ' '.join(request['prompt'].split())
    assert sent['files'] == request['attachments']
    record = web.read_record(request['folder'],item['id'])
    assert record['browser_service'] == 'gemini' and record['state'] == 'completed'
    assert record['highest_model_verified'] and record['reasoning_verified']
    assert record['sha256'] == hashlib.sha256(source.read_bytes()).hexdigest()
    cached = web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None)
    assert cached['cached'] and len(capture['sends']) == 1


def test_unverified_model_blocks_before_upload_or_send(site):
    store,item,request,_source,capture=site
    capture['broken']=True
    with pytest.raises(web.BrowserProblem) as exc:
        web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=10)
    assert exc.value.code=='model' and not capture['sends']
    assert web.read_record(request['folder'],item['id'])['state']=='prepared'


def test_unknown_send_receipt_recovers_same_gemini_chat_without_resending(site, monkeypatch):
    store,item,request,_source,capture=site
    receipt = web._receipt
    monkeypatch.setattr(web, '_receipt', lambda *args, **kwargs: False)
    with pytest.raises(web.BrowserProblem) as exc:
        web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=10)
    assert exc.value.code=='receipt' and len(capture['sends']) == 1
    assert web.read_record(request['folder'],item['id'])['state']=='submitting'
    monkeypatch.setattr(web, '_receipt', receipt)
    result = web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=10)
    assert result['url']=='https://gemini.google.com/app/fixtureGemini'
    assert len(capture['sends'])==1


def test_gemini_quality_correction_uses_same_chat_and_rechecks_highest_controls(site, monkeypatch):
    store,item,request,_source,capture=site
    config_path=Path(request['folder'])/'config.json'
    config=json.loads(config_path.read_text(encoding='utf-8'))
    config['quality_review_version']=1
    config_path.write_text(json.dumps(config),encoding='utf-8')
    def checked(path, folder, **kwargs):
        issues=[{'code':'english_body_missing','severity':'FAIL'}] if len(capture['sends'])<2 else []
        quality={'slides':[{'slide':1,'issues':issues,'status':'FAIL' if issues else 'CHECK'}]}
        return {'path':str(path),'inspection':{'slide_quality':quality},'changes':[]}
    monkeypatch.setattr('app.slide_refinement.repair_and_inspect',checked)
    result=web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=15)
    assert result['url']=='https://gemini.google.com/app/fixtureGemini'
    assert len(capture['sends'])==2
    assert all(s['model']=='Gemini 3.1 Pro' and s['effort']=='High' for s in capture['sends'])
    record=web.read_record(request['folder'],item['id'])
    assert record['quality_followups']==1 and record['quality_stop_reason']=='complete'
    assert 'SỬA BẢN SAO' in capture['sends'][1]['prompt']


def test_slow_gemini_model_menu_is_waited_for_before_upload(site, monkeypatch):
    store,item,request,_source,capture=site
    monkeypatch.setattr(sys.modules[__name__], 'HTML', HTML.replace('model.onclick=()=>choices(', 'model.onclick=()=>setTimeout(()=>choices(').replace("'Gemini 3.1 Pro'],model);", "'Gemini 3.1 Pro'],model),400);"))
    web.convert(store.get(item['id']),store.root,request['folder'],Event(),lambda *_:None,timeout=15)
    assert len(capture['sends']) == 1 and capture['sends'][0]['model']=='Gemini 3.1 Pro'


def test_gemini_quota_notice_does_not_mark_saved_session_logged_out(site, monkeypatch):
    store,item,_request,_source,capture=site
    monkeypatch.setattr(sys.modules[__name__], 'HTML', HTML.replace('<div id="history">', '<div role="alert">You have reached your daily limit</div><div id="history">'))
    result=web.check_session(store.get(item['id']),store.root,Event(),lambda *_:None)
    assert result['ready'] and result['quota_limited'] and not capture['sends']
