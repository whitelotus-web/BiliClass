import io
from threading import Event

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.universal_content import decorative_picture, math_text, source_blocks
from app.universal_templates import asset, build


def test_modern_math_frame_keeps_all_prose_and_ignores_transparent_source_marker(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape = page.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    shape.text = 'Phương trình của vận tốc là: \nVận tốc biến thiên tuần hoàn theo thời gian.'
    shape.text_frame.paragraphs[0]._p.append(etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>v = 2</m:t></m:r></m:oMath>'))
    hidden = shape.text_frame.paragraphs[1].add_run()
    hidden.text = ' hidden-marker-123'
    props = hidden._r.get_or_add_rPr()
    fill = etree.SubElement(props, '{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
    color = etree.SubElement(fill, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val='FFFFFF')
    etree.SubElement(color, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val='0')
    parent = shape._element.getparent()
    index = parent.index(shape._element)
    alternate = etree.Element('{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent')
    choice = etree.SubElement(alternate, '{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice', Requires='a14')
    choice.append(shape._element)
    parent.insert(index, alternate)
    source = tmp_path / 'source.pptx'
    deck.save(source)
    blocks = source_blocks(Presentation(source))
    assert len(blocks) == 2
    assert '$v = 2$' in blocks[0][1]
    assert blocks[1][1] == 'Vận tốc biến thiên tuần hoàn theo thời gian.'
    assert 'hidden-marker' not in str(blocks)


def test_inline_equations_remain_editable_and_keep_surrounding_bilingual_prose(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    page.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(1)).text = 'Vận tốc'
    shape = page.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    shape.text = 'Phương trình của vận tốc là: '
    shape.text_frame.paragraphs[0]._p.append(etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>v = 2</m:t></m:r></m:oMath>'))
    path = tmp_path / 'source.pptx'
    deck.save(path)
    pairs = [(key, vi, vi.replace('Vận tốc', 'Velocity').replace('Phương trình của vận tốc là:', 'The velocity equation is:'))
             for key, vi in source_blocks(deck)]
    output = tmp_path / 'out.pptx'
    build(output, asset('standard'), {'title': 'Vận tốc', 'english_title': 'Velocity',
          'preset': 'standard', 'conversion_format': 'parallel_columns'}, pairs, Event(), path)
    result = Presentation(output)
    text = '\n'.join(s.text for slide in result.slides for s in slide.shapes if s.has_text_frame)
    assert 'Phương trình của vận tốc là:' in text and 'The velocity equation is:' in text
    assert '$' not in text
    assert sum(len(s._element.xpath('.//*[local-name()="oMath"]')) for s in result.slides) == 2


def test_equation_fraction_does_not_flatten_to_concatenated_digits():
    node = etree.fromstring(b'<m:f xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:num><m:r><m:t>T</m:t></m:r></m:num><m:den><m:r><m:t>4</m:t></m:r></m:den></m:f>')
    assert math_text(node) == '(T)/(4)'


def test_ordinary_superscript_runs_keep_scientific_power():
    from app.universal_content import paragraph_text
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape = page.shapes.add_textbox(0, 0, Inches(8), Inches(1))
    p = shape.text_frame.paragraphs[0]
    p.add_run().text = 'A. 63,2 m/s'
    power = p.add_run()
    power.text = '2'
    power._r.get_or_add_rPr().set('baseline', '30000')
    assert paragraph_text(p._p) == 'A. 63,2 m/s²'
    power._r.get_or_add_rPr().set('baseline', '0')
    assert paragraph_text(p._p) == 'A. 63,2 m/s2'


def test_formula_options_and_prose_remain_with_the_question():
    from app.universal_templates import formula

    assert not formula('C. $3√(3)$ cm')
    assert not formula('D. −3√3 cm')
    assert not formula('The period is $T=2π/ω$.')
    assert formula('$T=2π/ω$')
    assert formula('v = ωA')


def test_empty_bordered_panel_is_removed_but_graph_and_solid_picture_are_kept():
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    for kind in ('panel', 'graph', 'solid'):
        image = Image.new('RGB', (240, 160), 'blue' if kind == 'solid' else 'white')
        draw = ImageDraw.Draw(image)
        if kind != 'solid':
            draw.rounded_rectangle((1, 1, 238, 158), radius=12, outline='navy', width=3)
        if kind == 'graph':
            draw.line((20, 120, 60, 35, 120, 110, 200, 45), fill='blue', width=3)
        stream = io.BytesIO()
        image.save(stream, 'PNG')
        shape = page.shapes.add_picture(io.BytesIO(stream.getvalue()), Inches(1), Inches(2))
        obj = {'shape': shape, 'x': shape.left, 'y': shape.top, 'w': shape.width, 'h': shape.height}
        assert decorative_picture(obj, (deck.slide_width, deck.slide_height)) == (kind == 'panel')


def test_only_exact_text_overlays_are_deduplicated():
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    for top in (2, 2, 4):
        page.shapes.add_textbox(Inches(1), Inches(top), Inches(5), Inches(1)).text = 'Vận tốc'
    assert [text for _, text in source_blocks(deck)] == ['Vận tốc', 'Vận tốc']


def test_source_rows_tolerate_import_rounding_and_keep_top_opaque_answer():
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    for left, top, text in [(1, 2, 'A. 2'), (7, 2-1/12700, 'B. 3'),
                            (1, 4, 'C. 4'), (7, 4, 'D. 6,31 m/s2'), (7, 4, 'D. 63,2 m/s2')]:
        shape = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(4), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.text = text
    assert [value for _, value in source_blocks(deck)] == ['A. 2', 'B. 3', 'C. 4', 'D. 63,2 m/s2']
    # A transparent overlay does not establish that the original is invisible.
    page.shapes[-1].fill.background()
    assert [value for _, value in source_blocks(deck)][-2:] == ['D. 6,31 m/s2', 'D. 63,2 m/s2']


def test_standalone_native_equation_keeps_units_outside_math_xml(tmp_path):
    from app.universal_templates import _media
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape = page.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    p = shape.text_frame.paragraphs[0]
    p._p.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>A=5</m:t></m:r></m:oMath>'))
    p.add_run().text = ' (cm).'
    native = _media(page)
    assert len(native) == 1 and '(cm).' in native[0]['shape'].text


def test_graph_and_standalone_equation_receive_readable_separate_pages(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    page.shapes.add_textbox(Inches(1), Inches(.3), Inches(9), Inches(.8)).text = 'Đồ thị'
    stream = io.BytesIO()
    Image.new('RGB', (400, 300), 'blue').save(stream, 'PNG')
    page.shapes.add_picture(io.BytesIO(stream.getvalue()), Inches(1), Inches(2), width=Inches(5))
    shape = page.shapes.add_textbox(Inches(8), Inches(2), Inches(4), Inches(1))
    p = shape.text_frame.paragraphs[0]
    p._p.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x=5</m:t></m:r></m:oMath>'))
    p.add_run().text = ' (cm)'
    source = tmp_path/'source.pptx'
    deck.save(source)
    target = tmp_path/'out.pptx'
    result = build(target, asset('practice'), {'title':'Đồ thị','english_title':'Graph',
                   'preset':'practice','conversion_format':'parallel_columns'},
                   [('Slide 1','Đồ thị','Graph')], Event(), source)
    output = Presentation(target)
    assert [p['layout'] for p in result['layout_plan']] == ['IMAGE_TEXT','FORMULA']
    assert len(output.slides[0]._element.xpath('.//*[local-name()="oMath"]')) == 0
    assert len(output.slides[1]._element.xpath('.//*[local-name()="oMath"]')) == 1
    math_shape = next(s for s in output.slides[1].shapes if s.name.startswith('BCL_SOURCE'))
    assert '(cm)' in math_shape.text
    assert min(int(n)/100 for n in math_shape._element.xpath('.//@sz')) >= 18


def test_table_translations_follow_reading_order_with_decorative_images(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    stream = io.BytesIO()
    Image.new('RGB', (30, 30), 'blue').save(stream, 'PNG')
    page.shapes.add_picture(io.BytesIO(stream.getvalue()), Inches(1), Inches(.1))
    page.shapes.add_textbox(Inches(1), Inches(.5), Inches(5), Inches(1)).text = 'Bảng'
    table = page.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = 'Mốc'
    table.cell(0, 1).text = 'Giá trị'
    source = tmp_path/'source.pptx'
    deck.save(source)
    pairs = [(loc, vi, {'Bảng': 'Table', 'Mốc': 'Point', 'Giá trị': 'Value'}[vi]) for loc, vi in source_blocks(deck)]
    target = tmp_path/'out.pptx'
    build(target, asset('standard'), {'title': 'Bài học', 'english_title': 'Lesson',
          'preset': 'standard', 'conversion_format': 'english_only'}, pairs, Event(), source)
    cells = [c.text for s in Presentation(target).slides for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
    assert cells == ['Point', 'Value']


def test_small_source_icons_are_preserved_without_taking_the_main_figure_column(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    stream = io.BytesIO()
    Image.new('RGB', (40, 40), 'blue').save(stream, 'PNG')
    page.shapes.add_picture(io.BytesIO(stream.getvalue()), Inches(1), Inches(2), width=Inches(.4))
    source = tmp_path/'source.pptx'
    deck.save(source)
    target = tmp_path/'out.pptx'
    result = build(target, asset('standard'), {'title': 'Bài học', 'english_title': 'Lesson',
                   'preset': 'standard', 'conversion_format': 'parallel_columns'},
                   [('Slide 1', 'Vận tốc', 'Velocity'), ('Slide 1 · ý 2', 'Vận tốc biến thiên.', 'Velocity varies.')], Event(), source)
    assert result['layout_plan'][0]['layout'] == 'THEORY'
    pictures = [sh for s in Presentation(target).slides for sh in s.shapes if int(sh.shape_type) == 13]
    assert len(pictures) == 1 and pictures[0].width <= Inches(.6)


def test_long_quiz_with_only_an_icon_does_not_add_an_empty_quiz_page(tmp_path):
    deck = Presentation()
    page = deck.slides.add_slide(deck.slide_layouts[6])
    stream = io.BytesIO()
    Image.new('RGB', (40, 40), 'blue').save(stream, 'PNG')
    page.shapes.add_picture(io.BytesIO(stream.getvalue()), Inches(1), Inches(2), width=Inches(.4))
    source = tmp_path/'source.pptx'
    deck.save(source)
    question = 'Tìm đáp án cho câu hỏi về vận tốc của dao động. ' * 8
    pairs = [('Slide 1', 'Câu 1', 'Question 1'), ('Slide 1 · câu hỏi', question, 'Find the answer about velocity. ' * 10)]
    pairs += [('Slide 1 · ' + c, c + '. ' + str(i), c + '. ' + str(i)) for i, c in enumerate('ABCD', 1)]
    target = tmp_path/'out.pptx'
    result = build(target, asset('standard'), {'title': 'Bài học', 'english_title': 'Lesson',
                   'preset': 'standard', 'conversion_format': 'parallel_columns'}, pairs, Event(), source)
    assert all(row['layout'] != 'QUIZ' for row in result['layout_plan'])
    pictures = [sh for s in Presentation(target).slides for sh in s.shapes if int(sh.shape_type) == 13]
    assert len(pictures) == 1
