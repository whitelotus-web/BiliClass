import zipfile
from contextlib import closing
from threading import Event

import pytest
from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image

from app import word_jobs as jobs
from app.library import Library
from app.word_document import NS, M, Package, digest, keyword_annotations, sentence_ranges, utf16_position


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading('Bài học Vật lý', 0)
    paragraph = document.add_paragraph()
    paragraph.add_run('Dao động điều hòa').bold = True
    paragraph.add_run(' có biên độ 2 cm. Gia tốc có đơn vị m/s².')
    document.add_paragraph('Giữ nguyên đoạn này.')
    document.add_paragraph('Giữ nguyên đoạn này.')
    equation = document.add_paragraph('Phương trình ')
    math = etree.SubElement(equation._p, '{' + M + '}oMath')
    run = etree.SubElement(math, '{' + M + '}r')
    etree.SubElement(run, '{' + M + '}t').text = 'x = 2t'
    equation.add_run(' mô tả chuyển động.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Vận tốc 12 m/s.'
    table.cell(0, 1).text = 'Không dịch ô này.'
    picture = tmp_path / 'picture.png'
    Image.new('RGB', (120, 80), '#237ab3').save(picture)
    document.add_picture(str(picture), width=Inches(1))
    document.sections[0].header.paragraphs[0].text = 'Trường thử nghiệm'
    document.sections[0].footer.paragraphs[0].text = 'Giữ nguyên chân trang'
    path = tmp_path / 'Bài học.docx'
    document.save(path)
    return path


def selected(package, text, start=0, end=None):
    row = next(row for row in package.rows if row['text'].startswith(text))
    region = package.selection([dict(paragraph=row['id'], start=start, end=len(row['text']) if end is None else end)])[0]
    region['en'] = region['wire'].replace('Dao động điều hòa', 'Simple harmonic motion').replace('có biên độ', 'has an amplitude of').replace('Vận tốc', 'Velocity').replace('Phương trình', 'Equation').replace('mô tả chuyển động', 'describes motion')
    return region


@pytest.mark.parametrize('mode', ['below', 'columns', 'english', 'keywords'])
def test_only_selected_regions_change_and_non_document_parts_stay_identical(source, tmp_path, mode):
    original = source.read_bytes()
    package = Package(source)
    first = selected(package, 'Dao động', end=len('Dao động điều hòa có biên độ 2 cm.'))
    second = selected(package, 'Vận tốc')
    if mode == 'keywords':
        first['annotations'] = keyword_annotations(first, {'dao động điều hòa': 'simple harmonic motion'})
        second['annotations'] = keyword_annotations(second, {'vận tốc': 'velocity'})
    output = tmp_path / (mode + '.docx')
    package.save(output, [first, second], dict(mode=mode))
    assert source.read_bytes() == original
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name != package.main)
    result = Package(output)
    text = '\n'.join(row['text'] for row in result.rows)
    assert text.count('Giữ nguyên đoạn này.') == 2
    assert 'Gia tốc có đơn vị m/s².' in text and 'Không dịch ô này.' in text
    assert 'simple harmonic motion' in text.lower() and 'velocity' in text.lower()
    assert len(result.tree.findall('.//m:oMath', NS)) == 1
    if mode == 'english':
        assert 'Dao động điều hòa' not in text and 'Vận tốc 12' not in text
    elif mode != 'keywords':
        assert 'Dao động điều hòa' in text and 'Vận tốc 12' in text
    else:
        assert 'Vận tốc (velocity) 12' in text
    if mode == 'columns':
        assert len(result.tree.findall('.//w:tbl', NS)) == 3


def test_partial_bold_run_and_second_duplicate_are_mapped_by_position(source, tmp_path):
    package = Package(source)
    duplicate = [row for row in package.rows if row['text'] == 'Giữ nguyên đoạn này.'][1]
    region = package.selection([dict(paragraph=duplicate['id'], start=0, end=len(duplicate['text']))])[0]
    region['en'] = 'Keep this paragraph.'
    output = tmp_path / 'duplicate.docx'
    package.save(output, [region], {'mode': 'english'})
    result = Document(output)
    text = [paragraph.text for paragraph in result.paragraphs]
    assert text.index('Giữ nguyên đoạn này.') < text.index('Keep this paragraph.')
    assert result.paragraphs[1].runs[0].bold


@pytest.mark.parametrize('mode,count', [('english', 1), ('below', 2), ('columns', 2)])
def test_equations_remain_native_editable_and_never_leak_tokens(source, tmp_path, mode, count):
    package = Package(source)
    region = selected(package, 'Phương trình')
    output = tmp_path / (mode + '-math.docx')
    package.save(output, [region], {'mode': mode})
    result = Package(output)
    assert len(result.tree.findall('.//m:oMath', NS)) == count
    assert not any('⟦BLC' in row['text'] for row in result.rows)
    assert 'x = 2t' in '\n'.join(row['text'] for row in result.rows)


def test_overlap_merge_unicode_positions_and_decimal_sentences(source):
    package = Package(source)
    row = package.rows[1]
    result = package.selection([dict(paragraph=row['id'], start=0, end=12), dict(paragraph=row['id'], start=5, end=20)])
    assert len(result) == 1 and result[0]['end'] == 20
    assert utf16_position('😀 Tiếng Việt', 3) == 2
    with pytest.raises(ValueError):
        utf16_position('😀 Tiếng Việt', 1)
    sentences = sentence_ranges('Giá trị 1.5 m. Xem Hình 2.3 trong bài. TS. An hướng dẫn.')
    assert len(sentences) == 3
    assert sentences[0]['text'] == 'Giá trị 1.5 m.'


def test_keyword_density_combining_accents_overlap_and_no_duplicate_annotation():
    row = dict(vi='Gia tốc của dao động điều hòa. Gia tốc luôn hướng về vị trí cân bằng.')
    terms = {'gia tốc': 'acceleration', 'dao động điều hòa': 'simple harmonic motion', 'dao động': 'oscillation', 'vị trí cân bằng': 'equilibrium position'}
    few = keyword_annotations(row, terms, 'few')
    assert len(few) == 2 and len({item['en'] for item in few}) == 2
    assert not keyword_annotations(dict(vi='Gia tốc (acceleration)'), terms)
    assert keyword_annotations(dict(vi='Gia tốc'), terms)[0]['en'] == 'acceleration'


def test_bad_numeric_edit_and_stale_selection_cannot_be_exported(source, tmp_path):
    package = Package(source)
    region = selected(package, 'Vận tốc')
    region['en'] = 'Velocity 999 m/s.'
    with pytest.raises(ValueError, match='số liệu'):
        package.save(tmp_path / 'bad.docx', [region], {'mode': 'english'})
    region['en'] = 'Velocity 12 m/s.'
    region['start'] = 1
    with pytest.raises(ValueError, match='thay đổi'):
        package.save(tmp_path / 'stale.docx', [region], {'mode': 'english'})


def test_job_library_export_is_idempotent_and_reopen_does_not_append_again(source, tmp_path):
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    package = Package(jobs.folder_for(directory, job['id']) / 'source.docx')
    region = selected(package, 'Vận tốc')
    job = jobs.selections(directory, job['id'], [region])
    job = jobs.edit(directory, job['id'], region['id'], 'Velocity 12 m/s.')
    before = source.read_bytes()
    first = jobs.export(directory, job['id'], Event())
    second = jobs.export(directory, job['id'], Event())
    assert first['export_path'] == second['export_path'] and first['output_sha256'] == second['output_sha256']
    assert source.read_bytes() == before
    assert jobs.create(directory, second['export_path'])['id'] == job['id']
    assert len(jobs.catalog(directory)) == 1
    assert digest(first['export_path']) == first['output_sha256']
    jobs.trash(directory, job['id'], True)
    assert not jobs.catalog(directory) and len(jobs.catalog(directory, True)) == 1
    jobs.trash(directory, job['id'], False)
    assert len(jobs.catalog(directory)) == 1


def test_teacher_edited_external_export_is_preserved(source, tmp_path):
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    region = selected(Package(source), 'Vận tốc')
    jobs.selections(directory, job['id'], [region])
    jobs.edit(directory, job['id'], region['id'], 'Velocity 12 m/s.')
    first = jobs.export(directory, job['id'], Event())
    external = Document(first['export_path'])
    external.add_paragraph('Teacher edit')
    external.save(first['export_path'])
    edited = open(first['export_path'], 'rb').read()
    second = jobs.export(directory, job['id'], Event())
    assert first['export_path'] != second['export_path']
    assert open(first['export_path'], 'rb').read() == edited
    third = jobs.export(directory, job['id'], Event())
    assert third['export_path'] == second['export_path'], 'Reuse the managed copy; do not create endless numbered files'


def test_undo_swaps_output_and_matching_preview_metadata(source, tmp_path):
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    region = selected(Package(source), 'Vận tốc')
    jobs.selections(directory, job['id'], [region])
    jobs.edit(directory, job['id'], region['id'], 'Velocity 12 m/s.')
    first = jobs.export(directory, job['id'], Event())
    jobs.edit(directory, job['id'], region['id'], 'A speed of 12 m/s.')
    second = jobs.export(directory, job['id'], Event())
    assert first['output_sha256'] != second['output_sha256']
    restored = jobs.restore_previous(directory, job['id'])
    assert restored['output_sha256'] == first['output_sha256']
    assert restored['regions'][0]['en'] == 'Velocity 12 m/s.'
    assert jobs.restore_previous(directory, job['id'])['output_sha256'] == second['output_sha256']


def test_word_backup_restore_and_source_change_guard(source, tmp_path):
    from app.storage import backup_library, restore_backup

    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    region = selected(Package(source), 'Vận tốc')
    jobs.selections(directory, job['id'], [region])
    jobs.edit(directory, job['id'], region['id'], 'Velocity 12 m/s.')
    exported = jobs.export(directory, job['id'], Event())
    archive = backup_library(directory)
    restored = tmp_path / 'restored'
    restore_backup(archive, restored)
    assert jobs.read(restored, job['id'])['output_sha256'] == exported['output_sha256']
    with zipfile.ZipFile(archive) as package:
        assert any(name.startswith('word_documents/') for name in package.namelist())
        assert not any(name.endswith('word.lock') for name in package.namelist())
    original = Document(source)
    original.add_paragraph('Changed source')
    original.save(source)
    with pytest.raises(ValueError, match='đã được sửa'):
        jobs.export(directory, job['id'], Event())


def test_keyword_jobs_are_offline_and_manual_changes_survive_preview(source, tmp_path):
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    region = selected(Package(source), 'Vận tốc')
    jobs.selections(directory, job['id'], [region])
    options = dict(mode='keywords', density='normal')
    job = jobs.translate(directory, job['id'], options, 'Vật lý', '11', Event(), lambda _: None)
    assert job['phase'] == 'review' and job['regions'][0]['annotations']
    job = jobs.edit_keyword(directory, job['id'], region['id'], 0, 'speed')
    job = jobs.translate(directory, job['id'], options, 'Vật lý', '11', Event(), lambda _: None)
    assert job['regions'][0]['annotations'][0]['en'] == 'speed'
    assert not (directory / 'translation_api' / 'keys.dpapi').exists()


def test_first_word_translation_without_keys_does_not_pin_an_empty_route(source, tmp_path):
    from app.google_translation import KeyStore, TranslationProblem

    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    job = jobs.create(directory, source)
    jobs.selections(directory, job['id'], [selected(Package(source), 'Vận tốc')])
    with pytest.raises(TranslationProblem, match='ít nhất một dịch vụ'):
        jobs.translate(directory, job['id'], dict(mode='below'), 'Vật lý', '11', Event(), lambda _: None)
    assert 'translation_config' not in jobs.read(directory, job['id'])
    KeyStore(directory).add('AIza' + 'z' * 35)
    # The next attempt sees the newly configured route; there is no permanently empty checkpoint.
    from app.translation_provider import snapshot

    assert snapshot(directory, include_fallback=True)['translation_fallback'][0]['translation_provider'] == 'google'


def test_upload_preferences_apply_to_new_word_and_do_not_overwrite_existing_work(source, tmp_path):
    directory = tmp_path / 'data'
    with closing(Library(directory)):
        pass
    preferences = dict(subject='Vật lý', grade='6', options=dict(mode='columns', color='666666', italic=False, density='deep'))
    first = jobs.create(directory, source, preferences=preferences)
    assert first['subject'] == 'Vật lý' and first['grade'] == '6'
    assert first['options'] == preferences['options']
    region = selected(Package(source), 'Vận tốc')
    jobs.selections(directory, first['id'], [region])
    jobs.edit(directory, first['id'], region['id'], 'Velocity 12 m/s.')
    resumed = jobs.create(directory, source, preferences=dict(subject='Toán', grade='11', options=dict(mode='english')))
    assert resumed['id'] == first['id'] and resumed['options'] == preferences['options']
    assert resumed['subject'] == 'Vật lý' and resumed['regions'][0]['en'] == 'Velocity 12 m/s.'
