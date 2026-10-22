from io import BytesIO

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.framed_text import _painted_image, constrain_generated, panels
from app.native_bilingual import assemble, extract
from app.slide_quality import inspect_quality
from app.slide_speech import deck_speech
from app.source_deck import NS


def panel_deck(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for n in range(3):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.6), Inches(1+n*1.8), Inches(12), Inches(1.4))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string('FFE580')
        text = slide.shapes.add_textbox(Inches(1.8), Inches(1.25+n*1.8), Inches(10), Inches(.6))
        text.text = 'Mục tiêu học tập: Nhận biết dao động qua đồ thị li độ theo thời gian và giải thích nguyên nhân làm tiêu hao cơ năng của hệ.'
        text.text_frame.paragraphs[0].font.size = Pt(26)
    deck.save(path)


def test_transparent_shadow_is_not_part_of_usable_panel():
    image = Image.new('RGBA', (500, 120), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 499, 119), fill=(0, 0, 0, 60))
    draw.rounded_rectangle((15, 12, 484, 94), radius=12, fill='#ffe580')
    stream = BytesIO()
    image.save(stream, format='PNG')
    left, top, right, bottom = _painted_image(stream.getvalue())
    assert .02 < left < .05 and .08 < top < .13
    assert .77 < bottom < .82 and .95 < right < .99
    draw.ellipse((0, 0, 499, 119), fill=(0, 0, 0, 0))
    stream = BytesIO()
    image.save(stream, format='PNG')
    assert _painted_image(stream.getvalue()) is None


def test_new_bilingual_columns_fit_inside_their_actual_panels(tmp_path):
    source, target = tmp_path/'source.pptx', tmp_path/'bilingual.pptx'
    panel_deck(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    translations = {u['id']: {'en': 'Learning objective: Identify oscillations through displacement-time graphs and explain the reasons for the mechanical energy loss of the system.'} for u in manifest['units']}
    assemble(source, target, manifest, translations, 'parallel_columns')
    deck = Presentation(target)
    report = inspect_quality(deck, deck_speech(target, deck), 'parallel_columns', layout_plan=manifest['layout_plan'])
    assert not any(i['code'] == 'panel_text_escape' for row in report['slides'] for i in row['issues'])
    assert len(manifest['layout_plan'][0]['pairs']) == 3
    bands = panels(deck.slides[0], deck.slide_width, deck.slide_height)
    assert len(bands) == 3
    before = etree.tostring(deck.slides[0]._element)
    # A deliberately escaped old generated frame is caught even if PowerPoint
    # thinks the glyphs fit into its oversized text box.
    shape = next(s for s in deck.slides[0].shapes if s.name.startswith('BCL_BODY '))
    shape.height = Inches(2)
    report = inspect_quality(deck, deck_speech(target, deck), 'parallel_columns', layout_plan=manifest['layout_plan'])
    assert any(i['code'] == 'panel_text_escape' for row in report['slides'] for i in row['issues'])
    tree = etree.fromstring(etree.tostring(deck.slides[0]._element))
    original_words = tree.xpath('.//a:t/text()', namespaces=NS)
    changes = constrain_generated(tree, deck.slides[0], deck.slide_width, deck.slide_height)
    assert changes and tree.xpath('.//a:t/text()', namespaces=NS) == original_words
    assert etree.tostring(tree) == before, 'Repair returns the frame to its intended panel'


def test_panel_repair_does_not_move_motion_targets_or_original_frames(tmp_path):
    source, target = tmp_path/'source.pptx', tmp_path/'generated.pptx'
    panel_deck(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': 'Learning objective: Identify energy loss of the system.'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    deck = Presentation(target)
    slide = deck.slides[0]
    moving = next(s for s in slide.shapes if s.name.startswith('BCL_BODY '))
    moving.height = Inches(2)
    tree = etree.fromstring(etree.tostring(slide._element))
    timing = etree.SubElement(tree, '{' + NS['p'] + '}timing')
    motion = etree.SubElement(timing, '{' + NS['p'] + '}animMotion')
    etree.SubElement(motion, '{' + NS['p'] + '}spTgt', spid=str(moving.shape_id))
    before = etree.tostring(tree)
    assert constrain_generated(tree, slide, deck.slide_width, deck.slide_height) == []
    assert etree.tostring(tree) == before
    original = Presentation(source)
    tree = etree.fromstring(etree.tostring(original.slides[0]._element))
    assert constrain_generated(tree, original.slides[0], original.slide_width, original.slide_height) == []


def test_body_in_shared_large_panel_never_moves_above_its_original_row(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(.7), Inches(8), Inches(6.5))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string('E6F2D8')
    heading = slide.shapes.add_textbox(Inches(6), Inches(1), Inches(6), Inches(1))
    heading.text = 'QUAN HỆ SONG SONG TRONG KHÔNG GIAN'
    body = slide.shapes.add_textbox(Inches(5.2), Inches(2.6), Inches(7.5), Inches(3.8))
    body.text = 'Đường thẳng và mặt phẳng trong không gian\nHai đường thẳng song song\nPhép chiếu song song'
    body.text_frame.paragraphs[0].font.size = Pt(26)
    source, target = tmp_path/'shared-panel.pptx', tmp_path/'paired.pptx'
    deck.save(source)
    original = source.read_bytes()
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': 'Lines and planes in space' if u['shape_id'] == heading.shape_id else 'Parallel lines and parallel projection'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    converted = Presentation(target)
    result = converted.slides[0]
    tree = etree.fromstring(etree.tostring(result._element))
    constrain_generated(tree, result, converted.slide_width, converted.slide_height)
    # Test both VI and EN columns, including the post-assembly repair pass.
    rows = tree.xpath('.//p:sp[starts-with(p:nvSpPr/p:cNvPr/@name,"BCL_BODY ")]/p:spPr/a:xfrm/a:off/@y', namespaces=NS)
    assert rows and all(int(y) >= body.top for y in rows)
    assert source.read_bytes() == original


def test_inherited_group_fill_constrains_the_solution_inside_its_panel(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    panel = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.3), Inches(4.5), Inches(9), Inches(2))
    properties = panel._element.find('p:spPr', NS)
    properties.append(etree.Element('{' + NS['a'] + '}grpFill'))
    inherited = etree.SubElement(group._element.find('p:grpSpPr', NS), '{' + NS['a'] + '}solidFill')
    etree.SubElement(inherited, '{' + NS['a'] + '}srgbClr', val='E6F2FF')
    body = group.shapes.add_textbox(Inches(1), Inches(4.7), Inches(7), Inches(1.2))
    body.text = 'Ta có ba vị trí tương đối giữa đường thẳng và mặt phẳng.\nĐường thẳng nằm trong mặt phẳng.\nĐường thẳng cắt mặt phẳng tại một điểm.'
    source, target = tmp_path/'group-fill.pptx', tmp_path/'paired.pptx'
    deck.save(source)
    assert any(row['shape_id'] == panel.shape_id for row in panels(slide, deck.slide_width, deck.slide_height))
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': 'A line lies in a plane or intersects the plane at a single point.'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    result = Presentation(target)
    from app.slide_speech import objects

    panel_box = next(row for row in panels(result.slides[0], result.slide_width, result.slide_height) if row['shape_id'] == panel.shape_id)
    columns = [row for row in objects(result.slides[0]) if row['name'].startswith('BCL_BODY ')]
    assert columns and all(row['y'] + row['h'] < panel_box['y'] + panel_box['h'] for row in columns)
    # An unpainted/transparent group cannot invent a usable panel.
    group._element.find('p:grpSpPr', NS).remove(inherited)
    assert not panels(slide, deck.slide_width, deck.slide_height)


def test_oversized_source_answer_does_not_grow_over_the_next_activity(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.3), Inches(1), Inches(9), Inches(5))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string('EFF6FF')
    answer = slide.shapes.add_textbox(Inches(.8), Inches(1.5), Inches(8), Inches(3))
    answer.text = 'Giải thích tại sao đường thẳng nằm trong mặt phẳng.\nHai điểm của đường thẳng thuộc mặt phẳng đó.'
    next_heading = slide.shapes.add_textbox(Inches(.8), Inches(3.7), Inches(8), Inches(.4))
    next_heading.text = 'Hoạt động 5'
    source, target = tmp_path/'oversized.pptx', tmp_path/'paired.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': u['text'].replace('Hoạt động', 'Activity') if u['shape_id']==next_heading.shape_id else 'The two points of the line lie in the plane, so the entire line lies in that plane.'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    result = Presentation(target)
    tree = etree.fromstring(etree.tostring(result.slides[0]._element))
    constrain_generated(tree, result.slides[0], result.slide_width, result.slide_height)
    for shape in tree.xpath('.//p:sp[starts-with(p:nvSpPr/p:cNvPr/@name,"BCL_BODY ")]', namespaces=NS):
        transform = shape.find('p:spPr/a:xfrm', NS)
        top = int(transform.find('a:off', NS).get('y'))
        bottom = top + int(transform.find('a:ext', NS).get('cy'))
        assert bottom < next_heading.top


def test_small_question_badge_keeps_its_frame_instead_of_expanding_into_question(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(26.666), Inches(15)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.2), Inches(2), Inches(26), Inches(1))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string('FFF8C4')
    badge = slide.shapes.add_textbox(Inches(.5), Inches(2.1), Inches(2), Inches(.7))
    badge.text = 'Câu 10'
    badge.text_frame.paragraphs[0].font.size = Pt(36)
    question = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(24), Inches(.7))
    question.text = 'Giao điểm của đường thẳng và mặt phẳng không thể là'
    source, target = tmp_path/'question.pptx', tmp_path/'paired.pptx'
    deck.save(source)
    manifest = extract(source, 'parallel_columns', same_slide=True)
    values = {u['id']: {'en': u['text'].replace('Câu', 'Question') if u['shape_id'] == badge.shape_id else 'The intersection of a line and a plane cannot be'} for u in manifest['units']}
    assemble(source, target, manifest, values, 'parallel_columns')
    result = Presentation(target)
    converted = next(s for s in result.slides[0].shapes if s.shape_id == badge.shape_id)
    assert converted.name.startswith('BCL_LABEL ')
    assert (converted.left, converted.top, converted.width, converted.height) == (badge.left, badge.top, badge.width, badge.height)
    stem = next(s for s in result.slides[0].shapes if s.shape_id == question.shape_id)
    assert stem.left >= badge.left + badge.width
    tree = etree.fromstring(etree.tostring(result.slides[0]._element))
    constrain_generated(tree, result.slides[0], result.slide_width, result.slide_height)
    node = tree.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@id=$id]]', namespaces=NS, id=str(badge.shape_id))[0]
    assert int(node.find('p:spPr/a:xfrm/a:off', NS).get('x')) == badge.left
