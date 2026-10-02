import hashlib
import time

from ..paths import output_root


def run() -> dict:
    import psutil
    import pythoncom
    import win32com.client
    from pptx import Presentation

    path = output_root() / "companion-probe.pptx"
    deck = Presentation()
    for number, title in enumerate(
        ("Nguồn bài giảng", "Lớp nội dung Anh–Việt", "Giáo viên duyệt nội dung"), 1
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"{number:02d} · {title}"
        slide.placeholders[1].text = "Tệp thử kỹ thuật đa môn · BiliClass M0\nKhông phải nội dung giáo án."
    deck.save(path)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    previous = [
        p.pid for p in psutil.process_iter(["name"]) if (p.info["name"] or "").lower() == "powerpnt.exe"
    ]
    pythoncom.CoInitialize()
    app = presentation = window = None
    started = time.perf_counter()
    trace = []
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
        presentation.SlideShowSettings.ShowType = 2  # ppShowTypeWindow; controlled experiment window
        window = presentation.SlideShowSettings.Run()
        for index in (1, 2, 3, 1):
            window.View.GotoSlide(index)
            time.sleep(0.1)
            pythoncom.PumpWaitingMessages()
            slide = window.View.Slide
            trace.append(
                {
                    "requested": index,
                    "actual": slide.SlideIndex,
                    "stable_id": slide.SlideID,
                    "text": slide.Shapes.Title.TextFrame.TextRange.Text,
                }
            )
        window.View.Previous()
        version = app.Version
        matched = all(item["requested"] == item["actual"] for item in trace)
        unchanged = hashlib.sha256(path.read_bytes()).hexdigest() == source_hash
        return {
            "status": "passed" if matched and unchanged else "failed",
            "office_version": version,
            "trace": trace,
            "source_unchanged": unchanged,
            "source_sha256": source_hash,
            "elapsed_ms": round((time.perf_counter() - started) * 1000),
            "pending": [
                "Full-screen + overlay trên projector",
                "Animation/build/video và custom show",
                "Thay DPI, rút màn hình, Presenter View và mất kết nối COM",
            ],
            "scope": "Thử COM thực với deck do probe tạo, chế độ cửa sổ; không mở bài của người dùng.",
        }
    finally:
        if window is not None:
            try:
                window.View.Exit()
            except Exception:
                pass
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None and not previous:
            try:
                app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
